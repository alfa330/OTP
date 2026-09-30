# -*- coding: utf-8 -*-
"""Выгрузка журнала выдач воды в xlsx.

Два листа: «Выдачи» (колонки ТЗ, раздел 5) и «Контекст» — на какую дату
собран файл, за какой период и с каким отбором. Без контекста через неделю по
файлу «Выдачи воды.xlsx» уже не понять, весь это журнал или один офис.

Соглашения те же, что у выгрузки «Посылок» (parcels/report.py):

* телефон и ID водителя — ТЕКСТОМ: числом номер теряет ведущие нули, а
  32-значный ID Excel округлит; зелёный уголок «Число сохранено как текст»
  гасится хелпером `<ignoredErrors>` из монолита (приходит аргументом);
* дата и время — настоящей датой с форматом, чтобы сортировалась;
* книга пишется потоково (`write_only=True`), стили — модульные константы.

Колонок ровно столько, сколько видно в журнале раздела: ИИН и полный снимок
CRM в базе лежат, но на экран не выводятся — значит, и в файл не идут.
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import rules

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
DATETIME_FORMAT = 'DD.MM.YYYY HH:MM'

SHEET_ISSUES = 'Выдачи'
SHEET_CONTEXT = 'Контекст'

# Сколько строк максимум в одном файле. Журнал растёт на десятки выдач в день,
# год — это единицы тысяч; потолок — страховка от случайной выгрузки «за всё».
EXPORT_LIMIT = 20000
# Период файла обязателен и не длиннее года: и потому, что файл — отчёт о
# выданном за период, и потому, что ручку зовут не только из интерфейса.
EXPORT_MAX_DAYS = 366

# Подписи видов выдачи — дословно как в waterMeta.js (KIND_LABELS).
KIND_LABELS = {
    'welcome': 'Приветственная',
    'activity': 'За активность',
}

BASIS_LABELS = {
    rules.BASIS_NEW: 'новый водитель',
    rules.BASIS_SINCE_LAST: 'с прошлой выдачи',
    rules.BASIS_WEEK: 'за 7 дней',
}

COLUMNS = (
    ('created_at', 'Дата и время', 17),
    ('city', 'Город', 14),
    ('office_name', 'Офис', 26),
    ('issued_by_name', 'Сотрудник', 26),
    ('driver_name', 'Водитель', 30),
    ('driver_phone', 'Телефон', 16),
    ('driver_account_id', 'ID водителя', 34),
    ('driver_park', 'Парк', 24),
    ('driver_tariffs', 'Тариф', 26),
    ('orders_counted', 'Заказы', 10),
    ('orders_basis', 'Заказы считаны', 18),
    ('kind', 'Вид выдачи', 16),
    ('blocks', 'Блоков', 9),
)

TEXT_COLUMNS = ('driver_phone', 'driver_account_id')


def kind_label(code):
    return KIND_LABELS.get(code) or code or ''


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def _as_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value or '').strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text[:19])
    except ValueError:
        return None


def period_label(period_from, period_to):
    if not period_from or not period_to:
        return ''
    if period_from == period_to:
        return period_from.strftime('%d.%m.%Y')
    return '%s — %s' % (period_from.strftime('%d.%m.%Y'), period_to.strftime('%d.%m.%Y'))


def report_filename(period_from, period_to):
    """«Выдачи воды 01.09.2026 — 30.09.2026.xlsx» — по выбранному периоду."""
    label = period_label(period_from, period_to)
    return ('Выдачи воды %s.xlsx' % label) if label else 'Выдачи воды.xlsx'


def _text(sheet, value, *, keep_format=False):
    """Строковая ячейка, которая ГАРАНТИРОВАННО остаётся строкой.

    openpyxl выводит тип из значения, и строка с ведущим «=» становится
    ФОРМУЛОЙ. ФИО и парк приходят из CRM, где их вводят сами парки и
    водители, а поиск и отбор на листе «Контекст» набирает человек: без явного
    типа `=HYPERLINK(…)` в имени парка открывался бы живой формулой. Тот же
    приём, что parcels/report.py::_text. Формат «текст» — только непустой
    ячейке телефона и ID: у пустой он превращает «нет данных» в ноль.
    """
    text = clean(value)
    cell = WriteOnlyCell(sheet, value=text)
    if text is not None:
        cell.data_type = 's'
        if keep_format and text:
            cell.number_format = '@'
    return cell


def _header_row(sheet, titles):
    row = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        row.append(cell)
    return row


def _issue_row(sheet, item):
    row = []
    for key, _title, _width in COLUMNS:
        value = item.get(key)
        if key == 'created_at':
            cell = WriteOnlyCell(sheet, value=_as_datetime(value))
            cell.number_format = DATETIME_FORMAT
        elif key == 'driver_tariffs':
            cell = _text(sheet, ', '.join(rules.tariff_label(code) for code in (value or [])))
        elif key == 'orders_basis':
            cell = _text(sheet, BASIS_LABELS.get(value) or '')
        elif key == 'kind':
            cell = _text(sheet, kind_label(value))
        elif key in TEXT_COLUMNS:
            cell = _text(sheet, str(value) if value else None, keep_format=True)
        elif isinstance(value, str):
            cell = _text(sheet, value)
        else:
            cell = WriteOnlyCell(sheet, value=value)
        row.append(cell)
    return row


def build_workbook(items, *, period_from, period_to, generated_by='', filters_note='',
                   total=None, generated_at=None, text_warning_patch=None):
    """Книга журнала. Возвращает (поток, сколько строк записано)."""
    workbook = Workbook(write_only=True)
    issues = workbook.create_sheet(SHEET_ISSUES)
    for index, (_key, _title, width) in enumerate(COLUMNS, start=1):
        issues.column_dimensions[get_column_letter(index)].width = width
    issues.freeze_panes = 'A2'
    issues.append(_header_row(issues, [title for _key, title, _width in COLUMNS]))
    written = 0
    blocks = 0
    for item in items:
        issues.append(_issue_row(issues, item))
        written += 1
        blocks += int(item.get('blocks') or 0)

    context = workbook.create_sheet(SHEET_CONTEXT)
    context.column_dimensions['A'].width = 26
    context.column_dimensions['B'].width = 70
    title = WriteOnlyCell(context, value='Учёт воды — журнал выдач')
    title.font = TITLE_FONT
    context.append([title])
    moment = generated_at or rules_now()
    lines = [
        ('Период', period_label(period_from, period_to)),
        ('Собран', moment.strftime('%d.%m.%Y %H:%M')),
        ('Собрал', generated_by or ''),
        ('Отбор', filters_note or 'весь журнал за период'),
        ('Выдач в файле', written),
        ('Блоков выдано', blocks),
    ]
    if total is not None and total > written:
        lines.append(('Внимание', 'В файл вошли первые %d выдач из %d — сузьте период или отбор'
                      % (written, total)))
    for label, value in lines:
        value_cell = _text(context, value) if isinstance(value, str) else WriteOnlyCell(context, value=value)
        value_cell.alignment = WRAP_TOP
        context.append([label, value_cell])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    if callable(text_warning_patch) and written:
        # «Выдачи» — первый лист книги, значит sheet1.xml. Текстовые колонки
        # разбросаны — sqref перечисляет их через пробел.
        sqref = ' '.join(
            '{0}2:{0}{1}'.format(get_column_letter(_column_index(key)), written + 1)
            for key in TEXT_COLUMNS)
        try:
            stream = text_warning_patch(stream, sqref, sheet_path='xl/worksheets/sheet1.xml')
        except Exception:  # noqa: BLE001 — значок в углу не повод не отдать файл
            stream.seek(0)
    return stream, written


def _column_index(key):
    return next(index for index, (name, _title, _width) in enumerate(COLUMNS, start=1)
                if name == key)


def rules_now():
    """Сейчас по Алматы — отдельной функцией, чтобы тест мог её подменить."""
    from .queries import now_almaty
    return now_almaty()
