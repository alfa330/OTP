# -*- coding: utf-8 -*-
"""Черновик акта в DOCX (ТЗ п. 9): «чтобы супервайзер не набирал его вручную».

Шаблон отдела загружает руководитель (DOCX с метками вида {{ФИО}}). Нет
шаблона — стандартный акт по образцу организации (владелец, 23.10.2026,
прислал пример акта: «из него сделай шаблон»): «Акт о совершении работником
дисциплинарного нарушения», город и дата, «Мною, супервайзером … составлен
настоящий акт о нижеследующем», нарушение, решение и испытательный срок,
предупреждение о повторном нарушении и три строки подписей — составил,
ознакомлен сотрудник, ознакомлен руководитель.

Что стандартный акт берёт из базы, а что оставляет руке:
* номер акта — в заголовке: по нему подписанный скан находит свой акт;
* составитель — тот, кто указан в акте, иначе тот, кто скачивает черновик;
  «супервайзером» или «руководителем» — по его месту в отделе акта;
* нарушение-основание — дата, время, тип и описание обстоятельств; все
  нарушения периода (ТЗ п. 9: «перечень всех нарушений периода с датами,
  типами и номерами записей, ссылки на объяснительные») — списком, без
  описаний: запись может быть вне доступа того, кто скачивает черновик;
* решение — мера из акта; пока её нет — линия, вписать от руки;
* испытательный срок — срок особого контроля, начатого актом (слово взято
  из образца организации);
* ФИО в строках подписей подставлены, подписи и даты ставят от руки.

Метки шаблона отдела:
    {{НОМЕР_АКТА}} {{ДАТА}} {{ФИО}} {{ДОЛЖНОСТЬ}} {{ОТДЕЛ}} {{ГРУППА}}
    {{ДАТА_ПРИЁМА}} {{НАРУШЕНИЯ}} {{ОБЪЯСНИТЕЛЬНЫЕ}} {{РЕШЕНИЕ}} {{СОСТАВИЛ}}

Табельного номера в портале нет (ни в одной карточке сотрудника его не
заводили), поэтому и метки под него нет — её отсутствие честнее пустой строки.

Метка в Word нередко разрезана на несколько «прогонов» с разным оформлением
(Word делит текст при проверке орфографии), поэтому замена идёт по тексту
абзаца целиком: найденный абзац переписывается в первый прогон. Оформление
внутри такого абзаца выравнивается по первому прогону — для бланка акта это
приемлемо, а не найти метку было бы хуже.
"""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from . import catalog, docx_layout as layout, files

DOCX_MIME = layout.DOCX_MIME

CITY = 'Алматы'
BLANK = '_' * 34
SIGNATURE = 'Подпись: _______'
# Ширины столбцов строк подписей: кто — подпись — ФИО (в сумме ширина текста).
# Посчитаны по метрикам Times New Roman 12: «С актом ознакомлен руководитель
# Тез КЦ» — 7,55 см, «Подпись: _______» — 3,26 см; длинное ФИО переносится.
SIGNATURE_WIDTHS_CM = (7.8, 3.5, 5.2)

POSITIONS = {
    'supervisor': ('Супервайзер', 'супервайзером'),
    'head': ('Руководитель', 'руководителем'),
}

# Первое слово названия отдела в родительном падеже: «Служба заботы о
# водителях» → «Службы заботы о водителях». Сокращение («Тез КЦ») не
# склоняется, поэтому незнакомое первое слово оставляет название как есть.
_GENITIVE_HEADS = {
    'служба': 'службы', 'отдел': 'отдела', 'департамент': 'департамента', 'центр': 'центра',
    'управление': 'управления', 'дирекция': 'дирекции', 'группа': 'группы', 'сектор': 'сектора',
    'колл-центр': 'колл-центра', 'контакт-центр': 'контакт-центра',
}

_fmt_date = layout.fmt_date


def _fmt_moment(value):
    return value.strftime('%d.%m.%Y %H:%M') if hasattr(value, 'strftime') else str(value or '')


def _explanation(record):
    if record.get('refused_at'):
        return 'отказ, акт об отказе от %s' % _fmt_date(record['refused_at'])
    if record.get('explanation_received_at'):
        return 'получена %s' % _fmt_date(record['explanation_received_at'])
    return 'не предоставлена'


def placeholders(act, operator, records, *, position, today, composed_by):
    lines = ['%d. %s — %s (запись %s)' % (index, _fmt_moment(r.get('occurred_at')), r.get('type_title') or '',
                                         r.get('number') or '')
             for index, r in enumerate(records, start=1)]
    explanations = ['%s — %s' % (r.get('number') or '', _explanation(r)) for r in records]
    return {
        '{{НОМЕР_АКТА}}': act.get('number') or '',
        '{{ДАТА}}': _fmt_date(act.get('composed_on') or today),
        '{{ФИО}}': operator.get('name') or '',
        '{{ДОЛЖНОСТЬ}}': position or '',
        '{{ОТДЕЛ}}': operator.get('department_name') or '',
        '{{ГРУППА}}': operator.get('group_name') or '',
        '{{ДАТА_ПРИЁМА}}': _fmt_date(operator.get('hire_date')) if operator.get('hire_date') else '',
        '{{НАРУШЕНИЯ}}': '\n'.join(lines),
        '{{ОБЪЯСНИТЕЛЬНЫЕ}}': '\n'.join(explanations),
        '{{РЕШЕНИЕ}}': act.get('decision_title') or '',
        '{{СОСТАВИЛ}}': composed_by or '',
    }


def _replace_in_paragraph(paragraph, values):
    text = paragraph.text
    if '{{' not in text:
        return
    replaced = text
    for key, value in values.items():
        replaced = replaced.replace(key, layout.clean(value))
    runs = paragraph.runs
    if replaced == text or not runs:
        return
    # Текст — в первый прогон, где он был: прогон-картинка перед ним (логотип
    # в начале строки) остаётся на своём месте.
    target = next((run for run in runs if run.text), runs[0])
    for run in runs:
        _set_run_text(run, replaced if run is target else '')


# Что в прогоне не текст и должно пережить подстановку: картинка, надпись
# (текстовое поле Word 2010+ лежит в mc:AlternateContent, старое — в w:pict),
# внедрённый объект.
_RUN_KEEP = frozenset((qn('w:drawing'), qn('w:pict'), qn('w:object'),
                       '{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent'))


def _set_run_text(run, text):
    """run.text = text, но картинки и надписи прогона остаются: присваивание
    python-docx стирает всё содержимое прогона, и надпись из абзаца с меткой
    («УТВЕРЖДАЮ» рядом с «АКТ № …») пропадала из черновика."""
    kept = [child for child in run._r if child.tag in _RUN_KEEP]
    run.text = text
    for child in kept:
        run._r.append(child)


def _walk_paragraphs(element, parent):
    """Все абзацы части — тело, таблицы любой вложенности, надписи, поля
    содержимого — и каждый ровно один раз. Обход через row.cells отдавал
    ячейку на N колонок сетки N раз, и вложенные таблицы проходились
    N^глубина раз: шаблон на 33 КБ держал сервер 74 с (замер 30.10.2026).
    Список собирается до подстановок: живой обход дерева, которое тут же
    правится, обрывался на первом абзаце с надписью."""
    for paragraph in list(element.iter(qn('w:p'))):
        yield Paragraph(paragraph, parent)


def fill_template(template_bytes, values):
    # Шаблон проверен при загрузке, но мог лечь в хранилище раньше этих
    # проверок, — без них не разбираем: тяжёлый шаблон занял бы память сервера,
    # а внешняя ссылка ушла бы в черновик. Скачивает черновик СВ — ему причина
    # и совет загрузчику ни к чему, ему — кто и что должен сделать.
    try:
        files.check_docx(template_bytes, name='Шаблон акта', template=True)
    except files.FileError:
        raise files.FileError('Шаблон акта отдела не проходит проверку — руководителю нужно загрузить его заново '
                              'в «Настройках». Пока шаблон не заменён, черновик не собирается',
                              code='DISCIPLINE_TEMPLATE_REFUSED', status=409) from None
    document = Document(BytesIO(template_bytes))
    for paragraph in _walk_paragraphs(document.element.body, document._body):
        _replace_in_paragraph(paragraph, values)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in _walk_paragraphs(part._element, part):
                _replace_in_paragraph(paragraph, values)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── Стандартный акт — по образцу организации ────────────────────────────────

def department_forms(name):
    """(полное в родительном падеже, краткое) из названия отдела.

    «СЗоВ — Служба заботы о водителях» → («Службы заботы о водителях», «СЗоВ»):
    в тексте акта — полное, как в образце («супервайзером Службы заботы о
    водителях»), в строках подписей — краткое («Супервайзер СЗоВ»). Без
    сокращения в названии краткое — то же полное («Супервайзер Отдела продаж»)."""
    name = ' '.join(str(name or '').split())
    short, full = '', name
    for dash in (' — ', ' – ', ' - '):
        if dash in name:
            short, full = (part.strip() for part in name.split(dash, 1))
            break
    head, _space, rest = full.partition(' ')
    genitive_head = _GENITIVE_HEADS.get(head.lower())
    if genitive_head:
        if head[:1].isupper():
            genitive_head = genitive_head[:1].upper() + genitive_head[1:]
        full = ('%s %s' % (genitive_head, rest)).strip()
    return full, short or full


def _lower_first(text):
    """«Опоздание к началу смены» → «опоздание …» посреди фразы; сокращение
    («ПК не включён») остаётся как есть."""
    text = str(text or '').strip()
    if len(text) > 1 and text[0].isupper() and not text[1].isupper():
        return text[0].lower() + text[1:]
    return text


def _sentence(text):
    text = ' '.join(str(text or '').split())
    if text and text[-1] not in '.!?…':
        text += '.'
    return text


def _verb(gender):
    return {'male': 'допустил', 'female': 'допустила'}.get(gender, 'допустил(а)')


def _days(count):
    return '%d %s' % (count, layout.plural(count, 'день', 'дня', 'дней'))


def _period_words(context, trigger):
    if context.get('period_kind') == catalog.PERIOD_MONTH and hasattr(trigger.get('occurred_at'), 'month'):
        moment = trigger['occurred_at']
        return '%s %d г.' % (layout.MONTHS_NOMINATIVE[moment.month - 1], moment.year)
    days = int(context.get('period_days') or 30)
    return 'последние %s' % _days(days)


def _record_tail(record):
    """Что с объяснительной по записи — для строки списка нарушений периода."""
    if record.get('refused_at'):
        return 'отказ от объяснения, акт об отказе от %s' % _fmt_date(record['refused_at'])
    if record.get('explanation_received_at'):
        return 'объяснительная получена %s' % _fmt_date(record['explanation_received_at'])
    if record.get('action') == catalog.STEP_TRAINING:
        return 'проведён тренинг' if record.get('training_id') else ''
    if record.get('needs_explanation_now', True):
        return 'объяснительная не предоставлена'
    return ''


def _decision(act):
    """«применить меру «Выговор»» — мера из справочника звучит существительным,
    поэтому она в кавычках после глагола; пока решения нет — линия."""
    title = str(act.get('decision_title') or '').strip()
    if not title:
        return BLANK
    text = 'применить меру «%s»' % title
    amount, percent = act.get('decision_amount'), act.get('decision_percent')
    if amount:
        text += ' в размере %s' % _money(amount)
    elif percent:
        text += ' в размере %s\u00a0%%' % ('%g' % float(percent)).replace('.', ',')
    return text


def _money(amount):
    """«20 000 ₸», «1 500,75 ₸»: тийыны не теряются, число и знак не разрываются."""
    value = float(amount)
    text = ('{:,.2f}' if value % 1 else '{:,.0f}').format(value)
    return text.replace(',', ' ').replace('.', ',').replace(' ', '\u00a0') + '\u00a0₸'


def standard_act(act, operator, records, *, today=None, composed_by=None, context=None, created=None):
    """Акт по образцу организации. records — неотменённые записи акта по
    времени; context — то, чего нет в акте и записях (см. service.act_draft_context):
    department_name, composer_name и composer_role ('supervisor' | 'head'),
    head_name, gender, control_days, control_end, period_kind, period_days."""
    context = context or {}
    composed_by = context.get('composer_name') or composed_by
    composed_on = act.get('composed_on') or today
    document = layout.new_document(12, title='Акт %s' % (act.get('number') or '').strip(), created=created)
    department, department_short = department_forms(context.get('department_name'))
    position, position_by = POSITIONS.get(context.get('composer_role'), POSITIONS['supervisor'])
    records = list(records or [])
    trigger = next((r for r in records if r.get('id') == act.get('trigger_record_id')),
                   records[-1] if records else {})

    title = 'Акт № %s' % act['number'] if act.get('number') else 'Акт'
    layout.paragraph(document, '%s\nо совершении работником дисциплинарного нарушения' % title, bold=True,
                     align='center', after=30)
    place = layout.paragraph(document, 'г. %s\t%s г.' % (CITY, _fmt_date(composed_on)), bold=True, after=30)
    place.paragraph_format.tab_stops.add_tab_stop(Cm(layout.TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)

    author = ' '.join(part for part in (position_by, department, composed_by or BLANK) if part)
    layout.paragraph(document, 'Мною, %s, составлен настоящий акт о нижеследующем:' % author, before=5, after=5)

    moment = trigger.get('occurred_at')
    if hasattr(moment, 'strftime'):
        when = '%s в %s' % (layout.day_words(moment, year=moment.year != getattr(composed_on, 'year', None)),
                            moment.strftime('%H:%M'))
    else:
        when = BLANK
    happened = 'Сотрудник %s %s %s нарушение трудовой дисциплины: %s' % (
        operator.get('name') or BLANK, when, _verb(context.get('gender')),
        _sentence(_lower_first(trigger.get('type_title')) or BLANK))
    description = _sentence(trigger.get('description'))
    layout.paragraph(document, '%s %s' % (happened, description) if description else happened, before=5, after=5)

    if len(records) > 1:
        layout.paragraph(document, 'Это %s нарушение сотрудника за %s:' % (
            catalog.ordinal_word(len(records)), _period_words(context, trigger)), before=5, after=2, keep_next=True)
        for index, record in enumerate(records, start=1):
            line = '%d. %s, %s — %s (запись %s)' % (
                index, _fmt_date(record.get('occurred_at')),
                record['occurred_at'].strftime('%H:%M') if hasattr(record.get('occurred_at'), 'strftime') else '',
                _lower_first(record.get('type_title')), record.get('number') or '')
            tail = _record_tail(record)
            line += (', %s' % tail if tail else '') + ('.' if index == len(records) else ';')
            layout.paragraph(document, line, left_cm=0.6, after=1, keep_next=index < len(records))

    control_days = context.get('control_days')
    probation = ''
    if control_days:
        probation = ' и установить испытательный срок продолжительностью %s%s для контроля соблюдения ' \
                    'трудовой дисциплины' % (_days(control_days), ' (до %s)' % _fmt_date(context['control_end'])
                                             if context.get('control_end') else '')
    layout.paragraph(document, 'В связи с данным нарушением принято решение %s%s.' % (_decision(act), probation),
                     before=5, after=5)
    layout.paragraph(document, 'В случае повторного нарушения%s сотруднику будет вынесено окончательное '
                               'предупреждение с последующим увольнением.' % (
                                   ' в течение %s' % _days(control_days) if control_days else ''),
                     before=5, after=5)

    layout.paragraph(document, 'Содержание настоящего Акта подтверждаем личными подписями:', before=36, after=12,
                     keep_next=True)
    layout.paragraph(document, 'Настоящий акт составил:', keep_next=True)
    table = layout.borderless_table(document, SIGNATURE_WIDTHS_CM)
    rows = (
        ('%s %s' % (position, department_short), composed_by),
        ('С актом ознакомлен сотрудник %s' % department_short, operator.get('name')),
        ('С актом ознакомлен руководитель %s' % department_short, context.get('head_name')),
    )
    for index, (who, name) in enumerate(rows):
        layout.table_row(table, (who.strip(), SIGNATURE, 'ФИО: %s' % (name or '____________________')),
                         SIGNATURE_WIDTHS_CM, before=0 if index == 0 else 14)
    # Документ не кончается таблицей: Word дописал бы абзац сам и сдвинул вёрстку.
    layout.paragraph(document)
    return layout.to_bytes(document)


def build(act, operator, records, *, template_bytes=None, position=None, today=None, composed_by=None,
          context=None, created=None):
    """Байты DOCX: по шаблону отдела, если он есть, иначе стандартный акт по образцу организации."""
    if template_bytes:
        values = placeholders(act, operator, records, position=position, today=today, composed_by=composed_by)
        return fill_template(template_bytes, values)
    return standard_act(act, operator, records, today=today, composed_by=composed_by, context=context,
                        created=created)
