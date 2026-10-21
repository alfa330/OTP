# -*- coding: utf-8 -*-
"""Отчёты и выгрузки раздела «Дисциплина» (ТЗ п. 4.2, 14, 15).

Отчёты считаются здесь из уже обогащённых записей (service.enrich): номер,
ступень и статус — вычислимые поля, и считать их второй раз в SQL значило бы
завести вторую копию правил лестницы.

Соглашения выгрузок портала: лист «Контекст» первым (на какую дату собрано и
что отобрано — через месяц иначе не понять), даты настоящими датами, книга
потоковая (write_only), стили — модульными константами.

«Выгрузка на удержания» (п. 3, 7.3, 10.4, 12) — это отчёт «По мерам»: все
утверждённые меры с финансовым эффектом, включая процент и «возможен» без
суммы. В расчёт зарплаты портал меру сам не пишет — удержание из заработка
решают бухгалтерия и юрист (см. отчёт владельцу).
"""

from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import catalog

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
LABEL_FONT = Font(bold=True)
WRAP = Alignment(vertical='top', wrap_text=True)
DATETIME_FORMAT = 'DD.MM.YYYY HH:MM'
DATE_FORMAT = 'DD.MM.YYYY'
XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

REPORTS = OrderedDict([
    ('operators', 'Сводка по операторам'),
    ('types', 'По типам нарушений'),
    ('supervisors', 'По супервайзерам'),
    ('groups', 'По группам и линиям'),
    ('measures', 'По мерам (удержания)'),
    ('controls', 'Особый контроль'),
])


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def _parse(value):
    if isinstance(value, str) and len(value) >= 10 and value[4] == '-' and value[7] == '-':
        try:
            if len(value) > 10:
                return datetime.strptime(value[:19].replace('T', ' '), '%Y-%m-%d %H:%M:%S')
            return datetime.strptime(value, '%Y-%m-%d').date()
        except ValueError:
            return value
    return value


def _cell(ws, value, *, font=None, fill=None, align=None, number_format=None):
    cell = WriteOnlyCell(ws, value=_clean(value))
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if number_format:
        cell.number_format = number_format
    return cell


def _write_context(wb, title, lines):
    ws = wb.create_sheet('Контекст')
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 70
    ws.append([_cell(ws, title, font=TITLE_FONT)])
    ws.append([])
    for label, value in lines:
        ws.append([_cell(ws, label, font=LABEL_FONT), _cell(ws, value, align=WRAP)])


def _write_table(wb, title, columns, rows):
    ws = wb.create_sheet(title[:31])
    for index, (_header, _key, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'A2'
    ws.append([_cell(ws, header, font=HEADER_FONT, fill=HEADER_FILL, align=HEADER_ALIGN)
               for header, _key, _width in columns])
    for row in rows:
        cells = []
        for _header, key, _width in columns:
            value = _parse(row.get(key))
            fmt = None
            if isinstance(value, datetime):
                fmt = DATETIME_FORMAT
            elif isinstance(value, date):
                fmt = DATE_FORMAT
            cells.append(_cell(ws, value, number_format=fmt, align=WRAP if isinstance(value, str) else None))
        ws.append(cells)


def workbook(title, context_lines, sheets):
    """sheets — [(название, колонки [(заголовок, ключ, ширина)], строки [dict])]."""
    wb = Workbook(write_only=True)
    _write_context(wb, title, context_lines)
    for sheet_title, columns, rows in sheets:
        _write_table(wb, sheet_title, columns, rows)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── Журнал нарушений ─────────────────────────────────────────────────────────

def _explanation_text(item):
    if item.get('refused_at'):
        return 'Отказ (акт об отказе)'
    if item.get('explanation_received_at'):
        return 'Получена'
    if item.get('needs_explanation_now'):
        return 'Просрочена' if item.get('explanation_overdue') else 'Запрошена'
    return '—'


def journal_rows(items):
    rows = []
    for item in items:
        act = item.get('act') or {}
        rows.append({
            'number': item.get('number'),
            'occurred_at': item.get('occurred_at'),
            'operator_name': item.get('operator_name'),
            'group_name': item.get('group_name'),
            'direction_name': item.get('direction_name'),
            'shift_label': item.get('shift_label'),
            'type_title': item.get('type_title'),
            'weight': item.get('weight'),
            'ordinal': item.get('ordinal'),
            'step_title': item.get('step_title'),
            'state_label': item.get('state_label'),
            'explanation': _explanation_text(item),
            'act': ('%s%s' % (act.get('number') or '', ' (подписан)' if act.get('has_file') else ' (черновик)'))
            if act else '',
            'measure_title': item.get('measure_title'),
            'measure_amount': float(item['measure_amount']) if item.get('measure_amount') is not None else None,
            'measure_percent': float(item['measure_percent']) if item.get('measure_percent') is not None else None,
            'created_by_name': item.get('created_by_name'),
            'closed_at': item.get('closed_at'),
            'cancelled_at': item.get('cancelled_at'),
            'cancel_reason': item.get('cancel_reason'),
            'description': item.get('description'),
        })
    return rows


JOURNAL_COLUMNS = [
    ('Номер', 'number', 14), ('Дата и время', 'occurred_at', 17), ('Оператор', 'operator_name', 28),
    ('Группа', 'group_name', 18), ('Линия', 'direction_name', 20), ('Смена', 'shift_label', 13),
    ('Тип нарушения', 'type_title', 34), ('Вес', 'weight', 6), ('№ в периоде', 'ordinal', 10),
    ('Ступень', 'step_title', 20), ('Статус', 'state_label', 22), ('Объяснительная', 'explanation', 18),
    ('Акт', 'act', 22), ('Мера', 'measure_title', 30), ('Сумма, ₸', 'measure_amount', 12),
    ('Процент', 'measure_percent', 10), ('Зафиксировал', 'created_by_name', 26), ('Закрыта', 'closed_at', 17),
    ('Отменена', 'cancelled_at', 17), ('Причина отмены', 'cancel_reason', 30), ('Описание', 'description', 60),
]


# ── Отчёты п. 15 ─────────────────────────────────────────────────────────────

def report_operators(items):
    groups = OrderedDict()
    for item in sorted(items, key=lambda i: (i.get('operator_name') or '', i.get('occurred_at') or '')):
        key = item['operator_id']
        row = groups.setdefault(key, {
            'operator_name': item.get('operator_name'), 'group_name': item.get('group_name'),
            'total': 0, 'types': Counter(), 'explanations': 0, 'refusals': 0, 'acts': 0, 'measures': Counter(),
        })
        row['total'] += 1
        row['types'][item.get('type_title') or ''] += 1
        if item.get('explanation_received_at'):
            row['explanations'] += 1
        if item.get('refused_at'):
            row['refusals'] += 1
        # Акт — с подписанным файлом: черновик, заведённый системой, ещё не акт.
        if (item.get('act') or {}).get('is_trigger') and (item.get('act') or {}).get('has_file'):
            row['acts'] += 1
        if item.get('measure_state') == 'approved' and item.get('measure_title'):
            row['measures'][item['measure_title']] += 1
    result = []
    for row in groups.values():
        result.append({
            'operator_name': row['operator_name'], 'group_name': row['group_name'], 'total': row['total'],
            'types': ', '.join('%s — %d' % (t, n) for t, n in row['types'].most_common()),
            'explanations': row['explanations'], 'refusals': row['refusals'], 'acts': row['acts'],
            'measures': ', '.join('%s — %d' % (m, n) for m, n in row['measures'].most_common()),
        })
    result.sort(key=lambda r: (-r['total'], r['operator_name'] or ''))
    return result


OPERATORS_COLUMNS = [
    ('Оператор', 'operator_name', 30), ('Группа', 'group_name', 20), ('Нарушений', 'total', 11),
    ('Типы', 'types', 60), ('Объяснительных', 'explanations', 15), ('Отказов', 'refusals', 10),
    ('Актов', 'acts', 8), ('Меры', 'measures', 40),
]


def report_types(items):
    months = sorted({(item.get('occurred_at') or '')[:7] for item in items if item.get('occurred_at')})
    table = defaultdict(Counter)
    for item in items:
        table[item.get('type_title') or ''][(item.get('occurred_at') or '')[:7]] += 1
    rows = []
    for title, counter in sorted(table.items(), key=lambda kv: -sum(kv[1].values())):
        row = {'type_title': title, 'total': sum(counter.values())}
        for month in months:
            row[month] = counter.get(month, 0)
        rows.append(row)
    columns = [('Тип нарушения', 'type_title', 40), ('Всего', 'total', 9)]
    columns += [(_month_label(month), month, 11) for month in months]
    return rows, columns


_MONTHS = ('янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек')


def _month_label(key):
    try:
        year, month = key.split('-')
        return '%s %s' % (_MONTHS[int(month) - 1], year)
    except (ValueError, IndexError):
        return key


def report_supervisors(items):
    """Сколько зафиксировал каждый: перекос в любую сторону — сигнал (п. 15.3)."""
    counter = Counter()
    operators = defaultdict(set)
    for item in items:
        name = item.get('created_by_name') or '—'
        counter[name] += 1
        operators[name].add(item['operator_id'])
    return [{'name': name, 'total': total, 'operators': len(operators[name])}
            for name, total in counter.most_common()]


SUPERVISORS_COLUMNS = [('Зафиксировал', 'name', 32), ('Нарушений', 'total', 11), ('Операторов', 'operators', 12)]


def report_groups(items):
    counter = Counter()
    operators = defaultdict(set)
    for item in items:
        key = (item.get('group_name') or 'Без группы', item.get('direction_name') or '—')
        counter[key] += 1
        operators[key].add(item['operator_id'])
    return [{'group_name': g, 'direction_name': d, 'total': total, 'operators': len(operators[(g, d)])}
            for (g, d), total in counter.most_common()]


GROUPS_COLUMNS = [('Группа', 'group_name', 24), ('Линия', 'direction_name', 26), ('Нарушений', 'total', 11),
                  ('Операторов', 'operators', 12)]


def report_measures(items, decisions=()):
    """Применённые меры: строки для выгрузки на удержания и итог по мерам.

    items — записи (их мера, если утверждена), decisions — меры решением по акту
    и по особому контролю (queries.decision_measures). Порядок — по дате
    утверждения: по ней мера попадает в период выгрузки."""
    rows = []
    for item in items:
        if item.get('measure_state') != 'approved' or not item.get('measure_title'):
            continue
        rows.append({
            'approved_at': item.get('measure_approved_at'), 'basis': 'Нарушение %s' % (item.get('number') or ''),
            'number': item.get('number'), 'occurred_at': item.get('occurred_at'),
            'operator_name': item.get('operator_name'), 'group_name': item.get('group_name'),
            'measure_title': item.get('measure_title'),
            'finance': catalog.FINANCE_LABELS.get(item.get('measure_finance') or 'none', ''),
            'measure_amount': float(item['measure_amount']) if item.get('measure_amount') is not None else None,
            'measure_percent': float(item['measure_percent']) if item.get('measure_percent') is not None else None,
            'approved_by': item.get('measure_approved_by_name') or ('без утверждения'
                                                                     if item.get('measure_approved_at') else None),
            'record_id': item.get('id'),
        })
    for row in decisions or ():
        basis = 'Акт %s' % row['number'] if row['source'] == 'act' else 'Решение по особому контролю'
        rows.append({
            'approved_at': row.get('applied_at'), 'basis': basis, 'number': row.get('number'),
            'occurred_at': row.get('occurred_at'), 'operator_name': row.get('operator_name'),
            'group_name': row.get('group_name'), 'measure_title': row.get('measure_title'),
            'finance': catalog.FINANCE_LABELS.get(row.get('measure_finance') or 'none', ''),
            'measure_amount': float(row['measure_amount']) if row.get('measure_amount') is not None else None,
            'measure_percent': float(row['measure_percent']) if row.get('measure_percent') is not None else None,
            'approved_by': row.get('applied_by_name'), 'record_id': row.get('record_id'),
        })
    rows.sort(key=lambda row: (row['approved_at'] is None, row['approved_at'] or 0))
    totals = OrderedDict()
    for row in rows:
        total = totals.setdefault(row['measure_title'], {'measure_title': row['measure_title'],
                                                         'finance': row['finance'], 'count': 0, 'amount': 0.0})
        total['count'] += 1
        total['amount'] += row['measure_amount'] or 0.0
    return rows, list(totals.values())


MEASURES_COLUMNS = [
    ('Утверждена', 'approved_at', 17), ('Основание', 'basis', 26), ('Оператор', 'operator_name', 28),
    ('Группа', 'group_name', 18), ('Мера', 'measure_title', 34), ('Финансовый эффект', 'finance', 14),
    ('Сумма, ₸', 'measure_amount', 12), ('Процент', 'measure_percent', 10), ('Дата нарушения', 'occurred_at', 17),
    ('Утвердил', 'approved_by', 26),
]
MEASURE_TOTALS_COLUMNS = [('Мера', 'measure_title', 40), ('Финансовый эффект', 'finance', 14),
                          ('Применено', 'count', 11), ('Сумма, ₸', 'amount', 14)]


def report_controls(controls, date_from, date_to):
    """Сколько поставлено, сколько прошло срок без нарушений, сколько дошло до
    расторжения (п. 15.6) — по контролям, начавшимся в периоде."""
    started = [c for c in controls if date_from <= (c.get('start_date') or '') < date_to]
    summary = {
        'started': len(started),
        'by_source': Counter(c.get('source_label') for c in started),
        'clean': sum(1 for c in started if c.get('decided_at') and not c.get('violations')),
        'termination': sum(1 for c in started if c.get('decision') == 'termination'),
        'active': sum(1 for c in started if not c.get('decided_at')),
    }
    return started, summary


CONTROLS_COLUMNS = [
    ('Оператор', 'operator_name', 28), ('Группа', 'group_name', 18), ('Основание', 'source_label', 18),
    ('Причина', 'reason', 40), ('Начало', 'start_date', 12), ('Окончание', 'end_date', 12),
    ('Нарушений в период', 'violations', 12), ('Решение', 'decision_label', 22),
    ('Подробности решения', 'decision_detail', 40), ('Принято', 'decided_at', 17), ('Кто решил', 'decided_by_name', 24),
    ('Ответственный', 'responsible_name', 24),
]


AUDIT_COLUMNS = [
    ('Когда', 'created_at', 17), ('Кто', 'actor_name', 26), ('Действие', 'action_label', 30),
    ('Сотрудник', 'operator_name', 26), ('Объект', 'entity', 12), ('id', 'entity_id', 8),
    ('Изменения', 'changes_text', 80),
]

AUDIT_ACTIONS = {
    'create_record': 'Зафиксировано нарушение',
    'update_record': 'Запись изменена',
    'cancel_record': 'Запись отменена',
    'close_record': 'Запись закрыта',
    'manager_comment': 'Комментарий руководителя',
    'upload_explanation': 'Загружена объяснительная',
    'refusal': 'Отказ от объяснения',
    'training': 'Отмечен тренинг',
    'measure_propose': 'Мера предложена',
    'measure_approve': 'Мера утверждена',
    'measure_apply': 'Мера применена (без утверждения)',
    'explanation_reset': 'Объяснительная снова нужна',
    'version_restored': 'Действующей снова стала прежняя версия файла',
    'act_draft': 'Черновик акта',
    'act_update': 'Акт изменён',
    'act_approve': 'Акт утверждён',
    'upload_act': 'Приложен подписанный акт',
    'act_docx': 'Выгружен черновик акта',
    'upload_evidence': 'Добавлено доказательство',
    'invalidate_file': 'Файл признан недействительным',
    'view_record': 'Просмотр записи',
    'view_card': 'Просмотр карточки сотрудника',
    'view_file': 'Открыт файл',
    'control_auto': 'Поставлен на контроль автоматически',
    'control_create': 'Поставлен на контроль',
    'control_extend': 'Контроль продлён',
    'control_act_linked': 'Акт привязан к контролю',
    'control_decide': 'Решение по контролю',
    'settings_update': 'Изменены настройки',
    'type_save': 'Изменён справочник типов',
    'measure_save': 'Изменён справочник мер',
    'template_upload': 'Загружен шаблон акта',
    'export': 'Выгрузка в Excel',
}


def audit_label(action):
    return AUDIT_ACTIONS.get(action, action)


# Подписи полей журнала действий по-русски. Служебные поля (id, хеши) в текст
# не попадают: человеку они ничего не говорят, а в выгрузке лишь шумят.
CHANGE_LABELS = {
    'number': 'Номер', 'type': 'Тип', 'type_title': 'Тип', 'occurred_at': 'Когда', 'step': 'Ступень',
    'ordinal': '№ в периоде', 'files': 'Файлов', 'measure': 'Мера', 'records': 'Записей в акте',
    'reason': 'Причина', 'start_date': 'Начало', 'end_date': 'Окончание', 'source': 'Основание',
    'file': 'Файл', 'witnesses': 'Свидетели', 'date': 'Дата', 'start': 'С', 'end': 'До',
    'description': 'Описание', 'shift_label': 'Смена', 'evidence_note': 'Тайм-код', 'manager_comment': 'Комментарий',
    'decision': 'Решение', 'comment': 'Комментарий', 'days': 'Дней', 'amount': 'Сумма', 'percent': 'Процент',
    'composed_on': 'Дата составления', 'decision_title': 'Решение', 'approved_at': 'Утверждён', 'weight': 'Вес',
    'needs_explanation': 'Нужна объяснительная', 'needs_evidence': 'Нужно доказательство', 'act': 'Акт',
    'report': 'Отчёт', 'rows': 'Строк', 'title': 'Название', 'severity': 'Тяжесть', 'is_other': '«Иное»',
    'is_archived': 'В архиве', 'needs_approval': 'Требует утверждения', 'finance': 'Финансовый эффект',
    'value_kind': 'Чем выражена', 'period_kind': 'Окно подсчёта', 'period_days': 'Длина окна, дней',
    'active_ladder': 'Действующая лестница', 'ladder': 'Лестница', 'ladders': 'Лестницы отдела',
    'control_days': 'Контроль после акта, дней', 'explanation_days': 'Срок объяснительной',
    'telegram': 'Telegram', 'version': None, 'sha256': None, 'training_id': None, 'file_id': None,
    'type_id': None, 'approved_by': None, 'composed_by': None, 'decision_measure_id': None,
}

_VALUE_LABELS = {
    'source': dict(catalog.CONTROL_SOURCES),
    'severity': catalog.SEVERITY_LABELS,
    'finance': catalog.FINANCE_LABELS,
    'value_kind': catalog.VALUE_LABELS,
    'period_kind': dict(catalog.PERIOD_KINDS),
}


def _human(key, value):
    if value in (None, ''):
        return '—'
    if isinstance(value, bool):
        return 'да' if value else 'нет'
    if key in _VALUE_LABELS and value in _VALUE_LABELS[key]:
        return _VALUE_LABELS[key][value]
    if key == 'ladder' and isinstance(value, list):
        return ', '.join(catalog.STEP_LABELS.get(step, step).lower() for step in value)
    if key == 'ladders' and isinstance(value, list):
        return ' · '.join('%s: %s' % (item.get('name'), _human('ladder', item.get('steps') or []))
                          for item in value if isinstance(item, dict))
    if key == 'telegram' and isinstance(value, dict):
        events = dict((code, label) for code, label, _who, _d in catalog.NOTIFY_EVENTS)
        return ', '.join(events.get(code, code) for code, on in value.items() if on) or 'выключен'
    if isinstance(value, list):
        return ', '.join(str(item) for item in value)
    text = str(value)
    if len(text) == 8 and text[2] == ':' and text[5] == ':':
        return text[:5]
    parsed = _parse(text)
    if isinstance(parsed, datetime):
        return parsed.strftime('%d.%m.%Y %H:%M')
    if isinstance(parsed, date):
        return parsed.strftime('%d.%m.%Y')
    return text[:300]


def changes_text(changes):
    """Изменения строкой для человека: «Смена: 09:00–18:00 → 10:00–19:00; Тип: …»."""
    if not changes:
        return ''
    if not isinstance(changes, dict):
        return str(changes)
    parts = []
    # JSONB хранит ключи в своём порядке — показываем в естественном (как в CHANGE_LABELS).
    order = {key: index for index, key in enumerate(CHANGE_LABELS)}
    for key in sorted(changes, key=lambda k: (order.get(k, len(order)), k)):
        value = changes[key]
        label = CHANGE_LABELS.get(key, key)
        if label is None:
            continue
        is_diff = isinstance(value, list) and len(value) == 2 and (
            key not in ('ladder', 'witnesses') or isinstance(value[0], list))
        if not is_diff and (value in (None, '') or (key == 'files' and not value)):
            continue
        if is_diff:
            parts.append('%s: %s → %s' % (label, _human(key, value[0]), _human(key, value[1])))
        else:
            parts.append('%s: %s' % (label, _human(key, value)))
    return '; '.join(parts)
