# -*- coding: utf-8 -*-
"""Общая вёрстка документов раздела: черновика акта и бланка объяснительной.

Лист A4 с полями бланков организации (слева 3 см под подшивку, справа 1,5,
сверху и снизу 2 — как в шаблоне объяснительной), Times New Roman, русский
язык проверки правописания, интервалы — только те, что заданы абзацу.

Шаблон python-docx — американский: лист Letter, Calibri из темы, интервал
после абзаца, эскиз чужой картинки в docProps и автор «python-docx» с датой
2013 года. Всё это снимается здесь: шрифт задаётся без атрибутов темы (они
сильнее явного имени, и Word показал бы Calibri), эскиз удаляется из пакета.
"""

import re
from datetime import timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
FONT = 'Times New Roman'

PAGE_CM = (21.0, 29.7)
MARGINS_CM = {'left': 3.0, 'right': 1.5, 'top': 2.0, 'bottom': 2.0}
TEXT_WIDTH_CM = PAGE_CM[0] - MARGINS_CM['left'] - MARGINS_CM['right']

ALIGN = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'both': WD_ALIGN_PARAGRAPH.JUSTIFY,
}

MONTHS_GENITIVE = ('января', 'февраля', 'марта', 'апреля', 'мая', 'июня', 'июля', 'августа', 'сентября',
                   'октября', 'ноября', 'декабря')
MONTHS_NOMINATIVE = ('январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль', 'август', 'сентябрь',
                     'октябрь', 'ноябрь', 'декабрь')

_THEME_FONT_ATTRS = ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme')
# Символов, которых нет в XML 1.0, python-docx не принимает («All strings must be XML
# compatible»), а приходят они вставкой из 1С и Excel — в ФИО, названии «Иное», шапке.
_NOT_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
# Время раздела — наивное по Алматы, а python-docx пишет дату файла как UTC.
ALMATY = ZoneInfo('Asia/Almaty')


def clean(text):
    """Текст для документа: без символов, которых нет в XML."""
    return _NOT_XML.sub('', str(text))


def fmt_date(value):
    return value.strftime('%d.%m.%Y') if hasattr(value, 'strftime') else str(value or '')


def day_words(value, *, year=True):
    """«9 октября 2026 г.» — как дата пишется в тексте документа."""
    if not hasattr(value, 'month'):
        return ''
    text = '%d %s' % (value.day, MONTHS_GENITIVE[value.month - 1])
    return '%s %d г.' % (text, value.year) if year else text


def plural(count, one, few, many):
    count = abs(int(count or 0))
    if count % 10 == 1 and count % 100 != 11:
        return one
    if 2 <= count % 10 <= 4 and not 12 <= count % 100 <= 14:
        return few
    return many


def _fonts(rPr):
    """Шрифт на все алфавиты одним именем, атрибуты темы — прочь."""
    fonts = rPr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rPr.insert(0, fonts)
    for attr in _THEME_FONT_ATTRS:
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), FONT)


def new_document(size_pt, *, title='', created=None):
    """Пустой документ: A4, поля бланков, Times New Roman size_pt, без интервалов."""
    document = Document()
    defaults = document.styles.element.find(qn('w:docDefaults'))
    rPr = defaults.find(qn('w:rPrDefault')).find(qn('w:rPr'))
    _fonts(rPr)
    lang = rPr.find(qn('w:lang'))
    if lang is not None:
        # Язык проверки — русский: иначе Word подчеркнул бы каждое слово.
        lang.set(qn('w:val'), 'ru-RU')
    normal = document.styles['Normal']
    normal.font.size = Pt(size_pt)
    _fonts(normal.element.get_or_add_rPr())
    fmt = normal.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # И в настройках по умолчанию — одинарный без отступа после: на них стоят
    # абзацы не от «Normal» (колонтитул), а шаблон даёт 1,15 и 10 пт.
    spacing = defaults.find(qn('w:pPrDefault')).find(qn('w:pPr')).find(qn('w:spacing'))
    if spacing is not None:
        for key, value in (('w:after', '0'), ('w:line', '240'), ('w:lineRule', 'auto')):
            spacing.set(qn(key), value)
    # Режим Word 2013+: с режимом шаблона (2010) Word пишет в заголовке окна
    # «Режим совместимости».
    compat = document.settings.element.find(qn('w:compat'))
    for setting in (compat.findall(qn('w:compatSetting')) if compat is not None else ()):
        if setting.get(qn('w:name')) == 'compatibilityMode':
            setting.set(qn('w:val'), '15')

    section = document.sections[0]
    section.page_width, section.page_height = Cm(PAGE_CM[0]), Cm(PAGE_CM[1])
    for side, value in MARGINS_CM.items():
        setattr(section, '%s_margin' % side, Cm(value))

    rels = document.part.package.rels
    for rel_id, rel in list(rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            del rels[rel_id]
            rels._target_parts_by_rId.pop(rel_id, None)
    props = document.core_properties
    props.title = title
    props.author = 'iCore'
    props.last_modified_by = ''
    props.comments = ''
    props.revision = 1
    if created is not None:
        moment = (created if created.tzinfo else created.replace(tzinfo=ALMATY)).astimezone(timezone.utc)
        props.created = props.modified = moment.replace(tzinfo=None)
    return document


def paragraph(container, text='', *, bold=False, size=None, align='left', before=0, after=0, underline=False,
              keep_next=False, indent_cm=None, left_cm=None):
    """Абзац одним прогоном. Перевод строки в тексте — разрыв строки Word.
    indent_cm — красная строка, left_cm — отступ всего абзаца слева."""
    item = container.add_paragraph()
    item.alignment = ALIGN[align]
    fmt = item.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if keep_next:
        fmt.keep_with_next = True
    if indent_cm is not None:
        fmt.first_line_indent = Cm(indent_cm)
    if left_cm is not None:
        fmt.left_indent = Cm(left_cm)
    if text:
        run(item, text, bold=bold, size=size, underline=underline)
    return item


def run(item, text, *, bold=False, size=None, underline=False):
    piece = item.add_run(clean(text))
    piece.bold = bold or None
    piece.underline = underline or None
    if size:
        piece.font.size = Pt(size)
    return piece


def borderless_table(document, widths_cm):
    """Таблица без рамок с жёсткими ширинами столбцов: строки подписей стоят
    ровно, а не по пробелам, как в образце. Слева у ячеек нет поля — текст
    начинается с края, как у абзацев над таблицей."""
    table = document.add_table(rows=0, cols=len(widths_cm))
    table.autofit = False
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    tblPr = table._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for side, twips in (('w:left', 0), ('w:right', 57)):
        edge = OxmlElement(side)
        edge.set(qn('w:w'), str(twips))
        edge.set(qn('w:type'), 'dxa')
        margins.append(edge)
    # По схеме tblCellMar идёт перед tblLook — иначе Word объявит файл повреждённым.
    look = tblPr.find(qn('w:tblLook'))
    if look is not None:
        look.addprevious(margins)
    else:
        tblPr.append(margins)
    return table


def table_row(table, texts, widths_cm, *, before=0, size=None):
    cells = table.add_row().cells
    for cell, text, width in zip(cells, texts, widths_cm):
        cell.width = Cm(width)
        first = cell.paragraphs[0]
        first.paragraph_format.space_before = Pt(before)
        if text:
            run(first, text, size=size)
    return cells


def to_bytes(document):
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
