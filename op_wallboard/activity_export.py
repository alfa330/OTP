# -*- coding: utf-8 -*-
"""Выгрузка «Активности операторов» в Excel (задача #392): те же 14 колонок, что на экране.

Строки — ровно те, что отфильтрованы на экране (`activity.filter_shifts`), посчитаны тем же
расчётом. Длительности и время записаны ЗНАЧЕНИЯМИ времени, а не текстом: в Excel их можно
складывать и сортировать, и зелёного уголка «число сохранено как текст» нет вовсе. Формат
длительности — `[h]:mm:ss`: у обычного `hh:mm:ss` сумма больше суток обнулялась бы по кругу.

Второй лист — «Как считается»: определения колонок словами, чтобы цифры можно было сверить,
не спрашивая разработчика.
"""

from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = (
    ('Дата', 12),
    ('Оператор', 32),
    ('Группа', 10),
    ('Вход в систему', 14),
    ('Среднее время разговора', 14),
    ('Общее время разговора', 14),
    ('Время на исходе / Активный', 15),
    ('Время в ожидании', 13),
    ('Время перерыва', 13),
    ('Время тех. перерыва', 13),
    ('Количество звонков (все попытки)', 15),
    ('Количество дозвонов (успешные разговоры)', 16),
    ('Время выхода из системы', 14),
    ('Фактическое отработанное время', 15),
)

DATE_FORMAT = 'DD.MM.YYYY'
CLOCK_FORMAT = 'HH:MM:SS'
DURATION_FORMAT = '[h]:mm:ss'

_HEADER_FILL = PatternFill('solid', fgColor='F1F5F9')
_HEADER_FONT = Font(bold=True, color='334155')
_MUTED_FONT = Font(italic=True, color='94A3B8')
_NOTE_FONT = Font(color='64748B', size=10)
_THIN = Side(style='thin', color='E2E8F0')


def _clock(stamp_value):
    if not stamp_value:
        return None
    moment = datetime.strptime(str(stamp_value)[:19], '%Y-%m-%dT%H:%M:%S')
    return time(moment.hour, moment.minute, moment.second)


def _duration(seconds):
    return None if seconds is None else timedelta(seconds=int(seconds))


def _day(value):
    return datetime.strptime(str(value), '%Y-%m-%d').date() if value else None


def rule_rows(talk_min_seconds):
    """Определения колонок словами — для листа «Как считается» и подсказок экрана."""
    return (
        ('Смена', 'События телефона iCORE Phone подряд. Новая смена — после выхода и паузы от '
                  '4 часов, а если оператор ждал звонка в рабочем статусе — после тишины от 8 часов. '
                  'Смена относится к дате своего начала: ночная — целиком к вечеру, с которого началась. '
                  'В отчёте только завершённые смены.'),
        ('Вход в систему', 'Первое событие телефона в смене.'),
        ('Среднее время разговора', 'Общее время разговора, делённое на все попытки — включая '
                                    'сбросы и недозвоны с нулевым разговором.'),
        ('Общее время разговора', 'Сумма чистого разговора по всем попыткам смены — от соединения до '
                                  'отбоя по данным станции, без набора и ожидания в очереди.'),
        ('Время на исходе / Активный', 'Всё время в рабочем статусе: «Активный» у Основы, «Исход» и '
                                       '«Автодозвон» у ЯР и Потока — вместе с наборами и разговорами внутри него.'),
        ('Время в ожидании', 'Часть предыдущей колонки без звонка: ни набора номера, ни разговора. '
                             'Уже входит во «Время на исходе / Активный» — складывать их нельзя.'),
        ('Время перерыва', 'Сумма периодов статуса «Перерыв».'),
        ('Время тех. перерыва', 'Сумма периодов статуса «Технический перерыв».'),
        ('Количество звонков (все попытки)', 'ЯР и Поток — все исходящие попытки, включая короткие и '
                                             'сброшенные. Основа — исходящие и входящие, поставленные '
                                             'оператору; входящий, который клиент сбросил, пока звонил '
                                             'телефон, оператору не засчитывается.'),
        ('Количество дозвонов (успешные разговоры)', 'Звонки, где после соединения разговор шёл не меньше '
                                                     '%d секунд.' % int(talk_min_seconds)),
        ('Время выхода из системы', 'Выход из телефона в конце смены. Если телефон выход не прислал, '
                                    'стоит время последнего события смены — серым курсивом.'),
        ('Фактическое отработанное время', 'Время в рабочем статусе плюс тренинг; без перерывов, '
                                           'тех. перерывов и офлайна. Каждая секунда смены лежит ровно '
                                           'в одном статусе — двойного учёта нет. По телефону за всю '
                                           'смену, без обрезки по графику, поэтому может быть больше, '
                                           'чем в «Учёте часов».'),
    )


def workbook(rows, *, month_label, group_label, generated_at, talk_min_seconds=10,
             calls_pending_days=()):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Смены'
    for index, (title, width) in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = Border(bottom=_THIN)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 44
    sheet.freeze_panes = 'A2'

    for offset, row in enumerate(rows, 2):
        values = (
            (_day(row.get('date')), DATE_FORMAT),
            (row.get('name') or '', None),
            (row.get('group_label') or '', None),
            (_clock(row.get('login_at')), CLOCK_FORMAT),
            (_duration(row.get('avg_talk_seconds')), DURATION_FORMAT),
            (_duration(row.get('talk_seconds')), DURATION_FORMAT),
            (_duration(row.get('active_seconds')), DURATION_FORMAT),
            (_duration(row.get('waiting_seconds')), DURATION_FORMAT),
            (_duration(row.get('break_seconds')), DURATION_FORMAT),
            (_duration(row.get('tech_seconds')), DURATION_FORMAT),
            (int(row.get('calls') or 0), '0'),
            (int(row.get('connected') or 0), '0'),
            (_clock(row.get('logout_at')), CLOCK_FORMAT),
            (_duration(row.get('worked_seconds')), DURATION_FORMAT),
        )
        for column, (value, number_format) in enumerate(values, 1):
            cell = sheet.cell(row=offset, column=column, value=value)
            if number_format and value is not None:
                cell.number_format = number_format
        if row.get('logout_missing'):
            sheet.cell(row=offset, column=13).font = _MUTED_FONT

    last_row = max(1, len(rows) + 1)
    if rows:
        sheet.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(COLUMNS)), last_row)
    notes = ['%s · %s · сформировано %s' % (month_label, group_label, generated_at)]
    if not rows:
        notes.insert(0, 'За выбранный период завершённых смен нет.')
    if any(row.get('logout_missing') for row in rows):
        notes.append('Серым курсивом — телефон не прислал выход: стоит время последнего события смены.')
    if calls_pending_days:
        notes.append('Звонки ещё догружаются за %d сут. — колонки звонков могут быть неполными.'
                     % len(calls_pending_days))
    for index, text in enumerate(notes):
        cell = sheet.cell(row=last_row + 2 + index, column=1, value=text)
        cell.font = _NOTE_FONT

    rules = book.create_sheet('Как считается')
    rules.column_dimensions['A'].width = 34
    rules.column_dimensions['B'].width = 110
    for index, (title, text) in enumerate(rule_rows(talk_min_seconds), 1):
        name = rules.cell(row=index, column=1, value=title)
        name.font = Font(bold=True, color='334155')
        name.alignment = Alignment(vertical='top')
        body = rules.cell(row=index, column=2, value=text)
        body.alignment = Alignment(wrap_text=True, vertical='top')

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def file_name(month_value):
    return 'aktivnost_operatorov_%s.xlsx' % (str(month_value or '').strip() or date.today().strftime('%Y-%m'))
