# -*- coding: utf-8 -*-
"""Выгрузка «Табло ОП · Чат» в Excel за период (задача #367).

Считает тот же `chat.py`, что и экран, — у показателя не может быть второй формулы. Раскладка
повторяет выгрузку чатов СЗоВ там, где это имеет смысл (сводка, дни, часы, разрезы «человек ×
час»), но без статусов: у Wazzup их нет, и листа «кто был на линии» здесь быть не может.

Потолок периода свой (`MAX_DAYS`), а не неделя, как у СЗоВ: там каждый прошедший день
качается у Chat2Desk за квоту компании, здесь всё уже лежит в своей базе. Верхняя граница —
хранение переписки (45 суток, `cleanup_wazzup_messages`) и то, что запрос синхронный.
"""

from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import chat as chat_mod

MAX_DAYS = 31
# Переписка хранится 45 суток (cleanup_wazzup_messages): день старше — это не «ноль чатов», а
# удалённые сообщения, и в файле он читался бы как день полного простоя.
RETENTION_DAYS = 45


def parse_period(date_from, date_to, today, max_days=MAX_DAYS):
    """Дни периода по возрастанию. ValueError — период не разобрался или слишком велик.

    Период не выбран — сегодня: самый частый случай (глянул на табло и забрал цифры)."""
    def _day(value, what):
        value = str(value or '').strip()
        if not value:
            return None
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except ValueError:
            raise ValueError(f'{what} периода должно быть в формате ГГГГ-ММ-ДД')

    start, end = _day(date_from, 'Начало'), _day(date_to, 'Конец')
    if start is None and end is None:
        start = end = today
    start, end = start or end, end or start
    if end < start:
        start, end = end, start
    if start > today:
        raise ValueError('Период начинается в будущем — выгружать нечего')
    earliest = today - timedelta(days=RETENTION_DAYS - 1)
    if start < earliest:
        raise ValueError(f'Переписка хранится {RETENTION_DAYS} суток — выберите период '
                         f'не раньше {earliest.strftime("%d.%m.%Y")}')
    # Хвост в будущем отрезаем: пустые сутки в файле читались бы как день полного простоя.
    end = min(end, today)
    length = (end - start).days + 1
    if length > max_days:
        raise ValueError(f'Период больше {max_days} суток — выберите короче')
    return [start + timedelta(days=offset) for offset in range(length)]


def file_name(days):
    """Имя файла: одна дата на однодневной выгрузке, обе — на периоде."""
    first, last = days[0].strftime('%Y%m%d'), days[-1].strftime('%Y%m%d')
    return f'op_wallboard_chat_{first}.xlsx' if first == last else f'op_wallboard_chat_{first}_{last}.xlsx'


def _minutes(seconds):
    return None if seconds is None else round(float(seconds) / 60, 1)


def workbook(blocks, *, first_target_seconds, inner_target_seconds, generated_at=None):
    """Файл .xlsx (bytes) по суткам из `chat.period_days`."""
    thin = Side(style='thin', color='E2E8F0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill('solid', fgColor='1F2937')
    total_fill = PatternFill('solid', fgColor='F1F5F9')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    title_font = Font(bold=True, size=15, color='0F172A')
    hint_font = Font(color='94A3B8', size=10)
    total_font = Font(bold=True, color='0F172A')
    good_font = Font(bold=True, color='15803D')
    bad_font = Font(bold=True, color='B91C1C')
    center = Alignment(horizontal='center', vertical='center')
    wrap = Alignment(vertical='top', wrap_text=True)
    targets = {'first': first_target_seconds, 'inner': inner_target_seconds}

    def _header(ws, row, titles, widths=None):
        for col, title in enumerate(titles, start=1):
            cell = ws.cell(row=row, column=col, value=title)
            cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border
        for col, width in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = ws.cell(row=row + 1, column=2)

    def _put(ws, row, col, value, *, kind=None, seconds=None, total=False):
        cell = ws.cell(row=row, column=col, value=value)
        cell.border = border
        if isinstance(value, (int, float)):
            cell.alignment = center
            if isinstance(value, float):
                cell.number_format = '0.0'
        if kind and seconds is not None:
            # Цвет — только у времени ответа и только относительно нормы, как на табло.
            cell.font = good_font if seconds <= targets[kind] else bad_font
        elif total:
            cell.font = total_font
        if total:
            cell.fill = total_fill
        return cell

    period = f"{blocks[0]['day']} — {blocks[-1]['day']}" if len(blocks) > 1 else blocks[0]['day']
    wb = Workbook()

    # ── Показатели: по дням и итог периода ────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Показатели'
    ws.cell(row=1, column=1, value='Табло ОП · чаты верификаторов').font = title_font
    ws.cell(row=2, column=1, value=(
        f"Период: {period}. Нормы: первый ответ {_minutes(first_target_seconds)} мин, "
        f"ответ внутри чата {_minutes(inner_target_seconds)} мин."
        + (f" Выгружено {generated_at}." if generated_at else ''))).font = hint_font
    ws.cell(row=3, column=1, value=(
        'Диалог — переписка с клиентом до шести часов тишины; день и час диалога — по его началу. '
        'Время ответа — только ответы верификаторов; среднее за период — по диалогам всех дней. '
        '«Верификаторов» — кто писал клиентам сам, рассылки не в счёт.')).font = hint_font
    _header(ws, 5, ['Дата', 'Чатов', 'Первый ответ, мин', 'Ответ внутри чата, мин', 'Верификаторов'],
            [14, 10, 18, 22, 15])
    row = 6
    everyone = set()
    for block in blocks:
        today = block['today']
        people = set(block['operators'])
        everyone |= people
        _put(ws, row, 1, block['day'])
        _put(ws, row, 2, today['chats'])
        _put(ws, row, 3, _minutes(today['first_reply_seconds']), kind='first',
             seconds=today['first_reply_seconds'])
        _put(ws, row, 4, _minutes(today['inner_reply_seconds']), kind='inner',
             seconds=today['inner_reply_seconds'])
        _put(ws, row, 5, len(people))
        row += 1
    if len(blocks) > 1:
        sums = chat_mod.sum_sums([block['today'] for block in blocks])
        first, inner = chat_mod.reply_average(sums, 'first'), chat_mod.reply_average(sums, 'inner')
        _put(ws, row, 1, 'Итого', total=True)
        _put(ws, row, 2, sum(block['today']['chats'] for block in blocks), total=True)
        _put(ws, row, 3, _minutes(first), kind='first', seconds=first, total=True)
        _put(ws, row, 4, _minutes(inner), kind='inner', seconds=inner, total=True)
        _put(ws, row, 5, len(everyone), total=True)

    # ── По часам: каждый час каждого дня ─────────────────────────────────────────────────
    ws = wb.create_sheet('По часам')
    _header(ws, 1, ['Дата', 'Час', 'Чатов начато', 'Первый ответ, мин', 'Ответ внутри чата, мин',
                    'Верификаторов', 'Кто писал (чатов за час)'], [12, 9, 13, 18, 22, 15, 70])
    row = 2
    for block in blocks:
        for hour in block['hourly']:
            if not hour['chats'] and not hour['verifiers']:
                continue
            _put(ws, row, 1, block['day'])
            _put(ws, row, 2, f"{hour['hour']:02d}–{(hour['hour'] + 1) % 24:02d}")
            _put(ws, row, 3, hour['chats'])
            _put(ws, row, 4, _minutes(hour['first_reply_seconds']), kind='first',
                 seconds=hour['first_reply_seconds'])
            _put(ws, row, 5, _minutes(hour['inner_reply_seconds']), kind='inner',
                 seconds=hour['inner_reply_seconds'])
            _put(ws, row, 6, hour['verifiers'])
            _put(ws, row, 7, ', '.join(f"{p['name']} {p['chats']}" for p in hour['operators'])).alignment = wrap
            row += 1

    # ── Верификаторы за период ───────────────────────────────────────────────────────────
    people = {}
    for block in blocks:
        for user_id, person in block['operators'].items():
            entry = people.setdefault(user_id, {'name': person['name'], 'chats': 0, 'days': 0,
                                                'sums': chat_mod.empty_sums()})
            entry['chats'] += person['chats']
            entry['days'] += 1 if person['chats'] else 0
            chat_mod.merge_sums(entry['sums'], person['reply_sums'])
    ws = wb.create_sheet('Верификаторы')
    _header(ws, 1, ['Верификатор', 'Чатов', 'Первый ответ, мин', 'Ответ внутри чата, мин', 'Дней с чатами'],
            [34, 10, 18, 22, 15])
    ws.cell(row=1, column=7, value=(
        'Чатов — пары «чат × сутки», где верификатор писал сам (как в «Чатах ОП»); чат, где писали '
        'двое, засчитан обоим, поэтому сумма по людям больше итога на листе «Показатели».')).font = hint_font
    row = 2
    for entry in sorted(people.values(), key=lambda item: (-item['chats'], item['name'])):
        first = chat_mod.reply_average(entry['sums'], 'first')
        inner = chat_mod.reply_average(entry['sums'], 'inner')
        _put(ws, row, 1, entry['name'])
        _put(ws, row, 2, entry['chats'])
        _put(ws, row, 3, _minutes(first), kind='first', seconds=first)
        _put(ws, row, 4, _minutes(inner), kind='inner', seconds=inner)
        _put(ws, row, 5, entry['days'])
        row += 1

    # ── Разрезы «верификатор × час суток» (за период — сумма по дням) ───────────────────────
    # Чаты — по часу, когда человек писал; время ответа — по часу начала диалога, как в итоге
    # часа (`operators_reply`). Ответ в 11:03 на диалог из 10:58 стоит в колонке «10–11».
    grid = {}
    for block in blocks:
        for hour in block['hourly']:
            for person in hour['operators']:
                cell = grid.setdefault((person['user_id'], hour['hour']),
                                       {'chats': 0, 'sums': chat_mod.empty_sums()})
                cell['chats'] += person['chats']
            for person in hour['operators_reply']:
                cell = grid.setdefault((person['user_id'], hour['hour']),
                                       {'chats': 0, 'sums': chat_mod.empty_sums()})
                chat_mod.merge_sums(cell['sums'], person)
    hours = sorted({hour for _user, hour in grid})
    ordered = sorted(people.items(), key=lambda item: (-item[1]['chats'], item[1]['name']))
    for title, value_of in (
            ('Чаты по часам', lambda cell: cell['chats'] or None),
            ('Первый ответ по часам', lambda cell: chat_mod.reply_average(cell['sums'], 'first')),
            ('Ответ внутри чата по часам', lambda cell: chat_mod.reply_average(cell['sums'], 'inner'))):
        ws = wb.create_sheet(title)
        kind = {'Первый ответ по часам': 'first', 'Ответ внутри чата по часам': 'inner'}.get(title)
        _header(ws, 1, ['Верификатор'] + [f'{hour:02d}–{(hour + 1) % 24:02d}' for hour in hours],
                [34] + [9] * len(hours))
        for index, (user_id, entry) in enumerate(ordered, start=2):
            _put(ws, index, 1, entry['name'])
            for col, hour in enumerate(hours, start=2):
                cell = grid.get((user_id, hour))
                value = value_of(cell) if cell else None
                if kind:
                    _put(ws, index, col, _minutes(value), kind=kind, seconds=value)
                else:
                    _put(ws, index, col, value)

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
