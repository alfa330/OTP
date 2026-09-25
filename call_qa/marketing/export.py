# -*- coding: utf-8 -*-
"""Выгрузка разборов «ИИ-оценки» в XLSX/CSV с атрибутами сделки (ТЗ #317, п. 2.1).

Единица строки — разбор (последняя оценка субъекта), как во вкладке «Звонки».
Отбор — ровно тот, что стоит в панели: выгрузка, которая игнорирует
выставленный рядом фильтр, читается как сломанная.

Колонки сделки идут после колонок разговора и остаются пустыми, если сделка не
нашлась: пустая ячейка честнее, чем выброшенная строка, — маркетолог видит, у
какой доли разборов вообще нет привязки к CRM.

Потолок строк — MAX_ROWS. Не защита от объёма (разборов сотни), а от случайно
снятого периода на выгрузке за всю историю через год.
"""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_ROWS = 5000

# Подписи видов субъекта для листа. Те же пять видов, что в `call_qa.config`;
# англоязычный код в колонке «Вид» маркетологу ничего не говорит.
KIND_LABELS = {
    'call': 'Звонок', 'imported_call': 'Звонок из АТС', 'wz_episode': 'Чат Wazzup',
    'c2d_snapshot': 'Заявка Chat2Desk', 'ca_episode': 'Чат ChatApp',
}

COLUMNS = (
    ("id", "№ разбора", 10),
    ("subject", "Вид", 14),
    ("subject_datetime", "Разговор", 14),
    ("direction", "Направление", 20),
    ("operator", "Сотрудник", 26),
    ("ai", "Балл ИИ", 9),
    ("human", "Оценка человека", 12),
    ("deal_id", "№ сделки", 12),
    ("channel", "Канал", 14),
    ("campaign", "Кампания", 24),
    ("lead_type", "Тип лида", 10),
    ("park", "Таксопарк", 16),
    ("stage", "Этап сделки", 26),
    ("stage_at_call", "Этап на момент разговора", 26),
    ("reason", "Причина отказа", 32),
    ("responsible", "Ответственный в CRM", 26),
)

_HEADER_FILL = PatternFill("solid", fgColor="F1F5F9")
_HEADER_FONT = Font(bold=True, color="0F172A")


def _flat(item):
    deal = item.get("deal") or {}
    return {
        "id": item.get("id"),
        "subject": KIND_LABELS.get(item.get("subject"), item.get("subject") or ""),
        "subject_datetime": item.get("subject_datetime") or "",
        "direction": item.get("direction") or "",
        "operator": item.get("operator") or "",
        "ai": item.get("ai"),
        "human": item.get("human"),
        "deal_id": deal.get("id") or "",
        "channel": deal.get("channel_title") or "",
        "campaign": deal.get("campaign") or "",
        "lead_type": deal.get("lead_type") or "",
        "park": deal.get("park_title") or "",
        "stage": deal.get("stage") or "",
        "stage_at_call": deal.get("stage_at_call") or "",
        "reason": deal.get("reason") or "",
        "responsible": deal.get("responsible") or "",
    }


def build_xlsx(items, *, title="Разборы"):
    """→ bytes. Числа кладутся числами, а не текстом: иначе Excel зелёным
    уголком помечает каждую ячейку «число сохранено как текст»."""
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.freeze_panes = "A2"
    for column, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_index, item in enumerate(items[:MAX_ROWS], start=2):
        flat = _flat(item)
        for column, (key, _, _) in enumerate(COLUMNS, start=1):
            value = flat.get(key)
            if key == "deal_id" and value:
                try:
                    value = int(value)
                except (TypeError, ValueError):
                    pass
            sheet.cell(row=row_index, column=column, value=value)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(2, len(items) + 1)}"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_csv(items):
    """→ bytes в UTF-8 с BOM: без метки Excel открывает кириллицу кракозябрами."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';', lineterminator='\r\n')
    writer.writerow([header for _, header, _ in COLUMNS])
    for item in items[:MAX_ROWS]:
        flat = _flat(item)
        writer.writerow(['' if flat.get(key) is None else flat.get(key) for key, _, _ in COLUMNS])
    return ('﻿' + buffer.getvalue()).encode('utf-8')


def file_name(extension, department=None):
    stamp = datetime.now().strftime('%Y-%m-%d_%H-%M')
    tail = f"_{department}" if department else ""
    return f"Разборы{tail}_{stamp}.{extension}"
