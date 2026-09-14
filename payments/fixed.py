"""Календарь фиксированных платежей: периодичность, автосоздание заявок, импорт из Excel.

Дополнение Дмитриевой, п. 11: перечень регулярных расходов (аренда, телеметрия…)
загружается файлом или заводится руками; система сама создаёт по нему заявки
в начале периода и ставит их «на согласование» с ответственным из настроек;
связь шаблон → заявка сохраняется (payment_requests.fixed_template_id), и такие
заявки участвуют в контроле просрочек.

Что здесь решено за постановку (она молчит):

* **Когда именно создавать.** «В начале соответствующего периода (например,
  месяца)» — для месячных, квартальных и годовых платежей заявка создаётся
  первого числа месяца, на который приходится срок оплаты. Для произвольного
  интервала — за `lead_days` до срока (по умолчанию неделя).
* **С какого шага начинается автозаявка.** Шаг 1 (реестр поставщиков, КП) для
  регулярного платежа смысла не имеет — поставщик выбран давно; заявка
  создаётся с закрытым системой шагом 1 и сразу ждёт руководителя (шаг 2), то
  есть стоит «на согласовании», как просит ТЗ. Дальше — стандартный маршрут.
* **Двойное создание исключено** сдвигом `next_due_on` вперёд сразу после
  создания и пометкой `last_generated_on`.

Чистые функции (даты, разбор файла) отделены от тех, что ходят в базу
(`generate`, `resolve_import`, `commit_import`) — первые проверяются тестами
без Postgres.
"""

import calendar
import io
import logging
import re
from datetime import date, datetime, timedelta
from decimal import Decimal

from . import notify, queries, workflow

PERIOD_LABELS = {
    'monthly': 'Ежемесячно',
    'quarterly': 'Ежеквартально',
    'yearly': 'Ежегодно',
    'custom': 'Свой интервал',
}


def add_months(value, months):
    """Дата через `months` месяцев с тем же числом; 31-е в коротком месяце — последний день."""
    month_index = value.month - 1 + months
    year = value.year + month_index // 12
    month = month_index % 12 + 1
    day = min(value.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def next_due(due_on, periodicity, interval_days=None):
    if periodicity == 'monthly':
        return add_months(due_on, 1)
    if periodicity == 'quarterly':
        return add_months(due_on, 3)
    if periodicity == 'yearly':
        return add_months(due_on, 12)
    days = int(interval_days or 0)
    if days <= 0:
        raise ValueError('Для своего интервала нужно число дней')
    return due_on + timedelta(days=days)


def generate_on(due_on, periodicity, lead_days=7):
    """Когда создавать заявку под этот срок."""
    if periodicity in ('monthly', 'quarterly', 'yearly'):
        return due_on.replace(day=1)
    return due_on - timedelta(days=max(0, int(lead_days or 0)))


def is_due(template, today):
    """Пора ли создавать заявку по шаблону сегодня."""
    if not template.get('is_active'):
        return False
    due_on = workflow._as_date(template.get('next_due_on'))
    if not due_on:
        return False
    start = generate_on(due_on, template.get('periodicity'), template.get('lead_days'))
    if start > today:
        return False
    last = workflow._as_date(template.get('last_generated_on'))
    # Уже создавали под ЭТОТ срок — next_due_on тогда сдвинут, и сюда не попадём;
    # проверка страхует ручной сдвиг срока назад в тот же день.
    return not (last and last >= start and template.get('open_request_id'))


def generate(cursor, *, today=None, actor=None, base_url=None):
    """Создаёт заявки по всем шаблонам, у которых наступило начало периода.

    Возвращает (созданные заявки, письма для Telegram). Письма отправляет
    вызывающий ПОСЛЕ коммита — иначе человек получит ссылку на заявку, которой
    ещё нет.
    """
    today = today or queries.today_almaty()
    actor = actor or {'id': None, 'name': 'Календарь платежей'}
    created, outbox = [], []
    members = queries.role_member_ids(cursor)
    if not members.get('founder') or not members.get('accounting'):
        # Без Учредителя и Бухгалтерии заявке некуда идти после шага 2 —
        # календарь молчит и говорит об этом в лог, а не плодит зависшие заявки.
        logging.warning('Оплата счетов: календарь не создаёт заявки — не назначены участники ролей')
        return created, outbox
    for template in queries.list_templates(cursor, include_inactive=False):
        if not is_due(template, today):
            continue
        locked = queries.read_template(cursor, template['id'], lock=True)
        if not locked or not is_due(locked, today):
            continue
        request_id = create_from_template(cursor, locked, actor=actor, today=today, outbox=outbox,
                                          base_url=base_url)
        created.append(request_id)
    return created, outbox


def create_from_template(cursor, template, *, actor, today, outbox, base_url=None):
    responsible = queries.user_brief(cursor, template['responsible_user_id'])
    if not responsible:
        raise ValueError('У платежа «%s» не найден ответственный' % template['name'])
    manager = queries.resolve_manager(cursor, responsible['id'])
    department = queries.department_brief(cursor, responsible.get('department_id'))
    due_on = workflow._as_date(template['next_due_on'])
    fields = {
        'department_id': (department or {}).get('id'),
        'department_name': (department or {}).get('name'),
        'project_id': template.get('project_id'),
        'branch': template.get('branch'),
        'expense_name': template['name'],
        'category_id': template.get('category_id'),
        'subcategory_id': template.get('subcategory_id'),
        'counterparty_id': template.get('counterparty_id'),
        'legal_entity_id': template.get('legal_entity_id'),
        'payment_source': template.get('payment_source'),
        'payment_type': 'fixed',
        'payment_period': _period_text(due_on, template.get('periodicity')),
        'notes': template.get('note'),
        'due_on': due_on,
        'fixed_template_id': template['id'],
    }
    items = [{'name': template['name'], 'quantity': 1, 'unit': None, 'unit_price': template['amount']}]
    route = workflow.build_route(initiator=responsible, manager=manager)
    comment = 'Создано по календарю фиксированных платежей: «%s», срок %s' % (
        template['name'], workflow.fmt_date(due_on))
    request_id = queries.create_request(
        cursor, fields=fields, items=items, actor=actor, initiator=responsible, manager=manager,
        route=route, first_step_done=True, system_comment=comment)
    queries.log_event(cursor, request_id, 'generated', actor, step_no=1,
                      payload={'template_id': template['id'], 'due_on': due_on})
    queries.advance_template(cursor, template['id'],
                             next_due_on=next_due(due_on, template['periodicity'], template.get('interval_days')),
                             generated_on=today)
    request = queries.read_request(cursor, request_id)
    link = notify.request_link(base_url, request_id)
    for recipient in queries.telegram_recipients(cursor, user_ids=[responsible['id']]):
        outbox.append((recipient['chat_id'], notify.generated_message(request, template['name']),
                       notify.reply_markup(link)))
    step_no = int(request['current_step'] or 0)
    if step_no > 1:
        current = next((s for s in queries.list_steps(cursor, request_id) if s['step_no'] == step_no), None)
        if current:
            recipients = queries.telegram_recipients(
                cursor, user_ids=[current['assignee_id']] if current.get('assignee_id') else (),
                role_code=None if current.get('assignee_id') else current['role_code'],
                exclude=[responsible['id']])
            for recipient in recipients:
                outbox.append((recipient['chat_id'], notify.step_message(request, step_no),
                               notify.reply_markup(link)))
    return request_id


def _period_text(due_on, periodicity):
    months = ('январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль', 'август',
              'сентябрь', 'октябрь', 'ноябрь', 'декабрь')
    if not due_on:
        return None
    if periodicity == 'monthly':
        return '%s %s' % (months[due_on.month - 1], due_on.year)
    if periodicity == 'quarterly':
        return '%s квартал %s' % ((due_on.month - 1) // 3 + 1, due_on.year)
    if periodicity == 'yearly':
        return '%s год' % due_on.year
    return 'до %s' % workflow.fmt_date(due_on)


# ─────────────────────────────────────────────────────────────────────────────
# Импорт календаря из Excel
# ─────────────────────────────────────────────────────────────────────────────

# Заголовки колонок по ТЗ п. 11.1 и их бытовые варианты. Сравнение — по
# нормализованной строке (нижний регистр, без пунктуации).
_HEADERS = {
    'name': ('наименование', 'название', 'платеж', 'платёж', 'расход', 'наименование расхода'),
    'amount': ('сумма', 'сумма расхода', 'стоимость'),
    'periodicity': ('периодичность', 'период', 'частота'),
    'due': ('дата оплаты', 'срок', 'крайний срок', 'дата', 'день оплаты', 'срок оплаты'),
    'project': ('проект', 'филиал', 'проект филиал', 'проект / филиал', 'таксопарк', 'город'),
    'responsible': ('ответственный', 'инициатор', 'фио', 'сотрудник'),
    'category': ('категория', 'категория расхода'),
    'subcategory': ('подкатегория', 'подкатегория расхода'),
    'counterparty': ('контрагент', 'поставщик'),
    'source': ('источник оплаты', 'источник'),
    'note': ('примечание', 'примечания', 'комментарий'),
}


def _norm(text):
    return re.sub(r'[^a-zа-яё0-9 ]+', ' ', str(text or '').lower().replace('ё', 'е')).strip()


def detect_columns(header_row):
    mapping = {}
    for index, cell in enumerate(header_row):
        key = _norm(cell)
        if not key:
            continue
        for field, variants in _HEADERS.items():
            if field in mapping:
                continue
            if key in {_norm(v) for v in variants}:
                mapping[field] = index
                break
    return mapping


def parse_periodicity(value):
    """Строка из файла → (periodicity, interval_days) либо (None, None)."""
    text = _norm(value)
    if not text:
        return None, None
    if any(word in text for word in ('месяц', 'ежемес', 'monthly')):
        return 'monthly', None
    if 'квартал' in text or 'quarter' in text:
        return 'quarterly', None
    if any(word in text for word in ('год', 'ежегод', 'year', 'annual')):
        return 'yearly', None
    match = re.search(r'(\d+)\s*(д|дн|день|дня|дней|day)', text)
    if match:
        return 'custom', int(match.group(1))
    if 'недел' in text or 'week' in text:
        return 'custom', 7
    return None, None


def parse_due(value, today):
    """Дата оплаты из ячейки: дата, «dd.mm.yyyy», либо просто число месяца («3»)."""
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.strip().isdigit()):
        day = int(float(value))
        if 1 <= day <= 31:
            candidate = date(today.year, today.month, min(day, calendar.monthrange(today.year, today.month)[1]))
            if candidate < today:
                candidate = add_months(candidate, 1)
                candidate = candidate.replace(day=min(day, calendar.monthrange(candidate.year, candidate.month)[1]))
            return candidate
        return None
    text = str(value).strip()
    for pattern in ('%d.%m.%Y', '%d.%m.%y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    match = re.match(r'^(\d{1,2})[ -]?(?:го|е|ое)?(?:\s+числа)?$', text.lower())
    if match:
        return parse_due(int(match.group(1)), today)
    return None


def parse_source(value):
    text = _norm(value)
    if not text:
        return None
    if 'тоо' in text or 'безнал' in text or 'счет' in text:
        return 'too'
    if 'кошел' in text or 'wallet' in text:
        return 'wallet'
    if 'нал' in text or 'cash' in text:
        return 'cash'
    return None


def parse_workbook(data, today=None):
    """Байты xlsx → строки предпросмотра с ошибками по каждой."""
    from openpyxl import load_workbook

    today = today or queries.today_almaty()
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    header_index = None
    mapping = {}
    for index, row in enumerate(rows[:20]):
        mapping = detect_columns(row or ())
        if 'name' in mapping and 'amount' in mapping:
            header_index = index
            break
    if header_index is None:
        raise ValueError('Не нашёл строку заголовков: нужны хотя бы колонки «Наименование» и «Сумма»')

    def cell(row, field):
        index = mapping.get(field)
        if index is None or index >= len(row):
            return None
        return row[index]

    result = []
    for row in rows[header_index + 1:]:
        if not row or all(value in (None, '') for value in row):
            continue
        name = str(cell(row, 'name') or '').strip()
        errors = []
        if not name:
            errors.append('нет наименования')
        amount = workflow.to_decimal(cell(row, 'amount'))
        if amount <= 0:
            errors.append('сумма не разобрана')
        periodicity, interval_days = parse_periodicity(cell(row, 'periodicity'))
        if not periodicity:
            periodicity, interval_days = 'monthly', None
        due_on = parse_due(cell(row, 'due'), today)
        if not due_on:
            errors.append('дата оплаты не разобрана')
        result.append({
            'name': name,
            'amount': amount,
            'periodicity': periodicity,
            'interval_days': interval_days,
            'next_due_on': due_on,
            'project': str(cell(row, 'project') or '').strip() or None,
            'responsible': str(cell(row, 'responsible') or '').strip() or None,
            'category': str(cell(row, 'category') or '').strip() or None,
            'subcategory': str(cell(row, 'subcategory') or '').strip() or None,
            'counterparty': str(cell(row, 'counterparty') or '').strip() or None,
            'payment_source': parse_source(cell(row, 'source')),
            'note': str(cell(row, 'note') or '').strip() or None,
            'errors': errors,
        })
    return result


def resolve_import(cursor, rows, fallback_responsible_id):
    """Сопоставляет строки файла со справочниками и людьми; ничего не пишет.

    Ответственный ищется по ФИО (без учёта регистра, по вхождению); не нашёлся —
    подставляется тот, кто загружает файл, с пометкой в строке.
    """
    users = queries.list_users(cursor)
    by_name = {_norm(u['name']): u for u in users}
    projects = {_norm(p['name']): p for p in queries.list_projects(cursor, include_inactive=True)}
    categories = queries.list_categories(cursor, include_inactive=True)
    roots = {_norm(c['name']): c for c in categories if not c['parent_id']}
    children = {}
    for item in categories:
        if item['parent_id']:
            children.setdefault(item['parent_id'], {})[_norm(item['name'])] = item
    counterparties = {_norm(c['name']): c for c in queries.list_counterparties(cursor, include_inactive=True)}
    fallback = next((u for u in users if u['id'] == int(fallback_responsible_id or 0)), None)

    for row in rows:
        notes = []
        responsible = None
        if row.get('responsible'):
            key = _norm(row['responsible'])
            responsible = by_name.get(key) or next(
                (u for norm_name, u in by_name.items() if key and (key in norm_name or norm_name in key)), None)
            if not responsible:
                notes.append('ответственный «%s» не найден — подставлен загружающий' % row['responsible'])
        if not responsible:
            responsible = fallback
        row['responsible_user_id'] = (responsible or {}).get('id')
        row['responsible_name'] = (responsible or {}).get('name')
        if not row['responsible_user_id']:
            row['errors'].append('не определён ответственный')

        project = projects.get(_norm(row.get('project'))) if row.get('project') else None
        row['project_id'] = (project or {}).get('id')
        row['project_new'] = bool(row.get('project') and not project)

        category = roots.get(_norm(row.get('category'))) if row.get('category') else None
        row['category_id'] = (category or {}).get('id')
        row['category_new'] = bool(row.get('category') and not category)
        sub = None
        if row.get('subcategory') and category:
            sub = children.get(category['id'], {}).get(_norm(row['subcategory']))
        row['subcategory_id'] = (sub or {}).get('id')
        row['subcategory_new'] = bool(row.get('subcategory') and not sub)

        cp = counterparties.get(_norm(row.get('counterparty'))) if row.get('counterparty') else None
        row['counterparty_id'] = (cp or {}).get('id')
        row['counterparty_new'] = bool(row.get('counterparty') and not cp)
        row['notes'] = notes
    return rows


def commit_import(cursor, rows, actor_id):
    """Заводит недостающие справочники и шаблоны. Строки с ошибками пропускаются."""
    created = []
    for row in rows:
        if row.get('errors'):
            continue
        project_id = row.get('project_id')
        if not project_id and row.get('project'):
            project_id, _ = queries.find_or_create_project(cursor, row['project'], actor_id)
        category_id = row.get('category_id')
        if not category_id and row.get('category'):
            category_id, _ = queries.find_or_create_category(cursor, row['category'], None, actor_id)
        subcategory_id = row.get('subcategory_id')
        if not subcategory_id and row.get('subcategory') and category_id:
            subcategory_id, _ = queries.find_or_create_category(cursor, row['subcategory'], category_id, actor_id)
        counterparty_id = row.get('counterparty_id')
        if not counterparty_id and row.get('counterparty'):
            counterparty_id, _ = queries.find_or_create_counterparty(cursor, row['counterparty'], actor_id)
        template_id = queries.upsert_template(cursor, fields={
            'name': row['name'],
            'amount': workflow.to_decimal(row['amount']),
            'periodicity': row['periodicity'],
            'interval_days': row.get('interval_days'),
            'next_due_on': workflow._as_date(row['next_due_on']),
            'lead_days': 7,
            'project_id': project_id,
            'branch': None,
            'responsible_user_id': int(row['responsible_user_id']),
            'category_id': category_id,
            'subcategory_id': subcategory_id,
            'counterparty_id': counterparty_id,
            'legal_entity_id': None,
            'payment_source': row.get('payment_source'),
            'note': row.get('note'),
            'is_active': True,
        }, actor_id=actor_id)
        created.append(template_id)
    return created


def template_sample_rows():
    """Строки-образцы для скачивания шаблона файла."""
    return [
        ('Наименование', 'Сумма', 'Периодичность', 'Дата оплаты', 'Проект / филиал', 'Ответственный',
         'Категория', 'Подкатегория', 'Контрагент', 'Источник оплаты', 'Примечание'),
        ('Аренда офиса Алматы', 450000, 'ежемесячно', '5', 'Таксопарк Алматы', 'Ядигаров Руслан',
         'Аренда', 'Офис', 'ТОО «Пример»', 'ТОО', 'по договору №12'),
    ]


def log_generation(created):
    if created:
        logging.info('Оплата счетов: по календарю фиксированных платежей создано заявок — %s', len(created))
