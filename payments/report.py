"""Выгрузка реестра заявок в xlsx (дополнение Дмитриевой, п. 8: «все поля, включая скрытые»).

Соглашения те же, что у выгрузок «Посылок» и «Касаний»:

* первым идёт лист «Контекст» — когда собрано, что отобрано, сколько строк;
* даты — датами с форматом, суммы — числами: иначе не отсортировать и не
  построить сводную;
* БИН, номер карты, номер счёта и договора — ТЕКСТОМ (число теряет ведущие нули
  и уезжает в экспоненту), а зелёный уголок «Число сохранено как текст» гасится
  тегом <ignoredErrors>; сама функция приходит аргументом, потому что живёт в
  монолите;
* подписи статусов, этапов и источников — те же слова, что на экране
  (paymentsMeta.js); второй словарь тех же кодов неизбежен, его сторожит тест.
"""

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import workflow

ALMATY = ZoneInfo('Asia/Almaty')

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')

DAY_FORMAT = 'DD.MM.YYYY'
DATETIME_FORMAT = 'DD.MM.YYYY HH:MM'
MONEY_FORMAT = '#,##0.00'

STATE_LABELS = {
    'active': 'В работе',
    'blocked': 'Ожидает договор',
    'overdue': 'Просрочена',
    'done': 'Закрыта',
    'rejected': 'Отклонена',
    'cancelled': 'Отменена',
}
SOURCE_LABELS = {'too': 'ТОО', 'wallet': 'Кошелёк', 'cash': 'Наличные'}
TYPE_LABELS = {'fixed': 'Фиксированный', 'monthly': 'Ежемесячный', 'one_time': 'Разовый'}

# (заголовок, ключ, вид) — вид: 'text' | 'money' | 'date' | 'datetime' | 'plain'
COLUMNS = (
    ('№', 'id', 'plain'),
    ('Дата заявки', 'created_at', 'datetime'),
    ('Состояние', 'state_label', 'plain'),
    ('Этап', 'step_label', 'plain'),
    ('Текущий ответственный', 'current_assignee_name', 'plain'),
    ('Инициатор', 'initiator_name', 'plain'),
    ('Руководитель', 'manager_name', 'plain'),
    ('Отдел', 'department_name', 'plain'),
    ('Проект', 'project_name', 'plain'),
    ('Таксопарк / филиал / регион', 'branch', 'plain'),
    ('Наименование расхода', 'expense_name', 'plain'),
    ('Позиции', 'items_text', 'plain'),
    ('Сумма расхода', 'amount', 'money'),
    ('Сумма возврата', 'refund_amount', 'money'),
    ('Дата возврата', 'refund_on', 'date'),
    ('Итоговая сумма', 'total_amount', 'money'),
    ('Категория', 'category_name', 'plain'),
    ('Подкатегория', 'subcategory_name', 'plain'),
    ('Контрагент', 'counterparty_name', 'plain'),
    ('БИН контрагента', 'counterparty_bin', 'text'),
    ('Юр. лицо (плательщик)', 'legal_entity_name', 'plain'),
    ('Договор', 'contract_number', 'text'),
    ('Источник оплаты', 'source_label', 'plain'),
    ('Тип оплаты', 'type_label', 'plain'),
    ('Номер карты', 'card_number', 'text'),
    ('Период оплаты', 'payment_period', 'plain'),
    ('Срок оплаты', 'due_on', 'date'),
    ('Дата платежа', 'paid_on', 'date'),
    ('Оплачено', 'paid_amount', 'money'),
    ('Номер счёта', 'invoice_number', 'text'),
    ('Дата счёта', 'invoice_date', 'date'),
    ('Реквизиты для счёта', 'invoice_requisites', 'plain'),
    ('Описание счёта', 'invoice_description', 'plain'),
    ('Проверка бухгалтерии', 'previous_payment_note', 'plain'),
    ('Согласующий счёта', 'approver_label', 'plain'),
    ('Приказ', 'approval_order_number', 'text'),
    ('Доверенность', 'needs_power_of_attorney_label', 'plain'),
    ('Платёжное поручение', 'needs_payment_order_label', 'plain'),
    ('Примечания', 'notes', 'plain'),
    ('Причина отклонения', 'rejected_reason', 'plain'),
    ('Закрыта', 'closed_at', 'datetime'),
)

WIDTHS = {
    '№': 7, 'Дата заявки': 16, 'Состояние': 16, 'Этап': 30, 'Наименование расхода': 34, 'Позиции': 40,
    'Сумма расхода': 15, 'Итоговая сумма': 15, 'Контрагент': 26, 'Описание счёта': 40,
    'Реквизиты для счёта': 30, 'Проверка бухгалтерии': 30, 'Примечания': 30,
}


def _clean(value):
    if value is None:
        return None
    text = str(value)
    text = ILLEGAL_CHARACTERS_RE.sub('', text)
    return text


def _text_cell(ws, value):
    cell = WriteOnlyCell(ws, value=_clean(value) if value not in (None, '') else None)
    cell.number_format = '@'
    cell.alignment = WRAP_TOP
    return cell


def _plain_cell(ws, value):
    text = _clean(value) if value not in (None, '') else None
    # Текст, начинающийся с «=», Excel считает формулой — оставляем текстом.
    cell = WriteOnlyCell(ws, value=text)
    if text and text.startswith('='):
        cell.number_format = '@'
        cell.data_type = 's'
    cell.alignment = WRAP_TOP
    return cell


def _money_cell(ws, value):
    cell = WriteOnlyCell(ws, value=float(value) if value not in (None, '') else None)
    cell.number_format = MONEY_FORMAT
    return cell


def _date_cell(ws, value, fmt):
    parsed = value
    if isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value)
        except ValueError:
            parsed = None
    cell = WriteOnlyCell(ws, value=parsed)
    cell.number_format = fmt
    return cell


def enrich(request, items):
    """Поля, которых нет в строке базы: подписи и производные суммы."""
    row = dict(request)
    row['state_label'] = STATE_LABELS.get(request.get('state'), request.get('state'))
    if request.get('status') == 'active':
        row['step_label'] = 'Шаг %s из %s · %s' % (request.get('current_step'), workflow.LAST_STEP,
                                                   workflow.step_title(request.get('current_step')))
    else:
        row['step_label'] = ''
    row['source_label'] = SOURCE_LABELS.get(request.get('payment_source'), '')
    row['type_label'] = TYPE_LABELS.get(request.get('payment_type'), '')
    amount = workflow.to_decimal(request.get('amount'))
    refund = workflow.to_decimal(request.get('refund_amount'))
    row['total_amount'] = amount - refund
    row['items_text'] = '\n'.join(
        '%s — %s%s × %s' % (item.get('name'), _qty(item.get('quantity')),
                            (' ' + item['unit']) if item.get('unit') else '',
                            workflow.fmt_money(item.get('unit_price')).replace(' ₸', ''))
        for item in (items or []))
    basis = request.get('route_basis') or {}
    if basis.get('approver_name'):
        row['approver_label'] = '%s (по Приказу №%s вместо Учредителя)' % (basis['approver_name'],
                                                                           basis.get('order_number') or '—')
    else:
        row['approver_label'] = 'Учредитель'
    row['needs_power_of_attorney_label'] = 'Да' if request.get('needs_power_of_attorney') else ''
    row['needs_payment_order_label'] = 'Да' if request.get('needs_payment_order') else ''
    return row


def _qty(value):
    number = workflow.to_decimal(value, default=None)
    if number is None:
        return ''
    if number == number.to_integral():
        return str(int(number))
    return str(number.normalize())


def build_workbook(requests, items_by_request, *, generated_at=None, filters_text='', total=None,
                   truncated=False, text_warning_patch=None):
    generated_at = generated_at or datetime.now(ALMATY)
    workbook = Workbook(write_only=True)

    context = workbook.create_sheet('Контекст')
    context.column_dimensions['A'].width = 28
    context.column_dimensions['B'].width = 80
    title = WriteOnlyCell(context, value='Реестр заявок на оплату')
    title.font = TITLE_FONT
    context.append([title])
    context.append(['Собрано', generated_at.strftime('%d.%m.%Y %H:%M')])
    context.append(['Отбор', filters_text or 'все заявки'])
    context.append(['Строк в файле', len(requests)])
    if total is not None and total != len(requests):
        context.append(['Всего по отбору', total])
    if truncated:
        context.append(['Внимание', 'Файл обрезан по потолку строк — сузьте отбор'])
    context.append([])
    context.append(['Как читать', 'Сумма расхода — из позиций заявки; итоговая сумма = сумма − возврат. '
                                  'БИН, номер карты и номера документов лежат текстом.'])

    sheet = workbook.create_sheet('Заявки')
    sheet.freeze_panes = 'C2'
    for index, (header, _key, _kind) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(header, 18)
    header_cells = []
    for header, _key, _kind in COLUMNS:
        cell = WriteOnlyCell(sheet, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        header_cells.append(cell)
    sheet.append(header_cells)

    for request in requests:
        row = enrich(request, items_by_request.get(request['id']) or [])
        cells = []
        for _header, key, kind in COLUMNS:
            value = row.get(key)
            if kind == 'text':
                cells.append(_text_cell(sheet, value))
            elif kind == 'money':
                cells.append(_money_cell(sheet, value))
            elif kind == 'date':
                cells.append(_date_cell(sheet, value, DAY_FORMAT))
            elif kind == 'datetime':
                cells.append(_date_cell(sheet, value, DATETIME_FORMAT))
            else:
                cells.append(_plain_cell(sheet, value))
        sheet.append(cells)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    if text_warning_patch and requests:
        text_columns = [get_column_letter(i) for i, (_h, _k, kind) in enumerate(COLUMNS, start=1)
                        if kind == 'text']
        last_row = len(requests) + 1
        sqref = ' '.join('%s2:%s%s' % (col, col, last_row) for col in text_columns)
        # Лист «Заявки» — второй по порядку создания.
        stream = text_warning_patch(stream, sqref, sheet_path='xl/worksheets/sheet2.xml')
    return stream


def file_name(generated_at=None):
    generated_at = generated_at or datetime.now(ALMATY)
    return 'Заявки на оплату %s.xlsx' % generated_at.strftime('%Y-%m-%d')
