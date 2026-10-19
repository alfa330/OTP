# -*- coding: utf-8 -*-
"""«Активность операторов» (задача #392): смены, статусы, звонки, журнал, выгрузка, права."""

import unittest
from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Flask, jsonify
from openpyxl import load_workbook

from op_wallboard import activity as A, activity_export as E, activity_routes as R


def at(text):
    return datetime.strptime(text, '%Y-%m-%d %H:%M:%S')


def ev(text, key):
    return (at(text), key)


def touch(started, call_type='Исходящий', talk=0, ext='6650', result='', queue='',
          measured=None, dial=None, answered=''):
    return {'started_at': at(started), 'answered_at': answered, 'ext': ext, 'call_type': call_type,
            'result': result, 'talk_seconds': talk, 'dial_seconds': dial if dial is not None else talk + 15,
            'queue': queue, 'talk_measured_seconds': measured}


NOW = at('2026-10-19 12:00:00')


class ShiftSplitTests(unittest.TestCase):
    def test_day_shift_ends_on_logout(self):
        shifts = A.split_shifts([ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 13:00:00', 'перерыв'),
                                 ev('2026-10-05 14:00:00', 'исход'), ev('2026-10-05 18:00:00', 'выключен')])
        self.assertEqual(len(shifts), 1)
        self.assertEqual(shifts[0][0][0], at('2026-10-05 09:00:00'))
        self.assertEqual(shifts[0][-1][1], 'выключен')
        self.assertTrue(A.shift_finished(shifts[0], NOW))

    def test_night_shift_is_one_shift_of_its_start_day(self):
        shifts = A.split_shifts([ev('2026-10-05 20:00:00', 'готов'), ev('2026-10-06 02:00:00', 'занят'),
                                 ev('2026-10-06 02:05:00', 'готов'), ev('2026-10-06 08:00:00', 'выключен')])
        self.assertEqual(len(shifts), 1)
        self.assertEqual(shifts[0][0][0].date(), date(2026, 10, 5))

    def test_quiet_night_in_active_does_not_split_the_shift(self):
        """Основа ночью ждёт звонка 4,5 ч — телефон всё это время молчит, смена одна."""
        shifts = A.split_shifts([ev('2026-10-05 20:00:00', 'готов'), ev('2026-10-05 23:00:00', 'занят'),
                                 ev('2026-10-05 23:03:00', 'готов'), ev('2026-10-06 03:40:00', 'занят'),
                                 ev('2026-10-06 03:45:00', 'готов'), ev('2026-10-06 08:00:00', 'выключен')])
        self.assertEqual(len(shifts), 1)

    def test_lost_logout_does_not_glue_two_days(self):
        """«Выключен» потерялся вечером — утренняя смена следующего дня остаётся своей. Повтор
        «Исхода» при утреннем входе сервер не хранит (тот же статус), поэтому утро начинается
        с первого набора."""
        shifts = A.split_shifts([ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 12:00:00', 'соединение'),
                                 ev('2026-10-05 12:00:40', 'исход'), ev('2026-10-05 15:00:00', 'соединение'),
                                 ev('2026-10-05 15:00:30', 'исход'), ev('2026-10-05 17:58:00', 'соединение'),
                                 ev('2026-10-05 17:59:00', 'исход'),
                                 ev('2026-10-06 09:02:00', 'соединение'), ev('2026-10-06 09:03:00', 'исход'),
                                 ev('2026-10-06 12:00:00', 'перерыв'), ev('2026-10-06 13:00:00', 'исход'),
                                 ev('2026-10-06 18:00:00', 'выключен')])
        self.assertEqual(len(shifts), 2)
        self.assertEqual(shifts[0][-1][0], at('2026-10-05 17:59:00'))
        self.assertNotEqual(A.classify(shifts[0][-1][1])[0], A.KIND_LOGOUT)
        self.assertEqual(shifts[1][0][0], at('2026-10-06 09:02:00'))

    def test_short_rest_after_lost_logout_still_splits_by_length(self):
        """Вечерняя до 02:00 без выхода и утренняя с 09:10: пауза 7 ч, но вместе вышло бы 25 ч."""
        shifts = A.split_shifts([ev('2026-10-05 17:00:00', 'готов'), ev('2026-10-05 20:00:00', 'занят'),
                                 ev('2026-10-05 20:03:00', 'готов'), ev('2026-10-05 23:00:00', 'занят'),
                                 ev('2026-10-05 23:02:00', 'готов'), ev('2026-10-06 01:58:00', 'занят'),
                                 ev('2026-10-06 02:00:00', 'готов'),
                                 ev('2026-10-06 09:10:00', 'занят'), ev('2026-10-06 09:15:00', 'готов'),
                                 ev('2026-10-06 12:00:00', 'занят'), ev('2026-10-06 12:05:00', 'готов'),
                                 ev('2026-10-06 15:00:00', 'перерыв'), ev('2026-10-06 15:30:00', 'готов'),
                                 ev('2026-10-06 18:00:00', 'выключен')])
        self.assertEqual(len(shifts), 2)
        self.assertEqual(shifts[0][0][0], at('2026-10-05 17:00:00'))
        self.assertEqual(shifts[1][0][0], at('2026-10-06 09:10:00'))

    def test_pause_after_break_splits_at_four_hours(self):
        shifts = A.split_shifts([ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 12:00:00', 'перерыв'),
                                 ev('2026-10-05 16:30:00', 'исход'), ev('2026-10-05 18:00:00', 'выключен')])
        self.assertEqual(len(shifts), 2)

    def test_lunch_logout_keeps_one_shift(self):
        shifts = A.split_shifts([ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 13:00:00', 'выключен'),
                                 ev('2026-10-05 14:00:00', 'исход'), ev('2026-10-05 18:00:00', 'выключен')])
        self.assertEqual(len(shifts), 1)

    def test_leading_logout_and_offline_only_blocks_are_not_shifts(self):
        shifts = A.split_shifts([ev('2026-10-05 08:59:50', 'выключен'), ev('2026-10-05 09:00:00', 'исход'),
                                 ev('2026-10-05 11:00:00', 'перерыв'), ev('2026-10-05 11:15:00', 'исход'),
                                 ev('2026-10-05 14:00:00', 'соединение'), ev('2026-10-05 14:01:00', 'исход'),
                                 ev('2026-10-05 18:00:00', 'выключен'),
                                 ev('2026-10-06 10:00:00', 'офлайн'), ev('2026-10-06 10:00:05', 'выключен')])
        self.assertEqual(len(shifts), 1)
        self.assertEqual(shifts[0][0], ev('2026-10-05 09:00:00', 'исход'))

    def test_same_second_keeps_the_last_event_and_repeats_collapse(self):
        events = A.canonical_events([ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 09:05:00', 'соединение'),
                                     ev('2026-10-05 09:05:00', 'исход'), ev('2026-10-05 09:06:00', 'исход'),
                                     ev('2026-10-05 09:07:00', 'перерыв')])
        self.assertEqual(events, [ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 09:07:00', 'перерыв')])

    def test_single_event_is_not_a_shift(self):
        self.assertEqual(A.split_shifts([ev('2026-10-05 09:00:00', 'исход')]), [])

    def test_running_shift_is_not_finished(self):
        shift = [ev('2026-10-19 09:00:00', 'готов'), ev('2026-10-19 11:00:00', 'готов')]
        self.assertFalse(A.shift_finished(A.canonical_events(shift), NOW))
        idle = [ev('2026-10-19 06:00:00', 'исход'), ev('2026-10-19 07:00:00', 'офлайн')]
        self.assertTrue(A.shift_finished(idle, NOW))


class StatusColumnsTests(unittest.TestCase):
    SHIFT = [ev('2026-10-05 09:00:00', 'исход'),        # ожидание 5 мин
             ev('2026-10-05 09:05:00', 'соединение'),   # набор 30 с
             ev('2026-10-05 09:05:30', 'занят'),        # разговор 2 мин
             ev('2026-10-05 09:07:30', 'перезвон'),     # старый ключ «Исхода» — ожидание 2,5 мин
             ev('2026-10-05 09:10:00', 'перерыв'),      # 15 мин
             ev('2026-10-05 09:25:00', 'tech.break'),   # 10 мин
             ev('2026-10-05 09:35:00', 'тренинг'),      # 30 мин
             ev('2026-10-05 10:05:00', 'офлайн'),       # 5 мин
             ev('2026-10-05 10:10:00', 'исход'),        # 50 мин
             ev('2026-10-05 11:00:00', 'выключен')]

    def test_columns(self):
        totals = A.status_totals(A.pieces(self.SHIFT, 'yar'))
        self.assertEqual(totals['active_seconds'], 300 + 30 + 120 + 150 + 3000)
        self.assertEqual(totals['waiting_seconds'], 300 + 150 + 3000)
        self.assertEqual(totals['break_seconds'], 900)
        self.assertEqual(totals['tech_seconds'], 600)
        self.assertEqual(totals['training_seconds'], 1800)
        self.assertEqual(totals['worked_seconds'], totals['active_seconds'] + 1800)
        self.assertEqual(totals['offline_seconds'], 300)
        # Ни одна секунда не потерялась и не посчиталась дважды.
        span = (self.SHIFT[-1][0] - self.SHIFT[0][0]).total_seconds()
        self.assertEqual(totals['worked_seconds'] + totals['break_seconds'] + totals['tech_seconds']
                         + totals['offline_seconds'], span)

    def test_journal_labels_and_sums(self):
        rows = A.journal(self.SHIFT, 'yar')
        labels = [(row['label'], row['note']) for row in rows]
        self.assertEqual(labels, [
            ('Ожидание', None), ('Исход', 'набор номера'), ('В разговоре', None), ('Ожидание', None),
            ('Перерыв', None), ('Технический перерыв', None), ('Тренинг', None),
            ('Офлайн', 'простой без звонков'), ('Ожидание', None), ('Офлайн', 'выход из телефона')])
        self.assertEqual(rows[-1]['seconds'], 0)
        self.assertEqual(rows[-1]['start_at'], rows[-1]['end_at'])
        totals = A.status_totals(A.pieces(self.SHIFT, 'yar'))
        self.assertEqual(sum(r['seconds'] for r in rows if r['kind'] in ('talk', 'dial', 'wait')),
                         totals['active_seconds'])
        self.assertEqual(sum(r['seconds'] for r in rows if r['kind'] == 'wait'), totals['waiting_seconds'])

    def test_dial_label_follows_the_working_status(self):
        rows = A.journal([ev('2026-10-05 09:00:00', 'готов'), ev('2026-10-05 09:01:00', 'соединение'),
                          ev('2026-10-05 09:01:20', 'готов')], 'osnova')
        self.assertEqual(rows[1]['label'], 'Активный')
        self.assertEqual(rows[1]['style'], 'active')
        rows = A.journal([ev('2026-10-05 09:00:00', 'автодозвон'), ev('2026-10-05 09:01:00', 'соединение'),
                          ev('2026-10-05 09:01:20', 'автодозвон')], 'potok')
        self.assertEqual(rows[1]['label'], 'Автодозвон')

    def test_adjacent_waits_merge(self):
        rows = A.journal([ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 09:01:00', 'перезвон'),
                          ev('2026-10-05 09:02:00', 'перерыв')], 'yar')
        self.assertEqual([r['label'] for r in rows], ['Ожидание'])
        self.assertEqual(rows[0]['seconds'], 120)

    def test_tech_variants(self):
        for key in ('тех причина', 'tech.break', 'тех_причина', 'Тех. причина'):
            self.assertEqual(A.classify(key)[0], A.KIND_TECH, key)


class CallColumnsTests(unittest.TestCase):
    def test_yar_counts_outgoing_only(self):
        touches = [touch('2026-10-05 09:00:00', talk=0), touch('2026-10-05 09:01:00', talk=9),
                   touch('2026-10-05 09:02:00', talk=10), touch('2026-10-05 09:05:00', talk=125),
                   touch('2026-10-05 09:10:00', 'Входящий', talk=60, queue='3001')]
        cols = A.call_columns(touches, 'yar')
        self.assertEqual(cols['calls'], 4)
        self.assertEqual(cols['connected'], 2)
        self.assertEqual(cols['talk_seconds'], 144)
        self.assertEqual(cols['avg_talk_seconds'], 36)   # 144 // 4 — нули в среднем есть

    def test_osnova_counts_incoming_by_the_board_rules(self):
        touches = [touch('2026-10-05 09:00:00', 'Входящий', talk=60, queue='3001'),
                   touch('2026-10-05 09:01:00', 'Входящий (не приняли)', queue='3001', result='Не ответил'),
                   touch('2026-10-05 09:02:00', 'Входящий (не приняли)', queue='3001', result='Клиент сбросил'),
                   touch('2026-10-05 09:03:00', 'Входящий', talk=30, queue=''),
                   touch('2026-10-05 09:04:00', talk=20)]
        cols = A.call_columns(touches, 'osnova')
        self.assertEqual(cols['calls'], 3)
        self.assertEqual(cols['connected'], 2)

    def test_measured_talk_wins_over_billsec(self):
        cols = A.call_columns([touch('2026-10-05 09:00:00', 'Входящий', talk=19, queue='3001', measured=8)],
                              'osnova')
        self.assertEqual(cols['connected'], 0)
        self.assertEqual(cols['talk_seconds'], 8)

    def test_no_calls_gives_no_average(self):
        self.assertIsNone(A.call_columns([], 'yar')['avg_talk_seconds'])

    def test_phone_talk_corrects_incoming_without_station_facts(self):
        incoming = touch('2026-10-05 09:00:00', 'Входящий', talk=49, queue='3001', dial=49)
        events = {'6650': ([at('2026-10-05 09:00:18'), at('2026-10-05 09:00:49')], [True, False])}
        fixed = A.attach_phone_talk([incoming], events)[0]
        self.assertEqual(fixed['talk_seconds'], 31)
        self.assertEqual(incoming['talk_seconds'], 49)
        measured = dict(incoming, talk_measured_seconds=31)
        self.assertIs(A.attach_phone_talk([measured], events)[0], measured)
        self.assertIs(A.attach_phone_talk([incoming], {})[0], incoming)


class BuildReportTests(unittest.TestCase):
    def setUp(self):
        self.memberships = {
            1: [{'id': 1, 'start_date': date(2026, 1, 1), 'end_date': None, 'group_id': 15, 'model': 'op_yandex_reg'}],
            2: [{'id': 2, 'start_date': date(2026, 1, 1), 'end_date': date(2026, 10, 9), 'group_id': 36, 'model': 'op_osnova'},
                {'id': 3, 'start_date': date(2026, 10, 10), 'end_date': None, 'group_id': 13, 'model': 'op_verificator'}],
        }
        self.events = {
            1: [ev('2026-10-05 09:00:00', 'исход'), ev('2026-10-05 09:00:10', 'соединение'),
                ev('2026-10-05 09:00:30', 'занят'), ev('2026-10-05 09:02:30', 'исход'),
                ev('2026-10-05 12:00:00', 'соединение'), ev('2026-10-05 12:00:20', 'исход'),
                ev('2026-10-05 18:00:00', 'выключен'),
                # Смена сентября: в отчёт за октябрь не идёт.
                ev('2026-09-30 20:00:00', 'исход'), ev('2026-10-01 08:00:00', 'выключен')],
            2: [ev('2026-10-06 20:00:00', 'готов'), ev('2026-10-07 01:00:00', 'занят'),
                ev('2026-10-07 01:02:00', 'готов'), ev('2026-10-07 08:00:00', 'выключен'),
                # Перевели в верификаторы — эта смена уже не ОП.
                ev('2026-10-12 09:00:00', 'готов'), ev('2026-10-12 18:00:00', 'выключен')],
        }
        self.touches = [touch('2026-10-05 09:00:10', talk=120), touch('2026-10-05 12:00:00', talk=5),
                        touch('2026-10-05 19:00:00', talk=60),                    # после смены
                        touch('2026-10-07 01:00:00', 'Входящий', talk=120, ext='6651', queue='3001'),
                        touch('2026-10-01 07:00:00', talk=30)]

    def report(self):
        return A.build_report(events_by_operator=self.events, names={1: 'Иванов', 2: 'Петрова'},
                              exts={1: '6650', 2: '6651'}, memberships=self.memberships,
                              touches=self.touches, day_from=date(2026, 10, 1),
                              day_to=date(2026, 10, 31), now=NOW)

    def test_rows(self):
        report = self.report()
        rows = report['shifts']
        self.assertEqual([(r['date'], r['name'], r['group_label']) for r in rows],
                         [('2026-10-05', 'Иванов', 'ЯР'), ('2026-10-06', 'Петрова', 'Основа')])
        yar, night = rows
        self.assertEqual((yar['calls'], yar['connected'], yar['talk_seconds']), (2, 1, 125))
        self.assertEqual(yar['login_at'], '2026-10-05T09:00:00')
        self.assertEqual(yar['logout_at'], '2026-10-05T18:00:00')
        self.assertFalse(yar['logout_missing'])
        self.assertEqual(night['calls'], 1)
        self.assertTrue(night['logout_next_day'])
        self.assertEqual([op['name'] for op in report['operators']], ['Иванов', 'Петрова'])

    def test_filters(self):
        rows = self.report()['shifts']
        self.assertEqual(len(A.filter_shifts(rows, group_key='osnova')), 1)
        self.assertEqual(len(A.filter_shifts(rows, operator_ids=[1])), 1)
        self.assertEqual(len(A.filter_shifts(rows, days=['2026-10-06'])), 1)
        self.assertEqual(len(A.filter_shifts(rows, group_key='yar', days=['2026-10-06'])), 0)


class PeriodTests(unittest.TestCase):
    def test_month_bounds(self):
        self.assertEqual(A.month_bounds('2026-02', date(2026, 10, 19)), (date(2026, 2, 1), date(2026, 2, 28)))
        self.assertEqual(A.month_bounds('2026-12', date(2026, 10, 19))[1], date(2026, 12, 31))
        self.assertEqual(A.month_bounds('', date(2026, 10, 19))[0], date(2026, 10, 1))
        with self.assertRaises(ValueError):
            A.month_bounds('2026-13', date(2026, 10, 19))

    def test_available_months_stop_at_retention(self):
        months = A.available_months(date(2026, 10, 19), 120)
        self.assertEqual([m['value'] for m in months], ['2026-10', '2026-09', '2026-08', '2026-07', '2026-06'])
        self.assertTrue(months[-1]['partial'])
        self.assertEqual(months[-1]['label'], 'Июнь 2026 · с 21.06')
        self.assertEqual(months[0]['label'], 'Октябрь 2026')

    def test_days_and_ids(self):
        self.assertEqual(A.parse_days('2026-10-03, 2026-10-01', date(2026, 10, 1), date(2026, 10, 31)),
                         ['2026-10-01', '2026-10-03'])
        with self.assertRaises(ValueError):
            A.parse_days('2026-11-01', date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(A.parse_ids('1, 2'), [1, 2])
        with self.assertRaises(ValueError):
            A.parse_ids('1;drop')


class ExportTests(unittest.TestCase):
    def test_workbook(self):
        rows = BuildReportTests('test_rows')
        rows.setUp()
        shifts = rows.report()['shifts']
        shifts[1]['logout_missing'] = True
        content = E.workbook(shifts, month_label='Октябрь 2026', group_label='Все группы',
                             generated_at='19.10.2026 12:00', calls_pending_days=['2026-10-18'])
        book = load_workbook(BytesIO(content))
        sheet = book['Смены']
        self.assertEqual([c.value for c in sheet[1]], [title for title, _ in E.COLUMNS])
        self.assertEqual(sheet['A2'].number_format, E.DATE_FORMAT)
        self.assertEqual(sheet['F2'].number_format, E.DURATION_FORMAT)
        self.assertEqual(sheet['F2'].value, timedelta(seconds=125))
        self.assertEqual(sheet['K2'].value, 2)
        self.assertTrue(sheet['M3'].font.italic)
        self.assertFalse(sheet['M2'].font.italic)
        self.assertIn('Как считается', book.sheetnames)
        notes = [sheet.cell(row=r, column=1).value for r in range(5, 9)]
        self.assertTrue(any('догружаются' in str(n) for n in notes))


class GuardTests(unittest.TestCase):
    def test_named_access_only_widens_and_never_skips_login(self):
        forbidden = lambda: (5, ('forbidden', 403))
        anonymous = lambda: (None, ('no auth', 401))
        self.assertEqual(R.make_activity_guard(forbidden, extra_user_ids={5})(), (5, None))
        self.assertEqual(R.make_activity_guard(forbidden, extra_user_ids={6})()[1], ('forbidden', 403))
        self.assertEqual(R.make_activity_guard(anonymous, extra_user_ids={5})()[1], ('no auth', 401))


class RoutesTests(unittest.TestCase):
    """Ручки без базы: проверки периода и права. Расчёт проверен выше."""

    def setUp(self):
        self.refusal = None

        def guard():
            return 10, self.refusal

        class _Db:
            def _get_cursor(self_inner):
                raise AssertionError('сюда отчёт с неверным периодом не доходит')

        app = Flask(__name__)
        app.register_blueprint(R.build_op_activity_blueprint(
            db=_Db(), require_api_key=lambda fn: fn, build_cors_preflight_response=lambda: ('', 204),
            guard=guard, retention_days=120))
        self.client = app.test_client()

    def test_refused(self):
        with Flask(__name__).app_context():
            self.refusal = ('forbidden', 403)
        self.assertEqual(self.client.get('/api/op_wallboard/activity?month=2026-10').status_code, 403)

    def test_bad_month(self):
        self.assertEqual(self.client.get('/api/op_wallboard/activity?month=2026-1').status_code, 400)
        self.assertEqual(self.client.get('/api/op_wallboard/activity?month=2025-01').status_code, 400)
        self.assertEqual(self.client.get('/api/op_wallboard/activity?month=2099-01').status_code, 400)

    def test_month_json_is_gzipped_when_the_client_accepts_it(self):
        import gzip as gzip_mod, json as json_mod
        app = Flask(__name__)
        payload = {'shifts': [{'name': 'Иванов', 'n': i} for i in range(200)]}
        with app.test_request_context(headers={'Accept-Encoding': 'gzip, deflate, br'}):
            response = R.json_response(payload)
            self.assertEqual(response.headers['Content-Encoding'], 'gzip')
            self.assertEqual(json_mod.loads(gzip_mod.decompress(response.get_data())), payload)
        with app.test_request_context():
            response = R.json_response(payload)
            self.assertNotIn('Content-Encoding', response.headers)
            self.assertEqual(json_mod.loads(response.get_data()), payload)

    def test_journal_window_is_checked(self):
        base = '/api/op_wallboard/activity/journal?operator_id=1'
        self.assertEqual(self.client.get(base + '&start=2026-10-05T09:00:00&end=2026-10-07T09:00:00').status_code, 400)
        self.assertEqual(self.client.get(base + '&start=2026-10-05T09:00:00&end=2026-10-05T08:00:00').status_code, 400)
        self.assertEqual(self.client.get('/api/op_wallboard/activity/journal?operator_id=x').status_code, 400)


if __name__ == '__main__':
    unittest.main()


class FrontendWiringTests(unittest.TestCase):
    """Где живёт отчёт (решения владельца 19.10.2026): не вкладка «Линия / Чат» и не отдельный раздел
    меню, а полноэкранное окно «Табло ОП», которое открывает кнопка «Отчёт по сменам» в шапке линии."""

    @classmethod
    def setUpClass(cls):
        from pathlib import Path
        root = Path(__file__).resolve().parents[1]
        monitoring = root / 'src' / 'components' / 'monitoring'
        cls.app = (root / 'src' / 'App.jsx').read_text(encoding='utf-8-sig')
        cls.shared = (monitoring / 'opWallboardShared.js').read_text(encoding='utf-8-sig')
        cls.view = (monitoring / 'OpWallboardView.jsx').read_text(encoding='utf-8-sig')
        cls.sheet = (monitoring / 'OpActivityReport.jsx').read_text(encoding='utf-8-sig')

    def test_not_a_tab_and_not_a_menu_section(self):
        self.assertNotIn("key: 'report'", self.shared)
        self.assertNotIn('op_activity', self.app)
        self.assertNotIn('OpActivityReport', self.app)

    def test_button_next_to_broadcast_in_both_headers_and_one_lazy_sheet(self):
        """Кнопка — рядом с «Отбивкой» и «Виджетом» в шапке любого направления; окно одно на табло."""
        line = self.view[self.view.index('function OpLineWallboard('):self.view.index('// ── Направление «Чат»')]
        chat = self.view[self.view.index('function OpChatBoard('):self.view.index('export default function OpWallboardView(')]
        root = self.view[self.view.index('export default function OpWallboardView('):]
        for header in (line, chat):
            self.assertEqual(header.count('{reportButton}'), 1)
            # Сразу за «Отбивкой» и до «Виджета».
            self.assertLess(header.index('<BroadcastControls'), header.index('{reportButton}'))
            self.assertLess(header.index('{reportButton}'), header.index('<WidgetButton'))
        self.assertIn('Отчёт по сменам', root)
        self.assertIn('onClick={openReport}', root)
        self.assertEqual(root.count('{reportSheet}'), 1)
        self.assertIn("const loadActivityReport = () => import('./OpActivityReport');", self.view)
        self.assertIn('const OpActivityReportSheet = lazy(loadActivityReport);', self.view)
        self.assertIn('<FullscreenSheet', self.sheet)
        self.assertIn("offsetLeft=\"var(--app-sidebar-offset, 0px)\"", self.sheet)
        self.assertIn('closeOnEscape={!journalShift}', self.sheet)
        # Окно — в корень страницы: внутри столбика space-y раздела оно встало бы ниже края.
        self.assertIn('createPortal(<FullscreenSheet', self.sheet)
