# -*- coding: utf-8 -*-
"""Две просьбы владельца от 23.09.2026 к разделу «Новости».

1. ВРЕМЯ НА ЧТЕНИЕ И НА ТЕСТ В ОКНЕ OKTELL. Дословно: «указывается время,
   после открытия новости идёт таймер… если имеется тест, то для него тоже
   ставится время, чтобы оператор не злоупотреблял им, данная настройка должна
   появляться только при отправке новости в Oktell». На уточнение: это ПОТОЛОК
   (сколько отведено), и время вышло — «ничего не прерывать, отметить».

2. «ВОЗМОЖНОСТЬ ПУБЛИКОВАТЬ СТАТЬЮ И КАК НОВОСТЬ» — из редактора статьи и из
   самой статьи.

Тесты чистые, как и остальные тесты раздела: правила — функциями, проводка —
чтением исходников.
"""

import io
import os
import re
import unittest

from news import access as news_access
from news import article_draft
from news import report_xlsx
from news import schema as news_schema

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _read(*parts):
    with open(os.path.join(ROOT, *parts), encoding='utf-8') as handle:
        return handle.read()


def _code_only(source):
    without_blocks = re.sub(r'"""[\s\S]*?"""', '', source)
    return re.sub(r'(?m)#.*$', '', without_blocks)


def _jsx_code_only(source):
    return re.sub(r'/\*[\s\S]*?\*/', '', source)


def _between(source, start, end):
    body = source[source.index(start):]
    return body[:body.index(end, len(start))]


# ─────────────────────────────────────────────────────────────────────────────
# 1. Время на чтение и на тест
# ─────────────────────────────────────────────────────────────────────────────

class TimeLimitRulesTests(unittest.TestCase):

    def test_limit_is_normalized_not_refused(self):
        normalize = news_access.normalize_time_limit
        # «Без ограничения» — всё пустое и неположительное.
        for raw in (None, '', 0, '0', -5, 'мусор', False):
            self.assertIsNone(normalize(raw), raw)
        self.assertEqual(normalize(120), 120)
        self.assertEqual(normalize('180'), 180)
        # Ближайшее допустимое, как у задержки кнопки.
        self.assertEqual(normalize(5), news_schema.MIN_TIME_LIMIT_SECONDS)
        self.assertEqual(normalize(10 ** 6), news_schema.MAX_TIME_LIMIT_SECONDS)

    def test_limits_exist_only_for_a_mandatory_news_in_oktell(self):
        """«Настройка должна появляться только при отправке новости в Oktell».
        Необязательное объявление в клиент АТС не уходит вовсе — лимит у него
        настраивал бы окно, которого не будет."""
        limits = news_access.time_limits
        self.assertEqual(limits(channel='icore', is_mandatory=True, has_quiz=True,
                                read_limit=180, quiz_limit=60), (None, None))
        self.assertEqual(limits(channel='oktell', is_mandatory=False, has_quiz=True,
                                read_limit=180, quiz_limit=60), (None, None))
        self.assertEqual(limits(channel='oktell', is_mandatory=True, has_quiz=True,
                                read_limit=180, quiz_limit=60), (180, 60))
        # Время на тест — только когда тест есть.
        self.assertEqual(limits(channel='oktell', is_mandatory=True, has_quiz=False,
                                read_limit=180, quiz_limit=60), (180, None))

    def test_reading_time_cannot_be_shorter_than_the_button_delay(self):
        refusal = news_access.time_limit_refusal
        self.assertTrue(refusal(read_limit=60, confirm_delay_seconds=120))
        self.assertIsNone(refusal(read_limit=120, confirm_delay_seconds=120))
        self.assertIsNone(refusal(read_limit=None, confirm_delay_seconds=600))
        self.assertIsNone(refusal(read_limit=60, confirm_delay_seconds=0))

    def test_overtime_is_nothing_to_compare_without_a_limit_or_a_measure(self):
        overtime = news_access.overtime
        self.assertIsNone(overtime(None, 180))       # старый агент — замера нет
        self.assertIsNone(overtime(500, None))       # лимита нет
        self.assertEqual(overtime(100, 180), 0)      # уложился
        self.assertEqual(overtime(252, 180), 72)     # +1:12

    def test_spent_time_from_the_agent_is_clamped(self):
        clean = news_access.clean_spent
        self.assertIsNone(clean(None))
        self.assertIsNone(clean('abc'))
        self.assertEqual(clean(-3), 0)
        self.assertEqual(clean(10 ** 9), news_schema.MAX_SPENT_SECONDS)
        self.assertEqual(clean('42'), 42)

    def test_summary_counts_people_over_the_time(self):
        rows = [
            {'in_audience': True, 'status': 'passed', 'read_over_seconds': 12},
            {'in_audience': True, 'status': 'passed', 'quiz_over_seconds': 5},
            {'in_audience': True, 'status': 'passed', 'read_over_seconds': 0},
            {'in_audience': True, 'status': 'pending'},
            # Выбывший из адресатов в сводку не идёт — как и в остальных счётчиках.
            {'in_audience': False, 'status': 'passed', 'read_over_seconds': 30},
        ]
        self.assertEqual(news_access.report_summary(rows)['overtime'], 2)


class TimeLimitSchemaTests(unittest.TestCase):

    def test_columns_arrive_idempotently_and_have_a_latch(self):
        ddl = '\n'.join(news_schema._STATEMENTS)
        for table, column in news_schema.LIMIT_COLUMNS:
            self.assertIn('ALTER TABLE %s ADD COLUMN IF NOT EXISTS %s' % (table, column), ddl)
        self.assertTrue(callable(news_schema.limits_ready))

    def test_every_reader_asks_the_latch(self):
        """Не доехавший DDL не должен ронять ни карточку, ни журнал, ни агента."""
        routes = _read('news', 'routes.py')
        self.assertIn('with_limits=_limits_ready(cursor)', routes)
        self.assertIn('limits_ready as schema_limits_ready', routes)
        guard = _read('oktell_guard', 'routes.py')
        self.assertIn('_news_limits_ready(cursor, news_limits_ready)', guard)


class TimeLimitRoutesTests(unittest.TestCase):

    def setUp(self):
        self.routes = _read('news', 'routes.py')

    def test_limits_are_checked_before_the_first_write(self):
        """Ответ 4xx здесь всё равно фиксирует транзакцию: отказ после записи
        оставил бы новость, которую автор считает отменённой."""
        create = _between(self.routes, 'def news_post_create(', 'def news_post_update(')
        self.assertLess(create.index('_limits_from_request('), create.index('queries.create_post('))
        self.assertIn('queries.set_time_limits(', create)
        update = _between(self.routes, 'def news_post_update(', 'def news_post_publish(')
        self.assertLess(update.index('_limits_from_request('), update.index('queries.update_post('))
        self.assertIn('queries.set_time_limits(', update)

    def test_limits_of_a_published_news_are_locked(self):
        helper = _between(self.routes, 'def _limits_from_request(', 'def _moment_or_none(')
        self.assertIn('NEWS_LIMITS_LOCKED', helper)
        self.assertIn("current.get('published_at')", helper)
        self.assertIn('NEWS_READ_LIMIT_SHORT', helper)
        self.assertIn('news_access.time_limits(', helper)

    def test_report_marks_overtime_with_the_single_rule(self):
        report = _between(self.routes, 'def _report_data(', '@news_route(')
        self.assertIn('news_access.overtime(', report)
        self.assertIn("with_limits=_limits_ready(cursor)", report)

    def test_audit_compares_the_limits(self):
        from news import audit as news_audit
        self.assertIn('read_limit_seconds', news_audit._COMPARED)
        self.assertIn('quiz_limit_seconds', news_audit._COMPARED)
        labels = _read('src', 'components', 'wiki', 'auditEvents.js')
        self.assertIn("read_limit_seconds: 'время на чтение'", labels)
        self.assertIn("quiz_limit_seconds: 'время на тест'", labels)


class TimeLimitAgentDoorsTests(unittest.TestCase):

    def setUp(self):
        self.guard = _code_only(_read('oktell_guard', 'routes.py'))

    def test_news_payload_carries_limits_and_what_is_already_spent(self):
        news = _between(self.guard, 'def oktell_guard_agent_news(', '@agent_route(')
        # ПОСЛЕ mark_shown: у только что показанного строка отметки уже есть.
        self.assertLess(news.index('mark_shown('), news.index('agent_time_state('))

    def test_progress_door_writes_only_the_measure(self):
        progress = _between(self.guard, 'def oktell_guard_agent_news_progress(', '@agent_route(')
        self.assertIn('record_time_spent(', progress)
        self.assertIn("owner['user_id']", progress)
        for forbidden in ('confirm_read', 'mark_shown', 'mark_trainer'):
            self.assertNotIn(forbidden, progress)

    def test_wrong_attempt_still_records_the_time(self):
        read = _between(self.guard, 'def oktell_guard_agent_news_read(', '@agent_route(')
        self.assertLess(read.index('record_time_spent('), read.index('confirm_read('))

    def test_state_door_stays_without_side_effects(self):
        state = _between(self.guard, 'def oktell_guard_agent_news_state(', '@agent_route(')
        self.assertNotIn('record_time_spent', state)

    def test_recording_takes_the_greatest_not_the_sum(self):
        """Окно шлёт НАРАСТАЮЩИЙ итог: повтор замера не должен раздувать время."""
        queries = _read('news', 'queries.py')
        record = _between(queries, 'def record_time_spent(', '\ndef ')
        self.assertIn('GREATEST(', record)
        self.assertNotIn('read_spent_seconds + ', record)


class TimeLimitAgentWindowTests(unittest.TestCase):

    def setUp(self):
        self.agent = _read('oktell_recall_guard', 'agent.py')

    def test_payload_of_the_window_has_limits_and_spent(self):
        builder = _between(self.agent, 'def build_news_js(', '\ndef ')
        for key in ('read_limit_seconds', 'quiz_limit_seconds',
                    'read_spent_seconds', 'quiz_spent_seconds'):
            self.assertIn('"%s"' % key, builder)

    def test_window_counts_on_and_never_closes_itself(self):
        template = _between(self.agent, 'NEWS_JS_TEMPLATE = r"""', '""".strip()')
        self.assertIn("'Осталось '", template)
        self.assertIn("'Время вышло", template)
        # Счёт продолжается с серверного итога, а не с нуля.
        self.assertIn('Number(data.read_spent_seconds', template)
        # Итог уходит вместе с ответами.
        self.assertIn('read_seconds: state.spent.read', template)
        # «Ничего не прерывать»: истёкшее время не нажимает и не закрывает.
        timer = _between(template, 'function paintTimer()', '\n  }\n')
        for forbidden in ('button.click', 'state.result', 'removeChild', 'disabled'):
            self.assertNotIn(forbidden, timer)

    def test_agent_reports_progress_only_for_news_with_limits(self):
        press = _between(self.agent, '    def handle_news_press(', '    def release_news(')
        self.assertIn('has_time_limits(active_news)', press)
        self.assertIn('link.news_progress(', press)
        self.assertLess(press.index('has_time_limits(active_news)'),
                        press.index('link.news_progress('))
        self.assertIn('link.news_read(pressed.get("id"), pressed.get("answers") or {}, pressed)',
                      press)
        # Статус оператора по-прежнему возвращает только release_news.
        self.assertNotIn('set_operator_status', press)

    def test_the_change_rides_the_unreleased_build(self):
        self.assertIn('VERSION = "1.0.32"', self.agent)

    def test_progress_helper_reads_without_side_effects(self):
        from oktell_recall_guard import agent
        script = agent.build_news_progress_js()
        self.assertIn('state.spent', script)
        self.assertNotIn('state.result = null', script)
        self.assertTrue(agent.has_time_limits({'read_limit_seconds': 60}))
        self.assertTrue(agent.has_time_limits({'quiz_limit_seconds': 60}))
        self.assertFalse(agent.has_time_limits({}))
        self.assertFalse(agent.has_time_limits(None))


class TimeLimitFormTests(unittest.TestCase):

    def setUp(self):
        self.form = _read('src', 'components', 'wiki', 'WikiNews.jsx')

    def test_rows_exist_only_for_oktell(self):
        code = _jsx_code_only(self.form)
        self.assertIn("const limitsShown = channel === 'oktell' && (mandatory || mustPass);", code)
        self.assertIn('{limitsShown && (', code)
        self.assertIn('{limitsShown && quiz.length > 0 && (', code)
        self.assertIn('read_limit_seconds: limitsShown ? readLimit : null', code)

    def test_choices_are_inside_the_server_bounds(self):
        for name in ('READ_LIMIT_OPTIONS', 'QUIZ_LIMIT_OPTIONS'):
            block = _between(self.form, 'const %s = [' % name, '];')
            values = [int(v) for v in re.findall(r'value:\s*(\d+)', block)]
            self.assertTrue(values)
            self.assertIn('value: null', block)
            for value in values:
                self.assertGreaterEqual(value, news_schema.MIN_TIME_LIMIT_SECONDS)
                self.assertLessEqual(value, news_schema.MAX_TIME_LIMIT_SECONDS)

    def test_limits_are_locked_with_the_channel(self):
        rows = _between(self.form, '{limitsShown && (', '<SettingRow label="Тест и тренажёр"')
        self.assertEqual(rows.count('disabled={channelLocked}'), 2)

    def test_report_marks_overtime(self):
        code = _jsx_code_only(self.form)
        self.assertIn('row.read_over_seconds > 0', code)
        self.assertIn('row.quiz_over_seconds > 0', code)
        self.assertIn('state.overtime > 0', code)


class TimeLimitExcelTests(unittest.TestCase):

    POST = {'published_at': '2026-09-23T09:00:00'}

    def _sheet(self, rows):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(report_xlsx.build(self.POST, rows).getvalue()))
        return book['Ознакомление']

    def _row(self, **extra):
        row = {'user_id': 7, 'name': 'Оператор', 'department_name': 'СЗоВ', 'role': 'operator',
               'shown_at': '2026-09-23T09:05:00', 'confirmed_at': '2026-09-23T09:10:00',
               'attempts': 1, 'status': 'done', 'in_audience': True}
        row.update(extra)
        return row

    def test_columns_appear_only_when_there_is_a_measure(self):
        headers = [cell.value for cell in self._sheet([self._row()])[1]]
        self.assertNotIn('На чтении', headers)
        self.assertNotIn('Сверх отведённого', headers)

    def test_overtime_is_written_and_marked(self):
        sheet = self._sheet([self._row(read_spent_seconds=252, quiz_spent_seconds=40,
                                       read_over_seconds=72, quiz_over_seconds=0)])
        headers = [cell.value for cell in sheet[1]]
        self.assertIn('На чтении', headers)
        self.assertEqual(sheet.cell(row=2, column=headers.index('На чтении') + 1).value, '4:12')
        over = sheet.cell(row=2, column=headers.index('Сверх отведённого') + 1)
        self.assertEqual(over.value, 'чтение +1:12')
        self.assertEqual(over.fill.fgColor.rgb[-6:], 'FEF3C7')

    def test_duration_label(self):
        self.assertEqual(report_xlsx._duration(None), '')
        self.assertEqual(report_xlsx._duration(5), '0:05')
        self.assertEqual(report_xlsx._duration(3725), '1:02:05')


# ─────────────────────────────────────────────────────────────────────────────
# 2. Статья как новость
# ─────────────────────────────────────────────────────────────────────────────

FILE_A = '/api/wiki/file/0a1b2c3d-1111-2222-3333-444455556666'
FILE_B = '/api/wiki/file/0a1b2c3d-1111-2222-3333-444455557777'


class ArticleDraftTests(unittest.TestCase):

    HTML = (
        '<div data-wiki-block="lead"><p>Вводка про <a href="?view=wiki&amp;article=tarify">'
        'тарифы</a> и <a href="https://pro.yandex.com" target="_blank">сайт</a>.</p></div>'
        '<h4 style="color:red">Глубокий</h4><p><img src="%s" data-width="50"></p>'
        '<div data-wiki-block="gallery"><img src="%s"><img src="%s"></div>'
        '<details data-wiki-collapsible="1"><summary>Подробнее</summary><p>Скрытый</p></details>'
        '<ul data-variant="steps"><li><p>Шаг <mark data-color="#fef3c7">один</mark></p></li></ul>'
        '<div data-wiki-trainer="order-app" class="wiki-trainer-embed">'
        '<span class="wiki-trainer-embed__label">Открыть тренажёр</span></div>'
        '<p><a href="#glava-2">к главе</a></p>'
        '<table><tbody><tr><td colspan="2" style="width:10px">A</td></tr></tbody></table>'
    ) % (FILE_A, FILE_B, FILE_A)

    def setUp(self):
        self.draft = article_draft.draft({'title': '  Тарифы  ', 'content': self.HTML})
        self.body = self.draft['body']

    def test_images_move_to_the_carousel_in_order_without_repeats(self):
        self.assertEqual(self.draft['images'], [FILE_A, FILE_B])
        self.assertNotIn('<img', self.body)

    def test_trainer_moves_to_the_news_trainer(self):
        self.assertEqual(self.draft['trainer_key'], 'order-app')
        self.assertNotIn('Открыть тренажёр', self.body)

    def test_wiki_links_lose_the_link_and_keep_the_text(self):
        self.assertIn('Вводка про тарифы', self.body)
        self.assertNotIn('article=', self.body)
        self.assertNotIn('#glava-2', self.body)
        self.assertIn('к главе', self.body)
        # Внешняя ссылка остаётся ссылкой.
        self.assertIn('href="https://pro.yandex.com"', self.body)

    def test_markup_is_reduced_to_what_the_news_understands(self):
        self.assertIn('<h3>Глубокий</h3>', self.body)
        self.assertIn('<p><strong>Подробнее</strong></p>', self.body)
        self.assertIn('<mark>один</mark>', self.body)
        self.assertIn('<td colspan="2">A</td>', self.body)
        for gone in ('data-wiki', 'style=', 'class=', '<div', '<details', '<summary',
                     'data-variant', 'data-color', '<p></p>'):
            self.assertNotIn(gone, self.body)

    def test_title_is_trimmed(self):
        self.assertEqual(self.draft['title'], 'Тарифы')

    def test_wiki_link_rule(self):
        self.assertTrue(article_draft.is_wiki_link('?view=wiki&article=tarify'))
        self.assertTrue(article_draft.is_wiki_link(
            'https://alfa330.github.io/OTP/?view=wiki&article=tarify'))
        self.assertTrue(article_draft.is_wiki_link('#chapter'))
        self.assertFalse(article_draft.is_wiki_link('https://pro.yandex.com/tariffs'))
        self.assertFalse(article_draft.is_wiki_link('mailto:hr@example.com'))

    def test_empty_article(self):
        draft = article_draft.draft({'title': None, 'content': None})
        self.assertEqual(draft, {'title': '', 'body': '', 'images': [], 'trainer_key': None})


class ArticleDraftDoorTests(unittest.TestCase):

    def setUp(self):
        self.route = _read('wiki', 'routes_news.py')
        self.code = _code_only(self.route)

    def test_door_is_registered_with_the_news_bucket(self):
        self.assertIn('routes_news.register(bp, wiki_route, db, _ip, gcs or {})',
                      _read('wiki', 'routes.py'))
        self.assertIn("'news_bucket_name': _news_bucket_name,", _read('bot_schedule2.py'))

    def test_the_article_must_be_visible_and_published(self):
        door = _between(self.code, 'def wiki_article_news_draft(', 'return jsonify({\n')
        self.assertIn('read_perimeter(cursor, ctx)', door)
        self.assertIn("article['id'] not in visible", door)
        self.assertIn("article.get('status') != 'published'", door)
        # Право выпуска — лестница новостей, а не права статьи.
        self.assertIn('news_access.publish_ceiling(', door)

    def test_files_follow_the_rule_of_the_file_door(self):
        allowed = _between(self.code, 'def _file_allowed(', 'def _image_bytes(')
        self.assertIn('owner_article in visible', allowed)
        self.assertIn("row.get('uploaded_by') == ctx['user_id']", allowed)

    def test_external_images_are_never_fetched(self):
        """Сервер, скачивающий по ссылке из тела статьи, — дверь в чужую сеть."""
        for forbidden in ('requests.', 'urlopen', 'urllib', 'httpx'):
            self.assertNotIn(forbidden, self.code)

    def test_photos_are_capped_like_a_regular_upload(self):
        copy = _between(self.code, 'def _copy_images(', '@wiki_route(')
        self.assertIn('MAX_PHOTOS_PER_POST', copy)
        self.assertIn('MAX_LOOSE_PHOTOS_PER_USER', copy)
        self.assertIn('news_photos.prepare(', copy)
        self.assertIn('insert_loose_photo(', copy)


class _Blob:
    def __init__(self, store, bucket, path):
        self.store, self.key = store, (bucket, path)

    def download_as_bytes(self):
        self.store['downloads'].append(self.key)
        return b'bytes-of-' + self.key[1].encode()


class _Bucket:
    def __init__(self, store, name):
        self.store, self.name = store, name

    def blob(self, path):
        return _Blob(self.store, self.name, path)


class _Client:
    def __init__(self, store):
        self.store = store

    def bucket(self, name):
        return _Bucket(self.store, name)


class ArticleDraftRouteTests(unittest.TestCase):
    """Дверь целиком — на Flask, с подменёнными базой и бакетом: какие картинки
    переезжают в новость, а какие нет, и кому дверь отказывает."""

    VISIBLE_ARTICLE = 10
    HIDDEN_ARTICLE = 99

    def _files(self):
        return [
            {'id': FILE_A.rsplit('/', 1)[1], 'article_id': self.VISIBLE_ARTICLE,
             'bucket': 'wiki-bucket', 'blob_path': 'a.webp', 'content_type': 'image/webp',
             'uploaded_by': 1},
            # Скопирована из статьи, которой этому человеку не видно.
            {'id': FILE_B.rsplit('/', 1)[1], 'article_id': self.HIDDEN_ARTICLE,
             'bucket': 'wiki-bucket', 'blob_path': 'b.webp', 'content_type': 'image/webp',
             'uploaded_by': 1},
        ]

    def _call(self, *, content, status='published', role='sv', storage=True,
              visible=True, loose=0):
        from unittest import mock
        from flask import Blueprint, Flask
        from wiki import routes_news

        store = {'downloads': [], 'prepared': [], 'inserted': []}
        article = {'id': self.VISIBLE_ARTICLE, 'title': 'Тарифы', 'content': content,
                   'status': status}
        bp = Blueprint('wiki_test', __name__)

        def wiki_route(rule, methods=('GET',)):
            def decorator(handler):
                def view(**kwargs):
                    return handler(cursor=object(), ctx={'user_id': 5}, **kwargs)
                bp.add_url_rule(rule, handler.__name__, view, methods=list(methods))
                return handler
            return decorator

        gcs = {'client': lambda: _Client(store)}
        if storage:
            gcs['news_bucket_name'] = lambda: 'news-bucket'
        routes_news.register(bp, wiki_route, db=None, log_ip=lambda: None, gcs=gcs)
        app = Flask(__name__)
        app.register_blueprint(bp)

        def prepare(data, *, filename, content_type):
            store['prepared'].append((data, content_type))
            return {'data': data, 'content_type': 'image/webp', 'original_name': filename,
                    'file_size': len(data), 'width': 1, 'height': 1}

        def insert(cursor, *, prepared, bucket, blob_path, uploaded_by):
            store['inserted'].append((bucket, uploaded_by))
            return {'id': 'p%d' % len(store['inserted']), 'bucket': bucket,
                    'blob_path': blob_path, 'content_type': 'image/webp'}

        patches = [
            mock.patch.object(routes_news.wiki_perimeter, 'read_perimeter',
                              return_value=(set(), set(),
                                            {self.VISIBLE_ARTICLE} if visible else set())),
            mock.patch.object(routes_news.wiki_articles, 'get_article', return_value=article),
            mock.patch.object(routes_news.wiki_articles, 'files_for_display',
                              side_effect=lambda cursor, ids: [
                                  row for row in self._files() if row['id'] in ids]),
            mock.patch.object(routes_news.news_schema, 'schema_is_ready', return_value=True),
            mock.patch.object(routes_news.news_schema, 'photos_ready', return_value=True),
            mock.patch.object(routes_news.news_queries, 'load_viewer_context',
                              return_value={'otp_role': role}),
            mock.patch.object(routes_news.news_queries, 'is_wiki_admin', return_value=False),
            mock.patch.object(routes_news.news_queries, 'count_loose_photos', return_value=loose),
            mock.patch.object(routes_news.news_queries, 'insert_loose_photo', side_effect=insert),
            mock.patch.object(routes_news.news_photos, 'prepare', side_effect=prepare),
            mock.patch.object(routes_news.news_photos, 'upload',
                              side_effect=lambda gcs, prepared: (gcs['bucket_name'](), 'x.webp')),
            mock.patch.object(routes_news.news_photos, 'sign_urls',
                              side_effect=lambda gcs, rows: [
                                  {'id': row['id'], 'url': 'https://signed/' + row['id']}
                                  for row in rows]),
        ]
        for patch in patches:
            patch.start()
        try:
            response = app.test_client().post('/articles/%d/news-draft' % self.VISIBLE_ARTICLE)
        finally:
            for patch in patches:
                patch.stop()
        return response, store

    def test_only_visible_files_and_inline_images_move_into_the_news(self):
        content = ('<p>Текст</p><img src="%s"><img src="%s">'
                   '<img src="data:image/png;base64,aGVsbG8=">'
                   '<img src="https://evil.example/x.png">') % (FILE_A, FILE_B)
        response, store = self._call(content=content)
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual([photo['id'] for photo in data['photos']], ['p1', 'p2'])
        # Чужой файл и внешний адрес — пропущены, и об этом сказано числом.
        self.assertEqual(data['skipped_photos'], 2)
        # Скачан только файл видимой статьи; внешний адрес сервер не трогал.
        self.assertEqual(store['downloads'], [('wiki-bucket', 'a.webp')])
        self.assertEqual(store['prepared'][1], (b'hello', 'image/png'))
        # Кадры — в бакет НОВОСТЕЙ и «ничьи» этого человека.
        self.assertEqual(store['inserted'], [('news-bucket', 5), ('news-bucket', 5)])
        self.assertEqual(data['body'], '<p>Текст</p>')
        self.assertEqual(data['title'], 'Тарифы')

    def test_carousel_cap_and_loose_ceiling(self):
        content = ''.join('<img src="data:image/png;base64,aGVsbG8%d">' % i for i in range(12))
        _response, store = self._call(content=content)
        self.assertEqual(len(store['inserted']), news_schema.MAX_PHOTOS_PER_POST)
        # Почти исчерпанный потолок «ничьих» режет ещё сильнее.
        response, store = self._call(content=content,
                                     loose=news_schema.MAX_LOOSE_PHOTOS_PER_USER - 3)
        self.assertEqual(len(store['inserted']), 3)
        self.assertEqual(response.get_json()['skipped_photos'], 9)

    def test_without_storage_text_still_comes(self):
        response, store = self._call(content='<p>Текст</p><img src="%s">' % FILE_A,
                                     storage=False)
        data = response.get_json()
        self.assertEqual(data['photos'], [])
        self.assertEqual(data['skipped_photos'], 1)
        self.assertEqual(data['body'], '<p>Текст</p>')
        self.assertEqual(store['downloads'], [])

    def test_refusals(self):
        self.assertEqual(self._call(content='<p>x</p>', visible=False)[0].status_code, 404)
        self.assertEqual(self._call(content='<p>x</p>', status='draft')[0].status_code, 409)
        # Оператор новости не публикует — лестница новостей, не права статьи.
        self.assertEqual(self._call(content='<p>x</p>', role='operator')[0].status_code, 403)

    def test_trainer_of_the_article_becomes_the_trainer_of_the_news(self):
        content = ('<p>x</p><div data-wiki-trainer="order-app" class="wiki-trainer-embed">'
                   '<span>Открыть</span></div>')
        data = self._call(content=content)[0].get_json()
        self.assertEqual(data['trainer_key'], 'order-app')


class ArticleAsNewsFrontendTests(unittest.TestCase):

    def test_view_opens_the_news_form_and_returns_to_the_article(self):
        view = _read('src', 'components', 'wiki', 'WikiView.jsx')
        compose = _between(view, 'const composeArticleNews = useCallback(', '}, [base, headers, showToast]);')
        self.assertIn('/news-draft', compose)
        self.assertIn("source: 'article'", compose)
        self.assertIn("setTab('news')", compose)
        finish = _between(view, 'const finishNewsCompose = useCallback(', '}, [base, headers, showToast]);')
        self.assertLess(finish.index("request?.source === 'article'"),
                        finish.index('request?.questionId'))
        self.assertIn('onPublishAsNews={features.news && canPublishNews', view)

    def test_news_form_takes_photos_and_trainer_from_the_draft(self):
        news = _read('src', 'components', 'wiki', 'WikiNews.jsx')
        self.assertIn('photos: compose.draft?.photos || []', news)
        self.assertIn('trainer_key: compose.draft?.trainer_key || null', news)
        # Таблицы статьи переживают редактор новости.
        self.assertIn("from '@tiptap/extension-table'", news)

    def test_article_button_only_for_a_published_article(self):
        article = _jsx_code_only(_read('src', 'components', 'wiki', 'WikiArticle.jsx'))
        self.assertIn("onPublishAsNews && article.status === 'published'", article)

    def test_editor_opens_the_news_only_after_the_article_is_out(self):
        editor = _read('src', 'components', 'wiki', 'WikiEditor.jsx')
        code = _jsx_code_only(editor)
        self.assertIn("if (applied === 'published')", code)
        self.assertIn('onPublishAsNews({', code)
        # Правка отвечает {"status": "ok"} — это не статус статьи. Прежнее
        # чтение говорило «сохранено черновиком» про опубликованную статью.
        self.assertNotIn('const applied = r.data?.status || status;', code)
        # «Сохранить» опубликованной статьи — «Сохранено», а не «опубликована».
        self.assertIn("status !== 'published' ? 'Сохранено'", code)
        # Переключатель — не свойство статьи: правкой не считается.
        self.assertIn('<IosToggle checked={alsoNews} onChange={setAlsoNews} />', code)

    def test_library_passes_the_door_to_both_places(self):
        library = _read('src', 'components', 'wiki', 'WikiLibrary.jsx')
        self.assertEqual(library.count('onPublishAsNews={publishAsNews}'), 2)
        # Вместе со статьёй уходит дверь, с которой в неё вошли: вернувшись из
        # формы новости, человек получает статью с прежним «назад».
        self.assertIn('onPublishAsNews(article, door)', library)
        view = _read('src', 'components', 'wiki', 'WikiView.jsx')
        self.assertIn('setSearchTarget({ slug: request.articleSlug, from: request.from || null })',
                      view)


if __name__ == '__main__':
    unittest.main()
