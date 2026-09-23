# -*- coding: utf-8 -*-
"""«Биллинг чатов» (задача #343): оценка, время ответа и первая реакция в минутах,
план чатов и сотрудники, выгрузки в формате отчётов СЗоВ — «Ежедневный отчёт по чатам»
и «Отчёт с группировкой по часам»."""
import ast
import math
import re
import unittest
from contextlib import contextmanager
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path

import xlsxwriter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from resource_fte.chat import (
    CHAT_BILLING_DETAIL_EXPORT_LIMIT,
    CHAT_BILLING_METRICS,
    CHAT_BILLING_PARK_SHARES,
    _chat_billing_phone_claims_tx,
    _chat_billing_plan_base_tx,
    _chat_billing_planned_shifts_tx,
    attach_chat_billing_grouping_staff,
    attach_chat_billing_plan,
    build_chat_billing_grouping,
    build_chat_billing_planned_staff,
    build_chat_billing_staff,
    get_chat_billing_details,
    get_chat_billing_grouping,
    get_chat_billing_grouping_by_park,
    get_chat_billing_operators,
    get_chat_billing_report,
    get_chat_billing_staff,
)
from tests import source_cache

ROOT = Path(__file__).resolve().parents[1]
BOT_PATH = ROOT / "bot_schedule2.py"
VIEW_PATH = ROOT / "src" / "components" / "resources" / "ResourceFteView.jsx"
METRICS_JS_PATH = ROOT / "src" / "components" / "resources" / "chatBillingMetrics.js"

FUNCTION_NAMES = (
    "_chat_billing_round_half_up",
    "_chat_billing_minutes",
    "_chat_billing_minutes_label",
    "_chat_billing_export_ratio",
    "_chat_billing_export_metrics",
    "_chat_billing_export_detail_values",
    "_chat_billing_sheet_title",
    "_chat_billing_day_comment",
    "_chat_billing_daily_values",
    "_chat_billing_daily_workbook",
    "_chat_billing_hourly_workbook",
    "_chat_billing_export_workbook",
)
CONST_NAMES = (
    "_CHAT_BILLING_EXPORT_PCT_FMT",
    "_CHAT_BILLING_EXPORT_RATING_FMT",
    "_CHAT_BILLING_EXPORT_SUMMARY_COLUMNS",
    "_CHAT_BILLING_EXPORT_DETAIL_COLUMNS",
    "_CHAT_BILLING_SHEET_FORBIDDEN",
    "_CHAT_BILLING_MONTHS_RU",
    "_CHAT_BILLING_DAILY_COLUMNS",
    "_CHAT_BILLING_DAILY_DAY_KEYS",
    "_CHAT_BILLING_DAILY_BASE_FMT",
    "_CHAT_BILLING_DAILY_RED_HEADER_KEYS",
    "CHAT_BILLING_LOW_RATING",
    "_CHAT_BILLING_HOURLY_COLUMNS",
    "_CHAT_BILLING_HOURLY_STAFF_KEYS",
)


def _namespace():
    module = source_cache.parse(BOT_PATH.read_text(encoding="utf-8"))
    selected = [node for node in module.body
                if isinstance(node, ast.FunctionDef) and node.name in FUNCTION_NAMES]
    missing = set(FUNCTION_NAMES) - {node.name for node in selected}
    if missing:
        raise AssertionError("Не найдены функции в bot_schedule2.py: %s" % sorted(missing))
    consts = [node for node in module.body
              if isinstance(node, ast.Assign)
              and any(getattr(t, "id", "") in CONST_NAMES for t in node.targets)]
    ns = {
        "math": math, "re": re, "datetime": datetime, "Workbook": Workbook,
        "xlsxwriter": xlsxwriter, "BytesIO": BytesIO,
        "Alignment": Alignment, "Border": Border, "Font": Font, "PatternFill": PatternFill,
        "Side": Side, "get_column_letter": get_column_letter,
        "CHAT_BILLING_DETAIL_EXPORT_LIMIT": CHAT_BILLING_DETAIL_EXPORT_LIMIT,
    }
    exec(compile(ast.Module(body=consts + selected, type_ignores=[]), str(BOT_PATH), "exec"), ns)
    return ns


def _aggregates(chats, answered, answered_sl, reply_sum, inner_sum, inner_replied,
                rating_sum, rated):
    return (chats, answered, answered_sl, reply_sum, inner_sum, inner_replied, rating_sum, rated)


class _RowsCursor:
    """Отдаёт заданные строки и помнит запросы — чтобы проверить и SQL, и разбор."""

    def __init__(self, rows):
        self.rows = rows
        self.calls = []

    def execute(self, sql, params=None):
        self.calls.append((" ".join(str(sql).split()), list(params or [])))

    def fetchall(self):
        return list(self.rows)

    def fetchone(self):
        return (len(self.rows), 99)


class _RowsDb:
    def __init__(self, rows):
        self.cursor = _RowsCursor(rows)

    @contextmanager
    def _get_cursor(self):
        yield self.cursor


GROUPING_ROWS = [
    (date(2026, 9, 16), 9, "Jana Taxi", *_aggregates(10, 9, 7, 540.0, 1200.0, 8, 13.0, 3)),
    (date(2026, 9, 16), 9, "Ноль такси", *_aggregates(4, 4, 4, 100.0, 300.0, 4, 10.0, 2)),
    (date(2026, 9, 16), 23, "Jana Taxi", *_aggregates(2, 1, 0, 200.0, 0.0, 0, 0.0, 0)),
    (date(2026, 9, 17), 0, "Jana Taxi", *_aggregates(1, 1, 1, 30.0, 60.0, 1, 5.0, 1)),
]


def _park_report():
    db = _RowsDb([
        (date(2026, 9, 16), "Jana Taxi", "", *_aggregates(60, 58, 50, 3480.0, 9000.0, 50, 90.0, 20)),
        (date(2026, 9, 16), "Ноль такси", "", *_aggregates(30, 30, 25, 900.0, 3000.0, 25, 36.0, 9)),
        (date(2026, 9, 17), "Jana Taxi", "", *_aggregates(40, 40, 38, 400.0, 2000.0, 20, 42.0, 10)),
    ])
    return get_chat_billing_report(db, date(2026, 9, 16), date(2026, 9, 17))


class ChatBillingMetricsTests(unittest.TestCase):
    def test_every_cut_parses_the_same_eight_aggregates(self):
        db = _RowsDb([(date(2026, 9, 16), "Jana Taxi", "",
                       *_aggregates(10, 9, 7, 540.4, 1200.6, 8, 13.0, 3))])
        report = get_chat_billing_report(db, date(2026, 9, 16), date(2026, 9, 16), sl_seconds=45)
        self.assertEqual(report["totals"], {
            "chats": 10, "answered": 9, "no_reply": 1, "answered_sl": 7,
            "first_reply_seconds": 540, "inner_reply_seconds": 1201, "inner_replied": 8,
            "rating_sum": 13.0, "rated": 3,
        })
        sql, params = db.cursor.calls[0]
        # Порог SL — первый параметр: он стоит в SELECT раньше условий WHERE.
        self.assertEqual(params[0], 45)
        for fragment in ("SUM(r.average_replies_time)", "COUNT(r.average_replies_time)",
                         "SUM(r.rating_score)", "COUNT(r.rating_score)"):
            self.assertIn(fragment, sql)
        # «Факт чатов» — обращения без автоматических опросов после чата: ручной отчёт
        # считал и строки request_type='rating', и завышал объём примерно вдвое.
        self.assertIn("r.request_type = %s", sql)
        self.assertIn("common", params)

        operators = get_chat_billing_operators(_RowsDb([
            (date(2026, 9, 16), "Асель", *_aggregates(5, 5, 4, 100.0, 250.0, 5, 24.0, 5), 30, 20),
        ]), date(2026, 9, 16), date(2026, 9, 16))
        self.assertEqual(operators["totals"]["rated"], 5)
        self.assertEqual(operators["totals"]["inner_reply_seconds"], 250)
        self.assertEqual(operators["totals"]["incoming_messages"], 30)
        self.assertEqual(operators["totals"]["outgoing_messages"], 20)

    def test_detail_row_carries_number_reply_time_and_rating(self):
        started = datetime(2026, 9, 16, 23, 59, 21)
        db = _RowsDb([(1, started, "Salam taxi", "whatsapp", "Бауржан", "77076825368",
                       "Ерланов Темирлан", 24, 37.5, 5.0, 6, 4)])
        row = get_chat_billing_details(db, date(2026, 9, 16), date(2026, 9, 16))["rows"][0]
        self.assertEqual(row["client"], "Бауржан")
        self.assertEqual(row["client_number"], "77076825368")
        self.assertEqual(row["inner_reply_seconds"], 37.5)
        self.assertEqual(row["rating"], 5.0)

    def test_metric_keys_are_the_ones_the_screen_reads(self):
        source = VIEW_PATH.read_text(encoding="utf-8") + METRICS_JS_PATH.read_text(encoding="utf-8")
        for key in ("inner_reply_seconds", "inner_replied", "rating_sum", "rated",
                    "plan_chats", "staff_planned", "staff_fact", "planned_hours", "fact_hours"):
            self.assertIn(key, source, f"витрина не читает {key}")
        for key in ("inner_reply_seconds", "inner_replied", "rating_sum", "rated"):
            self.assertIn(key, CHAT_BILLING_METRICS)


class ChatBillingGroupingTests(unittest.TestCase):
    def test_hours_cover_the_window_and_totals_match(self):
        report = build_chat_billing_grouping(GROUPING_ROWS)
        self.assertEqual([day["date"] for day in report["days"]], ["2026-09-16", "2026-09-17"])
        self.assertEqual([len(day["hours"]) for day in report["days"]], [24, 24])
        nine = report["days"][0]["hours"][9]
        self.assertEqual((nine["hour"], nine["chats"], nine["rated"]), (9, 14, 5))
        # Пустой час — нули, а не пропуск строки.
        self.assertEqual(report["days"][0]["hours"][10]["chats"], 0)
        day_sum = sum(day["totals"]["chats"] for day in report["days"])
        self.assertEqual(day_sum, report["totals"]["chats"])
        self.assertEqual(report["totals"]["chats"], 17)

    def test_park_filter_keeps_the_full_park_list(self):
        report = build_chat_billing_grouping(GROUPING_ROWS, park="Ноль такси")
        self.assertEqual(report["park"], "Ноль такси")
        self.assertEqual([item["park"] for item in report["parks"]], ["Jana Taxi", "Ноль такси"])
        self.assertEqual(report["totals"]["chats"], 4)
        self.assertEqual([day["date"] for day in report["days"]], ["2026-09-16"])

    def test_time_window_trims_hours(self):
        report = build_chat_billing_grouping(GROUPING_ROWS, minute_from=8 * 60, minute_to=22 * 60 + 59)
        self.assertEqual([item["hour"] for item in report["days"][0]["hours"]], list(range(8, 23)))
        self.assertEqual(report["totals"]["chats"], 14)

    def test_one_query_serves_every_park_sheet(self):
        db = _RowsDb(GROUPING_ROWS)
        reports = get_chat_billing_grouping_by_park(db, date(2026, 9, 16), date(2026, 9, 17))
        self.assertEqual([park for park, _ in reports], [None, "Jana Taxi", "Ноль такси"])
        self.assertEqual(len(db.cursor.calls), 1)
        self.assertEqual(sum(report["totals"]["chats"] for park, report in reports[1:]),
                         reports[0][1]["totals"]["chats"])
        single = get_chat_billing_grouping_by_park(_RowsDb(GROUPING_ROWS), date(2026, 9, 16),
                                                   date(2026, 9, 17), park="Jana Taxi")
        self.assertEqual([park for park, _ in single], ["Jana Taxi"])
        db = _RowsDb([])
        get_chat_billing_grouping(db, date(2026, 9, 16), date(2026, 9, 16))
        self.assertIn("EXTRACT(HOUR FROM r.request_start)", db.cursor.calls[0][0])


def _shift(operator_id, day, start, end):
    """Строка графика так, как её отдаёт psycopg2: дата и время объектами."""
    return (operator_id, date.fromisoformat(day), time.fromisoformat(start), time.fromisoformat(end))


def _phone(operator_id, day, start, end, piece_start=None, piece_end=None):
    return (*_shift(operator_id, day, start, end),
            time.fromisoformat(piece_start or start), time.fromisoformat(piece_end or end))


class _SequenceCursor(_RowsCursor):
    """Каждый запрос получает свою пачку строк — по порядку вызовов."""

    def __init__(self, batches):
        super().__init__([])
        self.batches = list(batches)

    def fetchall(self):
        return list(self.batches.pop(0))


class ChatBillingStaffTests(unittest.TestCase):
    def test_planned_staff_is_the_schedule_without_phone_time(self):
        plan = build_chat_billing_planned_staff(date(2026, 9, 12), date(2026, 9, 13), [
            _shift(1, "2026-09-12", "09:00", "18:00"),
            # Наложенная смена того же человека — часы не задваиваются, голова одна.
            _shift(1, "2026-09-12", "15:00", "21:30"),
            # Хвост ночной смены прошлых суток — утро первого дня периода.
            _shift(4, "2026-09-11", "20:00", "02:00"),
            # Смена, слитая публикацией из телефонной 08–17 и чатовой 17–01.
            _shift(5, "2026-09-12", "08:00", "01:00"),
            # Кусок телефонной смены после полуночи вычитается из следующих суток.
            _shift(6, "2026-09-12", "17:00", "02:00"),
        ], [
            _phone(5, "2026-09-12", "08:00", "17:00"),
            _phone(6, "2026-09-12", "17:00", "01:00", "00:00", "01:00"),
            # Телефон взял человек, которого в графике нет, — чужие смены не трогает.
            _phone(9, "2026-09-12", "09:00", "18:00"),
        ])
        day = plan["2026-09-12"]
        self.assertEqual(day["heads"][:3], [1, 1, 0])
        self.assertEqual(day["heads"][8:10], [0, 1])
        self.assertEqual(day["heads"][16:18], [1, 3])
        # В 21:00 первый ещё на смене (до 21:30), дальше — пятый и шестой.
        self.assertEqual(day["heads"][21:], [3, 2, 2])
        # 2 ч ночного хвоста, 12,5 ч наложенных смен, 7 ч чата у пятого, 7 ч у шестого.
        self.assertEqual(sum(day["hours"]), 2 + 12.5 + 7 + 7)
        next_day = plan["2026-09-13"]
        self.assertEqual(next_day["heads"][:3], [1, 1, 0])
        self.assertEqual(sum(next_day["hours"]), 2)
        self.assertEqual(sum(next_day["heads"]), 2)
        # В графике никого — прочерк, а не ноль людей.
        empty = build_chat_billing_planned_staff(date(2026, 9, 12), date(2026, 9, 12), [], [])
        self.assertEqual(empty, {"2026-09-12": None})

    def test_plan_reads_regular_shifts_of_chat_people_from_the_day_before(self):
        db = _RowsDb([])
        with db._get_cursor() as cursor:
            _chat_billing_planned_shifts_tx(cursor, date(2026, 9, 21), date(2026, 9, 27))
        sql, params = db.cursor.calls[0]
        self.assertIn("FROM work_shifts ws", sql)
        self.assertIn("d.name ILIKE %s", sql)
        self.assertIn("ws.shift_type = %s", sql)
        # Сутки раньше периода — ради ночной смены; телефоны и практика в план не идут.
        self.assertEqual(params, ["%чат%", "regular", date(2026, 9, 20), date(2026, 9, 27)])

    def test_phone_time_comes_from_whole_lots_and_pieces_of_the_chat_auction(self):
        db = _RowsDb([])
        with db._get_cursor() as cursor:
            _chat_billing_phone_claims_tx(cursor, date(2026, 9, 21), date(2026, 9, 27))
        sql, params = db.cursor.calls[0]
        self.assertIn("FROM shift_auction_test_lots l", sql)
        self.assertIn("FROM shift_auction_historical_claims hc", sql)
        self.assertIn("sh.meta->>'shiftKind' = %s", sql)
        # Целый лот с кусками не берётся: его claimed_by — лишь последний дозабравший.
        self.assertIn("NOT EXISTS", sql)
        self.assertEqual(params, ["chat", "phone", date(2026, 9, 20), date(2026, 9, 27)] * 2)

    def test_staff_is_three_queries_on_one_cursor(self):
        cursor = _SequenceCursor([
            [(date(2026, 9, 21), 9, 1, 3600.0)],
            [_shift(1, "2026-09-21", "08:00", "17:00"), _shift(2, "2026-09-21", "08:00", "17:00")],
            [_phone(2, "2026-09-21", "08:00", "17:00")],
        ])

        class _Db:
            @contextmanager
            def _get_cursor(self):
                yield cursor

        staff = get_chat_billing_staff(_Db(), date(2026, 9, 21), date(2026, 9, 21))
        self.assertEqual(len(cursor.calls), 3)
        day = staff["2026-09-21"]
        self.assertEqual(day["planned_heads"][9], 1)
        self.assertEqual(sum(day["planned_hours"]), 9)
        self.assertEqual(day["fact_heads"][9], 1)

    def test_fact_heads_need_a_minute_online(self):
        rows = [
            (date(2026, 9, 12), 9, 1, 3600.0),
            (date(2026, 9, 12), 9, 2, 59.0),
            (date(2026, 9, 12), 10, 2, 60.0),
            # Строка соседних суток из ночной смены — не наша.
            (date(2026, 9, 11), 23, 1, 3600.0),
        ]
        staff = build_chat_billing_staff(date(2026, 9, 12), date(2026, 9, 13), rows,
                                         {"2026-09-12": {"heads": [2] * 24, "hours": [2.0] * 24}})
        day = staff["2026-09-12"]
        self.assertEqual(day["fact_heads"][9], 1)
        self.assertEqual(day["fact_heads"][10], 1)
        self.assertAlmostEqual(sum(day["fact_hours"]), (3600 + 59 + 60) / 3600, places=3)
        # Нет ни статусов, ни графика — прочерк, а не ноль людей.
        self.assertIsNone(staff["2026-09-13"]["fact_heads"])
        self.assertIsNone(staff["2026-09-13"]["planned_heads"])

    def test_plan_is_park_share_times_same_day_last_week(self):
        staff = {"2026-09-16": {"planned_hours": [2.0] * 24, "fact_hours": [1.5] * 24}}
        report = attach_chat_billing_plan(_park_report(), {"2026-09-16": 1000}, staff)
        day = report["days"][0]
        fixed = [name for name, _ in CHAT_BILLING_PARK_SHARES]
        # Сначала парки с долей в порядке файла, за ними парк без доли.
        self.assertEqual([item["park"] for item in day["parks"]], fixed + ["Ноль такси"])
        by_park = {item["park"]: item for item in day["parks"]}
        self.assertEqual(by_park["Jana Taxi"]["plan_share"], 0.1469962424736294)
        self.assertAlmostEqual(by_park["Jana Taxi"]["plan_chats"], 146.9962424736294)
        # Парк с долей без обращений остаётся строкой с планом и нулевым фактом, как в файле.
        self.assertEqual(by_park["Tenge Taxi"]["chats"], 0)
        self.assertAlmostEqual(by_park["Tenge Taxi"]["plan_chats"], 69.03888813436553)
        # У парка без доли плана нет, факт на месте.
        self.assertEqual([by_park["Ноль такси"][key] for key in ("plan_share", "plan_chats", "chats")],
                         [None, None, 30])
        self.assertEqual(day["totals"]["plan_base_chats"], 1000)
        # Доли файла в сумме 0,817 — план дня ниже прошлой недели, как в файле.
        self.assertAlmostEqual(day["totals"]["plan_chats"], 816.8319072841684, places=9)
        self.assertEqual(day["staff"], {"planned_hours": 48.0, "fact_hours": 36.0})
        # Второй день: чатов неделей раньше в базе нет — плана нет, доли на месте.
        second = report["days"][1]
        self.assertEqual([item["park"] for item in second["parks"]], fixed + ["Ноль такси"])
        self.assertIsNone(second["totals"]["plan_chats"])
        self.assertTrue(all(item["plan_chats"] is None for item in second["parks"]))
        self.assertEqual(second["parks"][0]["plan_share"], 0.30798134818235323)
        self.assertEqual(report["days"][1]["staff"], {"planned_hours": None, "fact_hours": None})
        self.assertEqual(report["totals"]["plan_chats"], 816.83)
        self.assertEqual(report["parks"][1]["plan_chats"], 147.0)
        self.assertEqual(report["staff"], {"planned_hours": 48.0, "fact_hours": 36.0})

    def test_plan_matches_the_sample_file(self):
        # «Ежедневный отчет Техподдержка чат», 01.09.2026: в «Нужно удалить» итога дня 1457,
        # в «План (чатов)» Excel посчитал 1190,1240889130333.
        report = attach_chat_billing_plan(_park_report(), {"2026-09-16": 1457}, {})
        self.assertEqual(report["days"][0]["totals"]["plan_chats"], 1190.1240889130333)

    def test_plan_base_is_the_same_day_last_week_in_the_same_window(self):
        db = _RowsDb([(date(2026, 9, 9), 1636), (date(2026, 9, 10), 1420)])
        with db._get_cursor() as cursor:
            base = _chat_billing_plan_base_tx(cursor, date(2026, 9, 16), date(2026, 9, 17),
                                              8 * 60, 19 * 60 + 59)
        self.assertEqual(base, {"2026-09-16": 1636, "2026-09-17": 1420})
        sql, params = db.cursor.calls[0]
        # Тот же WHERE, что у факта: без опросов после чата и в том же окне времени.
        self.assertIn("r.request_type = %s", sql)
        self.assertIn("EXTRACT(HOUR FROM r.request_start)", sql)
        self.assertEqual(params, ["common", date(2026, 9, 9), date(2026, 9, 10), 480, 1199])

    def test_grouping_staff_is_per_hour(self):
        report = build_chat_billing_grouping(GROUPING_ROWS)
        staff = {"2026-09-16": {"planned_heads": list(range(24)), "fact_heads": [3] * 24,
                                "planned_hours": [1.0] * 24, "fact_hours": [0.5] * 24}}
        attach_chat_billing_grouping_staff(report, staff)
        nine = report["days"][0]["hours"][9]
        self.assertEqual((nine["staff_planned"], nine["staff_fact"]), (9, 3))
        self.assertIsNone(report["days"][1]["hours"][0]["staff_planned"])
        self.assertEqual(report["days"][0]["staff"], {"planned_hours": 24.0, "fact_hours": 12.0})


class ChatBillingExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ns = _namespace()
        cls.params = {
            "start_day": date(2026, 9, 16), "end_day": date(2026, 9, 17),
            "minute_from": 0, "minute_to": 1439, "sl_seconds": 60,
        }

    def _load(self, workbook):
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return load_workbook(output)

    def test_minutes_round_half_up_like_the_screen(self):
        # Те же значения проверяет tests/chat_billing_metrics.test.mjs.
        minutes = self.ns["_chat_billing_minutes"]
        cases = {135: 2.3, 150: 2.5, 420: 7.0, 720: 12.0, 3: 0.1, 2: 0.0, 1047: 17.5}
        for seconds, expected in cases.items():
            self.assertEqual(minutes(seconds), expected, seconds)
        self.assertIsNone(minutes(None))
        self.assertEqual(self.ns["_chat_billing_round_half_up"](4.25, 1), 4.3)
        self.assertEqual(self.ns["_chat_billing_minutes_label"](288), "4,8")
        self.assertEqual(self.ns["_chat_billing_minutes_label"](420), "7")

    def test_operator_export_follows_the_screen_columns(self):
        titles = [title for _, title, _, _ in self.ns["_CHAT_BILLING_EXPORT_SUMMARY_COLUMNS"]]
        source = VIEW_PATH.read_text(encoding="utf-8")
        block = source[source.index("const CHAT_BILLING_COLUMN_SETS = {"):]
        block = block[block.index("  operator: ["):]
        block = block[:block.index("\n  ],")]
        self.assertEqual(re.findall(r"\['[a-z_]+', '([^']+)'\]", block), titles)
        values = self.ns["_chat_billing_export_metrics"]({
            "chats": 10, "answered": 8, "no_reply": 2, "answered_sl": 5,
            "first_reply_seconds": 1200, "inner_reply_seconds": 900, "inner_replied": 6,
            "rating_sum": 13.0, "rated": 3,
        })
        self.assertEqual(values, [10, 8, 2, 2.5, 0.5, 2.5, 4.3])

    def test_detail_values_are_minutes_and_keep_the_number(self):
        titles = [title for _, title, _ in self.ns["_CHAT_BILLING_EXPORT_DETAIL_COLUMNS"]]
        row = {"started_at": "2026-09-16 10:00:00", "park": "Jana Taxi", "client_number": "77000000108",
               "client": "Dimash", "operator": "Асель", "first_reply_seconds": 150,
               "answered_sl": 0, "inner_reply_seconds": 45, "rating": 5.0,
               "incoming_messages": 3, "outgoing_messages": 2}
        values = dict(zip(titles, self.ns["_chat_billing_export_detail_values"](row)))
        self.assertEqual(values["Номер клиента"], "77000000108")
        self.assertEqual(values["Реакция на 1 сообщение, мин"], 2.5)
        self.assertEqual(values["Время ответа внутри чата, мин"], 0.8)
        self.assertEqual(values["Оценка"], 5.0)

    def test_sheet_titles_are_valid_and_unique(self):
        title = self.ns["_chat_billing_sheet_title"]
        used = set()
        self.assertEqual(title("Такси24/Нур: [VIP]*", used), "Такси24 Нур   VIP")
        self.assertEqual(title("A" * 40, used), "A" * 31)
        self.assertEqual(title("a" * 40, used), "a" * 27 + " (2)")
        self.assertEqual(title("", used), "Без парка")

    def _daily_report(self):
        staff = {"2026-09-16": {"planned_hours": [2.5] * 24, "fact_hours": [2.0] * 24}}
        return attach_chat_billing_plan(_park_report(), {"2026-09-16": 1000}, staff)

    def _daily_books(self, report):
        """Книга дважды: с формулами и со значениями, сохранёнными при сборке."""
        content = self.ns["_chat_billing_daily_workbook"](self.params, report).getvalue()
        return load_workbook(BytesIO(content)), load_workbook(BytesIO(content), data_only=True)

    def test_daily_workbook_follows_the_szov_sample(self):
        formulas, values = self._daily_books(self._daily_report())
        self.assertEqual(values.sheetnames, ["Сентябрь 2026"])
        ws = values.active
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        self.assertEqual(rows[0], [
            "Дата", "Название очереди", "Нужно удалить", "План (чатов)", "факт (чатов)",
            "% совпадения прогноза",
            "Среднее время реакции на 1 сообщение оператора (в минутах) в разрезе парков",
            "Среднее время ответа внутри чата (в минутах) в разрезе парков",
            "Сред оценка водителей в разрезе парков", "План по часам", "Факт по часам",
            "% Пере/Недоработки", "Комментарии",
        ])
        # Заголовок служебной колонки красный, как в файле.
        self.assertEqual(ws["C1"].font.color.rgb[-6:], "FF0000")
        # Блок дня: 14 парков с долей, парк без доли, итог в 17-й строке.
        self.assertEqual([row[1] for row in rows[1:16]],
                         [name for name, _ in CHAT_BILLING_PARK_SHARES] + ["Ноль такси"])
        self.assertEqual(rows[1][0], datetime(2026, 9, 16))
        jana = rows[2]
        self.assertEqual(jana[1:3], ["Jana Taxi", 0.1469962424736294])
        self.assertAlmostEqual(jana[3], 146.9962424736294)
        self.assertEqual(jana[4], 60)
        self.assertAlmostEqual(jana[5], 60 / 146.9962424736294)
        # 3480 сек на 58 отвеченных = 1 мин; 9000 на 50 = 3 мин; 90 на 20 = 4,5.
        self.assertEqual(jana[6:9], [1.0, 3.0, 4.5])
        self.assertEqual(ws["C3"].number_format, "0.0")
        self.assertEqual(rows[1][9:12], [60.0, 48.0, -0.2])
        self.assertIn("Ежедневный отчет по чатам за 16.09.2026\nПлан чатов: 817\nФакт чатов: 90",
                      rows[1][12])
        tenge = rows[4]
        self.assertEqual(tenge[1:3], ["Tenge Taxi", 0.06903888813436553])
        self.assertEqual((tenge[4], tenge[5]), (None, 0))
        zero = rows[15]
        self.assertEqual(zero[1:6], ["Ноль такси", None, None, 30, None])
        total = rows[16]
        self.assertEqual(total[:3], [datetime(2026, 9, 16), None, 1000])
        self.assertEqual(ws["C17"].number_format, "0")
        self.assertAlmostEqual(total[3], 816.8319072841684, places=9)
        self.assertEqual(total[4], 90)
        self.assertAlmostEqual(total[5], 90 / 816.8319072841684)
        # Итог дня — по обращениям всех парков: (3480+900) / 88 отвеченных.
        self.assertEqual(total[6], round((3480 + 900) / 88 / 60, 1))
        merged = [str(item) for item in ws.merged_cells.ranges]
        self.assertIn("A2:A16", merged)
        self.assertIn("M2:M16", merged)
        self.assertEqual(ws["A17"].fill.fgColor.rgb[-6:], "9DC3E6")
        self.assertEqual(ws["D17"].fill.fgColor.rgb[-6:], "9DC3E6")

        # В ячейках — формулы файла: правка «Нужно удалить» пересчитывает план.
        wf = formulas.active
        self.assertEqual(wf["D3"].value, "=C3*$C$17")
        self.assertEqual(wf["F3"].value, "=E3/D3")
        self.assertEqual(wf["D17"].value, "=SUM(D2:D16)")
        self.assertEqual(wf["F17"].value, "=E17/D17")
        self.assertEqual(wf["C17"].value, 1000)
        self.assertIsNone(wf["D16"].value)

        # Второй день: чатов неделей раньше нет — доли стоят, плана и формул нет.
        second = rows[17]
        self.assertEqual(second[:2], [datetime(2026, 9, 17), "Техподдержка iTaxi"])
        # xlsxwriter пишет число 16 значащими цифрами — 17-й знак доли файла теряется.
        self.assertAlmostEqual(second[2], 0.30798134818235323, places=15)
        self.assertEqual(second[3], None)
        self.assertIsNone(wf["D18"].value)
        self.assertEqual(rows[32][2:6], [None, None, 40, None])
        self.assertIn("План чатов: —", rows[17][12])
        # Плана часов нет — пусто, а не ноль.
        self.assertEqual(second[9:12], [None, None, None])

    def test_daily_workbook_marks_low_rating(self):
        report = self._daily_report()
        report["days"][0]["parks"][-1]["rating_sum"] = 26.0  # Ноль такси: 26 / 9 = 2,9
        ws = self._daily_books(report)[1].active
        self.assertEqual(ws["I16"].value, 2.9)
        self.assertEqual(ws["I16"].font.color.rgb[-6:], "9C0006")
        self.assertNotEqual(str(getattr(ws["I3"].font.color, "rgb", ""))[-6:], "9C0006")

    def test_daily_workbook_keeps_park_names_as_text(self):
        report = self._daily_report()
        report["days"][0]["parks"][-1]["park"] = "=HYPERLINK(\"http://x\")"
        ws = self._daily_books(report)[0].active
        self.assertEqual(ws["B16"].data_type, "s")

    def test_hourly_workbook_follows_the_szov_sample(self):
        overall = build_chat_billing_grouping(GROUPING_ROWS)
        attach_chat_billing_grouping_staff(overall, {
            "2026-09-16": {"planned_heads": [2] * 24, "fact_heads": [1] * 9 + [2] + [1] * 14,
                           "planned_hours": [2.0] * 24, "fact_hours": [1.0] * 24},
        })
        reports = [(None, overall)] + [
            (item["park"], build_chat_billing_grouping(GROUPING_ROWS, park=item["park"]))
            for item in overall["parks"]]
        workbook = self._load(self.ns["_chat_billing_hourly_workbook"](self.params, reports))
        self.assertEqual(workbook.sheetnames, ["Все таксопарки", "Jana Taxi", "Ноль такси"])

        ws = workbook["Все таксопарки"]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        self.assertEqual(rows[0][1], datetime(2026, 9, 16))
        self.assertEqual(rows[0][7], datetime(2026, 9, 17))
        self.assertEqual(rows[1][:7], ["Время", "Поступившие чаты", "План(По сотрудникам)",
                                       "Факт(По сотрудникам)", "Разница",
                                       "Ср вр реакции на 1 сообщение", "Ср время ответа внутри чата"])
        self.assertIn("B1:G1", [str(item) for item in ws.merged_cells.ranges])
        nine = rows[2 + 9]
        # 640 сек на 13 отвеченных = 0,8 мин; 1500 сек на 12 ответивших = 2,1 мин.
        self.assertEqual(nine[:7], [9, 14, 2, 2, 0, 0.8, 2.1])
        eight = rows[2 + 8]
        self.assertEqual(eight[:5], [8, 0, 2, 1, -1])
        self.assertEqual(ws["E11"].fill.fgColor.rgb[-6:], "FF0000")
        self.assertIsNone(ws["E12"].fill.fgColor.rgb if ws["E12"].fill.fill_type else None)
        # Второй день без графика и статусов — прочерков нет, ячейки пустые.
        self.assertEqual(rows[2][8:11], [None, None, None])
        total = rows[26]
        self.assertEqual(total[:2], ["Общий итог", 16])

        park_rows = [list(row) for row in workbook["Jana Taxi"].iter_rows(values_only=True)]
        # У парка сотрудников нет — три колонки на день.
        self.assertEqual(park_rows[1][:4], ["Время", "Поступившие чаты", "Ср вр реакции на 1 сообщение",
                                            "Ср время ответа внутри чата"])
        self.assertEqual(park_rows[0][4], datetime(2026, 9, 17))

    def test_export_route_uses_the_szov_formats(self):
        backend = BOT_PATH.read_text(encoding="utf-8")
        block = backend[backend.index("def api_resource_fte_chat_billing_export"):]
        block = block[:block.index("@app.route('/api/resource_fte/chat/day/")]
        self.assertIn("mode not in ('park', 'operator', 'detail', 'grouping')", block)
        self.assertIn("_chat_billing_hourly_workbook(params, reports)", block)
        self.assertIn("_chat_billing_daily_workbook(params, get_chat_billing_daily(", block)
        self.assertIn("attach_chat_billing_grouping_staff(", block)
        self.assertIn("park=_chat_billing_park_arg()", block)
        # В модуле нет `import io` — только `from io import BytesIO`. Выгрузка биллинга
        # чата падала NameError на io.BytesIO() в любом разрезе с самого появления.
        self.assertNotIn("io.BytesIO", block)
        self.assertIn("output = BytesIO()", block)
        self.assertIn("\nfrom io import BytesIO", backend)
        route = backend[backend.index("def api_resource_fte_chat_billing_grouping"):]
        route = route[:route.index("# Выгрузка биллинга чата.")]
        self.assertIn("if not park:", route)
        self.assertIn("get_chat_billing_daily(db,", backend[backend.index("def api_resource_fte_chat_billing():"):])


if __name__ == "__main__":
    unittest.main()
