# -*- coding: utf-8 -*-
"""Калибровка «Журнала оценок» — руководителям и супервайзерам.

Просьба владельца 27.10.2026: «в журнале оценок калибровка на данный момент
исчезла, добавь её для руководителей и супервайзеров». С 16.06.2026 вкладку
видели только глобальный админ и СВ, а все главы отделов на проде — админы,
поэтому у руководителей она и пропала.

Ручки берём настоящие, из bot_schedule2.py, и гоняем на двойниках базы и
запроса: импортировать монолит в тестах нельзя (см. tests/source_cache.py).
Главу проверяем и с базовой ролью оператора или тренера — руководитель
определяется назначением, а не ролью. Ловушку try/except в ручках обходим
точными кодами ответа: NameError внутри ручки дал бы 500, а не 200.
"""

import ast
import copy
import json
import logging
import re
import unittest
from contextlib import contextmanager
from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tests import source_cache

ROOT = Path(__file__).resolve().parents[1]
BOT_PATH = ROOT / "bot_schedule2.py"
JOURNAL_PATH = ROOT / "src" / "call_evaluation" / "main.jsx"

CONSTANTS = ("ROLE_HIERARCHY", "CALIBRATION_FORBIDDEN_ERROR")
FUNCTIONS = (
    "_normalize_user_role",
    "_get_role_level",
    "_has_min_role",
    "_is_admin_role",
    "_is_super_admin_role",
    "_is_supervisor_role",
    "_headed_department_id",
    "_is_global_admin_requester",
    "_calibration_role",
    "_normalize_calibration_score_value",
    "_calibration_label_score",
    "_build_calibration_results",
    "list_call_calibration_rooms",
    "create_call_calibration_room",
    "update_calibration_room",
    "add_call_to_calibration_room",
    "get_call_calibration_room",
    "export_call_calibration_room_excel",
    "update_call_calibration_etalon",
    "update_call_calibration_general_comment",
    "get_calibration_room_history",
    "join_call_calibration_room",
    "evaluate_call_calibration_room",
    "delete_calibration_room",
    "delete_call_from_calibration_room",
)

# Люди стенда: id → (роль, возглавляемый отдел или None).
ADMIN = 169            # админ без отдела — глобальный
ADMIN_HEAD = 1         # глава СЗоВ, роль admin — как все главы на проде
OPERATOR_HEAD = 700    # глава отдела с базовой ролью оператора
TRAINER_HEAD = 701     # глава отдела с базовой ролью тренера
SV = 55
SV_HEAD = 702          # глава отдела с базовой ролью СВ
OPERATOR = 126
TRAINER = 206
OBSERVER = 474         # наблюдатель «Маркетинга»: журнал смотрит, калибровку — нет
CALL_OPERATOR = 330    # оператор, чей звонок в комнате

PEOPLE = {
    ADMIN: ("admin", None),
    ADMIN_HEAD: ("admin", 1),
    OPERATOR_HEAD: ("operator", 2008),
    TRAINER_HEAD: ("trainer", 1933),
    SV: ("sv", None),
    SV_HEAD: ("sv", 367),
    OPERATOR: ("operator", None),
    TRAINER: ("trainer", None),
    OBSERVER: ("marketing_manager", None),
    CALL_OPERATOR: ("operator", None),
    448: ("super_admin", None),
}

ROOM_ID = 10
ROOM_CREATOR = 202     # СВ, создавший комнату (не участник тестов)


def _norm(sql):
    return " ".join(str(sql).split())


class _Cursor:
    """Отвечает на SQL по первому подходящему правилу (подстрока → строки)."""

    def __init__(self, rules):
        self.rules = list(rules)
        self.executed = []
        self._rows = []

    def execute(self, sql, params=None):
        text = _norm(sql)
        self.executed.append(text)
        for needle, rows in self.rules:
            if needle in text:
                self._rows = list(rows)
                return
        self._rows = []

    def fetchone(self):
        return self._rows[0] if self._rows else None

    def fetchall(self):
        return list(self._rows)

    def ran(self, needle):
        return any(needle in sql for sql in self.executed)


class _DB:
    def __init__(self, rules=()):
        self.cursor = _Cursor(rules)

    def get_user(self, *, id):
        person = PEOPLE.get(int(id))
        if person is None:
            return None
        return (int(id), None, f"User {id}", person[0])

    def headed_department_id_for_user(self, user_id):
        person = PEOPLE.get(int(user_id))
        return person[1] if person else None

    @contextmanager
    def _get_cursor(self):
        yield self.cursor


class _Request:
    def __init__(self, *, args=None, body=None):
        self.args = dict(args or {})
        self.form = {}
        self.files = {}
        self.method = "GET"
        self._body = body

    def get_json(self, silent=False):
        return self._body


def _load():
    module = source_cache.tree(BOT_PATH)
    nodes = [
        copy.deepcopy(node)
        for node in module.body
        if isinstance(node, ast.Assign)
        and any(isinstance(t, ast.Name) and t.id in CONSTANTS for t in node.targets)
    ]
    for name in FUNCTIONS:
        node = source_cache.function_copy(BOT_PATH, name)
        node.decorator_list = []
        nodes.append(node)
    namespace = {
        "json": json,
        "re": re,
        "logging": logging,
        "datetime": datetime,
        "BytesIO": BytesIO,
        "Workbook": Workbook,
        "PatternFill": PatternFill,
        "Font": Font,
        "Alignment": Alignment,
        "Border": Border,
        "Side": Side,
        "get_column_letter": get_column_letter,
        # Выгрузка отдаёт файл; в тесте нужен сам файл, а не ответ Flask.
        "send_file": lambda output, **kwargs: {"xlsx": output.getvalue(), "name": kwargs.get("download_name")},
        "jsonify": lambda payload: payload,
        "_compute_total_score_from_criteria": lambda criteria, scores: (0.0, []),
        "_build_signed_audio_url": lambda path: None,
    }
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(BOT_PATH), "exec"), namespace)
    return namespace


NS = _load()


def _call(route, user_id, *args, rules=(), query=None, body=None, scope_ok=True):
    """Позвать ручку от имени user_id; вернуть (статус, ответ, курсор)."""
    db = _DB(rules)
    NS["db"] = db
    NS["g"] = SimpleNamespace(user_id=user_id)
    NS["request"] = _Request(args=query, body=body)
    NS["_authorize_operator_scope"] = lambda requester, requester_id, operator_id: scope_ok
    with _quiet_route_errors():
        result = NS[route](*args)
    payload, status = result if isinstance(result, tuple) else (result, 200)
    return status, payload, db.cursor


@contextmanager
def _quiet_route_errors():
    # Ручки пишут трейсбек 500 в лог; в тестах он только шумит — код ответа и так проверяем.
    logger = logging.getLogger()
    previous = logger.level
    logger.setLevel(logging.CRITICAL)
    try:
        yield
    finally:
        logger.setLevel(previous)


ROOM_ROW = (ROOM_ID, "2026-10", "Октябрь", datetime(2026, 10, 27, 9, 0), ROOM_CREATOR, "СВ-автор")
NOT_MEMBER = ("FROM calibration_room_members", [])
MEMBER = ("FROM calibration_room_members", [(1,)])
ROOM_FOR_DETAIL = ("WHERE r.id = %s", [ROOM_ROW])
ROOM_FOR_UPDATE = ("SELECT id, created_by_admin_id FROM calibration_rooms", [(ROOM_ID, ROOM_CREATOR)])
CALL_LIST = ("AS evaluated_count", [
    (5, CALL_OPERATOR, "Оператор", "+77001112233", None, 9.5, None, 87, "Линия", 2, False),
])
SELECTED_CALL = ("c.general_comment_updated_at", [(
    5, ROOM_ID, "2026-10", "+77001112233", None, 9.5, "", None,
    ["Correct"], [""], None, None, CALL_OPERATOR, "Оператор", 87, "Линия",
    [{"name": "Приветствие"}], False, ["Correct"], [""], None, None, None, "", None, None, None,
)])
# Выгрузка: звонки комнаты и оценки по ним (схема — как в SQL ручки).
EXPORT_CALLS = ("d.criteria FROM calibration_room_calls c", [(
    5, CALL_OPERATOR, "Оператор", "+77001112233", None, 9.5, "Линия",
    ["Correct"], [""], ["Correct"], [""], None, None, [{"name": "Приветствие"}],
)])
EXPORT_EVALUATIONS = ("WHERE e.room_call_id = ANY(%s)", [
    (5, SV, "СВ-участник", 9.5, "", ["Correct"], [""]),
])
ETALON_CALL = ("SELECT d.criteria, c.operator_id, r.created_by_admin_id",
               [([{"name": "Приветствие"}], CALL_OPERATOR, ROOM_CREATOR)])
COMMENT_CALL = ("SELECT c.operator_id, r.created_by_admin_id", [(CALL_OPERATOR, ROOM_CREATOR)])
ADD_CALL_RULES = (
    ("SELECT id, month, created_by_admin_id FROM calibration_rooms", [(ROOM_ID, "2026-10", ROOM_CREATOR)]),
    ("FROM directions", [(87, [{"name": "Приветствие"}], False)]),
    ("INSERT INTO calibration_room_calls", [(5,)]),
)
ADD_CALL_BODY = {
    "operator_id": CALL_OPERATOR, "phone_number": "+77001112233",
    "appeal_date": "2026-10-20T10:00:00", "direction": 87,
    "scores": ["Correct"], "criterion_comments": [""], "month": "2026-10",
}


class CalibrationRoleTests(unittest.TestCase):
    """Кто ведущий, кто участник, кто вне калибровки."""

    def role_of(self, user_id):
        NS["db"] = _DB()
        NS["g"] = SimpleNamespace(user_id=user_id)
        return NS["_calibration_role"](user_id, NS["db"].get_user(id=user_id))

    def test_managers_lead_the_room(self):
        for user_id in (448, ADMIN, ADMIN_HEAD, OPERATOR_HEAD, TRAINER_HEAD):
            with self.subTest(user_id=user_id):
                self.assertEqual("admin", self.role_of(user_id))

    def test_supervisors_take_part_even_when_heading_a_department(self):
        self.assertEqual("sv", self.role_of(SV))
        self.assertEqual("sv", self.role_of(SV_HEAD))

    def test_everyone_else_stays_out(self):
        for user_id in (OPERATOR, TRAINER, OBSERVER, 999):
            with self.subTest(user_id=user_id):
                self.assertIsNone(self.role_of(user_id))


class CalibrationRoutesForManagersTests(unittest.TestCase):
    """Глава отдела с любой базовой ролью работает в калибровке как ведущий."""

    def test_room_list_opens_and_counts_the_head_as_inside(self):
        rows = [(ROOM_ID, "2026-10", "Октябрь", datetime(2026, 10, 27), ROOM_CREATOR, "СВ-автор",
                 None, 1, 0, 0, 0)]
        for head in (ADMIN_HEAD, OPERATOR_HEAD, TRAINER_HEAD):
            with self.subTest(head=head):
                status, payload, _ = _call(
                    "list_call_calibration_rooms", head,
                    rules=[("FROM calibration_rooms r", rows)], query={"month": "2026-10"},
                )
                self.assertEqual(200, status)
                self.assertTrue(payload["rooms"][0]["joined"])

    def test_head_creates_a_room(self):
        status, payload, cursor = _call(
            "create_call_calibration_room", OPERATOR_HEAD,
            rules=[("INSERT INTO calibration_rooms", [(ROOM_ID,)])],
            body={"month": "2026-10", "room_title": "Калибровка ОП"},
        )
        self.assertEqual(201, status)
        self.assertEqual(ROOM_ID, payload["room_id"])
        self.assertTrue(cursor.ran("INSERT INTO calibration_rooms"))

    def test_head_opens_a_room_without_joining_and_sees_results(self):
        status, payload, cursor = _call(
            "get_call_calibration_room", OPERATOR_HEAD, ROOM_ID,
            rules=[ROOM_FOR_DETAIL, NOT_MEMBER, CALL_LIST, SELECTED_CALL],
        )
        self.assertEqual(200, status)
        self.assertTrue(payload["joined"])
        self.assertTrue(payload["can_view_results"])
        self.assertTrue(payload["can_edit_call"])
        # Оценивает только участник: ведущий ставит эталон, а не оценку.
        self.assertFalse(payload["can_evaluate"])
        self.assertFalse(cursor.ran("FROM calibration_room_members"))

    def test_foreign_call_opens_read_only_for_the_head(self):
        # Список комнат общий (решение K1), но править эталон чужого отдела
        # нельзя — признак говорит журналу не рисовать редактор.
        status, payload, _ = _call(
            "get_call_calibration_room", ADMIN_HEAD, ROOM_ID,
            rules=[ROOM_FOR_DETAIL, CALL_LIST, SELECTED_CALL], scope_ok=False,
        )
        self.assertEqual(200, status)
        self.assertTrue(payload["can_view_results"])
        self.assertFalse(payload["can_edit_call"])

    def test_head_renames_a_room_he_has_not_joined(self):
        status, payload, _ = _call(
            "update_calibration_room", TRAINER_HEAD, ROOM_ID,
            rules=[ROOM_FOR_UPDATE, NOT_MEMBER, ("UPDATE calibration_rooms", [(ROOM_ID, "Новое имя")])],
            body={"room_title": "Новое имя"},
        )
        self.assertEqual(200, status)
        self.assertEqual("Новое имя", payload["room"]["room_title"])

    def test_head_adds_calls_only_for_operators_in_scope(self):
        rules = [ADD_CALL_RULES[0], NOT_MEMBER, *ADD_CALL_RULES[1:]]
        status, payload, cursor = _call(
            "add_call_to_calibration_room", OPERATOR_HEAD, ROOM_ID, rules=rules, body=ADD_CALL_BODY,
        )
        self.assertEqual(201, status)
        self.assertEqual(5, payload["room_call_id"])
        self.assertFalse(cursor.ran("FROM calibration_room_members"))
        # Граница отдела главы — та же, что у журнала: чужого оператора не добавить.
        status, _, cursor = _call(
            "add_call_to_calibration_room", OPERATOR_HEAD, ROOM_ID,
            rules=rules, body=ADD_CALL_BODY, scope_ok=False,
        )
        self.assertEqual(403, status)
        self.assertFalse(cursor.ran("INSERT INTO calibration_room_calls"))

    def test_head_edits_etalon_and_general_comment_without_joining(self):
        status, payload, cursor = _call(
            "update_call_calibration_etalon", OPERATOR_HEAD, ROOM_ID, 5,
            rules=[ETALON_CALL, NOT_MEMBER, ("UPDATE calibration_room_calls", [(5,)])],
            body={"scores": ["Incorrect"], "criterion_comments": ["не поздоровался"]},
        )
        self.assertEqual(200, status)
        self.assertEqual(5, payload["call_id"])
        self.assertFalse(cursor.ran("FROM calibration_room_members"))

        status, _, cursor = _call(
            "update_call_calibration_general_comment", TRAINER_HEAD, ROOM_ID, 5,
            rules=[COMMENT_CALL, NOT_MEMBER, ("UPDATE calibration_room_calls", [(5,)])],
            body={"general_comment": "Разобрали на планёрке"},
        )
        self.assertEqual(200, status)
        self.assertTrue(cursor.ran("UPDATE calibration_room_calls"))
        self.assertFalse(cursor.ran("FROM calibration_room_members"))

    def test_head_cannot_rewrite_etalon_or_comment_of_another_department(self):
        # Комнаты общие, но эталон по чужому оператору глава не перепишет —
        # та же граница, что при добавлении звонка и у СВ.
        for route, rules, body in (
            ("update_call_calibration_etalon", [ETALON_CALL, ("UPDATE calibration_room_calls", [(5,)])],
             {"scores": ["Incorrect"], "criterion_comments": ["x"]}),
            ("update_call_calibration_general_comment", [COMMENT_CALL, ("UPDATE calibration_room_calls", [(5,)])],
             {"general_comment": "x"}),
        ):
            for head in (ADMIN_HEAD, OPERATOR_HEAD):
                with self.subTest(route=route, head=head):
                    status, _, cursor = _call(route, head, ROOM_ID, 5, rules=rules, body=body, scope_ok=False)
                    self.assertEqual(403, status)
                    self.assertFalse(cursor.ran("UPDATE calibration_room_calls"))

    def test_head_exports_a_room_he_has_not_joined(self):
        status, payload, cursor = _call(
            "export_call_calibration_room_excel", OPERATOR_HEAD, ROOM_ID,
            rules=[ROOM_FOR_DETAIL, NOT_MEMBER, EXPORT_CALLS, EXPORT_EVALUATIONS],
        )
        self.assertEqual(200, status)
        self.assertFalse(cursor.ran("FROM calibration_room_members"))
        sheet = load_workbook(BytesIO(payload["xlsx"]))["Критерии"]
        self.assertIn("СВ-участник: Корректно (совпало)", [cell.value for cell in sheet["J"]])

    def test_head_reaches_history(self):
        status, payload, _ = _call(
            "get_calibration_room_history", OPERATOR_HEAD, ROOM_ID,
            rules=[("JOIN users u ON u.id = r.created_by_admin_id",
                    [(ROOM_ID, datetime(2026, 10, 27, 9, 0), "Октябрь", "СВ-автор")])],
        )
        self.assertEqual(200, status)
        self.assertEqual("room_created", payload["events"][0]["type"])

    def test_head_neither_joins_nor_evaluates(self):
        status, _, _ = _call("join_call_calibration_room", OPERATOR_HEAD, ROOM_ID,
                             rules=[("SELECT id FROM calibration_rooms", [(ROOM_ID,)])])
        self.assertEqual(403, status)
        status, _, _ = _call("evaluate_call_calibration_room", TRAINER_HEAD, ROOM_ID,
                             body={"call_id": 5, "scores": [], "criterion_comments": []})
        self.assertEqual(403, status)

    def test_deleting_stays_with_the_global_admin(self):
        rules = [("SELECT id FROM calibration_room", [(ROOM_ID,)])]
        for route, args in (("delete_calibration_room", (ROOM_ID,)),
                            ("delete_call_from_calibration_room", (ROOM_ID, 5))):
            # Глава с ролью admin тоже не удаляет: в журнале у него нет корзин,
            # и сервер держит ту же границу.
            for head in (ADMIN_HEAD, OPERATOR_HEAD, SV):
                with self.subTest(route=route, user=head):
                    status, _, cursor = _call(route, head, *args, rules=rules)
                    self.assertEqual(403, status)
                    self.assertFalse(cursor.ran("DELETE FROM"))
            for admin in (ADMIN, 448):
                with self.subTest(route=route, user=admin):
                    status, _, cursor = _call(route, admin, *args, rules=rules)
                    self.assertEqual(200, status)
                    self.assertTrue(cursor.ran("DELETE FROM"))


class CalibrationRoutesForSupervisorsTests(unittest.TestCase):
    """У супервайзера калибровка прежняя: войти в комнату, оценить вслепую."""

    def test_room_list_shows_whether_he_joined(self):
        rows = [(ROOM_ID, "2026-10", "Октябрь", datetime(2026, 10, 27), ROOM_CREATOR, "СВ-автор",
                 None, 1, 0, 0, 0)]
        status, payload, _ = _call("list_call_calibration_rooms", SV,
                                   rules=[("FROM calibration_rooms r", rows)])
        self.assertEqual(200, status)
        self.assertFalse(payload["rooms"][0]["joined"])

    def test_room_stays_closed_until_he_joins(self):
        status, payload, _ = _call("get_call_calibration_room", SV, ROOM_ID,
                                   rules=[ROOM_FOR_DETAIL, NOT_MEMBER])
        self.assertEqual(403, status)
        self.assertEqual("Join the room first", payload["error"])

    def test_member_evaluates_before_seeing_results(self):
        status, payload, _ = _call("get_call_calibration_room", SV, ROOM_ID,
                                   rules=[ROOM_FOR_DETAIL, MEMBER, CALL_LIST, SELECTED_CALL])
        self.assertEqual(200, status)
        self.assertTrue(payload["can_evaluate"])
        self.assertFalse(payload["can_view_results"])

    def test_actions_inside_a_room_wait_for_joining(self):
        cases = (
            ("add_call_to_calibration_room", (ROOM_ID,), [ADD_CALL_RULES[0], NOT_MEMBER, *ADD_CALL_RULES[1:]],
             ADD_CALL_BODY, "INSERT INTO calibration_room_calls"),
            ("export_call_calibration_room_excel", (ROOM_ID,),
             [ROOM_FOR_DETAIL, NOT_MEMBER, EXPORT_CALLS, EXPORT_EVALUATIONS], None, "AS evaluator_name"),
            ("update_call_calibration_etalon", (ROOM_ID, 5),
             [ETALON_CALL, NOT_MEMBER, ("UPDATE calibration_room_calls", [(5,)])],
             {"scores": ["Correct"], "criterion_comments": [""]}, "UPDATE calibration_room_calls"),
            ("update_call_calibration_general_comment", (ROOM_ID, 5),
             [COMMENT_CALL, NOT_MEMBER, ("UPDATE calibration_room_calls", [(5,)])],
             {"general_comment": "x"}, "UPDATE calibration_room_calls"),
        )
        for route, args, rules, body, write in cases:
            with self.subTest(route=route):
                status, payload, cursor = _call(route, SV, *args, rules=rules, body=body)
                self.assertEqual(403, status)
                self.assertEqual("Join the room first", payload["error"])
                self.assertFalse(cursor.ran(write))

    def test_member_works_only_with_operators_in_scope(self):
        for route, rules, body in (
            ("update_call_calibration_etalon", [ETALON_CALL, MEMBER, ("UPDATE calibration_room_calls", [(5,)])],
             {"scores": ["Correct"], "criterion_comments": [""]}),
            ("update_call_calibration_general_comment", [COMMENT_CALL, MEMBER, ("UPDATE calibration_room_calls", [(5,)])],
             {"general_comment": "x"}),
        ):
            with self.subTest(route=route, scope="свой оператор"):
                status, _, cursor = _call(route, SV, ROOM_ID, 5, rules=rules, body=body)
                self.assertEqual(200, status)
                self.assertTrue(cursor.ran("UPDATE calibration_room_calls"))
            with self.subTest(route=route, scope="чужой оператор"):
                status, _, cursor = _call(route, SV, ROOM_ID, 5, rules=rules, body=body, scope_ok=False)
                self.assertEqual(403, status)
                self.assertFalse(cursor.ran("UPDATE calibration_room_calls"))

    def test_member_exports_the_room(self):
        status, payload, _ = _call(
            "export_call_calibration_room_excel", SV, ROOM_ID,
            rules=[ROOM_FOR_DETAIL, MEMBER, EXPORT_CALLS, EXPORT_EVALUATIONS],
        )
        self.assertEqual(200, status)
        self.assertTrue(payload["xlsx"].startswith(b"PK"))

    def test_supervisor_heading_a_department_still_joins_first(self):
        status, _, _ = _call("update_calibration_room", SV_HEAD, ROOM_ID,
                             rules=[ROOM_FOR_UPDATE, NOT_MEMBER], body={"room_title": "x"})
        self.assertEqual(403, status)
        status, _, cursor = _call("join_call_calibration_room", SV_HEAD, ROOM_ID,
                                  rules=[("SELECT id FROM calibration_rooms", [(ROOM_ID,)])])
        self.assertEqual(200, status)
        self.assertTrue(cursor.ran("INSERT INTO calibration_room_members"))


class CalibrationClosedForOthersTests(unittest.TestCase):
    """Оператор, тренер и наблюдатель «Маркетинга» калибровку не получают."""

    def test_every_route_answers_403_with_one_message(self):
        self.assertEqual("Калибровка доступна руководителям и супервайзерам",
                         NS["CALIBRATION_FORBIDDEN_ERROR"])
        routes = (
            ("list_call_calibration_rooms", ()),
            ("create_call_calibration_room", ()),
            ("update_calibration_room", (ROOM_ID,)),
            ("add_call_to_calibration_room", (ROOM_ID,)),
            ("get_call_calibration_room", (ROOM_ID,)),
            ("export_call_calibration_room_excel", (ROOM_ID,)),
            ("update_call_calibration_etalon", (ROOM_ID, 5)),
            ("update_call_calibration_general_comment", (ROOM_ID, 5)),
            ("get_calibration_room_history", (ROOM_ID,)),
        )
        for user_id in (OPERATOR, TRAINER, OBSERVER):
            for route, args in routes:
                with self.subTest(user_id=user_id, route=route):
                    status, payload, cursor = _call(route, user_id, *args, rules=[ROOM_FOR_DETAIL])
                    self.assertEqual(403, status)
                    self.assertEqual(NS["CALIBRATION_FORBIDDEN_ERROR"], payload["error"])
                    self.assertEqual([], cursor.executed)


class JournalMirrorTests(unittest.TestCase):
    """Фронт журнала держит ту же границу, что сервер (_calibration_role)."""

    def setUp(self):
        self.source = source_cache.read(JOURNAL_PATH)

    def test_tab_and_room_management_follow_managers_and_supervisors(self):
        self.assertIn("const canUseCalibration = isAdminRole || isSupervisorRole;", self.source)
        self.assertIn("const canManageCalibrationRooms = isAdminRole || isSupervisorRole;", self.source)
        self.assertIn("const isAdminRole = isBaseAdminRole || isDepartmentHead;", self.source)

    def test_delete_buttons_stay_with_the_global_admin(self):
        for handler in ("handleDeleteCalibrationRoom(e, room.id)", "handleDeleteCalibrationCall(e, call.id)"):
            with self.subTest(handler=handler):
                at = self.source.index(handler)
                guard = self.source.rindex("{isGlobalAdminRole && (", 0, at)
                self.assertLess(at - guard, 400)

    def test_separate_tab_keeps_the_head_after_reload(self):
        # Журнал в отдельной вкладке после F5 берёт пользователя только из
        # sessionStorage: без отдела глава превращался в глобального админа.
        self.assertIn(
            "user: { id: userId, role: userRole, name: userName, headed_department_id: headedDepartmentId },",
            self.source,
        )

    def test_editors_follow_the_server_flag(self):
        # Ответ комнаты кладётся в состояние поле за полем: без этой строки
        # признак терялся, и редактор чужого звонка оставался открытым (поймал стенд).
        payload_at = self.source.index("const applyCalibrationDetailPayload")
        payload_end = self.source.index("}, []);", payload_at)
        self.assertIn("can_edit_call: d?.can_edit_call !== false,", self.source[payload_at:payload_end])
        self.assertIn(
            "const canEditCalibrationCall = canManageCalibrationRooms && calibrationDetail?.can_edit_call !== false;",
            self.source,
        )
        start = self.source.index("const canEditCalibrationCall =")
        body = self.source[start:]
        # Эталон (таблица, кнопка, проверки, сохранение) и общий комментарий —
        # только по признаку.
        self.assertEqual(4, body.count("canEditCalibrationCall && "))
        self.assertEqual(2, body.count("{canEditCalibrationCall ? ("))
        self.assertIn("if (!canEditCalibrationCall || !calibrationCall?.id || !activeCalibrationRoomId) return;", body)
        self.assertNotIn("canManageCalibrationRooms && calibrationCall", body)

    def test_room_creation_needs_someone_to_calibrate(self):
        self.assertIn(
            "const canFillCalibrationRooms = canManageCalibrationRooms"
            " && (isGlobalAdminRole || isSupervisorRole || supervisors.length > 0);",
            self.source,
        )
        create_at = self.source.index("setShowCalibrationCreateModal(true)")
        self.assertLess(create_at - self.source.rindex("{canFillCalibrationRooms && (", 0, create_at), 300)
        add_at = self.source.index("setEvalModalMode('calibration_add_call');\n                                                            setShowEvalModal(true);")
        self.assertLess(add_at - self.source.rindex("{canFillCalibrationRooms && (", 0, add_at), 1200)


if __name__ == "__main__":
    unittest.main()
