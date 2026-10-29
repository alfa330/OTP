# -*- coding: utf-8 -*-
"""Раздел «Удаленный КЦ»: детализация звонков за период в Excel (владелец, 15.10.2026).

Что защищаем:
  * звонок называется в файле так же, как в журнале и карточке водителя (исход, приписки);
  * номер водителя — только маской, текст ошибки АТС (в нём бывает номер) — нигде; ИИН и ID
    водителя — полностью и текстом (владелец, 29.10.2026);
  * итоги операторов считаются по тем же строкам, что на листе «Звонки»;
  * период: потолок, будущее, перепутанные границы — словами для руководителя;
  * ручка только для руководителя своего отдела, пустой период — 404 с кодом, а не пустой файл;
  * зелёный уголок «Число сохранено как текст» гасится у линии, и проводка от монолита цела.

Настоящий SQL на настоящей схеме — tests/test_dial_list_calls_export_postgres.py.
"""
import ast
import io
import os
import re
import sys
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dial_list import calls_export  # noqa: E402
from dial_list import routes as dial_routes  # noqa: E402
from dial_list import service as dial_service  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
UTC = timezone.utc
TODAY = date(2026, 10, 15)
PHONE = '77011234567'
# Выдуманные: ИИН с ведущим нулём (родился в 2005-м) — числом Excel его бы потерял.
IIN = '050312500017'
YANDEX_ID = '0f' * 16


def row(**over):
    """Строка CALLS_SQL: по умолчанию — разговор 2:05 с итогом и комментарием."""
    base = dict(
        requested_at=datetime(2026, 10, 14, 4, 5, 6, tzinfo=UTC), operator_id=520, operator='Своев Сава',
        line='902', state='finished', disposition='ANSWER', billsec=125, cancelled=False, auto_hangup='',
        talk_detected=None, ring_timeout_at=None, api_error='', driver='Иванов Иван', phone=PHONE,
        outcome='Отказ', subtype='Дорого', comment='перезвонить после 18', success=False, iin=IIN,
        driver_ids=YANDEX_ID)
    base.update(over)
    return tuple(base[key] for key in (
        'requested_at', 'operator_id', 'operator', 'line', 'state', 'disposition', 'billsec', 'cancelled',
        'auto_hangup', 'talk_detected', 'ring_timeout_at', 'api_error', 'driver', 'phone', 'outcome',
        'subtype', 'comment', 'success', 'iin', 'driver_ids'))


def call(**over):
    return calls_export.call_from_row(row(**over))


def book(calls, **kwargs):
    kwargs.setdefault('department', 'Удаленный КЦ')
    kwargs.setdefault('day_from', date(2026, 10, 1))
    kwargs.setdefault('day_to', date(2026, 10, 14))
    kwargs.setdefault('generated_by', 'Руководителева Раиса')
    kwargs.setdefault('generated_at', datetime(2026, 10, 15, 9, 30))
    stream, written = calls_export.build_workbook(iter(calls), **kwargs)
    stream.seek(0)
    return load_workbook(stream), written, stream


def values(sheet):
    return [[cell.value for cell in line] for line in sheet.iter_rows()]


class PeriodTests(unittest.TestCase):
    def days(self, a, b):
        return calls_export.parse_days(a, b, today=TODAY)

    def test_plain_and_single_day(self):
        self.assertEqual(self.days('2026-10-01', '2026-10-15'), (date(2026, 10, 1), TODAY))
        self.assertEqual(self.days('2026-10-03', ''), (date(2026, 10, 3), date(2026, 10, 3)))
        self.assertEqual(self.days('', '2026-10-03'), (date(2026, 10, 3), date(2026, 10, 3)))

    def test_swapped_and_future_tail(self):
        self.assertEqual(self.days('2026-10-10', '2026-10-01'), (date(2026, 10, 1), date(2026, 10, 10)))
        # Хвост в будущем отрезается: звонков там нет, «по 31.10» в файле было бы неправдой.
        self.assertEqual(self.days('2026-10-10', '2026-10-31'), (date(2026, 10, 10), TODAY))

    def test_ceiling_is_a_month(self):
        self.assertEqual(calls_export.MAX_DAYS, 31)
        self.assertEqual(self.days('2026-09-15', '2026-10-15'), (date(2026, 9, 15), TODAY))  # 31 день
        with self.assertRaisesRegex(ValueError, r'^Не больше 31 дня за раз'):
            self.days('2026-09-14', '2026-10-15')
        # Отрезанный будущий хвост в потолок не считается.
        self.assertEqual(self.days('2026-09-15', '2026-11-30'), (date(2026, 9, 15), TODAY))
        with self.assertRaisesRegex(ValueError, r'^Не больше 7 дней'):
            calls_export.parse_days('2026-10-01', '2026-10-08', today=TODAY, max_days=7)

    def test_refusals_in_words(self):
        for args, text in ((('', ''), 'Выберите период'), (('2026-10-16', '2026-10-20'), 'ещё не наступил'),
                           (('15.10.2026', ''), 'ГГГГ-ММ-ДД'), (('2026-13-01', ''), 'ГГГГ-ММ-ДД')):
            with self.assertRaisesRegex(ValueError, text):
                self.days(*args)

    def test_file_name(self):
        self.assertEqual(calls_export.report_filename(date(2026, 10, 1), TODAY),
                         'Детализация звонков 01.10.2026 — 15.10.2026.xlsx')
        self.assertEqual(calls_export.report_filename(TODAY, TODAY), 'Детализация звонков 15.10.2026.xlsx')


class CallRowTests(unittest.TestCase):
    def test_answered_call(self):
        c = call()
        self.assertEqual(c['at'], datetime(2026, 10, 14, 9, 5, 6))  # Алматы, без пояса
        self.assertEqual((c['result'], c['talk_sec'], c['outcome']), ('answered', 125, 'Отказ · Дорого'))
        self.assertEqual((c['phone'], c['line'], c['note']), ('+7 ••• ••• 45 67', '902', ''))
        self.assertEqual((c['iin'], c['driver_ids']), (IIN, YANDEX_ID))
        # Sapar ещё не спрашивали (запрос отдаёт '') или водителя он не знает — пусто, а не «None».
        self.assertEqual(call(driver_ids='')['driver_ids'], '')
        self.assertEqual(call(subtype=None)['outcome'], 'Отказ')
        self.assertEqual(call(outcome=None, subtype=None)['outcome'], '')

    def test_result_is_the_journal_rule(self):
        """Исход — общий помощник журнала и карточки, с теми же поправками."""
        cases = [
            (dict(), 'answered'),
            (dict(disposition='BUSY', billsec=0), 'busy'),
            (dict(disposition='NOANSWER', billsec=0), 'no_answer'),
            (dict(disposition='CHANUNAVAIL', billsec=0), 'other'),
            (dict(state='failed', disposition=''), 'failed'),
            (dict(state='ended', disposition=''), ''),
            (dict(state='requested', disposition=''), ''),
            (dict(disposition='CANCEL', billsec=0, cancelled=True), 'cancelled'),
            # Сообщение сети с ANSWER (автоответчик) и ответ, где телефон разговора не слышал, — недозвон.
            (dict(auto_hangup='announcement', billsec=40), 'no_answer'),
            (dict(talk_detected=False, billsec=0), 'no_answer'),
            (dict(talk_detected=False, billsec=3), 'answered'),
        ]
        for over, expected in cases:
            with self.subTest(over=over):
                c = call(**over)
                self.assertEqual(c['result'], expected)
                r = row(**over)
                self.assertEqual(c['result'], dial_service.attempt_outcome(r[4], r[5], r[6], r[8], r[9],
                                                                          cancelled=r[7]))
                # Разговор — только у дозвона: у автоответчика Binotel пишет свою длительность.
                self.assertEqual(c['talk_sec'], r[6] if expected == 'answered' else 0)

    def test_auto_hangup_notes_follow_the_card(self):
        ring_at = datetime(2026, 10, 14, 4, 5, 36, tzinfo=UTC)
        no = dict(disposition='NOANSWER', billsec=0)
        self.assertEqual(call(**no, ring_timeout_at=ring_at)['note'], calls_export.RING_TIMEOUT_NOTE)
        self.assertEqual(call(**no, auto_hangup='ring_timeout')['note'], calls_export.RING_TIMEOUT_NOTE)
        # Сообщение сети сильнее гудков.
        self.assertEqual(call(**no, auto_hangup='announcement', ring_timeout_at=ring_at)['note'],
                         calls_export.ANNOUNCEMENT_NOTE)
        # Водитель успел ответить в последний момент — это разговор, приписка была бы неправдой.
        self.assertEqual(call(ring_timeout_at=ring_at)['note'], '')
        self.assertEqual(call(disposition='CANCEL', billsec=0, cancelled=True, ring_timeout_at=ring_at)['note'],
                         calls_export.CANCELLED_NOTE)
        self.assertEqual(call(**no)['note'], '')
        self.assertEqual(call(state='ended', disposition='')['note'], '')

    def test_twins_on_the_site(self):
        """Подписи исхода и приписки — те же слова, что в журнале и карточке водителя."""
        journal = (ROOT / 'src/components/dial_list/DialListJournal.jsx').read_text(encoding='utf-8')
        block = journal[journal.index('const RESULT_LABEL = {'):]
        block = block[:block.index('};')]
        site = dict(re.findall(r"(\w+): '([^']+)'", block))
        self.assertEqual(site, {k: v for k, v in calls_export.RESULT_LABELS.items() if k})
        hints = (ROOT / 'src/components/dial_list/attemptHints.js').read_text(encoding='utf-8')
        self.assertIn(f"RING_TIMEOUT_HINT = ' · {calls_export.RING_TIMEOUT_NOTE}'", hints)
        self.assertIn(f"ANNOUNCEMENT_HINT = ' · {calls_export.ANNOUNCEMENT_NOTE}'", hints)

    def test_failure_note_in_words_without_binotel_text(self):
        raw = "Binotel calls/internal-number-to-external-number status='error' code={}: Something {}"
        cases = [
            (raw.format(150, "Can't call to the ext"), 'линия оператора не на связи'),
            (raw.format(104, 'went wrong for 77011234567'), 'АТС отказала (код 104)'),
            (raw.format(105, 'went wrong'), 'АТС не приняла звонок (код 105)'),
            (raw.format(106, 'too frequent'), 'слишком частые запросы к АТС'),
            ('Read timed out. (read timeout=40)', 'нет связи с АТС'),
            (dial_service.BINOTEL_NOT_CONFIGURED, 'подключение к Binotel не настроено'),
            ('', 'АТС отказала'),
        ]
        for api_error, expected in cases:
            with self.subTest(api_error=api_error):
                note = call(state='failed', disposition='', billsec=0, api_error=api_error)['note']
                self.assertEqual(note, expected)
                self.assertNotIn('1234567', note)


class WorkbookTests(unittest.TestCase):
    def test_sheets_columns_and_formats(self):
        wb, written, _ = book([call(comment='=HYPERLINK("http://x")', success=True),
                               call(requested_at=datetime(2026, 10, 14, 5, 0, tzinfo=UTC), disposition='NOANSWER',
                                    billsec=0, outcome=None, subtype=None, comment='')])
        self.assertEqual(written, 2)
        self.assertEqual(wb.sheetnames, ['Звонки', 'Операторы', 'Контекст'])
        ws = wb['Звонки']
        self.assertEqual([c.value for c in ws[1]], [title for _k, title, _w in calls_export.COLUMNS])
        self.assertEqual([c.value for c in ws[1]][3:7], ['Водитель', 'ID водителя', 'ИИН', 'Номер'])
        self.assertEqual((ws.freeze_panes, ws.auto_filter.ref), ('A2', 'A1:M3'))
        first = {title: cell for (_k, title, _w), cell in zip(calls_export.COLUMNS, ws[2])}
        # ИИН и ID — текстом: у ИИН числом пропал бы ведущий ноль.
        for title, expected in (('ИИН', IIN), ('ID водителя', YANDEX_ID)):
            self.assertEqual((first[title].value, first[title].data_type, first[title].number_format),
                             (expected, 's', '@'), title)
        self.assertEqual(first['Дата и время'].value, datetime(2026, 10, 14, 9, 5, 6))
        self.assertEqual(first['Дата и время'].number_format, 'DD.MM.YYYY HH:MM:SS')
        # Разговор — длительность: «2:05» на экране и 125 секунд в сумме.
        self.assertEqual(first['Разговор'].number_format, '[m]:ss')
        self.assertEqual(first['Разговор'].value, timedelta(seconds=125))
        self.assertEqual((first['Линия'].value, first['Линия'].data_type, first['Линия'].number_format),
                         ('902', 's', '@'))
        # Комментарий с «=» — строка, а не живая формула.
        self.assertEqual((first['Комментарий оператора'].value, first['Комментарий оператора'].data_type),
                         ('=HYPERLINK("http://x")', 's'))
        self.assertEqual((first['Результат'].value, first['Успешка'].value), ('Дозвонились', 'да'))
        second = {title: cell.value for (_k, title, _w), cell in zip(calls_export.COLUMNS, ws[3])}
        # Пусто там, где нечего сказать: ни «0:00» разговора, ни «нет» успешки.
        self.assertEqual((second['Результат'], second['Разговор'], second['Итог оператора'], second['Успешка']),
                         ('Не поднял трубку', None, None, None))

    def test_several_ids_wrap_and_unknown_is_empty(self):
        """Аккаунтов несколько — все, каждый своей строкой в ячейке; Sapar водителя не знает — пусто."""
        several_ids = f'{"a" * 32}, {"b" * 32}'
        wb, _, _ = book([call(driver_ids=several_ids), call(driver_ids=''), call()])
        column = [title for _k, title, _w in calls_export.COLUMNS].index('ID водителя')
        several, unknown, single = (wb['Звонки'][index][column] for index in (2, 3, 4))
        self.assertEqual((several.value, several.alignment.wrap_text), (several_ids, True))
        self.assertIsNone(unknown.value)
        # Один ID — без переноса: строка не растёт зря.
        self.assertEqual(single.value, YANDEX_ID)
        self.assertFalse(single.alignment.wrap_text)

    def test_operators_are_counted_from_the_same_rows(self):
        calls = [
            call(),                                                                    # Своев: разговор 2:05
            call(billsec=55, success=True),                                            # Своев: 0:55, успешка
            call(disposition='NOANSWER', billsec=0),                                   # Своев: недозвон
            call(operator_id=701, operator='Гостев Глеб', disposition='CANCEL', billsec=0, cancelled=True),
            call(operator_id=701, operator='Гостев Глеб', state='failed', disposition='', billsec=0),
            call(operator_id=701, operator='Гостев Глеб', billsec=30),
            call(operator_id=701, operator='Гостев Глеб', auto_hangup='announcement', billsec=12),
        ]
        wb, written, _ = book(calls)
        ws = wb['Операторы']
        self.assertEqual([c.value for c in ws[1]], [title for title, _w in calls_export.OPERATOR_COLUMNS])
        rows = values(ws)[1:]
        # Больше звонков — выше.
        self.assertEqual([r[0] for r in rows], ['Гостев Глеб', 'Своев Сава', 'Итого'])
        gostev, svoev, total = rows
        self.assertEqual(gostev[1:4], [4, 1, 0.25])
        self.assertEqual((gostev[4], gostev[5]), (timedelta(seconds=30), timedelta(seconds=30)))
        self.assertEqual(gostev[6:], [0, 1, 1])
        self.assertEqual(svoev[1:3], [3, 2])
        self.assertAlmostEqual(svoev[3], 2 / 3)
        self.assertEqual((svoev[4], svoev[5]), (timedelta(seconds=180), timedelta(seconds=90)))
        self.assertEqual(svoev[6:], [1, 0, 0])
        self.assertEqual(total[1:3], [7, 3])
        self.assertEqual(total[4], timedelta(seconds=210))
        self.assertEqual(total[6:], [1, 1, 1])
        # Сумма «Разговор» на листе «Звонки» равна итогу — это одни и те же строки.
        talk_column = [title for _k, title, _w in calls_export.COLUMNS].index('Разговор')
        talk = sum((r[talk_column] for r in values(wb['Звонки'])[1:] if r[talk_column]), timedelta())
        self.assertEqual(talk, timedelta(seconds=210))
        self.assertEqual(wb['Операторы']['D2'].number_format, '0%')
        self.assertEqual(wb['Операторы']['E2'].number_format, '[h]:mm:ss')
        self.assertTrue(wb['Операторы']['A4'].font.bold)

    def test_single_operator_has_no_total_row(self):
        wb, _, _ = book([call(), call(disposition='BUSY', billsec=0)])
        self.assertEqual([r[0] for r in values(wb['Операторы'])], ['Оператор', 'Своев Сава'])

    def test_context_names_the_period_and_the_rules(self):
        wb, _, _ = book([call()])
        context = {r[0]: r[1] for r in values(wb['Контекст']) if r and r[0]}
        self.assertEqual(context['Отдел'], 'Удаленный КЦ')
        self.assertEqual(context['Период'], '01.10.2026 — 14.10.2026')
        self.assertEqual(context['Собрано'], '15.10.2026 09:30 (Алматы)')
        self.assertEqual(context['Собрал'], 'Руководителева Раиса')
        self.assertEqual(context['Звонков'], 1)
        for label in ('Результат', 'Отменён до ответа', 'Ошибка АТС', 'Успешка', 'ID водителя', 'ИИН', 'Номер',
                      'Тестовые номера'):
            self.assertTrue(context.get(label), label)
        self.assertIn('Sapar', context['ID водителя'])

    def test_no_full_phone_or_binotel_text_anywhere_iin_and_id_in_full(self):
        """Файл — тоже ручка раздела: полный номер и текст АТС не выходят ни на один лист. ИИН и ID
        водителя — полностью (владелец, 29.10.2026: «в детализацию и в отчеты добавить id и иин»)."""
        calls = [call(), call(state='failed', disposition='', billsec=0,
                              api_error="status='error' code=104: bad number 77011234567")]
        _, _, stream = book(calls)
        stream.seek(0)
        with ZipFile(stream) as archive:
            xml = ''.join(archive.read(name).decode('utf-8') for name in archive.namelist()
                          if name.endswith('.xml'))
        for forbidden in ('1234567', '77011234567', "status='error'", 'bad number'):
            self.assertNotIn(forbidden, xml, forbidden)
        for present in ('45 67', IIN, YANDEX_ID):
            self.assertIn(present, xml, present)

    def test_route_does_not_read_phone_outside_the_mask(self):
        src = (ROOT / 'dial_list/calls_export.py').read_text(encoding='utf-8')
        sql = calls_export.CALLS_SQL
        # ИИН и ID — из строки водителя; ещё не спрошенный у Sapar ID приходит пустым, а не NULL.
        self.assertIn("l.iin, COALESCE(l.yandex_driver_ids, '')", sql)
        # Номер читается ровно один раз — чтобы тут же стать маской.
        self.assertEqual(src.count('phone_norm'), 3)  # SQL, распаковка строки, mask_phone(phone_norm)
        self.assertIn('"phone": service.mask_phone(phone_norm)', src)
        # Звонки тестовых номеров не входят (как в сводке «Операторы»).
        self.assertIn(dial_service.test_keys.DIAL_LIST_LEAD_NOT_TEST_SQL, sql)


class TextWarningTests(unittest.TestCase):
    def calls(self):
        recorded = []

        def patch(stream, sqref, sheet_path=None):
            recorded.append({'sqref': sqref, 'sheet_path': sheet_path})
            return stream
        return recorded, patch

    def test_text_columns_on_the_calls_sheet(self):
        recorded, patch = self.calls()
        book([call(), call()], text_warning_patch=patch)
        self.assertEqual(recorded, [{'sqref': 'C2:C3 E2:E3 F2:F3', 'sheet_path': 'xl/worksheets/sheet1.xml'}])
        titles = [title for _k, title, _w in calls_export.COLUMNS]
        self.assertEqual((titles[2], titles[4], titles[5]), ('Линия', 'ID водителя', 'ИИН'))

    def test_real_patch_lands_and_book_opens(self):
        from tests import source_cache
        node = next(n for n in source_cache.tree(ROOT / 'bot_schedule2.py').body
                    if isinstance(n, ast.FunctionDef) and n.name == '_excel_suppress_number_as_text_warning')
        namespace = {}
        exec(compile(ast.Module(body=[node], type_ignores=[]), 'bot', 'exec'),
             {'BytesIO': io.BytesIO, 'ZipFile': ZipFile, 'ZIP_DEFLATED': ZIP_DEFLATED}, namespace)
        _, _, stream = book([call()], text_warning_patch=namespace['_excel_suppress_number_as_text_warning'])
        stream.seek(0)
        xml = ZipFile(stream).read('xl/worksheets/sheet1.xml').decode('utf-8')
        self.assertIn('<ignoredError sqref="C2:C2 E2:E2 F2:F2" numberStoredAsText="1"/>', xml)
        self.assertLess(xml.index('<pageMargins'), xml.index('<ignoredErrors'))
        stream.seek(0)
        self.assertEqual(load_workbook(stream).sheetnames[0], 'Звонки')

    def test_broken_patch_still_gives_the_file(self):
        def explode(_stream, _sqref, sheet_path=None):
            raise RuntimeError('перепаковка не удалась')
        wb, written, _ = book([call()], text_warning_patch=explode)
        self.assertEqual((written, wb.sheetnames[0]), (1, 'Звонки'))

    def test_monolith_hands_the_helper_over(self):
        bot = (ROOT / 'bot_schedule2.py').read_text(encoding='utf-8-sig')
        block = bot.split('_dial_list_bp = build_dial_list_blueprint(')[1].split('\n    )\n')[0]
        self.assertIn('excel_text_warning=_excel_suppress_number_as_text_warning', block)
        routes = (ROOT / 'dial_list/routes.py').read_text(encoding='utf-8')
        signature = routes.split('def build_dial_list_blueprint(')[1].split('):')[0]
        # Без хелпера раздел не падает — файл просто с уголком.
        self.assertIn('excel_text_warning=None', signature)


RCC, FRONT_OFFICE, OP, SZOV = 1954, 909, 367, 1
REGIONS = {"id": 37, "name": "Группа Регионы", "active": True}


class _Svc:
    """Сервис раздела без базы. Зона руководителя — как на проде (manager_scope): админ и глава
    СЗоВ — весь раздел, глава удалённого КЦ — свой отдел, остальные — ничего. Группы отчётов —
    у главы «Фронт офисов» (report_groups)."""

    def __init__(self, written=2):
        self.seen = []
        self.calls = []
        self.written = written

    def manager_scope(self, is_admin, heads, login=None, supervisor_department_code=None):
        if is_admin or SZOV in heads:
            return None
        return [d for d in heads if d == RCC]

    def report_groups(self, heads):
        self.calls.append(('report_groups', list(heads)))
        return [dict(REGIONS)] if FRONT_OFFICE in list(heads) else []

    def list_departments(self, ids=None, profile=None):
        self.calls.append(('list_departments', ids, profile))
        rows = [{"department_id": RCC, "department_name": "Удаленный КЦ", "enabled": True, "leads_total": 340,
                 "leads_open": 120, "portion_size": 20, "code": "remote_cc", "provider": "binotel",
                 "configured": True, "profile": "remote_cc"},
                {"department_id": 560, "department_name": "Тез КЦ", "enabled": False, "leads_total": 10,
                 "leads_open": 5, "portion_size": 20, "code": "tez", "provider": "binotel",
                 "configured": True, "profile": "tez"}]
        return [r for r in rows if (ids is None or r["department_id"] in ids)
                and (profile is None or r["profile"] == profile)]

    def candidate_departments(self, profile=None):
        return [{"department_id": 2008}]

    def calls_export(self, department_id, day_from, day_to, generated_by='', text_warning_patch=None, groups=None):
        self.seen.append((department_id, day_from, day_to, generated_by, text_warning_patch, groups))
        if not self.written:
            return None, 0
        stream, written = calls_export.build_workbook(
            iter([call()] * self.written), department='Удаленный КЦ', day_from=day_from, day_to=day_to,
            groups=None if groups is None else [g["name"] for g in groups])
        return stream, written

    def __getattr__(self, name):
        # Любая другая ручка раздела дошла до сервиса — значит, страж её пропустил.
        def touched(*args, **kwargs):
            self.calls.append(('touched', name))
            return {}
        return touched


def _client(svc, role='admin', heads=(), warning=None, name='Руководителева Раиса'):
    import flask
    requester = (5, 111, name, role, None, None, None, 'login5')
    bp = dial_routes.build_dial_list_blueprint(
        db=None, require_api_key=lambda fn: fn, build_cors_preflight_response=lambda: ('', 204),
        resolve_requester=lambda: (5, requester, None), is_admin_role=lambda r: r in ('admin', 'super_admin'),
        is_super_admin_role=lambda r: r == 'super_admin', headed_department_ids=lambda rid: list(heads),
        excel_text_warning=warning, service=svc)
    app = flask.Flask(f'dial_list_calls_export_{role}_{"_".join(map(str, heads))}')
    app.register_blueprint(bp)
    return app.test_client()


class RouteTests(unittest.TestCase):
    URL = f'/api/dial_list/departments/{RCC}/calls/export'
    DAYS = '?date_from=2026-10-01&date_to=2026-10-03'

    def test_file_with_requester_and_helper(self):
        svc = _Svc()
        helper = object()
        resp = _client(svc, warning=helper).get(self.URL + self.DAYS)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.mimetype, dial_routes.XLSX_MIME)
        self.assertIn('attachment', resp.headers['Content-Disposition'])
        # Админ всей компании — весь отдел: групп нет.
        self.assertEqual(svc.seen, [(RCC, date(2026, 10, 1), date(2026, 10, 3), 'Руководителева Раиса', helper, None)])
        self.assertEqual(load_workbook(io.BytesIO(resp.data)).sheetnames, ['Звонки', 'Операторы', 'Контекст'])

    def test_empty_period_is_said_in_words(self):
        resp = _client(_Svc(written=0)).get(self.URL + self.DAYS)
        self.assertEqual(resp.status_code, 404)
        self.assertEqual(resp.get_json(), {"error": "За выбранный период звонков не было",
                                           "code": "DIAL_LIST_CALLS_EXPORT_EMPTY"})

    def test_bad_period_is_400_and_service_not_called(self):
        svc = _Svc()
        client = _client(svc)
        for query in ('', '?date_from=2026-01-01&date_to=2026-12-31', '?date_from=01.10.2026'):
            with self.subTest(query=query):
                resp = client.get(self.URL + query)
                self.assertEqual(resp.status_code, 400)
                self.assertTrue(resp.get_json()['error'])
        self.assertEqual(svc.seen, [])

    def test_only_managers_of_this_department(self):
        svc = _Svc()
        # Глава удалённого КЦ (сам не админ всей компании) — свой отдел целиком, чужой — нет.
        self.assertEqual(_client(svc, role='sv', heads=[RCC]).get(self.URL + self.DAYS).status_code, 200)
        self.assertEqual(svc.seen[-1][-1], None)
        resp = _client(svc, role='sv', heads=[RCC]).get('/api/dial_list/departments/560/calls/export' + self.DAYS)
        self.assertEqual((resp.status_code, resp.get_json()), (403, {"error": "Это не ваш отдел"}))
        # Рядовой и СВ чужого отдела — мимо.
        for role, heads in (('operator', []), ('sv', [])):
            self.assertEqual(_client(svc, role=role, heads=heads).get(self.URL + self.DAYS).status_code, 403)
        self.assertEqual(len(svc.seen), 1)


class GroupReportHeadTests(unittest.TestCase):
    """Глава «Регионов» (владелец, 15.10.2026: «мог выгружать отчеты только по своей группе»):
    «Регионы» — группа отдела «Фронт офисы»; глава отдела — роль admin. Раздел ей — одна
    «Детализация» по своей группе, остальные ручки — нет; админом раздела её роль не делает."""

    URL = f'/api/dial_list/departments/{RCC}/calls/export?date_from=2026-10-01&date_to=2026-10-03'

    def test_head_of_regions_exports_only_her_group(self):
        svc = _Svc()
        resp = _client(svc, role='admin', heads=[FRONT_OFFICE], name='Глава Регионов').get(self.URL)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(svc.seen, [(RCC, date(2026, 10, 1), date(2026, 10, 3), 'Глава Регионов', None, [REGIONS])])
        context = {r[0]: r[1] for r in values(load_workbook(io.BytesIO(resp.data))['Контекст']) if r and r[0]}
        self.assertTrue(context['Группа'].startswith('Группа Регионы'))

    def test_head_of_regions_sees_departments_without_base_numbers(self):
        svc = _Svc()
        client = _client(svc, role='admin', heads=[FRONT_OFFICE])
        body = client.get('/api/dial_list/departments').get_json()
        # Только «Удаленный КЦ» (профиль remote_cc) и без чисел базы.
        self.assertEqual(body, {"status": "success", "profile": "remote_cc", "candidates": [],
                                "report_groups": [REGIONS],
                                "departments": [{"department_id": RCC, "department_name": "Удаленный КЦ",
                                                 "profile": "remote_cc"}]})
        # «Автообзвон Тез» — не её раздел.
        self.assertEqual(client.get('/api/dial_list/departments?profile=tez').status_code, 403)
        # Выгрузка отдела Тез (другой профиль) — тоже нет.
        resp = client.get('/api/dial_list/departments/560/calls/export?date_from=2026-10-01&date_to=2026-10-01')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(svc.seen, [])

    def test_head_of_regions_gets_nothing_else(self):
        svc = _Svc()
        client = _client(svc, role='admin', heads=[FRONT_OFFICE])
        for method, url in (('get', '/api/dial_list/overview'), ('get', f'/api/dial_list/departments/{RCC}/leads'),
                            ('get', f'/api/dial_list/departments/{RCC}/lines'),
                            ('get', f'/api/dial_list/departments/{RCC}/settings'),
                            ('get', f'/api/dial_list/departments/{RCC}/analytics'),
                            ('get', f'/api/dial_list/departments/{RCC}/users'),
                            ('get', f'/api/dial_list/departments/{RCC}/leads/summary'),
                            ('post', f'/api/dial_list/departments/{RCC}/enroll'),
                            ('get', '/api/dial_list/leads/00000000-0000-0000-0000-000000000001')):
            with self.subTest(url=url):
                self.assertEqual(getattr(client, method)(url).status_code, 403)
        # Карточка сперва спрашивает отдел водителя (lead_department) — дальше стража не прошла ни одна.
        self.assertEqual([c for c in svc.calls if c[0] == 'touched' and c[1] != 'lead_department'], [])
        # И выгрузка — только по отделам раздела.
        resp = client.get('/api/dial_list/departments/77/calls/export?date_from=2026-10-01&date_to=2026-10-01')
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(svc.seen, [])

    def test_head_of_another_department_with_admin_role_is_not_a_section_admin(self):
        svc = _Svc()
        client = _client(svc, role='admin', heads=[OP])
        self.assertEqual(client.get('/api/dial_list/departments').status_code, 403)
        self.assertEqual(client.get(self.URL).status_code, 403)
        self.assertEqual(svc.seen, [])

    def test_admins_and_overseers_keep_the_whole_department(self):
        cases = (('admin', []), ('super_admin', [FRONT_OFFICE]), ('admin', [SZOV, RCC]),
                 ('admin', [FRONT_OFFICE, RCC]))  # глава и «Регионов», и удалённого КЦ — весь отдел
        for role, heads in cases:
            with self.subTest(role=role, heads=heads):
                svc = _Svc()
                self.assertEqual(_client(svc, role=role, heads=heads).get(self.URL).status_code, 200)
                self.assertIsNone(svc.seen[-1][-1])
                body = _client(svc, role=role, heads=heads).get('/api/dial_list/departments').get_json()
                self.assertNotIn('report_groups', body)

    def test_rule_is_one_list_and_the_query_filters_by_membership_on_the_call_day(self):
        self.assertEqual(dial_service.DIAL_LIST_GROUP_REPORT_HEAD_CODES, frozenset({'front_office'}))
        sql = ' '.join(calls_export.calls_sql([37]).split())
        self.assertIn('gm.group_id = ANY(%(group_ids)s)', sql)
        self.assertIn("gm.start_date <= (t.requested_at AT TIME ZONE 'Asia/Almaty')::date", sql)
        self.assertIn("gm.end_date >= (t.requested_at AT TIME ZONE 'Asia/Almaty')::date", sql)
        # Весь отдел — без фильтра групп; пустой список групп — фильтр есть (ничего), а не «весь отдел».
        self.assertNotIn('group_operator_memberships', calls_export.calls_sql(None))
        self.assertIn('group_operator_memberships', calls_export.calls_sql([]))
        # Тот же список в App.jsx сверяет tests/dial_list_section_access.test.mjs.


if __name__ == '__main__':
    unittest.main()
