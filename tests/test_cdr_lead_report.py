# -*- coding: utf-8 -*-
"""Книга «сделки + касания» (cdr/lead_report.py) собирается и читается обратно.

Проверяется форма, к которой привыкли получатели файлов: «Контекст» первым
листом, «Лиды» с двухэтажной шапкой и блоками «Касание N», «Касания» по звонку в
строке, «Операторы», «Сводка»; телефоны текстом; полный адрес записи в блоках,
когда блоков мало. JSON-двойник начинается с ключа КОНТЕКСТ_ДЛЯ_ИИ.

Телефоны учебные (7XX555XXXX), ссылки на записи — учебный хост.
"""

import json
import unittest
from datetime import datetime

from openpyxl import load_workbook

from cdr import lead_report, leads as L


def make_rows():
    lead_a = {
        'key': '101', 'source': 'amo', 'stream_type': 0, 'work_day': '2026-09-10',
        'moment': datetime(2026, 9, 10, 10, 0, 0), 'full_name': 'Заявка', 'phone': '7015550001',
        'phones': ['7015550001'], 'owner': 'Жупан Аружан', 'owner_ext': '6474',
        'park': 'iTaxi', 'city': 'Астана', 'base_title': '', 'stage': 'ПРОШЕЛ РЕГИСТРАЦИЮ',
        'reason': '', 'comment': '', 'tags': 'forma_itaxi_google', 'utm_source': 'google',
        'lead_type': 'форма', 'registered': True, 'updated_at': datetime(2026, 9, 10, 10, 5, 0),
        'dup': '',
    }
    touches_a = [
        {'started_at': datetime(2026, 9, 10, 9, 59, 50), 'answered_at': None, 'phone': '7015550001',
         'ext': '6474', 'operator': 'Жупан Аружан', 'direction': 'Основа ОП',
         'call_type': 'Входящий', 'result': 'Разговор', 'talk_seconds': 75, 'dial_seconds': 80,
         'queue': '3034', 'recording_url': 'http://recordings.test/a.wav', 'linkedid': '1.1',
         'legs': 1},
        {'started_at': datetime(2026, 9, 10, 10, 30, 0), 'answered_at': None, 'phone': '7015550001',
         'ext': '6656', 'operator': 'Зинеден Аружан', 'direction': 'Основа ОП',
         'call_type': 'Исходящий', 'result': 'Не ответил', 'talk_seconds': 0, 'dial_seconds': 20,
         'queue': '', 'recording_url': 'http://recordings.test/b.wav', 'linkedid': '1.2', 'legs': 1},
    ]
    lead_b = dict(lead_a, key='102', phone='7025550002', phones=['7025550002'], full_name='Другая',
                  lead_type='wz', stage='Диалоги', moment=datetime(2026, 9, 11, 12, 0, 0),
                  updated_at=None)
    rows = [
        {'lead': lead_a, 'touches': touches_a, 'agg': L.aggregate(lead_a, touches_a),
         'owner_called': 'да', 'lead_type': 'форма', 'stage': 'ПРОШЕЛ РЕГИСТРАЦИЮ', 'park': 'iTaxi'},
        {'lead': lead_b, 'touches': [], 'agg': L.aggregate(lead_b, []),
         'owner_called': '', 'lead_type': 'wz', 'stage': 'Диалоги', 'park': 'iTaxi'},
    ]
    return rows


class WorkbookTests(unittest.TestCase):

    def setUp(self):
        self.rows = make_rows()
        self.summary = L.summarize(self.rows)
        self.blocks, self.full_url = L.blocks_for(self.summary['max_touches'])
        stream, written = lead_report.build_workbook(
            self.rows, source_info=dict(L.SOURCES['amo'], key='amo'),
            period_from='2026-09-10', period_to='2026-09-11', touches_to='2026-09-12',
            summary=self.summary, blocks=self.blocks, full_url=self.full_url,
            filters_note='парк = iTaxi', generated_by='Тест')
        self.written = written
        self.book = load_workbook(stream)

    def test_листы_в_привычном_порядке(self):
        self.assertEqual(self.book.sheetnames, ['Контекст', 'Лиды', 'Касания', 'Операторы', 'Сводка'])
        self.assertEqual(self.written, 2)

    def test_лист_лидов_шапка_и_блоки(self):
        sheet = self.book['Лиды']
        top = [cell.value for cell in sheet[1]]
        second = [cell.value for cell in sheet[2]]
        self.assertIn('СВОДКА ПО КАСАНИЯМ', top)
        self.assertIn('КАСАНИЕ 1', top)
        self.assertIn('КАСАНИЕ 2', top)
        self.assertNotIn('КАСАНИЕ 3', top)          # блоков ровно столько, сколько касаний максимум
        self.assertIn('Ответственный звонил', second)
        self.assertIn('Ссылка на запись', second)    # блоков мало — полный адрес
        self.assertEqual(sheet.freeze_panes, 'C3')
        rows = list(sheet.iter_rows(min_row=3, values_only=True))
        self.assertEqual(len(rows), 2)
        first = dict(zip(second, rows[0]))
        self.assertEqual(first['Телефон'], '7015550001')          # текстом
        self.assertEqual(first['Касаний всего'], 2)
        self.assertEqual(first['Разговоров'], 1)
        self.assertEqual(first['Ответственный звонил'], 'да')
        self.assertEqual(first['Касаний до заявки'], 1)
        self.assertEqual(first['Реакция ОП, мин'], 30.0)
        self.assertEqual(first['Зарегистрирован'], 'да')

    def test_полный_адрес_записи_в_блоке_кликабелен(self):
        sheet = self.book['Лиды']
        header = [cell.value for cell in sheet[2]]
        index = header.index('Ссылка на запись') + 1
        cell = sheet.cell(row=3, column=index)
        self.assertEqual(cell.value, 'http://recordings.test/a.wav')
        self.assertEqual(cell.hyperlink.target, 'http://recordings.test/a.wav')

    def test_лист_касаний_по_звонку_в_строке(self):
        sheet = self.book['Касания']
        header = [cell.value for cell in sheet[1]]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 2)
        first = dict(zip(header, rows[0]))
        self.assertEqual(first['ID сделки'], '101')
        self.assertEqual(first['Относительно заявки'], 'до заявки')
        self.assertEqual(first['№ касания'], 1)
        self.assertEqual(first['Вн. номер'], '6474')

    def test_операторы_и_сводка(self):
        operators = list(self.book['Операторы'].iter_rows(min_row=2, values_only=True))
        self.assertEqual({row[0] for row in operators}, {'Жупан Аружан', 'Зинеден Аружан'})
        summary_cells = [row[0] for row in self.book['Сводка'].iter_rows(values_only=True)]
        self.assertIn('ГЛАВНОЕ', summary_cells)
        self.assertIn('ПО ТИПУ ЛИДА', summary_cells)
        self.assertIn('КАСАНИЯ ПО ДНЯМ', summary_cells)

    def test_контекст_первым_и_с_ловушками(self):
        context = [row[0] for row in self.book['Контекст'].iter_rows(values_only=True) if row]
        self.assertIn('ПРОЧТИ ПЕРВЫМ', context)
        self.assertIn('ЛОВУШКИ — ОБЯЗАТЕЛЬНО УЧЕСТЬ', context)
        text = ' '.join(str(cell) for row in self.book['Контекст'].iter_rows(values_only=True)
                        for cell in row if cell)
        self.assertIn('парк = iTaxi', text)
        self.assertIn('2 минуты', text)


class JsonTests(unittest.TestCase):

    def test_первый_ключ_контекст_и_касания_внутри_записей(self):
        rows = make_rows()
        summary = L.summarize(rows)
        payload = lead_report.build_json(
            rows, source_info=dict(L.SOURCES['amo'], key='amo'),
            period_from='2026-09-10', period_to='2026-09-11', touches_to='2026-09-12',
            summary=summary, blocks=2, full_url=True)
        self.assertEqual(list(payload)[0], 'КОНТЕКСТ_ДЛЯ_ИИ')
        text = json.dumps(payload, ensure_ascii=False, default=str)
        self.assertIn('Жупан Аружан', text)
        first = payload['записи'][0]
        self.assertEqual(first['касаний'], 2)
        self.assertEqual(first['касания'][0]['результат'], 'Разговор')
        self.assertEqual(first['moment'], '2026-09-10 10:00:00')

    def test_имя_файла(self):
        info = dict(L.SOURCES['crm_paid_hire'], key='crm_paid_hire')
        self.assertEqual(lead_report.report_filename(info, '2026-09-01', '2026-09-14'),
                         'Платный найм 01.09.2026-14.09.2026 + касания ОП.xlsx')
        self.assertEqual(lead_report.report_filename(info, '2026-09-14', '2026-09-14', 'json'),
                         'Платный найм 14.09.2026 + касания ОП.json')


if __name__ == '__main__':
    unittest.main()
