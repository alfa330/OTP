# -*- coding: utf-8 -*-
"""«Дисциплина» на настоящем Postgres: DISCIPLINE_TEST_PORT — порт локального кластера.

Чистые правила (лестница, плашка, статусы, права) проверяет
tests/test_discipline_rules.py. Здесь — что они ЗНАЧАТ на схеме раздела:
кто кого видит и кому может зафиксировать нарушение, как лестница ведёт
запись от тренинга до акта, что акт и особый контроль появляются сами, что
отмена сдвигает номера соседей и будит «проверить основание акта», что
напоминания не задваиваются, а журнал действий не правится.

Ни боевых учёток, ни импорта приложения: только 127.0.0.1, своя схема со
случайным именем, в конце она удаляется. Люди выдуманы.

    DISCIPLINE_TEST_PORT=55461 python -m pytest tests/test_discipline_postgres.py
"""
import contextlib
import os
import uuid
from datetime import date, datetime, timedelta

import pytest

from discipline import access, catalog, queries, report, schema, service, act_docx

PORT = os.environ.get('DISCIPLINE_TEST_PORT')
pytestmark = pytest.mark.skipif(not PORT, reason='needs local PostgreSQL (DISCIPLINE_TEST_PORT)')


@pytest.fixture(autouse=True)
def section_open_for_all(monkeypatch):
    """Здесь проверяются правила ТЗ для всех ролей. Выкладка «пока только
    супер-админам» (access.SUPER_ADMIN_ONLY) — отдельными тестами."""
    monkeypatch.setattr(access, 'SUPER_ADMIN_ONLY', False)

SUPER, ADMIN, HEAD_SZOV, HEAD_OP, SV_A, SV_B, OP_A, OP_B, OP_OP, TRAINEE_A, FIRED = (
    1, 2, 100, 200, 300, 301, 400, 401, 402, 403, 404)
SZOV, OP, TEZ, REMOTE = 1, 2, 3, 4
NOW = datetime(2026, 10, 20, 12, 0)
TODAY = NOW.date()

BASE_DDL = '''
    CREATE TABLE departments (id integer PRIMARY KEY, code text UNIQUE, name text NOT NULL,
                              head_user_id integer, is_active boolean NOT NULL DEFAULT TRUE);
    CREATE TABLE users (id integer PRIMARY KEY, name text NOT NULL, login text, role text NOT NULL,
                        status text NOT NULL DEFAULT 'working', department_id integer, hire_date date,
                        telegram_id bigint, avatar_bucket text, avatar_blob_path text);
    CREATE TABLE directions (id integer PRIMARY KEY, name text NOT NULL);
    CREATE TABLE groups (id integer PRIMARY KEY, name text NOT NULL, department_id integer,
                         direction_id integer, status text NOT NULL DEFAULT 'active');
    CREATE TABLE group_operator_memberships (id serial PRIMARY KEY, group_id integer, operator_id integer,
                                             start_date date NOT NULL, end_date date);
    CREATE TABLE group_supervisor_memberships (id serial PRIMARY KEY, group_id integer, supervisor_id integer,
                                               start_date date NOT NULL, end_date date);
    CREATE TABLE trainings (id serial PRIMARY KEY, operator_id integer NOT NULL, training_date date NOT NULL,
                            start_time time NOT NULL, end_time time NOT NULL, reason text NOT NULL, comment text,
                            created_by integer, count_in_hours boolean NOT NULL DEFAULT TRUE,
                            UNIQUE (operator_id, training_date, start_time, end_time));
    CREATE TABLE work_shifts (id serial PRIMARY KEY, operator_id integer, shift_date date, start_time time,
                              end_time time);
    INSERT INTO departments (id, code, name, head_user_id) VALUES
        (1, 'szov', 'СЗоВ', 100), (2, 'op', 'Отдел продаж', 200), (3, 'tez', 'Тез КЦ', NULL),
        (4, 'remote_cc', 'Удалённый КЦ', 100);
    INSERT INTO users (id, name, login, role, status, department_id, telegram_id) VALUES
        (1, 'Супер Админов', 'super', 'super_admin', 'working', NULL, NULL),
        (2, 'Админ Простой', 'admin', 'admin', 'working', 1, NULL),
        (100, 'Глава Сзовова', 'head1', 'admin', 'working', 1, 9100),
        (200, 'Глава Продаж', 'head2', 'sv', 'working', 2, NULL),
        (300, 'Супервайзер Альфа', 'sva', 'sv', 'working', 1, NULL),
        (301, 'Супервайзер Бета', 'svb', 'sv', 'working', 1, NULL),
        (400, 'Оператор Анна', 'opa', 'operator', 'working', 1, NULL),
        (401, 'Оператор Борис', 'opb', 'operator', 'working', 1, NULL),
        (402, 'Оператор Продаж', 'opop', 'operator', 'working', 2, NULL),
        (403, 'Стажёр Вера', 'tra', 'trainee', 'working', 1, NULL),
        (404, 'Уволенный Глеб', 'opf', 'operator', 'fired', 1, NULL),
        (500, 'Свидетель Один', 'w1', 'operator', 'working', 1, NULL),
        (501, 'Свидетель Два', 'w2', 'operator', 'working', 1, NULL);
    INSERT INTO directions (id, name) VALUES (7, 'Линия водителей');
    INSERT INTO groups (id, name, department_id, direction_id) VALUES
        (10, 'Группа 1', 1, 7), (11, 'Группа 2', 1, 7), (20, 'Группа ОП', 2, NULL);
    INSERT INTO group_operator_memberships (group_id, operator_id, start_date) VALUES
        (10, 400, '2026-01-01'), (11, 401, '2026-01-01'), (20, 402, '2026-01-01'), (10, 403, '2026-01-01'),
        (10, 404, '2026-01-01');
    INSERT INTO group_supervisor_memberships (group_id, supervisor_id, start_date) VALUES
        (10, 300, '2026-01-01'), (11, 301, '2026-01-01');
    INSERT INTO work_shifts (operator_id, shift_date, start_time, end_time) VALUES
        (400, '2026-10-19', '22:00', '06:00'), (400, '2026-10-20', '09:00', '18:00');
'''


class _Db:
    def __init__(self, name):
        self.name = name

    def connect(self):
        import psycopg2
        conn = psycopg2.connect(host='127.0.0.1', port=int(PORT), user='postgres', dbname='postgres',
                                connect_timeout=3)
        with conn.cursor() as cur:
            cur.execute('SET search_path TO ' + self.name)
            cur.execute("SET statement_timeout TO '20s'")
        conn.commit()
        return conn

    @contextlib.contextmanager
    def cursor(self):
        conn = self.connect()
        cur = conn.cursor()
        try:
            yield cur
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()


@pytest.fixture
def db():
    name = 't_discipline_' + uuid.uuid4().hex
    base = _Db(name)
    admin = base.connect()
    try:
        with admin.cursor() as cur:
            cur.execute('CREATE SCHEMA ' + name)
            cur.execute('SET search_path TO ' + name)
            cur.execute(BASE_DDL)
            schema.init_discipline_schema(cur)
            # Повторный запуск схемы — как каждый рестарт прода: идемпотентно.
            schema.init_discipline_schema(cur)
        admin.commit()
    except Exception:
        admin.rollback()
        admin.close()
        raise
    try:
        yield base
    finally:
        with admin.cursor() as cur:
            cur.execute('DROP SCHEMA ' + name + ' CASCADE')
        admin.commit()
        admin.close()


def ctx_of(cur, user_id):
    return queries.load_access_context(cur, user_id, TODAY)


def upload(name='scan.pdf'):
    return {'name': name, 'content_type': 'application/pdf', 'size': 10, 'sha256': 'a' * 64,
            'bucket': 'bucket', 'blob_path': 'discipline/test/%s' % uuid.uuid4().hex}


def type_id(cur, department_id, title='Сон на рабочем месте'):
    queries.ensure_department_defaults(cur, department_id)
    cur.execute('SELECT id FROM discipline_types WHERE department_id = %s AND title = %s', (department_id, title))
    return cur.fetchone()[0]


def measure_id(cur, department_id, title):
    queries.ensure_department_defaults(cur, department_id)
    cur.execute('SELECT id FROM discipline_measures WHERE department_id = %s AND title = %s',
                (department_id, title))
    return cur.fetchone()[0]


DESCRIPTION = 'Оператор спал на рабочем месте во время смены, видно по камере.'


def use_sample_ladder(cur, department=SZOV):
    """Лестница со скриншота заказчика — как её выбрал бы руководитель до первой записи."""
    settings = catalog.normalize_settings({'ladder': list(catalog.LADDER_SAMPLE)})
    queries.save_settings(cur, department, settings, HEAD_SZOV)


def create(cur, actor, operator, at, *, title='Сон на рабочем месте', department=SZOV, **extra):
    ctx = ctx_of(cur, actor)
    payload = {'operator_id': operator, 'type_id': type_id(cur, department, title),
               'occurred_at': at.strftime('%Y-%m-%d %H:%M'), 'description': DESCRIPTION}
    if service.shift_label_at(cur, operator, at) is None:
        payload['shift_label'] = '09:00–18:00'
    payload.update(extra)
    return service.create_record(cur, ctx, payload, [], now=NOW, today=TODAY)


def enriched(cur, actor, record_id):
    ctx = ctx_of(cur, actor)
    record = service.load_record(cur, ctx, record_id, TODAY)
    return service.enrich(cur, ctx, [record], NOW)[0]


def bell_titles(cur, user_id):
    cur.execute('SELECT title FROM discipline_notifications WHERE user_id = %s AND read_at IS NULL ORDER BY id',
                (user_id,))
    return [row[0] for row in cur.fetchall()]


# ── Доступ ───────────────────────────────────────────────────────────────────

def test_section_circle(db):
    with db.cursor() as cur:
        assert access.can_open_section(ctx_of(cur, SUPER))
        assert access.scope(ctx_of(cur, SUPER)) == access.SCOPE_ALL
        # «Если админ не руководитель этих трёх отделов — он не может просматривать раздел».
        assert not access.can_open_section(ctx_of(cur, ADMIN))
        head = ctx_of(cur, HEAD_SZOV)
        assert access.can_open_section(head) and access.managed_department_ids(head) == {SZOV}
        # Глава ОП с ролью СВ — руководитель своего отдела, а не рядовой СВ.
        head_op = ctx_of(cur, HEAD_OP)
        assert access.is_manager(head_op) and not access.is_supervisor(head_op)
        assert access.scope(ctx_of(cur, SV_A)) == access.SCOPE_GROUPS
        assert access.scope(ctx_of(cur, OP_A)) == access.SCOPE_OWN
        assert access.requires_sensitive_qr(ctx_of(cur, OP_A))
        assert not access.requires_sensitive_qr(ctx_of(cur, SV_A))


def test_supervisor_sees_only_own_group(db):
    with db.cursor() as cur:
        found = {row['id'] for row in queries.search_operators(cur, ctx_of(cur, SV_A), TODAY)}
        # Своя группа: оператор и стажёр; уволенный и чужая группа — нет.
        assert found == {OP_A, TRAINEE_A}
        # Глава СЗоВ ищет по всему своему отделу; оператор ОП в выдачу не попадает.
        assert {row['id'] for row in queries.search_operators(cur, ctx_of(cur, HEAD_SZOV), TODAY, 'опер')} == {
            OP_A, OP_B}
        with pytest.raises(service.DisciplineError) as error:
            create(cur, SV_B, OP_A, NOW - timedelta(hours=1))
        assert error.value.status == 403
        with pytest.raises(service.DisciplineError):
            create(cur, ADMIN, OP_A, NOW - timedelta(hours=1))


def test_head_cannot_touch_other_department(db):
    with db.cursor() as cur:
        with pytest.raises(service.DisciplineError) as error:
            create(cur, HEAD_SZOV, OP_OP, NOW - timedelta(hours=1), department=OP)
        assert error.value.status == 403
        record_id = create(cur, HEAD_OP, OP_OP, NOW - timedelta(hours=1), department=OP)
        # Глава СЗоВ чужую запись не видит, супер-админ — видит.
        with pytest.raises(service.DisciplineError):
            service.load_record(cur, ctx_of(cur, HEAD_SZOV), record_id, TODAY)
        assert service.load_record(cur, ctx_of(cur, SUPER), record_id, TODAY)['id'] == record_id


# ── Лестница по умолчанию — как в ТЗ: критерии приёмки 2–4 ─────────────────

def test_default_ladder_passes_acceptance_2_3_4(db):
    with db.cursor() as cur:
        first = create(cur, SV_A, OP_A, NOW - timedelta(days=10))
        assert enriched(cur, SV_A, first)['step_title'] == 'Объясн. №1'
        assert bell_titles(cur, OP_A) == ['Нужна объяснительная']
        second = create(cur, SV_A, OP_A, NOW - timedelta(days=8))
        assert enriched(cur, SV_A, second)['step_title'] == 'Объясн. №2'
        # Критерий 2: при двух нарушениях — красная плашка «потребуется акт».
        step = service.operator_context(cur, ctx_of(cur, SV_A), OP_A, NOW, TODAY)
        assert step['plate']['tone'] == 'red' and 'потребуется акт' in step['plate']['text']
        third = create(cur, SV_A, OP_A, NOW - timedelta(days=6))
        item = enriched(cur, SV_A, third)
        assert (item['ordinal'], item['action'], item['state']) == (3, 'act', 'explanation_requested')
        # Критерий 4: контроль — сам, без действий человека.
        control = queries.active_controls(cur, [OP_A])[OP_A]
        assert control['source'] == 'act' and control['act_id'] == item['act']['id']
        # Критерий 3: без подписанного акта не закрыть, даже с объяснительной.
        service.add_explanation(cur, ctx_of(cur, OP_A), service.load_record(cur, ctx_of(cur, OP_A), third, TODAY),
                                upload(), now=NOW)
        assert [b['code'] for b in enriched(cur, HEAD_SZOV, third)['close_blockers']] == ['act']
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), third, TODAY, for_update=True)
        with pytest.raises(service.DisciplineError) as error:
            service.close_record(cur, ctx_of(cur, HEAD_SZOV), record, now=NOW)
        assert error.value.code == 'CLOSE_BLOCKED'


# ── Лестница со скриншота: тренинг → объясн. №1 → №2 → №3 + акт ─────────────

def test_ladder_walks_to_act_and_control(db):
    with db.cursor() as cur:
        use_sample_ladder(cur)
        first = create(cur, SV_A, OP_A, NOW - timedelta(days=10))
        item = enriched(cur, SV_A, first)
        assert (item['ordinal'], item['action'], item['state']) == (1, 'training', 'training_required')
        assert item['step_title'] == 'Тренинг'
        # Оператору — «Зафиксировано нарушение», объяснительная на ступени тренинга не нужна.
        assert bell_titles(cur, OP_A) == ['Зафиксировано нарушение']

        service.record_training(cur, ctx_of(cur, SV_A), service.load_record(cur, ctx_of(cur, SV_A), first, TODAY),
                                {'date': (TODAY - timedelta(days=9)).isoformat(), 'start': '10:00', 'end': '10:30'},
                                now=NOW, today=TODAY)
        cur.execute('SELECT reason, count_in_hours FROM trainings')
        assert cur.fetchall() == [('Дисциплинарный тренинг', True)]
        assert enriched(cur, SV_A, first)['state'] == 'training_done'

        second = create(cur, SV_A, OP_A, NOW - timedelta(days=8))
        item = enriched(cur, SV_A, second)
        assert (item['ordinal'], item['action'], item['explanation_no']) == (2, 'explanation', 1)
        assert item['state'] == 'explanation_requested'
        assert item['step_title'] == 'Объясн. №1'
        service.add_explanation(cur, ctx_of(cur, OP_A), service.load_record(cur, ctx_of(cur, OP_A), second, TODAY),
                                upload(), now=NOW)
        assert enriched(cur, SV_A, second)['state'] == 'explanation_received'

        third = create(cur, SV_A, OP_A, NOW - timedelta(days=6))
        assert enriched(cur, SV_A, third)['step_title'] == 'Объясн. №2'

        fourth = create(cur, SV_A, OP_A, NOW - timedelta(days=4))
        item = enriched(cur, SV_A, fourth)
        assert (item['ordinal'], item['action'], item['step_title']) == (4, 'act', 'Объясн. №3 + акт')
        assert item['act'] and item['act']['is_trigger'] and not item['act']['has_file']
        # Черновик акта связан со всеми записями окна (п. 9).
        act = queries.get_act(cur, item['act']['id'])
        assert sorted(act['record_ids']) == sorted([first, second, third, fourth])
        # Особый контроль поставлен сам (критерий 4) и виден в сводке.
        control = queries.active_controls(cur, [OP_A])[OP_A]
        assert control['source'] == 'act' and control['act_id'] == act['id']
        assert control['end_date'] == (NOW - timedelta(days=4)).date() + timedelta(days=30)
        # Руководителю — «нужен акт».
        assert any('нужен акт' in title for title in bell_titles(cur, HEAD_SZOV))

        # Закрыть без подписанного акта нельзя (критерий 3).
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), fourth, TODAY, for_update=True)
        with pytest.raises(service.DisciplineError) as error:
            service.close_record(cur, ctx_of(cur, HEAD_SZOV), record, now=NOW)
        assert error.value.code == 'CLOSE_BLOCKED'


def test_pre_act_warns_group_supervisor(db):
    with db.cursor() as cur:
        for days in (10, 8):
            create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=days))
        # Вторая запись — за шаг до акта: СВ группы получает предупреждение (п. 7.3).
        assert any('следующее нарушение — акт' in title for title in bell_titles(cur, SV_A))
        step = service.operator_context(cur, ctx_of(cur, SV_A), OP_A, NOW, TODAY)
        assert step['plate']['tone'] == 'red'
        assert 'потребуется акт' in step['plate']['text']


def test_violation_under_control_requires_measure(db):
    with db.cursor() as cur:
        for days in (12, 10, 8):
            create(cur, SV_A, OP_A, NOW - timedelta(days=days))
        context = service.operator_context(cur, ctx_of(cur, SV_A), OP_A, NOW, TODAY)
        assert context['measure_required'] and context['plate']['title'].startswith('Оператор на особом контроле')
        with pytest.raises(service.DisciplineError) as error:
            create(cur, SV_A, OP_A, NOW - timedelta(hours=2))
        assert error.value.code == 'MEASURE_REQUIRED'
        fine = measure_id(cur, SZOV, 'Лишение переменной части (фиксированная сумма)')
        with pytest.raises(service.DisciplineError):
            create(cur, SV_A, OP_A, NOW - timedelta(hours=2), measure_id=fine)
        fifth = create(cur, SV_A, OP_A, NOW - timedelta(hours=2), measure_id=fine, measure_amount='15000')
        item = enriched(cur, SV_A, fifth)
        assert item['under_control'] and item['measure_state'] == 'proposed'
        # Четвёртое по лестнице снова «акт», но под контролем после акта второй акт не
        # составляется: ступень — объяснительная, закрывает утверждённая мера (п. 7.3, 18.2).
        assert item['ladder_action'] == 'act' and item['action'] == 'explanation'
        assert item['act'] is None
        cur.execute('SELECT COUNT(*) FROM discipline_acts')
        assert cur.fetchone()[0] == 1
        assert [b['code'] for b in item['close_blockers']] == ['explanation', 'measure']
        assert item['permissions']['can_close'] is False
        # СВ закрыть не может, руководитель утверждает — и мера уходит в выгрузку на удержания.
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), fifth, TODAY, for_update=True)
        service.approve_measure(cur, ctx_of(cur, HEAD_SZOV), record, now=NOW)
        item = enriched(cur, HEAD_SZOV, fifth)
        assert item['measure_state'] == 'approved'
        rows, totals = report.report_measures([item])
        assert rows[0]['measure_amount'] == 15000.0 and totals[0]['count'] == 1


def test_cancel_shifts_numbers_and_flags_act_basis(db):
    with db.cursor() as cur:
        ids = [create(cur, SV_A, OP_A, NOW - timedelta(days=days)) for days in (12, 10, 8)]
        act_id = enriched(cur, SV_A, ids[2])['act']['id']
        record = service.load_record(cur, ctx_of(cur, SV_A), ids[1], TODAY, for_update=True)
        with pytest.raises(service.DisciplineError) as error:
            service.cancel_record(cur, ctx_of(cur, SV_A), record, 'ошибка', now=NOW)
        assert error.value.status == 403
        service.cancel_record(cur, ctx_of(cur, HEAD_SZOV), record, 'Ошибочно зафиксировано', now=NOW)
        item = enriched(cur, HEAD_SZOV, ids[2])
        # Номер третьей записи сдвинулся на второй, акт остался, но основание под вопросом.
        assert item['ordinal'] == 2 and item['action'] == 'explanation'
        assert item['act']['id'] == act_id and item['act']['basis_check'] is True
        assert [t for t in bell_titles(cur, HEAD_SZOV) if t.startswith('Проверить основание акта')] == [
            'Проверить основание акта %s' % item['act']['number']]
        # Повторная отмена другой записи о том же акте второй раз не будит.
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), ids[0], TODAY, for_update=True)
        service.cancel_record(cur, ctx_of(cur, HEAD_SZOV), record, 'Ошибочно зафиксировано', now=NOW)
        assert len([t for t in bell_titles(cur, HEAD_SZOV) if t.startswith('Проверить основание акта')]) == 1


def test_supervisor_cannot_edit_colleagues_record(db):
    with db.cursor() as cur:
        record_id = create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=1))
        record = service.load_record(cur, ctx_of(cur, SV_A), record_id, TODAY, for_update=True)
        with pytest.raises(service.DisciplineError) as error:
            service.edit_record(cur, ctx_of(cur, SV_A), record, {'description': DESCRIPTION + ' Дополнено.'}, now=NOW)
        assert error.value.status == 403
        own = create(cur, SV_A, OP_A, NOW - timedelta(hours=3))
        record = service.load_record(cur, ctx_of(cur, SV_A), own, TODAY, for_update=True)
        # Смена подставилась из графика сама, а править её свою запись СВ может.
        assert record['shift_label'] == '09:00–18:00'
        assert service.edit_record(cur, ctx_of(cur, SV_A), record, {'shift_label': '10:00–19:00'}, now=NOW)
        cur.execute("SELECT changes FROM discipline_audit WHERE action = 'update_record'")
        assert cur.fetchone()[0]['shift_label'] == ['09:00–18:00', '10:00–19:00']


def test_refusal_needs_two_witnesses_and_counts(db):
    with db.cursor() as cur:
        create(cur, SV_A, OP_A, NOW - timedelta(days=5))
        second = create(cur, SV_A, OP_A, NOW - timedelta(days=3))
        record = service.load_record(cur, ctx_of(cur, SV_A), second, TODAY, for_update=True)
        with pytest.raises(service.DisciplineError):
            service.record_refusal(cur, ctx_of(cur, SV_A), record, upload(), [500], now=NOW)
        service.record_refusal(cur, ctx_of(cur, SV_A), record, upload(), [500, 501], now=NOW)
        item = enriched(cur, SV_A, second)
        assert item['state'] == 'explanation_refused' and item['ordinal'] == 2
        assert any('отказ от объяснения' in title for title in bell_titles(cur, HEAD_SZOV))


def test_files_are_versioned_not_deleted(db):
    with db.cursor() as cur:
        create(cur, SV_A, OP_A, NOW - timedelta(days=5))
        second = create(cur, SV_A, OP_A, NOW - timedelta(days=3))
        for _ in range(2):
            record = service.load_record(cur, ctx_of(cur, SV_A), second, TODAY, for_update=True)
            service.add_explanation(cur, ctx_of(cur, SV_A), record, upload(), now=NOW)
        rows = queries.files_for(cur, record_ids=[second])
        assert len(rows) == 2 and sum(1 for r in rows if r['superseded_at'] is None) == 1


def test_audit_is_append_only(db):
    with db.cursor() as cur:
        create(cur, SV_A, OP_A, NOW - timedelta(days=1))
        cur.execute('SAVEPOINT a')
        with pytest.raises(Exception):
            cur.execute('UPDATE discipline_audit SET actor_name = %s', ('Подделка',))
        cur.execute('ROLLBACK TO SAVEPOINT a')
        with pytest.raises(Exception):
            cur.execute('DELETE FROM discipline_audit')
        cur.execute('ROLLBACK TO SAVEPOINT a')


def test_reminders_are_not_duplicated(db):
    with db.cursor() as cur:
        create(cur, SV_A, OP_A, NOW - timedelta(days=10))
        late = create(cur, SV_A, OP_A, NOW - timedelta(days=9))
        cur.execute('UPDATE discipline_records SET created_at = %s WHERE id = %s', (NOW - timedelta(days=9), late))
        first = service.run_reminders(cur, now=NOW, today=TODAY)
        second = service.run_reminders(cur, now=NOW, today=TODAY)
        assert first > 0 and second == 0
        assert any('не предоставлена в срок' in title for title in bell_titles(cur, HEAD_SZOV))
        assert enriched(cur, SV_A, late)['explanation_overdue'] is True


def test_manual_control_and_decisions(db):
    with db.cursor() as cur:
        control_id = service.create_control(cur, ctx_of(cur, HEAD_SZOV), {
            'operator_id': OP_B, 'source': 'probation', 'end_date': (TODAY + timedelta(days=2)).isoformat()},
            now=NOW, today=TODAY)
        with pytest.raises(service.DisciplineError) as error:
            service.create_control(cur, ctx_of(cur, HEAD_SZOV), {
                'operator_id': OP_B, 'source': 'manual', 'reason': 'Повторно',
                'end_date': (TODAY + timedelta(days=5)).isoformat()}, now=NOW, today=TODAY)
        assert error.value.code == 'CONTROL_EXISTS'
        rows = service.control_rows(cur, ctx_of(cur, HEAD_SZOV), service.visible_controls(
            cur, ctx_of(cur, HEAD_SZOV), TODAY), TODAY)
        # До конца 2 дня — жёлтая строка (п. 10.2).
        assert rows[0]['tone'] == 'yellow' and rows[0]['source_label'] == 'Испытательный срок'
        # СВ чужой группы контроль оператора не видит.
        assert service.visible_controls(cur, ctx_of(cur, SV_A), TODAY) == []
        assert len(service.visible_controls(cur, ctx_of(cur, SV_B), TODAY)) == 1
        service.decide_control(cur, ctx_of(cur, HEAD_SZOV), control_id, {'decision': 'extended', 'days': 10},
                               now=NOW, today=TODAY)
        assert queries.get_control(cur, control_id)['end_date'] == TODAY + timedelta(days=12)
        service.decide_control(cur, ctx_of(cur, HEAD_SZOV), control_id, {'decision': 'removed'},
                               now=NOW, today=TODAY)
        assert queries.active_controls(cur, [OP_B]) == {}


def test_operator_sees_only_public_fields(db):
    with db.cursor() as cur:
        record_id = create(cur, SV_A, OP_A, NOW - timedelta(days=1))
        op_ctx = ctx_of(cur, OP_A)
        item = enriched(cur, OP_A, record_id)
        public = service.public_record(item, op_ctx)
        assert 'description' not in public and 'created_by_name' not in public
        assert public['type_title'] == 'Сон на рабочем месте'
        assert [r['id'] for r in queries.list_records(cur, ctx_of(cur, OP_B), TODAY, {})] == []


def test_shift_label_covers_night_shift(db):
    with db.cursor() as cur:
        assert service.shift_label_at(cur, OP_A, datetime(2026, 10, 20, 3, 0)) == '22:00–06:00'
        assert service.shift_label_at(cur, OP_A, datetime(2026, 10, 20, 10, 0)) == '09:00–18:00'


def test_act_docx_builds(db):
    with db.cursor() as cur:
        ids = [create(cur, SV_A, OP_A, NOW - timedelta(days=days)) for days in (12, 10, 8)]
        act = queries.get_act(cur, enriched(cur, SV_A, ids[2])['act']['id'])
        operator = queries.operator_profiles(cur, [OP_A], TODAY)[OP_A]
        records = [queries.get_record(cur, rid, TODAY) for rid in act['record_ids']]
        data = act_docx.build(act, operator, records, position='Оператор', today=TODAY, composed_by='Глава')
        assert data[:2] == b'PK' and len(data) > 3000
        # Шаблон отдела: метки, разрезанные Word на прогоны, тоже находятся.
        from io import BytesIO
        from docx import Document
        template = Document()
        paragraph = template.add_paragraph()
        paragraph.add_run('Акт {{НОМЕР')
        paragraph.add_run('_АКТА}} на {{ФИО}}')
        buffer = BytesIO()
        template.save(buffer)
        filled = Document(BytesIO(act_docx.build(act, operator, records, template_bytes=buffer.getvalue(),
                                                 today=TODAY)))
        assert filled.paragraphs[0].text == 'Акт %s на Оператор Анна' % act['number']


def test_concurrent_violations_on_act_step_are_serialized(db):
    """Два СВ одновременно фиксируют нарушение одному человеку на ступени акта.
    Без очереди на человека обе записи стали бы «третьими», акт составился бы
    дважды, а второй контроль упал бы на уникальном индексе (500)."""
    import threading

    with db.cursor() as cur:
        for days in (12, 10):
            create(cur, SV_A, OP_A, NOW - timedelta(days=days))
    results, barrier = [], threading.Barrier(2)

    def worker(actor):
        conn = db.connect()
        cur = conn.cursor()
        try:
            barrier.wait()
            record_id = create(cur, actor, OP_A, NOW - timedelta(hours=1))
            conn.commit()
            results.append(('ok', record_id))
        except service.DisciplineError as exc:
            conn.rollback()
            results.append(('refused', exc.code))
        except Exception as exc:  # noqa: BLE001
            conn.rollback()
            results.append(('crash', type(exc).__name__))
        finally:
            conn.close()

    threads = [threading.Thread(target=worker, args=(actor,)) for actor in (SV_A, HEAD_SZOV)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(30)
    assert sorted(kind for kind, _ in results) == ['ok', 'refused'], results
    # Второй ждал первого и увидел: человек уже на контроле — нужна мера.
    assert [code for kind, code in results if kind == 'refused'] == ['MEASURE_REQUIRED']
    with db.cursor() as cur:
        cur.execute('SELECT COUNT(*) FROM discipline_acts')
        assert cur.fetchone()[0] == 1
        cur.execute('SELECT COUNT(*) FROM discipline_controls')
        assert cur.fetchone()[0] == 1


def test_refusal_only_where_explanation_is_needed(db):
    with db.cursor() as cur:
        use_sample_ladder(cur)
        first = create(cur, SV_A, OP_A, NOW - timedelta(days=2))
        record = service.load_record(cur, ctx_of(cur, SV_A), first, TODAY, for_update=True)
        with pytest.raises(service.DisciplineError):
            service.record_refusal(cur, ctx_of(cur, SV_A), record, upload(), [500, 501], now=NOW)


# ── Лестница на живых данных ─────────────────────────────────────────────────

def _create_at(cur, actor, operator, now, at, **extra):
    """Фиксация «в другой момент»: now — когда нажали «Зафиксировать»."""
    payload = {'operator_id': operator, 'type_id': type_id(cur, SZOV, 'Сон на рабочем месте'),
               'occurred_at': at.strftime('%Y-%m-%d %H:%M'), 'description': DESCRIPTION,
               'shift_label': '09:00–18:00'}
    payload.update(extra)
    return service.create_record(cur, ctx_of(cur, actor), payload, [], now=now, today=now.date())


def _change_ladder(cur, ladder, now):
    before = queries.get_settings(cur, SZOV)
    after = catalog.normalize_settings({'ladder': list(ladder)}, base=before)
    assert service.save_counting_change(cur, SZOV, before, after, now=now, user_id=HEAD_SZOV)
    queries.save_settings(cur, SZOV, after, HEAD_SZOV)


def test_ladder_change_applies_to_new_records_only(db):
    """Смена лестницы не переписывает прошлое: запись, зафиксированная по старой
    лестнице, не требует акта, которого не было, а у прошлого акта не загорается
    «проверить основание». Новая запись идёт уже по новой."""
    with db.cursor() as cur:
        use_sample_ladder(cur)
        ids = [create(cur, SV_A, OP_A, NOW - timedelta(days=days)) for days in (12, 10, 8)]
        assert enriched(cur, SV_A, ids[2])['step_title'] == 'Объясн. №2'
        _change_ladder(cur, catalog.LADDER_TZ, NOW + timedelta(hours=1))
        item = enriched(cur, HEAD_SZOV, ids[2])
        assert (item['action'], item['act']) == ('explanation', None)
        assert 'act' not in [b['code'] for b in item['close_blockers']]
        later = NOW + timedelta(hours=2)
        fourth = _create_at(cur, SV_A, OP_A, later, later - timedelta(minutes=10))
        item = enriched(cur, HEAD_SZOV, fourth)
        assert item['action'] == 'act' and item['act']['is_trigger']
        assert queries.active_controls(cur, [OP_A])[OP_A]['act_id'] == item['act']['id']
        cur.execute('SELECT rules, valid_from FROM discipline_counting_rules ORDER BY valid_from')
        rows = cur.fetchall()
        assert [r[0]['ladder'] for r in rows] == [list(catalog.LADDER_SAMPLE), list(catalog.LADDER_TZ)]


def test_ladder_change_keeps_past_act_basis(db):
    with db.cursor() as cur:
        ids = [create(cur, SV_A, OP_A, NOW - timedelta(days=days)) for days in (12, 10, 8)]
        assert enriched(cur, HEAD_SZOV, ids[2])['act']['basis_check'] is False
        _change_ladder(cur, catalog.LADDER_SAMPLE, NOW + timedelta(hours=1))
        item = enriched(cur, HEAD_SZOV, ids[2])
        assert item['action'] == 'act' and item['act']['basis_check'] is False


def test_backdated_record_moves_later_one_to_act_with_control(db):
    with db.cursor() as cur:
        create(cur, SV_A, OP_A, NOW - timedelta(days=6))
        later = create(cur, SV_A, OP_A, NOW - timedelta(days=4))
        assert queries.active_controls(cur, [OP_A]) == {}
        # Бумага о нарушении пришла позже: оно случилось раньше обоих.
        create(cur, SV_A, OP_A, NOW - timedelta(days=8))
        item = enriched(cur, HEAD_SZOV, later)
        assert item['action'] == 'act' and item['act']['is_trigger']
        control = queries.active_controls(cur, [OP_A])[OP_A]
        assert control['act_id'] == item['act']['id'] and control['source'] == 'act'
        assert control['start_date'] == (NOW - timedelta(days=4)).date()
        assert any('нужен акт' in title for title in bell_titles(cur, HEAD_SZOV))
        # Следующее — уже нарушение на контроле: нужна мера, второго акта нет.
        with pytest.raises(service.DisciplineError) as error:
            create(cur, SV_A, OP_A, NOW - timedelta(hours=1))
        assert error.value.code == 'MEASURE_REQUIRED'


def test_act_during_probation_is_linked_to_that_control(db):
    with db.cursor() as cur:
        control_id = service.create_control(cur, ctx_of(cur, HEAD_SZOV), {
            'operator_id': OP_A, 'source': 'probation', 'start_date': (TODAY - timedelta(days=20)).isoformat(),
            'end_date': (TODAY + timedelta(days=1)).isoformat()}, now=NOW, today=TODAY)
        remark = measure_id(cur, SZOV, 'Устное замечание')
        ids = [create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=days), measure_id=remark) for days in (6, 4, 2)]
        item = enriched(cur, HEAD_SZOV, ids[2])
        assert item['action'] == 'act' and item['act']['is_trigger']
        control = queries.get_control(cur, control_id)
        # Контроль один, теперь он держится на акте — не короче срока «после акта»,
        # и основание «После акта»: снять его, не заметив акт, нельзя.
        assert control['act_id'] == item['act']['id']
        assert control['end_date'] == (NOW - timedelta(days=2)).date() + timedelta(days=30)
        assert control['source'] == 'act' and control['reason'].startswith('Акт ' + item['act']['number'])
        assert 'Испытательный срок' in control['reason']
        # Следующее нарушение — мера, а не второй акт.
        fourth = create(cur, HEAD_SZOV, OP_A, NOW - timedelta(hours=1), measure_id=remark)
        assert enriched(cur, HEAD_SZOV, fourth)['action'] == 'explanation'
        cur.execute('SELECT COUNT(*) FROM discipline_acts')
        assert cur.fetchone()[0] == 1


def test_edit_that_drops_act_basis_tells_the_head(db):
    with db.cursor() as cur:
        ids = [create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=days)) for days in (12, 10, 8)]
        act = enriched(cur, HEAD_SZOV, ids[2])['act']
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), ids[0], TODAY, for_update=True)
        service.edit_record(cur, ctx_of(cur, HEAD_SZOV), record,
                            {'occurred_at': (NOW - timedelta(days=40)).strftime('%Y-%m-%d %H:%M')}, now=NOW)
        item = enriched(cur, HEAD_SZOV, ids[2])
        assert item['act']['basis_check'] is True
        # Задача руководителю — та же, что после отмены (п. 7.2), и одна.
        titles = bell_titles(cur, HEAD_SZOV)
        assert titles.count('Проверить основание акта %s' % act['number']) == 1


def test_plate_number_matches_journal(db):
    """Плашка в форме и журнал после сохранения называют одну и ту же объяснительную."""
    with db.cursor() as cur:
        use_sample_ladder(cur)
        for days in (40, 25, 10):
            create(cur, SV_A, OP_A, NOW - timedelta(days=days))
        preview = service.operator_context(cur, ctx_of(cur, SV_A), OP_A, NOW, TODAY)['step']
        record_id = create(cur, SV_A, OP_A, NOW)
        item = enriched(cur, SV_A, record_id)
        assert (preview['ordinal'], preview['explanation_no']) == (item['ordinal'], item['explanation_no'])
        assert preview['title'] == item['step_title']


# ── Меры ─────────────────────────────────────────────────────────────────────

def _under_control_record(cur, actor=SV_A, **extra):
    """Запись на особом контроле после акта: три нарушения, затем четвёртое с мерой."""
    for days in (12, 10, 8):
        create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=days))
    return create(cur, actor, OP_A, NOW - timedelta(hours=2), **extra)


def test_supervisor_cannot_overwrite_approved_measure(db):
    with db.cursor() as cur:
        reprimand = measure_id(cur, SZOV, 'Выговор')
        record_id = _under_control_record(cur, measure_id=reprimand)
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), record_id, TODAY, for_update=True)
        service.approve_measure(cur, ctx_of(cur, HEAD_SZOV), record, now=NOW)
        remark = measure_id(cur, SZOV, 'Устное замечание')
        record = service.load_record(cur, ctx_of(cur, SV_A), record_id, TODAY, for_update=True)
        assert enriched(cur, SV_A, record_id)['permissions']['can_propose_measure'] is False
        with pytest.raises(service.DisciplineError) as error:
            service.set_measure(cur, ctx_of(cur, SV_A), record, {'measure_id': remark}, now=NOW)
        assert error.value.status == 403
        item = enriched(cur, HEAD_SZOV, record_id)
        assert (item['measure_title'], item['measure_state']) == ('Выговор', 'approved')


def test_operator_sees_measure_only_when_applied(db):
    with db.cursor() as cur:
        fine = measure_id(cur, SZOV, 'Лишение переменной части (фиксированная сумма)')
        record_id = _under_control_record(cur, measure_id=fine, measure_amount='15000')
        public = service.public_record(enriched(cur, OP_A, record_id), ctx_of(cur, OP_A))
        assert (public['measure_title'], public['measure_amount'], public['measure_state']) == (None, None, None)
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), record_id, TODAY, for_update=True)
        service.approve_measure(cur, ctx_of(cur, HEAD_SZOV), record, now=NOW)
        public = service.public_record(enriched(cur, OP_A, record_id), ctx_of(cur, OP_A))
        assert public['measure_title'].startswith('Лишение') and public['measure_amount'] == 15000


def test_measure_without_approval_applies_at_once_but_money_waits(db):
    with db.cursor() as cur:
        remark = measure_id(cur, SZOV, 'Устное замечание')
        record_id = _under_control_record(cur, measure_id=remark)
        item = enriched(cur, SV_A, record_id)
        # П. 5.2: устное замечание утверждения не требует; закрывает запись на
        # контроле всё равно руководитель (п. 10.4).
        assert item['measure_state'] == 'approved' and item['state'] == 'explanation_requested'
        assert item['permissions']['can_close'] is False
        # Мера с деньгами ждёт руководителя, даже если в справочнике сняли «требует утверждения».
        fine = measure_id(cur, SZOV, 'Лишение переменной части (фиксированная сумма)')
        cur.execute('UPDATE discipline_measures SET needs_approval = FALSE WHERE id = %s', (fine,))
        other = create(cur, SV_A, OP_A, NOW - timedelta(hours=1), measure_id=fine, measure_amount='5000')
        assert enriched(cur, SV_A, other)['measure_state'] == 'proposed'


def test_measure_by_control_decision_reaches_deductions(db):
    with db.cursor() as cur:
        control_id = service.create_control(cur, ctx_of(cur, HEAD_SZOV), {
            'operator_id': OP_B, 'source': 'manual', 'reason': 'Решение руководителя',
            'end_date': TODAY.isoformat()}, now=NOW, today=TODAY)
        fine = measure_id(cur, SZOV, 'Лишение переменной части (фиксированная сумма)')
        with pytest.raises(service.DisciplineError) as error:
            service.decide_control(cur, ctx_of(cur, HEAD_SZOV), control_id,
                                   {'decision': 'measure', 'measure_id': fine}, now=NOW, today=TODAY)
        assert error.value.code == 'MEASURE_AMOUNT'
        service.decide_control(cur, ctx_of(cur, HEAD_SZOV), control_id,
                               {'decision': 'measure', 'measure_id': fine, 'measure_amount': '20000'},
                               now=NOW, today=TODAY)
        decisions = queries.decision_measures(cur, [SZOV], date_from=NOW - timedelta(days=1),
                                              date_to=NOW + timedelta(days=1))
        rows, totals = report.report_measures([], decisions)
        assert [(r['basis'], r['operator_name'], r['measure_amount']) for r in rows] == [
            ('Решение по особому контролю', 'Оператор Борис', 20000.0)]
        assert totals[0]['amount'] == 20000.0


def test_deductions_period_is_the_approval_date(db):
    with db.cursor() as cur:
        fine = measure_id(cur, SZOV, 'Лишение переменной части (фиксированная сумма)')
        record_id = _under_control_record(cur, measure_id=fine, measure_amount='15000')
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), record_id, TODAY, for_update=True)
        approved_at = NOW + timedelta(days=12)  # утвердили уже в ноябре
        service.approve_measure(cur, ctx_of(cur, HEAD_SZOV), record, now=approved_at)
        october = queries.list_records(cur, ctx_of(cur, HEAD_SZOV), TODAY, {
            'approved_from': datetime(2026, 10, 1), 'approved_to': datetime(2026, 11, 1)})
        november = queries.list_records(cur, ctx_of(cur, HEAD_SZOV), TODAY, {
            'approved_from': datetime(2026, 11, 1), 'approved_to': datetime(2026, 12, 1)})
        assert [r['id'] for r in october] == [] and [r['id'] for r in november] == [record_id]


# ── Файлы ────────────────────────────────────────────────────────────────────

def test_invalidated_explanation_asks_for_a_new_one(db):
    with db.cursor() as cur:
        record_id = create(cur, SV_A, OP_A, NOW - timedelta(days=3))
        record = service.load_record(cur, ctx_of(cur, SV_A), record_id, TODAY, for_update=True)
        file_id = service.add_explanation(cur, ctx_of(cur, SV_A), record, upload('wrong.pdf'), now=NOW)
        assert enriched(cur, SV_A, record_id)['state'] == 'explanation_received'
        row = queries.get_file(cur, file_id)
        record = service.load_record(cur, ctx_of(cur, HEAD_SZOV), record_id, TODAY, for_update=True)
        service.invalidate_file(cur, ctx_of(cur, HEAD_SZOV), row, record, 'Чужая объяснительная', now=NOW)
        item = enriched(cur, HEAD_SZOV, record_id)
        assert item['state'] == 'explanation_requested' and item['explanation_received_at'] is None
        assert [b['code'] for b in item['close_blockers']] == ['explanation']
        # Файл не удалён — он в истории с пометкой.
        files = queries.files_for(cur, record_ids=[record_id])
        assert len(files) == 1 and files[0]['invalid_at'] is not None
        # Прежнее «нужна объяснительная» погасло при загрузке; новое — с причиной.
        cur.execute("SELECT body FROM discipline_notifications WHERE user_id = %s AND read_at IS NULL"
                    " AND kind = 'explanation_required'", (OP_A,))
        bodies = [row[0] for row in cur.fetchall()]
        assert len(bodies) == 1 and 'признан недействительным: Чужая объяснительная' in bodies[0]


# ── Перевод оператора и список контролей ─────────────────────────────────────

def test_transferred_operator_history_moves_to_the_new_supervisor(db):
    with db.cursor() as cur:
        record_id = create(cur, SV_A, OP_A, NOW - timedelta(days=3))
        cur.execute("UPDATE group_operator_memberships SET end_date = %s WHERE operator_id = %s",
                    (TODAY - timedelta(days=1), OP_A))
        cur.execute("INSERT INTO group_operator_memberships (group_id, operator_id, start_date) VALUES (11, %s, %s)",
                    (OP_A, TODAY))
        # Владелец: «СВ видит только свою группу» — и записи, которые завёл сам.
        assert queries.list_records(cur, ctx_of(cur, SV_A), TODAY) == []
        with pytest.raises(service.DisciplineError) as error:
            service.load_record(cur, ctx_of(cur, SV_A), record_id, TODAY)
        assert error.value.status == 403
        assert [r['id'] for r in queries.list_records(cur, ctx_of(cur, SV_B), TODAY)] == [record_id]


def test_controls_list_keeps_fresh_closed_and_filters_supervisor_in_sql(db):
    with db.cursor() as cur:
        for index in range(3):
            control_id = service.create_control(cur, ctx_of(cur, HEAD_SZOV), {
                'operator_id': OP_B, 'source': 'manual', 'reason': 'Решение %d' % index,
                'start_date': (TODAY - timedelta(days=30 - index)).isoformat(),
                'end_date': (TODAY - timedelta(days=20 - index)).isoformat()}, now=NOW, today=TODAY)
            service.decide_control(cur, ctx_of(cur, HEAD_SZOV), control_id, {'decision': 'removed'},
                                   now=NOW + timedelta(minutes=index), today=TODAY)
        closed = queries.list_controls(cur, [SZOV], state='closed', limit=2)
        assert [c['reason'] for c in closed] == ['Решение 2', 'Решение 1']
        service.create_control(cur, ctx_of(cur, HEAD_SZOV), {
            'operator_id': OP_A, 'source': 'manual', 'reason': 'Группа Альфа',
            'end_date': (TODAY + timedelta(days=5)).isoformat()}, now=NOW, today=TODAY)
        assert [c['operator_id'] for c in service.visible_controls(cur, ctx_of(cur, SV_A), TODAY)] == [OP_A]
        assert [c['operator_id'] for c in service.visible_controls(cur, ctx_of(cur, SV_B), TODAY, state='all')] == [
            OP_B, OP_B, OP_B]


def test_after_transfer_bell_and_reminders_follow_the_operator(db):
    with db.cursor() as cur:
        create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=10))
        late = create(cur, HEAD_SZOV, OP_A, NOW - timedelta(days=9))
        total, rows = queries.bell(cur, ctx_of(cur, SV_A), 10, TODAY)
        assert total == 1 and 'следующее нарушение — акт' in rows[0]['title']
        cur.execute("UPDATE group_operator_memberships SET end_date = %s WHERE operator_id = %s",
                    (TODAY - timedelta(days=1), OP_A))
        cur.execute("INSERT INTO group_operator_memberships (group_id, operator_id, start_date) VALUES (11, %s, %s)",
                    (OP_A, TODAY))
        # Колокол прежнего СВ о человеке молчит: запись ему больше не откроется.
        assert queries.bell(cur, ctx_of(cur, SV_A), 10, TODAY) == (0, [])
        cur.execute('UPDATE discipline_records SET created_at = %s WHERE id = %s', (NOW - timedelta(days=9), late))
        service.run_reminders(cur, now=NOW, today=TODAY)
        overdue = 'Оператор Анна: объяснительная не предоставлена в срок'
        assert overdue in bell_titles(cur, SV_B) and overdue not in bell_titles(cur, SV_A)


def test_shift_is_required_when_schedule_has_none(db):
    with db.cursor() as cur:
        with pytest.raises(service.DisciplineError) as error:
            create(cur, SV_A, OP_A, NOW - timedelta(days=5), shift_label='')
        assert error.value.code == 'SHIFT_REQUIRED'
        # В графике смена есть — подставляется сама.
        record_id = create(cur, SV_A, OP_A, NOW - timedelta(hours=1), shift_label='')
        assert queries.get_record(cur, record_id, TODAY)['shift_label'] == '09:00–18:00'
