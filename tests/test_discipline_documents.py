# -*- coding: utf-8 -*-
"""«Дисциплина»: документы по шаблонам организации (владелец, 23.10.2026).

Стандартный акт — по присланному примеру акта («из него сделай шаблон»),
бланк объяснительной — по присланному шаблону. Здесь — что в документах
стоит слово в слово и что Word откроет их без «файл повреждён»: порядок
элементов по схеме, жёсткие ширины строк подписей, A4 и Times New Roman без
следов шаблона python-docx. Базы не нужно.
"""
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest

docx = pytest.importorskip('docx')
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

from discipline import act_docx, catalog, docx_layout, explanation_docx  # noqa: E402

DESCRIPTION = 'Ушла на перерыв в 14:32 вне графика и вернулась в 14:41'
RECORDS = [
    {'id': 12, 'number': 'Д-2026-0012', 'occurred_at': datetime(2026, 10, 3, 9, 15),
     'type_title': 'Опоздание к началу смены', 'explanation_received_at': datetime(2026, 10, 4, 12, 0),
     'action': 'explanation', 'needs_explanation_now': True, 'description': 'Скрытое описание первой записи'},
    {'id': 20, 'number': 'Д-2026-0020', 'occurred_at': datetime(2026, 10, 6, 18, 40),
     'type_title': 'Уход со смены раньше времени без согласования', 'refused_at': datetime(2026, 10, 7, 10, 0),
     'action': 'explanation', 'needs_explanation_now': True},
    {'id': 33, 'number': 'Д-2026-0031', 'occurred_at': datetime(2026, 10, 9, 14, 32),
     'type_title': 'Отсутствие на рабочем месте сверх регламента перерывов', 'description': DESCRIPTION,
     'action': 'act', 'needs_explanation_now': True},
]
ACT = {'id': 7, 'number': 'А-2026-0003', 'trigger_record_id': 33, 'composed_on': date(2026, 10, 11),
       'decision_title': 'Лишение переменной части (фиксированная сумма)', 'decision_amount': 20000}
OPERATOR = {'name': 'Иванова Анна Сергеевна', 'role': 'operator'}
CONTEXT = {'department_name': 'СЗоВ — Служба заботы о водителях', 'composer_role': 'supervisor',
           'head_name': 'Петрова Мария Ивановна', 'gender': 'female', 'control_days': 30,
           'control_end': date(2026, 11, 8), 'period_kind': 'rolling', 'period_days': 30}
CREATED = datetime(2026, 10, 23, 12, 0)


def _act(act=None, records=None, context=None, operator=None, composed_by='Сидорова Жанна'):
    data = act_docx.build(dict(ACT, **(act or {})), operator or OPERATOR, RECORDS if records is None else records,
                          today=date(2026, 10, 11), composed_by=composed_by,
                          context=CONTEXT if context is None else context, created=CREATED)
    return Document(BytesIO(data))


def _texts(document):
    return [p.text for p in document.paragraphs]


def _rows(document):
    return [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]


# ── Названия отделов и шапка объяснительной ─────────────────────────────────

@pytest.mark.parametrize('name, forms', [
    ('СЗоВ — Служба заботы о водителях', ('Службы заботы о водителях', 'СЗоВ')),
    ('Отдел продаж', ('Отдела продаж', 'Отдела продаж')),
    ('Тез КЦ', ('Тез КЦ', 'Тез КЦ')),
    ('ОП - отдел продаж', ('отдела продаж', 'ОП')),
    ('  Служба   заботы о водителях ', ('Службы заботы о водителях', 'Службы заботы о водителях')),
    (None, ('', '')),
])
def test_department_forms(name, forms):
    # Боевые названия на 23.10.2026: «СЗоВ — Служба заботы о водителях», «Отдел продаж», «Тез КЦ».
    assert act_docx.department_forms(name) == forms


def test_addressee_is_normalized_and_bounded():
    assert catalog.normalize_addressee('  Заказчику:  ИП «Пример» \n\n  Примеровой А. Б. \n') == \
        'Заказчику: ИП «Пример»\nПримеровой А. Б.'
    assert catalog.normalize_addressee(None) == ''
    with pytest.raises(catalog.SettingsError):
        catalog.normalize_addressee('\n'.join('строка %d' % i for i in range(catalog.ADDRESSEE_MAX_LINES + 1)))
    with pytest.raises(catalog.SettingsError):
        catalog.normalize_addressee('я' * (catalog.ADDRESSEE_MAX_CHARS + 1))
    with pytest.raises(catalog.SettingsError):
        catalog.normalize_addressee(['Заказчику'])
    # Сохранённое раньше не роняет чтение настроек — обрезается.
    stored = catalog.normalize_addressee('\n'.join('строка %d' % i for i in range(9)), strict=False)
    assert stored.count('\n') == catalog.ADDRESSEE_MAX_LINES - 1


def test_addressee_lives_in_department_settings():
    settings = catalog.normalize_settings({})
    assert settings['explanation_addressee'] == ''
    saved = catalog.normalize_settings({'explanation_addressee': 'Заказчику: ИП «Пример»'}, base=settings)
    assert saved['explanation_addressee'] == 'Заказчику: ИП «Пример»'
    # Правка других настроек шапку не стирает.
    again = catalog.normalize_settings({'period_days': 45}, base=saved)
    assert again['explanation_addressee'] == 'Заказчику: ИП «Пример»'
    with pytest.raises(catalog.SettingsError):
        catalog.normalize_settings({'explanation_addressee': 'я' * 400}, base=saved)
    # Шапка — не правило счёта: её правка не заводит новую версию окна и лестницы.
    assert 'explanation_addressee' not in catalog.COUNTING_KEYS


# ── Стандартный акт ─────────────────────────────────────────────────────────

def test_act_follows_the_organization_sample_word_for_word():
    document = _act()
    texts = _texts(document)
    assert texts[0] == 'Акт № А-2026-0003\nо совершении работником дисциплинарного нарушения'
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.bold for run in document.paragraphs[0].runs)
    assert texts[1] == 'г. Алматы\t11.10.2026 г.'
    assert texts[2] == ('Мною, супервайзером Службы заботы о водителях Сидорова Жанна, составлен настоящий акт '
                        'о нижеследующем:')
    assert texts[3] == ('Сотрудник Иванова Анна Сергеевна 9 октября в 14:32 допустила нарушение трудовой '
                        'дисциплины: отсутствие на рабочем месте сверх регламента перерывов. %s.' % DESCRIPTION)
    assert texts[4] == 'Это третье нарушение сотрудника за последние 30 дней:'
    assert texts[5:8] == [
        '1. 03.10.2026, 09:15 — опоздание к началу смены (запись Д-2026-0012), объяснительная получена 04.10.2026;',
        '2. 06.10.2026, 18:40 — уход со смены раньше времени без согласования (запись Д-2026-0020), '
        'отказ от объяснения, акт об отказе от 07.10.2026;',
        '3. 09.10.2026, 14:32 — отсутствие на рабочем месте сверх регламента перерывов (запись Д-2026-0031), '
        'объяснительная не предоставлена.',
    ]
    assert texts[8] == ('В связи с данным нарушением принято решение применить меру «Лишение переменной части '
                        '(фиксированная сумма)» в размере 20\u00a0000\u00a0₸ и установить испытательный срок '
                        'продолжительностью 30 дней (до 08.11.2026) для контроля соблюдения трудовой дисциплины.')
    assert texts[9] == ('В случае повторного нарушения в течение 30 дней сотруднику будет вынесено окончательное '
                        'предупреждение с последующим увольнением.')
    assert texts[10:12] == ['Содержание настоящего Акта подтверждаем личными подписями:', 'Настоящий акт составил:']
    assert _rows(document) == [
        ['Супервайзер СЗоВ', act_docx.SIGNATURE, 'ФИО: Сидорова Жанна'],
        ['С актом ознакомлен сотрудник СЗоВ', act_docx.SIGNATURE, 'ФИО: Иванова Анна Сергеевна'],
        ['С актом ознакомлен руководитель СЗоВ', act_docx.SIGNATURE, 'ФИО: Петрова Мария Ивановна'],
    ]
    # Описания прочих записей периода в акт не попадают: запись может быть вне
    # доступа того, кто скачивает черновик.
    assert 'Скрытое описание' not in '\n'.join(texts)


def test_act_date_sits_on_the_right_edge():
    place = _act().paragraphs[1]
    stops = list(place.paragraph_format.tab_stops)
    assert len(stops) == 1 and stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
    assert abs(stops[0].position - Cm(docx_layout.TEXT_WIDTH_CM)) < Cm(0.01)
    assert all(run.bold for run in place.runs)


@pytest.mark.parametrize('gender, verb', [('female', 'допустила'), ('male', 'допустил'), (None, 'допустил(а)')])
def test_act_verb_follows_gender_from_the_card(gender, verb):
    texts = _texts(_act(context=dict(CONTEXT, gender=gender)))
    assert texts[3].startswith('Сотрудник Иванова Анна Сергеевна 9 октября в 14:32 %s нарушение' % verb)


def test_act_head_composer_and_other_department_names():
    document = _act(context=dict(CONTEXT, composer_role='head', department_name='Отдел продаж'),
                    composed_by='Глава Продаж')
    texts = _texts(document)
    assert texts[2] == 'Мною, руководителем Отдела продаж Глава Продаж, составлен настоящий акт о нижеследующем:'
    assert [row[0] for row in _rows(document)] == [
        'Руководитель Отдела продаж', 'С актом ознакомлен сотрудник Отдела продаж',
        'С актом ознакомлен руководитель Отдела продаж']


def test_act_without_decision_control_or_names_leaves_lines_for_the_pen():
    document = _act(act={'decision_title': None, 'decision_amount': None},
                    context={'department_name': None}, operator={'name': ''}, composed_by=None)
    texts = _texts(document)
    blank = act_docx.BLANK
    assert texts[2] == 'Мною, супервайзером %s, составлен настоящий акт о нижеследующем:' % blank
    assert texts[3].startswith('Сотрудник %s 9 октября в 14:32 допустил(а) нарушение' % blank)
    assert 'В связи с данным нарушением принято решение %s.' % blank in texts
    assert ('В случае повторного нарушения сотруднику будет вынесено окончательное предупреждение с последующим '
            'увольнением.') in texts
    assert _rows(document) == [
        ['Супервайзер', act_docx.SIGNATURE, 'ФИО: ____________________'],
        ['С актом ознакомлен сотрудник', act_docx.SIGNATURE, 'ФИО: ____________________'],
        ['С актом ознакомлен руководитель', act_docx.SIGNATURE, 'ФИО: ____________________'],
    ]


def test_act_decision_in_percent_and_month_window():
    texts = _texts(_act(act={'decision_title': 'Лишение переменной части (процент)', 'decision_amount': None,
                             'decision_percent': 15}, context=dict(CONTEXT, period_kind='month')))
    assert texts[4] == 'Это третье нарушение сотрудника за октябрь 2026 г.:'
    assert texts[8].startswith('В связи с данным нарушением принято решение применить меру «Лишение переменной '
                               'части (процент)» в размере 15\u00a0% и установить')


def test_act_on_a_single_record_has_no_period_list():
    single = [dict(RECORDS[2], description='Спал на рабочем месте!')]
    texts = _texts(_act(records=single, context=dict(CONTEXT, gender='male')))
    assert texts[3] == ('Сотрудник Иванова Анна Сергеевна 9 октября в 14:32 допустил нарушение трудовой '
                        'дисциплины: отсутствие на рабочем месте сверх регламента перерывов. Спал на рабочем месте!')
    assert not any(text.startswith('Это ') for text in texts)
    assert texts[4].startswith('В связи с данным нарушением')


def test_act_names_the_year_when_the_violation_is_from_another_year():
    december = [dict(RECORDS[2], occurred_at=datetime(2025, 12, 29, 21, 5), description=None)]
    texts = _texts(_act(act={'composed_on': date(2026, 1, 12)}, records=december))
    assert texts[1] == 'г. Алматы\t12.01.2026 г.'
    assert texts[3] == ('Сотрудник Иванова Анна Сергеевна 29 декабря 2025 г. в 21:05 допустила нарушение '
                        'трудовой дисциплины: отсутствие на рабочем месте сверх регламента перерывов.')


def test_act_period_list_tails():
    records = [
        dict(RECORDS[0], explanation_received_at=None, action='training', training_id=5),
        dict(RECORDS[1], refused_at=None, action='training', training_id=None),
        dict(RECORDS[2], type_title='ПК не выключен после смены', needs_explanation_now=False, action='act'),
    ]
    texts = _texts(_act(records=records))
    assert texts[5].endswith('(запись Д-2026-0012), проведён тренинг;')
    assert texts[6].endswith('(запись Д-2026-0020);')
    # Сокращение в начале типа остаётся заглавным, объяснительная не нужна — хвоста нет.
    assert texts[7] == '3. 09.10.2026, 14:32 — ПК не выключен после смены (запись Д-2026-0031).'
    assert texts[3].startswith('Сотрудник Иванова Анна Сергеевна 9 октября в 14:32 допустила нарушение трудовой '
                               'дисциплины: ПК не выключен после смены.')


def test_act_signature_rows_have_fixed_widths_in_schema_order():
    table = _act().tables[0]
    tblPr = table._tbl.tblPr
    children = [child.tag for child in tblPr]
    assert tblPr.find(qn('w:tblLayout')).get(qn('w:type')) == 'fixed'
    # tblCellMar по схеме — после tblLayout и до tblLook.
    assert children.index(qn('w:tblLayout')) < children.index(qn('w:tblCellMar')) < children.index(qn('w:tblLook'))
    margins = tblPr.find(qn('w:tblCellMar'))
    assert margins.find(qn('w:left')).get(qn('w:w')) == '0'
    assert margins.find(qn('w:right')).get(qn('w:w')) == '57'
    widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid.findall(qn('w:gridCol'))]
    assert [round(w / 1440 * 2.54, 1) for w in widths] == list(act_docx.SIGNATURE_WIDTHS_CM)
    assert abs(sum(act_docx.SIGNATURE_WIDTHS_CM) - docx_layout.TEXT_WIDTH_CM) < 1e-9
    for row in table.rows:
        assert [round(cell.width / 360000, 1) for cell in row.cells] == list(act_docx.SIGNATURE_WIDTHS_CM)
    # Документ не кончается таблицей.
    body = table._tbl.getparent()
    assert body[-2].tag == qn('w:p') and body[-1].tag == qn('w:sectPr')


def test_act_signature_texts_fit_their_columns():
    """Ширины посчитаны по Times New Roman 12: подпись строки и «Подпись: ___»
    не переносятся у отделов раздела (у «Отдела продаж» строка длиннее — ей можно)."""
    font_path = '/System/Library/Fonts/Supplemental/Times New Roman.ttf'
    ImageFont = pytest.importorskip('PIL.ImageFont')
    try:
        font = ImageFont.truetype(font_path, 1200)
    except OSError:
        pytest.skip('нет Times New Roman')

    def width_cm(text):
        return font.getlength(text) / 100.0 / 72 * 2.54

    usable = [width - 0.1 for width in act_docx.SIGNATURE_WIDTHS_CM]  # правое поле ячейки 57 twip ≈ 0,1 см
    for short in ('СЗоВ', 'Тез КЦ'):
        assert width_cm('С актом ознакомлен руководитель %s' % short) < usable[0]
    assert width_cm(act_docx.SIGNATURE) < usable[1]


# ── Бланк объяснительной ────────────────────────────────────────────────────

def _blank(record=None, operator=None, addressee=None):
    data = explanation_docx.build(record or RECORDS[2], OPERATOR if operator is None else operator,
                                  addressee=addressee, created=CREATED)
    return Document(BytesIO(data))


def test_blank_follows_the_organization_template():
    document = _blank(addressee='Заказчику: Индивидуальному предпринимателю «Пример»\n\nПримеровой А. Б.')
    paragraphs = document.paragraphs
    assert [p.text for p in paragraphs[:3]] == [
        'Заказчику: Индивидуальному предпринимателю «Пример»', 'Примеровой А. Б.',
        'от Оператора: Иванова Анна Сергеевна']
    for item in paragraphs[:3]:
        assert item.alignment == WD_ALIGN_PARAGRAPH.RIGHT and all(run.bold for run in item.runs)
    title = paragraphs[3]
    assert title.text == 'Объяснительная' and title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].font.size == Pt(16)
    writing = paragraphs[4:4 + explanation_docx.WRITING_LINES]
    # Отступ хранится в twip — сравнение с допуском.
    assert all(p.text == '' and abs(p.paragraph_format.first_line_indent - Cm(1.25)) < Cm(0.01) for p in writing)
    assert paragraphs[4 + explanation_docx.WRITING_LINES].text == '«____» ___________________ 20___ г.'
    nbsp = explanation_docx.NBSP
    assert _rows(document) == [['Иванова Анна Сергеевна', nbsp, nbsp], ['Ф.И.О', nbsp, 'подпись']]
    footer = document.sections[0].footer.paragraphs[0].text
    assert document.sections[0].footer.paragraphs[0].runs[0].font.size == Pt(9)
    assert footer == ('Запись о нарушении Д-2026-0031 от 09.10.2026: '
                      'Отсутствие на рабочем месте сверх регламента перерывов')
    # Описание обстоятельств оператору не показывается (ТЗ п. 3) — и в бланке его нет.
    everything = '\n'.join([p.text for p in paragraphs] + [footer])
    assert DESCRIPTION not in everything


def test_blank_without_addressee_or_name_leaves_lines_and_names_the_trainee():
    document = _blank(operator={'name': '', 'role': 'trainee'}, addressee='   \n ')
    assert [p.text for p in document.paragraphs[:2]] == [explanation_docx.DEFAULT_ADDRESSEE,
                                                         'от Стажёра: ' + '_' * 30]
    assert _blank(operator={'name': 'Тестқызы Әлия Нұрланқызы', 'role': 'trainee'}).paragraphs[1].text == \
        'от Стажёра: Тестқызы Әлия Нұрланқызы'


def test_blank_signature_lines_are_cell_borders_in_schema_order():
    table = _blank().tables[0]
    name, gap, sign = table.rows[0].cells
    # Справа, как в шаблоне, — прижатием таблицы, без пустой ячейки-отступа.
    assert table.alignment == WD_TABLE_ALIGNMENT.RIGHT
    assert abs(sum(explanation_docx.SIGNATURE_WIDTHS_CM) - 12.5) < 1e-9
    for cell in table.rows[0].cells + table.rows[1].cells:
        assert cell.text and cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    for cell in (name, sign):
        tcPr = cell._tc.tcPr
        tags = [child.tag for child in tcPr]
        # tcBorders по схеме — сразу за tcW, иначе Word объявит файл повреждённым.
        assert tags.index(qn('w:tcBorders')) == tags.index(qn('w:tcW')) + 1
        bottom = tcPr.find(qn('w:tcBorders')).find(qn('w:bottom'))
        assert bottom.get(qn('w:val')) == 'single'
    # Между линией ФИО и линией подписи — зазор без линии.
    assert gap._tc.tcPr.find(qn('w:tcBorders')) is None
    tblPr = table._tbl.tblPr
    children = [child.tag for child in tblPr]
    assert children.index(qn('w:jc')) < children.index(qn('w:tblLayout')) < children.index(qn('w:tblCellMar'))


# ── Пакет: A4, шрифт, язык, без следов шаблона python-docx ─────────────────

@pytest.mark.parametrize('make', [_act, _blank])
def test_package_is_clean(make):
    document = make()
    section = document.sections[0]
    assert abs(section.page_width - Cm(21)) < Cm(0.01) and abs(section.page_height - Cm(29.7)) < Cm(0.01)
    assert [round(getattr(section, '%s_margin' % side) / 360000, 1) for side in ('left', 'right', 'top', 'bottom')] \
        == [3.0, 1.5, 2.0, 2.0]
    doc_defaults = document.styles.element.find(qn('w:docDefaults'))
    defaults = doc_defaults.find(qn('w:rPrDefault')).find(qn('w:rPr'))
    fonts = defaults.find(qn('w:rFonts'))
    assert fonts.get(qn('w:ascii')) == fonts.get(qn('w:hAnsi')) == docx_layout.FONT
    assert all(fonts.get(qn(attr)) is None for attr in docx_layout._THEME_FONT_ATTRS)
    assert defaults.find(qn('w:lang')).get(qn('w:val')) == 'ru-RU'
    # Одинарный интервал без отступов — и у «Normal», и в настройках по умолчанию
    # (на них стоят абзацы не от «Normal», например колонтитул).
    spacing = doc_defaults.find(qn('w:pPrDefault')).find(qn('w:pPr')).find(qn('w:spacing'))
    assert [spacing.get(qn(key)) for key in ('w:after', 'w:line', 'w:lineRule')] == ['0', '240', 'auto']
    normal = document.styles['Normal']
    assert normal.font.name == docx_layout.FONT
    fmt = normal.paragraph_format
    assert (fmt.space_before, fmt.space_after, fmt.line_spacing_rule) == (Pt(0), Pt(0), WD_LINE_SPACING.SINGLE)
    rels = document.part.package.rels
    assert not [rel for rel in rels.values() if rel.reltype == RT.THUMBNAIL]
    # Режим Word 2013+: с режимом шаблона Word пишет в заголовке «Режим совместимости».
    compat = document.settings.element.find(qn('w:compat'))
    assert [item.get(qn('w:val')) for item in compat.findall(qn('w:compatSetting'))
            if item.get(qn('w:name')) == 'compatibilityMode'] == ['15']
    props = document.core_properties
    assert props.author == 'iCore' and props.last_modified_by == ''
    # Время раздела — по Алматы (UTC+5), а в свойствах файла — UTC.
    assert props.created.replace(tzinfo=None) == props.modified.replace(tzinfo=None) == CREATED - timedelta(hours=5)


@pytest.mark.parametrize('make', [_act, _blank])
def test_own_documents_pass_the_upload_check(make):
    """Оператор заполняет наш бланк и прикладывает его объяснительной, глава
    делает шаблон акта из нашего черновика: проверка DOCX (files.check_docx)
    свои же документы не отклоняет."""
    from discipline import files
    buffer = BytesIO()
    make().save(buffer)
    files.check_docx(buffer.getvalue(), name='документ.docx')
    files.check_docx(buffer.getvalue(), name='шаблон.docx', template=True)


def test_act_composer_window_and_term_come_from_the_context():
    """Составитель — из контекста (СВ записи-основания), окно и срок — свои числа,
    а не значения по умолчанию (30 — и запасное значение в коде)."""
    texts = _texts(_act(context=dict(CONTEXT, composer_name='Сидорова Жанна', period_days=45, control_days=60,
                                     control_end=date(2026, 12, 8)), composed_by='Глава Отдела'))
    assert texts[2].startswith('Мною, супервайзером Службы заботы о водителях Сидорова Жанна, составлен')
    assert texts[4] == 'Это третье нарушение сотрудника за последние 45 дней:'
    assert 'и установить испытательный срок продолжительностью 60 дней (до 08.12.2026) для контроля' in texts[8]
    assert texts[9].startswith('В случае повторного нарушения в течение 60 дней сотруднику')


def test_act_money_keeps_tiyn_and_stays_on_one_line():
    texts = _texts(_act(act={'decision_amount': Decimal('1500.75')}))
    assert 'в размере 1\u00a0500,75\u00a0₸ и установить' in texts[8]
    texts = _texts(_act(act={'decision_title': 'Лишение переменной части (процент)', 'decision_amount': None,
                             'decision_percent': Decimal('12.5')}))
    assert 'в размере 12,5\u00a0% и установить' in texts[8]


def test_documents_drop_characters_xml_cannot_hold():
    """Вставка из 1С и Excel приносит управляющие символы — документ собирается
    без них, а не падает 500 («All strings must be XML compatible»)."""
    dirty = [dict(RECORDS[2], type_title='Отсутствие\x08 на месте', description='Ушла на перерыв\x00')]
    texts = _texts(_act(records=dirty, operator={'name': 'Иванова\x1f Анна'}))
    assert texts[3] == ('Сотрудник Иванова Анна 9 октября в 14:32 допустила нарушение трудовой дисциплины: '
                        'отсутствие на месте. Ушла на перерыв.')
    blank = _blank(record=dict(RECORDS[2], type_title='Отсутствие\x08 на месте'),
                   operator={'name': 'Иванова\x1f Анна', 'role': 'operator'}, addressee='Заказчику:\x07 ИП')
    assert [p.text for p in blank.paragraphs[:2]] == ['Заказчику: ИП', 'от Оператора: Иванова Анна']
    assert blank.sections[0].footer.paragraphs[0].text.endswith(': Отсутствие на месте')


def test_day_words_and_plurals():
    assert docx_layout.day_words(date(2026, 10, 9)) == '9 октября 2026 г.'
    assert docx_layout.day_words(date(2026, 5, 1), year=False) == '1 мая'
    assert docx_layout.day_words(None) == ''
    assert [docx_layout.plural(n, 'день', 'дня', 'дней') for n in (1, 2, 5, 11, 21, 22, 25, 111)] == [
        'день', 'дня', 'дней', 'дней', 'день', 'дня', 'дней', 'дней']
