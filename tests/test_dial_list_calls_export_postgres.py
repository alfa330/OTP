# -*- coding: utf-8 -*-
"""Необязательная проверка настоящим SQL: DIAL_LIST_TEST_PORT — порт локального Postgres.

Детализация звонков раздела «Удаленный КЦ» (владелец, 15.10.2026): запрос dial_list/calls_export.py
исполняется на настоящей схеме раздела (dial_list.schema), книга собирается настоящим
DialListService.calls_export и читается обратно. Строковые проверки в
tests/test_dial_list_calls_export.py ловят, что запрос НАПИСАН, а не что он ЗНАЧИТ: границы суток
по Алматы, отдел базы, а не человека, тестовые номера, итог с типом, успешка.

Ни боевых учёток, ни импорта приложения: только 127.0.0.1, своя схема со случайным именем, в конце
она удаляется. Запуск на локальном кластере:

    DIAL_LIST_TEST_PORT=55441 python -m pytest tests/test_dial_list_calls_export_postgres.py

Люди и номера выдуманы.
"""
import contextlib
import os
import sys
import uuid
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dial_list import schema as dial_schema  # noqa: E402
from dial_list import service as dial_service  # noqa: E402

PORT = os.environ.get('DIAL_LIST_TEST_PORT')
pytestmark = pytest.mark.skipif(not PORT, reason='needs local PostgreSQL (DIAL_LIST_TEST_PORT)')

RCC, TEZ = 1954, 560
OWN, GUEST, TEZ_OP = 520, 701, 703
ALMATY = timezone(timedelta(hours=5))

BASE_DDL = '''
    CREATE TABLE departments (id integer PRIMARY KEY, name text NOT NULL, code text,
                              is_active boolean DEFAULT TRUE);
    CREATE TABLE test_phone_numbers (phone_key varchar(10) UNIQUE);
    CREATE TABLE users (id integer PRIMARY KEY, name text NOT NULL, login text, role text,
                        status text NOT NULL DEFAULT 'working', department_id integer,
                        sip_number text);
    CREATE TABLE sip_department_config (department_id integer PRIMARY KEY, provider text,
                                        sip_server text, ring_timeout_rings smallint);
    CREATE TABLE groups (id integer PRIMARY KEY, name text NOT NULL, department_id integer,
                         status text NOT NULL DEFAULT 'active');
    CREATE TABLE group_operator_memberships (id serial PRIMARY KEY, group_id integer NOT NULL,
                                             operator_id integer NOT NULL, start_date date NOT NULL,
                                             end_date date);
    INSERT INTO departments (id, name, code) VALUES (1, 'СЗоВ', 'szov'), (560, 'Тез КЦ', 'tez'),
                                                    (909, 'Фронт офисы', 'front_office'),
                                                    (1954, 'Удаленный КЦ', 'remote_cc');
    INSERT INTO users (id, name, login, role, department_id, sip_number) VALUES
        (520, 'Своев Сава', 'sava', 'operator', 1954, '902'),
        (701, 'Гостев Глеб', 'gleb', 'operator', 1, '1024'),
        (703, 'Тезов Тимур', 'timur', 'operator', 560, '913'),
        (421, 'Регионова Рая', 'raya', 'operator', 909, NULL),
        (422, 'Областной Олжас', 'olzhas', 'operator', 909, NULL);
    INSERT INTO groups (id, name, department_id, status) VALUES
        (37, 'Группа Регионы', 909, 'active'), (38, 'Старая группа офисов', 909, 'archived'),
        (39, 'Тестовая группа', 1954, 'active');
    INSERT INTO test_phone_numbers (phone_key) VALUES ('7019999999');
'''


class _Db:
    """_get_cursor на своей схеме: отдельное соединение, фиксация в конце блока."""

    def __init__(self, schema):
        self.schema = schema

    def connect(self):
        import psycopg2
        conn = psycopg2.connect(host='127.0.0.1', port=int(PORT), user='postgres', dbname='postgres',
                                connect_timeout=3)
        with conn.cursor() as cur:
            cur.execute('SET search_path TO ' + self.schema)
            cur.execute("SET statement_timeout TO '10s'")
        conn.commit()
        return conn

    @contextlib.contextmanager
    def _get_cursor(self):
        conn = self.connect()
        cur = conn.cursor()
        try:
            yield cur
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()


@pytest.fixture
def db():
    schema = 't_dial_export_' + uuid.uuid4().hex
    base = _Db(schema)
    admin = base.connect()
    admin.autocommit = True
    with admin.cursor() as cur:
        cur.execute('CREATE SCHEMA ' + schema)
        cur.execute('SET search_path TO ' + schema)
        cur.execute(BASE_DDL)
        dial_schema.init_dial_list_schema(cur)
    try:
        yield base
    finally:
        with admin.cursor() as cur:
            cur.execute('DROP SCHEMA ' + schema + ' CASCADE')
        admin.close()


class _Seed:
    def __init__(self, db):
        self.db = db

    def one(self, sql, params=None):
        with self.db._get_cursor() as cur:
            cur.execute(sql, params)
            return cur.fetchone()

    def outcome(self, department, name, subtype=None):
        outcome_id = self.one("INSERT INTO dial_list_outcomes (department_id, name) VALUES (%s, %s) RETURNING id",
                              (department, name))[0]
        subtype_id = None
        if subtype:
            subtype_id = self.one("INSERT INTO dial_list_outcome_subtypes (outcome_id, department_id, name) "
                                  "VALUES (%s, %s, %s) RETURNING id", (outcome_id, department, subtype))[0]
        return str(outcome_id), (str(subtype_id) if subtype_id else None)

    def lead(self, department, name, phone):
        """Водитель базы отдела; ИИН — выдуманный, свой у каждого (он уникален в базе месяца)."""
        return str(self.one("INSERT INTO dial_list_leads (department_id, phone_norm, full_name, iin) "
                            "VALUES (%s, %s, %s, %s) RETURNING id", (department, phone, name, iin(phone)))[0])

    def call(self, lead_id, operator, at, line='902', state='finished', disposition='ANSWER', billsec=0,
             outcome=(None, None), comment='', cancelled=False, api_error=''):
        """Попытка оператора по водителю в момент at (Алматы) — своей строкой выдачи."""
        portion = self.one("INSERT INTO dial_list_portions (operator_id, department_id, size) VALUES (%s, %s, 1) "
                           "RETURNING id", (operator, RCC))[0]
        assignment = self.one("INSERT INTO dial_list_assignments (portion_id, lead_id, operator_id, position, state) "
                              "VALUES (%s, %s, %s, 1, 'done') RETURNING id", (portion, lead_id, operator))[0]
        return str(self.one("""
            INSERT INTO dial_list_attempts (assignment_id, operator_id, internal_number, requested_at, state,
                                            disposition, billsec, outcome_id, outcome_subtype_id, operator_comment,
                                            cancelled, api_error)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING id
        """, (assignment, operator, line, at, state, disposition, billsec, outcome[0], outcome[1], comment,
              cancelled, api_error))[0])


def _at(day, hh, mm=0, ss=0):
    return datetime(2026, 10, day, hh, mm, ss, tzinfo=ALMATY)


def iin(phone):
    return '9901013' + phone[-5:]


def _rows(stream):
    wb = load_workbook(stream)
    return {name: [[c.value for c in line] for line in wb[name].iter_rows()] for name in wb.sheetnames}


def test_period_department_tests_and_columns(db):
    seed = _Seed(db)
    svc = dial_service.DialListService(db)
    refusal = seed.outcome(RCC, 'Отказ', 'Дорого')
    ivanov = seed.lead(RCC, 'Иванов Иван', '77011234567')
    petrov = seed.lead(RCC, 'Петров Пётр', '77017654321')
    tester = seed.lead(RCC, 'Тестовый номер', '77019999999')      # номер из «Реестра тестовых номеров»
    stranger = seed.lead(TEZ, 'Чужой водитель', '77015550000')    # база другого отдела

    # Сутки — по Алматы: 23:59:59 13.10 — ещё не период, 00:00 14.10 — уже он.
    seed.call(ivanov, OWN, _at(13, 23, 59, 59), billsec=20)
    first = seed.call(ivanov, OWN, _at(14, 0, 0, 0), billsec=125, outcome=refusal, comment='дорого, подумает')
    seed.call(petrov, GUEST, _at(14, 9, 30), line='903', disposition='NOANSWER')
    seed.call(petrov, GUEST, _at(15, 23, 59, 59), line='903', state='failed', disposition='',
              api_error="status='error' code=104: number 77017654321")
    seed.call(petrov, GUEST, _at(16, 0, 0, 0), line='903', billsec=40)       # уже после периода
    seed.call(tester, OWN, _at(14, 12, 0), billsec=60)
    seed.call(stranger, TEZ_OP, _at(14, 12, 0), line='913', billsec=60)
    # Успешка засчитана первому разговору с Ивановым. ID водителя — у Иванова два аккаунта; Петрова Sapar
    # ещё не спрашивали (NULL) — в файле пусто.
    ivan_ids = 'a' * 32 + ', ' + 'b' * 32
    with db._get_cursor() as cur:
        cur.execute("UPDATE dial_list_leads SET signed_at = %s, success_attempt_id = %s, success_operator_id = %s, "
                    "yandex_driver_ids = %s WHERE id = %s", (_at(14, 10, 0), first, OWN, ivan_ids, ivanov))

    stream, written = svc.calls_export(RCC, date(2026, 10, 14), date(2026, 10, 15), generated_by='Руководителева Раиса')
    assert written == 3
    book = _rows(stream)
    calls = book['Звонки']
    assert calls[0][:8] == ['Дата и время', 'Оператор', 'Линия', 'Водитель', 'ID водителя', 'ИИН', 'Номер',
                            'Результат']
    assert [r[0] for r in calls[1:]] == [datetime(2026, 10, 14, 0, 0, 0), datetime(2026, 10, 14, 9, 30),
                                         datetime(2026, 10, 15, 23, 59, 59)]
    ivan, petr_no, petr_fail = calls[1:]
    assert ivan[1:9] == ['Своев Сава', '902', 'Иванов Иван', ivan_ids, iin('77011234567'), '+7 ••• ••• 45 67',
                         'Дозвонились', timedelta(seconds=125)]
    assert ivan[9:13] == ['Отказ · Дорого', 'дорого, подумает', 'да', None]
    # Гость с линии удалённого КЦ (сотрудник СЗоВ) — в детализации отдела базы.
    assert petr_no[1:8] == ['Гостев Глеб', '903', 'Петров Пётр', None, iin('77017654321'), '+7 ••• ••• 43 21',
                            'Не поднял трубку']
    assert petr_fail[7] == 'Ошибка АТС' and petr_fail[12] == 'АТС отказала (код 104)'
    flat = repr(book)
    for forbidden in ('77011234567', '7654321', 'Тестовый номер', 'Чужой водитель', 'Тезов'):
        assert forbidden not in flat, forbidden

    operators = book['Операторы']
    assert [r[0] for r in operators[1:]] == ['Гостев Глеб', 'Своев Сава', 'Итого']
    assert operators[1][1:4] == [2, 0, 0]
    assert operators[2][1:4] == [1, 1, 1]
    assert operators[3][1:3] == [3, 1] and operators[3][6:] == [1, 0, 1]
    context = {r[0]: r[1] for r in book['Контекст'] if r and r[0]}
    assert (context['Отдел'], context['Период'], context['Звонков']) == ('Удаленный КЦ', '14.10.2026 — 15.10.2026', 3)


def test_no_calls_gives_no_file(db):
    seed = _Seed(db)
    svc = dial_service.DialListService(db)
    lead = seed.lead(RCC, 'Иванов Иван', '77011234567')
    seed.call(lead, OWN, _at(10, 12, 0), billsec=30)
    assert svc.calls_export(RCC, date(2026, 10, 11), date(2026, 10, 12)) == (None, 0)
    # Звонки только тестового номера — тоже «звонков не было».
    tester = seed.lead(RCC, 'Тестовый номер', '77019999999')
    seed.call(tester, OWN, _at(11, 12, 0), billsec=30)
    assert svc.calls_export(RCC, date(2026, 10, 11), date(2026, 10, 12)) == (None, 0)
    stream, written = svc.calls_export(RCC, date(2026, 10, 10), date(2026, 10, 10))
    assert written == 1 and stream is not None


def test_head_of_regions_groups_are_her_department_groups(db):
    svc = dial_service.DialListService(db)
    assert svc.report_groups([909]) == [{'id': 37, 'name': 'Группа Регионы', 'active': True},
                                        {'id': 38, 'name': 'Старая группа офисов', 'active': False}]
    # Глава удалённого КЦ (код не из списка) и никто — групп отчётов нет.
    assert svc.report_groups([RCC]) == []
    assert svc.report_groups([]) == []
    assert [g['id'] for g in svc.report_groups([909, RCC])] == [37, 38]


def test_group_report_takes_calls_by_membership_on_the_call_day(db):
    seed = _Seed(db)
    svc = dial_service.DialListService(db)
    with db._get_cursor() as cur:
        cur.execute("""INSERT INTO group_operator_memberships (group_id, operator_id, start_date, end_date) VALUES
            (37, 421, '2026-10-14', NULL),            -- Рая в группе с 14-го
            (37, 422, '2026-10-01', '2026-10-14'),    -- Олжас ушёл из группы после 14-го
            (39, 422, '2026-10-15', NULL),
            (39, 520, '2026-09-01', NULL)""")
    lead = seed.lead(RCC, 'Иванов Иван', '77011234567')
    seed.call(lead, 421, _at(13, 23, 59), line='929', billsec=30)            # до вступления — нет
    seed.call(lead, 421, _at(14, 0, 30), line='929', billsec=61)             # 14-е по Алматы — да
    seed.call(lead, 422, _at(14, 18, 0), line='943', disposition='NOANSWER')  # ещё в группе — да
    seed.call(lead, 422, _at(15, 9, 0), line='943', billsec=20)              # уже ушёл — нет
    seed.call(lead, OWN, _at(14, 12, 0), billsec=45)                          # чужая группа — нет

    groups = svc.report_groups([909])
    stream, written = svc.calls_export(RCC, date(2026, 10, 13), date(2026, 10, 15), groups=groups)
    assert written == 2
    book = _rows(stream)
    assert [(r[0], r[1]) for r in book['Звонки'][1:]] == [(datetime(2026, 10, 14, 0, 30), 'Регионова Рая'),
                                                         (datetime(2026, 10, 14, 18, 0), 'Областной Олжас')]
    assert [r[0] for r in book['Операторы'][1:]] == ['Областной Олжас', 'Регионова Рая', 'Итого']
    context = {r[0]: r[1] for r in book['Контекст'] if r and r[0]}
    # В «Контексте» — действующая группа; архивная участвует в отборе, но названием не шумит.
    assert context['Группа'] == 'Группа Регионы — звонки тех, кто в день звонка состоял в группе'
    # Весь отдел без групп — все пять звонков.
    assert svc.calls_export(RCC, date(2026, 10, 13), date(2026, 10, 15))[1] == 5
    # Пустой список групп — ничего, а не «весь отдел».
    assert svc.calls_export(RCC, date(2026, 10, 13), date(2026, 10, 15), groups=[]) == (None, 0)


def test_query_uses_the_time_index(db):
    """Период режется по индексу времени звонка, а не перебором всех попыток."""
    from dial_list import calls_export
    with db._get_cursor() as cur:
        cur.execute('SET enable_seqscan = off')
        cur.execute('EXPLAIN ' + calls_export.CALLS_SQL,
                    {'department_id': RCC, 'date_from': '2026-10-14', 'date_to': '2026-10-15'})
        plan = '\n'.join(r[0] for r in cur.fetchall())
    assert 'idx_dial_list_attempts_requested' in plan or 'idx_dial_list_attempts_live' in plan, plan
