# -*- coding: utf-8 -*-
"""«Табло ОП» · «Чат»: чаты верификаторов в Wazzup (задача #367).

Что закреплено:
  * правила ответа те же, что у строителя обращений СЗоВ (`build_chat_webhook_request_rows`):
    сверяется настоящей функцией из bot_schedule2 на одном и том же диалоге;
  * задержка засчитывается тому, кто ответил; в итог направления идут только верификаторы, а
    диалог, который вёл только сотрудник другой группы, чатом верификаторов не считается;
  * рассылки (бот, автор без привязки, отправка без автора) не отвечают клиенту и не начинают диалог;
  * диалог — до шести часов тишины, относится к дню и часу своего начала; «чатов у человека» —
    пары «чат × сутки», где он писал сам (как «Чаты ОП»);
  * «в работе» — последнее сообщение не старше окна, чат у того, кто писал последним;
    «ждут ответа» — не дольше часа (решения владельца 25.09.2026);
  * тишина потока судится по часу: ночью полтора часа — норма, днём двадцать минут — обрыв;
  * ручки под гардом «Табло ОП», выгрузка — xlsx за период с потолком;
  * отбивка: отклонение = то, что красное на табло; направление зарегистрировано везде;
  * экран: переключатель «Линия / Чат» как у СЗоВ, опрашивается только открытое направление,
    статусов людей на экране чатов нет.
"""

import unittest
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from unittest import mock

from flask import Flask
from openpyxl import load_workbook

from op_wallboard import chat as C, chat_export as X, routes as op_routes
from tests.test_szov_wallboard import _load_names

ROOT = Path(__file__).resolve().parents[1]
BOT_SOURCE = (ROOT / 'bot_schedule2.py').read_text(encoding='utf-8-sig')
DB_SOURCE = (ROOT / 'database.py').read_text(encoding='utf-8-sig')
MONITORING = ROOT / 'src' / 'components' / 'monitoring'

DAY = date(2026, 9, 25)
VER_A, VER_B, OTHER = 328, 302, 287
NAMES = {VER_A: 'Иванова Айгерим', VER_B: 'Петров Данияр'}


def at(hh, mm=0, ss=0, day=DAY):
    return datetime.combine(day, datetime.min.time()) + timedelta(hours=hh, minutes=mm, seconds=ss)


def client(when, chat='c1', channel='ch'):
    return {'channel_id': channel, 'chat_id': chat, 'at': when, 'is_echo': False,
            'user_id': None, 'is_bot': False, 'message_id': f'{chat}-{when.isoformat()}-in'}


def reply(when, user_id, chat='c1', channel='ch', is_bot=False):
    return {'channel_id': channel, 'chat_id': chat, 'at': when, 'is_echo': True,
            'user_id': user_id, 'is_bot': is_bot, 'message_id': f'{chat}-{when.isoformat()}-{user_id}'}


def day_block(rows, day=DAY, now=None):
    messages = C.classify(rows, NAMES)
    return C.day_block(C.split_episodes(messages), messages, day, names=NAMES, now=now)


class ReplyRulesTests(unittest.TestCase):
    def test_first_and_inner_reply_follow_the_szov_rules(self):
        """Ожидание — от первой реплики клиента после ответа; второе сообщение подряд — не ответ."""
        rows = [client(at(10)), reply(at(10, 1), VER_A),           # первый ответ 60 с
                client(at(10, 5)), client(at(10, 6)),               # ждёт с 10:05, а не с 10:06
                reply(at(10, 9), VER_A), reply(at(10, 10), VER_A)]  # 240 с; второе подряд — не ответ
        today = day_block(rows)['today']
        self.assertEqual(today['chats'], 1)
        self.assertEqual(today['first_reply_seconds'], 60.0)
        self.assertEqual(today['inner_reply_seconds'], 150.0)       # (60 + 240) / 2 — первый входит

    def test_same_numbers_as_the_real_szov_request_builder(self):
        """Одна формула на два табло: тот же диалог через строитель обращений СЗоВ даёт те же цифры."""
        ns = {'datetime': datetime, 'timedelta': timedelta}
        _load_names(BOT_SOURCE, {
            'CHAT2DESK_WEBHOOK_CLIENT_TYPE', 'CHAT2DESK_WEBHOOK_OPERATOR_TYPE',
            'CHAT2DESK_WEBHOOK_COMMON_REQUEST_TYPE', 'CHAT_HOURLY_REQUEST_TYPE',
            'build_chat_webhook_request_rows', '_chat_hourly_number', '_chat_hourly_response_sums',
        }, ns)
        rows = [client(at(9)), client(at(9, 0, 20)), reply(at(9, 2), VER_A), reply(at(9, 3), VER_A),
                client(at(9, 30)), reply(at(9, 37), VER_A), client(at(9, 50)), reply(at(9, 51), VER_A)]
        events = [{'request_id': 'r1', 'event_at': row['at'], 'hook_type': 'inbox' if not row['is_echo'] else 'outbox',
                   'message_type': (ns['CHAT2DESK_WEBHOOK_OPERATOR_TYPE'] if row['is_echo']
                                    else ns['CHAT2DESK_WEBHOOK_CLIENT_TYPE']),
                   'c2d_operator_id': row['user_id'], 'request_type': 'common'} for row in rows]
        c2d = ns['build_chat_webhook_request_rows'](events, DAY.isoformat(), {})
        first_sum, first_count, inner_sum, inner_count = ns['_chat_hourly_response_sums'](c2d)
        sums = day_block(rows)['today']['reply_sums']
        self.assertEqual((sums['first_sum'], sums['first_count']), (first_sum, first_count))
        self.assertAlmostEqual(sums['inner_sum'], round(inner_sum, 1), places=1)
        self.assertEqual(sums['inner_count'], inner_count)

    def test_delay_is_credited_to_whoever_answered(self):
        """В чате пишут по двое: первый ответ — первому ответившему, задержки — каждому своя."""
        rows = [client(at(11)), reply(at(11, 0, 30), VER_A), client(at(11, 5)), reply(at(11, 7), VER_B)]
        block = day_block(rows)
        self.assertEqual(block['today']['first_reply_seconds'], 30.0)
        self.assertEqual(block['today']['inner_reply_seconds'], 75.0)      # (30 + 120) / 2
        a, b = block['operators'][VER_A], block['operators'][VER_B]
        self.assertEqual((a['first_reply_seconds'], a['inner_reply_seconds']), (30.0, 30.0))
        self.assertEqual((b['first_reply_seconds'], b['inner_reply_seconds']), (None, 120.0))
        # Итог направления — по диалогам, как у СЗоВ по обращениям: диалог один.
        self.assertEqual(block['today']['reply_sums']['inner_count'], 1)

    def test_team_inner_is_the_dialog_average_not_the_average_of_people(self):
        """А ответил десять раз по 10 с, Б один раз за 20 мин: у диалога среднее ~2 мин, как
        посчитал бы строитель СЗоВ, а не 10 мин — «среднее двух людей»."""
        rows = []
        for i in range(10):
            rows += [client(at(9, i * 2)), reply(at(9, i * 2, 10), VER_A)]
        rows += [client(at(9, 30)), reply(at(9, 50), VER_B)]
        block = day_block(rows)
        self.assertEqual(block['today']['inner_reply_seconds'], round((100 + 1200) / 11))
        self.assertEqual(block['operators'][VER_B]['inner_reply_seconds'], 1200)

    def test_unlinked_author_closes_the_wait_but_is_nobody_s_answer(self):
        """Новый сотрудник, которого ещё не привязали в карте авторов: клиенту ответили, «ждущим» он
        не висит, но и верификаторам ответ не засчитан. Шаблон такого автора — рассылка."""
        unlinked = reply(at(15, 1), None)
        unlinked.update(author_id='13950986', type='text')
        rows = [client(at(15)), unlinked, client(at(15, 3)), reply(at(15, 5), VER_A)]
        block = day_block(rows)
        self.assertIsNone(block['today']['first_reply_seconds'])
        self.assertEqual(block['today']['inner_reply_seconds'], 120)
        broadcast = reply(at(16, 1), None, chat='c2')
        broadcast.update(author_id='8403694', type='wapi_template')
        rows = [client(at(16), chat='c2'), broadcast, reply(at(16, 4), VER_A, chat='c2')]
        self.assertEqual(day_block(rows)['today']['first_reply_seconds'], 240)

    def test_other_group_answers_do_not_count_for_verifiers(self):
        """Сотрудник «Основы» ответил первым: первый ответ не верификаторов, диалог — их (писали)."""
        rows = [client(at(12)), reply(at(12, 0, 10), OTHER), client(at(12, 2)), reply(at(12, 5), VER_A)]
        block = day_block(rows)
        self.assertEqual(block['today']['chats'], 1)
        self.assertIsNone(block['today']['first_reply_seconds'])
        self.assertEqual(block['today']['inner_reply_seconds'], 180.0)
        self.assertNotIn(OTHER, block['operators'])

    def test_dialog_run_only_by_another_group_is_not_a_verifier_chat(self):
        block = day_block([client(at(12)), reply(at(12, 1), OTHER)])
        self.assertEqual(block['today']['chats'], 0)

    def test_broadcasts_neither_answer_nor_start_a_dialog(self):
        """Бот и отправка без привязанного автора — рассылка: клиент по-прежнему ждёт."""
        rows = [client(at(13)), reply(at(13, 0, 5), None), reply(at(13, 0, 6), 999, is_bot=True),
                reply(at(13, 4), VER_A)]
        self.assertEqual(day_block(rows)['today']['first_reply_seconds'], 240.0)
        self.assertEqual(day_block([reply(at(13), None)])['today']['chats'], 0)

    def test_unanswered_client_is_a_chat_without_reply_time(self):
        today = day_block([client(at(14))])['today']
        self.assertEqual(today['chats'], 1)
        self.assertIsNone(today['first_reply_seconds'])
        self.assertIsNone(today['inner_reply_seconds'])


class DialogBoundaryTests(unittest.TestCase):
    def test_night_wait_of_an_unanswered_client_is_one_dialog(self):
        """Написал в 02:00, ответили в 09:00 — первый ответ через семь часов, а не два диалога."""
        block = day_block([client(at(2)), reply(at(9), VER_A)])
        self.assertEqual(block['today']['chats'], 1)
        self.assertEqual(block['today']['first_reply_seconds'], 7 * 3600)
        self.assertEqual(block['hourly'][2]['chats'], 1)                 # диалог — час своего начала
        # Больше двенадцати часов — уже новое обращение, а не ответ на то сообщение.
        block = day_block([client(at(2, day=DAY - timedelta(days=1))), reply(at(9), VER_A),
                           client(at(9, 1)), reply(at(9, 2), VER_A)])
        self.assertEqual(block['today']['first_reply_seconds'], 60)
        # Если клиенту уже ответили, «Рахмет» и через шесть часов новое обращение — как раньше.
        rows = [client(at(8)), reply(at(8, 1), VER_A), client(at(8, 2)), reply(at(14, 30), VER_A)]
        self.assertEqual(day_block(rows)['today']['chats'], 2)

    def test_six_hours_of_silence_split_the_dialog(self):
        rows = [client(at(8)), reply(at(8, 1), VER_A), client(at(14, 1)), reply(at(14, 2), VER_A)]
        self.assertEqual(day_block(rows)['today']['chats'], 2)
        rows = [client(at(8)), reply(at(8, 1), VER_A), client(at(13, 59)), reply(at(14), VER_A)]
        self.assertEqual(day_block(rows)['today']['chats'], 1)

    def test_dialog_belongs_to_the_day_it_started(self):
        """Начатый вчера в 23:50 — вчерашний; но чат человека считается по дню его сообщения."""
        yesterday = DAY - timedelta(days=1)
        rows = [client(at(23, 50, day=yesterday)), reply(at(0, 5), VER_A)]
        block = day_block(rows)
        self.assertEqual(block['today']['chats'], 0)
        self.assertIsNone(block['today']['first_reply_seconds'])
        self.assertEqual(block['operators'][VER_A]['chats'], 1)

    def test_person_chats_are_chat_by_day_pairs(self):
        rows = [client(at(9)), reply(at(9, 1), VER_A), reply(at(9, 2), VER_A),
                client(at(9, 3), chat='c2'), reply(at(9, 4), VER_A, chat='c2'),
                reply(at(18), VER_A)]                                   # тот же c1 вечером
        block = day_block(rows)
        self.assertEqual(block['operators'][VER_A]['chats'], 2)
        self.assertEqual(block['today']['chats'], 3)                    # вечер c1 — новый диалог

    def test_equal_timestamps_put_the_client_first(self):
        """База отдаёт строки без сортировки, а message_id — UUID: в одну миллисекунду реплика
        клиента идёт раньше ответа, каким бы ни был id."""
        first = client(at(10))
        first['message_id'] = 'zzz'
        answer = reply(at(10), VER_A)
        answer['message_id'] = 'aaa'
        self.assertEqual(day_block([answer, first])['today']['first_reply_seconds'], 0)

    def test_reply_in_the_next_hour_is_kept_by_the_hour_the_dialog_started(self):
        """Клиент в 10:58, ответ в 11:03: время ответа — в часе 10 и у человека, писал он в 11."""
        block = day_block([client(at(10, 58)), reply(at(11, 3), VER_A)])
        ten, eleven = block['hourly'][10], block['hourly'][11]
        self.assertEqual(ten['first_reply_seconds'], 300)
        self.assertEqual([p['user_id'] for p in ten['operators_reply']], [VER_A])
        self.assertEqual(ten['operators'], [])
        self.assertEqual([p['user_id'] for p in eleven['operators']], [VER_A])
        self.assertEqual(eleven['operators_reply'], [])

    def test_hours_run_to_now_and_the_current_one_is_partial(self):
        rows = [client(at(9, 10)), reply(at(9, 12), VER_A), client(at(10, 5)), reply(at(10, 6), VER_B)]
        hourly = day_block(rows, now=at(10, 30))['hourly']
        self.assertEqual(len(hourly), 11)
        self.assertEqual((hourly[9]['chats'], hourly[9]['verifiers']), (1, 1))
        self.assertEqual(hourly[10]['operators'][0]['name'], NAMES[VER_B])
        self.assertNotIn('first_sum', hourly[10]['operators'][0])       # кто писал — без сумм
        self.assertTrue(hourly[10]['partial'])
        self.assertFalse(hourly[9]['partial'])
        self.assertEqual(len(day_block(rows)['hourly']), 24)           # прошедшие сутки — целиком


class NowTests(unittest.TestCase):
    def snapshot(self, rows, now, last=None):
        return C.assemble(rows, now, verifier_ids=NAMES, names=NAMES, last_message_at=last or now)

    def test_in_work_window_and_owner_is_the_last_writer(self):
        now = at(15)
        rows = [client(at(14, 50)), reply(at(14, 52), VER_A),                       # в работе у А
                client(at(14, 30), chat='c2'), reply(at(14, 31), VER_B, chat='c2'),  # 29 мин — нет
                client(at(14, 55), chat='c3'), reply(at(14, 56), VER_A, chat='c3'),
                reply(at(14, 57), OTHER, chat='c3')]                                # теперь чат «Основы»
        current = self.snapshot(rows, now)['now']
        self.assertEqual(current['chats_in_work'], 1)
        self.assertEqual(current['verifiers_in_work'], 1)
        people = {p['user_id']: p for p in current['operators']}
        self.assertEqual(people[VER_A]['in_work'], 1)
        self.assertEqual(people[VER_B]['in_work'], 0)
        self.assertEqual(current['operators'][0]['user_id'], VER_A)            # с работой — выше

    def test_chat_is_not_kept_on_someone_who_left_hours_ago(self):
        """Ответил в 14:01 и ушёл; «Рахмет» в 19:50 — чат в работе, но уже не его."""
        current = self.snapshot([client(at(14)), reply(at(14, 1), VER_A), client(at(19, 50))],
                                at(19, 55))['now']
        self.assertEqual(current['chats_in_work'], 1)
        self.assertEqual(current['verifiers_in_work'], 0)

    def test_waiting_is_capped_at_an_hour(self):
        """«Рахмет» пятичасовой давности не должен висеть на стене весь день."""
        now = at(15)
        rows = [client(at(14, 30)), reply(at(14, 31), VER_A), client(at(14, 40)),   # ждёт 20 мин
                client(at(13, 20), chat='c2'), reply(at(13, 21), VER_A, chat='c2'),
                client(at(13, 25), chat='c2')]                                      # ждёт 95 мин
        current = self.snapshot(rows, now)['now']
        self.assertEqual(current['chats_waiting'], 1)
        self.assertEqual(current['longest_wait_seconds'], 20 * 60)

    def test_stream_silence_is_judged_by_the_hour_it_started(self):
        self.assertTrue(C.stream_state(at(14, 30), at(15))['silent'])       # день: 30 мин — обрыв
        self.assertFalse(C.stream_state(at(3, 30), at(5))['silent'])        # ночь: полтора часа — норма
        self.assertFalse(C.stream_state(at(8, 40), at(9, 5))['silent'])     # ночная пауза не будит в 09:00
        self.assertTrue(C.stream_state(None, at(5))['silent'])

    def test_snapshot_carries_owner_norms(self):
        snap = self.snapshot([client(at(9))], at(9, 5))
        self.assertEqual((snap['first_target_seconds'], snap['inner_target_seconds']), (60, 240))
        self.assertEqual((snap['in_work_minutes'], snap['waiting_max_minutes']), (15, 60))
        self.assertNotIn('status', str(sorted(snap['now'])))


class ExportTests(unittest.TestCase):
    def test_period_rules(self):
        today = DAY
        self.assertEqual(X.parse_period('', '', today), [today])
        self.assertEqual(len(X.parse_period('2026-09-20', '2026-09-18', today)), 3)      # перепутали
        self.assertEqual(X.parse_period('2026-09-24', '2026-09-30', today)[-1], today)  # хвост отрезан
        with self.assertRaises(ValueError):
            X.parse_period('2026-10-01', '', today)
        with self.assertRaises(ValueError):
            X.parse_period('2026-08-01', '2026-09-25', today)
        with self.assertRaises(ValueError):
            X.parse_period('25.09.2026', '', today)
        with self.assertRaises(ValueError) as refused:                  # переписка хранится 45 суток
            X.parse_period('2026-08-01', '2026-08-20', today)
        self.assertIn('45 суток', str(refused.exception))
        self.assertEqual(X.parse_period('2026-08-12', '2026-08-12', today)[0], date(2026, 8, 12))

    def test_roster_is_taken_per_day(self):
        """Перешёл в верификаторы 25-го: 24-го его ответы — ответы другой группы."""
        yesterday = DAY - timedelta(days=1)
        rows = [client(at(10, day=yesterday)), reply(at(10, 2, day=yesterday), VER_B),
                client(at(11)), reply(at(11, 1), VER_B)]
        blocks = C.period_days(rows, [yesterday, DAY], verifier_ids=NAMES, names=NAMES,
                               verifiers_by_day={yesterday: {VER_A}, DAY: {VER_A, VER_B}})
        self.assertEqual(blocks[0]['today']['chats'], 0)
        self.assertEqual(blocks[1]['today']['first_reply_seconds'], 60)

    def test_workbook_has_the_sheets_and_the_period_total(self):
        yesterday = DAY - timedelta(days=1)
        rows = [client(at(10, day=yesterday)), reply(at(10, 2, day=yesterday), VER_A),
                client(at(11)), reply(at(11, 1), VER_B)]
        blocks = C.period_days(rows, [yesterday, DAY], verifier_ids=NAMES, names=NAMES)
        book = load_workbook(BytesIO(X.workbook(blocks, first_target_seconds=60, inner_target_seconds=240)))
        self.assertEqual(book.sheetnames, ['Показатели', 'По часам', 'Верификаторы', 'Чаты по часам',
                                           'Первый ответ по часам', 'Ответ внутри чата по часам'])
        summary = book['Показатели']
        total = [row for row in summary.iter_rows(values_only=True) if row[0] == 'Итого'][0]
        self.assertEqual(total[1], 2)
        self.assertEqual(total[2], 1.5)                                  # (120 + 60) / 2 с — по диалогам
        # Ответ, пришедший в следующем часу, стоит в сетке по часу начала диалога.
        grid = [row for row in book['Первый ответ по часам'].iter_rows(values_only=True)]
        self.assertIn('10–11', grid[0])
        self.assertEqual(X.file_name([yesterday, DAY]), 'op_wallboard_chat_20260924_20260925.xlsx')


class _FakeDb:
    def __init__(self):
        self.departments = [{'id': 7, 'code': 'op'}, {'id': 8, 'code': 'szov'}]

    def get_departments(self):
        return self.departments

    def get_user_department_id(self, user_id):
        return None

    def _get_cursor(self):
        @contextmanager
        def cursor():
            yield object()
        return cursor()


def _passthrough_cache(*, fetch, before=None, after=None, **_ignored):
    return dict(fetch(), stale=False, age_seconds=0)


class RouteTests(unittest.TestCase):
    def setUp(self):
        self.requester = {'role': 'operator', 'headed': None}
        now = datetime.now().replace(microsecond=0)
        self.rows = [client(now - timedelta(minutes=5)), reply(now - timedelta(minutes=4), VER_A)]
        self.loaded = []
        for name, value in (
                ('load_chat_verifiers', lambda cursor, day_from, day_to=None: dict(NAMES)),
                ('load_chat_verifier_days', lambda cursor, day_from, day_to=None: (dict(NAMES), {})),
                ('load_chat_messages',
                 lambda cursor, since, until, account='op': self.loaded.append((since, until)) or list(self.rows)),
                ('load_chat_last_message_at', lambda cursor, account='op': now)):
            patcher = mock.patch.object(op_routes, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        db = _FakeDb()
        department_id = op_routes.make_department_resolver(db)
        guard = op_routes.make_guard(
            db=db, department_id=department_id,
            get_authenticated_requester=lambda: (10, (None, None, None, self.requester['role']), None),
            normalize_user_role=lambda role: role,
            is_global_admin_requester=lambda role, uid: role == 'admin' and self.requester['headed'] is None,
            headed_department_id=lambda uid: self.requester['headed'],
            is_supervisor_role=lambda role: role == 'sv')
        app = Flask(__name__)
        self.bp = op_routes.build_op_wallboard_blueprint(
            db=db, require_api_key=lambda fn: fn, build_cors_preflight_response=lambda: ('', 204),
            guard=guard, department_id=department_id, snapshot_with_cache=_passthrough_cache,
            restore_cache=lambda *a, **k: None, persist_cache=lambda *a, **k: None,
            status_entry=lambda key: ('—', 'unknown', 90), load_people=lambda dept, day: [],
            live_statuses=lambda ids: {})
        app.register_blueprint(self.bp)
        self.client = app.test_client()

    def test_operator_is_refused(self):
        self.assertEqual(self.client.get('/api/op_wallboard/chat_snapshot').status_code, 403)
        self.assertEqual(self.client.get('/api/op_wallboard/chat_export').status_code, 403)

    def test_head_of_sales_gets_the_chat_board(self):
        self.requester.update(role='admin', headed=7)
        response = self.client.get('/api/op_wallboard/chat_snapshot')
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual(body['today']['chats'], 1)
        self.assertEqual(body['now']['operators'][0]['name'], NAMES[VER_A])
        self.requester.update(headed=8)
        self.assertEqual(self.client.get('/api/op_wallboard/chat_snapshot').status_code, 403)

    def test_export_is_an_xlsx_and_a_bad_period_is_a_400(self):
        self.requester.update(role='admin')
        response = self.client.get('/api/op_wallboard/chat_export')
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response.headers['Content-Type'])
        load_workbook(BytesIO(response.data))
        bad = self.client.get('/api/op_wallboard/chat_export?date_from=2020-01-01&date_to=2020-12-31')
        self.assertEqual(bad.status_code, 400)
        self.assertIn('суток', bad.get_json()['error'])

    def test_export_reads_past_midnight_of_the_last_day(self):
        """Ответ в 00:03 на диалог из 23:50 последних суток — этих суток: читаем с запасом паузы."""
        self.requester.update(role='admin')
        day = (datetime.now() - timedelta(days=2)).date()
        self.client.get(f'/api/op_wallboard/chat_export?date_from={day}&date_to={day}')
        since, until = self.loaded[-1]
        self.assertEqual(since, datetime.combine(day, datetime.min.time()) - timedelta(hours=6))
        self.assertEqual(until, datetime.combine(day + timedelta(days=1), datetime.min.time())
                         + timedelta(hours=6))

    def test_broadcast_reads_the_same_cached_snapshot(self):
        self.assertTrue(callable(self.bp.chat_snapshot))
        self.assertIn('_op_chat_wallboard_snapshot = _op_wallboard_bp.chat_snapshot', BOT_SOURCE)
        self.assertIn('_op_chat_wallboard_day = _op_wallboard_bp.chat_day', BOT_SOURCE)
        self.assertEqual(self.bp.chat_day(DAY)['day'], DAY.isoformat())


class LoaderSqlTests(unittest.TestCase):
    """SQL проверен EXPLAIN на боевой схеме 25.09.2026; здесь — то, что легко сломать правкой."""

    def test_messages_query_uses_almaty_bounds_and_skips_the_db_sort(self):
        import inspect
        source = inspect.getsource(op_routes.load_chat_messages)
        source = source[source.index('cursor.execute('):]          # сам SQL, без докстринга
        self.assertEqual(source.count("::timestamp AT TIME ZONE 'Asia/Almaty'"), 2)
        self.assertIn('NOT w.is_deleted', source)
        self.assertNotIn('ORDER BY', source)

    def test_verifiers_are_the_group_model_not_the_direction(self):
        self.assertEqual(C.VERIFIER_GROUP_MODEL, 'op_verificator')
        from op_wallboard import snapshot as S
        self.assertIn(C.VERIFIER_GROUP_MODEL, S.OFF_BOARD_GROUP_MODELS)   # с телефонного табло не вернуть


class BroadcastTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import os
        cls.ns = _load_names(BOT_SOURCE, {
            '_env_int', 'OP_CHAT_BROADCAST_MIN_CHATS', '_szov_wallboard_int', '_szov_plural',
            '_szov_format_seconds_mmss', '_szov_format_age_ru', '_szov_broadcast_stale_note',
            '_op_chat_broadcast_duration', '_op_chat_broadcast_stream_note',
            '_op_chat_broadcast_deviations', '_op_chat_broadcast_notes',
        }, {'os': os, 'logging': __import__('logging')})

    def data(self, chats=20, first=50, inner=200, silent=False, waiting=0):
        return {'today': {'chats': chats, 'first_reply_seconds': first, 'inner_reply_seconds': inner},
                'now': {'chats_in_work': 4, 'chats_waiting': waiting, 'longest_wait_seconds': 90},
                'stream': {'silent': silent, 'silent_seconds': 1800 if silent else 5},
                'first_target_seconds': 60, 'inner_target_seconds': 240,
                'snapshot_stale': False, 'snapshot_age_seconds': 0}

    def test_deviation_is_what_is_red_on_the_board(self):
        deviations = self.ns['_op_chat_broadcast_deviations']
        self.assertEqual(deviations(self.data()), [])
        notes = deviations(self.data(first=130, inner=300))
        self.assertEqual(len(notes), 2)
        self.assertIn('первый ответ дольше нормы', notes[0])
        self.assertIn('ответ внутри чата дольше нормы', notes[1])

    def test_small_night_sample_does_not_wake_anyone(self):
        self.assertEqual(self.ns['_op_chat_broadcast_deviations'](self.data(chats=3, first=400)), [])

    def test_frozen_snapshot_is_blamed_on_our_database_not_wazzup(self):
        data = self.data()
        data.update(snapshot_stale=True, snapshot_age_seconds=900)
        note = self.ns['_op_chat_broadcast_deviations'](data)[-1]
        self.assertIn('базы iCORE', note)
        self.assertNotIn('Wazzup', note)

    def test_midnight_report_is_about_yesterday_as_a_whole(self):
        """В 00:00 снимок уже живёт новыми сутками: плановая отбивка берёт итог вчера целиком."""
        from zoneinfo import ZoneInfo
        ns = dict(self.ns)
        _load_names(BOT_SOURCE, {'OP_BROADCAST_TIMEZONE', '_op_chat_broadcast_collect'}, ns)
        ns.update(datetime=datetime, timedelta=timedelta, ZoneInfo=ZoneInfo,
                  _op_chat_wallboard_snapshot=lambda: {'today': {'chats': 0}, 'now': {}},
                  _op_chat_wallboard_day=lambda day: {'today': {'chats': 480, 'first_reply_seconds': 90}})
        midnight = datetime(2026, 9, 26, 0, 0, tzinfo=ZoneInfo('Asia/Almaty'))
        data = ns['_op_chat_broadcast_collect'](True, midnight)
        self.assertEqual((data['period_label'], data['today']['chats']), ('25.09', 480))
        notes = self.ns['_op_chat_broadcast_notes'](dict(self.data(), **{k: data[k] for k in ('today', 'period_label')}))
        self.assertIn('За 25.09: 480 чатов', notes[-1])
        self.assertIsNone(ns['_op_chat_broadcast_collect'](False, midnight)['period_label'])

    def test_silent_wazzup_is_a_deviation_and_goes_last(self):
        data = self.data(silent=True, waiting=2)
        self.assertEqual(len(self.ns['_op_chat_broadcast_deviations'](data)), 1)
        notes = self.ns['_op_chat_broadcast_notes'](data)
        self.assertIn('Wazzup не присылает сообщения', notes[-1])
        self.assertIn('4 чата в работе, 2 ждут ответа', notes[0])

    def test_direction_is_registered_everywhere(self):
        """Список направлений лежит в трёх местах плюс CHECK — 16.09.2026 его уже забывали."""
        self.assertIn("SZOV_BROADCAST_DIRECTION_OP_CHAT = 'op_chat'", BOT_SOURCE)
        self.assertIn('SZOV_BROADCAST_DIRECTION_OP_CHAT)', BOT_SOURCE)
        self.assertIn("SZOV_BROADCAST_DIRECTIONS = ('osnova', 'chat', 'op', 'tez', 'op_chat')", DB_SOURCE)
        self.assertEqual(DB_SOURCE.count("CHECK (direction IN ('osnova', 'chat', 'op', 'tez', 'op_chat'))"), 2)
        self.assertIn("position('''op_chat''' IN pg_get_constraintdef(oid)) = 0", DB_SOURCE)
        guard = BOT_SOURCE[BOT_SOURCE.index('def _szov_broadcast_guard():'):]
        guard = guard[:guard.index('\ndef ', 10)]
        self.assertIn('if direction in (SZOV_BROADCAST_DIRECTION_OP, SZOV_BROADCAST_DIRECTION_OP_CHAT):', guard)
        for fragment in ('times = _op_chat_broadcast_send_times()', 'return _op_chat_broadcast_preview()',
                         'SZOV_BROADCAST_DIRECTION_OP_CHAT: _op_chat_broadcast_send,',
                         "id=f'op_chat_wallboard_broadcast_{_hour:02d}{_minute:02d}'"):
            self.assertIn(fragment, BOT_SOURCE, fragment)


class PersonalBroadcastTests(unittest.TestCase):
    """Отбивка «Чата» лично себе — как у «Тез КЦ»: супер-админ, админ без отдела, админ — глава
    отдела продаж; у каждого табло своя подписка."""

    def test_admin_gets_the_row_and_it_reads_the_chat_subscription(self):
        from tests.test_tez_break_violations import _PersonalDb, _owner_namespace
        db = _PersonalDb({300: {'enabled': True, 'mode': 'deviations'}})
        ns = _owner_namespace((300, 777, 'Глава ОП', 'admin'), db=db)
        self.assertEqual(ns['_szov_broadcast_personal_state']('op_chat'),
                         {'enabled': True, 'mode': 'deviations', 'telegram_connected': True})
        self.assertEqual(db.last_direction, 'op_chat')
        # У линии ОП личной отбивки по-прежнему нет.
        self.assertIsNone(ns['_szov_broadcast_personal_state']('op'))

    def test_supervisor_gets_it_only_by_name_and_only_for_the_chat(self):
        """СВ по общему правилу отбивку не получает; поимённо допущенный (постановщик #367) — да,
        и только у «Чата»: у линии ОП и у «Тез КЦ» он остаётся СВ."""
        from tests.test_tez_break_violations import _owner_namespace
        named = _owner_namespace((402, 777, 'СВ верификаторов', 'sv'))
        self.assertIsNotNone(named['_szov_broadcast_personal_owner']('op_chat'))
        self.assertIsNone(named['_szov_broadcast_personal_owner']('tez'))
        other = _owner_namespace((403, 777, 'Другой СВ', 'sv'))
        self.assertIsNone(other['_szov_broadcast_personal_owner']('op_chat'))

    def _guard(self, requester_id, role, direction):
        from types import SimpleNamespace
        flags = SimpleNamespace()
        ns = {
            'jsonify': lambda payload: payload, 'g': flags,
            '_get_authenticated_requester': lambda: (requester_id, (requester_id, 777, None, role), None),
            '_normalize_user_role': lambda value: str(value or '').strip().lower(),
            '_is_global_admin_requester': lambda r, rid: r in ('admin', 'super_admin'),
            '_headed_department_id': lambda rid: None,
            '_op_wallboard_department_id': lambda: 367,
            '_szov_wallboard_department_id': lambda: 1,
            'request': type('Req', (), {'method': 'GET',
                                        'get_json': staticmethod(lambda silent=True: None)})(),
            '_szov_broadcast_direction_arg': lambda payload=None: direction,
        }
        _load_names(BOT_SOURCE, {'SZOV_BROADCAST_DIRECTION_LINE', 'SZOV_BROADCAST_DIRECTION_OP',
                                 'SZOV_BROADCAST_DIRECTION_TEZ', 'SZOV_BROADCAST_DIRECTION_OP_CHAT',
                                 'OP_CHAT_BROADCAST_PERSONAL_ONLY_USER_IDS', '_szov_broadcast_guard',
                                 '_szov_broadcast_is_personal_only'}, ns)
        _, refusal = ns['_szov_broadcast_guard']()
        return refusal, getattr(flags, 'szov_broadcast_personal_only', False)

    def test_named_supervisor_passes_the_chat_guard_marked_personal_only(self):
        self.assertEqual(self._guard(402, 'sv', 'op_chat'), (None, True))
        refusal, _ = self._guard(402, 'sv', 'op')            # линия ОП — по-прежнему нет
        self.assertEqual(refusal[1], 403)
        refusal, _ = self._guard(403, 'sv', 'op_chat')       # другой СВ — нет
        self.assertEqual(refusal[1], 403)
        self.assertEqual(self._guard(5, 'admin', 'op_chat'), (None, False))   # админу — всё

    def test_personal_only_request_never_touches_groups(self):
        handler = BOT_SOURCE[BOT_SOURCE.index('def api_szov_wallboard_broadcast():'):]
        handler = handler[:handler.index('\n@app.route')]
        self.assertIn('personal_only = _szov_broadcast_is_personal_only()', handler)
        self.assertLess(handler.index('return jsonify({"error": SZOV_BROADCAST_PERSONAL_ONLY_ERROR}), 403'),
                        handler.index("if request.method in ('POST', 'DELETE'):"))
        only = handler[handler.index('    if personal_only:\n        # Групп'):]
        only = only[:only.index('    return jsonify({\n        "direction": direction,\n        "groups_allowed": True')]
        self.assertIn('"groups_allowed": False', only)
        self.assertIn('"recipients": [], "history": [], "chats": [], "groups": []', only)
        self.assertNotIn('get_szov_broadcast_chats', only)
        test_send = BOT_SOURCE[BOT_SOURCE.index('def api_szov_wallboard_broadcast_test():'):]
        test_send = test_send[:test_send.index('\n@app.route')]
        self.assertIn('elif _szov_broadcast_is_personal_only():\n        return jsonify({"error": '
                      'SZOV_BROADCAST_PERSONAL_ONLY_ERROR}), 403', test_send)

    def test_named_people_are_personal_recipients_and_the_lists_match(self):
        import re
        self.assertIn('extra_user_ids=OP_CHAT_BROADCAST_PERSONAL_ONLY_USER_IDS', BOT_SOURCE)
        self.assertIn('OR u.id = ANY(%s::int[])', DB_SOURCE)
        app = (ROOT / 'src' / 'App.jsx').read_text(encoding='utf-8-sig')
        backend = set(map(int, re.search(r'OP_CHAT_BROADCAST_PERSONAL_ONLY_USER_IDS = \{([\d, ]+)\}',
                                         BOT_SOURCE).group(1).split(',')))
        frontend = set(map(int, re.search(r'OP_CHAT_BROADCAST_PERSONAL_ONLY_USER_IDS = new Set\(\[([\d, ]+)\]\)',
                                          app).group(1).split(',')))
        self.assertEqual(backend, frontend)
        self.assertIn('canPersonalChatBroadcast={canReceiveOpChatBroadcastPersonallyForUser(user)}', app)

    def test_scheduled_job_adds_personal_recipients_of_the_sales_department(self):
        calls = []

        class Db:
            def get_tez_broadcast_personal_recipients(self, department_id=None, direction='tez',
                                                      extra_user_ids=None):
                calls.append((department_id, direction, sorted(extra_user_ids or ())))
                return [{'id': 1, 'telegram_id': 777, 'mode': 'always'}]

        ns = _load_names(BOT_SOURCE, {'SZOV_BROADCAST_DIRECTION_OP_CHAT',
                                      'OP_CHAT_BROADCAST_PERSONAL_ONLY_USER_IDS',
                                      '_op_chat_broadcast_personal_recipients'},
                         {'db': Db(), 'logging': __import__('logging'),
                          '_op_wallboard_department_id': lambda: 367})
        self.assertEqual(len(ns['_op_chat_broadcast_personal_recipients']()), 1)
        self.assertEqual(calls, [(367, 'op_chat', [402])])
        job = BOT_SOURCE[BOT_SOURCE.index('async def op_chat_broadcast_job():'):]
        job = job[:job.index('\n\n\n')]
        self.assertIn('personal=_op_chat_broadcast_personal_recipients', job)

    def test_subscription_has_its_own_columns(self):
        self.assertIn("'op_chat': ('op_chat_broadcast_personal_enabled', 'op_chat_broadcast_personal_mode'),",
                      DB_SOURCE)
        self.assertIn('ALTER TABLE admin_profiles ADD COLUMN IF NOT EXISTS '
                      'op_chat_broadcast_personal_enabled BOOLEAN NOT NULL DEFAULT FALSE;', DB_SOURCE)
        self.assertIn('ALTER TABLE admin_profiles ADD COLUMN IF NOT EXISTS '
                      "op_chat_broadcast_personal_mode VARCHAR(16) NOT NULL DEFAULT 'always';", DB_SOURCE)
        self.assertIn('raise ValueError("У этого табло нет отбивки лично себе")', DB_SOURCE)
        self.assertIn('SZOV_BROADCAST_PERSONAL_DIRECTIONS = (SZOV_BROADCAST_DIRECTION_TEZ, '
                      'SZOV_BROADCAST_DIRECTION_OP_CHAT)', BOT_SOURCE)


class FrontendTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.view = (MONITORING / 'OpWallboardView.jsx').read_text(encoding='utf-8-sig')
        cls.board = (MONITORING / 'OpChatWallboard.jsx').read_text(encoding='utf-8-sig')
        cls.shared = (MONITORING / 'opWallboardShared.js').read_text(encoding='utf-8-sig')
        cls.szov_view = (MONITORING / 'SzovWallboardView.jsx').read_text(encoding='utf-8-sig')
        cls.app = (ROOT / 'src' / 'App.jsx').read_text(encoding='utf-8-sig')

    def test_switch_is_line_and_chat_like_szov(self):
        self.assertIn("{ key: 'line', label: 'Линия'", self.shared)
        self.assertIn("{ key: 'chat', label: 'Чат'", self.shared)
        self.assertIn('<SegmentedSwitch value={direction} options={OP_WALLBOARD_VIEWS} onChange={changeDirection} />',
                      self.view)
        self.assertIn('otp:op-wallboard-direction${userId ? `:${userId}` : \'\'}', self.view)
        self.assertIn('user={user}\n                                    apiBaseUrl={API_BASE_URL}', self.app)

    def test_only_the_open_direction_is_polled(self):
        """Хук линии живёт в компоненте линии, хук чатов — в компоненте чатов; монтируется один."""
        line = self.view[self.view.index('function OpLineWallboard('):
                         self.view.index('// ── Направление «Чат»')]
        self.assertIn('useOpWallboardSnapshot({ apiBaseUrl, withAccessTokenHeader })', line)
        chat = self.view[self.view.index('function OpChatBoard('):
                         self.view.index('export default function OpWallboardView(')]
        self.assertIn('useOpChatWallboardSnapshot({ apiBaseUrl, withAccessTokenHeader })', chat)
        root = self.view[self.view.index('export default function OpWallboardView('):]
        self.assertNotIn('useOpWallboardSnapshot', root)
        self.assertNotIn('useOpChatWallboardSnapshot', root)
        self.assertIn("if (direction === 'chat') {\n        return <OpChatBoard", root)
        self.assertEqual(self.view.count('useOpWallboardSnapshot({'), 1)
        self.assertEqual(self.view.count('useOpChatWallboardSnapshot({'), 1)

    def test_chat_board_has_export_broadcast_and_widget(self):
        chat = self.view[self.view.index('function OpChatBoard('):]
        self.assertIn('exportPath="/api/op_wallboard/chat_export"', chat)
        self.assertIn('maxDays={OP_CHAT_EXPORT_MAX_DAYS}', chat)
        self.assertIn('const OP_CHAT_EXPORT_MAX_DAYS = 31;', self.view)
        self.assertIn('direction="op_chat"', chat)
        self.assertIn('<WidgetButton direction="op_chat"', chat)
        self.assertIn('export const ChatExportControls', self.szov_view)
        # Виджет «Чата» — по правам раздела «Табло ОП», а не СЗоВ.
        self.assertIn("(szovWallboardWidget === 'op' || szovWallboardWidget === 'op_chat')", self.app)
        registry = self.shared[self.shared.index('export const OP_WALLBOARD_DIRECTIONS'):]
        for field in ("key: 'op_chat'", 'useSnapshot: useOpChatWallboardSnapshot',
                      'metrics: OP_CHAT_METRICS', 'freshnessNotice: opChatFreshnessNotice'):
            self.assertIn(field, registry, field)

    def test_personal_only_user_sees_the_button_and_only_the_personal_row(self):
        chat = self.view[self.view.index('function OpChatBoard('):]
        self.assertIn('{canManageBroadcast || canPersonalChatBroadcast ? (', chat)
        self.assertIn("const groupsAllowed = state?.groups_allowed !== false;", self.szov_view)
        self.assertEqual(self.szov_view.count('{groupsAllowed ? ('), 2)   # группы и «Кто менял»

    def test_export_limit_reads_as_russian(self):
        self.assertIn("pluralRu(maxDays, 'сутки', 'суток', 'суток')", self.szov_view)
        self.assertIn("const OP_CHAT_SOURCE_LABEL = 'База iCORE';", self.view)

    def test_no_people_statuses_on_the_chat_board(self):
        """Решение владельца 25.09.2026: у Wazzup статусов нет — и на экране их нет."""
        for word in ('Онлайн', 'Перерыв', 'Тренинг', 'Не в системе', 'CHAT_STATUS_STYLE', 'status_key'):
            self.assertNotIn(word, self.board, word)
        chat_metrics = self.shared[self.shared.index('export const OP_CHAT_METRICS'):
                                   self.shared.index('export const OP_CHAT_METRIC_MAP')]
        self.assertNotIn('Онлайн', chat_metrics)

    def test_only_reply_time_is_coloured(self):
        chat_metrics = self.shared[self.shared.index('export const OP_CHAT_METRICS'):
                                   self.shared.index('export const OP_CHAT_METRIC_MAP')]
        self.assertEqual(chat_metrics.count('tone:'), 2)
        self.assertIn('chatReplyTone(s.today?.first_reply_seconds, s.first_target_seconds)', chat_metrics)
        self.assertIn('chatReplyTone(s.today?.inner_reply_seconds, s.inner_target_seconds)', chat_metrics)


if __name__ == '__main__':
    unittest.main()
