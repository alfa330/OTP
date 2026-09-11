# -*- coding: utf-8 -*-
"""Выгрузка табло Тез КЦ за период: раскладка книги и арифметика периода.

Модуль `tez_wallboard_export` намеренно чистый — ни Flask, ни database, — поэтому здесь он
импортируется обычным import и проверяется настоящая книга, а не её описание в исходнике.
"""
import unittest
from io import BytesIO

from openpyxl import load_workbook

import tez_wallboard_export as export


def _person(name, **stats):
    base = {
        'served': 0, 'missed': 0, 'talk_seconds': 0, 'wait_seconds': 0,
        'avg_talk_seconds': None, 'avg_wait_seconds': None,
        'outgoing_total': 0, 'outgoing_success': 0, 'outgoing_talk_seconds': 0,
        'avg_outgoing_talk_seconds': None, 'outgoing_success_percent': None,
    }
    base.update(stats)
    return {'operator_id': abs(hash(name)) % 1000, 'name': name, 'stats': base}


def _payload(direction='tp', days=None):
    return {
        'direction': direction,
        'direction_label': 'ОП' if direction == 'op' else 'ТП',
        'date_from': '2026-09-01',
        'date_to': '2026-09-02',
        'generated_at': '2026-09-11 12:00',
        'days': days if days is not None else [
            {'day': '2026-09-01', 'people': [
                _person('Первый', served=10, talk_seconds=1000, wait_seconds=100,
                        outgoing_total=5, outgoing_success=2, outgoing_talk_seconds=200),
                _person('Второй', served=30, talk_seconds=900, wait_seconds=300,
                        outgoing_total=1, outgoing_success=1, outgoing_talk_seconds=100),
            ], 'error': None},
            {'day': '2026-09-02', 'people': [
                _person('Первый', served=2, talk_seconds=40, wait_seconds=10),
            ], 'error': None},
        ],
    }


class ColumnsTests(unittest.TestCase):
    """Колонки разные у направлений — как и на самом табло."""

    def test_op_has_no_incoming_columns(self):
        """У отдела продаж входящих нет вовсе: «Принято» стояло бы пустым столбцом."""
        headers = [column[0] for column in export.columns_for('op')]
        self.assertNotIn('Принято', headers)
        self.assertNotIn('Пропущено', headers)
        self.assertIn('Набрано', headers)
        self.assertIn('Дозвонились', headers)

    def test_tp_keeps_the_line_columns(self):
        headers = [column[0] for column in export.columns_for('tp')]
        for expected in ('Принято', 'Пропущено', 'Ср. разговор, сек', 'Ср. ожидание, сек'):
            self.assertIn(expected, headers)

    def test_both_start_with_day_and_person(self):
        for direction in ('tp', 'op'):
            headers = [column[0] for column in export.columns_for(direction)]
            self.assertEqual(headers[:2], ['Дата', 'Сотрудник'], direction)


class DayTotalsTests(unittest.TestCase):
    """Свёртка дня: суммы складываются, средние пересчитываются заново."""

    def test_average_is_counted_from_sums_not_from_averages(self):
        """У одного 10 звонков, у другого 30 — их средние складывать пополам нельзя.

        1900 секунд на 40 звонков = 47, а среднее их средних дало бы (100 + 30) / 2 = 65."""
        people = _payload()['days'][0]['people']
        totals = export.day_totals('tp', people)
        self.assertEqual(totals['served'], 40)
        self.assertEqual(totals['talk_seconds'], 1900)
        self.assertEqual(totals['avg_talk_seconds'], 47)
        self.assertEqual(totals['avg_wait_seconds'], 10)

    def test_zero_calls_give_a_dash_not_a_zero(self):
        """Ноль в средней длительности читался бы как мгновенный разговор."""
        totals = export.day_totals('tp', [_person('Тихий')])
        self.assertIsNone(totals['avg_talk_seconds'])
        self.assertIsNone(totals['avg_outgoing_talk_seconds'])
        self.assertIsNone(totals['outgoing_success_percent'])
        self.assertEqual(totals['served'], 0)

    def test_people_counts_only_those_with_calls(self):
        """«Людей со звонками» — про работу, а не про состав: состав виден на листе поимённо."""
        totals = export.day_totals('tp', [_person('Тихий'), _person('Звонивший', served=3)])
        self.assertEqual(totals['people'], 1)

    def test_dial_rate_is_a_percent_of_dialed(self):
        totals = export.day_totals('op', [_person('Продавец', outgoing_total=8, outgoing_success=2)])
        self.assertEqual(totals['outgoing_success_percent'], 25.0)


class WorkbookTests(unittest.TestCase):
    """Книга целиком: два листа, заголовки, итог периода, помеченные дни."""

    def _book(self, payload):
        return load_workbook(BytesIO(export.build_workbook(payload)))

    def test_two_sheets_with_expected_titles(self):
        book = self._book(_payload())
        self.assertEqual(book.sheetnames, ['Поимённо', 'По дням'])

    def test_person_rows_carry_the_day_and_the_name(self):
        sheet = self._book(_payload())['Поимённо']
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=5)]
        self.assertEqual([row[0] for row in rows], ['2026-09-01', '2026-09-01', '2026-09-02'])
        self.assertEqual(sorted({row[1] for row in rows}), ['Второй', 'Первый'])
        # Принято у «Первого» за 1 сентября — 10, и это его собственная строка
        first = next(row for row in rows if row[0] == '2026-09-01' and row[1] == 'Первый')
        self.assertEqual(first[2], 10)

    def test_period_total_is_counted_over_all_rows(self):
        """Итог периода — от всех строк, а не сложением дневных средних.

        1940 секунд на 42 звонка = 46; сложение средних по дням дало бы (47 + 20) / 2."""
        sheet = self._book(_payload())['По дням']
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=5)]
        total = next(row for row in rows if row[0] == 'Итого')
        self.assertEqual(total[2], 42)
        self.assertEqual(total[5], 46)

    def test_failed_day_is_marked_not_zeroed(self):
        """День, который кабинет не отдал, и день без звонков — разные вещи."""
        payload = _payload(days=[{'day': '2026-09-01', 'people': [], 'error': 'кабинет не ответил'}])
        book = self._book(payload)
        for name in ('Поимённо', 'По дням'):
            values = [cell.value for row in book[name].iter_rows(min_row=5) for cell in row]
            self.assertIn('2026-09-01', values)
            self.assertTrue(any('кабинет не ответил' in str(value) for value in values if value), name)
            self.assertNotIn(0, values, name)

    def test_op_book_has_no_incoming_columns(self):
        payload = _payload(direction='op')
        headers = [cell.value for cell in next(self._book(payload)['Поимённо'].iter_rows(min_row=4, max_row=4))]
        self.assertNotIn('Принято', headers)
        self.assertIn('Дозвон, %', headers)

    def test_header_names_the_direction_and_period(self):
        sheet = self._book(_payload())['Поимённо']
        self.assertIn('ТП', sheet['A1'].value)
        self.assertIn('2026-09-01 — 2026-09-02', sheet['A1'].value)
        self.assertIn('Binotel', sheet['A2'].value)

    def test_durations_are_numbers_not_text(self):
        """Секундами и числом: «2:11» строкой и сортируется как текст, и зажигает в Excel
        зелёный уголок «Число сохранено как текст» — владелец читает его как брак файла."""
        sheet = self._book(_payload())['Поимённо']
        row = next(row for row in sheet.iter_rows(min_row=5, max_row=5))
        for cell in row[2:]:
            if cell.value is not None:
                self.assertIsInstance(cell.value, (int, float), cell.coordinate)


class FileNameTests(unittest.TestCase):
    def test_name_carries_direction_and_period(self):
        self.assertEqual(export.export_file_name(_payload()),
                         'tez_wallboard_tp_20260901-20260902.xlsx')

    def test_single_day_is_not_doubled(self):
        payload = dict(_payload(direction='op'), date_to='2026-09-01')
        self.assertEqual(export.export_file_name(payload), 'tez_wallboard_op_20260901.xlsx')


if __name__ == '__main__':
    unittest.main()
