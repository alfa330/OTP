# -*- coding: utf-8 -*-
"""Пропущенные входящие → amoCRM: тело сделки, поиск контакта, метка в «Касаниях» и SQL.

Что закреплено:
  * сделка — «Новая заявка» воронки «Отдел продаж», тег «Пропущенный входящий», номер
    маской 77XXXXXXXXX; найденный контакт привязывается по id, иначе заводится новый;
  * `metadata` в теле нет — иначе сделка уехала бы в «Неразобранное»;
  * контакт выбирается только при точном совпадении десяти цифр: поиск amoCRM подстрочный;
  * тег берётся по id, если он в справочнике уже есть, — двойник по имени не заводится;
  * повтор после сбоя находит только НАШУ сделку, заведённую после попытки;
  * метка «в amoCRM» в таблице и колонка в Excel читаются из журнала робота, а фильтр
    «только переданные» не множит строк и не меняет итоги;
  * новый SQL разбирается настоящей грамматикой Postgres.

Телефоны учебные (7XX555XXXX).
"""

import re
import unittest
from datetime import datetime

from openpyxl import Workbook

from cdr import missed_amo, missed_config, missed_queries, queries, report, schema

PHONE = '7015550001'


class _Client(object):
    """Двойник AmoClient: GET по пути отдаёт заготовленное тело и запоминает запросы."""

    def __init__(self, bodies=None, fail=None):
        self.bodies = bodies or {}
        self.fail = fail
        self.calls = []

    def get(self, path, params=None):
        self.calls.append((path, params))
        if self.fail:
            raise self.fail
        return self.bodies.get(path)


def contact(contact_id, *phones, updated=0, leads=()):
    return {'id': contact_id, 'updated_at': updated,
            'custom_fields_values': [{'field_id': 892223, 'field_code': 'PHONE',
                                      'values': [{'value': p} for p in phones]}],
            '_embedded': {'leads': [{'id': lead} for lead in leads]}}


class BuildLeadTests(unittest.TestCase):
    def test_new_contact_is_created_with_the_phone(self):
        lead = missed_amo.build_lead(PHONE, {'name': 'Пропущенный входящий'})
        self.assertEqual(lead['name'], '77015550001 - Пропущенный входящий')
        self.assertEqual((lead['pipeline_id'], lead['status_id']),
                         (missed_config.PIPELINE_ID, missed_config.STATUS_ID))
        self.assertEqual(lead['responsible_user_id'], missed_config.RESPONSIBLE_USER_ID)
        self.assertEqual(lead['_embedded']['tags'], [{'name': 'Пропущенный входящий'}])
        field = lead['_embedded']['contacts'][0]['custom_fields_values'][0]
        self.assertEqual(field['values'][0]['value'], '77015550001')
        self.assertNotIn('metadata', lead['_embedded'])

    def test_existing_contact_is_linked_by_id(self):
        lead = missed_amo.build_lead(PHONE, {'id': 725001}, contact_id=75000001)
        self.assertEqual(lead['_embedded']['contacts'], [{'id': 75000001}])
        self.assertEqual(lead['_embedded']['tags'], [{'id': 725001}])

    def test_no_phone_no_lead(self):
        with self.assertRaises(missed_amo.AmoWriteError):
            missed_amo.build_lead('', {'name': 'x'})


class WriterReadTests(unittest.TestCase):
    def test_contact_is_chosen_by_exact_number_not_by_substring(self):
        client = _Client({'/api/v4/contacts': {'_embedded': {'contacts': [
            contact(1, '+77015550001'),             # тот же номер
            contact(2, '87015550001', updated=5),    # тот же номер, изменён позже — он
            contact(3, '+77015550001999'),           # подстрока в чужом номере
        ]}}})
        writer = missed_amo.MissedCallWriter(client)
        self.assertEqual(writer.find_contact_id(PHONE), 2)

    def test_contact_search_failure_means_a_new_contact(self):
        writer = missed_amo.MissedCallWriter(_Client(fail=RuntimeError('502')))
        self.assertIsNone(writer.find_contact_id(PHONE))

    def test_tag_is_taken_by_id_when_it_exists(self):
        client = _Client({'/api/v4/leads/tags': {'_embedded': {'tags': [
            {'id': 1, 'name': 'Пропущенный входящий звонок'},
            {'id': 725001, 'name': 'Пропущенный входящий'},
        ]}}})
        writer = missed_amo.MissedCallWriter(client)
        self.assertEqual(writer.tag_ref(), {'id': 725001, 'name': 'Пропущенный входящий'})
        writer.tag_ref()
        self.assertEqual(len(client.calls), 1, 'найденный id запоминается')

    def test_missing_tag_goes_by_name_and_is_looked_up_again(self):
        client = _Client({'/api/v4/leads/tags': None})
        writer = missed_amo.MissedCallWriter(client)
        self.assertEqual(writer.tag_ref(), {'name': 'Пропущенный входящий'})
        writer.tag_ref()
        self.assertEqual(len(client.calls), 2)

    def test_lead_state_of_a_deleted_lead_is_none(self):
        writer = missed_amo.MissedCallWriter(_Client({'/api/v4/leads/57000001': None}))
        self.assertIsNone(writer.lead_state(57000001))
        gone = RuntimeError('amoCRM GET /api/v4/leads/57000001 -> 404: {"title":"Not Found"}')
        writer = missed_amo.MissedCallWriter(_Client(fail=gone))
        self.assertIsNone(writer.lead_state(57000001))
        deleted = {'/api/v4/leads/57000001': {'is_deleted': True, 'pipeline_id': 1}}
        self.assertIsNone(missed_amo.MissedCallWriter(_Client(deleted)).lead_state(57000001))

    def test_transient_failure_is_not_mistaken_for_a_missing_lead(self):
        # 429 от лимита, который робот делит с роботом OLX, — не «сделки нет»: иначе
        # клиенту завелась бы вторая при живой первой.
        for text in ('amoCRM GET /api/v4/leads/57000001 -> 429: Too Many Requests',
                     'amoCRM GET /api/v4/leads/57000001 -> 502: Bad Gateway',
                     'Connection aborted.'):
            writer = missed_amo.MissedCallWriter(_Client(fail=RuntimeError(text)))
            with self.assertRaises(missed_amo.LeadReadError, msg=text):
                writer.lead_state(57000001)

    def test_lead_state_reads_pipeline_and_stage(self):
        writer = missed_amo.MissedCallWriter(_Client({'/api/v4/leads/57000001': {
            'pipeline_id': 5524684, 'status_id': 48846277}}))
        self.assertEqual(writer.lead_state(57000001), (5524684, 48846277))

    def test_recent_own_lead_needs_our_tag_and_a_fresh_date(self):
        client = _Client({
            '/api/v4/contacts': {'_embedded': {'contacts': [
                contact(1, '77015550001', leads=(10, 11, 12))]}},
            '/api/v4/leads': {'_embedded': {'leads': [
                {'id': 10, 'created_at': 1790000500,
                 '_embedded': {'tags': [{'name': 'call_jana'}]}},              # чужая
                {'id': 11, 'created_at': 1789000000,
                 '_embedded': {'tags': [{'name': 'Пропущенный входящий'}]}},  # старая
                {'id': 12, 'created_at': 1790000400,
                 '_embedded': {'tags': [{'name': 'Пропущенный входящий'}]}},  # наша
            ]}},
        })
        writer = missed_amo.MissedCallWriter(client)
        self.assertEqual(writer.recent_own_lead(PHONE, 1790000000), 12)


def touch_row(status=None, lead=None, reason=None, error=None):
    """Строка выборки: семнадцать колонок _COLUMNS и хвост журнала робота."""
    return (datetime(2026, 9, 28, 10, 0, 0), None, PHONE, '6728', 'Входящий (не приняли)',
            'Занято', 0, 47, '3010', None, '1790571600.1', 1, datetime(2026, 9, 28, 10, 0, 7),
            40, None, '', '7475550078', status, lead, reason, error)


class TouchMarkTests(unittest.TestCase):
    def test_transferred_call_carries_the_lead_and_its_link(self):
        touch = queries._row_to_touch(touch_row('created', 57009999, 'За минуту…'))
        self.assertEqual(touch['amo_state'], 'sent')
        self.assertEqual(touch['amo_lead_id'], 57009999)
        self.assertTrue(touch['amo_url'].endswith('/leads/detail/57009999'))

    def test_chained_call_with_the_lead_is_also_sent(self):
        touch = queries._row_to_touch(touch_row('chained', 57009999))
        self.assertEqual(touch['amo_state'], 'sent')

    def test_failure_is_visible_with_its_text(self):
        touch = queries._row_to_touch(touch_row('failed', None, 'За минуту…', 'передать не удалось'))
        self.assertEqual((touch['amo_state'], touch['amo_note']), ('error', 'передать не удалось'))

    def test_error_still_being_retried_shows_nothing(self):
        # Красная метка «не передан» зовёт человека перезванивать руками — пока робот ещё
        # повторяет, через пару минут сделка может появиться сама.
        touch = queries._row_to_touch(touch_row('error', None, 'За минуту…', 'HTTP 502'))
        self.assertEqual(touch['amo_state'], '')

    def test_answered_explains_why_and_sending_shows_nothing(self):
        self.assertEqual(queries._row_to_touch(touch_row('answered', None, 'Клиент…'))['amo_state'],
                         'answered')
        self.assertEqual(queries._row_to_touch(touch_row('sending'))['amo_state'], '')

    def test_row_without_the_tail_is_a_plain_touch(self):
        touch = queries._row_to_touch(touch_row()[:17])
        self.assertEqual((touch['amo_state'], touch['amo_lead_id'], touch['amo_url']),
                         ('', None, ''))

    def test_filter_joins_the_journal_by_its_primary_key(self):
        sql = ' '.join(queries._FILTER_SQL.split())
        self.assertIn('LEFT JOIN cdr_missed_leads ml ON ml.linkedid = t.linkedid '
                      'AND ml.phone = t.phone', sql)
        self.assertIn('(NOT %(amo_only)s OR ml.amo_lead_id IS NOT NULL)', sql)
        self.assertIn('PRIMARY KEY (linkedid, phone)', schema.CDR_SCHEMA_SQL.split(
            'cdr_missed_leads')[1], 'без первичного ключа JOIN размножил бы строки итогов')

    def test_amo_filter_is_off_by_default(self):
        self.assertFalse(queries._params(None, None, {})['amo_only'])
        self.assertTrue(queries._params(None, None, {'amo_only': True})['amo_only'])

    def test_summary_counts_transferred(self):
        class Cursor(object):
            def execute(self, sql, params=None):
                self.sql = sql

            def fetchone(self):
                return (10, 5, 3, 4, 3, 600, 2, 9, 5, 1, 2)
        cursor = Cursor()
        self.assertEqual(queries.summary(cursor, None, None)['amo_transferred'], 2)
        self.assertIn('ml.amo_lead_id IS NOT NULL', cursor.sql)


class ExportTests(unittest.TestCase):
    def test_lead_column_is_a_link_and_lines_up(self):
        touch = queries._row_to_touch(touch_row('created', 57009999))
        sheet = Workbook(write_only=True).create_sheet('Касания')
        values = report._touch_row(sheet, touch)
        self.assertEqual(len(values), len(report.COLUMNS))
        index = [key for key, _t, _w in report.COLUMNS].index('amo_lead_id')
        self.assertEqual(values[index].value, '57009999')
        self.assertTrue(values[index].hyperlink.target.endswith('/leads/detail/57009999'))
        self.assertIn('amo_lead_id', report.TEXT_COLUMNS)

    def test_empty_when_not_transferred(self):
        touch = queries._row_to_touch(touch_row())
        sheet = Workbook(write_only=True).create_sheet('Касания')
        index = [key for key, _t, _w in report.COLUMNS].index('amo_lead_id')
        self.assertIsNone(report._touch_row(sheet, touch)[index])


class StateTests(unittest.TestCase):
    class Cursor(object):
        def __init__(self, stored):
            self.stored = stored
            self.sql = []

        def execute(self, sql, params=None):
            self.sql.append(' '.join(sql.split()))
            self.params = params

        def fetchone(self):
            if self.sql[-1].startswith('SELECT'):
                return (self.stored,) if self.stored else None
            return (self.params[0],)

    def test_known_start_is_only_read(self):
        # Цикл раз в пятнадцать секунд: писать строку состояния на каждом незачем.
        cursor = self.Cursor(datetime(2026, 9, 28, 9, 0))
        self.assertEqual(missed_queries.ensure_state(cursor, datetime(2026, 9, 28, 12, 0)),
                         datetime(2026, 9, 28, 9, 0))
        self.assertEqual(len(cursor.sql), 1)

    def test_first_cycle_writes_now(self):
        cursor = self.Cursor(None)
        now = datetime(2026, 9, 28, 12, 0)
        self.assertEqual(missed_queries.ensure_state(cursor, now), now)
        self.assertTrue(cursor.sql[-1].startswith('INSERT INTO cdr_missed_state'))


def _pg_parse(sql):
    import pglast
    # Именованные и позиционные параметры psycopg2 → литералы: грамматике нужен SQL.
    text = re.sub(r'%\((\w+)\)s', "'1'", sql)
    text = text.replace('%s', "'1'").replace('%%', '%')
    return pglast.parse_sql(text)


class SqlGrammarTests(unittest.TestCase):
    """Всё новое SQL — через настоящую грамматику Postgres (pglast): опечатка в запросе,
    который крутится раз в пятнадцать секунд, иначе всплыла бы только на проде."""

    def setUp(self):
        try:
            import pglast  # noqa: F401
        except ImportError:  # pragma: no cover
            self.skipTest('pglast не установлен')

    def test_schema_parses(self):
        _pg_parse(schema.CDR_SCHEMA_SQL)
        for statement in schema.CDR_SCHEMA_MIGRATIONS:
            _pg_parse(statement)

    def test_robot_queries_parse(self):
        import inspect
        source = inspect.getsource(missed_queries)
        statements = re.findall(r'cursor\.execute\("""(.*?)"""', source, flags=re.S)
        self.assertGreaterEqual(len(statements), 12)
        for statement in statements:
            _pg_parse(statement)

    def test_touch_selects_with_the_journal_parse(self):
        _pg_parse('SELECT ' + queries._COLUMNS + queries._MISSED_COLUMNS + queries._FILTER_SQL
                  + ' ORDER BY t.started_at')
        _pg_parse('SELECT COUNT(*) ' + queries._FILTER_SQL + queries.COUNTED_SQL)
        _pg_parse('UPDATE cdr_agent_state SET queue_calls = %s, queue_calls_at = NOW() WHERE id = 1')


if __name__ == '__main__':
    unittest.main()
