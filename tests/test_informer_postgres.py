# -*- coding: utf-8 -*-
"""Автоинформатор на настоящем SQL: INFORMER_TEST_PORT (или DIAL_LIST_TEST_PORT) — порт локального
Postgres. Очередь звонков для диспетчера, темп закона, ёмкость линий, повторы, потерянные итоги,
аналитика по частям записи и выгрузка — тем самым InformerService, что работает на портале.

Своя схема со случайным именем, в конце удаляется; номера выдуманы.

    INFORMER_TEST_PORT=55441 python -m pytest tests/test_informer_postgres.py
"""
import contextlib
import hashlib
import io
import os
import struct
import sys
import uuid
from datetime import datetime, timedelta, timezone

import numpy as np
import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from informer import export as informer_export  # noqa: E402
from informer import rules  # noqa: E402
from informer import schema as informer_schema  # noqa: E402
from informer import telephone as informer_telephone  # noqa: E402
from informer import wav as informer_wav  # noqa: E402
from informer.service import InformerError, InformerService  # noqa: E402

PORT = os.environ.get('INFORMER_TEST_PORT') or os.environ.get('DIAL_LIST_TEST_PORT')
pytestmark = pytest.mark.skipif(not PORT, reason='needs local PostgreSQL (INFORMER_TEST_PORT)')

DEPT, OTHER_DEPT, BOSS = 1954, 560, 2
# 15.10.2026 12:00 по Алматы — внутри окна 10:00–20:00.
NOON = datetime(2026, 10, 15, 7, 0, tzinfo=timezone.utc)

BASE_DDL = '''
    CREATE TABLE departments (id integer PRIMARY KEY, name text NOT NULL);
    CREATE TABLE users (id integer PRIMARY KEY, name text NOT NULL);
    INSERT INTO departments VALUES (1954, 'Удаленный КЦ'), (560, 'Тез КЦ');
    INSERT INTO users VALUES (2, 'Владелец');
    -- База «Обзвона» того же отдела: по ней считается «подписал после звонка» (колонки — как в dial_list).
    CREATE TABLE dial_list_leads (id serial PRIMARY KEY, department_id integer NOT NULL, phone_norm text NOT NULL,
                                  signed_at timestamptz, success_attempt_id integer, sign_checked_at timestamptz);
'''


def master_wav_bytes(seconds=25.52):
    """Оригинал: WAV 24 кГц, моно, 16 бит — тон 440 Гц, чтобы телефонной версии было что сводить."""
    t = np.arange(int(informer_wav.MASTER_RATE * seconds)) / informer_wav.MASTER_RATE
    pcm = np.rint(0.3 * np.sin(2 * np.pi * 440 * t) * 32767).astype('<i2').tobytes()
    return informer_wav.encode(pcm, informer_wav.MASTER_RATE)


def wav_bytes(seconds=25.52):
    """WAV 8 кГц, моно, 16 бит: тишина нужной длины — запись прежней страницы, без оригинала."""
    samples = int(8000 * seconds)
    data = b'\x00\x00' * samples
    fmt = struct.pack('<HHIIHH', 1, 1, 8000, 16000, 2, 16)
    return (b'RIFF' + struct.pack('<I', 36 + len(data)) + b'WAVE' + b'fmt ' + struct.pack('<I', 16) + fmt
            + b'data' + struct.pack('<I', len(data)) + data)


PARTS = [{'lang': 'kz', 'start_ms': 500, 'end_ms': 13050}, {'lang': 'ru', 'start_ms': 13650, 'end_ms': 25220}]
LINES_939 = {'939': {'capacity': 2, 'active': 0}}


class _Db:
    def __init__(self, schema):
        self.schema = schema

    def connect(self):
        conn = __import__('psycopg2').connect(host='127.0.0.1', port=int(PORT), user='postgres', dbname='postgres')
        with conn.cursor() as cur:
            cur.execute('SET search_path TO ' + self.schema)
        conn.commit()
        return conn

    @contextlib.contextmanager
    def _get_cursor(self):
        conn = self.connect()
        cur = conn.cursor()
        try:
            yield cur
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()


class Clock:
    def __init__(self, now):
        self.now = now

    def __call__(self):
        return self.now

    def tick(self, **delta):
        self.now = self.now + timedelta(**delta)


@pytest.fixture
def env():
    schema = 't_informer_' + uuid.uuid4().hex
    db = _Db(schema)
    admin = db.connect()
    admin.autocommit = True
    with admin.cursor() as cur:
        cur.execute('CREATE SCHEMA ' + schema)
        cur.execute('SET search_path TO ' + schema)
        cur.execute(BASE_DDL)
        informer_schema.apply(cur)
    informer_schema._reset_for_tests()
    clock = Clock(NOON)
    svc = InformerService(db, clock=clock)
    try:
        yield svc, clock, db
    finally:
        informer_schema._reset_for_tests()
        with admin.cursor() as cur:
            cur.execute('DROP SCHEMA ' + schema + ' CASCADE')
        admin.close()


def make_campaign(svc, phones, **kw):
    rows = [(i + 2, f'Водитель {i}', p) for i, p in enumerate(phones)]
    options = dict(name='Акция', wav_bytes=wav_bytes(), parts=PARTS, recipients={'rows': rows, 'invalid': 0,
                                                                                    'duplicates': 0})
    options.update(kw)
    return svc.create_campaign(DEPT, BOSS, **options)


def phones(n, start=0):
    return [f'7701{start + i:07d}' for i in range(n)]


def poll(svc, lines=None, ami_ok=True):
    return svc.agent_poll({'agent_id': 'gw-1', 'version': 't', 'ami_ok': ami_ok,
                           'lines': LINES_939 if lines is None else lines}, 'kid')


def answer(svc, call, listened_ms, played=False):
    return svc.agent_result({'call_id': call['call_id'], 'result': rules.RESULT_ANSWERED,
                             'dialed_at': NOON.timestamp(), 'answered_at': NOON.timestamp() + 7,
                             'ended_at': NOON.timestamp() + 7 + listened_ms / 1000, 'ring_ms': 7000,
                             'listened_ms': listened_ms, 'played_to_end': played,
                             'hangup_by': 'robot' if played else 'driver'})


def fail(svc, call, result):
    return svc.agent_result({'call_id': call['call_id'], 'result': result, 'dialed_at': NOON.timestamp()})


def test_draft_does_not_call_and_start_issues_up_to_line_capacity(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(5))
    assert campaign['status'] == 'draft' and campaign['targets_total'] == 5
    assert poll(svc)['calls'] == []                      # черновик не звонит
    svc.campaign_action(DEPT, campaign['id'], 'start')
    first = poll(svc)
    assert len(first['calls']) == 2                       # линия держит два звонка
    assert {c['line'] for c in first['calls']} == {'939'}
    assert all(c['audio_sha256'] and c['audio_ms'] == 25520 for c in first['calls'])
    # Диспетчер сообщил, что обе заняты — новых не даём.
    assert poll(svc, {'939': {'capacity': 2, 'active': 2}})['calls'] == []


NINE_LINES = {str(n): {'capacity': 2, 'active': 0} for n in (902, 907, 919, 923, 925, 928, 939, 942, 944)}


def test_at_most_twenty_calls_at_once_no_minute_cap(env):
    """Владелец, 16.10.2026: «сколько тянут линии, без предела» — общего потолка в минуту нет; но после того
    как залп 46 в минуту уронил исходящие всему кабинету — предел звонков разом («по 3 максимум», вечером
    поднят до 10, 15, 20). Двенадцать линий держат 24 — предел 20 ниже."""
    svc, clock, _ = env
    twelve = {str(n): {'capacity': 2, 'active': 0} for n in range(902, 914)}
    a = make_campaign(svc, phones(40), max_per_minute=600)
    b = make_campaign(svc, phones(8, start=100), max_per_minute=600)
    svc.campaign_action(DEPT, a['id'], 'start')
    svc.campaign_action(DEPT, b['id'], 'start')
    clock.tick(minutes=25)                                # прогрев кончился — полный темп
    issued = poll(svc, twelve)['calls']
    assert len(issued) == 20                              # двадцать мест, а не все 24 места двенадцати линий
    assert len({call['line'] for call in issued}) == 12   # раскладываем по всем линиям
    active = {}
    for call in issued:
        active[call['line']] = active.get(call['line'], 0) + 1
    busy = {line: {'capacity': 2, 'active': active.get(line, 0)} for line in twelve}
    assert poll(svc, busy)['calls'] == []                 # двадцать идут — новых не даём, хоть места у линий есть
    one = dict(busy, **{issued[0]['line']: {'capacity': 2, 'active': active[issued[0]['line']] - 1}})
    assert len(poll(svc, one)['calls']) == 1              # один кончился — одно место
    assert len(poll(svc, twelve)['calls']) == 20          # все кончились — снова двадцать, без минутного потолка


def test_warmup_from_start_and_after_robot_pause(env):
    """Владелец, 17.10.2026: «каждое автоинформирование начинать через прогрев — сначала 3, 7, 10, 15, 20».
    Ступени по 5 минут от запуска; пауза робота из-за отказов сети (мест ≤ 1) начинает прогрев заново."""
    svc, clock, _ = env
    twelve = {str(n): {'capacity': 2, 'active': 0} for n in range(902, 914)}
    campaign = make_campaign(svc, phones(120), max_per_minute=20)
    svc.campaign_action(DEPT, campaign['id'], 'start')
    first = poll(svc, twelve)
    assert len(first['calls']) == 3 and first['dial_gap_ms'] == 20000       # ступень 3 в минуту — раз в 20 с
    assert svc.campaign_detail(DEPT, campaign['id'])['warmup_rate'] == 3
    assert poll(svc, twelve)['calls'] == []                                   # в эту минуту больше не даём
    clock.tick(minutes=5)
    assert len(poll(svc, twelve)['calls']) == 7                               # ступень 7
    clock.tick(minutes=5)
    assert len(poll(svc, twelve)['calls']) == 10
    clock.tick(minutes=5)
    assert len(poll(svc, twelve)['calls']) == 15
    clock.tick(minutes=5)
    assert len(poll(svc, twelve)['calls']) == 20                              # ступень 20 = темп рассылки
    assert svc.campaign_detail(DEPT, campaign['id'])['warmup_rate'] is None   # прогрев дошёл до темпа
    # Робот ушёл в паузу (мест нет), потом проверяет связь одним звонком — прогрев с начала.
    clock.tick(minutes=5)
    paused = {line: {'capacity': 0, 'active': 0} for line in twelve}
    assert poll(svc, paused)['calls'] == []
    clock.tick(minutes=5)
    probe = dict(paused, **{'902': {'capacity': 1, 'active': 0}})
    assert len(poll(svc, probe)['calls']) == 1                                # пробный звонок
    clock.tick(seconds=30)
    assert len(poll(svc, twelve)['calls']) == 2                               # связь вернулась: 3 в минуту, 1 уже был
    assert svc.campaign_detail(DEPT, campaign['id'])['warmup_rate'] == 3
    # Простой дольше 30 минут (ночь между окнами, долгая пауза) — прогрев заново с 3 в минуту.
    clock.tick(minutes=25)
    assert len(poll(svc, twelve)['calls']) == 20                              # прогрев после паузы робота прошёл
    clock.tick(minutes=31)
    assert len(poll(svc, twelve)['calls']) == 3
    assert svc.campaign_detail(DEPT, campaign['id'])['warmup_rate'] == 3
    clock.tick(minutes=5)
    assert len(poll(svc, twelve)['calls']) == 7
    # Рассылка с темпом ниже ступени — прогрев её не трогает.
    small = make_campaign(svc, phones(5, start=500), max_per_minute=2)
    svc.campaign_action(DEPT, campaign['id'], 'pause')
    svc.campaign_action(DEPT, small['id'], 'start')
    clock.tick(minutes=1)
    assert len(poll(svc, twelve)['calls']) == 2
    assert svc.campaign_detail(DEPT, small['id'])['warmup_rate'] is None


def test_default_rate_is_what_robot_fits(env):
    svc, _, _ = env
    assert make_campaign(svc, phones(1))['max_per_minute'] == 3              # линий робота не видно — как одна
    poll(svc, NINE_LINES)                                                     # робот сообщил девять линий
    assert make_campaign(svc, phones(1))['max_per_minute'] == rules.line_rate(18, 25520) == 27  # 18 мест, предел 20
    assert make_campaign(svc, phones(1), max_per_minute=5)['max_per_minute'] == 5


def test_campaign_rate_and_window(env):
    svc, clock, _ = env
    clock.now = datetime(2026, 10, 15, 5, 30, tzinfo=timezone.utc)      # 10:30 по Алматы
    campaign = make_campaign(svc, phones(10), max_per_minute=3, window_start='10:00', window_end='11:00')
    svc.campaign_action(DEPT, campaign['id'], 'start')
    big = {'939': {'capacity': 4, 'active': 0}, '940': {'capacity': 4, 'active': 0}}
    answer = poll(svc, big)
    issued = answer['calls']
    assert len(issued) == 3                               # темп рассылки — 3 в минуту
    assert {c['line'] for c in issued} == {'939', '940'}  # раскладываем по линиям
    assert answer['dial_gap_ms'] == 20000                 # и ровно по минуте: раз в 20 с, а не пачкой
    assert poll(svc, big)['dial_gap_ms'] == 0             # нечего набирать — и шага нет
    clock.tick(hours=1)                                   # 11:30 — окно 10–11 кончилось
    assert poll(svc, big)['calls'] == []


def test_answered_and_stages(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(3))
    svc.campaign_action(DEPT, campaign['id'], 'start')
    calls = poll(svc, {'939': {'capacity': 4, 'active': 0}})['calls']
    assert len(calls) == 3
    answer(svc, calls[0], 10000)                          # положил на казахской (до 13,05 с)
    answer(svc, calls[1], 20000)                          # на русской
    answer(svc, calls[2], 25300, played=True)             # дослушал
    detail = svc.campaign_detail(DEPT, campaign['id'])
    assert detail['status'] == 'done'                     # очередь кончилась — рассылка завершена
    assert detail['analytics']['answered'] == 3
    assert detail['analytics']['stages'] == {'kz': 1, 'ru': 1, 'full': 1}
    stages = sorted(t['stage'] for t in detail['targets'])
    assert stages == ['full', 'kz', 'ru']
    assert all('•••' in t['phone'] for t in detail['targets'])         # номер — маской
    # Повтор того же итога (сеть, перезапуск) ничего не меняет.
    assert answer(svc, calls[0], 99999)['applied'] is False


def test_line_busy_is_not_an_attempt(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(1), max_attempts=1)
    svc.campaign_action(DEPT, campaign['id'], 'start')
    call = poll(svc)['calls'][0]
    fail(svc, call, rules.RESULT_LINE_BUSY)
    detail = svc.campaign_detail(DEPT, campaign['id'])
    target = detail['targets'][0]
    assert target['state'] == 'pending' and target['attempts'] == 0     # попытка не сгорела
    assert poll(svc)['calls'] == []                       # повтор — через полминуты
    clock.tick(seconds=31)
    again = poll(svc)['calls']
    assert len(again) == 1 and again[0]['call_id'] != call['call_id']


def test_no_link_is_not_an_attempt_and_kept_apart(env):
    """16.10.2026: вызов не дошёл до Binotel — водителю не звонили; попытка цела, повтор через 5 минут,
    в долях итогов «нет связи» не участвует (иначе «не ответил» врёт про водителей)."""
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(1), max_attempts=1)
    svc.campaign_action(DEPT, campaign['id'], 'start')
    call = poll(svc)['calls'][0]
    fail(svc, call, rules.RESULT_NO_LINK)
    detail = svc.campaign_detail(DEPT, campaign['id'])
    assert detail['targets'][0]['state'] == 'pending' and detail['targets'][0]['attempts'] == 0
    calls = detail['analytics']['calls']
    assert calls['no_link'] == 1 and calls['results'] == {}
    clock.tick(seconds=299)
    assert poll(svc)['calls'] == []                       # повтор не раньше пяти минут
    clock.tick(seconds=2)
    assert len(poll(svc)['calls']) == 1


def test_retry_then_give_up(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(1), max_attempts=2, retry_after_min=30)
    svc.campaign_action(DEPT, campaign['id'], 'start')
    fail(svc, poll(svc)['calls'][0], rules.RESULT_NO_ANSWER)
    assert poll(svc)['calls'] == []
    clock.tick(minutes=31)
    second = poll(svc)['calls']
    assert len(second) == 1
    fail(svc, second[0], rules.RESULT_BUSY)
    detail = svc.campaign_detail(DEPT, campaign['id'])
    assert detail['targets'][0]['state'] == 'given_up'
    assert detail['status'] == 'done'
    assert detail['analytics']['calls']['results'] == {'no_answer': 1, 'busy': 1}


def test_stale_call_is_lost_then_late_result_wins(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(1), max_attempts=3, retry_after_min=5)
    svc.campaign_action(DEPT, campaign['id'], 'start')
    call = poll(svc)['calls'][0]
    clock.tick(seconds=rules.STALE_CALL_SECONDS + 1)
    poll(svc, {'939': {'capacity': 2, 'active': 2}})
    # Запись 25,5 с ещё может играть: «потерян» — только после «запись + запас».
    assert svc.campaign_detail(DEPT, campaign['id'])['targets'][0]['state'] == 'calling'
    clock.tick(seconds=30)
    poll(svc, {'939': {'capacity': 2, 'active': 2}})      # опрос подбирает потерянные
    detail = svc.campaign_detail(DEPT, campaign['id'])
    assert detail['targets'][0]['last_result'] == 'lost' and detail['targets'][0]['state'] == 'pending'
    # Настоящий итог пришёл позже — он точнее догадки.
    assert answer(svc, call, 25000, played=True)['applied'] is True
    assert svc.campaign_detail(DEPT, campaign['id'])['targets'][0]['state'] == 'done'


def test_finish_stops_queue_and_ami_down_gives_nothing(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(4))
    svc.campaign_action(DEPT, campaign['id'], 'start')
    assert poll(svc, ami_ok=False)['calls'] == []         # робот не на связи с Asterisk — не звоним
    in_flight = poll(svc)['calls']
    assert len(in_flight) == 2
    svc.campaign_action(DEPT, campaign['id'], 'finish')
    assert poll(svc, {'939': {'capacity': 4, 'active': 0}})['calls'] == []
    states = svc.campaign_detail(DEPT, campaign['id'])['analytics']['targets']
    assert states == {'stopped': 2, 'calling': 2}         # кому не звонили — «завершили раньше»
    fail(svc, in_flight[0], rules.RESULT_NO_ANSWER)       # звонки в полёте — итог всё равно записан
    answer(svc, in_flight[1], 4000)
    states = svc.campaign_detail(DEPT, campaign['id'])['analytics']['targets']
    assert states == {'stopped': 3, 'done': 1}            # недозвон на завершённой — не в очередь
    with pytest.raises(InformerError):
        svc.campaign_action(DEPT, campaign['id'], 'start')   # завершённую — новой рассылкой


def test_validation_and_department_scope(env):
    svc, clock, _ = env
    with pytest.raises(InformerError):
        make_campaign(svc, phones(1), parts=list(reversed(PARTS)))      # порядок: казахская первой
    for wrong in (0, rules.MAX_PER_MINUTE + 1, 'много'):
        with pytest.raises(InformerError):
            make_campaign(svc, phones(1), max_per_minute=wrong)         # верх поля — проверка ввода
    assert make_campaign(svc, phones(1), max_per_minute=40)['max_per_minute'] == 40   # выше 10 — можно
    with pytest.raises(InformerError):
        make_campaign(svc, phones(1), wav_bytes=b'RIFF0000WAVEjunk')
    campaign = make_campaign(svc, phones(1))
    with pytest.raises(InformerError):
        svc.campaign_detail(OTHER_DEPT, campaign['id'])                 # чужой отдел рассылку не видит


def test_export_workbook(env):
    svc, clock, _ = env
    campaign = make_campaign(svc, phones(2))
    svc.campaign_action(DEPT, campaign['id'], 'start')
    calls = poll(svc)['calls']
    answer(svc, calls[0], 5000)
    fail(svc, calls[1], rules.RESULT_UNAVAILABLE)
    detail, targets, call_rows = svc.export_rows(DEPT, campaign['id'])
    wb = load_workbook(io.BytesIO(informer_export.build(detail, targets, call_rows).getvalue()))
    assert wb.sheetnames == ['Водители', 'Звонки', 'Контекст']
    drivers = list(wb['Водители'].iter_rows(min_row=2, values_only=True))
    assert len(drivers) == 2 and all('•••' in row[2] for row in drivers)
    assert all('7701' not in str(cell) for row in drivers for cell in row)        # полного номера нет
    call_sheet = list(wb['Звонки'].iter_rows(min_row=2, values_only=True))
    assert {row[5] for row in call_sheet} == {'Дозвонились', 'Недоступен'}


def test_signed_after_robot_call(env):
    """Владелец, 17.10.2026: «отображать, сколько человек подписали именно после звонка автоинформатора».
    Подписи — из базы «Обзвона» того же отдела по номеру: позже первого ответа роботу — «после звонка»,
    без успешки оператора — «сам»; подпись раньше звонка — «подписал до звонка»."""
    svc, clock, db = env
    campaign = make_campaign(svc, phones(4), max_per_minute=10)
    svc.campaign_action(DEPT, campaign['id'], 'start')
    clock.tick(minutes=25)                                         # прогрев позади — все четыре разом
    calls = poll(svc, {'939': {'capacity': 4, 'active': 0}})['calls']
    assert len(calls) == 4
    for call in calls[:3]:
        answer(svc, call, 20000)                                   # ответили в NOON+7 с
    fail(svc, calls[3], rules.RESULT_NO_ANSWER)
    by_call = {c['call_id']: c['phone'] for c in calls}
    signed_after, signed_self, signed_before, not_called = (by_call[calls[0]['call_id']], by_call[calls[1]['call_id']],
                                                            by_call[calls[2]['call_id']], by_call[calls[3]['call_id']])
    with db._get_cursor() as cur:
        cur.execute("""INSERT INTO dial_list_leads (department_id, phone_norm, signed_at, success_attempt_id, sign_checked_at)
                       VALUES (1954, %s, %s, 77, %s), (1954, %s, %s, NULL, %s), (1954, %s, %s, NULL, %s),
                              (1954, %s, %s, NULL, %s), (560, %s, %s, NULL, %s)""",
                    (signed_after, NOON + timedelta(hours=1), NOON + timedelta(hours=3),      # после звонка, успешка оператора
                     signed_self, NOON + timedelta(minutes=30), NOON + timedelta(hours=3),    # после звонка, сам
                     signed_before, NOON - timedelta(days=1), NOON + timedelta(hours=3),      # подписал до звонка
                     not_called, NOON + timedelta(hours=2), NOON + timedelta(hours=3),        # подписал, робот не дозвонился
                     signed_self, NOON + timedelta(hours=1), NOON + timedelta(hours=3)))      # другой отдел — не наш
    detail = svc.campaign_detail(DEPT, campaign['id'])
    signing = detail['analytics']['signing']
    assert datetime.fromisoformat(signing.pop('checked_at')) == NOON + timedelta(hours=3)
    assert signing == {'matched': 4, 'after_call': 2, 'after_call_self': 1, 'before_call': 1}
    rows = {t['phone']: t for t in detail['targets']}
    marks = {phone: (t['signed_after_call'], bool(t['signed_at'])) for phone, t in rows.items()}
    assert sorted(marks.values()) == [(False, True), (False, True), (True, True), (True, True)]
    _, targets, _ = svc.export_rows(DEPT, campaign['id'])
    assert sorted(t['signed_after_call'] for t in targets) == [False, False, True, True]
    wb = load_workbook(io.BytesIO(informer_export.build(detail, targets, []).getvalue()))
    drivers = list(wb['Водители'].iter_rows(min_row=1, values_only=True))
    assert drivers[0][8:10] == ('Подписал (Sapar)', 'Подпись после звонка')
    assert sorted(row[9] for row in drivers[1:]) == ['да', 'да', 'подписал до звонка', 'подписал до звонка']
    context = {row[0]: row[1] for row in wb['Контекст'].iter_rows(min_row=2, values_only=True)}
    assert context['Подписали после звонка робота'] == '2' and context['…из них без разговора с оператором'] == '1'
    # Базы «Обзвона» нет (другой стенд) — сводки нет, остальное работает.
    with db._get_cursor() as cur:
        cur.execute("DROP TABLE dial_list_leads")
    detail = svc.campaign_detail(DEPT, campaign['id'])
    assert detail['analytics']['signing'] is None and detail['analytics']['answered'] == 3
    assert all(t['signed_at'] is None for t in detail['targets'])


def test_robot_status(env):
    svc, clock, _ = env
    assert svc.robot_status()['online'] is False
    poll(svc)
    status = svc.robot_status()
    assert status['online'] is True and status['lines'] == LINES_939
    clock.tick(seconds=120)
    assert svc.robot_status()['online'] is False


def test_master_for_listening_phone_version_for_robot(env):
    """Владелец, 15.10.2026: голос в полном качестве. Руководитель слушает оригинал 24 кГц,
    робот получает телефонную версию 8 кГц — её сводит сервер, а не браузер."""
    svc, _, _ = env
    original = master_wav_bytes()
    campaign = make_campaign(svc, phones(1), wav_bytes=original)
    assert campaign['audio_ms'] == 25520
    sha, robot_audio = svc.agent_audio(campaign['id'])
    assert informer_wav.inspect(robot_audio)['sample_rate'] == informer_wav.PHONE_RATE
    assert robot_audio == informer_telephone.from_master(original)
    assert sha == hashlib.sha256(robot_audio).hexdigest()
    assert svc.campaign_audio(DEPT, campaign['id']) == original
    svc.campaign_action(DEPT, campaign['id'], 'start')
    call = poll(svc)['calls'][0]
    assert call['audio_sha256'] == sha and call['audio_ms'] == 25520


def test_phone_only_recording_of_previous_page_still_works(env):
    svc, _, _ = env
    old = wav_bytes()
    campaign = make_campaign(svc, phones(1), wav_bytes=old)
    assert svc.agent_audio(campaign['id']) == (hashlib.sha256(old).hexdigest(), old)
    assert svc.campaign_audio(DEPT, campaign['id']) == old
    with pytest.raises(InformerError, match='24 кГц или 8 кГц'):
        make_campaign(svc, phones(1), wav_bytes=informer_wav.encode(b'\x00\x00' * 48000, 48000))


def test_existing_table_gets_master_column():
    """Боевая таблица создана до оригиналов: повторный apply добавляет колонку, рассылки целы."""
    schema = 't_informer_' + uuid.uuid4().hex
    db = _Db(schema)
    admin = db.connect()
    admin.autocommit = True
    try:
        with admin.cursor() as cur:
            cur.execute('CREATE SCHEMA ' + schema)
            cur.execute('SET search_path TO ' + schema)
            cur.execute(BASE_DDL)
            lines = informer_schema.DDL[0].splitlines()
            old_ddl = '\n'.join(line for line in lines if 'audio_master' not in line and 'оригинал (24' not in line)
            cur.execute(old_ddl)
            cur.execute("SELECT count(*) FROM information_schema.columns WHERE table_schema = %s"
                        " AND table_name = 'informer_campaigns' AND column_name = 'audio_master'", (schema,))
            assert cur.fetchone()[0] == 0
            cur.execute("INSERT INTO informer_campaigns (department_id, name, audio_wav, audio_sha256, audio_ms)"
                        " VALUES (1954, 'Старая', '\\x00', 'sha', 1000)")
            informer_schema.apply(cur)
            cur.execute("SELECT name, audio_master FROM informer_campaigns")
            assert cur.fetchall() == [('Старая', None)]
    finally:
        with admin.cursor() as cur:
            cur.execute('DROP SCHEMA ' + schema + ' CASCADE')
        admin.close()
