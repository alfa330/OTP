# -*- coding: utf-8 -*-
"""Раздел «Термокороба» (#363): права, проверка чисел, ручки и проводка.

Что сторожится:
  * периметр: фронт-офисы правят, СЗоВ смотрит, памятку правит СВ, тренер не входит;
  * «только целые числа, без отрицательных значений» — и у сервера, и у фронта;
  * сохранение пачки ложится целиком или не ложится вовсе: ни одной записи
    до того, как проверено всё (db._get_cursor коммитит и то, что записано
    до ответа 4xx);
  * правка коллеги не затирается молча — 409 по версии строки;
  * пустое сохранение (те же числа) — не изменение: ни истории, ни «Обновлено»;
  * подписи экрана, файла и истории — одни и те же строки на фронте и сервере;
  * раздел подключён: схема, Blueprint, пункт меню, гард видимости.

SQL здесь не проверяется — его проверяет прогон на стенде с базой.
"""

import re
import sys
import unittest
from contextlib import contextmanager
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    from flask import Flask
except ImportError:  # pragma: no cover
    Flask = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

from thermoboxes import access, queries, report, routes, rules, schema  # noqa: E402

APP_JSX = ROOT / 'src' / 'App.jsx'
META_JS = ROOT / 'src' / 'components' / 'thermoboxes' / 'thermoboxMeta.js'
DATABASE_PY = ROOT / 'database.py'
BOT_PY = ROOT / 'bot_schedule2.py'


def _read(path):
    return path.read_text(encoding='utf-8-sig')


def person(role='operator', department_code='front_office', headed_codes=(), user_id=10):
    return {
        'user_id': user_id, 'name': 'Сотрудник %d' % user_id, 'role': role,
        'department_id': 909, 'department_code': department_code, 'city': 'Алматы',
        'headed_department_ids': [909] if headed_codes else [],
        'headed_department_codes': list(headed_codes),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Права
# ─────────────────────────────────────────────────────────────────────────────

class PilotTests(unittest.TestCase):
    """Пилот «только супер-админ» (02.10.2026): сервер и меню закрыты вместе."""

    def test_only_super_admin_while_pilot(self):
        self.assertTrue(access.PILOT_SUPER_ADMIN_ONLY)
        self.assertTrue(access.can_open_section(person(role='super_admin', department_code=None)))
        for kwargs in ({}, {'department_code': 'szov'}, {'role': 'admin', 'department_code': None},
                       {'role': 'admin', 'department_code': None, 'headed_codes': ('front_office',)}):
            ctx = person(**kwargs)
            self.assertFalse(access.can_open_section(ctx), kwargs)
            self.assertFalse(any(access.capabilities(ctx).values()), kwargs)

    def test_pilot_flag_is_the_same_on_both_sides(self):
        app = _read(APP_JSX)
        self.assertIn('const THERMOBOXES_PILOT_SUPER_ADMIN_ONLY = %s;' % str(access.PILOT_SUPER_ADMIN_ONLY).lower(), app)
        predicate = app.split('const canAccessThermoboxesSectionForUser = (userLike) => {')[1].split('\n};')[0]
        self.assertLess(predicate.index("if (role === 'super_admin') return true;"),
                        predicate.index('if (THERMOBOXES_PILOT_SUPER_ADMIN_ONLY) return false;'))


@mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False)
class AccessTests(unittest.TestCase):
    """Периметр после пилота."""

    def caps(self, **kwargs):
        return access.capabilities(person(**kwargs))

    def test_front_office_edits_data_but_not_memo_or_composition(self):
        self.assertEqual(self.caps(), {
            'can_open': True, 'can_edit': True, 'can_manage': False, 'can_edit_memo': False})

    def test_call_center_only_looks(self):
        self.assertEqual(self.caps(department_code='szov'), {
            'can_open': True, 'can_edit': False, 'can_manage': False, 'can_edit_memo': False})

    def test_supervisor_edits_memo_but_not_numbers(self):
        caps = self.caps(role='sv', department_code='szov')
        self.assertTrue(caps['can_edit_memo'])
        self.assertFalse(caps['can_edit'])
        self.assertTrue(self.caps(role='supervisor', department_code='szov')['can_edit_memo'])

    def test_supervisor_of_another_department_is_not_let_in(self):
        caps = self.caps(role='sv', department_code='tez')
        self.assertFalse(caps['can_open'])
        self.assertFalse(caps['can_edit_memo'])

    def test_trainer_is_not_let_in(self):
        self.assertFalse(self.caps(role='trainer', department_code='szov')['can_open'])

    def test_other_departments_are_not_let_in(self):
        for code in ('op', 'tez', None):
            self.assertFalse(self.caps(department_code=code)['can_open'], code)

    def test_global_admin_does_everything(self):
        for role in ('super_admin', 'admin'):
            caps = self.caps(role=role, department_code=None)
            self.assertTrue(all(caps.values()), role)

    def test_front_office_head_manages_but_does_not_write_memo(self):
        caps = self.caps(role='admin', department_code=None, headed_codes=('front_office',))
        self.assertTrue(caps['can_edit'])
        self.assertTrue(caps['can_manage'])
        self.assertFalse(caps['can_edit_memo'])

    def test_call_center_head_only_looks(self):
        caps = self.caps(role='admin', department_code=None, headed_codes=('szov',))
        self.assertTrue(caps['can_open'])
        self.assertFalse(caps['can_edit'])
        self.assertFalse(caps['can_manage'])


# ─────────────────────────────────────────────────────────────────────────────
# Числа и подписи
# ─────────────────────────────────────────────────────────────────────────────

class RulesTests(unittest.TestCase):

    def test_whole_numbers_only(self):
        self.assertEqual(rules.clean_value('free_boxes', 11), 11)
        self.assertEqual(rules.clean_value('free_boxes', 0), 0)
        self.assertEqual(rules.clean_value('free_boxes', 6.0), 6)
        for bad in (-1, 1.5, '5', None, True, [], 100001):
            with self.assertRaises(rules.FieldError, msg=repr(bad)):
                rules.clean_value('free_boxes', bad)

    def test_negative_says_so(self):
        with self.assertRaises(rules.FieldError) as caught:
            rules.clean_value('thermo_bags', -3)
        self.assertIn('меньше нуля', caught.exception.message)
        self.assertEqual(caught.exception.field, 'thermo_bags')

    def test_period_is_at_least_a_day(self):
        with self.assertRaises(rules.FieldError):
            rules.clean_value('period_days', 0)
        self.assertEqual(rules.clean_value('period_days', 7), 7)

    def test_tariff_is_from_the_list(self):
        self.assertEqual(rules.clean_value('tariff', 'all'), 'all')
        with self.assertRaises(rules.FieldError):
            rules.clean_value('tariff', 'Все тарифы')

    def test_special_condition_is_trimmed_and_empty_is_none(self):
        self.assertEqual(rules.clean_value('special_condition', '  НЕ   новички '), 'НЕ новички')
        self.assertIsNone(rules.clean_value('special_condition', '   '))
        with self.assertRaises(rules.FieldError):
            rules.clean_value('special_condition', 'я' * 201)

    def test_only_sent_fields_are_cleaned(self):
        self.assertEqual(rules.clean_fields({'free_boxes': 3, 'id': 1, 'version': 2}), {'free_boxes': 3})

    def test_same_numbers_are_not_a_change(self):
        row = {'free_boxes': 11, 'thermo_bags': 0, 'tariff': 'all'}
        self.assertEqual(rules.changes(row, {'free_boxes': 11, 'tariff': 'all'}), [])
        self.assertEqual(rules.changes(row, {'free_boxes': 9}),
                         [{'field': 'free_boxes', 'from': 11, 'to': 9}])

    def test_labels_read_like_the_sheet(self):
        self.assertEqual(rules.orders_label(15), '15+ заказов')
        self.assertEqual(rules.orders_label(21), '21+ заказов')
        self.assertEqual(rules.orders_label(0), 'Без нормы')
        self.assertEqual(rules.period_label(7), 'Неделя (7 дней)')
        self.assertEqual(rules.period_label(3), '3 дня')
        self.assertEqual(rules.deposit_label(5000), '5 000 тенге')
        self.assertEqual(rules.deposit_label(0), 'Нет депозита')
        self.assertEqual(rules.tariff_label('auto_couriers'), 'Выдаются авто-курьерам')

    def test_memo_drops_empty_items_and_keeps_line_breaks(self):
        memo = rules.clean_memo({'title': ' Правила ', 'items': [
            {'text': 'Донгелек не обслуживается.\nНЕ отправлять курьеров', 'owner': 'kc'},
            {'text': '   ', 'owner': 'regions'},
            {'text': 'Сверять с таблицей', 'owner': None},
        ]})
        self.assertEqual(memo['title'], 'Правила')
        self.assertEqual(memo['items'], [
            {'text': 'Донгелек не обслуживается.\nНЕ отправлять курьеров', 'owner': 'kc', 'kind': None},
            {'text': 'Сверять с таблицей', 'owner': None, 'kind': None},
        ])

    def test_memo_kind_is_from_the_list(self):
        memo = rules.clean_memo({'items': [{'text': 'Нельзя', 'owner': 'kc', 'kind': 'forbidden'}]})
        self.assertEqual(memo['items'][0]['kind'], 'forbidden')
        with self.assertRaises(rules.FieldError):
            rules.clean_memo({'items': [{'text': 'x', 'kind': 'skull'}]})

    def test_memo_rejects_unknown_owner_and_too_many_items(self):
        with self.assertRaises(rules.FieldError):
            rules.clean_memo({'items': [{'text': 'x', 'owner': 'boss'}]})
        with self.assertRaises(rules.FieldError):
            rules.clean_memo({'items': [{'text': str(n)} for n in range(rules.MEMO_ITEMS_MAX + 1)]})
        with self.assertRaises(rules.FieldError):
            rules.clean_memo({'items': 'текст'})

    def test_filter_is_by_city_and_every_word(self):
        row = {'city': 'Алматы', 'name': 'Офис', 'address': '7-й микрорайон, 5', 'special_condition': 'НЕ новички'}
        self.assertTrue(rules.matches(row, 'алматы', ''))
        self.assertFalse(rules.matches(row, 'Астана', ''))
        self.assertTrue(rules.matches(row, '', 'микрорайон новички'))
        self.assertFalse(rules.matches(row, '', 'микрорайон депозит'))
        self.assertTrue(rules.matches({'city': 'Семей', 'address': 'Каюма Мухамедханова'}, '', 'мухамедханова'))


class FrontMirrorTests(unittest.TestCase):
    """Подписи экрана и сервера — одни и те же строки."""

    def setUp(self):
        self.meta = _read(META_JS)

    def _object(self, name):
        block = self.meta.split('export const %s = {' % name)[1].split('};')[0]
        return dict(re.findall(r"(\w+):\s*'([^']*)'", block))

    def test_field_labels(self):
        self.assertEqual(self._object('FIELD_LABELS'), rules.FIELD_LABELS)

    def test_tariff_labels(self):
        self.assertEqual(self._object('TARIFF_LABELS'), rules.TARIFF_LABELS)
        self.assertEqual(set(self._object('TARIFF_SHORT')), set(rules.TARIFF_LABELS))

    def test_memo_owner_labels(self):
        self.assertEqual(self._object('MEMO_OWNER_LABELS'), rules.MEMO_OWNER_LABELS)

    def test_memo_kind_labels(self):
        self.assertEqual(self._object('MEMO_KIND_LABELS'), rules.MEMO_KIND_LABELS)

    def test_limits_match(self):
        self.assertIn('export const SPECIAL_CONDITION_MAX = %d;' % rules.SPECIAL_CONDITION_MAX, self.meta)
        self.assertIn('export const MEMO_TEXT_MAX = %d;' % rules.MEMO_TEXT_MAX, self.meta)
        self.assertIn('export const MEMO_ITEMS_MAX = %d;' % rules.MEMO_ITEMS_MAX, self.meta)
        block = self.meta.split('const LIMITS = {')[1].split('};')[0]
        limits = {name: (int(low), int(high))
                  for name, low, high in re.findall(r'(\w+):\s*\[(\d+),\s*(\d+)\]', block)}
        self.assertEqual(limits, rules._LIMITS)

    def test_export_file_name(self):
        self.assertEqual(report.file_name(date(2026, 10, 1)), 'Термокороба 01.10.2026.xlsx')
        self.assertIn('`Термокороба ${pad(day)}.${pad(month)}.${year}.xlsx`', self.meta)


# ─────────────────────────────────────────────────────────────────────────────
# Схема
# ─────────────────────────────────────────────────────────────────────────────

class _Cursor:
    def __init__(self):
        self.statements = []

    def execute(self, sql, params=None):
        self.statements.append(' '.join(sql.split()))


class SchemaTests(unittest.TestCase):

    def test_counts_cannot_go_negative_in_the_database(self):
        ddl = ' '.join(' '.join(statement.split()) for statement in schema._STATEMENTS)
        for column in ('free_boxes', 'thermo_bags', 'used_boxes', 'min_orders', 'deposit_tenge'):
            self.assertIn('CHECK (%s >= 0)' % column, ddl)
        self.assertIn('office_id INTEGER UNIQUE REFERENCES wiki_offices(id)', ddl)

    def test_tables_go_first_and_indexes_last(self):
        cursor = _Cursor()
        schema.init_thermoboxes_schema(cursor)
        kinds = ['table' if 'CREATE TABLE' in sql else 'other' for sql in cursor.statements]
        self.assertEqual(kinds[:3], ['table'] * 3)
        self.assertNotIn('table', kinds[3:])
        self.assertTrue(all('IF NOT EXISTS' in sql or 'ON CONFLICT' in sql for sql in cursor.statements))


# ─────────────────────────────────────────────────────────────────────────────
# Ручки — на подменённом слое запросов
# ─────────────────────────────────────────────────────────────────────────────

def _row(row_id, city, **fields):
    row = {
        'id': row_id, 'office_id': 40 + row_id, 'city': city, 'name': 'Офис %s' % city,
        'address': 'ул. %s, %d' % (city, row_id),
        'free_boxes': 10, 'thermo_bags': 0, 'used_boxes': 0,
        'tariff': 'auto_couriers', 'min_orders': 15, 'period_days': 7,
        'deposit_tenge': 5000, 'special_condition': None,
        'is_active': True, 'version': 1, 'updated_by_name': None,
        'updated_at': '2026-10-01T10:00:00', 'created_at': '2026-10-01T10:00:00',
    }
    row.update(fields)
    return row


class Store:
    """Память вместо базы: ровно те функции queries, что зовут роуты."""

    def __init__(self):
        self.rows = {1: _row(1, 'Алматы'), 2: _row(2, 'Астана'), 3: _row(3, 'Атырау', is_active=False)}
        self.events = []
        self.updates = []
        self.memo = {'title': None, 'items': [], 'updated_by_name': None, 'updated_at': None}
        self.directory_rows = [
            {'id': 41, 'city': 'Алматы', 'name': 'Офис Алматы', 'address': 'ул. Алматы, 1'},
            {'id': 50, 'city': 'Актобе', 'name': 'Офис Актобе', 'address': '11-й мкр, 3д'},
        ]

    def section_spaces(self, cursor):
        return [11]

    def list_rows(self, cursor, *, spaces, include_hidden=False):
        return [dict(row) for row in self.rows.values() if include_hidden or row['is_active']]

    def read_row(self, cursor, row_id, *, spaces):
        row = self.rows.get(int(row_id))
        return dict(row) if row else None

    def lock_rows(self, cursor, row_ids, *, spaces):
        return {int(x): dict(self.rows[int(x)]) for x in row_ids if int(x) in self.rows}

    def update_row(self, cursor, row_id, fields, actor):
        self.updates.append((row_id, dict(fields)))
        self.rows[row_id].update(fields)
        self.rows[row_id]['version'] += 1
        self.rows[row_id]['updated_by_name'] = actor['name']

    def set_active(self, cursor, row_id, active, actor):
        self.rows[row_id]['is_active'] = active
        self.rows[row_id]['version'] += 1

    def create_row(self, cursor, office, fields, actor):
        row_id = max(self.rows) + 1
        values = dict(rules.NEW_ROW_DEFAULTS)
        values.update(fields)
        self.rows[row_id] = _row(row_id, office['city'], office_id=office['id'], **values)
        return row_id

    def taken_office_ids(self, cursor):
        return {row['office_id'] for row in self.rows.values()}

    def directory(self, cursor, *, spaces):
        return list(self.directory_rows)

    def insert_event(self, cursor, row_id, kind, changes, actor):
        self.events.append((row_id, kind, changes, actor['user_id']))

    def list_events(self, cursor, row_id, limit=100):
        return [{'id': n, 'kind': kind, 'changes': changes, 'actor_name': 'x', 'created_at': None}
                for n, (rid, kind, changes, _) in enumerate(self.events) if rid == row_id]

    def get_memo(self, cursor):
        return dict(self.memo)

    def save_memo(self, cursor, memo, actor):
        self.memo = dict(memo, updated_by_name=actor['name'], updated_at='2026-10-01T12:00:00')
        return dict(self.memo)


class FakeDb:
    def __init__(self):
        self.commits = 0
        self.rollbacks = 0

    @contextmanager
    def _get_cursor(self):
        try:
            yield object()
        except Exception:
            self.rollbacks += 1
            raise
        self.commits += 1


@unittest.skipIf(Flask is None, 'Flask не установлен')
class _Base(unittest.TestCase):

    def setUp(self):
        self.store = Store()
        self.db = FakeDb()
        self.viewer = person()
        patches = {'load_access_context': lambda cursor, user_id: dict(self.viewer)}
        for name in ('section_spaces', 'list_rows', 'read_row', 'lock_rows', 'update_row', 'set_active',
                     'create_row', 'taken_office_ids', 'directory', 'insert_event', 'list_events',
                     'get_memo', 'save_memo'):
            patches[name] = getattr(self.store, name)
        patcher = mock.patch.multiple(queries, **patches)
        patcher.start()
        self.addCleanup(patcher.stop)
        ready = mock.patch.object(schema, 'schema_is_ready', lambda cursor: True)
        ready.start()
        self.addCleanup(ready.stop)

        def read_office(cursor, office_id, *, space_ids):
            return next((office for office in self.store.directory_rows if office['id'] == office_id), None)

        office = mock.patch.object(routes.parcels_queries, 'read_office', read_office)
        office.start()
        self.addCleanup(office.stop)
        today = mock.patch.object(routes.parcels_queries, 'today_almaty', lambda: date(2026, 10, 1))
        today.start()
        self.addCleanup(today.stop)

        app = Flask(__name__)
        app.register_blueprint(routes.build_thermoboxes_blueprint(
            db=self.db,
            require_api_key=lambda f: f,
            build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=lambda: (self.viewer['user_id'], None, None),
        ))
        app.config['TESTING'] = True
        self.client = app.test_client()
        # Ручки проверяются с периметром ПОСЛЕ пилота; сам пилот — PilotTests.
        pilot = mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False)
        pilot.start()
        self.addCleanup(pilot.stop)

    def save(self, *items):
        return self.client.put('/api/thermoboxes/rows', json={'items': list(items)})


class ScreenTests(_Base):

    def test_screen_is_one_request(self):
        data = self.client.get('/api/thermoboxes').get_json()
        self.assertTrue(data['schema_ready'])
        self.assertEqual([row['id'] for row in data['rows']], [1, 2])
        self.assertEqual(data['capabilities']['can_edit'], True)
        self.assertNotIn('directory', data)

    def test_manager_sees_hidden_rows_and_free_directory_offices(self):
        self.viewer = person(role='admin', department_code=None, headed_codes=('front_office',))
        data = self.client.get('/api/thermoboxes').get_json()
        self.assertEqual([row['id'] for row in data['rows']], [1, 2, 3])
        # Офис 41 уже в таблице (строка 1) — в «добавить» его нет.
        self.assertEqual([office['id'] for office in data['directory']], [50])

    def test_other_department_is_closed(self):
        self.viewer = person(department_code='op')
        response = self.client.get('/api/thermoboxes')
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'THERMO_SECTION_CLOSED')


class SaveTests(_Base):

    def test_batch_saves_changes_and_history(self):
        response = self.save({'id': 1, 'version': 1, 'free_boxes': 8, 'thermo_bags': 2},
                             {'id': 2, 'version': 1, 'free_boxes': 10})
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        # Строка 2 прислана с теми же числами — не изменение.
        self.assertEqual(data['saved'], 1)
        self.assertEqual([row['id'] for row in data['rows']], [1])
        self.assertEqual(self.store.updates, [(1, {'free_boxes': 8, 'thermo_bags': 2})])
        self.assertEqual(self.store.events, [(1, 'edited', [
            {'field': 'free_boxes', 'from': 10, 'to': 8},
            {'field': 'thermo_bags', 'from': 0, 'to': 2},
        ], 10)])
        self.assertEqual(self.store.rows[1]['version'], 2)
        self.assertEqual(self.store.rows[2]['version'], 1)

    def test_call_center_cannot_save(self):
        self.viewer = person(department_code='szov', role='sv')
        response = self.save({'id': 1, 'version': 1, 'free_boxes': 8})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'THERMO_EDIT_FORBIDDEN')
        self.assertEqual(self.store.updates, [])

    def test_one_bad_number_writes_nothing(self):
        response = self.save({'id': 1, 'version': 1, 'free_boxes': 8},
                             {'id': 2, 'version': 1, 'used_boxes': -1})
        self.assertEqual(response.status_code, 400)
        data = response.get_json()
        self.assertEqual(data['code'], 'THERMO_FIELD_INVALID')
        self.assertEqual((data['row_id'], data['field']), (2, 'used_boxes'))
        self.assertIn('Астана', data['error'])
        self.assertEqual(self.store.updates, [])
        self.assertEqual(self.store.events, [])

    def test_stale_version_is_a_conflict_and_writes_nothing(self):
        self.store.rows[2]['version'] = 5
        self.store.rows[2]['free_boxes'] = 3
        response = self.save({'id': 1, 'version': 1, 'free_boxes': 8},
                             {'id': 2, 'version': 4, 'free_boxes': 9})
        self.assertEqual(response.status_code, 409)
        data = response.get_json()
        self.assertEqual(data['code'], 'THERMO_CONFLICT')
        self.assertIn('Астана', data['error'])
        self.assertEqual([(row['id'], row['free_boxes']) for row in data['rows']], [(2, 3)])
        self.assertEqual(self.store.updates, [])

    def test_unknown_and_hidden_rows_are_refused(self):
        self.assertEqual(self.save({'id': 99, 'version': 1, 'free_boxes': 1}).status_code, 404)
        response = self.save({'id': 3, 'version': 1, 'free_boxes': 1})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['code'], 'THERMO_ROW_HIDDEN')
        self.assertEqual(self.store.updates, [])

    def test_malformed_batches_are_refused(self):
        self.assertEqual(self.client.put('/api/thermoboxes/rows', json={'items': []}).status_code, 400)
        self.assertEqual(self.save({'id': 1, 'free_boxes': 1}).status_code, 400)
        duplicate = self.save({'id': 1, 'version': 1, 'free_boxes': 1}, {'id': 1, 'version': 1, 'free_boxes': 2})
        self.assertEqual(duplicate.get_json()['code'], 'THERMO_ITEM_DUPLICATE')
        self.assertEqual(self.store.updates, [])


class CompositionTests(_Base):

    def setUp(self):
        super().setUp()
        self.viewer = person(role='admin', department_code=None, headed_codes=('front_office',))

    def test_add_office_from_directory(self):
        response = self.client.post('/api/thermoboxes/rows', json={'office_id': 50, 'free_boxes': 4})
        self.assertEqual(response.status_code, 201)
        row = response.get_json()['row']
        self.assertEqual((row['city'], row['free_boxes'], row['tariff']), ('Актобе', 4, 'auto_couriers'))
        self.assertEqual(self.store.events[-1][1], 'created')

    def test_office_cannot_be_added_twice_or_from_nowhere(self):
        self.assertEqual(self.client.post('/api/thermoboxes/rows', json={'office_id': 41}).status_code, 409)
        self.assertEqual(self.client.post('/api/thermoboxes/rows', json={'office_id': 999}).status_code, 404)

    def test_regular_front_office_cannot_add(self):
        self.viewer = person()
        response = self.client.post('/api/thermoboxes/rows', json={'office_id': 50})
        self.assertEqual(response.status_code, 403)

    def test_hide_and_show_write_history_once(self):
        url = '/api/thermoboxes/rows/1/visibility'
        self.assertFalse(self.client.patch(url, json={'is_active': False}).get_json()['row']['is_active'])
        self.client.patch(url, json={'is_active': False})
        self.assertEqual([event[1] for event in self.store.events], ['hidden'])
        self.assertEqual(self.client.patch(url, json={'is_active': 'нет'}).status_code, 400)


class MemoAndHistoryTests(_Base):

    def test_operator_cannot_write_memo(self):
        response = self.client.put('/api/thermoboxes/memo', json={'items': [{'text': 'x'}]})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'THERMO_MEMO_FORBIDDEN')

    def test_supervisor_writes_memo(self):
        self.viewer = person(role='sv', department_code='szov')
        response = self.client.put('/api/thermoboxes/memo', json={
            'title': 'Правила', 'items': [{'text': 'В офисах снимать деньги нельзя', 'owner': 'kc'}]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.store.memo['items'],
                         [{'text': 'В офисах снимать деньги нельзя', 'owner': 'kc', 'kind': None}])

    def test_hidden_row_history_is_for_managers_only(self):
        self.assertEqual(self.client.get('/api/thermoboxes/rows/3/events').status_code, 404)
        self.assertEqual(self.client.get('/api/thermoboxes/rows/1/events').status_code, 200)


@unittest.skipIf(load_workbook is None, 'openpyxl не установлен')
class ExportTests(_Base):

    def test_export_follows_the_screen_filter_and_writes_numbers(self):
        self.store.rows[2].update(deposit_tenge=0, special_condition='НЕ новички')
        response = self.client.get('/api/thermoboxes/export?city=Астана')
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename*=UTF-8', response.headers.get('Content-Disposition', ''))
        sheet = load_workbook(BytesIO(response.data)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][:5], ('Город', 'Адрес', 'Бесплатные термокороба', 'Термопакет', 'Б/У короб'))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], 'Астана')
        self.assertEqual(rows[1][2], 10)
        self.assertIsInstance(rows[1][2], int)
        self.assertEqual(rows[1][8], 'НЕ новички')
        self.assertEqual(rows[1][9], 'Нет депозита')


# ─────────────────────────────────────────────────────────────────────────────
# Проводка
# ─────────────────────────────────────────────────────────────────────────────

class ImportScriptTests(unittest.TestCase):
    """Перенос листа: значок правила из эмодзи и склейка переносов ячейки."""

    def setUp(self):
        sys.path.insert(0, str(ROOT / 'scripts'))
        self.addCleanup(sys.path.remove, str(ROOT / 'scripts'))
        import import_thermoboxes
        self.script = import_thermoboxes

    def test_sheet_emoji_become_memo_kinds(self):
        kinds = [self.script.memo_kind(mark) for mark in ('⛔', '🚫', '💰', '📊', '⚠️', '⚠', '✅', '❓', '', None)]
        self.assertEqual(kinds, ['forbidden', 'forbidden', 'money', 'data', 'warning', 'warning',
                                 'check', 'question', None, None])
        self.assertTrue(set(kind for kind in kinds if kind) <= set(rules.MEMO_KIND_LABELS))

    def test_cell_line_breaks_are_joined_inside_a_sentence(self):
        self.assertEqual(self.script.memo_text('проверьте условиям\nи отправьте данные'),
                         'проверьте условиям и отправьте данные')
        self.assertEqual(self.script.memo_text('НЕ ОБСЛУЖИВАЕТСЯ.\n НЕ ОТПРАВЛЯТЬ курьеров'),
                         'НЕ ОБСЛУЖИВАЕТСЯ.\nНЕ ОТПРАВЛЯТЬ курьеров')

    def _sheet_row(self, line, city, address, free='5'):
        raw = {'free_boxes': free, 'thermo_bags': '0', 'used_boxes': '0', 'tariff': 'Выдаются авто-курьерам',
               'min_orders': '15+заказов', 'period_days': 'Неделя (7 дней)', 'special_condition': '',
               'deposit_tenge': '5 000 тенге'}
        return {'line': line, 'city': city, 'address': address, 'raw': raw}

    def test_plan_takes_manual_office_and_parses_sheet_values(self):
        offices = [{'id': 63, 'city': 'Туркестан', 'address': 'улица Нышанов 15'}]
        plan, problems = self.script.build_plan(
            [self._sheet_row(14, 'Туркестан', 'ул. Нышанов, 14, БЦ «Балнур»', free='55')], offices,
            {self.script._fold('Туркестан'): 63})
        self.assertEqual(problems, [])
        (_, office, fields, note), = plan
        self.assertEqual((office['id'], note), (63, 'указан вручную'))
        self.assertEqual(fields, {'free_boxes': 55, 'thermo_bags': 0, 'used_boxes': 0, 'tariff': 'auto_couriers',
                                  'min_orders': 15, 'period_days': 7, 'deposit_tenge': 5000,
                                  'special_condition': None})

    def test_plan_refuses_two_sheet_rows_on_one_office(self):
        offices = [{'id': 49, 'city': 'Астана', 'address': 'Проспект Сарыарка, 31'}]
        override = {self.script._fold('Астана'): 49}
        plan, problems = self.script.build_plan(
            [self._sheet_row(5, 'Астана', 'пр. Сарыарка, 31'), self._sheet_row(6, 'Астана', 'Сарыарка 31')],
            offices, override)
        self.assertEqual(len(problems), 2)
        self.assertTrue(all('двух строк' in problem for problem in problems))

    def test_offices_match_by_house_number(self):
        offices = [
            {'id': 47, 'city': 'Алматы', 'address': 'Улица Жамбыла, 172В угол улицы Байзакова'},
            {'id': 48, 'city': 'Алматы', 'address': '7-й микрорайон, 5'},
            {'id': 44, 'city': 'Астана', 'address': 'проспект Сарыарка 31'},
            {'id': 49, 'city': 'Астана', 'address': 'Проспект Сарыарка, 31 угол улицы Алиби Жангельдин'},
        ]
        office, reason, note = self.script.match_office('Алматы', '7-й микрорайон, 5', offices)
        self.assertEqual((office['id'], reason, note), (48, None, None))
        office, _, note = self.script.match_office('Алматы', 'ул. Жамбыла, 172', offices)
        self.assertEqual((office['id'], note), (47, 'дом совпал без литеры'))
        office, reason, _ = self.script.match_office('Астана', 'пр. Сарыарка, 31', offices)
        self.assertIsNone(office)
        self.assertIn('несколько', reason)


class WiringTests(unittest.TestCase):

    def setUp(self):
        self.app = _read(APP_JSX)

    def test_schema_is_deployed_after_wiki_and_water(self):
        source = _read(DATABASE_PY)
        water = source.index('            self._init_water_schema_tx(cursor)\n')
        thermo = source.index('            self._init_thermoboxes_schema_tx(cursor)\n')
        wiki = source.index('self._init_wiki_schema_tx(cursor)')
        self.assertLess(wiki, thermo)
        self.assertLess(water, thermo)
        self.assertIn('def _init_thermoboxes_schema_tx(self, cursor):', source)
        self.assertIn('SAVEPOINT thermoboxes_schema', source)

    def test_blueprint_is_registered(self):
        source = _read(BOT_PY)
        self.assertIn('from thermoboxes.routes import build_thermoboxes_blueprint', source)
        self.assertIn('app.register_blueprint(build_thermoboxes_blueprint(', source)

    def test_lazy_view_without_qr_gate(self):
        self.assertIn(
            "const ThermoboxesView = lazyWithRetry(() => import('./components/thermoboxes/ThermoboxesView'));",
            self.app)
        block = self.app.split('{view === "thermoboxes" && canAccessThermoboxesSection && (')[1].split('</Suspense>')[0]
        self.assertIn('<ThermoboxesView', block)
        self.assertNotIn('SensitiveSectionGate', block)

    def test_menu_item_is_declared_once_in_the_common_part(self):
        self.assertEqual(self.app.count("handleSidebarViewNavigation(e, 'thermoboxes')"), 1)
        item = self.app.split("handleSidebarViewNavigation(e, 'thermoboxes')")[1][:600]
        self.assertIn('Термокороба', item)
        self.assertLess(self.app.index("handleSidebarViewNavigation(e, 'water')"),
                        self.app.index("handleSidebarViewNavigation(e, 'thermoboxes')"))

    def test_predicate_mirrors_the_backend_perimeter(self):
        self.assertIn("const THERMOBOXES_SECTION_DEPARTMENT_CODES = ['front_office', 'szov'];", self.app)
        self.assertEqual(set(access.SECTION_DEPARTMENT_CODES), {'front_office', 'szov'})
        predicate = self.app.split('const canAccessThermoboxesSectionForUser = (userLike) => {')[1].split('\n};')[0]
        self.assertIn("if (role === 'trainer') return false;", predicate)
        self.assertIn("role === 'admin' && !isDepartmentHead(userLike)", predicate)

    def test_prototype_variants_are_gone(self):
        """Пять вариантов «заметнее» были прототипом для выбора (01.10.2026);
        выбран «Светофор» — переключатель и папка вариантов в сборку не идут."""
        folder = ROOT / 'src' / 'components' / 'thermoboxes'
        self.assertFalse((folder / 'variants').exists())
        view = _read(folder / 'ThermoboxesView.jsx')
        self.assertNotIn('variants', view)
        self.assertNotIn('thermo_variant', view)

    def test_menu_icons_are_components_not_elements(self):
        """IosMenu рисует `<Icon />` сам: элемент вместо компонента роняет всю
        страницу (React #130) при открытии меню «···» — так было на стенде."""
        view = _read(ROOT / 'src' / 'components' / 'thermoboxes' / 'ThermoboxesView.jsx')
        menu = view.split('const menuItems = [')[1].split('];')[0]
        self.assertNotRegex(menu, r'icon:\s*<')
        self.assertIn('icon: Download', menu)

    def test_guards_and_registries(self):
        self.assertIn("if (view === 'thermoboxes' && canAccessThermoboxesSection) return;", self.app)
        self.assertIn("    thermoboxes: ['front_office', 'szov'],", self.app)
        self.assertIn("    thermoboxes: 'Thermoboxes',", self.app)
        self.assertIn("canAccessThermoboxesSection && deptAllowsInner('thermoboxes'),", self.app)
        trainer = self.app.split('const TRAINER_ALLOWED_VIEWS = Object.freeze([')[1].split(']);')[0]
        self.assertNotIn("'thermoboxes'", trainer)
        # Флаг доступа обязан стоять в зависимостях гарда видимости: иначе
        # пришедший позже профиль не пустит человека в раздел до перезагрузки.
        deps = self.app.split("if (view === 'thermoboxes' && canAccessThermoboxesSection) return;")[1]
        deps = deps.split('}, [')[1].split(']);')[0]
        self.assertIn('canAccessThermoboxesSection', deps)


if __name__ == '__main__':
    unittest.main()
