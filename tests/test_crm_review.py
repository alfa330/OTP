# -*- coding: utf-8 -*-
"""Проверка супервайзером ДО группы: «Сотрудничество с Яндексом».

Возврат задачи #297 (постановщик, 06.10.2026): «Сотрудничество с
Яндексом не нужно отправлять в группу. Это обращение должен проверить СВ или
оно должно только фиксироваться». Решение владельца 07.10.2026 — проверка СВ,
как у жалоб на Яндекс: супервайзер жмёт «Решено» с итогом либо «Отправить в
группу»; само в группу такое обращение не уходит никогда.

Что здесь сторожится — по слоям, как устроен раздел:

* правило       — кто проходит проверку: ветка «С Яндексом» и только она;
* права         — кто решает, и что нельзя делать с обращением, пока не решили;
* запросы       — «ждёт МОЕЙ проверки» считается одной формулой на счётчик,
                  колокол, фильтр и бейдж строки;
* сервис        — в группу обращение не уходит никаким путём, кроме решения;
* роуты         — отказ доходит до отказа: в базу и в Telegram ничего не ушло;
* файл          — приложенный к такому обращению файл не теряется;
* выгрузка и интерфейс — называют это состояние своими словами.

Ни одно из правил не падает само, если сломается: обращение просто уйдёт в
группу, в которую его просили не отправлять.
"""

import io
import re
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from complaints import queries as complaint_queries
from crm import access, files, queries, report, scenarios as sc, schema, service

from tests.test_crm_access import RecordingCursor
from tests.test_crm_routing import (
    ALL_CHATS, ALL_QUEUES, KASPI_CHAT, FakeCursor, admin_ctx, build_client, chat_row,
    queue_row, route_row,
)
from tests.test_crm_service import FakeDb, FakeQueries, FakeTransport, PoolViolation

ROOT = Path(__file__).resolve().parents[1]

KEY = 'cooperation'
COOP_CHAT = -5200000001

# Как на проде: у очереди «Сотрудничество» своего чата нет, тема уведена в
# группу маршрутом.
COOP_QUEUE = queue_row(6, 'cooperation', 'Сотрудничество', chat_id=None)
COOP_CHATS = ALL_CHATS + [chat_row(COOP_CHAT, 'Сотрудничество')]
COOP_ROUTE = route_row(KEY, COOP_CHAT, 'Сотрудничество')

YANDEX = {
    'coop_channel': 'Звонок', 'coop_park': 'iTaxi', 'coop_city': 'Алматы',
    'coop_type': sc.COOP_WITH_YANDEX, 'coop_phone': '+7 701 123 45 67',
    'coop_comment': 'Менеджер по партнёрам, просит связаться с директором',
}
PARKS = {
    'coop_channel': 'Звонок', 'coop_park': 'iTaxi', 'coop_city': 'Алматы',
    'coop_type': sc.COOP_WITH_PARKS, 'coop_company': 'ТОО Альфа',
    'coop_phone': '+7 701 123 45 67', 'coop_service': 'Подключение водителей',
}


def viewer(role='operator', user_id=10, department_id=1, headed=(), groups=(),
           department_code='szov'):
    return {
        'user_id': user_id, 'name': 'Проверяющий', 'role': role,
        'department_id': department_id, 'department_code': department_code,
        'headed_department_ids': list(headed),
        'headed_department_codes': ['szov'] if headed else [],
        'group_ids': list(groups),
    }


def pending(created_by=158, department_id=1, author_group_ids=(6,), **changes):
    """Обращение на проверке — в том виде, в каком его отдаёт get_ticket."""
    item = {
        'id': 110, 'subject': 'Сотрудничество с Яндексом', 'body': 'Комментарий: …',
        'status': 'open', 'created_by': created_by, 'created_by_name': 'Оператор',
        'department_id': department_id, 'queue_department_id': None,
        'author_group_ids': list(author_group_ids), 'scenario_key': KEY,
        'delivery_status': 'pending', 'tg_chat_id': None, 'tg_message_id': None,
        'review_state': schema.REVIEW_PENDING, 'answers': dict(YANDEX),
    }
    item.update(changes)
    return item


# ─────────────────────────────────────────────────────────────────────────────
# Правило
# ─────────────────────────────────────────────────────────────────────────────

class RuleTest(unittest.TestCase):
    """Проверку проходит ветка «С Яндексом» — и больше ничего."""

    def test_yandex_branch_goes_to_the_supervisor(self):
        self.assertTrue(sc.needs_review(KEY, YANDEX))

    def test_parks_branch_goes_straight_to_the_group(self):
        """«С таксопарками» правка не трогает: возврат был только про Яндекс."""
        self.assertFalse(sc.needs_review(KEY, PARKS))

    def test_no_other_topic_is_reviewed(self):
        """Правило владельца узкое, и расширять его «по смыслу» нельзя: ни одна
        другая тематика проверку не проходит, что бы в ответах ни стояло."""
        self.assertEqual([item['key'] for item in sc.SCENARIOS if item.get('review_when')], [KEY])
        for item in sc.SCENARIOS:
            if item['key'] != KEY:
                self.assertFalse(sc.needs_review(item['key'], YANDEX), item['key'])
        self.assertFalse(sc.needs_review('нет такой тематики', YANDEX))

    def test_unanswered_type_is_not_a_review(self):
        self.assertFalse(sc.needs_review(KEY, {}))
        self.assertFalse(sc.needs_review(KEY, None))

    def test_the_rule_names_a_real_question_and_a_real_option(self):
        """Правило — данными, как depends_on. Опечатка в ключе или в варианте
        не упала бы: обращение просто ушло бы в группу."""
        for item in sc.SCENARIOS:
            rule = item.get('review_when')
            if not rule:
                continue
            step = next((s for s in item['steps'] if s['key'] == rule[0]), None)
            self.assertIsNotNone(step, '%s: правило ссылается на несуществующий вопрос' % item['key'])
            self.assertEqual(step['kind'], sc.CHOICE)
            self.assertIn(rule[1], step['options'])
            self.assertFalse(step.get('depends_on'),
                             'вопрос правила обязан задаваться всегда, иначе оно недостижимо')

    def test_the_answer_of_a_branch_left_behind_does_not_decide(self):
        """Оператор выбрал «С Яндексом», потом «С таксопарками»: решает
        последний выбор, а хвост первой ветки сервер отбрасывает раньше."""
        answers = dict(PARKS, coop_comment='хвост ветки Яндекса')
        self.assertFalse(sc.needs_review(KEY, sc.visible_answers(KEY, answers)))

    def test_review_does_not_change_what_is_asked(self):
        """Проверка меняет адресата, а не обращение: вопросы, обязательность и
        текст остаются прежними — супервайзер читает то же, что читала группа."""
        verdict = sc.evaluate(KEY, YANDEX, has_attachment=False)
        self.assertEqual(verdict['outcome'], sc.READY)
        self.assertEqual(sc.render_subject(KEY, YANDEX), 'Сотрудничество с Яндексом')
        self.assertIn('Комментарий: Менеджер по партнёрам', sc.render_body(KEY, YANDEX))

    def test_the_choice_explains_where_each_option_leads(self):
        """По подписи «С Яндексом» не видно, кто получит обращение — это
        написано под «i» у самого выбора, а не строкой на экране."""
        step = next(s for s in sc.get(KEY)['steps'] if s['key'] == 'coop_type')
        self.assertIn('супервайзер', step['hint'])
        self.assertIn('в группу', step['hint'])

    def test_the_catalog_carries_no_second_copy_of_the_rule(self):
        """Куда пойдёт обращение, мастер узнаёт из ответа сервера. Правило в
        каталоге стало бы его второй копией — на клиенте."""
        item = next(i for i in sc.public_catalog() if i['key'] == KEY)
        self.assertNotIn('review_when', item)
        self.assertTrue(item['sends_to_group'], 'ветка таксопарков по-прежнему уходит в группу')


# ─────────────────────────────────────────────────────────────────────────────
# Права
# ─────────────────────────────────────────────────────────────────────────────

class AccessTest(unittest.TestCase):

    def test_any_supervisor_of_the_authors_department_decides(self):
        """Владелец, 09.10.2026: «чтобы мог проверить любой супервайзер
        отдела» — и супервайзер группы автора, и супервайзер соседней группы,
        и супервайзер без группы вовсе."""
        for me in (viewer('sv', user_id=55, groups=(6,)), viewer('sv', user_id=56, groups=(9,)),
                   viewer('supervisor', user_id=57)):
            self.assertTrue(access.can_review(me, pending()), me['user_id'])

    def test_supervisor_of_another_department_does_not(self):
        """Даже если автор — в его группе: проверяет отдел автора."""
        stranger = viewer('sv', user_id=77, department_id=2, department_code='op', groups=(6,))
        self.assertTrue(access.can_view_ticket(stranger, pending()), 'видит — но не решает')
        self.assertFalse(access.can_review(stranger, pending()))
        self.assertFalse(access.can_review(viewer('sv', user_id=78, department_id=None), pending()))

    def test_supervisor_does_not_decide_his_own_ticket(self):
        """Своё обращение решает коллега, иначе проверки не было бы вовсе."""
        self.assertFalse(access.can_review(viewer('sv', user_id=55, groups=(6,)),
                                           pending(created_by=55)))

    def test_head_of_the_authors_department_decides(self):
        """И когда у автора нет группы с супервайзером, и когда тот в отпуске."""
        head = viewer('admin', user_id=1, headed=(1,))
        self.assertTrue(access.can_review(head, pending(author_group_ids=())))
        self.assertTrue(access.can_review(head, pending()))

    def test_head_of_another_department_does_not(self):
        other = viewer('admin', user_id=3, headed=(7,))
        self.assertFalse(access.can_review(other, pending()))

    def test_global_admin_decides(self):
        self.assertTrue(access.can_review(viewer('super_admin', user_id=2), pending()))

    def test_operator_never_decides(self):
        """Ни своё, ни чужое: иначе проверки не было бы вовсе."""
        self.assertFalse(access.can_review(viewer('operator', user_id=158), pending()))
        self.assertFalse(access.can_review(viewer('operator', user_id=190), pending()))
        self.assertFalse(access.can_review(viewer('trainer', user_id=30), pending()))

    def test_only_a_pending_ticket_is_decided(self):
        boss = viewer('super_admin', user_id=2)
        for state in (None, schema.REVIEW_SENT, schema.REVIEW_RESOLVED):
            self.assertFalse(access.can_review(boss, pending(review_state=state)), state)

    def test_nobody_writes_to_a_ticket_that_is_not_in_a_group_yet(self):
        for me in (viewer('operator', user_id=158), viewer('sv', user_id=55, groups=(6,)),
                   viewer('super_admin', user_id=2)):
            self.assertFalse(access.can_reply(me, pending()), me['role'])

    def test_status_buttons_do_not_bypass_the_review(self):
        """«Вопрос решён» закрыл бы обращение в обход проверки, а «Вернуть в
        работу» сделал бы из решённого супервайзером открытое обращение,
        которое никто не получал."""
        for me in (viewer('operator', user_id=158), viewer('sv', user_id=55, groups=(6,)),
                   viewer('super_admin', user_id=2)):
            self.assertFalse(access.can_change_status(me, pending()), me['role'])
            self.assertFalse(access.can_change_status(
                me, pending(review_state=schema.REVIEW_RESOLVED, status='resolved')), me['role'])

    def test_a_ticket_sent_by_the_supervisor_lives_like_any_other(self):
        sent = pending(review_state=schema.REVIEW_SENT, tg_message_id=555)
        me = viewer('operator', user_id=158)
        self.assertTrue(access.can_reply(me, sent))
        self.assertTrue(access.can_change_status(me, sent))

    def test_ordinary_tickets_are_untouched(self):
        ordinary = pending(review_state=None)
        me = viewer('operator', user_id=158)
        self.assertTrue(access.can_reply(me, ordinary))
        self.assertTrue(access.can_change_status(me, ordinary))


# ─────────────────────────────────────────────────────────────────────────────
# Запросы
# ─────────────────────────────────────────────────────────────────────────────

def squash(sql):
    return ' '.join(str(sql).split())


class ReviewerRuleTest(unittest.TestCase):
    """«Ждёт МОЕЙ проверки» — одна формула на всё, что её показывает."""

    def test_same_rule_as_the_complaints_about_yandex(self):
        """Это одна и та же проверка у двух видов обращений. Разойдись формулы,
        супервайзер увидел бы жалобу оператора, но не его обращение."""
        ours = squash(queries.reviewer_sql('viewer_id'))
        theirs = squash(complaint_queries.reviewer_sql('viewer_id'))
        self.assertEqual(ours, theirs.replace('c.created_by', 't.created_by')
                         .replace('c.creator_department_id', 't.department_id'))

    def test_any_working_supervisor_of_the_department_reviews(self):
        """Любой работающий СВ отдела автора; других СВ нет — глава отдела.
        Уволенные и «на увольнении» не в счёт ни там, ни там: иначе восемь
        уволенных СВ навсегда закрыли бы главе задачу."""
        rule = squash(queries.reviewer_sql('viewer_id'))
        self.assertIn('rv.department_id = t.department_id', rule)
        self.assertIn("lower(COALESCE(rv.role, '')) IN ('sv', 'supervisor')", rule)
        self.assertEqual(rule.count("COALESCE(rv.status, 'working') NOT IN ('fired', 'dismissal')"), 2)
        self.assertIn('AND rv.id = %(viewer_id)s', rule)
        self.assertIn('AND rv.id IS DISTINCT FROM t.created_by', rule)
        self.assertIn('d.head_user_id = %(viewer_id)s AND d.is_active', rule)
        self.assertNotIn('group_operator_memberships', rule)

    def test_the_task_is_pending_and_not_my_own(self):
        task = squash(queries.review_task_sql('viewer_id'))
        self.assertIn("t.review_state = 'pending'", task)
        self.assertIn('t.created_by IS DISTINCT FROM %(viewer_id)s', task)
        self.assertIn(squash(queries.reviewer_sql('viewer_id')), task)

    def test_counter_bell_filter_and_badge_share_it(self):
        import inspect
        for function in (queries.counters, queries.review_for_bell, queries.list_tickets):
            self.assertIn('review_task_sql(', inspect.getsource(function), function.__name__)

    def test_the_pending_queue_has_a_partial_index(self):
        ddl = squash(' '.join(schema._STATEMENTS))
        self.assertIn("idx_crm_tickets_review_pending ON crm_tickets(created_by, created_at DESC) "
                      "WHERE review_state = 'pending'", ddl)

    def test_review_columns_are_migrated_without_a_check(self):
        """Расширять CHECK на живой таблице больно — значения держит код."""
        migration = next(sql for sql in schema._MIGRATIONS if 'review_state' in sql)
        for column in ('review_state', 'review_by', 'review_by_name', 'review_at', 'review_note'):
            self.assertIn('ADD COLUMN IF NOT EXISTS %s' % column, migration)
        self.assertNotIn('CHECK', migration.upper())

    def test_states_match_the_complaints_word_for_word(self):
        from complaints import catalog
        self.assertEqual((schema.REVIEW_PENDING, schema.REVIEW_SENT, schema.REVIEW_RESOLVED),
                         (catalog.REVIEW_PENDING, catalog.REVIEW_SENT, catalog.REVIEW_RESOLVED))


class TicketColumnsTest(unittest.TestCase):
    def test_width_matches_the_column_list(self):
        """Запросы дописывают свои столбцы следом и читают их по этому числу:
        забытое +1 прочитало бы review_note вместо «ждёт моей проверки»."""
        columns = [part for part in queries._TICKET_COLUMNS.split(',') if part.strip()]
        self.assertEqual(len(columns), queries._TICKET_WIDTH)

    def test_row_carries_the_review(self):
        row = [None] * queries._TICKET_WIDTH
        row[0], row[12] = 110, 158
        row[35], row[36], row[38] = 'resolved', 'Проверова Асель', 'Передала контакты директору'
        row[37] = datetime(2026, 10, 7, 11, 30)
        item = queries._ticket_row(row, viewer_id=158)
        self.assertEqual((item['review_state'], item['review_by_name'], item['review_note']),
                         ('resolved', 'Проверова Асель', 'Передала контакты директору'))
        self.assertEqual(item['review_at'], '2026-10-07T11:30:00')


class ListTest(unittest.TestCase):
    def test_review_filter_asks_for_my_tasks(self):
        cursor = RecordingCursor()
        queries.list_tickets(cursor, viewer('sv', user_id=55, groups=(6,)), review_only=True)
        where = squash(cursor.queries[0]).split(' WHERE ', 1)[1]
        self.assertIn(squash(queries.review_task_sql('viewer_id')), where)

    def test_review_filter_ignores_mine_status_and_unread(self):
        """С любым из них список был бы пуст всегда: обращения на проверке
        чужие по определению и все в одном состоянии."""
        cursor = RecordingCursor()
        queries.list_tickets(cursor, viewer('sv', user_id=55, groups=(6,)), review_only=True,
                             mine=True, status=['resolved'], unread_only=True)
        where = squash(cursor.queries[0]).split(' WHERE ', 1)[1]
        self.assertNotIn('t.created_by = %(viewer_id)s', where)
        self.assertNotIn('t.status = ANY', where)
        self.assertNotIn('author_unread_at IS NOT NULL', where)

    def test_ordinary_list_is_not_narrowed(self):
        cursor = RecordingCursor()
        queries.list_tickets(cursor, viewer('sv', user_id=55, groups=(6,)))
        where = squash(cursor.queries[0]).split(' WHERE ', 1)[1].split(' ORDER BY ')[0]
        self.assertNotIn('review_state', where)

    def test_row_says_whether_the_ticket_waits_for_me(self):
        """Бейдж «Ждёт проверки» горит у того, чья это задача, — а не у
        каждого, кому обращение видно."""
        mine = [None] * queries._TICKET_WIDTH + [True]
        theirs = [None] * queries._TICKET_WIDTH + [False]
        mine[0], theirs[0] = 110, 111
        cursor = RecordingCursor(rows=[tuple(mine), tuple(theirs)])
        items, _more = queries.list_tickets(cursor, viewer('sv', user_id=55, groups=(6,)))
        self.assertEqual([(item['id'], item['review_mine']) for item in items],
                         [(110, True), (111, False)])
        self.assertIn('AS review_mine', cursor.queries[0])

    def test_the_badge_column_is_computed_only_for_pending_rows(self):
        cursor = RecordingCursor()
        queries.list_tickets(cursor, viewer('super_admin'))
        self.assertIn("CASE WHEN t.review_state = 'pending' THEN", squash(cursor.queries[0]))


class WritesTest(unittest.TestCase):
    def test_new_ticket_can_be_born_pending(self):
        cursor = RecordingCursor(rows=[(110,)])
        queries.create_ticket(
            cursor, queue_id=6, topic_id=None, subject='Сотрудничество с Яндексом', body='…',
            priority='normal', source='manual', client_name=None, client_phone='+7701',
            created_by=158, created_by_name='Оператор', department_id=1,
            scenario_key=KEY, answers=YANDEX, review_state=schema.REVIEW_PENDING)
        self.assertIn('review_state', cursor.queries[0])
        self.assertEqual(cursor.params[-1], 'pending')

    def test_decision_lands_only_on_a_pending_ticket(self):
        """Два супервайзера, нажавшие разное одновременно, не перезаписывают
        решение друг друга: второй получает «уже не ждёт проверки»."""
        cursor = RecordingCursor()
        cursor.rowcount = 0
        self.assertFalse(queries.set_review(cursor, 110, state=schema.REVIEW_SENT,
                                            actor_id=55, actor_name='СВ', chat_id=COOP_CHAT))
        self.assertIn("WHERE id = %(id)s AND review_state = 'pending'", squash(cursor.queries[0]))
        cursor.rowcount = 1
        self.assertTrue(queries.set_review(cursor, 110, state=schema.REVIEW_RESOLVED,
                                           actor_id=55, actor_name='СВ', note='итог'))

    def test_resolving_closes_the_ticket_and_tells_the_author(self):
        cursor = RecordingCursor()
        queries.set_review(cursor, 110, state=schema.REVIEW_RESOLVED, actor_id=55,
                           actor_name='Проверова Асель', note='Передала контакты')
        self.assertTrue(cursor.params['resolved'])
        self.assertEqual(cursor.params['note'], 'Передала контакты')
        self.assertEqual(cursor.params['kind'], queries.UNREAD_REVIEWED)
        sql = squash(cursor.queries[0])
        self.assertIn("status = CASE WHEN %(resolved)s THEN 'resolved' ELSE status END", sql)
        # Себе не звоним: автор сам и проверял (глава отдела, заведший обращение).
        self.assertIn('created_by IS DISTINCT FROM %(actor_id)s', sql)

    def test_sending_stamps_the_address_and_keeps_the_ticket_open(self):
        due = datetime(2026, 10, 7, 13, 0)
        cursor = RecordingCursor()
        queries.set_review(cursor, 110, state=schema.REVIEW_SENT, actor_id=55, actor_name='СВ',
                           chat_id=COOP_CHAT, chat_title='Сотрудничество', due_at=due)
        self.assertFalse(cursor.params['resolved'])
        self.assertEqual((cursor.params['chat_id'], cursor.params['chat_title'],
                          cursor.params['due_at']), (COOP_CHAT, 'Сотрудничество', due))
        self.assertIsNone(cursor.params['note'])

    def test_every_new_query_composes(self):
        """Лишний процент в SQL ловится оператором % — без базы и без сети."""
        cursor = RecordingCursor()
        me = viewer('sv', user_id=55, groups=(6,))
        queries.review_for_bell(cursor, 55, 5)
        queries.counters(cursor, me)
        queries.stored_attachments(cursor, 110)
        queries.drop_message(cursor, 110, 3)
        queries.delivery_payload(cursor, 110)
        for search in (None, 'яндекс', '110', '100%'):
            queries.list_tickets(cursor, me, review_only=True, search=search, queue_id=6)
        self.assertGreater(len(cursor.queries), 8)

    def test_bell_rows_are_counted_by_the_same_condition_as_the_counter(self):
        cursor = RecordingCursor()
        queries.review_for_bell(cursor, 55, 5)
        self.assertIn(squash(queries.review_task_sql('user_id')), squash(cursor.queries[0]))
        self.assertIn('COUNT(*) OVER ()', cursor.queries[0])

    def test_delivery_knows_about_the_review(self):
        cursor = RecordingCursor(rows=[tuple([None] * 20 + ['pending'])])
        self.assertEqual(queries.delivery_payload(cursor, 110)['review_state'], 'pending')


# ─────────────────────────────────────────────────────────────────────────────
# Сервис
# ─────────────────────────────────────────────────────────────────────────────

class ReviewQueries(FakeQueries):
    """SQL-слой сервиса с проверкой: то же, что у остальных тестов сервиса,
    плюс новые вызовы. resolve_route — настоящая: адрес считает она."""

    UNREAD_REVIEWED = 'reviewed'

    def __init__(self, db, *, route_context=None, review_lands=True, stored=(), **kwargs):
        super().__init__(db, **kwargs)
        self._route_context = route_context
        self._review_lands = review_lands
        self._stored = list(stored)

    def routing_context(self, _cursor):
        self._record('routing_context')
        return self._route_context

    resolve_route = staticmethod(queries.resolve_route)

    def set_review(self, _cursor, ticket_id, **kwargs):
        self._record('set_review', ticket_id=ticket_id, **kwargs)
        if self._review_lands and self._payload is not None:
            # Решение легло — обращение больше не «на проверке».
            self._payload['review_state'] = kwargs['state']
            if kwargs.get('chat_id'):
                self._payload['chat_id'] = kwargs['chat_id']
        return self._review_lands

    def stored_attachments(self, _cursor, ticket_id):
        self._record('stored_attachments', ticket_id=ticket_id)
        return list(self._stored)

    def drop_message(self, _cursor, ticket_id, message_id):
        self._record('drop_message', ticket_id=ticket_id, message_id=message_id)


def coop_context(routes=(COOP_ROUTE,), chats=None):
    return queries.routing_context(
        FakeCursor(ALL_QUEUES + [COOP_QUEUE], routes, COOP_CHATS if chats is None else chats))


def coop_payload(**changes):
    payload = {
        'subject': 'Сотрудничество с Яндексом', 'body': sc.render_body(KEY, YANDEX),
        'priority': 'normal', 'status': 'open', 'due_at': None,
        'client_name': None, 'client_phone': '+7 701 123 45 67',
        'created_by': 158, 'created_by_name': 'Оператор',
        'delivery_status': 'pending', 'tg_message_id': None, 'chat_id': None,
        'queue_title': 'Сотрудничество', 'topic_title': None, 'department_name': 'СЗоВ',
        'scenario_key': KEY, 'answers': dict(YANDEX), 'flags': [],
        'queue_mentions': [], 'review_state': schema.REVIEW_PENDING,
    }
    payload.update(changes)
    return payload


class FakeBlob:
    def __init__(self, store, name):
        self.store, self.name = store, name

    def upload_from_string(self, data, content_type=None):
        self.store.objects[self.name] = bytes(data)
        self.store.types[self.name] = content_type

    def exists(self):
        return self.name in self.store.objects

    def download_as_bytes(self):
        return self.store.objects[self.name]

    def delete(self):
        del self.store.objects[self.name]


class FakeStorage:
    """Бакет в памяти — тем же контрактом, что приходит в раздел фабрикой."""

    def __init__(self, bucket='otp-tasks'):
        self.bucket_name = bucket
        self.objects = {}
        self.types = {}

    def bucket(self, _name):
        return self

    def blob(self, name):
        return FakeBlob(self, name)

    def gcs(self):
        return {'bucket_name': lambda: self.bucket_name, 'client': lambda: self}


class ServiceCase(unittest.TestCase):
    def setUp(self):
        self.db = FakeDb()
        self._queries, self._transport = service.queries, service.transport
        self.addCleanup(self._restore)

    def _restore(self):
        service.queries, service.transport = self._queries, self._transport

    def wire(self, *, payload=None, ticket=None, error=None, file_error=None, **extra):
        service.queries = ReviewQueries(self.db, payload=payload, ticket=ticket, **extra)
        service.transport = FakeTransport(self.db, error=error, file_error=file_error)
        return service.queries, service.transport


class GroupIsClosedTest(ServiceCase):
    """В группу обращение на проверке не уходит НИКАКИМ путём."""

    def test_delivery_refuses_a_pending_ticket(self):
        """Одна отправка на все входы: создание, «Отправить ещё раз», решение.
        Запрет стоит в ней — роут, забывший проверку, ничего не отнесёт."""
        fake, transport = self.wire(payload=coop_payload(chat_id=COOP_CHAT))
        ok, error = service.deliver_ticket(self.db, 110)
        self.assertFalse(ok)
        self.assertIn('проверки', error)
        self.assertEqual((transport.sent, transport.files), ([], []))
        self.assertNotIn('set_delivery', fake.kinds(), 'отказ — не «сбой доставки»')

    def test_ordinary_ticket_is_delivered_as_before(self):
        _fake, transport = self.wire(payload=coop_payload(
            chat_id=COOP_CHAT, review_state=None, subject='Сотрудничество с таксопарками · ТОО'))
        ok, error = service.deliver_ticket(self.db, 110)
        self.assertTrue(ok, error)
        self.assertEqual(len(transport.sent), 1)


class SendDecisionTest(ServiceCase):
    CTX = {'user_id': 55, 'name': 'Проверова Асель'}

    def test_supervisor_sends_the_ticket_to_the_group_of_the_topic(self):
        fake, transport = self.wire(payload=coop_payload(), ticket=pending(),
                                    route_context=coop_context())
        delivered, error = service.review_send(self.db, 110, ctx=self.CTX)
        self.assertTrue(delivered, error)
        decision = fake.find('set_review')[0]
        self.assertEqual((decision['state'], decision['actor_id'], decision['chat_id'],
                          decision['chat_title']),
                         (schema.REVIEW_SENT, 55, COOP_CHAT, 'Сотрудничество'))
        self.assertEqual(transport.sent[0]['chat_id'], COOP_CHAT)
        # Группа получает то же сообщение, что получала до правки: просьба
        # тематики первой строкой и данные обращения.
        self.assertIn('Просьба связаться по вопросу сотрудничества', transport.sent[0]['text'])
        self.assertIn('<b>Тип сотрудничества:</b> С Яндексом', transport.sent[0]['text'])
        self.assertIn('review_sent', [call['kind'] for call in fake.find('add_event')])

    def test_decision_is_written_before_the_network(self):
        """И сеть — без открытого курсора: пул делят колокол и аукцион."""
        fake, _transport = self.wire(payload=coop_payload(), ticket=pending(),
                                     route_context=coop_context())
        service.review_send(self.db, 110, ctx=self.CTX)
        self.assertLess(fake.kinds().index('set_review'), fake.kinds().index('delivery_payload'))
        self.assertEqual(self.db.max_depth, 1)

    def test_address_is_the_one_the_topic_has_now(self):
        """Обращение ждало проверки, и маршрут за это время могли поменять:
        уходит оно туда, куда тема идёт сейчас."""
        moved = route_row(KEY, KASPI_CHAT, 'Sapar/Kaspi - отмена')
        fake, transport = self.wire(payload=coop_payload(), ticket=pending(),
                                    route_context=coop_context(routes=(moved,)))
        service.review_send(self.db, 110, ctx=self.CTX)
        self.assertEqual(fake.find('set_review')[0]['chat_id'], KASPI_CHAT)
        self.assertEqual(transport.sent[0]['chat_id'], KASPI_CHAT)

    def test_no_group_means_no_decision(self):
        """Отправлять некуда — обращение остаётся ждать, а не числится
        отправленным в никуда."""
        fake, transport = self.wire(payload=coop_payload(), ticket=pending(),
                                    route_context=coop_context(routes=()))
        with self.assertRaises(service.ReviewError) as caught:
            service.review_send(self.db, 110, ctx=self.CTX)
        self.assertEqual(caught.exception.status, 400)
        self.assertNotIn('set_review', fake.kinds())
        self.assertEqual(transport.sent, [])

    def test_a_group_the_bot_left_is_named(self):
        fake, _transport = self.wire(payload=coop_payload(), ticket=pending(),
                                     route_context=coop_context(chats=ALL_CHATS))
        with self.assertRaises(service.ReviewError) as caught:
            service.review_send(self.db, 110, ctx=self.CTX)
        self.assertIn('«Сотрудничество»', str(caught.exception))
        self.assertNotIn('set_review', fake.kinds())

    def test_second_supervisor_is_told_the_ticket_is_decided(self):
        fake, transport = self.wire(payload=coop_payload(), ticket=pending(),
                                    route_context=coop_context(), review_lands=False)
        with self.assertRaises(service.ReviewError) as caught:
            service.review_send(self.db, 110, ctx=self.CTX)
        self.assertEqual(caught.exception.status, 409)
        self.assertEqual(transport.sent, [], 'в группу второй раз ничего не ушло')
        self.assertNotIn('add_event', fake.kinds())

    def test_telegram_refusal_keeps_the_decision(self):
        """Решение принято, доставка — нет: дальше обычная «Отправить ещё раз»."""
        fake, _transport = self.wire(payload=coop_payload(), ticket=pending(),
                                     route_context=coop_context(), error='бота нет в группе')
        delivered, error = service.review_send(self.db, 110, ctx=self.CTX)
        self.assertFalse(delivered)
        self.assertEqual(error, 'бота нет в группе')
        self.assertEqual(fake.find('set_review')[0]['state'], schema.REVIEW_SENT)
        self.assertEqual(fake.find('set_delivery')[0]['status'], 'failed')

    def test_missing_ticket_is_404(self):
        self.wire(ticket=None, route_context=coop_context())
        with self.assertRaises(service.ReviewError) as caught:
            service.review_send(self.db, 999, ctx=self.CTX)
        self.assertEqual(caught.exception.status, 404)


class ResolveDecisionTest(ServiceCase):
    CTX = {'user_id': 55, 'name': 'Проверова Асель'}

    def test_the_outcome_is_mandatory(self):
        """Итог — то, что узнает оператор. Без него «Решено» ничего не говорит."""
        fake, _transport = self.wire(payload=coop_payload())
        for empty in (None, '', '   ', '\n'):
            with self.assertRaises(service.ReviewError):
                service.review_resolve(self.db, 110, empty, ctx=self.CTX)
        self.assertEqual(fake.calls, [], 'пустой итог до базы не доходит')

    def test_outcome_closes_the_ticket_and_lands_in_the_thread(self):
        fake, transport = self.wire(payload=coop_payload())
        service.review_resolve(self.db, 110, '  Передала контакты директору  ', ctx=self.CTX)
        decision = fake.find('set_review')[0]
        self.assertEqual((decision['state'], decision['note'], decision['actor_name']),
                         (schema.REVIEW_RESOLVED, 'Передала контакты директору', 'Проверова Асель'))
        note = fake.find('add_message')[0]
        self.assertEqual((note['direction'], note['body'], note['author_user_id'],
                          note['author_name']),
                         ('note', 'Передала контакты директору', 55, 'Проверова Асель'))
        self.assertIn('review_resolved', [call['kind'] for call in fake.find('add_event')])
        # В группу не уходит ничего: обращение там не было, и отбивать нечего.
        self.assertEqual((transport.sent, transport.files), ([], []))

    def test_everything_is_one_transaction(self):
        self.wire(payload=coop_payload())
        service.review_resolve(self.db, 110, 'итог', ctx=self.CTX)
        self.assertEqual(self.db.opened, 1)

    def test_decided_ticket_takes_no_second_outcome(self):
        fake, _transport = self.wire(payload=coop_payload(), review_lands=False)
        with self.assertRaises(service.ReviewError) as caught:
            service.review_resolve(self.db, 110, 'итог', ctx=self.CTX)
        self.assertEqual(caught.exception.status, 409)
        self.assertNotIn('add_message', fake.kinds())

    def test_the_outcome_has_a_ceiling(self):
        fake, _transport = self.wire(payload=coop_payload())
        service.review_resolve(self.db, 110, 'я' * 9000, ctx=self.CTX)
        self.assertEqual(len(fake.find('set_review')[0]['note']), service.REVIEW_NOTE_LIMIT)


# ─────────────────────────────────────────────────────────────────────────────
# Файл обращения, которому пока некуда уйти
# ─────────────────────────────────────────────────────────────────────────────

class FileStoreTest(unittest.TestCase):
    def test_store_load_drop(self):
        storage = FakeStorage()
        ref = files.store(storage.gcs(), data=b'%PDF-1.7', filename='КП для Яндекса.pdf',
                          content_type='application/pdf')
        self.assertTrue(files.is_stored(ref))
        self.assertTrue(ref.startswith('gs://otp-tasks/crm/attachments/'))
        self.assertTrue(ref.endswith('_КП для Яндекса.pdf'), 'кириллица в имени сохраняется')
        self.assertEqual(files.load(storage.gcs(), ref), b'%PDF-1.7')
        self.assertEqual(list(storage.types.values()), ['application/pdf'])
        files.drop(storage.gcs(), [ref])
        self.assertIsNone(files.load(storage.gcs(), ref))

    def test_telegram_file_id_is_not_ours(self):
        for file_id in ('AgACAgIAAxkBAAIB', '', None, 'gs:/кривая', 'https://x/y'):
            self.assertFalse(files.is_stored(file_id), repr(file_id))
        self.assertIsNone(files.load(FakeStorage().gcs(), 'AgACAgIAAxkBAAIB'))

    def test_without_storage_the_file_is_refused_not_lost(self):
        for gcs in (None, {}, {'bucket_name': lambda: '', 'client': lambda: None}):
            self.assertFalse(files.storage_ready(gcs))
            with self.assertRaises(files.StorageError):
                files.store(gcs, data=b'x', filename='a.pdf')

    def test_storage_failure_is_a_refusal_too(self):
        class Broken(FakeStorage):
            def blob(self, name):
                raise RuntimeError('503 от хранилища')

        with self.assertLogs(level='WARNING'), self.assertRaises(files.StorageError):
            files.store(Broken().gcs(), data=b'x', filename='a.pdf')

    def test_name_cannot_escape_the_prefix(self):
        self.assertEqual(files.safe_name('..\\..\\etc/passwd'), 'passwd')
        self.assertEqual(files.safe_name('  '), 'file')
        self.assertEqual(files.safe_name('a:b*c?.pdf'), 'abc.pdf')

    def test_dropping_tolerates_a_missing_object(self):
        storage = FakeStorage()
        with self.assertLogs(level='WARNING'):
            files.drop(storage.gcs(), ['gs://otp-tasks/crm/attachments/нет'])
        files.drop(storage.gcs(), ['AgACAgIAAxkBAAIB', None])
        files.drop(None, ['gs://otp-tasks/x'])


class StoredFileFollowsTheTicketTest(ServiceCase):
    """Супервайзер отправил обращение в группу — файл уходит следом, как у
    любого другого обращения, и больше у нас не лежит."""

    def stored(self, storage):
        ref = files.store(storage.gcs(), data=b'%PDF-1.7', filename='КП.pdf',
                          content_type='application/pdf')
        return {'id': 31, 'ref': ref, 'name': 'КП.pdf', 'mime': 'application/pdf',
                'author_user_id': 158, 'author_name': 'Оператор'}

    def test_file_goes_to_the_group_and_leaves_the_storage(self):
        storage = FakeStorage()
        item = self.stored(storage)
        fake, transport = self.wire(payload=coop_payload(chat_id=COOP_CHAT, review_state='sent'),
                                    stored=[item])
        ok, error = service.deliver_ticket(self.db, 110, gcs=storage.gcs())
        self.assertTrue(ok, error)
        sent = transport.files[0]
        self.assertEqual((sent['chat_id'], sent['file_name'], sent['reply_to_message_id']),
                         (COOP_CHAT, 'КП.pdf', 555))
        self.assertEqual(sent['stream'].read(), b'%PDF-1.7')
        # В нити вместо нашей строки — обычная, с file_id Telegram.
        self.assertEqual(fake.find('drop_message'), [{'ticket_id': 110, 'message_id': 31}])
        attached = [call for call in fake.find('add_message') if call.get('attachment')]
        self.assertEqual(attached[0]['attachment']['file_id'], 'ph1')
        self.assertEqual(attached[0]['author_user_id'], 158, 'файл остаётся файлом оператора')
        self.assertEqual(storage.objects, {})

    def test_refused_file_stays_with_us(self):
        """Отказ Telegram на файле обращения не отменяет — и файл не теряется:
        он по-прежнему открывается из карточки."""
        storage = FakeStorage()
        item = self.stored(storage)
        fake, _transport = self.wire(payload=coop_payload(chat_id=COOP_CHAT, review_state='sent'),
                                     stored=[item], file_error='file is too big')
        with self.assertLogs(level='WARNING'):
            ok, _error = service.deliver_ticket(self.db, 110, gcs=storage.gcs())
        self.assertTrue(ok, 'обращение в группе, отказ только по файлу')
        self.assertEqual(fake.find('drop_message'), [])
        self.assertEqual(len(storage.objects), 1)

    def test_vanished_file_does_not_stop_the_ticket(self):
        storage = FakeStorage()
        item = self.stored(storage)
        storage.objects.clear()
        fake, transport = self.wire(payload=coop_payload(chat_id=COOP_CHAT, review_state='sent'),
                                    stored=[item])
        with self.assertLogs(level='WARNING'):
            ok, _error = service.deliver_ticket(self.db, 110, gcs=storage.gcs())
        self.assertTrue(ok)
        self.assertEqual(transport.files, [])
        self.assertEqual(fake.find('drop_message'), [])

    def test_storage_is_not_read_under_a_cursor(self):
        storage = FakeStorage()
        item = self.stored(storage)
        db = self.db

        class Guarded(FakeStorage):
            def blob(inner, name):  # noqa: N805
                if db.depth:
                    raise PoolViolation('хранилище вызвано с открытым курсором')
                return FakeBlob(storage, name)

        self.wire(payload=coop_payload(chat_id=COOP_CHAT, review_state='sent'), stored=[item])
        ok, _error = service.deliver_ticket(self.db, 110, gcs=Guarded().gcs())
        self.assertTrue(ok)

    def test_without_storage_nothing_is_looked_up(self):
        """Обычная отправка за файлами не ходит: лишний запрос на каждое
        обращение ради случая, которого у него быть не может."""
        fake, _transport = self.wire(payload=coop_payload(chat_id=COOP_CHAT, review_state=None))
        service.deliver_ticket(self.db, 110)
        self.assertNotIn('stored_attachments', fake.kinds())


# ─────────────────────────────────────────────────────────────────────────────
# Роуты: боевой Blueprint на подменённом SQL-слое
# ─────────────────────────────────────────────────────────────────────────────

def client_for(ctx, gcs=None, routes=(COOP_ROUTE,), chats=None, coop_queue=COOP_QUEUE):
    from flask import Flask
    from crm import routes as crm_routes
    from tests.test_crm_routing import FakeDb as RoutingDb

    cursor = FakeCursor(ALL_QUEUES + [coop_queue], routes, COOP_CHATS if chats is None else chats)
    app = Flask(__name__)
    app.register_blueprint(crm_routes.build_crm_blueprint(
        db=RoutingDb(cursor),
        require_api_key=lambda handler: handler,
        build_cors_preflight_response=lambda: ('', 204),
        resolve_requester=lambda: (ctx['user_id'], {'id': ctx['user_id']}, None),
        sensitive_access_granted=lambda user_id: True,
        gcs=gcs,
    ))
    return app.test_client(), cursor


class CreateTest(unittest.TestCase):
    def setUp(self):
        self.ctx = admin_ctx(role='operator')
        self.created = {}
        self.messages = []
        self.events = []
        self.delivered = []
        patches = [
            mock.patch.object(queries, 'load_access_context', lambda cursor, user_id: self.ctx),
            mock.patch.object(queries, 'create_ticket', self._create),
            mock.patch.object(queries, 'add_message',
                              lambda cursor, **kw: self.messages.append(kw) or 1),
            mock.patch.object(queries, 'add_event',
                              lambda cursor, **kw: self.events.append(kw)),
            mock.patch.object(queries, 'get_ticket',
                              lambda cursor, ticket_id, viewer_id=None: {'id': ticket_id}),
            mock.patch.object(service, 'deliver_ticket', self._deliver),
        ]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)

    def _create(self, cursor, **kwargs):
        self.created = kwargs
        return 110

    def _deliver(self, db, ticket_id, attachment=None, gcs=None):
        self.delivered.append(ticket_id)
        return True, None

    def post(self, answers, **kwargs):
        client, _cursor = client_for(self.ctx, **kwargs)
        return client.post('/api/crm/tickets', json={'scenario_key': KEY, 'answers': answers})

    def test_yandex_ticket_is_saved_and_not_sent(self):
        response = self.post(YANDEX)
        self.assertEqual(response.status_code, 201, response.get_json())
        data = response.get_json()
        self.assertTrue(data['review'])
        self.assertFalse(data['delivered'])
        self.assertIsNone(data['delivery_error'], 'это не сбой доставки — так и должно быть')
        self.assertEqual(self.delivered, [], 'в группу не ушло ничего')
        self.assertEqual(self.created['review_state'], schema.REVIEW_PENDING)

    def test_pending_ticket_has_no_address_and_no_group_deadline(self):
        """Оно никуда не ушло: адрес и срок ответа группы появятся, только
        если супервайзер его туда отправит."""
        self.post(YANDEX)
        self.assertIsNone(self.created['tg_chat_id'])
        self.assertIsNone(self.created['tg_chat_title'])
        self.assertIsNone(self.created['due_at'])
        self.assertEqual(self.created['queue_id'], 6, 'числится за своей тематикой')
        created = next(event for event in self.events if event['kind'] == 'created')
        self.assertTrue(created['payload']['review'])
        self.assertIsNone(created['payload']['chat'])

    def test_group_deadline_does_not_tick_while_the_supervisor_checks(self):
        """У тематики есть срок ответа группы — но обращению на проверке он не
        ставится: иначе оно «просрочилось» бы, ещё не дойдя до группы. Срок
        считается в момент «Отправить в группу»."""
        with_sla = COOP_QUEUE[:6] + (60,) + COOP_QUEUE[7:]
        self.post(YANDEX, coop_queue=with_sla)
        self.assertIsNone(self.created['due_at'])
        # Тот же срок у таксопарка ставится сразу — значит, срок у тематики
        # действительно есть и проверка выше не пустая.
        self.post(PARKS, coop_queue=with_sla)
        self.assertIsNotNone(self.created['due_at'])

    def test_parks_ticket_goes_to_the_group_as_before(self):
        response = self.post(PARKS)
        self.assertEqual(response.status_code, 201, response.get_json())
        data = response.get_json()
        self.assertFalse(data['review'])
        self.assertTrue(data['delivered'])
        self.assertEqual(self.delivered, [110])
        self.assertIsNone(self.created['review_state'])
        self.assertEqual((self.created['tg_chat_id'], self.created['tg_chat_title']),
                         (COOP_CHAT, 'Сотрудничество'))

    def test_the_server_decides_not_the_client(self):
        """Клиент не может ни отправить обращение Яндекса в группу, ни
        приписать проверку чужой ветке: решают только ответы."""
        client, _cursor = client_for(self.ctx)
        client.post('/api/crm/tickets', json={'scenario_key': KEY, 'answers': YANDEX,
                                              'review': False, 'review_state': None})
        self.assertEqual(self.delivered, [])
        self.assertEqual(self.created['review_state'], schema.REVIEW_PENDING)

    def test_branch_switched_at_the_last_moment_is_decided_by_the_last_choice(self):
        self.post(dict(YANDEX, coop_type=sc.COOP_WITH_PARKS, coop_company='ТОО Альфа',
                       coop_service='Подключение'))
        self.assertEqual(self.delivered, [110])
        self.assertIsNone(self.created['review_state'])

    def test_topic_without_a_group_is_closed_for_both_branches(self):
        """Мастер такую тему не предлагает вовсе; сервер отвечает тем же."""
        response = self.post(YANDEX, routes=())
        self.assertEqual(response.status_code, 400)
        self.assertFalse(self.created)

    def test_preview_tells_the_wizard_where_the_ticket_goes(self):
        client, _cursor = client_for(self.ctx)
        for answers, expected in ((YANDEX, True), (PARKS, False)):
            data = client.post('/api/crm/scenarios/%s/evaluate' % KEY,
                               json={'answers': answers}).get_json()
            self.assertEqual(data['outcome'], sc.READY)
            self.assertIs(data['review'], expected)

    def test_unfinished_ticket_is_not_called_a_review(self):
        client, _cursor = client_for(self.ctx)
        data = client.post('/api/crm/scenarios/%s/evaluate' % KEY,
                           json={'answers': {'coop_type': sc.COOP_WITH_YANDEX}}).get_json()
        self.assertEqual(data['outcome'], sc.INCOMPLETE)
        self.assertNotIn('review', data)


class CreateWithFileTest(CreateTest):
    def post_file(self, answers, gcs):
        import json
        client, _cursor = client_for(self.ctx, gcs=gcs)
        return client.post('/api/crm/tickets', data={
            'scenario_key': KEY, 'answers': json.dumps(answers),
            'attachment': (io.BytesIO(b'%PDF-1.7'), 'КП.pdf', 'application/pdf'),
        }, content_type='multipart/form-data')

    def test_file_of_a_pending_ticket_is_kept_by_us(self):
        storage = FakeStorage()
        response = self.post_file(YANDEX, storage.gcs())
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(list(storage.objects.values()), [b'%PDF-1.7'])
        row = self.messages[0]
        self.assertEqual((row['ticket_id'], row['direction'], row['author_user_id']),
                         (110, 'out', 7))
        attachment = row['attachment']
        self.assertTrue(files.is_stored(attachment['file_id']))
        self.assertEqual((attachment['kind'], attachment['name'], attachment['mime'],
                          attachment['size']), ('document', 'КП.pdf', 'application/pdf', 8))
        self.assertEqual(self.delivered, [])

    def test_picture_is_stored_as_a_picture(self):
        import json
        storage = FakeStorage()
        client, _cursor = client_for(self.ctx, gcs=storage.gcs())
        client.post('/api/crm/tickets', data={
            'scenario_key': KEY, 'answers': json.dumps(YANDEX),
            'attachment': (io.BytesIO(b'\x89PNG'), 'чат.png', 'image/png'),
        }, content_type='multipart/form-data')
        self.assertEqual(self.messages[0]['attachment']['kind'], 'photo')

    def test_without_storage_the_ticket_is_not_created_at_all(self):
        """Иначе оператор считал бы, что приложил файл, которого нигде нет."""
        response = self.post_file(YANDEX, None)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.get_json()['code'], 'CRM_FILE_NOT_STORED')
        self.assertFalse(self.created)

    def test_file_is_not_left_behind_when_the_ticket_is_refused(self):
        import json
        storage = FakeStorage()
        client, _cursor = client_for(self.ctx, gcs=storage.gcs(), routes=())
        response = client.post('/api/crm/tickets', data={
            'scenario_key': KEY, 'answers': json.dumps(YANDEX),
            'attachment': (io.BytesIO(b'%PDF-1.7'), 'КП.pdf', 'application/pdf'),
        }, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(storage.objects, {}, 'файл обращения, которого нет, не хранится')

    def test_file_is_not_left_behind_when_the_base_fails(self):
        storage = FakeStorage()

        def broken(cursor, **kwargs):
            raise RuntimeError('база отвалилась')

        with mock.patch.object(queries, 'create_ticket', broken), \
                self.assertLogs(level='ERROR'):
            response = self.post_file(YANDEX, storage.gcs())
        self.assertEqual(response.status_code, 500)
        self.assertEqual(storage.objects, {})

    def test_parks_file_still_goes_to_telegram_not_to_us(self):
        storage = FakeStorage()
        response = self.post_file(PARKS, storage.gcs())
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(storage.objects, {})
        self.assertEqual(self.messages, [])
        self.assertEqual(self.delivered, [110])


class DecisionEndpointTest(unittest.TestCase):
    def setUp(self):
        self.ctx = viewer('sv', user_id=55, groups=(6,))
        self.ticket = pending()
        self.calls = []
        patches = [
            mock.patch.object(queries, 'load_access_context', lambda cursor, user_id: self.ctx),
            mock.patch.object(queries, 'get_ticket',
                              lambda cursor, ticket_id, viewer_id=None: self.ticket),
            mock.patch.object(queries, 'list_messages', lambda cursor, ticket_id: []),
            mock.patch.object(service, 'review_send', self._send),
            mock.patch.object(service, 'review_resolve', self._resolve),
        ]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)
        self.send_result = (True, None)

    def _send(self, db, ticket_id, *, ctx, gcs=None):
        self.calls.append(('send', ticket_id, ctx['user_id']))
        return self.send_result

    def _resolve(self, db, ticket_id, note, *, ctx):
        self.calls.append(('resolve', ticket_id, note, ctx['user_id']))

    def post(self, payload, ticket_id=110):
        client, _cursor = client_for(self.ctx)
        return client.post('/api/crm/tickets/%s/review' % ticket_id, json=payload)

    def test_supervisor_resolves_with_an_outcome(self):
        response = self.post({'decision': 'resolve', 'note': 'Передала контакты'})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(self.calls, [('resolve', 110, 'Передала контакты', 55)])
        data = response.get_json()
        self.assertEqual(set(data), {'item', 'messages', 'permissions', 'delivered',
                                     'delivery_error'})
        self.assertIsNone(data['delivered'], '«Решено» про доставку ничего не говорит')

    def test_supervisor_sends_to_the_group(self):
        response = self.post({'decision': 'send'})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(self.calls, [('send', 110, 55)])
        self.assertTrue(response.get_json()['delivered'])

    def test_undelivered_send_is_still_a_decision(self):
        self.send_result = (False, 'бота нет в группе')
        data = self.post({'decision': 'send'}).get_json()
        self.assertEqual((data['delivered'], data['delivery_error']), (False, 'бота нет в группе'))

    def test_author_is_refused_and_nothing_happens(self):
        self.ctx = viewer('operator', user_id=158)
        for decision in ('send', 'resolve'):
            response = self.post({'decision': decision, 'note': 'сам решил'})
            self.assertEqual(response.status_code, 403, decision)
        self.assertEqual(self.calls, [])

    def test_supervisor_of_another_group_of_the_department_decides(self):
        self.ctx = viewer('sv', user_id=56, groups=(9,))
        response = self.post({'decision': 'send'})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(self.calls, [('send', 110, 56)])

    def test_supervisor_is_refused_on_his_own_ticket(self):
        self.ctx = viewer('sv', user_id=56, groups=(9,))
        self.ticket = pending(created_by=56)
        response = self.post({'decision': 'send'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['error'], 'Проверяет обращение супервайзер отдела')
        self.assertEqual(self.calls, [])

    def test_outsider_is_stopped_at_the_door_of_the_section(self):
        """Раздел закрыт — закрыта и проверка, даже у супервайзера группы
        автора: периметр держит вход в раздел, и он стоит раньше любого права."""
        self.ctx = viewer('sv', user_id=77, groups=(6,), department_code='op')
        response = self.post({'decision': 'resolve', 'note': 'итог'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'CRM_SECTION_CLOSED')
        self.assertEqual(self.calls, [])

    def test_decided_ticket_is_409_before_anyone_is_asked(self):
        for state in (schema.REVIEW_SENT, schema.REVIEW_RESOLVED, None):
            self.ticket = pending(review_state=state)
            self.assertEqual(self.post({'decision': 'send'}).status_code, 409, state)
        self.assertEqual(self.calls, [])

    def test_missing_ticket_is_404(self):
        self.ticket = None
        self.assertEqual(self.post({'decision': 'send'}, 9999).status_code, 404)

    def test_unknown_decision_is_400(self):
        for payload in ({}, {'decision': 'cancel'}, {'decision': ''}):
            self.assertEqual(self.post(payload).status_code, 400)
        self.assertEqual(self.calls, [])

    def test_refusals_of_the_service_keep_their_codes(self):
        def refuse(db, ticket_id, note, *, ctx):
            raise service.ReviewError('Напишите итог — что выяснили и что сделали')

        with mock.patch.object(service, 'review_resolve', refuse):
            response = self.post({'decision': 'resolve', 'note': ''})
        self.assertEqual(response.status_code, 400)
        self.assertIn('итог', response.get_json()['error'])

        def raced(db, ticket_id, *, ctx, gcs=None):
            raise service.ReviewError('Обращение уже не ждёт проверки — обновите карточку', 409)

        with mock.patch.object(service, 'review_send', raced):
            self.assertEqual(self.post({'decision': 'send'}).status_code, 409)


class PendingTicketIsLockedTest(unittest.TestCase):
    """Обходные дороги в группу и к статусу закрыты у обращения на проверке."""

    def setUp(self):
        self.ctx = viewer('operator', user_id=158)
        self.ticket = pending()
        self.forbidden = []
        patches = [
            mock.patch.object(queries, 'load_access_context', lambda cursor, user_id: self.ctx),
            mock.patch.object(queries, 'get_ticket',
                              lambda cursor, ticket_id, viewer_id=None: self.ticket),
            mock.patch.object(queries, 'list_messages', lambda cursor, ticket_id: []),
            mock.patch.object(queries, 'mark_seen_by_author', lambda *args: False),
            mock.patch.object(service, 'deliver_ticket', self._trap('deliver')),
            mock.patch.object(service, 'post_operator_reply', self._trap('reply')),
            mock.patch.object(service, 'change_status_from_system', self._trap('status')),
        ]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)

    def _trap(self, name):
        def trapped(*args, **kwargs):
            self.forbidden.append(name)
            return True, None
        return trapped

    def test_resend_does_not_push_it_to_the_group(self):
        """delivery_status у него «не отправлено» — как у сбоя доставки. Без
        запрета «Отправить ещё раз» отнесла бы его в группу мимо супервайзера."""
        for self.ctx in (viewer('operator', user_id=158), viewer('sv', user_id=55, groups=(6,)),
                         viewer('super_admin', user_id=2)):
            client, _cursor = client_for(self.ctx)
            self.assertEqual(client.post('/api/crm/tickets/110/resend').status_code, 403,
                             self.ctx['role'])
        self.assertEqual(self.forbidden, [])

    def test_nobody_writes_and_nobody_closes(self):
        client, _cursor = client_for(self.ctx)
        self.assertEqual(client.post('/api/crm/tickets/110/messages',
                                     json={'body': 'есть ответ?'}).status_code, 403)
        for status in ('resolved', 'cancelled', 'in_progress'):
            self.assertEqual(client.post('/api/crm/tickets/110/status',
                                         json={'status': status}).status_code, 403, status)
        self.assertEqual(self.forbidden, [])

    def test_resolved_by_the_supervisor_is_not_reopened(self):
        self.ticket = pending(review_state=schema.REVIEW_RESOLVED, status='resolved')
        client, _cursor = client_for(viewer('super_admin', user_id=2))
        self.assertEqual(client.post('/api/crm/tickets/110/status',
                                     json={'status': 'open'}).status_code, 403)
        self.assertEqual(self.forbidden, [])

    def test_card_tells_who_may_decide(self):
        client, _cursor = client_for(self.ctx)
        mine = client.get('/api/crm/tickets/110').get_json()['permissions']
        self.assertEqual((mine['can_review'], mine['can_reply'], mine['can_change_status']),
                         (False, False, False))
        self.ctx = viewer('sv', user_id=55, groups=(6,))
        client, _cursor = client_for(self.ctx)
        self.assertTrue(client.get('/api/crm/tickets/110').get_json()['permissions']['can_review'])

    def test_list_filter_reaches_the_query(self):
        seen = {}

        def listed(cursor, ctx, **kwargs):
            seen.update(kwargs)
            return [], False

        with mock.patch.object(queries, 'list_tickets', listed):
            client, _cursor = client_for(viewer('sv', user_id=55, groups=(6,)))
            client.get('/api/crm/tickets?review=1&mine=1')
            self.assertTrue(seen['review_only'])
            client.get('/api/crm/tickets?mine=1')
            self.assertFalse(seen['review_only'])


class StoredFileRouteTest(unittest.TestCase):
    def setUp(self):
        self.ctx = viewer('sv', user_id=55, groups=(6,))
        self.storage = FakeStorage()
        self.ref = files.store(self.storage.gcs(), data=b'%PDF-1.7', filename='КП.pdf',
                               content_type='application/pdf')
        self.found = {'file_id': self.ref, 'name': 'КП.pdf', 'mime': 'application/pdf',
                      'kind': 'document'}
        patches = [
            mock.patch.object(queries, 'load_access_context', lambda cursor, user_id: self.ctx),
            mock.patch.object(queries, 'get_ticket',
                              lambda cursor, ticket_id, viewer_id=None: pending()),
            mock.patch.object(queries, 'find_message_attachment',
                              lambda cursor, ticket_id, message_id: self.found),
        ]
        for item in patches:
            item.start()
            self.addCleanup(item.stop)

    def test_supervisor_opens_the_file_before_deciding(self):
        from crm import transport

        def telegram(file_id):
            raise AssertionError('файл обращения на проверке в Telegram не лежит')

        with mock.patch.object(transport, 'fetch_file', telegram):
            client, _cursor = client_for(self.ctx, gcs=self.storage.gcs())
            response = client.get('/api/crm/tickets/110/attachments/31')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, b'%PDF-1.7')
        self.assertEqual(response.mimetype, 'application/pdf')

    def test_lost_file_is_an_honest_error(self):
        self.storage.objects.clear()
        client, _cursor = client_for(self.ctx, gcs=self.storage.gcs())
        self.assertEqual(client.get('/api/crm/tickets/110/attachments/31').status_code, 502)

    def test_telegram_files_are_served_as_before(self):
        from crm import transport
        self.found = dict(self.found, file_id='AgACAgIAAxkBAAIB')
        with mock.patch.object(transport, 'fetch_file', lambda file_id: (b'tg', None)):
            client, _cursor = client_for(self.ctx, gcs=self.storage.gcs())
            self.assertEqual(client.get('/api/crm/tickets/110/attachments/31').data, b'tg')

    def test_deleting_the_ticket_takes_its_file_along(self):
        self.ctx = viewer('super_admin', user_id=2)
        with mock.patch.object(queries, 'stored_attachments',
                               lambda cursor, ticket_id: [{'id': 31, 'ref': self.ref}]):
            client, cursor = client_for(self.ctx, gcs=self.storage.gcs())
            response = client.delete('/api/crm/tickets/110')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.storage.objects, {})
        self.assertTrue(any(sql.startswith('DELETE FROM crm_tickets')
                            for sql, _params in cursor.executed))


# ─────────────────────────────────────────────────────────────────────────────
# Колокол
# ─────────────────────────────────────────────────────────────────────────────

class BellTriggerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        database = (ROOT / 'database.py').read_text(encoding='utf-8')
        start = database.index('def _init_bell_notify_schema_tx(self, cursor):')
        cls.block = database[start:database.index('def _init_amo_leads_schema_tx', start)]
        branch = cls.block[cls.block.index("TG_TABLE_NAME = 'crm_tickets'"):]
        cls.branch = branch[:branch.index("ELSIF TG_TABLE_NAME = 'complaints'")]

    def test_reviewers_are_woken_for_a_reviewed_ticket_only(self):
        """Обычное обращение по-прежнему будит одного автора: звонить
        супервайзеру по чужой переписке — шум."""
        self.assertIn('targets := ARRAY[NEW.created_by]', self.branch)
        tail = self.branch[self.branch.index('IF NEW.review_state IS NOT NULL THEN'):]
        # Все работающие супервайзеры отдела автора и глава — а не только
        # супервайзер его группы (владелец, 09.10.2026).
        self.assertIn('u.department_id = NEW.department_id', tail)
        self.assertIn("COALESCE(u.status, 'working') NOT IN ('fired', 'dismissal')", tail)
        self.assertIn("lower(COALESCE(u.role, '')) IN ('sv', 'supervisor')", tail)
        self.assertIn('d.id = NEW.department_id', tail)
        self.assertNotIn('gsm.supervisor_id', tail)

    def test_complaints_wake_the_same_circle(self):
        """Жалобы на Яндекс проверяет тот же круг — и будит его так же."""
        branch = self.block[self.block.index("ELSIF TG_TABLE_NAME = 'complaints'"):]
        branch = branch[:branch.index("ELSIF TG_TABLE_NAME = 'tasks'")]
        review = branch[branch.index('IF NEW.review_state IS NOT NULL THEN'):]
        review = review[:review.index('END IF;')]
        self.assertIn('u.department_id = NEW.creator_department_id', review)
        self.assertIn("NOT IN ('fired', 'dismissal')", review)
        self.assertIn('d.id = NEW.creator_department_id', review)
        self.assertNotIn('gsm.supervisor_id', review)

    def test_new_pending_ticket_wakes_the_supervisor_at_once(self):
        """Обычная вставка колокол не будит (автору нечего читать), а вставшее
        на проверку — будит: у супервайзера появилась задача."""
        self.assertRegex(self.block, r"'trg_bell_crm_tickets_review', 'crm_tickets', 'AFTER INSERT',"
                                     r"\s+\"\"\"WHEN \(NEW\.review_state IS NOT NULL\)\"\"\"")

    def test_decision_wakes_everyone_who_was_asked(self):
        self.assertIn("'AFTER UPDATE OF author_unread_at, status, delivery_status, review_state'",
                      self.block)
        self.assertIn('OLD.review_state IS DISTINCT FROM NEW.review_state', self.block)

    def test_plain_thread_writes_still_wake_nobody(self):
        update = self.block[self.block.index("'trg_bell_crm_tickets',"):]
        update = update[:update.index('),')]
        self.assertNotIn('last_message_at', update)


# ─────────────────────────────────────────────────────────────────────────────
# Выгрузка
# ─────────────────────────────────────────────────────────────────────────────

def exported(**changes):
    item = {
        'id': 110, 'subject': 'Сотрудничество с Яндексом', 'body': 'Комментарий: …',
        'status': 'open', 'queue_title': 'Сотрудничество', 'topic_title': None,
        'scenario_key': KEY, 'created_by_name': 'Оператор', 'department_name': 'СЗоВ',
        'client_name': None, 'client_phone': '87011234567', 'answers': dict(YANDEX),
        'tg_chat_title': None, 'queue_chat_title': 'Сотрудничество', 'delivery_status': 'pending',
        'flags': [], 'created_at': '2026-10-07T11:07:00', 'first_reply_at': None,
        'resolved_at': None, 'resolved_by_name': None,
        'review_state': schema.REVIEW_PENDING, 'review_by_name': None, 'review_at': None,
        'review_note': None,
    }
    item.update(changes)
    return item


class ExportTest(unittest.TestCase):
    def row(self, item):
        columns = report.columns_for([item])
        return dict(zip([key for _t, key, _w in columns], report.row_values(item, columns)))

    def test_pending_ticket_is_not_called_sent(self):
        row = self.row(exported())
        self.assertEqual(row['status_title'], 'На проверке')
        self.assertEqual(row['review_title'], 'Ждёт проверки')
        # В группе его не было: ни группы, ни «ждёт отправки».
        self.assertEqual((row['group_title'], row['delivery_title']), ('', ''))

    def test_resolved_by_the_supervisor_carries_the_outcome(self):
        row = self.row(exported(status='resolved', review_state=schema.REVIEW_RESOLVED,
                                review_by_name='Проверова Асель', review_at='2026-10-07T11:30:00',
                                review_note='Передала контакты директору',
                                resolved_at='2026-10-07T11:30:00', resolved_by_name='Проверова Асель'))
        self.assertEqual((row['status_title'], row['review_title'], row['review_by_name'],
                          row['review_note']),
                         ('Решено', 'Решено супервайзером', 'Проверова Асель',
                          'Передала контакты директору'))
        self.assertEqual((row['group_title'], row['delivery_title']), ('', ''))

    def test_sent_by_the_supervisor_reads_like_a_sent_ticket(self):
        row = self.row(exported(review_state=schema.REVIEW_SENT, tg_chat_title='Сотрудничество',
                                delivery_status='sent', review_by_name='Проверова Асель'))
        self.assertEqual((row['status_title'], row['review_title'], row['group_title'],
                          row['delivery_title']),
                         ('Отправлено', 'Отправлено в группу', 'Сотрудничество', 'Доставлено'))

    def test_review_columns_appear_only_when_there_is_a_review(self):
        """Проверку проходят единицы обращений — четыре пустые колонки в каждой
        остальной выгрузке были бы шумом."""
        ordinary = exported(review_state=None, tg_chat_title='Сотрудничество',
                            delivery_status='sent')
        self.assertIs(report.columns_for([ordinary]), report.COLUMNS)
        self.assertIs(report.columns_for([]), report.COLUMNS)
        titles = [title for title, _k, _w in report.columns_for([ordinary, exported()])]
        self.assertEqual(titles[titles.index('Кто закрыл') + 1:titles.index('Метки')],
                         ['Проверка супервайзера', 'Кто проверил', 'Когда проверил',
                          'Итог проверки'])

    def test_workbook_carries_the_columns_and_explains_them(self):
        seen = {}

        def patch(stream, sqref, sheet_path):
            seen['sqref'] = sqref
            return stream

        items = [exported(), exported(id=111, review_state=schema.REVIEW_RESOLVED,
                                      status='resolved', review_at='2026-10-07T11:30:00',
                                      review_note='Передала контакты')]
        book = load_workbook(report.build_workbook(
            items, date_from='2026-10-01', date_to='2026-10-07', generated_by='СВ',
            generated_at=datetime(2026, 10, 7, 12, 0), text_warning_patch=patch))
        sheet = book['Обращения']
        header = [cell.value for cell in sheet[1]]
        second = {title: cell for title, cell in zip(header, sheet[3])}
        self.assertEqual(second['Итог проверки'].value, 'Передала контакты')
        self.assertIsInstance(second['Когда проверил'].value, datetime)
        context = ' '.join(str(cell.value) for row in book['Контекст'].iter_rows() for cell in row)
        self.assertIn('Проверка супервайзера', context)
        # Телефон и ИИН остались на своих местах — уголок гасится там же.
        self.assertEqual(seen['sqref'], 'J2:K3')

    def test_ordinary_export_is_unchanged(self):
        ordinary = exported(review_state=None, tg_chat_title='Сотрудничество',
                            delivery_status='sent')
        book = load_workbook(report.build_workbook(
            [ordinary], date_from='2026-10-01', date_to='2026-10-07', generated_by='СВ',
            generated_at=datetime(2026, 10, 7, 12, 0)))
        header = [cell.value for cell in book['Обращения'][1]]
        self.assertEqual(header, [title for title, _k, _w in report.COLUMNS])
        context = ' '.join(str(cell.value) for row in book['Контекст'].iter_rows() for cell in row)
        self.assertNotIn('Проверка супервайзера', context)


# ─────────────────────────────────────────────────────────────────────────────
# Интерфейс: проводка, которую не видит ни один тест логики
# ─────────────────────────────────────────────────────────────────────────────

class FrontendWiringTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        crm = ROOT / 'src' / 'components' / 'crm'
        cls.view = (crm / 'CrmTicketsView.jsx').read_text(encoding='utf-8')
        cls.wizard = (crm / 'TicketWizard.jsx').read_text(encoding='utf-8')
        cls.rows = (crm / 'ticketList.js').read_text(encoding='utf-8')
        cls.complaint = (ROOT / 'src' / 'components' / 'complaints'
                         / 'ComplaintCard.jsx').read_text(encoding='utf-8')
        cls.shared = (ROOT / 'src' / 'components' / 'common'
                      / 'SupervisorReview.jsx').read_text(encoding='utf-8')
        card = cls.view[cls.view.index('const TicketCard = ('):]
        cls.card = card[:card.index('const ComplaintThreadCard = (')]

    def test_decision_panel_stands_where_the_reply_field_was(self):
        """По праву с сервера, а не по роли — и раньше поля ответа."""
        bottom = self.card[self.card.index('{/* Ответ и действия */}'):]
        self.assertLess(bottom.index('permissions.can_review ? ('),
                        bottom.index('permissions.can_reply ? ('))
        self.assertIn('<ReviewPanel', bottom)
        self.assertNotIn("role === 'sv'", self.card)

    def test_both_buttons_call_the_review_endpoint(self):
        self.assertIn('/api/crm/tickets/${ticketId}/review', self.card)
        self.assertIn("onSend={() => review('send')}", self.card)
        self.assertIn("onResolve={(note) => review('resolve', note)}", self.card)

    def test_one_review_panel_for_complaints_and_tickets(self):
        """Проверка у них одна и та же — и окно решения одно: два разных окна
        про одно решение человек прочитал бы как два разных действия."""
        marker = "from '../common/SupervisorReview'"
        self.assertIn(marker, self.view)
        self.assertIn(marker, self.complaint)
        self.assertNotIn('const ResolveModal', self.complaint)
        for label in ('Проверка супервайзером', 'Отправить в группу', 'Решено'):
            self.assertIn(label, self.shared)
            self.assertEqual(self.card.count('>%s<' % label), 0, label)

    def test_resolve_window_says_what_closes(self):
        """Окно общее, слова свои: без подписи окно молчало бы о главном —
        решение закроет карточку, и в группу она не уйдёт."""
        self.assertIn('subtitle="Жалоба закроется с вашим итогом и в группу не уйдёт"',
                      self.complaint)
        self.assertIn('subtitle="Обращение закроется с вашим итогом и в группу не уйдёт"',
                      self.card)

    def test_states_are_the_same_words_as_on_the_server(self):
        for name, value in (('REVIEW_PENDING', schema.REVIEW_PENDING),
                            ('REVIEW_SENT', schema.REVIEW_SENT),
                            ('REVIEW_RESOLVED', schema.REVIEW_RESOLVED)):
            self.assertIn("export const %s = '%s';" % (name, value), self.rows)

    def test_wizard_names_the_destination_by_the_servers_word(self):
        self.assertIn('review: Boolean(data.review)', self.wizard)
        self.assertIn('deliveryWording(Boolean(preview?.review))', self.wizard)
        self.assertIn('createdToast(created)', self.wizard)
        self.assertNotIn('Так обращение увидят в группе', self.wizard)
        self.assertNotIn('Подтвердить и отправить', self.wizard)

    def test_review_segment_asks_the_server(self):
        self.assertIn("params.set('review', '1')", self.view)
        self.assertIn('stateFilters(STATE_FILTERS,', self.view)

    def test_review_queue_is_not_narrowed_to_complaints(self):
        """Жалоб в очереди проверки не бывает: строка «Жалобы» в фильтре групп
        оставила бы под сегментом с числом пустой список. Вход в сегмент
        сбрасывает фильтр, а сама строка в нём не предлагается."""
        self.assertIn("if (item.key === REVIEW_FILTER) setQueueFilter('');", self.view)
        self.assertRegex(self.view, r"complaintsEnabled && !reviewing\s+"
                                    r"\? \[\{ value: COMPLAINTS_FILTER, label: 'Жалобы' \}\] : \[\]")

    def test_badge_of_the_section_counts_what_the_bell_counts(self):
        """Число в меню ставят и колокол, и раздел: разойдись суммы — оно
        прыгало бы между двумя значениями."""
        self.assertRegex(self.view, r"Number\(counters\.unread\) \|\| 0\) \+ "
                                    r"\(Number\(counters\.review\) \|\| 0\)\s+\+ complaintsUnread")

    def test_history_survives_on_a_reviewed_ticket(self):
        self.assertIn('|| Boolean(ticket.review_state)) && (', self.card)

    def test_blueprint_receives_the_storage(self):
        bot = (ROOT / 'bot_schedule2.py').read_text(encoding='utf-8')
        start = bot.index('app.register_blueprint(build_crm_blueprint(')
        self.assertIn("gcs={'bucket_name': _crm_bucket_name, 'client': get_gcs_client}",
                      bot[start:start + 900])
        self.assertRegex(bot, re.compile(
            r"def _crm_bucket_name\(\):.*?GOOGLE_CLOUD_STORAGE_BUCKET_CRM.*?"
            r"GOOGLE_CLOUD_STORAGE_BUCKET_TASKS", re.S))


if __name__ == '__main__':
    unittest.main()
