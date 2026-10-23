# -*- coding: utf-8 -*-
"""«Дисциплина»: ручки Blueprint целиком — Flask test client, настоящий Postgres,
хранилище файлов подменено словарём. DISCIPLINE_TEST_PORT — порт кластера.

Ловит то, чего не видят тесты сервиса: разбор формы и файлов, коды ответов,
сериализацию дат и сумм в JSON, выгрузки xlsx и DOCX, вход рядовых без QR.

    DISCIPLINE_TEST_PORT=55461 python -m pytest tests/test_discipline_routes_postgres.py
"""
import io
import os
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

import pytest

from discipline import catalog, report
from tests.test_discipline_postgres import (  # noqa: F401 — фикстура db
    ADMIN, HEAD_SZOV, OP_A, OP_B, SUPER, SV_A, SV_B, SZOV, TRAINEE_A, PORT, db, section_open_for_all,
    use_sample_ladder)

flask = pytest.importorskip('flask')
pytestmark = pytest.mark.skipif(not PORT, reason='needs local PostgreSQL (DISCIPLINE_TEST_PORT)')

PDF = b'%PDF-1.4\n' + b'0' * 200
PNG = b'\x89PNG\r\n\x1a\n' + b'0' * 200


class _Blob:
    def __init__(self, store, path):
        self.store, self.path = store, path

    def upload_from_string(self, data, content_type=None):
        self.store[self.path] = data

    def exists(self):
        return self.path in self.store

    def download_as_bytes(self):
        return self.store[self.path]

    def delete(self):
        self.store.pop(self.path, None)


class _Bucket:
    def __init__(self, store):
        self.store = store

    def blob(self, path):
        return _Blob(self.store, path)


class _Client:
    def __init__(self, store):
        self.store = store

    def bucket(self, _name):
        return _Bucket(self.store)


class _Sqlite:
    """db._get_cursor() поверх тестовой схемы."""

    def __init__(self, base):
        self.base = base

    def _get_cursor(self):
        return self.base.cursor()


@pytest.fixture
def client(db):
    from flask import Flask, request

    from discipline.routes import build_discipline_blueprint

    store = {}
    sent = []

    def resolve_requester():
        user_id = request.headers.get('X-User-Id')
        if not user_id:
            return None, None, ('Unauthorized', 401)
        return int(user_id), None, None

    app = Flask(__name__)
    app.register_blueprint(build_discipline_blueprint(
        db=_Sqlite(db), require_api_key=lambda fn: fn,
        build_cors_preflight_response=lambda: ('', 204), resolve_requester=resolve_requester,
        send_telegram=lambda chat_id, text, **kw: sent.append((chat_id, text)),
        web_app_base_url='https://portal.example', gcs={'bucket_name': lambda: 'bucket',
                                                         'client': lambda: _Client(store)}))
    test_client = app.test_client()
    test_client.store = store
    test_client.sent = sent
    return test_client


def h(user_id):
    return {'X-User-Id': str(user_id)}


def _type_id(client, user_id, operator_id, title='Сон на рабочем месте'):
    response = client.get('/api/discipline/operator-context?operator_id=%d' % operator_id, headers=h(user_id))
    assert response.status_code == 200, response.get_json()
    return next(t['id'] for t in response.get_json()['types'] if t['title'] == title)


def _create(client, user_id, operator_id, at, files=None, **extra):
    data = {'operator_id': str(operator_id), 'type_id': str(_type_id(client, user_id, operator_id)),
            'occurred_at': at.strftime('%Y-%m-%dT%H:%M'), 'shift_label': '09:00–18:00',
            'description': 'Оператор спал на рабочем месте во время смены, видно по камере.'}
    data.update(extra)
    if files:
        data['files'] = files
    return client.post('/api/discipline/records', data=data, headers=h(user_id),
                       content_type='multipart/form-data')


def test_gates(client):
    assert client.get('/api/discipline/ping').status_code == 401
    response = client.get('/api/discipline/ping', headers=h(ADMIN))
    assert response.status_code == 403 and response.get_json()['code'] == 'DISCIPLINE_SECTION_CLOSED'
    # Рядовой — без QR (владелец, 23.10.2026: «для операторов сделать доступ без qr»):
    # права, число своих записей и сами записи отдаются сразу.
    payload = client.get('/api/discipline/ping', headers=h(OP_A)).get_json()
    assert payload['capabilities']['scope'] == 'own' and payload['own_total'] == 0
    assert 'qr_granted' not in payload and 'requires_qr' not in payload['capabilities']
    assert client.get('/api/discipline/records', headers=h(OP_A)).status_code == 200
    payload = client.get('/api/discipline/ping', headers=h(HEAD_SZOV)).get_json()
    assert payload['capabilities']['can_settings'] and payload['departments'][0]['id'] == SZOV
    assert client.get('/api/discipline/settings?department_id=%d' % SZOV, headers=h(SV_A)).status_code == 403


def test_full_flow(client, db):
    now = datetime.now().replace(second=0, microsecond=0)
    # Лестница со скриншота: тренинг → объясн. №1 → №2 → №3 + акт — выбрана до первой записи.
    with db.cursor() as cur:
        use_sample_ladder(cur)
    response = _create(client, SV_A, OP_A, now - timedelta(days=3),
                       files=[(io.BytesIO(PNG), 'фото.png')])
    assert response.status_code == 201, response.get_json()
    first = response.get_json()['id']
    assert response.get_json()['number'].startswith('Д-')
    assert len(client.store) == 1
    # Неверный файл: расширение и содержимое спорят — ничего не пишется и не остаётся в хранилище.
    bad = _create(client, SV_A, OP_A, now - timedelta(days=2), files=[(io.BytesIO(b'MZ' + b'0' * 50), 'акт.pdf')])
    assert bad.status_code == 400 and len(client.store) == 1

    second = _create(client, SV_A, OP_A, now - timedelta(days=2)).get_json()['id']
    journal = client.get('/api/discipline/records', headers=h(SV_A)).get_json()
    assert [i['id'] for i in journal['items']] == [second, first]
    assert journal['items'][0]['step_title'] == 'Объясн. №1'
    # СВ чужой группы этих записей не видит.
    assert client.get('/api/discipline/records', headers=h(SV_B)).get_json()['items'] == []
    assert client.get('/api/discipline/records/%d' % first, headers=h(SV_B)).status_code == 403

    card = client.get('/api/discipline/records/%d' % first, headers=h(SV_A)).get_json()
    assert card['files'][0]['kind'] == 'evidence' and card['history'][0]['action'] == 'create_record'
    file_id = card['files'][0]['id']
    preview = client.get('/api/discipline/files/%d?inline=1' % file_id, headers=h(SV_A))
    assert preview.status_code == 200 and preview.data == PNG
    assert preview.headers['Content-Disposition'].startswith('inline')
    assert preview.headers['X-Content-Type-Options'] == 'nosniff'

    # Оператор загружает объяснительную сам — и видит только свои поля.
    own = client.get('/api/discipline/records', headers=h(OP_A)).get_json()['items']
    assert {i['id'] for i in own} == {first, second} and 'description' not in own[0]
    response = client.post('/api/discipline/records/%d/explanation' % second, headers=h(OP_A),
                           data={'file': (io.BytesIO(PDF), 'объяснительная.pdf')},
                           content_type='multipart/form-data')
    assert response.status_code == 200, response.get_json()
    # Чужое доказательство оператору не отдаётся.
    assert client.get('/api/discipline/files/%d' % file_id, headers=h(OP_A)).status_code == 403
    item = client.get('/api/discipline/records/%d' % second, headers=h(SV_A)).get_json()
    assert item['state'] == 'explanation_received'
    # Без QR оператор открывает и свою объяснительную, и свою карточку: снятый
    # замок вернуть точечно в одну из этих ручек — тест краснеет (разбор, 23.10.2026).
    explanation_id = next(f['id'] for f in item['files'] if f['kind'] == 'explanation')
    assert client.get('/api/discipline/files/%d' % explanation_id, headers=h(OP_A)).status_code == 200
    assert client.get('/api/discipline/operators/%d' % OP_A, headers=h(OP_A)).status_code == 200

    # Тренинг на первой ступени.
    response = client.post('/api/discipline/records/%d/training' % first, headers=h(SV_A), json={
        'date': (now - timedelta(days=1)).strftime('%Y-%m-%d'), 'start': '10:00', 'end': '10:30'})
    assert response.status_code == 200, response.get_json()
    response = client.post('/api/discipline/records/%d/close' % first, headers=h(SV_A))
    assert response.status_code == 200, response.get_json()

    third = _create(client, SV_A, OP_A, now - timedelta(days=1)).get_json()['id']
    fourth = _create(client, SV_A, OP_A, now - timedelta(hours=5)).get_json()['id']
    item = client.get('/api/discipline/records/%d' % fourth, headers=h(HEAD_SZOV)).get_json()
    assert item['action'] == 'act' and item['act_full']['number'].startswith('А-')
    act_id = item['act_full']['id']
    summary = client.get('/api/discipline/summary', headers=h(HEAD_SZOV)).get_json()
    # Черновик, заведённый системой, — ещё не составленный акт.
    assert summary['controls'][0]['operator_id'] == OP_A and summary['counters']['acts'] == 0

    draft = client.get('/api/discipline/acts/%d/draft.docx' % act_id, headers=h(SV_A))
    assert draft.status_code == 200 and zipfile.is_zipfile(io.BytesIO(draft.data))
    assert client.post('/api/discipline/acts/%d/file' % act_id, headers=h(SV_A),
                       data={'file': (io.BytesIO(PDF), 'акт.pdf')},
                       content_type='multipart/form-data').status_code == 403
    response = client.post('/api/discipline/acts/%d/file' % act_id, headers=h(HEAD_SZOV),
                           data={'file': (io.BytesIO(PDF), 'акт.pdf')}, content_type='multipart/form-data')
    assert response.status_code == 200, response.get_json()
    summary = client.get('/api/discipline/summary', headers=h(HEAD_SZOV)).get_json()
    assert summary['counters']['acts'] == 1
    response = client.patch('/api/discipline/acts/%d' % act_id, headers=h(HEAD_SZOV), json={'approve': True})
    assert response.status_code == 400 and response.get_json()['code'] == 'DECISION_MEASURE'  # решения ещё нет
    measures = client.get('/api/discipline/controls', headers=h(HEAD_SZOV)).get_json()['measures']
    fine = next(m['id'] for m in measures if m['title'] == 'Лишение переменной части (фиксированная сумма)')
    # Решение с деньгами — без суммы не утверждается: ошибка встаёт у поля суммы.
    response = client.patch('/api/discipline/acts/%d' % act_id, headers=h(HEAD_SZOV),
                            json={'decision_measure_id': fine, 'approve': True})
    assert response.status_code == 400 and response.get_json()['code'] == 'MEASURE_AMOUNT'
    assert client.patch('/api/discipline/acts/%d' % act_id, headers=h(HEAD_SZOV),
                        json={'decision_measure_id': fine, 'decision_amount': '20000',
                              'approve': True}).status_code == 200
    # Мера решением по акту — в выгрузке на удержания и во вкладке «Меры» карточки.
    report_rows = client.get('/api/discipline/reports/measures', headers=h(HEAD_SZOV)).get_json()['rows']
    assert [(r['basis'], r['measure_amount']) for r in report_rows] == [
        ('Акт %s' % item['act_full']['number'], 20000.0)]

    # Отчёты и выгрузки — JSON и xlsx.
    for kind in ('operators', 'types', 'supervisors', 'groups', 'measures', 'controls'):
        response = client.get('/api/discipline/reports/%s' % kind, headers=h(HEAD_SZOV))
        assert response.status_code == 200, (kind, response.get_json())
        xlsx = client.get('/api/discipline/reports/%s?format=xlsx' % kind, headers=h(HEAD_SZOV))
        assert xlsx.status_code == 200 and zipfile.is_zipfile(io.BytesIO(xlsx.data))
    export = client.get('/api/discipline/records/export', headers=h(SV_A))
    assert export.status_code == 200 and zipfile.is_zipfile(io.BytesIO(export.data))
    audit = client.get('/api/discipline/audit', headers=h(HEAD_SZOV)).get_json()
    actions = {row['action'] for row in audit['items']}
    assert {'create_record', 'view_file', 'upload_explanation', 'upload_act', 'export'} <= actions
    assert client.get('/api/discipline/audit', headers=h(SV_A)).status_code == 403

    card = client.get('/api/discipline/operators/%d' % OP_A, headers=h(HEAD_SZOV)).get_json()
    assert card['control'] and len(card['records']) == 4 and card['documents']
    assert {third, fourth} <= {r['id'] for r in card['records']}
    assert [m['measure_amount'] for m in card['measures']] == [20000.0]


def test_settings_roundtrip(client):
    payload = client.get('/api/discipline/settings?department_id=%d' % SZOV, headers=h(HEAD_SZOV)).get_json()
    assert payload['settings']['ladder'] == ['explanation', 'explanation', 'act']
    response = client.put('/api/discipline/settings', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'settings': {'ladder': ['training', 'explanation', 'explanation', 'act'],
                                            'period_days': 45}})
    assert response.status_code == 200 and response.get_json()['settings']['period_days'] == 45
    bad = client.put('/api/discipline/settings', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'settings': {'ladder': ['nap']}})
    assert bad.status_code == 400
    response = client.post('/api/discipline/types', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'title': 'Курение в неположенном месте', 'severity': 'light', 'weight': 1,
        'needs_explanation': True})
    assert response.status_code == 200
    duplicate = client.post('/api/discipline/types', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'title': 'курение в неположенном месте', 'severity': 'light', 'weight': 1})
    assert duplicate.status_code == 400 and 'уже есть' in duplicate.get_json()['error']
    response = client.post('/api/discipline/settings/template', headers=h(HEAD_SZOV),
                           data={'department_id': str(SZOV), 'file': (io.BytesIO(b'PK\x03\x04' + b'0' * 50),
                                                                       'шаблон.docx')},
                           content_type='multipart/form-data')
    assert response.status_code == 200, response.get_json()


def test_repeated_submission_returns_the_same_record(client, db):
    """Ответ на «Зафиксировать» потерялся в сети, СВ нажал ещё раз: та же метка
    формы — та же запись. Второе нарушение лестницу бы сдвинуло."""
    now = datetime.now().replace(second=0, microsecond=0)
    token = 'f0e1d2c3-b4a5-4697-8899-aabbccddeeff'
    first = _create(client, SV_A, OP_A, now - timedelta(hours=2), client_token=token,
                    files=[(io.BytesIO(PNG), 'фото.png')])
    assert first.status_code == 201, first.get_json()
    again = _create(client, SV_A, OP_A, now - timedelta(hours=2), client_token=token,
                    files=[(io.BytesIO(PNG), 'фото.png')])
    assert again.status_code == 200, again.get_json()
    assert again.get_json()['id'] == first.get_json()['id']
    assert again.get_json()['number'] == first.get_json()['number']
    # Копия файла из повтора в хранилище не осталась.
    assert len(client.store) == 1
    # Форму после сбоя переписали на другого человека — это уже другая запись.
    other = _create(client, SV_A, TRAINEE_A, now - timedelta(hours=1), client_token=token)
    assert other.status_code == 201 and other.get_json()['id'] != first.get_json()['id']
    with db.cursor() as cur:
        cur.execute('SELECT operator_id, client_token FROM discipline_records ORDER BY id')
        assert cur.fetchall() == [(OP_A, token), (TRAINEE_A, None)]


# ── Человек между отделами: что видит руководитель ───────────────────────────

def _cross_department_act(db):
    """Две записи заведены в ОП, человека перевели в СЗоВ, третья (акт) — в СЗоВ,
    затем человек ушёл в Тез. Глава СЗоВ видит основание акта, но не записи ОП."""
    from tests.test_discipline_postgres import NOW as T0, OP as OP_DEPT, TODAY as D0, create

    with db.cursor() as cur:
        r1 = create(cur, 200, 402, T0 - timedelta(days=10), department=OP_DEPT)
        r2 = create(cur, 200, 402, T0 - timedelta(days=8), department=OP_DEPT)
        cur.execute('UPDATE users SET department_id = %s WHERE id = 402', (SZOV,))
        cur.execute('UPDATE group_operator_memberships SET end_date = %s WHERE operator_id = 402',
                    (D0 - timedelta(days=6),))
        cur.execute('INSERT INTO group_operator_memberships (group_id, operator_id, start_date) VALUES (10, 402, %s)',
                    (D0 - timedelta(days=5),))
        r3 = create(cur, HEAD_SZOV, 402, T0 - timedelta(days=2))
        cur.execute('UPDATE users SET department_id = 3 WHERE id = 402')
        cur.execute('UPDATE group_operator_memberships SET end_date = %s WHERE operator_id = 402 AND end_date IS NULL',
                    (D0 - timedelta(days=1),))
        cur.execute('SELECT id, number FROM discipline_records ORDER BY id')
        numbers = dict(cur.fetchall())
        cur.execute('SELECT id FROM discipline_acts WHERE trigger_record_id = %s', (r3,))
        act_id = cur.fetchone()[0]
    return r1, r2, r3, act_id, numbers


def test_hidden_act_records_are_stubs_but_the_act_draft_lists_all(client, db):
    r1, r2, r3, act_id, numbers = _cross_department_act(db)
    assert client.get('/api/discipline/records/%d' % r1, headers=h(HEAD_SZOV)).status_code == 403
    item = client.get('/api/discipline/records/%d' % r3, headers=h(HEAD_SZOV)).get_json()
    records = item['act_full']['records']
    assert sorted(r['id'] for r in records if r.get('hidden')) == sorted([r1, r2])
    assert all(set(r) == {'id', 'number', 'occurred_at', 'hidden'} for r in records if r.get('hidden'))
    # Черновик акта — документ по всем нарушениям периода: в нём только номер,
    # дата и тип, без описаний, поэтому записи вне доступа в нём остаются.
    from docx import Document
    draft = client.get('/api/discipline/acts/%d/draft.docx' % act_id, headers=h(HEAD_SZOV))
    assert draft.status_code == 200
    document = Document(io.BytesIO(draft.data))
    text = '\n'.join([p.text for p in document.paragraphs]
                     + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert all(numbers[rid] in text for rid in (r1, r2, r3))


def test_card_does_not_show_other_departments_control(client, db):
    from tests.test_discipline_postgres import NOW as T0, create

    with db.cursor() as cur:
        create(cur, HEAD_SZOV, OP_A, T0 - timedelta(days=3))
        cur.execute('UPDATE users SET department_id = 2 WHERE id = %s', (OP_A,))
        cur.execute('UPDATE group_operator_memberships SET end_date = %s WHERE operator_id = %s',
                    (T0.date() - timedelta(days=1), OP_A))
        cur.execute('INSERT INTO group_operator_memberships (group_id, operator_id, start_date) VALUES (20, %s, %s)',
                    (OP_A, T0.date()))
    response = client.post('/api/discipline/controls', headers=h(200), json={
        'operator_id': OP_A, 'source': 'manual', 'reason': 'Решение главы ОП',
        'end_date': (T0.date() + timedelta(days=10)).isoformat()})
    assert response.status_code in (200, 201), response.get_json()
    card = client.get('/api/discipline/operators/%d' % OP_A, headers=h(HEAD_SZOV)).get_json()
    assert card['control'] is None and card['controls'] == []
    assert len(card['records']) == 1


def test_record_payload_brings_its_departments_dictionaries(client):
    now = datetime.now().replace(second=0, microsecond=0)
    rid = _create(client, SV_A, OP_A, now - timedelta(days=1)).get_json()['id']
    item = client.get('/api/discipline/records/%d' % rid, headers=h(HEAD_SZOV)).get_json()
    assert item['measures'] and {m['department_id'] for m in item['measures']} == {SZOV}
    assert item['types'] and {t['department_id'] for t in item['types']} == {SZOV}
    # Оператору справочники не отдаются.
    own = client.get('/api/discipline/records/%d' % rid, headers=h(OP_A)).get_json()
    assert 'measures' not in own and 'types' not in own


def test_deductions_by_approval_date_through_the_route(client, db):
    from tests.test_discipline_postgres import NOW as T0, TODAY as D0, create, ctx_of, measure_id
    from discipline import service as svc

    with db.cursor() as cur:
        for day in (20, 22, 24):
            create(cur, HEAD_SZOV, OP_A, datetime(2026, 9, day, 10, 0))
        fine = measure_id(cur, SZOV, 'Лишение переменной части (фиксированная сумма)')
        rid = create(cur, SV_A, OP_A, datetime(2026, 9, 28, 10, 0), measure_id=fine, measure_amount='15000')
        head = ctx_of(cur, HEAD_SZOV)
        svc.approve_measure(cur, head, svc.load_record(cur, head, rid, D0, for_update=True),
                            now=datetime(2026, 10, 5, 11, 0))
        control_id = svc.create_control(cur, head, {
            'operator_id': OP_B, 'source': 'manual', 'reason': 'Решение руководителя', 'end_date': D0.isoformat()},
            now=T0, today=D0)
        svc.decide_control(cur, head, control_id, {'decision': 'measure', 'measure_id': fine,
                                                   'measure_amount': '20000'}, now=datetime(2026, 10, 6, 9, 0),
                           today=D0)
    september = client.get('/api/discipline/reports/measures?date_from=2026-09-01&date_to=2026-09-30',
                           headers=h(HEAD_SZOV)).get_json()
    october = client.get('/api/discipline/reports/measures?date_from=2026-10-01&date_to=2026-10-31',
                         headers=h(HEAD_SZOV)).get_json()
    assert september['rows'] == []
    assert [(r['basis'].split()[0], r['measure_amount']) for r in october['rows']] == [
        ('Нарушение', 15000.0), ('Решение', 20000.0)]
    assert october['columns'][0]['key'] == 'approved_at'


def test_act_records_in_card_follow_time_not_number(client):
    """Запись задним числом встаёт в списке акта на своё место по времени — как
    в черновике DOCX, — а не в конец по номеру."""
    now = datetime.now().replace(second=0, microsecond=0)
    first = _create(client, SV_A, OP_A, now - timedelta(days=2)).get_json()['id']
    second = _create(client, SV_A, OP_A, now - timedelta(days=1)).get_json()['id']
    backdated = _create(client, SV_A, OP_A, now - timedelta(days=5)).get_json()['id']
    item = client.get('/api/discipline/records/%d' % second, headers=h(HEAD_SZOV)).get_json()
    assert [r['id'] for r in item['act_full']['records']] == [backdated, first, second]


def test_shift_and_timecode_are_times_only(client):
    """Смена и тайм-код — только время: набранное приводится к одному виду,
    свободный текст получает отказ с кодом поля, а не ложится в запись."""
    at = datetime.now().replace(second=0, microsecond=0) - timedelta(days=1)
    for extra, code in (({'shift_label': '1234'}, 'SHIFT_FORMAT'), ({'evidence_note': 'видео 10:15'}, 'EVIDENCE_NOTE')):
        refused = _create(client, SV_A, OP_A, at, **extra)
        assert (refused.status_code, refused.get_json()['code']) == (400, code)
    created = _create(client, SV_A, OP_A, at, shift_label='20:00-8:00', evidence_note='14:32-14:41; 15:05')
    assert created.status_code == 201, created.get_json()
    item = client.get('/api/discipline/records/%d' % created.get_json()['id'], headers=h(SV_A)).get_json()
    assert (item['shift_label'], item['evidence_note']) == ('20:00–08:00', '14:32–14:41, 15:05')


def test_timecode_alone_is_evidence(client, db):
    type_id = _type_id(client, SV_A, OP_A)   # справочник отдела заводится первым запросом
    with db.cursor() as cur:
        cur.execute('UPDATE discipline_types SET needs_evidence = TRUE WHERE id = %s', (type_id,))
        assert cur.rowcount == 1
    at = datetime.now().replace(second=0, microsecond=0) - timedelta(days=1)
    missing = _create(client, SV_A, OP_A, at)
    assert (missing.status_code, missing.get_json()['code']) == (400, 'EVIDENCE')
    assert _create(client, SV_A, OP_A, at, evidence_note='14:32').status_code == 201


def test_edit_keeps_old_free_text_but_checks_new(client, db):
    """Запись со сменой и тайм-кодом свободным текстом (до полей времени)
    правится, не трогая их; новое значение — только время."""
    at = datetime.now().replace(second=0, microsecond=0) - timedelta(days=1)
    record_id = _create(client, SV_A, OP_A, at).get_json()['id']
    with db.cursor() as cur:
        cur.execute("UPDATE discipline_records SET shift_label = 'Линия 2', evidence_note = 'видео 10:15'"
                    " WHERE id = %s", (record_id,))
    url = '/api/discipline/records/%d' % record_id
    kept = client.patch(url, headers=h(SV_A), json={
        'shift_label': 'Линия 2', 'evidence_note': 'видео 10:15',
        'description': 'Оператор спал на рабочем месте около двадцати минут, видно по камере зала.'})
    assert kept.status_code == 200, kept.get_json()
    for body, code in (({'shift_label': 'с утра'}, 'SHIFT_FORMAT'), ({'evidence_note': 'видео 11:00'}, 'EVIDENCE_NOTE')):
        refused = client.patch(url, headers=h(SV_A), json=body)
        assert (refused.status_code, refused.get_json()['code']) == (400, code)
    assert client.patch(url, headers=h(SV_A), json={'shift_label': '10:00-19:00', 'evidence_note': '10:15'}).status_code == 200
    item = client.get(url, headers=h(SV_A)).get_json()
    assert (item['shift_label'], item['evidence_note']) == ('10:00–19:00', '10:15')


def test_rollout_only_super_admins(client, db, monkeypatch):
    """Пока раздел открыт одним супер-админам (владелец, 20.10.2026): остальным —
    отказ на ручках и пусто в колоколе, а пробная запись на живого оператора не
    шлёт ни колокола, ни Telegram ни ему, ни его СВ, ни главе отдела."""
    from discipline import access
    from notifications import sources

    monkeypatch.setattr(access, 'SUPER_ADMIN_ONLY', True)
    for user in (ADMIN, HEAD_SZOV, SV_A, OP_A, TRAINEE_A):
        refused = client.get('/api/discipline/ping', headers=h(user))
        assert (refused.status_code, refused.get_json()['code']) == (403, 'DISCIPLINE_SECTION_CLOSED'), user
    assert client.get('/api/discipline/ping', headers=h(SUPER)).status_code == 200
    created = _create(client, SUPER, OP_A, datetime.now().replace(second=0, microsecond=0) - timedelta(days=1))
    assert created.status_code == 201, created.get_json()
    with db.cursor() as cur:
        cur.execute('SELECT COUNT(*) FROM discipline_notifications')
        assert cur.fetchone()[0] == 0
        assert sources.discipline(cur, {'user_id': SV_A}, 10) == (0, [])
    assert client.sent == []
    # Флаг снят — те же люди снова по правилам ТЗ.
    monkeypatch.setattr(access, 'SUPER_ADMIN_ONLY', False)
    assert client.get('/api/discipline/ping', headers=h(SV_A)).status_code == 200


def test_department_ladders_through_the_route(client, db):
    """Глава отдела называет лестницы и заводит свои; действует одна. Новая
    действующая — новая версия правил счёта, переименование — нет. СВ настроек
    не видит и не меняет."""
    url = '/api/discipline/settings'
    got = client.get(url + '?department_id=%d' % SZOV, headers=h(HEAD_SZOV)).get_json()
    assert [(l['name'], l['steps']) for l in got['settings']['ladders']] == [('Основная', list(catalog.LADDER_TZ))]
    assert 'presets' not in got['catalog'] and got['catalog']['ladders']['max'] == catalog.MAX_LADDERS
    assert client.get(url + '?department_id=%d' % SZOV, headers=h(SV_A)).status_code == 403

    ladders = got['settings']['ladders'] + [{'id': 'l7', 'name': 'Строгая', 'steps': ['explanation', 'act']}]
    saved = client.put(url, headers=h(HEAD_SZOV), json={'department_id': SZOV,
                                                        'settings': {'ladders': ladders, 'active_ladder': 'l7'}})
    assert saved.status_code == 200, saved.get_json()
    assert saved.get_json()['settings']['ladder'] == ['explanation', 'act']
    refused = client.put(url, headers=h(SV_A), json={'department_id': SZOV, 'settings': {'active_ladder': 'main'}})
    assert refused.status_code == 403
    duplicate = client.put(url, headers=h(HEAD_SZOV), json={'department_id': SZOV, 'settings': {
        'ladders': ladders + [{'id': 'l8', 'name': 'строгая', 'steps': ['act']}], 'active_ladder': 'l7'}})
    assert duplicate.status_code == 400 and 'уже есть' in duplicate.get_json()['error']

    renamed = [dict(l, name='Жёсткая') if l['id'] == 'l7' else l for l in ladders]
    assert client.put(url, headers=h(HEAD_SZOV), json={'department_id': SZOV, 'settings': {
        'ladders': renamed, 'active_ladder': 'l7'}}).status_code == 200
    with db.cursor() as cur:
        cur.execute('SELECT rules FROM discipline_counting_rules WHERE department_id = %s ORDER BY valid_from', (SZOV,))
        assert [row[0]['ladder'] for row in cur.fetchall()] == [list(catalog.LADDER_TZ), ['explanation', 'act']]
        cur.execute("SELECT changes FROM discipline_audit WHERE action = 'settings_update' ORDER BY id")
        texts = [report.changes_text(row[0]) for row in cur.fetchall()]
    assert texts[0].startswith('Действующая лестница: Основная → Строгая; Лестницы отдела: Основная: ')
    assert 'Жёсткая: объяснительная, объяснительная + акт' in texts[1]


def test_settings_saved_before_named_ladders_keep_their_ladder(client, db):
    with db.cursor() as cur:
        cur.execute("INSERT INTO discipline_settings (department_id, settings) VALUES (%s, %s::jsonb)",
                    (SZOV, '{"ladder": ["training", "explanation", "explanation", "act"], "period_days": 30}'))
    got = client.get('/api/discipline/settings?department_id=%d' % SZOV, headers=h(HEAD_SZOV)).get_json()
    assert got['settings']['ladders'] == [{'id': 'main', 'name': 'Основная', 'steps': list(catalog.LADDER_SAMPLE)}]
    assert got['settings']['ladder'] == list(catalog.LADDER_SAMPLE)



# ── Документы по шаблонам организации (владелец, 23.10.2026) ────────────────

def _docx_text(data):
    from docx import Document
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    for section in document.sections:
        parts += [p.text for p in section.footer.paragraphs]
    return '\n'.join(parts)


def test_explanation_blank_by_the_organization_template(client):
    now = datetime.now().replace(second=0, microsecond=0)
    rid = _create(client, SV_A, OP_A, now - timedelta(hours=3),
                  description='Спал на рабочем месте с 14:00 до 14:20, видно по камере у окна.').get_json()['id']
    number = client.get('/api/discipline/records/%d' % rid, headers=h(SV_A)).get_json()['number']
    url = '/api/discipline/records/%d/explanation.docx' % rid

    # Пока шапки нет — линия, вписать от руки.
    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    assert text.startswith('Заказчику: ______') and '\nот Оператора: Оператор Анна\n' in text

    # Шапка — настройка отдела записи: сервер приводит её к виду и не берёт лишнего.
    payload = client.get('/api/discipline/settings?department_id=%d' % SZOV, headers=h(HEAD_SZOV)).get_json()
    assert payload['settings']['explanation_addressee'] == ''
    assert payload['catalog']['addressee'] == {'max_lines': catalog.ADDRESSEE_MAX_LINES,
                                               'max_chars': catalog.ADDRESSEE_MAX_CHARS}
    saved = client.put('/api/discipline/settings', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'settings': {'explanation_addressee': ' Заказчику:  ИП «Пример»\n\n Примеровой А. Б. '}})
    assert saved.status_code == 200
    assert saved.get_json()['settings']['explanation_addressee'] == 'Заказчику: ИП «Пример»\nПримеровой А. Б.'
    too_long = client.put('/api/discipline/settings', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'settings': {'explanation_addressee': 'строка\n' * 5}})
    assert too_long.status_code == 400 and 'не больше 4 строк' in too_long.get_json()['error']
    # Шапку меняет руководитель отдела, СВ — нет.
    assert client.put('/api/discipline/settings', headers=h(SV_A), json={
        'department_id': SZOV, 'settings': {'explanation_addressee': 'Кому-то'}}).status_code == 403

    blank = client.get(url, headers=h(SV_A))
    assert blank.status_code == 200
    assert blank.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert 'attachment' in blank.headers['Content-Disposition']
    text = _docx_text(blank.data)
    assert text.startswith('Заказчику: ИП «Пример»\nПримеровой А. Б.\nот Оператора: Оператор Анна\n')
    assert 'Запись о нарушении %s от %s: Сон на рабочем месте' % (number, (now - timedelta(hours=3)).strftime(
        '%d.%m.%Y')) in text
    # Описание обстоятельств оператору не показывается (ТЗ п. 3) — и в бланке его нет.
    assert 'видно по камере у окна' not in text
    # Время файла — UTC (сервер живёт по Алматы, python-docx пишет дату как UTC).
    from docx import Document
    created = Document(io.BytesIO(blank.data)).core_properties.created.replace(tzinfo=None)
    assert abs(created - datetime.now(timezone.utc).replace(tzinfo=None)) < timedelta(minutes=2)

    # Оператор — свой бланк; чужой человек и СВ чужой группы — нет.
    assert client.get(url, headers=h(OP_A)).status_code == 200
    assert client.get(url, headers=h(OP_B)).status_code == 403
    assert client.get(url, headers=h(SV_B)).status_code == 403
    # Кто скачивал бланк — в журнале действий (персональные данные, ТЗ п. 14).
    audit = client.get('/api/discipline/audit', headers=h(HEAD_SZOV)).get_json()['items']
    rows = [row for row in audit if row['action'] == 'explanation_docx']
    assert len(rows) == 3 and {row['action_label'] for row in rows} == {'Выгружен бланк объяснительной'}
    assert {row['actor_name'] for row in rows} == {'Супервайзер Альфа', 'Оператор Анна'}
    assert any(row['action'] == 'settings_update' and 'Бланк объяснительной: кому' in row['changes_text']
               for row in audit)
    # В истории записи скачиваний нет — это чтение, оно только в журнале действий.
    history = client.get('/api/discipline/records/%d' % rid, headers=h(SV_A)).get_json()['history']
    assert 'explanation_docx' not in {row['action'] for row in history}

    # Запись закрыта — бланк больше не выдаётся.
    assert client.post('/api/discipline/records/%d/explanation' % rid, headers=h(SV_A),
                       data={'file': (io.BytesIO(PDF), 'объяснительная.pdf')},
                       content_type='multipart/form-data').status_code == 200
    assert client.post('/api/discipline/records/%d/close' % rid, headers=h(SV_A)).status_code == 200
    assert client.get(url, headers=h(SV_A)).status_code == 403


def test_act_draft_follows_the_organization_sample(client):
    now = datetime.now().replace(second=0, microsecond=0)
    descriptions = ('Первое: проспал начало смены и пришёл в 09:40.',
                    'Второе: ушёл со смены в 17:20 без согласования.',
                    'Третье: спал на рабочем месте с 14:00 до 14:20, видно по камере')
    ids = [_create(client, SV_A, OP_A, now - timedelta(days=days), description=text).get_json()['id']
           for days, text in zip((3, 2, 1), descriptions)]
    item = client.get('/api/discipline/records/%d' % ids[2], headers=h(HEAD_SZOV)).get_json()
    act = item['act_full']
    url = '/api/discipline/acts/%d/draft.docx' % act['id']
    control = next(c for c in client.get('/api/discipline/controls', headers=h(HEAD_SZOV)).get_json()['items']
                   if c['operator_id'] == OP_A)
    end = datetime.strptime(control['end_date'], '%Y-%m-%d').strftime('%d.%m.%Y')

    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    # Дата черновика — сегодня по Алматы (как у сервера), а не по часам машины с тестами.
    from discipline import service
    assert text.startswith('Акт № %s\nо совершении работником дисциплинарного нарушения\nг. Алматы\t%s г.\n' % (
        act['number'], service.today_almaty().strftime('%d.%m.%Y')))
    # Составитель — СВ, зафиксировавший нарушение-основание (как в образце организации).
    assert 'Мною, супервайзером СЗоВ Супервайзер Альфа, составлен настоящий акт о нижеследующем:' in text
    # Пол — из карточки (у Оператора Анны — женский).
    assert ('допустила нарушение трудовой дисциплины: сон на рабочем месте. %s.' % descriptions[2]) in text
    assert 'Это третье нарушение сотрудника за последние 30 дней:' in text
    numbers = [r['number'] for r in act['records']]
    assert all('(запись %s), объяснительная не предоставлена' % n in text for n in numbers)
    # Описания прочих записей периода в акт не идут.
    assert descriptions[0] not in text and descriptions[1] not in text
    assert 'принято решение %s и установить испытательный срок продолжительностью 30 дней (до %s)' % (
        '_' * 34, end) in text
    assert 'В случае повторного нарушения в течение 30 дней сотруднику' in text
    assert 'Супервайзер СЗоВ\nПодпись: _______\nФИО: Супервайзер Альфа' in text
    assert 'С актом ознакомлен сотрудник СЗоВ\nПодпись: _______\nФИО: Оператор Анна' in text
    # Руководитель — глава отдела акта, пока акт не утверждён.
    assert 'С актом ознакомлен руководитель СЗоВ\nПодпись: _______\nФИО: Глава Сзовова' in text

    # Глава поставил дату и утвердил решение, черновик скачивает сам — составитель
    # всё равно СВ: дата в акте не делает главу составителем, иначе одна фамилия
    # стояла бы в двух строках подписей из трёх.
    measures = client.get('/api/discipline/controls', headers=h(HEAD_SZOV)).get_json()['measures']
    reprimand = next(m['id'] for m in measures if m['title'] == 'Выговор')
    assert client.patch('/api/discipline/acts/%d' % act['id'], headers=h(HEAD_SZOV), json={
        'composed_on': now.strftime('%Y-%m-%d'), 'decision_measure_id': reprimand, 'approve': True}).status_code == 200
    text = _docx_text(client.get(url, headers=h(HEAD_SZOV)).data)
    assert 'Мною, супервайзером СЗоВ Супервайзер Альфа, составлен настоящий акт о нижеследующем:' in text
    assert 'принято решение применить меру «Выговор» и установить испытательный срок' in text
    assert 'Супервайзер СЗоВ\nПодпись: _______\nФИО: Супервайзер Альфа' in text
    assert 'С актом ознакомлен руководитель СЗоВ\nПодпись: _______\nФИО: Глава Сзовова' in text
    # Скачивания документов — в журнале действий, но не в истории записи.
    history = client.get('/api/discipline/records/%d' % ids[2], headers=h(HEAD_SZOV)).get_json()['history']
    assert 'act_docx' not in {row['action'] for row in history}
    assert any(row['action'] == 'act_docx' for row in
               client.get('/api/discipline/audit', headers=h(HEAD_SZOV)).get_json()['items'])

    # Свой шаблон отдела главнее стандартного акта.
    from docx import Document
    template = Document()
    template.add_paragraph('Акт {{НОМЕР_АКТА}}: {{ФИО}}, решение — {{РЕШЕНИЕ}}')
    buffer = io.BytesIO()
    template.save(buffer)
    assert client.post('/api/discipline/settings/template', headers=h(HEAD_SZOV),
                       data={'department_id': str(SZOV), 'file': (io.BytesIO(buffer.getvalue()), 'акт.docx')},
                       content_type='multipart/form-data').status_code == 200
    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    assert text.strip() == 'Акт %s: Оператор Анна, решение — Выговор' % act['number']


def test_act_draft_keeps_the_rules_of_its_day_and_the_approver(client):
    """После акта отдел сменил окно и срок контроля — акт пишет срок и окно
    своего дня. Утвердил не глава отдела — в строке руководителя утвердивший."""
    now = datetime.now().replace(second=0, microsecond=0)
    ids = [_create(client, SV_A, OP_A, now - timedelta(days=days)).get_json()['id'] for days in (3, 2, 1)]
    act = client.get('/api/discipline/records/%d' % ids[2], headers=h(HEAD_SZOV)).get_json()['act_full']
    assert client.put('/api/discipline/settings', headers=h(HEAD_SZOV), json={
        'department_id': SZOV, 'settings': {'period_kind': 'month', 'control_days': 45}}).status_code == 200
    measures = client.get('/api/discipline/controls', headers=h(SUPER)).get_json()['measures']
    reprimand = next(m['id'] for m in measures if m['title'] == 'Выговор' and m['department_id'] == SZOV)
    assert client.patch('/api/discipline/acts/%d' % act['id'], headers=h(SUPER), json={
        'composed_on': now.strftime('%Y-%m-%d'), 'decision_measure_id': reprimand, 'approve': True}).status_code == 200
    text = _docx_text(client.get('/api/discipline/acts/%d/draft.docx' % act['id'], headers=h(SV_A)).data)
    assert 'Это третье нарушение сотрудника за последние 30 дней:' in text
    assert 'испытательный срок продолжительностью 30 дней' in text and 'в течение 30 дней' in text
    assert 'Мною, супервайзером СЗоВ Супервайзер Альфа, составлен' in text
    assert 'С актом ознакомлен руководитель СЗоВ\nПодпись: _______\nФИО: Супер Админов' in text


def test_explanation_blank_is_addressed_by_the_department_of_the_record(client, db):
    from tests.test_discipline_postgres import NOW as T0, create
    from discipline import queries

    with db.cursor() as cur:
        rid = create(cur, SV_A, OP_A, T0 - timedelta(hours=2))
        for department, text in ((SZOV, 'Заказчику: адресат СЗоВ'), (2, 'Заказчику: адресат ОП')):
            queries.save_settings(cur, department, catalog.normalize_settings({'explanation_addressee': text}), SUPER)
        # Человека перевели в ОП: запись осталась записью СЗоВ.
        cur.execute('UPDATE users SET department_id = 2 WHERE id = %s', (OP_A,))
    text = _docx_text(client.get('/api/discipline/records/%d/explanation.docx' % rid, headers=h(HEAD_SZOV)).data)
    assert text.startswith('Заказчику: адресат СЗоВ\n') and 'адресат ОП' not in text


def test_act_draft_head_with_supervisor_role_writes_as_head(client):
    """Глава ОП с ролью СВ — руководитель своего отдела: «руководителем Отдела
    продаж», а не «супервайзером» (название без сокращения склоняется целиком)."""
    from tests.test_discipline_postgres import HEAD_OP, OP_OP

    now = datetime.now().replace(second=0, microsecond=0)
    ids = [_create(client, HEAD_OP, OP_OP, now - timedelta(days=days)).get_json()['id'] for days in (3, 2, 1)]
    act = client.get('/api/discipline/records/%d' % ids[2], headers=h(HEAD_OP)).get_json()['act_full']
    text = _docx_text(client.get('/api/discipline/acts/%d/draft.docx' % act['id'], headers=h(HEAD_OP)).data)
    assert 'Мною, руководителем Отдела продаж Глава Продаж, составлен настоящий акт о нижеследующем:' in text
    # Пол в карточке не указан — форма с «(а)».
    assert 'Сотрудник Оператор Продаж ' in text and ' допустил(а) нарушение трудовой дисциплины' in text
    assert 'Руководитель Отдела продаж\nПодпись: _______\nФИО: Глава Продаж' in text
    # Составитель он же руководитель отдела — его строка «ознакомлен» остаётся линией.
    assert ('С актом ознакомлен руководитель Отдела продаж\nПодпись: _______\nФИО: ____________________'
            in text)


def test_act_draft_refuses_when_its_basis_is_cancelled(client):
    """Запись-основание отменена — черновик не собирается: документ описал бы
    другое нарушение (п. 7.2: «проверить основание акта»)."""
    now = datetime.now().replace(second=0, microsecond=0)
    ids = [_create(client, SV_A, OP_A, now - timedelta(days=days)).get_json()['id'] for days in (3, 2, 1)]
    act = client.get('/api/discipline/records/%d' % ids[2], headers=h(HEAD_SZOV)).get_json()['act_full']
    assert client.post('/api/discipline/records/%d/cancel' % ids[2], headers=h(HEAD_SZOV),
                       json={'reason': 'Ошибочная запись: не тот оператор'}).status_code == 200
    response = client.get('/api/discipline/acts/%d/draft.docx' % act['id'], headers=h(HEAD_SZOV))
    assert response.status_code == 409 and response.get_json()['code'] == 'ACT_BASIS_CANCELLED'


def test_act_draft_period_list_follows_each_records_step(client, db):
    """Лестница с тренингом первой ступенью: у первой записи объяснительную не
    требовали — в списке нарушений акта у неё нет «объяснительная не
    предоставлена»; проведён тренинг — так и написано."""
    with db.cursor() as cur:
        use_sample_ladder(cur)
    now = datetime.now().replace(second=0, microsecond=0)
    ids = [_create(client, SV_A, OP_A, now - timedelta(days=days)).get_json()['id'] for days in (4, 3, 2, 1)]
    act = client.get('/api/discipline/records/%d' % ids[3], headers=h(HEAD_SZOV)).get_json()['act_full']
    numbers = {r['id']: r['number'] for r in act['records']}
    url = '/api/discipline/acts/%d/draft.docx' % act['id']
    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    assert 'Это четвёртое нарушение сотрудника за последние 30 дней:' in text
    assert '(запись %s);' % numbers[ids[0]] in text
    assert '(запись %s), объяснительная не предоставлена;' % numbers[ids[1]] in text
    assert client.post('/api/discipline/records/%d/training' % ids[0], headers=h(SV_A), json={
        'date': (now - timedelta(days=1)).strftime('%Y-%m-%d'), 'start': '10:00', 'end': '10:30'}).status_code == 200
    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    assert '(запись %s), проведён тренинг;' % numbers[ids[0]] in text


def test_act_draft_window_and_term_follow_department_settings(client, db):
    """Окно 45 дней и контроль 40 — в акте эти числа, а не 30 по умолчанию; у акта
    нет контроля — срок из настроек отдела, без даты окончания."""
    from discipline import queries

    with db.cursor() as cur:
        queries.save_settings(cur, SZOV, catalog.normalize_settings({'period_days': 45, 'control_days': 40}),
                              HEAD_SZOV)
    now = datetime.now().replace(second=0, microsecond=0)
    ids = [_create(client, SV_A, OP_A, now - timedelta(days=days)).get_json()['id'] for days in (40, 20, 1)]
    act = client.get('/api/discipline/records/%d' % ids[2], headers=h(HEAD_SZOV)).get_json()['act_full']
    url = '/api/discipline/acts/%d/draft.docx' % act['id']
    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    assert 'Это третье нарушение сотрудника за последние 45 дней:' in text
    assert 'испытательный срок продолжительностью 40 дней (до ' in text and 'в течение 40 дней' in text
    with db.cursor() as cur:
        cur.execute('DELETE FROM discipline_controls WHERE act_id = %s', (act['id'],))
    text = _docx_text(client.get(url, headers=h(SV_A)).data)
    assert 'испытательный срок продолжительностью 40 дней для контроля соблюдения' in text


def test_act_draft_term_runs_to_the_end_of_a_linked_probation(client):
    """Акт во время испытательного срока привязывается к нему: срок в акте — от
    дня нарушения до конца этого контроля, а не срок «после акта» из настроек и
    не от начала испытательного."""
    from discipline import docx_layout

    now = datetime.now().replace(second=0, microsecond=0)
    start, end = (now - timedelta(days=10)).date(), (now + timedelta(days=50)).date()
    assert client.post('/api/discipline/controls', headers=h(HEAD_SZOV), json={
        'operator_id': OP_A, 'source': 'probation', 'start_date': start.isoformat(),
        'end_date': end.isoformat()}).status_code in (200, 201)
    # Справочник мер отдела — из контекста оператора (он же заводит справочники по умолчанию).
    measures = client.get('/api/discipline/operator-context?operator_id=%d' % OP_A,
                          headers=h(SV_A)).get_json()['measures']
    oral = next(m['id'] for m in measures if m['title'] == 'Устное замечание')
    created = [_create(client, SV_A, OP_A, now - timedelta(days=days), measure_id=str(oral)) for days in (3, 2, 1)]
    assert all(r.status_code == 201 for r in created), [r.get_json() for r in created]
    act = client.get('/api/discipline/records/%d' % created[2].get_json()['id'],
                     headers=h(HEAD_SZOV)).get_json()['act_full']
    text = _docx_text(client.get('/api/discipline/acts/%d/draft.docx' % act['id'], headers=h(SV_A)).data)
    days = (end - (now - timedelta(days=1)).date()).days
    assert 'испытательный срок продолжительностью %d %s (до %s)' % (
        days, docx_layout.plural(days, 'день', 'дня', 'дней'), end.strftime('%d.%m.%Y')) in text


def test_rollout_named_operator_works_as_an_operator(client, db, monkeypatch):
    """Владелец, 23.10.2026: раздел открыт поимённо «как оператору», пока он закрыт
    остальным. Поимённый входит по обычным правилам (только свои записи),
    колокол ему приходит, а его СВ и главе — по-прежнему ничего."""
    from discipline import access
    from notifications import sources

    monkeypatch.setattr(access, 'SUPER_ADMIN_ONLY', True)
    monkeypatch.setattr(access, 'ROLLOUT_USER_IDS', frozenset({OP_A}))
    assert client.get('/api/discipline/ping', headers=h(OP_A)).status_code == 200
    for user in (SV_A, OP_B, HEAD_SZOV):
        refused = client.get('/api/discipline/ping', headers=h(user))
        assert (refused.status_code, refused.get_json()['code']) == (403, 'DISCIPLINE_SECTION_CLOSED'), user
    mine = _create(client, SUPER, OP_A, datetime.now().replace(second=0, microsecond=0) - timedelta(hours=2))
    other = _create(client, SUPER, OP_B, datetime.now().replace(second=0, microsecond=0) - timedelta(hours=1))
    assert (mine.status_code, other.status_code) == (201, 201)
    with db.cursor() as cur:
        cur.execute('SELECT DISTINCT user_id FROM discipline_notifications')
        assert {row[0] for row in cur.fetchall()} == {OP_A}
        total, rows = sources.discipline(cur, {'user_id': OP_A}, 10)
        assert total >= 1 and rows
        assert sources.discipline(cur, {'user_id': SV_A}, 10) == (0, [])
    # Записи — только свои.
    items = client.get('/api/discipline/records', headers=h(OP_A)).get_json()['items']
    assert [item['id'] for item in items] == [mine.get_json()['id']]
    assert client.get('/api/discipline/records/%d' % other.get_json()['id'], headers=h(OP_A)).status_code == 403
    blank = client.get('/api/discipline/records/%d/explanation.docx' % mine.get_json()['id'], headers=h(OP_A))
    assert blank.status_code == 200


def test_audit_pages_by_the_requested_size(client, db):
    """«Отчёты» листают журнал действий по 10 строк (владелец, 23.10.2026: «чтобы
    не уходило вниз»): сервер отдаёт страницу заданного размера в границах, без
    размера — прежние 50; страницы не пересекаются."""
    from discipline import queries, routes
    from tests.test_discipline_postgres import ctx_of

    with db.cursor() as cur:
        head = ctx_of(cur, HEAD_SZOV)
        for index in range(23):
            queries.audit(cur, head, 'export', 'journal', department_id=SZOV, changes={'rows': index})
    url = '/api/discipline/audit'
    first = client.get(url + '?page_size=10', headers=h(HEAD_SZOV)).get_json()
    second = client.get(url + '?page_size=10&page=2', headers=h(HEAD_SZOV)).get_json()
    third = client.get(url + '?page_size=10&page=3', headers=h(HEAD_SZOV)).get_json()
    assert (first['total'], first['page_size'], len(first['items'])) == (23, 10, 10)
    assert (len(second['items']), len(third['items'])) == (10, 3)
    ids = [row['id'] for page in (first, second, third) for row in page['items']]
    assert len(set(ids)) == 23
    sizes = {raw: client.get(url + '?page_size=%s' % raw, headers=h(HEAD_SZOV)).get_json()['page_size']
             for raw in ('1', '100000', 'abc', '')}
    assert sizes == {'1': routes.AUDIT_PAGE_MIN, '100000': routes.AUDIT_PAGE_MAX, 'abc': 50, '': 50}
    # Выгрузка в Excel — весь журнал, размер страницы на неё не влияет.
    xlsx = client.get(url + '?format=xlsx&page_size=10', headers=h(HEAD_SZOV))
    assert xlsx.status_code == 200 and zipfile.is_zipfile(io.BytesIO(xlsx.data))
    import openpyxl
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.data))['Журнал действий']
    assert sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if any(row)) == 23
