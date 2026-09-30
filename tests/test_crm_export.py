# -*- coding: utf-8 -*-
"""Выгрузка «Обращений» в Excel за период (владелец, 30.09.2026).

Что сторожится:

* кому выгрузка открыта — СВ, главе и админу, не оператору: в файле ИИН и
  телефоны всех водителей периода разом;
* период — как у интерфейса (exportPeriod.js): обе даты включительно, не
  больше года, «по» не раньше «с»;
* строка файла: тема по сценарию, город и парк из разных вопросов тематик,
  куда обращение ушло на самом деле, время до первого ответа;
* ИИН и телефон — текстом, иначе ИИН уезжает в экспоненту;
* SQL периода — по дате создания и в периметре зрителя.
"""

import re
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from crm import access, queries, report, routes

ROOT = Path(__file__).resolve().parents[1]


def ticket(**changes):
    item = {
        'id': 12, 'subject': 'Статус работы офиса · Алматы', 'body': 'Со слов водителя офис закрыт',
        'status': 'resolved', 'queue_title': 'Регионы', 'topic_title': None,
        'scenario_key': 'office_status', 'created_by_name': 'Оператор', 'department_name': 'СЗоВ',
        'client_name': 'Сериков', 'client_phone': '87011234567',
        'answers': {'office_city': 'Алматы', 'iin': '060606202020'},
        'tg_chat_title': 'iTaxi Вопросы/ответы', 'queue_chat_title': 'Регионы — старый',
        'delivery_status': 'sent', 'flags': ['mass_outage'],
        'created_at': '2026-09-29T10:00:00', 'first_reply_at': '2026-09-29T10:42:30',
        'resolved_at': '2026-09-29T11:00:00', 'resolved_by_name': 'Супервайзер',
    }
    item.update(changes)
    return item


def values(item):
    return dict(zip([key for _t, key, _w in report.COLUMNS], report.row_values(item)))


class AccessTest(unittest.TestCase):
    def test_export_is_for_supervisors_and_heads(self):
        self.assertTrue(access.can_export({'role': 'sv', 'department_code': 'szov'}))
        self.assertTrue(access.can_export({'role': 'operator', 'headed_department_ids': [1]}))
        self.assertTrue(access.can_export({'role': 'super_admin'}))
        self.assertFalse(access.can_export({'role': 'operator', 'department_code': 'szov'}))
        self.assertFalse(access.can_export({'role': 'trainer', 'department_code': 'szov'}))

    def test_capability_reaches_the_interface(self):
        self.assertIn('can_export', access.capabilities({'role': 'sv', 'department_code': 'szov'}))


class PeriodTest(unittest.TestCase):
    def test_period_rules(self):
        self.assertEqual(routes._period('2026-09-01', '2026-09-30'), ('2026-09-01', '2026-09-30', None))
        self.assertEqual(routes._period('2026-09-30', '2026-09-30')[2], None)
        self.assertIsNotNone(routes._period('', '2026-09-30')[2])
        self.assertIsNotNone(routes._period('2026-09-30', '2026-09-01')[2])
        self.assertIsNone(routes._period('2025-10-01', '2026-09-30')[2], 'год целиком')
        self.assertIsNotNone(routes._period('2024-01-01', '2026-09-30')[2])

    def test_same_limit_as_the_interface(self):
        source = (ROOT / 'src' / 'components' / 'crm' / 'exportPeriod.js').read_text(encoding='utf-8')
        found = re.search(r'EXPORT_MAX_DAYS = (\d+);', source)
        self.assertEqual(int(found.group(1)), routes.EXPORT_MAX_DAYS)


class RowTest(unittest.TestCase):
    def test_row_reads_like_the_section(self):
        row = values(ticket())
        self.assertEqual(row['topic'], 'Статус работы офиса')
        self.assertEqual(row['status_title'], 'Решено')
        self.assertEqual(row['city'], 'Алматы', 'город офиса — в общую колонку «Город»')
        self.assertEqual(row['iin'], '060606202020')
        self.assertEqual(row['group_title'], 'iTaxi Вопросы/ответы',
                         'куда ушло на самом деле, а не чат тематики')
        self.assertEqual(row['reply_minutes'], 42)
        self.assertEqual(row['delivery_title'], 'Доставлено')
        self.assertEqual(row['flags_title'], 'Возможный массовый сбой')

    def test_missing_parts_stay_empty(self):
        row = values(ticket(first_reply_at=None, tg_chat_title=None, answers={}, flags=[],
                            scenario_key=None, topic_title='Своя тема'))
        self.assertIsNone(row['reply_minutes'])
        self.assertEqual(row['group_title'], 'Регионы — старый')
        self.assertEqual((row['city'], row['iin'], row['park']), ('', '', ''))
        self.assertEqual(row['topic'], 'Своя тема')

    def test_cooperation_city_and_park(self):
        row = values(ticket(scenario_key='cooperation',
                            answers={'coop_city': 'Астана', 'coop_park': 'iTaxi',
                                     'coop_email': 'partner@company.kz'}))
        self.assertEqual((row['city'], row['park']), ('Астана', 'iTaxi'))


class WorkbookTest(unittest.TestCase):
    def build(self, items, **kwargs):
        return report.build_workbook(items, date_from='2026-09-01', date_to='2026-09-30',
                                     generated_by='Супервайзер',
                                     generated_at=datetime(2026, 9, 30, 12, 0), **kwargs)

    def test_context_first_and_numbers_as_text(self):
        book = load_workbook(self.build([ticket()]))
        self.assertEqual(book.sheetnames, ['Контекст', 'Обращения'])
        context = [[cell.value for cell in row] for row in book['Контекст'].iter_rows()]
        self.assertIn(['Период', '01.09.2026 — 30.09.2026 (по дате создания обращения)'], context)
        sheet = book['Обращения']
        header = [cell.value for cell in sheet[1]]
        data = {title: cell for title, cell in zip(header, sheet[2])}
        self.assertEqual(data['ИИН'].value, '060606202020')
        self.assertEqual(data['ИИН'].number_format, '@')
        self.assertEqual(data['Телефон'].number_format, '@')
        self.assertIsInstance(data['Создано'].value, datetime)

    def test_truncated_export_says_so(self):
        book = load_workbook(self.build([ticket()], truncated=True))
        labels = [row[0].value for row in book['Контекст'].iter_rows()]
        self.assertIn('Внимание', labels)

    def test_green_corner_is_suppressed_on_the_text_columns(self):
        seen = {}

        def patch(stream, sqref, sheet_path):
            seen.update(sqref=sqref, sheet_path=sheet_path)
            return stream

        self.build([ticket(), ticket(id=13)], text_warning_patch=patch)
        # Телефон и ИИН — соседние колонки J и K, строки данных 2–3.
        self.assertEqual(seen, {'sqref': 'J2:K3', 'sheet_path': 'xl/worksheets/sheet2.xml'})

    def test_empty_period_is_a_valid_file(self):
        stream = self.build([])
        self.assertTrue(zipfile.is_zipfile(stream))

    def test_file_name(self):
        self.assertEqual(report.filename('2026-09-01', '2026-09-30'),
                         'Обращения 01.09.2026–30.09.2026.xlsx')


class ExportQueryTest(unittest.TestCase):
    class Cursor:
        def __init__(self):
            self.sql = None
            self.params = None

        def execute(self, sql, params=None):
            self.sql, self.params = sql, params

        def fetchall(self):
            return []

    def test_period_by_creation_date_inside_the_viewer_perimeter(self):
        cursor = self.Cursor()
        items, truncated = queries.export_tickets(
            cursor, {'user_id': 5, 'role': 'sv', 'department_code': 'szov'},
            date_from='2026-09-01', date_to='2026-09-30')
        self.assertEqual((items, truncated), ([], False))
        self.assertIn('t.created_at >= %(date_from)s::date', cursor.sql)
        self.assertIn("t.created_at < (%(date_to)s::date + INTERVAL '1 day')", cursor.sql)
        self.assertEqual(cursor.params['limit'], queries.EXPORT_LIMIT + 1)

    def test_blueprint_receives_the_excel_helper(self):
        bot = (ROOT / 'bot_schedule2.py').read_text(encoding='utf-8')
        start = bot.index('app.register_blueprint(build_crm_blueprint(')
        block = bot[start:start + 700]
        self.assertIn('excel_text_warning=_excel_suppress_number_as_text_warning', block)


if __name__ == '__main__':
    unittest.main()
