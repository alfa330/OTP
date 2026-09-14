# -*- coding: utf-8 -*-
"""Календарь фиксированных платежей и выгрузка реестра «Оплаты счетов».

Календарь (дополнение Дмитриевой, п. 11): периодичность сдвигает срок без
потери числа месяца, заявка создаётся в начале периода, файл с перечнем
разбирается по бытовым заголовкам. Выгрузка (п. 8): все поля, БИН и номера —
текстом, суммы — числами, подписи совпадают с экраном.

Ни базы, ни сети: `generate` здесь не вызывается, проверяются чистые функции.
"""

import io
import re
import sys
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from payments import fixed, report, workflow  # noqa: E402

META_PATH = ROOT / 'src' / 'components' / 'payments' / 'paymentsMeta.js'


class PeriodTests(unittest.TestCase):
    def test_add_months_keeps_the_day_and_clamps(self):
        self.assertEqual(fixed.add_months(date(2026, 1, 31), 1), date(2026, 2, 28))
        self.assertEqual(fixed.add_months(date(2026, 11, 5), 3), date(2027, 2, 5))
        self.assertEqual(fixed.next_due(date(2026, 3, 10), 'quarterly'), date(2026, 6, 10))
        self.assertEqual(fixed.next_due(date(2026, 3, 10), 'yearly'), date(2027, 3, 10))
        self.assertEqual(fixed.next_due(date(2026, 3, 10), 'custom', 14), date(2026, 3, 24))
        with self.assertRaises(ValueError):
            fixed.next_due(date(2026, 3, 10), 'custom', 0)

    def test_request_is_generated_at_the_start_of_the_period(self):
        self.assertEqual(fixed.generate_on(date(2026, 9, 25), 'monthly'), date(2026, 9, 1))
        self.assertEqual(fixed.generate_on(date(2026, 9, 25), 'custom', lead_days=7), date(2026, 9, 18))

    def test_is_due_respects_activity_and_period_start(self):
        template = {'is_active': True, 'next_due_on': date(2026, 9, 25), 'periodicity': 'monthly',
                    'lead_days': 7, 'last_generated_on': None, 'open_request_id': None}
        self.assertTrue(fixed.is_due(template, date(2026, 9, 1)))
        self.assertFalse(fixed.is_due(template, date(2026, 8, 31)))
        self.assertFalse(fixed.is_due(dict(template, is_active=False), date(2026, 9, 1)))
        # Уже создали под этот срок и заявка ещё живая — второй раз не создаём.
        self.assertFalse(fixed.is_due(dict(template, last_generated_on=date(2026, 9, 1), open_request_id=5),
                                      date(2026, 9, 2)))


class ImportParsingTests(unittest.TestCase):
    def test_headers_are_detected_in_everyday_wording(self):
        mapping = fixed.detect_columns(('Наименование', 'Сумма', 'Периодичность', 'Дата оплаты',
                                        'Проект / филиал', 'Ответственный', 'Категория', 'Подкатегория',
                                        'Контрагент', 'Источник оплаты', 'Примечание'))
        self.assertEqual(set(mapping), {'name', 'amount', 'periodicity', 'due', 'project', 'responsible',
                                        'category', 'subcategory', 'counterparty', 'source', 'note'})

    def test_periodicity_words(self):
        self.assertEqual(fixed.parse_periodicity('ежемесячно'), ('monthly', None))
        self.assertEqual(fixed.parse_periodicity('Раз в квартал'), ('quarterly', None))
        self.assertEqual(fixed.parse_periodicity('годовой'), ('yearly', None))
        self.assertEqual(fixed.parse_periodicity('каждые 10 дней'), ('custom', 10))
        self.assertEqual(fixed.parse_periodicity(''), (None, None))

    def test_due_day_number_rolls_to_the_next_month_when_passed(self):
        today = date(2026, 9, 14)
        self.assertEqual(fixed.parse_due(20, today), date(2026, 9, 20))
        self.assertEqual(fixed.parse_due('5', today), date(2026, 10, 5))
        self.assertEqual(fixed.parse_due('05.11.2026', today), date(2026, 11, 5))
        self.assertEqual(fixed.parse_due('3-го числа', today), date(2026, 10, 3))
        self.assertIsNone(fixed.parse_due('когда-нибудь', today))

    def test_workbook_is_parsed_with_errors_per_row(self):
        workbook = Workbook()
        sheet = workbook.active
        for row in fixed.template_sample_rows():
            sheet.append(list(row))
        sheet.append(['Без суммы', None, 'ежемесячно', '1', None, None, None, None, None, None, None])
        stream = io.BytesIO()
        workbook.save(stream)
        rows = fixed.parse_workbook(stream.getvalue(), today=date(2026, 9, 14))
        self.assertEqual(len(rows), 2)
        good, bad = rows
        self.assertEqual(good['name'], 'Аренда офиса Алматы')
        self.assertEqual(good['amount'], Decimal('450000'))
        self.assertEqual(good['periodicity'], 'monthly')
        self.assertEqual(good['next_due_on'], date(2026, 10, 5))
        self.assertEqual(good['payment_source'], 'too')
        self.assertEqual(good['errors'], [])
        self.assertIn('сумма не разобрана', bad['errors'])

    def test_missing_headers_raise_a_readable_error(self):
        workbook = Workbook()
        workbook.active.append(['Что-то', 'Другое'])
        stream = io.BytesIO()
        workbook.save(stream)
        with self.assertRaises(ValueError):
            fixed.parse_workbook(stream.getvalue())


def request_row(**overrides):
    row = {
        'id': 17, 'created_at': '2026-09-14T10:00:00', 'status': 'active', 'state': 'active', 'current_step': 7,
        'current_assignee_name': 'Ядигаров Руслан', 'initiator_name': 'Ядигаров Руслан', 'manager_name': None,
        'department_name': 'СЗоВ', 'project_name': 'Проект А', 'branch': 'Алматы', 'expense_name': 'Бумага А4',
        'amount': Decimal('2950'), 'refund_amount': Decimal('450'), 'refund_on': date(2026, 9, 20),
        'category_name': 'Канцелярия', 'subcategory_name': None, 'counterparty_name': 'ТОО «Пример»',
        'counterparty_bin': '000000000123', 'legal_entity_name': 'ТОО «Наше»', 'contract_number': '0025',
        'payment_source': 'too', 'payment_type': 'one_time', 'card_number': '4400000000001234',
        'payment_period': 'сентябрь 2026', 'due_on': date(2026, 9, 30), 'paid_on': None, 'paid_amount': None,
        'invoice_number': '00012', 'invoice_date': date(2026, 9, 10), 'invoice_requisites': 'IBAN…',
        'invoice_description': '=что-то с равенства', 'previous_payment_note': None,
        'route_basis': {'approver_name': 'Директор по развитию', 'order_number': '15'},
        'approval_order_number': '15', 'needs_power_of_attorney': True, 'needs_payment_order': False,
        'notes': None, 'rejected_reason': None, 'closed_at': None,
    }
    row.update(overrides)
    return row


class ReportTests(unittest.TestCase):
    def build(self, **overrides):
        rows = [request_row(**overrides)]
        items = {17: [{'name': 'Бумага А4', 'quantity': Decimal('1'), 'unit': 'ед.', 'unit_price': Decimal('2500')},
                      {'name': 'Карандаш', 'quantity': Decimal('3'), 'unit': 'шт', 'unit_price': Decimal('150')}]}
        return report.build_workbook(rows, items, filters_text='тест', total=1)

    def test_context_sheet_comes_first_and_all_columns_are_present(self):
        book = load_workbook(self.build())
        self.assertEqual(book.sheetnames, ['Контекст', 'Заявки'])
        sheet = book['Заявки']
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, [column[0] for column in report.COLUMNS])
        self.assertIn('Номер карты', headers)
        self.assertIn('Приказ', headers)

    def test_numbers_are_text_where_leading_zeros_matter(self):
        sheet = load_workbook(self.build())['Заявки']
        headers = [cell.value for cell in sheet[1]]
        row = {header: sheet.cell(row=2, column=index + 1) for index, header in enumerate(headers)}
        self.assertEqual(row['БИН контрагента'].value, '000000000123')
        self.assertEqual(row['БИН контрагента'].number_format, '@')
        self.assertEqual(row['Договор'].value, '0025')
        self.assertEqual(row['Номер счёта'].value, '00012')
        self.assertEqual(row['Номер карты'].value, '4400000000001234')
        self.assertEqual(row['Сумма расхода'].value, 2950.0)
        self.assertEqual(row['Итоговая сумма'].value, 2500.0, 'сумма минус возврат')
        self.assertEqual(row['Описание счёта'].value, '=что-то с равенства')
        self.assertNotEqual(row['Описание счёта'].data_type, 'f', 'текст с «=» не должен стать формулой')
        self.assertEqual(row['Срок оплаты'].value.date(), date(2026, 9, 30))

    def test_labels_match_the_frontend(self):
        source = META_PATH.read_text(encoding='utf-8')
        for code, label in report.SOURCE_LABELS.items():
            self.assertIn("%s: { label: '%s' }" % (code, label), source, code)
        for code, label in report.TYPE_LABELS.items():
            self.assertIn("%s: { label: '%s' }" % (code, label), source, code)
        for code, label in report.STATE_LABELS.items():
            self.assertRegex(source, r"%s: \{ label: '%s'" % (code, re.escape(label)), code)

    def test_items_and_approver_are_described(self):
        row = report.enrich(request_row(), [{'name': 'Бумага А4', 'quantity': Decimal('1'), 'unit': 'ед.',
                                              'unit_price': Decimal('2500')}])
        self.assertEqual(row['items_text'], 'Бумага А4 — 1 ед. × 2 500')
        self.assertIn('Директор по развитию', row['approver_label'])
        self.assertIn('№15', row['approver_label'])
        self.assertEqual(row['step_label'], 'Шаг 7 из 12 · %s' % workflow.step_title(7))
        self.assertEqual(row['needs_power_of_attorney_label'], 'Да')

    def test_text_warning_patch_targets_the_second_sheet(self):
        calls = []

        def patch(stream, sqref, sheet_path):
            calls.append((sqref, sheet_path))
            return stream

        rows = [request_row()]
        report.build_workbook(rows, {}, text_warning_patch=patch)
        self.assertEqual(len(calls), 1)
        sqref, sheet_path = calls[0]
        self.assertEqual(sheet_path, 'xl/worksheets/sheet2.xml')
        self.assertIn('2:', sqref)


if __name__ == '__main__':
    unittest.main()
