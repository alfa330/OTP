# -*- coding: utf-8 -*-
"""Статья вики файлом Word: дверь, периметр, журнал, перевод разметки, кнопка.

Решение владельца 08.10.2026: «статьи в вики сделай так, чтобы я мог их
скачивать… по роли админам и выше», «в ворд формате». Что здесь сторожится:

  1. ЛЕСТНИЦА. Файл отдаётся с должности администратора и выше, по РОЛИ: ни
     способность администратора вики, ни правило раздела круг не расширяют.
  2. ДВЕРЬ. Отказ по должности — 403 и один на все номера (существования
     статьи не раскрывает); невидимая статья — 404, как сама статья по слагу.
  3. КАРТИНКИ. В файл идут только те кадры, что читателю и так показывают
     (правило ручки /file/<id>): файл чужой статьи не скачивается из бакета.
     Чужой http-адрес сервер не ходит качать вовсе (SSRF).
  4. ЖУРНАЛ. Скачивание — дверь, через которую текст покидает портал; запись
     article.export обязательна и подписана в auditEvents.js.
  5. ПЕРЕВОД РАЗМЕТКИ. Белый список санитайзера (wiki/sanitize.py) — полный
     перечень того, что бывает в теле; на каждый пункт здесь есть проверка:
     заголовки, списки с перезапуском нумерации, таблицы с объединёнными
     ячейками, картинки (WebP → PNG/JPEG), ссылки, плашки, карточки,
     галереи, раскрывашки, тренажёр.
  6. КНОПКА. Рисуется по признаку сервера can_download, файл приходит blob'ом
     и сохраняется в этом же окне — без window.open и новых вкладок.
  7. ВИД (уточнение владельца 09.10.2026: «выглядит прям как иишный, сделай…
     как в статье… в стиле ios/macos»). Документ собран по правилам статьи,
     а не шаблона Word: системный шрифт с подсказкой подстановки, палитра и
     кегли статьи, тема без синего Office, бейджи и оглавление как в шапке
     статьи, блоки один в один (значки плашек, сетка карточек с зазорами,
     показатели, чипы, галочки, шаги с номерами), явная ширина таблиц,
     скруглённые картинки, подвал с номерами страниц. Порядок детей в
     w:pPr/w:rPr/w:tcPr/w:tblPr/w:trPr — по схеме, иначе Word объявляет
     файл повреждённым.

Часть проверок читает фронт ТЕКСТОМ — так в репозитории сторожат решения по
интерфейсу (см. tests/test_wiki_copy_protection.py).
"""

import base64
import io
import re
import sys
import unittest
import zipfile
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    from flask import Flask
except ImportError:  # pragma: no cover
    Flask = None

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.run import Run
except ImportError:  # pragma: no cover
    Document = None
    qn = None
    Run = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

from wiki import access as wiki_access  # noqa: E402
from wiki import article_access as wiki_article_access  # noqa: E402
from wiki import articles as wiki_articles  # noqa: E402
from wiki import docx_export  # noqa: E402
from wiki import perimeter as wiki_perimeter  # noqa: E402
from wiki import queries  # noqa: E402
from wiki import structure as wiki_structure  # noqa: E402
from wiki.access import collect_subjects  # noqa: E402
from wiki.routes import build_wiki_blueprint  # noqa: E402

SRC = ROOT / 'src' / 'components' / 'wiki'
ARTICLE_JSX = (SRC / 'WikiArticle.jsx').read_text(encoding='utf-8')
DOWNLOAD_JS = (SRC / 'articleDownload.js').read_text(encoding='utf-8')
AUDIT_JS = (SRC / 'auditEvents.js').read_text(encoding='utf-8')
REQUIREMENTS = (ROOT / 'requirements.txt').read_text(encoding='utf-8')

FILE_A = '11111111-1111-1111-1111-111111111111'   # картинка своей статьи
FILE_B = '22222222-2222-2222-2222-222222222222'   # картинка ЧУЖОЙ статьи
FILE_C = '33333333-3333-3333-3333-333333333333'   # не привязана, грузил сам


def strip_comments(source):
    """Исходник без комментариев: проверять КОД по тексту с комментариями нельзя."""
    source = re.sub(r'/\*.*?\*/', '', source, flags=re.S)
    source = re.sub(r'(^|[^:])//[^\n]*', r'\1', source)
    return source


def image_bytes(fmt='WEBP', size=(400, 300), alpha=False, frames=1):
    """Кадр для проверок. alpha=True — с НАСТОЯЩЕЙ прозрачностью: полностью
    непрозрачный RGBA-кадр кодек WebP честно сохраняет без альфа-канала, и
    Pillow открывает его как RGB."""
    if Image is None:
        return b''
    mode = 'RGBA' if alpha else 'RGB'
    color = (67, 56, 202, 0) if alpha else (67, 56, 202)
    img = Image.new(mode, size, color)
    if alpha:
        img.paste((200, 30, 30, 255), (0, 0, max(1, size[0] // 2), size[1]))
    out = io.BytesIO()
    if frames > 1:
        extra = [Image.new(mode, size, (200, 30, 30, 255) if alpha else (200, 30, 30))
                 for _ in range(frames - 1)]
        img.save(out, format=fmt, save_all=True, append_images=extra, duration=100)
    else:
        img.save(out, format=fmt)
    return out.getvalue()


def build(html, *, fetch=None, portal='', **fields):
    article = {
        'id': 1, 'slug': 'статья-1', 'title': 'Статья', 'summary': '', 'content': html,
        'status': 'published', 'author_name': '', 'updated_at': None,
    }
    article.update(fields)
    stream, stats = docx_export.build_document(
        article, fetch_image=fetch, portal_url=portal, section_paths=(), exported_at=None)
    return Document(stream), stats


def texts(doc):
    return [p.text for p in doc.paragraphs]


def styles(doc):
    return [p.style.name for p in doc.paragraphs]


HEADER_STYLES = ('Wiki Badges', 'Title', 'Subtitle', 'Wiki Meta', 'TOC Heading',
                 'toc 1', 'toc 2', 'toc 3', 'Wiki Spacer')


def body_paragraphs(doc):
    """Абзацы тела — после шапки (бейджи, название, аннотация, сводка) и оглавления."""
    out = list(doc.paragraphs)
    while out and out[0].style.name in HEADER_STYLES:
        out.pop(0)
    return out


def by_style(doc, name):
    return [p for p in doc.paragraphs if p.style.name == name]


def box_cell(doc, index=0):
    """Содержимое блока-рамки (вводка, цитата, код, раскрывашка): её единственная ячейка."""
    return doc.tables[index].cell(0, 0)


def note_cell(doc, index=0):
    """Ячейка текста плашки — вторая: первая держит значок тона."""
    return doc.tables[index].cell(0, 1)


def table_xml(table):
    return table._tbl.tblPr.xml


def numbering_xml(doc):
    return doc.part.numbering_part.numbering_definitions._numbering.xml


def all_runs(paragraph):
    """Все run'ы абзаца, включая те, что внутри гиперссылок (paragraph.runs их не отдаёт)."""
    return [Run(r, paragraph) for r in paragraph._p.xpath('.//w:r')]


def near(test, value, expected, delta=700):
    """Отступы python-docx хранит в twip и отдаёт в EMU с округлением (~45 EMU)."""
    test.assertAlmostEqual(int(value), int(expected), delta=delta)


def num_id(paragraph):
    pPr = paragraph._p.pPr
    if pPr is None or pPr.numPr is None or pPr.numPr.numId is None:
        return None
    return pPr.numPr.numId.val


def start_override(doc, number):
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    for num in numbering.num_lst:
        if num.numId == number:
            for override in num.lvlOverride_lst:
                if override.startOverride is not None:
                    return override.startOverride.val
    return None


def hyperlinks(doc, part=None):
    part = part or doc.part
    return sorted(rel.target_ref for rel in part.rels.values() if 'hyperlink' in rel.reltype)


def runs_xml(paragraph):
    return paragraph._p.xml


def cell_fill(cell):
    shd = cell._tc.tcPr.find(qn('w:shd')) if cell._tc.tcPr is not None else None
    return shd.get(qn('w:fill')) if shd is not None else None


def media_types(doc):
    """Типы картинок ТЕЛА документа. Эскиз docProps/thumbnail.jpeg из шаблона
    python-docx — не картинка статьи, поэтому берём только word/media."""
    return sorted(part.content_type for part in doc.part.package.iter_parts()
                  if str(part.partname).startswith('/word/media/'))


# ── 1. Лестница ──────────────────────────────────────────────────────────────

class ExportLadderTests(unittest.TestCase):
    def test_admin_and_above_only(self):
        for role in ('admin', 'super_admin', 'superadmin', 'Super Admin', ' ADMIN '):
            self.assertTrue(wiki_access.may_export(role), role)
        for role in ('operator', 'trainee', 'trainer', 'sv', 'supervisor', 'hr_manager',
                     'accounting_manager', 'marketing_manager', '', None, 'wiki_admin'):
            self.assertFalse(wiki_access.may_export(role), repr(role))

    def test_threshold_is_the_admin_level(self):
        """«Админам и выше» — порог ровно на admin, а не «выше супервайзера»."""
        self.assertEqual(wiki_access.EXPORT_MIN_LEVEL, wiki_access.ROLE_LEVELS['admin'])
        self.assertGreater(wiki_access.EXPORT_MIN_LEVEL, wiki_access.ROLE_LEVELS['sv'])


# ── 2. Имя файла, адрес портала, дата ────────────────────────────────────────

class NamingTests(unittest.TestCase):
    def test_file_name_keeps_cyrillic_and_drops_forbidden_characters(self):
        cases = {
            'Регламент выплат': 'Регламент выплат.docx',
            'Тарифы: Алматы / Астана?': 'Тарифы Алматы Астана.docx',
            '  Много   пробелов  ': 'Много пробелов.docx',
            'Точки в конце...': 'Точки в конце.docx',
            '': 'Статья.docx',
            None: 'Статья.docx',
            '<>:"|?*\\/': 'Статья.docx',
            'a\tb\nc': 'a b c.docx',
        }
        for title, expected in cases.items():
            self.assertEqual(docx_export.file_name(title), expected, repr(title))

    def test_file_name_is_capped(self):
        name = docx_export.file_name('ы' * 300)
        self.assertEqual(name, 'ы' * docx_export.FILENAME_MAX + '.docx')

    def test_portal_address_accepts_only_plain_http_origin(self):
        ok = 'https://alfa330.github.io/OTP/'
        self.assertEqual(docx_export.portal_address(ok), ok)
        self.assertEqual(docx_export.portal_address('http://localhost:5173/'), 'http://localhost:5173/')
        for bad in ('javascript:alert(1)', 'https://x.kz/?view=wiki', 'https://x.kz/#a',
                    'ftp://x.kz/', '', None, 'https://' + 'a' * 300):
            self.assertEqual(docx_export.portal_address(bad), '', repr(bad))

    def test_date_words(self):
        self.assertEqual(docx_export.date_words(datetime(2026, 10, 8, 12, 0)), '8 октября 2026')
        self.assertEqual(docx_export.date_words('2026-03-01T09:00:00'), '1 марта 2026')
        self.assertEqual(docx_export.date_words('не дата'), '')
        self.assertEqual(docx_export.date_words(None), '')


# ── 3. Перевод разметки ──────────────────────────────────────────────────────

@unittest.skipIf(Document is None, 'python-docx не установлен')
class DocumentShellTests(unittest.TestCase):
    def test_title_summary_meta_and_footer(self):
        article = {
            'id': 1, 'slug': 'тарифы', 'title': 'Тарифы', 'summary': 'Коротко',
            'content': '<p>Текст</p>', 'status': 'draft', 'author_name': 'Иванов Иван',
            'updated_at': datetime(2026, 10, 7, 15, 30),
        }
        stream, _ = docx_export.build_document(
            article, section_paths=[['Тез КЦ', 'Супервайзер'], ['Таксопарки']],
            portal_url='https://alfa330.github.io/OTP/', exported_at=datetime(2026, 10, 8))
        doc = Document(stream)
        # Шапка как на сайте: бейджи, название, аннотация, тихая строка-сводка.
        self.assertEqual(styles(doc)[:4], ['Wiki Badges', 'Title', 'Subtitle', 'Wiki Meta'])
        self.assertEqual(texts(doc)[0].strip('\xa0 '), 'Черновик')
        self.assertEqual(texts(doc)[1:3], ['Тарифы', 'Коротко'])
        meta = texts(doc)[3]
        self.assertNotIn('Черновик', meta, 'статус — бейдж, а не слово в сводке')
        self.assertIn('Разделы: Тез КЦ › Супервайзер; Таксопарки', meta)
        self.assertIn('Обновлено 7 октября 2026', meta)
        self.assertIn('Автор: Иванов Иван', meta)
        self.assertIn('<w:bottom', by_style(doc, 'Wiki Meta')[0]._p.pPr.xml, 'линия под шапкой')
        footer = doc.sections[0].footer
        self.assertIn('«Тарифы» · вики iCORE · выгружено 8 октября 2026', footer.paragraphs[0].text)
        self.assertIn('\tСтр. ', footer.paragraphs[0].text)
        footer_xml = footer.paragraphs[0]._p.xml
        # Номера страниц — поля с подставленным значением: видны и там, где поля не пересчитывают.
        self.assertRegex(footer_xml, re.compile(r'<w:fldSimple w:instr="PAGE">.*?<w:t>1</w:t>', re.S))
        self.assertRegex(footer_xml, re.compile(r'<w:fldSimple w:instr="NUMPAGES">.*?<w:t>1</w:t>', re.S))
        self.assertEqual(hyperlinks(doc, footer.part),
                         ['https://alfa330.github.io/OTP/?view=wiki&article=%D1%82%D0%B0%D1%80%D0%B8%D1%84%D1%8B'])
        self.assertEqual(doc.core_properties.title, 'Тарифы')
        self.assertEqual(doc.core_properties.author, 'Иванов Иван')
        self.assertEqual(doc.core_properties.subject, 'Коротко')
        self.assertEqual(doc.core_properties.created.replace(tzinfo=None), datetime(2026, 10, 8))

    def test_badges_show_type_status_and_tags_like_the_site(self):
        doc, _ = build('<p>Текст</p>', status='draft', article_type='regulation', tags=['выплаты', ' '])
        badges = by_style(doc, 'Wiki Badges')
        self.assertEqual(len(badges), 1)
        words = [run.text.strip('\xa0') for run in badges[0].runs if run.text.strip('\xa0 ')]
        self.assertEqual(words, ['Черновик', 'Регламент', 'выплаты'])
        self.assertIn('<w:bdr', badges[0]._p.xml, 'бейдж — ярлык в рамке, как IosBadge')
        # Обычная опубликованная статья без ярлыков — без строки бейджей, как на сайте.
        doc, _ = build('<p>Текст</p>', article_type='general')
        self.assertEqual(by_style(doc, 'Wiki Badges'), [])
        self.assertEqual(doc.paragraphs[0].style.name, 'Title')

    def test_published_article_without_places_has_no_status_word(self):
        doc, _ = build('<p>Текст</p>', updated_at=datetime(2026, 1, 2))
        self.assertEqual(texts(doc)[1], 'Обновлено 2 января 2026')
        self.assertNotIn('Черновик', texts(doc)[1])

    def test_no_portal_means_no_links_to_the_portal(self):
        doc, _ = build('<p><a href="?view=wiki&article=x">внутренняя</a></p>')
        self.assertEqual(hyperlinks(doc), [])
        self.assertEqual(hyperlinks(doc, doc.sections[0].footer.part), [])
        self.assertIn('внутренняя', texts(doc)[-1])

    def test_page_is_a4_with_two_centimetre_margins(self):
        doc, _ = build('<p>Текст</p>')
        section = doc.sections[0]
        self.assertEqual(round(section.page_width.cm, 1), 21.0)
        self.assertEqual(round(section.page_height.cm, 1), 29.7)
        self.assertEqual(round(section.left_margin.cm, 1), 2.0)

    def test_trailing_empty_paragraphs_are_dropped_and_doubles_collapsed(self):
        doc, _ = build('<p>Один</p><p></p><p></p><p></p><p>Два</p><p></p><p></p>')
        body = [p.text for p in body_paragraphs(doc)]
        self.assertEqual(body, ['Один', '', 'Два'])

    def test_file_is_a_valid_package(self):
        doc, _ = build('<h2>Заголовок</h2><p>Текст</p>')
        stream = io.BytesIO()
        doc.save(stream)
        with zipfile.ZipFile(stream) as package:
            self.assertIsNone(package.testzip())
            names = package.namelist()
        self.assertIn('word/document.xml', names)
        self.assertIn('word/numbering.xml', names)


@unittest.skipIf(Document is None, 'python-docx не установлен')
class TextMarkupTests(unittest.TestCase):
    def test_headings_keep_their_levels(self):
        doc, _ = build('<h1>Один</h1><h2>Два</h2><h3>Три</h3><h4>Четыре</h4><h5>Пять</h5><h6>Шесть</h6>')
        self.assertEqual([p.style.name for p in body_paragraphs(doc)],
                         ['Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Heading 5', 'Heading 6'])

    def test_inline_marks_land_on_runs(self):
        doc, _ = build('<p><strong>ж</strong><em>к</em><u>п</u><s>з</s>H<sub>2</sub>O x<sup>2</sup>'
                       '<code>код</code><mark data-color="#fde68a">цвет</mark><mark>жёлтый</mark>'
                       '<span style="color: rgb(220, 38, 38); font-size: 14px; font-family: Arial">стиль</span></p>')
        paragraph = body_paragraphs(doc)[0]
        by_text = {run.text: run for run in paragraph.runs}
        self.assertTrue(by_text['ж'].font.bold)
        self.assertTrue(by_text['к'].font.italic)
        self.assertTrue(by_text['п'].font.underline)
        self.assertTrue(by_text['з'].font.strike)
        twos = [run for run in paragraph.runs if run.text == '2']
        self.assertEqual([bool(run.font.subscript) for run in twos], [True, False])
        self.assertEqual([bool(run.font.superscript) for run in twos], [False, True])
        self.assertEqual(by_text['код'].font.name, docx_export.MONO_FONT)
        self.assertIn('w:fill="FDE68A"', by_text['цвет']._r.xml)
        # Выделение без цвета — янтарь статьи (.wiki-prose mark), а не кислотный маркер Word.
        self.assertIn('w:fill="%s"' % docx_export.MARK_FILL, by_text['жёлтый']._r.xml)
        self.assertIsNone(by_text['жёлтый'].font.highlight_color)
        self.assertEqual(str(by_text['стиль'].font.color.rgb), 'DC2626')
        self.assertEqual(by_text['стиль'].font.size.pt, 10.5)
        self.assertEqual(by_text['стиль'].font.name, 'Arial')

    def test_links_become_hyperlinks_and_internal_ones_get_the_portal(self):
        html = ('<p><a href="https://example.com/a?b=1">внешняя</a> '
                '<a href="?view=wiki&article=тарифы-2026">внутренняя</a> '
                '<a href="mailto:a@b.kz">почта</a> <a href="#x">якорь</a> '
                '<a href="/api/wiki/file/' + FILE_A + '">файл</a></p>')
        doc, _ = build(html, portal='https://alfa330.github.io/OTP/')
        self.assertEqual(hyperlinks(doc), [
            'https://alfa330.github.io/OTP/?view=wiki&article=%D1%82%D0%B0%D1%80%D0%B8%D1%84%D1%8B-2026',
            'https://example.com/a?b=1',
            'mailto:a@b.kz',
        ])
        paragraph = body_paragraphs(doc)[0]
        self.assertEqual(paragraph.text, 'внешняя внутренняя почта якорь файл')
        self.assertEqual(len(paragraph._p.findall(qn('w:hyperlink'))), 3)

    def test_whitespace_collapses_like_a_browser(self):
        doc, _ = build('<p>   Один\n   два <strong> три </strong> четыре   <br>  пять  </p>')
        paragraph = body_paragraphs(doc)[0]
        self.assertEqual(paragraph.text, 'Один два три четыре\nпять')
        self.assertEqual(len(paragraph._p.findall('.//' + qn('w:br'))), 1)

    def test_text_alignment_is_kept(self):
        doc, _ = build('<p style="text-align: center">центр</p><p style="text-align: right">право</p>')
        first, second = body_paragraphs(doc)[:2]
        self.assertEqual(str(first.alignment), 'CENTER (1)')
        self.assertEqual(str(second.alignment), 'RIGHT (2)')

    def test_blockquote_pre_and_rule(self):
        """Цитата и код — блоки с подложкой (ячейки: заливку абзаца предпросмотр
        macOS не рисует), разделитель — тонкая линия."""
        doc, _ = build('<blockquote><p>Цитата</p></blockquote><pre><code>a = 1\nb = 2</code></pre><hr><p>дальше</p>')
        quote_cell, code_cell = box_cell(doc, 0), box_cell(doc, 1)
        quote = quote_cell.paragraphs[0]
        self.assertEqual(quote.text, 'Цитата')
        self.assertEqual(quote.style.name, 'Quote')
        self.assertTrue(quote.runs[0].font.italic)
        self.assertEqual(cell_fill(quote_cell), docx_export.LINE_SOFT)
        self.assertIn('<w:left w:val="single" w:sz="24" w:space="0" w:color="%s"' % docx_export.ACCENT_LINE,
                      quote_cell._tc.xml, 'акцентная грань слева, как у цитаты на сайте')
        code = code_cell.paragraphs[0]
        self.assertEqual(code.text, 'a = 1\nb = 2')
        self.assertEqual(code.style.name, 'Wiki Code')
        self.assertEqual(code.runs[0].font.name, docx_export.MONO_FONT)
        self.assertEqual(str(code.runs[0].font.color.rgb), docx_export.CODE_TEXT)
        self.assertEqual(cell_fill(code_cell), docx_export.CODE_FILL)
        rules = [p for p in body_paragraphs(doc) if p._p.pPr is not None and '<w:bottom' in p._p.pPr.xml]
        self.assertEqual(len(rules), 1)
        self.assertEqual(body_paragraphs(doc)[-1].text, 'дальше')

    def test_trailing_rule_is_dropped_with_the_empty_tail(self):
        doc, _ = build('<p>Текст</p><hr>')
        self.assertEqual([p.text for p in body_paragraphs(doc)], ['Текст'])

    def test_stray_comments_and_scripts_do_not_leak(self):
        doc, _ = build('<p>Текст<!-- тайна --></p>')
        self.assertEqual(body_paragraphs(doc)[0].text, 'Текст')


@unittest.skipIf(Document is None, 'python-docx не установлен')
class ListTests(unittest.TestCase):
    def test_bullets_nest_up_to_three_levels(self):
        doc, _ = build('<ul><li>а<ul><li>б<ul><li>в<ul><li>г</li></ul></li></ul></li></ul></li></ul>')
        self.assertEqual([p.style.name for p in body_paragraphs(doc)],
                         ['List Bullet', 'List Bullet 2', 'List Bullet 3', 'List Bullet 3'])

    def test_every_ordered_list_restarts_and_honours_start(self):
        """Два списка подряд — счёт заново; start=3 — с тройки; вложенный — свой."""
        doc, _ = build('<ol start="3"><li>три</li><li>четыре<ol><li>вложенный</li></ol></li><li>пять</li></ol>'
                       '<ol><li>один</li><li>два</li></ol>')
        paragraphs = body_paragraphs(doc)
        ids = [num_id(p) for p in paragraphs]
        self.assertTrue(all(ids), 'у каждого пункта должна быть нумерация')
        first, nested, second = ids[0], ids[2], ids[4]
        self.assertEqual(ids, [first, first, nested, first, second, second])
        self.assertEqual(len({first, nested, second}), 3)
        self.assertEqual(start_override(doc, first), 3)
        self.assertEqual(start_override(doc, nested), 1)
        self.assertEqual(start_override(doc, second), 1)
        self.assertEqual([p.style.name for p in paragraphs],
                         ['List Number', 'List Number', 'List Number 2', 'List Number',
                          'List Number', 'List Number'])

    def test_extra_paragraphs_in_an_item_continue_the_item(self):
        doc, _ = build('<ul><li><p>Первый</p><p>Второй абзац</p></li></ul>')
        self.assertEqual([(p.style.name, p.text) for p in body_paragraphs(doc)],
                         [('List Bullet', 'Первый'), ('List Continue', 'Второй абзац')])

    def test_list_variants_look_like_the_site(self):
        """Шаги — номер белым на акцентной плашке, галочки и крестики — цветные
        знаки вместо точки, чипы — строка ярлыков, а не столбец пунктов."""
        doc, _ = build('<ol data-variant="steps"><li>шаг</li></ol><ul data-variant="checks"><li>да</li></ul>'
                       '<ul data-variant="crosses"><li>нет</li></ul><ul data-variant="chips"><li>чип</li><li>ещё</li></ul>')
        step, check, cross, chips = body_paragraphs(doc)[:4]
        self.assertEqual([p.style.name for p in (step, check, cross, chips)],
                         ['List Number', 'List Bullet', 'List Bullet', 'Wiki Chips'])
        numbering = numbering_xml(doc)
        self.assertIn('w:val="\xa0%1\xa0"', numbering)
        self.assertRegex(numbering, re.compile(r'w:val="\xa0%1\xa0".*?w:fill="' + docx_export.ACCENT + '"', re.S))
        self.assertIn('w:val="✓"', numbering)
        self.assertIn('w:val="✕"', numbering)
        self.assertRegex(numbering, re.compile(r'w:val="✓".*?w:color w:val="' + docx_export.CHECK_COLOR, re.S))
        self.assertRegex(numbering, re.compile(r'w:val="✕".*?w:color w:val="' + docx_export.CROSS_COLOR, re.S))
        self.assertNotEqual(num_id(check), num_id(cross))
        self.assertEqual(chips.text, '\xa0чип\xa0  \xa0ещё\xa0')
        self.assertIn('<w:bdr', chips._p.xml)
        self.assertIn('w:fill="%s"' % docx_export.SURFACE_ALT, chips._p.xml)
        self.assertIsNone(num_id(chips))

    def test_nested_list_in_a_step_hangs_from_the_step_text(self):
        """Вложенный список стоит от текста шага, а не от его номера; номер —
        в выносе на ширину плашки."""
        doc, _ = build('<ol data-variant="steps"><li><h4>Шаг</h4><p>пояснение</p><ul><li>вложенный</li></ul></li></ol>')
        head, note, nested = body_paragraphs(doc)[:3]
        self.assertEqual([p.text for p in (head, note, nested)], ['Шаг', 'пояснение', 'вложенный'])
        self.assertTrue(head.runs[0].font.bold)
        self.assertEqual(head.style.name, 'List Number')
        near(self, head.paragraph_format.left_indent, docx_export.STEP_INDENT)
        near(self, head.paragraph_format.first_line_indent, -docx_export.STEP_INDENT)
        self.assertEqual(note.style.name, 'List Continue')
        near(self, note.paragraph_format.left_indent, docx_export.STEP_INDENT)
        near(self, note.paragraph_format.first_line_indent, 0)
        self.assertEqual(nested.style.name, 'List Bullet 2')
        near(self, nested.paragraph_format.left_indent, docx_export.STEP_INDENT + docx_export.LIST_INDENT)
        near(self, nested.paragraph_format.first_line_indent, -docx_export.LIST_INDENT)


@unittest.skipIf(Document is None, 'python-docx не установлен')
class TableTests(unittest.TestCase):
    HTML = ('<table><colgroup><col style="width: 100px"><col style="width: 300px"></colgroup>'
            '<thead><tr><th>А</th><th>Б</th></tr></thead><tbody>'
            '<tr><td rowspan="2">слито вниз</td><td style="text-align: center; background-color: #fef2f2">центр</td></tr>'
            '<tr><td>низ</td></tr>'
            '<tr><td colspan="2"><strong>на две</strong> колонки</td></tr>'
            '<tr><td>список</td><td><ul><li>раз</li><li>два</li></ul></td></tr>'
            '</tbody></table>')

    def test_grid_merges_and_header(self):
        doc, _ = build(self.HTML)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (5, 2))
        self.assertEqual(table.style.name, 'Wiki Table')
        self.assertEqual(cell_fill(table.cell(0, 0)), docx_export.HEADER_FILL)
        head_run = table.cell(0, 0).paragraphs[0].runs[0]
        self.assertTrue(head_run.font.bold)
        self.assertTrue(head_run.font.all_caps, 'шапка капителью, как th на сайте')
        self.assertIn('<w:tblHeader', table.rows[0]._tr.xml, 'шапка повторяется на каждой странице')
        self.assertNotIn('<w:tblHeader', table.rows[1]._tr.xml)
        # Чередование строк тела: вторая строка тела (индекс 2) подкрашена, первая — нет.
        self.assertEqual(cell_fill(table.cell(2, 1)), docx_export.SURFACE_ALT)
        self.assertIsNone(cell_fill(table.cell(1, 0)))
        self.assertEqual(table.cell(1, 1).paragraphs[0].style.name, 'Wiki Table Text')
        self.assertEqual(table.cell(1, 0).text, table.cell(2, 0).text)
        self.assertIn('<w:vMerge', table.cell(1, 0)._tc.xml)
        self.assertIn('w:gridSpan', table.cell(3, 0)._tc.xml)
        self.assertEqual(table.cell(3, 0).text, 'на две колонки')
        self.assertEqual(cell_fill(table.cell(1, 1)), 'FEF2F2')
        self.assertEqual(str(table.cell(1, 1).paragraphs[0].alignment), 'CENTER (1)')
        self.assertEqual([p.style.name for p in table.cell(4, 1).paragraphs],
                         ['List Bullet', 'List Bullet'])

    def test_column_widths_follow_the_editor(self):
        doc, _ = build(self.HTML)
        table = doc.tables[0]
        total = table.columns[0].width + table.columns[1].width
        self.assertAlmostEqual(table.columns[0].width / total, 0.25, places=2)
        self.assertAlmostEqual(total / docx_export.Cm(17), 1.0, places=1)

    def test_colwidth_from_cells_when_colgroup_is_missing(self):
        doc, _ = build('<table><tr><td colwidth="100">а</td><td colwidth="100">б</td><td colwidth="200">в</td></tr></table>')
        table = doc.tables[0]
        widths = [column.width for column in table.columns]
        self.assertAlmostEqual(widths[2] / widths[0], 2.0, places=2)

    def test_cell_properties_are_in_schema_order(self):
        """tcBorders раньше shd, оба после vMerge/gridSpan — иначе Word
        объявляет файл повреждённым."""
        doc, _ = build('<div data-wiki-block="note" data-tone="warn"><p>текст</p></div>' + self.HTML)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    tcPr = cell._tc.tcPr
                    if tcPr is None:
                        continue
                    tags = [child.tag.split('}')[1] for child in tcPr]
                    order = ['tcW', 'gridSpan', 'vMerge', 'tcBorders', 'shd', 'vAlign']
                    seen = [tag for tag in tags if tag in order]
                    self.assertEqual(seen, sorted(seen, key=order.index), tags)

    def test_tables_do_not_glue_together(self):
        """Две таблицы подряд без абзаца между ними Word сливает в одну."""
        doc, _ = build('<table><tr><td>а</td></tr></table><table><tr><td>б</td></tr></table>')
        body = [child.tag.split('}')[1] for child in doc.element.body]
        first = body.index('tbl')
        self.assertEqual(body[first:first + 3], ['tbl', 'p', 'tbl'])


@unittest.skipIf(Document is None or Image is None, 'python-docx или Pillow не установлены')
class ImageTests(unittest.TestCase):
    def fetcher(self, calls=None):
        store = {
            FILE_A: (image_bytes('WEBP', (800, 400)), 'image/webp'),
            FILE_C: (image_bytes('PNG', (120, 120), alpha=True), 'image/png'),
        }

        def fetch(file_id):
            if calls is not None:
                calls.append(file_id)
            return store.get(file_id)
        return fetch

    def test_bucket_webp_is_embedded_as_png_or_jpeg(self):
        calls = []
        doc, stats = build('<img src="/api/wiki/file/%s" data-width="50" data-align="center">' % FILE_A,
                           fetch=self.fetcher(calls))
        self.assertEqual(calls, [FILE_A])
        self.assertEqual(stats, {'images': 1, 'missing': 0, 'external': 0})
        self.assertEqual(len(doc.inline_shapes), 1)
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.width / docx_export.Cm(17), 0.5, places=2)
        self.assertAlmostEqual(shape.width / shape.height, 2.0, places=1)
        self.assertEqual(media_types(doc), ['image/jpeg'])
        picture_paragraph = [p for p in doc.paragraphs if p._p.findall('.//' + qn('w:drawing'))][0]
        self.assertEqual(str(picture_paragraph.alignment), 'CENTER (1)')

    def test_transparent_webp_becomes_png(self):
        data = image_bytes('WEBP', (50, 50), alpha=True)
        doc, _ = build('<img src="/api/wiki/file/%s">' % FILE_A,
                       fetch=lambda _id: (data, 'image/webp'))
        self.assertEqual(media_types(doc), ['image/png'])

    def test_natural_size_is_capped_by_the_column(self):
        data = image_bytes('WEBP', (3000, 1000))
        doc, _ = build('<img src="/api/wiki/file/%s">' % FILE_A,
                       fetch=lambda _id: (data, 'image/webp'))
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.width / docx_export.Cm(17), 1.0, places=2)

    def test_tall_image_is_capped_by_height(self):
        data = image_bytes('PNG', (300, 3000))
        doc, _ = build('<img src="/api/wiki/file/%s" data-width="100">' % FILE_A,
                       fetch=lambda _id: (data, 'image/png'))
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.height / docx_export.Cm(docx_export.MAX_IMAGE_HEIGHT_CM), 1.0, places=2)
        self.assertLess(shape.width, docx_export.Cm(3))

    def test_legacy_pixel_width_and_data_uri(self):
        png = base64.b64encode(image_bytes('PNG', (200, 100))).decode()
        doc, stats = build('<p><img src="data:image/png;base64,%s" width="96"> подпись</p>' % png)
        self.assertEqual(stats['images'], 1)
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.width / docx_export.Cm(2.54), 1.0, places=2)
        self.assertEqual(body_paragraphs(doc)[0].text, ' подпись')

    def test_external_image_is_never_fetched_and_becomes_a_link(self):
        calls = []
        doc, stats = build('<img src="https://example.com/pic.png" alt="кадр">', fetch=self.fetcher(calls))
        self.assertEqual(calls, [])
        self.assertEqual(stats, {'images': 0, 'missing': 0, 'external': 1})
        self.assertEqual(hyperlinks(doc), ['https://example.com/pic.png'])
        self.assertIn('[изображение по ссылке]', texts(doc))

    def test_missing_and_failing_files_leave_a_note(self):
        def fetch(file_id):
            if file_id == FILE_B:
                raise RuntimeError('bucket down')
            return None
        doc, stats = build('<img src="/api/wiki/file/%s"><img src="/api/wiki/file/%s">' % (FILE_A, FILE_B),
                           fetch=fetch)
        self.assertEqual(stats, {'images': 0, 'missing': 2, 'external': 0})
        self.assertEqual([t for t in texts(doc) if 'недоступно' in t], ['[изображение недоступно]'] * 2)

    def test_same_file_is_fetched_once(self):
        calls = []
        build('<img src="/api/wiki/file/%s"><img src="/api/wiki/file/%s">' % (FILE_A, FILE_A.upper()),
              fetch=self.fetcher(calls))
        self.assertEqual(calls, [FILE_A])

    def test_oversized_frame_is_refused(self):
        self.assertIsNone(docx_export.prepare_image(image_bytes('PNG', (9000, 5000)), 'image/png'))

    def test_prepare_passes_small_png_through_untouched(self):
        data = image_bytes('PNG', (300, 200))
        self.assertEqual(docx_export.prepare_image(data, 'image/png'), (data, 300, 200))

    def test_animated_webp_takes_the_first_frame(self):
        data = image_bytes('WEBP', (60, 40), frames=3)
        prepared = docx_export.prepare_image(data, 'image/webp')
        self.assertIsNotNone(prepared)
        self.assertEqual(prepared[1:], (60, 40))

    def test_garbage_is_not_an_image(self):
        self.assertIsNone(docx_export.prepare_image(b'not an image', 'image/webp'))
        self.assertIsNone(docx_export.prepare_image(b'', 'image/webp'))

    def test_figure_caption_and_image_in_table_cell(self):
        doc, stats = build('<figure><img src="/api/wiki/file/%s" data-width="20"><figcaption>Подпись</figcaption></figure>'
                           '<table><tr><td><img src="/api/wiki/file/%s" data-width="50"></td></tr></table>' % (FILE_A, FILE_A),
                           fetch=self.fetcher())
        self.assertEqual(stats['images'], 2)
        self.assertIn(('Caption', 'Подпись'), [(p.style.name, p.text) for p in doc.paragraphs])
        cell = doc.tables[0].cell(0, 0)
        self.assertTrue(cell._tc.findall('.//' + qn('w:drawing')))


@unittest.skipIf(Document is None or Image is None, 'python-docx или Pillow не установлены')
class BlockTests(unittest.TestCase):
    def test_note_is_a_toned_box_with_an_icon(self):
        """Плашка как на сайте: значок тона слева, рамка и фон тона, заголовок чернилами тона."""
        doc, _ = build('<div data-wiki-block="note" data-tone="warn"><h4>Внимание</h4><p>Текст</p></div>')
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (1, 2))
        fill, line, ink = docx_export.TONES['warn']
        icon, cell = table.cell(0, 0), table.cell(0, 1)
        self.assertTrue(icon._tc.findall('.//' + qn('w:drawing')), 'значок тона — картинка в первой ячейке')
        self.assertEqual(cell_fill(icon), fill)
        self.assertEqual(cell_fill(cell), fill)
        self.assertIn('<w:left w:val="single" w:sz="4" w:space="0" w:color="%s"' % line, table_xml(table))
        self.assertIn('<w:insideV w:val="nil"', table_xml(table), 'между значком и текстом линии нет')
        heading, body = cell.paragraphs[0], cell.paragraphs[1]
        self.assertEqual(heading.text, 'Внимание')
        self.assertTrue(heading.runs[0].font.bold)
        self.assertEqual(str(heading.runs[0].font.color.rgb), ink)
        self.assertNotEqual(heading.style.name[:7], 'Heading', 'заголовок плашки не должен попадать в оглавление')
        self.assertEqual(body.text, 'Текст')

    def test_icons_follow_the_tone(self):
        """Значки рисуются Pillow в цвете тона; разные тона — разные картинки."""
        pngs = {kind: docx_export.icon_png(kind, docx_export.TONES.get(kind, docx_export.TONES['info'])[2])
                for kind in ('info', 'ok', 'warn', 'danger', 'tip', 'neutral', 'dark', 'play')}
        self.assertEqual(len(set(pngs.values())), len(pngs))
        image = Image.open(io.BytesIO(pngs['warn']))
        self.assertEqual((image.size, image.mode), ((docx_export.ICON_PX, docx_export.ICON_PX), 'RGBA'))
        opaque = {pixel[:3] for pixel in image.getdata() if pixel[3] == 255}
        self.assertEqual(opaque, {(0xB4, 0x53, 0x09)}, 'краска — чернила тона warn')
        self.assertTrue(any(pixel[3] == 0 for pixel in image.getdata()), 'фон прозрачный')

    def test_dark_note_has_light_text_strong_and_links(self):
        doc, _ = build('<div data-wiki-block="note" data-tone="dark"><p>Пример <strong>важно</strong> '
                       '<a href="https://x.kz">ссылка</a></p></div>')
        runs = {run.text: run for run in all_runs(note_cell(doc).paragraphs[0])}
        self.assertEqual(str(runs['Пример '].font.color.rgb), docx_export.DARK_TEXT)
        self.assertEqual(str(runs['важно'].font.color.rgb), docx_export.DARK_STRONG)
        self.assertEqual(str(runs['ссылка'].font.color.rgb), docx_export.DARK_LINK)

    def test_unknown_tone_falls_back_to_info(self):
        doc, _ = build('<div data-wiki-block="note" data-tone="purple"><p>Текст</p></div>')
        self.assertEqual(cell_fill(doc.tables[0].cell(0, 0)), docx_export.TONES['info'][0])

    def test_cards_grid_numbers_and_columns(self):
        doc, _ = build('<div data-wiki-block="cards" data-cols="2" data-numbered="true">'
                       '<div data-wiki-block="card" data-tone="ok"><h4>Первый</h4><p>а</p></div>'
                       '<div data-wiki-block="card"><h4>Второй</h4><p>б</p></div>'
                       '<div data-wiki-block="card"><h4>Третий</h4><p>в</p></div></div>')
        table = doc.tables[0]
        # Сетка с зазорами: между карточками пустые столбец и строка, как gap на сайте.
        self.assertEqual((len(table.rows), len(table.columns)), (3, 3))
        self.assertEqual(table.cell(0, 0).paragraphs[0].text, '\xa01\xa0  Первый')
        self.assertEqual(table.cell(0, 2).paragraphs[0].text, '\xa02\xa0  Второй')
        self.assertEqual(table.cell(2, 0).paragraphs[0].text, '\xa03\xa0  Третий')
        badge = table.cell(0, 0).paragraphs[0].runs[0]
        self.assertEqual((badge.text, str(badge.font.color.rgb)), ('\xa01\xa0', 'FFFFFF'))
        self.assertIn('w:fill="%s"' % docx_export.ACCENT, badge._r.xml, 'номер белым на акцентной плашке')
        self.assertEqual(table.cell(0, 1).text, '')
        self.assertEqual(table.cell(1, 0).text, '')
        self.assertIn('w:hRule="exact"', table.rows[1]._tr.xml, 'строка-зазор фиксированной высоты')
        self.assertEqual(cell_fill(table.cell(0, 0)), docx_export.TONES['ok'][0])
        self.assertIn('<w:left w:val="single" w:sz="18"', table.cell(0, 0)._tc.xml, 'цветная грань тона')
        self.assertEqual(cell_fill(table.cell(0, 2)), docx_export.SURFACE_ALT)
        self.assertIsNone(cell_fill(table.cell(0, 1)))

    def test_single_card_grid_does_not_leave_empty_columns(self):
        doc, _ = build('<div data-wiki-block="cards" data-cols="3">'
                       '<div data-wiki-block="card"><h4>Одна</h4></div></div>')
        self.assertEqual(len(doc.tables[0].columns), 1)

    def test_stats_are_large_numbers_with_captions(self):
        doc, _ = build('<div data-wiki-block="stats" data-cols="3">'
                       '<div data-wiki-block="stat"><h4>10 минут</h4><p>подпись</p></div>'
                       '<div data-wiki-block="stat"><h4>4,75</h4><p>оценка</p></div></div>')
        table = doc.tables[0]
        self.assertEqual(len(table.columns), 3, 'два показателя и зазор между ними')
        number, caption = table.cell(0, 0).paragraphs[:2]
        self.assertEqual(number.runs[0].font.size.pt, docx_export.STAT_SIZE)
        self.assertEqual(str(number.runs[0].font.color.rgb), docx_export.ACCENT)
        self.assertEqual(caption.runs[0].font.size.pt, 9)
        self.assertEqual(table.cell(0, 2).paragraphs[0].text, '4,75')
        self.assertEqual(cell_fill(table.cell(0, 2)), docx_export.SURFACE_ALT)

    def test_gallery_is_a_strip_of_frames(self):
        data = image_bytes('WEBP', (720, 1560))
        html = '<div data-wiki-block="gallery">' + ''.join(
            '<img src="/api/wiki/file/%s" alt="Шаг %d">' % (FILE_A, i) for i in range(1, 5)) + '</div>'
        doc, stats = build(html, fetch=lambda _id: (data, 'image/webp'))
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (2, 3))
        self.assertEqual(stats['images'], 4)
        self.assertEqual(len(doc.inline_shapes), 4)
        for shape in doc.inline_shapes:
            self.assertLessEqual(shape.height, docx_export.Cm(9))
            self.assertLessEqual(shape.width, docx_export.Cm(17) // 3)
        self.assertEqual(table.cell(0, 0).paragraphs[1].text, 'Шаг 1')

    def test_lead_is_larger_on_an_accent_backing(self):
        """Вводка как на сайте: крупнее, чернилами, на акцентной подложке с гранью слева."""
        doc, _ = build('<div data-wiki-block="lead"><p>Вводка</p><p>Ещё</p></div>')
        cell = box_cell(doc)
        self.assertEqual([p.text for p in cell.paragraphs], ['Вводка', 'Ещё'])
        run = cell.paragraphs[0].runs[0]
        self.assertEqual(run.font.size.pt, docx_export.LEAD_SIZE)
        self.assertEqual(str(run.font.color.rgb), docx_export.INK)
        self.assertEqual(cell.paragraphs[0].style.name, 'Wiki Lead')
        self.assertEqual(cell_fill(cell), docx_export.ACCENT_SOFT)
        self.assertIn('<w:left w:val="single" w:sz="24" w:space="0" w:color="%s"' % docx_export.ACCENT, cell._tc.xml)

    def test_collapsibles_open_into_a_framed_title_and_body(self):
        doc, _ = build('<details><summary>Раз</summary><p>Тело раз</p></details>'
                       '<div data-wiki-collapsible data-title="Два" data-default-open="true"><p>Тело два</p></div>')
        self.assertEqual(len(doc.tables), 2)
        pairs = [(p.text, bool(p.runs and p.runs[0].font.bold))
                 for index in (0, 1) for p in box_cell(doc, index).paragraphs]
        self.assertEqual(pairs, [('Раз', True), ('Тело раз', False), ('Два', True), ('Тело два', False)])
        title = box_cell(doc).paragraphs[0]
        self.assertIn('<w:bottom', title._p.pPr.xml, 'линия под заголовком, как на сайте')
        self.assertIn('w:color="%s"' % docx_export.LINE, table_xml(doc.tables[0]))

    def test_trainer_button_becomes_an_accent_card_with_a_link(self):
        doc, _ = build('<div data-wiki-trainer="sapar" data-label="Подписание" data-width="60">Подписание</div>',
                       portal='https://alfa330.github.io/OTP/', slug='статья-1')
        cell = box_cell(doc)
        self.assertEqual(cell_fill(cell), docx_export.ACCENT)
        head, line = cell.paragraphs[:2]
        self.assertEqual(head.text, '  Подписание')
        self.assertTrue(head._p.findall('.//' + qn('w:drawing')), 'треугольник «плей», как у кнопки')
        self.assertEqual(line.text, 'Тренажёр проходят в статье на портале · открыть')
        self.assertTrue(line._p.findall(qn('w:hyperlink')))
        self.assertAlmostEqual(doc.tables[0].columns[0].width / docx_export.Cm(17), 0.6, places=2)

    def test_plain_div_is_transparent(self):
        doc, _ = build('<div><p>Внутри</p></div>')
        self.assertEqual([p.text for p in body_paragraphs(doc)], ['Внутри'])


# ── 3а. Вид документа: по правилам статьи, а не шаблона Word ─────────────────

@unittest.skipIf(Document is None or Image is None, 'python-docx или Pillow не установлены')
class LookTests(unittest.TestCase):
    def setUp(self):
        self.doc, _ = build('<h1>Раз</h1><p>Текст</p><h2>Два</h2><p>Ещё</p><h3>Три</h3><h4>Четыре</h4>')

    def test_system_font_with_substitution_hints(self):
        """Helvetica Neue на всех алфавитах, без атрибутов темы (они сильнее имени);
        подсказки подстановки для Windows; язык проверки — русский."""
        styles_xml = self.doc.styles.element.xml
        normal = self.doc.styles['Normal']
        self.assertEqual(normal.font.name, docx_export.BODY_FONT)
        rFonts = normal.element.rPr.find(qn('w:rFonts'))
        self.assertEqual(rFonts.get(qn('w:cs')), docx_export.BODY_FONT)
        self.assertEqual(rFonts.get(qn('w:eastAsia')), docx_export.BODY_FONT)
        for style in ('Title', 'Heading 1', 'Heading 2', 'Caption', 'Quote', 'Footer'):
            element = self.doc.styles[style].element.rPr.find(qn('w:rFonts'))
            self.assertIsNone(element.get(qn('w:asciiTheme')), style)
            self.assertEqual(element.get(qn('w:ascii')), docx_export.BODY_FONT, style)
        self.assertNotIn('w:asciiTheme', styles_xml.split('<w:style ')[0], 'docDefaults без шрифта темы')
        self.assertIn('w:lang w:val="ru-RU"', styles_xml.split('<w:style ')[0])
        fonts = self.doc.part.part_related_by(docx_export.RT.FONT_TABLE).blob.decode('utf-8')
        self.assertRegex(fonts, r'w:name="Helvetica Neue">\s*<w:altName w:val="Segoe UI"/>')
        self.assertIn('w:panose1 w:val="%s"' % docx_export.BODY_PANOSE, fonts)
        self.assertRegex(fonts, r'w:name="Menlo">\s*<w:altName w:val="Consolas"/>')
        theme = self.doc.part.part_related_by(docx_export.RT.THEME).blob.decode('utf-8')
        self.assertEqual(theme.count('<a:latin typeface="Helvetica Neue"'), 2, 'шрифты темы — основной и заголовков')
        self.assertNotIn('typeface="Calibri"', theme)
        self.assertRegex(theme, r'<a:accent1>\s*<a:srgbClr val="%s"' % docx_export.ACCENT)
        self.assertRegex(theme, r'<a:hlink>\s*<a:srgbClr val="%s"' % docx_export.ACCENT)

    def test_template_look_is_gone(self):
        """Ни синих заголовков Office, ни черты под названием, ни курсивного подзаголовка."""
        styles = self.doc.styles
        self.assertEqual(str(styles['Title'].font.color.rgb), docx_export.INK)
        self.assertIsNone(styles['Title'].element.pPr.find(qn('w:pBdr')))
        self.assertEqual(styles['Title'].font.size.pt, docx_export.TITLE_SIZE)
        self.assertIsNone(styles['Subtitle'].font.italic)
        self.assertEqual(str(styles['Subtitle'].font.color.rgb), docx_export.MUTED)
        for level, size in docx_export.HEADING_SIZES.items():
            heading = styles['Heading %d' % level]
            self.assertEqual(heading.font.size.pt, size, level)
            self.assertTrue(heading.font.bold, level)
            self.assertIn(str(heading.font.color.rgb), (docx_export.INK, docx_export.INK_SOFT), level)
            self.assertTrue(heading.paragraph_format.keep_with_next, level)
        h1 = styles['Heading 1'].element.pPr.find(qn('w:pBdr'))
        self.assertIsNotNone(h1, 'у h1 — метка слева и линия снизу, как на сайте')
        self.assertEqual(h1.find(qn('w:left')).get(qn('w:color')), docx_export.ACCENT)
        self.assertEqual(h1.find(qn('w:bottom')).get(qn('w:color')), docx_export.LINE)
        self.assertIsNone(styles['Caption'].font.bold)
        self.assertEqual(str(styles['Caption'].font.color.rgb), docx_export.MUTED)
        for word in ('4F81BD', '365F91', '17365D', 'Calibri'):
            for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
                self.assertNotIn(word, styles[name].element.xml, (word, name))

    def test_links_are_indigo_with_a_pale_line(self):
        doc, _ = build('<p><a href="https://x.kz">ссылка</a></p>')
        run = all_runs(body_paragraphs(doc)[0])[0]
        self.assertEqual(run.text, 'ссылка')
        self.assertEqual(str(run.font.color.rgb), docx_export.ACCENT)
        self.assertIn('<w:u w:val="single" w:color="%s"' % docx_export.ACCENT_LINE, run._r.xml)

    def test_template_thumbnail_is_not_shipped(self):
        names = [str(part.partname) for part in self.doc.part.package.iter_parts()]
        self.assertNotIn('/docProps/thumbnail.jpeg', names)

    def test_toc_lists_h1_to_h3_when_there_are_two_or_more(self):
        """Оглавление — как на витрине: по h1–h3, при двух и больше, ссылками на
        заголовки; обёрнуто в поле TOC, чтобы «Обновить поле» дало номера страниц."""
        doc = self.doc
        label = by_style(doc, 'TOC Heading')
        self.assertEqual([p.text for p in label], ['Содержание'])
        entries = by_style(doc, 'toc 1') + by_style(doc, 'toc 2') + by_style(doc, 'toc 3')
        self.assertEqual(sorted(p.text for p in entries), ['Два', 'Раз', 'Три'])
        self.assertEqual([p.style.name for p in doc.paragraphs
                          if p.style.name.startswith('toc ')], ['toc 1', 'toc 2', 'toc 3'])
        first, last = by_style(doc, 'toc 1')[0], by_style(doc, 'toc 3')[0]
        self.assertIn('w:fldCharType="begin"', first._p.xml)
        self.assertIn('TOC \\o "1-3" \\h \\z \\u', first._p.xml)
        self.assertIn('w:fldCharType="end"', last._p.xml)
        self.assertIn('w:anchor="_wiki_h1"', first._p.xml)
        body_xml = doc.element.body.xml
        for number in (1, 2, 3):
            self.assertIn('w:name="_wiki_h%d"' % number, body_xml)
        self.assertNotIn('w:name="_wiki_h4"', body_xml, 'h4 в оглавление не идёт')
        headings = [p for p in body_paragraphs(doc) if p.style.name.startswith('Heading')]
        self.assertIn('<w:bookmarkStart', headings[0]._p.xml)
        # Один заголовок — оглавления нет, как и на витрине (toc.length > 1).
        single, _ = build('<h2>Один</h2><p>Текст</p>')
        self.assertEqual(by_style(single, 'TOC Heading'), [])
        self.assertNotIn('bookmarkStart', single.element.body.xml)

    def test_tables_have_explicit_width_and_fixed_layout(self):
        """Без w:tblW и фиксированной раскладки Word считает ширину «по
        содержимому», и плашка на три слова выходит шириной в три слова."""
        doc, _ = build('<div data-wiki-block="note"><p>Текст</p></div>'
                       '<table><tr><th>А</th></tr><tr><td>б</td></tr></table>')
        for table in doc.tables:
            xml = table_xml(table)
            self.assertIn('<w:tblW w:w="9638" w:type="dxa"/>', xml)
            self.assertIn('<w:tblLayout w:type="fixed"/>', xml)
            self.assertIn('<w:tblCellMar>', xml)
        self.assertIn('<w:cantSplit/>', doc.tables[0].rows[0]._tr.xml)

    def test_pictures_are_rounded_with_a_hairline(self):
        data = image_bytes('PNG', (400, 300))
        doc, _ = build('<img src="/api/wiki/file/%s">' % FILE_A, fetch=lambda _id: (data, 'image/png'))
        xml = doc.inline_shapes[0]._inline.xml
        self.assertIn('prst="roundRect"', xml)
        self.assertRegex(xml, r'<a:gd name="adj" fmla="val \d+"/>')
        self.assertIn('<a:ln w="6350">', xml)
        self.assertIn('<a:srgbClr val="%s"/>' % docx_export.LINE, xml)
        picture = [p for p in doc.paragraphs if p._p.findall('.//' + qn('w:drawing'))][0]
        self.assertTrue(picture.paragraph_format.keep_with_next, 'кадр не отрывается от подписи')

    def test_gallery_of_landscape_frames_is_one_per_row(self):
        data = image_bytes('WEBP', (1600, 900))
        html = '<div data-wiki-block="gallery">' + ''.join(
            '<img src="/api/wiki/file/%s">' % FILE_A for _ in range(2)) + '</div>'
        doc, _ = build(html, fetch=lambda _id: (data, 'image/webp'))
        table = doc.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (2, 1))
        self.assertEqual(cell_fill(table.cell(0, 0)), docx_export.SURFACE_ALT)

    def test_children_of_pPr_rPr_tcPr_tblPr_and_trPr_follow_the_schema(self):
        """Иначе Word объявляет файл повреждённым. Проверяется ВСЁ тело документа
        и таблица стилей, не только ячейки."""
        html = ('<div data-wiki-block="lead"><p>Вводка</p></div><h1>Раз</h1>'
                '<p style="text-align: center"><mark>а</mark> <code>б</code> <a href="https://x.kz">в</a></p>'
                '<ol data-variant="steps"><li><h4>Шаг</h4><ul><li>вложенный</li></ul></li></ol>'
                '<ul data-variant="checks"><li>да</li></ul><ul data-variant="chips"><li>чип</li></ul>'
                '<blockquote><p>Цитата</p></blockquote><pre>код</pre><hr>'
                '<div data-wiki-block="note" data-tone="dark"><h4>Т</h4><p><strong>ж</strong></p></div>'
                '<div data-wiki-block="cards" data-numbered="true"><div data-wiki-block="card" data-tone="ok"><h4>К</h4></div>'
                '<div data-wiki-block="card"><p>Б</p></div></div>'
                '<div data-wiki-block="stats"><div data-wiki-block="stat"><h4>1</h4><p>п</p></div></div>'
                '<details><summary>С</summary><p>Т</p></details>'
                '<div data-wiki-trainer="sapar" data-label="Т">Т</div>'
                '<table><thead><tr><th>А</th><th>Б</th></tr></thead><tr><td colspan="2">в</td></tr></table>'
                '<h2>Два</h2>')
        doc, _ = build(html, portal='https://alfa330.github.io/OTP/')
        orders = {
            'pPr': docx_export._PPR_ORDER, 'rPr': docx_export._RPR_ORDER, 'tcPr': docx_export._TCPR_ORDER,
            'tblPr': docx_export._TBLPR_ORDER, 'trPr': docx_export._TRPR_ORDER,
            'lvl': docx_export._LVL_ORDER, 'style': docx_export._STYLE_ORDER,
        }
        roots = [doc.element.body, doc.styles.element, doc.sections[0].footer._element,
                 doc.part.numbering_part.numbering_definitions._numbering]
        checked = 0
        for root in roots:
            for element in root.iter():
                tag = element.tag.split('}')[1]
                if tag not in orders:
                    continue
                names = [qn(name) for name in orders[tag]]
                ranks = [names.index(child.tag) for child in element if child.tag in names]
                self.assertEqual(ranks, sorted(ranks), (tag, [c.tag.split('}')[1] for c in element]))
                checked += 1
        self.assertGreater(checked, 100)

    def test_file_reopens_cleanly(self):
        html = ('<div data-wiki-block="lead"><p>Вводка</p></div><h1>Раз</h1><p>Текст</p>'
                '<div data-wiki-block="note" data-tone="tip"><p>Совет</p></div><h2>Два</h2>'
                '<ol data-variant="steps"><li>шаг</li></ol>')
        doc, _ = build(html)
        stream = io.BytesIO()
        doc.save(stream)
        with zipfile.ZipFile(stream) as package:
            self.assertIsNone(package.testzip())
            for name in package.namelist():
                if name.endswith('.xml'):
                    package.read(name).decode('utf-8')
        again = Document(io.BytesIO(stream.getvalue()))
        self.assertEqual(again.paragraphs[0].text, 'Статья')


# ── 4. Дверь ─────────────────────────────────────────────────────────────────

def make_context(otp_role='operator'):
    return {
        'user_id': 42, 'otp_role': otp_role, 'department_id': None, 'direction_id': None,
        'headed_department_ids': [], 'group_ids': [], 'wiki_roles': [], 'access_mode': 'auto',
    }


ARTICLE = {
    'id': 1, 'slug': 'reglament', 'title': 'Регламент: выплаты', 'summary': None,
    'content': ('<p>Текст</p><img src="/api/wiki/file/%s"><img src="/api/wiki/file/%s">'
                '<img src="/api/wiki/file/%s">' % (FILE_A, FILE_B, FILE_C)),
    'article_type': 'article', 'status': 'published', 'visibility_mode': 'inherit',
    'strict_mode': False, 'ai_opt_out': False, 'copy_protected': False, 'historical': False,
    'toc': None, 'views': 0, 'author_id': 7, 'author_name': 'Автор', 'owner_user_id': None,
    'updated_by': None, 'updated_at': datetime(2026, 10, 7), 'created_at': None,
    'published_at': None, 'review_due_at': None, 'cross_department': None,
    'source_article_id': None, 'source_article_title': None, 'section_ids': [3], 'tags': [],
}

FILE_ROWS = [
    {'id': FILE_A, 'article_id': 1, 'bucket': 'b', 'blob_path': 'a.webp',
     'content_type': 'image/webp', 'uploaded_by': 7},
    {'id': FILE_B, 'article_id': 9, 'bucket': 'b', 'blob_path': 'b.webp',
     'content_type': 'image/webp', 'uploaded_by': 7},
    {'id': FILE_C, 'article_id': None, 'bucket': 'b', 'blob_path': 'c.png',
     'content_type': 'image/png', 'uploaded_by': 42},
]


class _FakeBlob:
    def __init__(self, path, downloaded):
        self.path = path
        self.downloaded = downloaded

    def download_as_bytes(self):
        self.downloaded.append(self.path)
        return image_bytes('PNG', (40, 30))


class _FakeBucket:
    def __init__(self, downloaded):
        self.downloaded = downloaded

    def blob(self, path):
        return _FakeBlob(path, self.downloaded)


class _FakeClient:
    def __init__(self, downloaded):
        self.downloaded = downloaded

    def bucket(self, _name):
        return _FakeBucket(self.downloaded)


@unittest.skipIf(Flask is None or Document is None or Image is None,
                 'flask, python-docx или Pillow не установлены')
class DocxRouteTests(unittest.TestCase):
    def build_client(self, context, *, visible=(1, 2), article=ARTICLE, files=FILE_ROWS):
        cursor = MagicMock()
        cursor.calls = []
        cursor.execute.side_effect = lambda sql, params=None: cursor.calls.append((sql, params))
        cursor.fetchone.return_value = None
        cursor.fetchall.return_value = []
        db = MagicMock()

        @contextmanager
        def _get_cursor():
            yield cursor

        db._get_cursor = _get_cursor
        self.logged = []
        self.downloaded = []

        def fake_perimeter(_cursor, _ctx, **_kwargs):
            return collect_subjects(user_id=42, otp_role=context['otp_role']), {3}, set(visible)

        patches = [
            (queries, 'load_access_context', lambda _c, _u: dict(context)),
            (queries, 'granted_rule_rights', lambda _c, _s, _u: ({}, [])),
            (queries, 'manage_section_ids', lambda _c, _ctx, _s: frozenset()),
            (queries, 'log_action', lambda *a, **k: self.logged.append(k)),
            (wiki_perimeter, 'read_perimeter', fake_perimeter),
            (wiki_articles, 'get_article', lambda _c, **k: dict(article) if article else None),
            (wiki_articles, 'files_for_display', lambda _c, _ids: [dict(row) for row in files]),
            (wiki_article_access, 'section_chains', lambda _c, _ids: {}),
        ]
        for module, name, replacement in patches:
            original = getattr(module, name)
            setattr(module, name, replacement)
            self.addCleanup(setattr, module, name, original)

        app = Flask(__name__)
        app.register_blueprint(build_wiki_blueprint(
            db=db, require_api_key=lambda f: f,
            build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=lambda: (42, None, None),
            sensitive_access_granted=lambda _user_id, cursor=None: True,
            client_ip=lambda: '127.0.0.1',
            gcs={'signed_url': lambda *a, **k: 'https://x', 'bucket_name': lambda: 'b',
                 'client': lambda: _FakeClient(self.downloaded)},
        ))
        app.config['TESTING'] = True
        return app.test_client(), cursor

    def test_below_admin_is_refused_whatever_the_article(self):
        for role in ('operator', 'trainee', 'trainer', 'sv', 'supervisor', 'hr_manager'):
            client, _ = self.build_client(make_context(role))
            for article_id in (1, 2, 777):
                response = client.get('/api/wiki/articles/%d/docx' % article_id)
                self.assertEqual(response.status_code, 403, (role, article_id))
                self.assertEqual(response.get_json().get('code'), 'WIKI_FORBIDDEN')
                self.assertEqual(response.get_json().get('required'), 'admin')
            self.assertEqual(self.logged, [], 'отказ в журнал статьи не пишется')

    def test_admin_gets_the_file(self):
        for role in ('admin', 'super_admin', 'superadmin'):
            client, _ = self.build_client(make_context(role))
            response = client.get('/api/wiki/articles/1/docx')
            self.assertEqual(response.status_code, 200, role)
            self.assertEqual(response.mimetype, docx_export.DOCX_MIME)
            disposition = response.headers.get('Content-Disposition', '')
            self.assertIn('attachment', disposition)
            self.assertIn('.docx', disposition)
            doc = Document(io.BytesIO(response.data))
            self.assertEqual(doc.paragraphs[0].text, 'Регламент: выплаты')

    def test_invisible_article_is_not_found_not_forbidden(self):
        client, _ = self.build_client(make_context('admin'), visible=(2,))
        response = client.get('/api/wiki/articles/1/docx')
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.logged, [])

    def test_missing_article_is_not_found(self):
        client, _ = self.build_client(make_context('admin'), article=None)
        self.assertEqual(client.get('/api/wiki/articles/1/docx').status_code, 404)

    def test_only_readable_files_leave_the_bucket(self):
        """Файл чужой статьи (9 вне периметра) не скачивается; свой и
        непривязанный собственный — да. То же правило, что у /file/<id>."""
        client, _ = self.build_client(make_context('admin'))
        response = client.get('/api/wiki/articles/1/docx')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(sorted(self.downloaded), ['a.webp', 'c.png'])
        doc = Document(io.BytesIO(response.data))
        self.assertEqual(len(doc.inline_shapes), 2)
        self.assertEqual([t for t in texts(doc) if 'недоступно' in t], ['[изображение недоступно]'])

    def test_export_is_written_to_the_journal(self):
        client, _ = self.build_client(make_context('admin'))
        client.get('/api/wiki/articles/1/docx')
        self.assertEqual(len(self.logged), 1)
        entry = self.logged[0]
        self.assertEqual(entry['action'], 'article.export')
        self.assertEqual((entry['entity_type'], entry['entity_id'], entry['actor_id']), ('article', 1, 42))
        self.assertEqual(entry['details'], {'slug': 'reglament', 'format': 'docx', 'images': 2})

    def test_portal_address_reaches_the_links_only_when_sane(self):
        client, _ = self.build_client(make_context('admin'))
        good = client.get('/api/wiki/articles/1/docx?portal=https://alfa330.github.io/OTP/')
        footer = Document(io.BytesIO(good.data)).sections[0].footer
        self.assertEqual(hyperlinks(None, footer.part),
                         ['https://alfa330.github.io/OTP/?view=wiki&article=reglament'])
        bad = client.get('/api/wiki/articles/1/docx?portal=javascript:alert(1)')
        self.assertEqual(bad.status_code, 200)
        footer = Document(io.BytesIO(bad.data)).sections[0].footer
        self.assertEqual(hyperlinks(None, footer.part), [])

    def test_missing_dependency_is_a_clear_503(self):
        client, _ = self.build_client(make_context('admin'))
        original = docx_export.Document
        docx_export.Document = None
        self.addCleanup(setattr, docx_export, 'Document', original)
        response = client.get('/api/wiki/articles/1/docx')
        self.assertEqual(response.status_code, 503)
        self.assertIn('Word', response.get_json()['error'])

    def test_article_payload_tells_the_button_whether_to_show(self):
        """Признак can_download едет вместе со статьёй и считается той же
        лестницей, что дверь: кнопка не появится у того, кому дверь откажет."""
        for role, expected in (('admin', True), ('super_admin', True), ('sv', False),
                               ('operator', False)):
            client, _ = self.build_client(make_context(role))
            response = client.get('/api/wiki/articles/reglament')
            self.assertEqual(response.status_code, 200, role)
            self.assertIs(response.get_json().get('can_download'), expected, role)

    def test_action_is_filterable_and_labelled(self):
        self.assertIn('article.export', wiki_structure.AUDIT_GROUPS['articles'])
        self.assertRegex(AUDIT_JS, r"'article\.export':\s*\{\s*label:\s*'[^']+'")


# ── 5. Кнопка и зависимость ─────────────────────────────────────────────────

class FrontendTests(unittest.TestCase):
    def setUp(self):
        self.article = strip_comments(ARTICLE_JSX)
        self.helper = strip_comments(DOWNLOAD_JS)

    def test_button_is_gated_by_the_server_flag(self):
        block = re.search(r'\{article\.can_download && \((.*?)\n\s*\)\}', self.article, re.S)
        self.assertIsNotNone(block, 'кнопка «Скачать» не стоит за признаком can_download')
        self.assertIn('Скачать', block.group(1))
        self.assertIn('onClick={download}', block.group(1))
        self.assertIn('<Download size={14} />', block.group(1))
        self.assertIn("import { docxFileName, downloadErrorText, portalAddress, saveBlob } from './articleDownload';",
                      ARTICLE_JSX)

    def test_no_role_ladder_on_the_client(self):
        """Лестница ролей — ОДНА и на сервере; второй экземпляр во фронте разошёлся бы молча."""
        for needle in ("role === 'admin'", "role === 'super_admin'", 'isAdminLikeRole'):
            self.assertNotIn(needle, self.article, needle)

    def test_download_goes_through_the_api_as_a_blob_in_the_same_window(self):
        start = self.article.index('const download = () => {')
        end = self.article.index('};', start)
        body = self.article[start:end]
        self.assertIn('/articles/${article.id}/docx', body)
        self.assertIn("responseType: 'blob'", body)
        self.assertIn('portal: portalAddress()', body)
        self.assertIn('saveBlob(r.data, docxFileName(article.title))', body)
        self.assertIn('downloadErrorText(e', body)
        for forbidden in ('window.open', '_blank', 'location.href', 'location.assign'):
            self.assertNotIn(forbidden, body, forbidden)
            self.assertNotIn(forbidden, self.helper, forbidden)

    def test_helper_names_the_file_like_the_server(self):
        self.assertIn("'Статья'", self.helper)
        self.assertIn('.docx', self.helper)
        self.assertIn('FILENAME_MAX = %d' % docx_export.FILENAME_MAX, self.helper)

    def test_dependency_is_pinned(self):
        self.assertRegex(REQUIREMENTS, r'(?m)^python-docx==\d+\.\d+\.\d+$')


if __name__ == '__main__':  # pragma: no cover
    unittest.main()
