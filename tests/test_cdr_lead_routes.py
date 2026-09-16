# -*- coding: utf-8 -*-
"""Ручки режима «Сделки» раздела «Касания»: /api/cdr/leads*.

Сети и базы нет: слои `queries` и `lead_queries` подменяются двойниками, Flask
поднимается своим приложением с фальшивыми зависимостями — тот же приём, что в
tests/test_cdr_agent_routes.py. Проверяется склейка ответа: записи получают свои
звонки, фильтр «без звонков» работает после привязки, страницы режутся после
фильтров, выгрузка отдаёт книгу, догрузка идёт в фоне и не запускается дважды.

Телефоны учебные (7XX555XXXX).
"""

import json
import unittest
from datetime import date, datetime
from unittest import mock

from openpyxl import load_workbook

try:
    from flask import Flask
except ImportError:  # pragma: no cover
    Flask = None

from cdr import routes as cdr_routes

TODAY = date(2026, 9, 16)
ADMIN = {'user_id': 1, 'name': 'Админ', 'role': 'super_admin', 'department_id': None,
         'department_code': None, 'headed_department_ids': [], 'headed_department_codes': []}
OPERATOR = dict(ADMIN, role='operator', department_code='op')


class _FakeCursor:
    def execute(self, sql, params=None):
        pass

    def fetchone(self):
        return (1,)

    def fetchall(self):
        return []


class _FakeDb:
    def _get_cursor(self):
        from contextlib import contextmanager

        @contextmanager
        def scope():
            yield _FakeCursor()
        return scope()


class _Queries:
    """Двойник cdr.queries: доступ, состояние моста, справочник."""

    def __init__(self, ctx=ADMIN):
        self.ctx = ctx
        self.enqueued = []

    def load_access_context(self, cursor, user_id):
        return dict(self.ctx)

    def day_states(self, cursor, day_from, day_to):
        return []

    def enqueue_days(self, cursor, days, requested_by=None):
        self.enqueued.extend(days)
        return len(days)

    def agent_state(self, cursor):
        return {'connected': True, 'last_seen_at': '2026-09-16T10:00:00'}

    def load_directory(self, cursor):
        return {'6474': {'periods': [{'since': None, 'name': 'Жупан Аружан',
                                      'direction': 'Основа ОП'}]},
                '6656': {'periods': [{'since': None, 'name': 'Зинеден Аружан',
                                      'direction': 'Основа ОП'}]}}

    def directory_updated_at(self, cursor):
        return datetime(2026, 9, 16, 9, 0, 0)

    def today_almaty(self):
        return TODAY

    def now_almaty(self):
        return datetime(2026, 9, 16, 12, 0, 0)


def _lead(key, phone, moment, **over):
    row = {
        'source': 'amo', 'stream_type': 0, 'lead_key': key, 'work_day': moment.date(),
        'user_id': 7, 'owner_name': 'Жупан Аружан', 'owner_ext': '6474',
        'owner_external_name': None, 'owner_raw': '12666570',
        'full_name': 'Заявка %s' % key, 'phone': phone, 'phones': phone,
        'park_name': 'iTaxi', 'city': 'Астана', 'base_title': '',
        'stage_raw': 'ПРОШЕЛ РЕГИСТРАЦИЮ', 'call_status': '', 'dialog_status': '',
        'reason_raw': '', 'sub_reason_raw': '', 'comment': '', 'tags': 'forma_itaxi_google',
        'utm_source': 'google', 'lead_type': 'форма', 'registered': 1,
        'created_at': moment, 'taken_at': moment, 'updated_at': moment,
    }
    row.update(over)
    return row


def _touch(phone, started, **over):
    row = {'started_at': started, 'answered_at': None, 'phone': phone, 'ext': '6474',
           'call_type': 'Исходящий', 'result': 'Не ответил', 'talk_seconds': 0,
           'dial_seconds': 15, 'queue': '', 'recording_url': 'http://recordings.test/x.wav',
           'linkedid': '1.%d' % int(started.timestamp()), 'legs': 1}
    row.update(over)
    return row


class _LeadQueries:
    """Двойник cdr.lead_queries с двумя сделками и тремя звонками."""

    def __init__(self):
        self.leads = [
            _lead('101', '7015550001', datetime(2026, 9, 10, 10, 0, 0)),
            _lead('102', '7025550002', datetime(2026, 9, 11, 12, 0, 0)),
        ]
        self.touches = [
            # за 10 секунд до появления сделки — её звонок (допуск)
            _touch('7015550001', datetime(2026, 9, 10, 9, 59, 50), call_type='Входящий',
                   result='Разговор', talk_seconds=70),
            _touch('7015550001', datetime(2026, 9, 10, 10, 30, 0), ext='6656'),
            # чужой номер — никуда
            _touch('7995550009', datetime(2026, 9, 10, 11, 0, 0)),
        ]
        self.last_run = None
        self.without_phone = 0

    def select_leads(self, cursor, source, direction, day_from, day_to, filters=None):
        return [dict(row) for row in self.leads]

    def select_touches_for_phones(self, cursor, phones, day_from, day_to):
        return [dict(row) for row in self.touches if row['phone'] in phones]

    def filter_values(self, cursor, source, direction, day_from, day_to):
        return {'parks': ['iTaxi'], 'cities': ['Астана'], 'stages': ['ПРОШЕЛ РЕГИСТРАЦИЮ'],
                'lead_types': ['форма'], 'owners': [{'key': 'user:7', 'label': 'Жупан Аружан',
                                                     'count': 2}]}

    def lead_days(self, cursor, source, direction, day_from, day_to):
        return {'2026-09-10': 1, '2026-09-11': 1}

    def count_without_phones(self, cursor, source, direction, day_from, day_to):
        return self.without_phone

    def last_lead_run(self, cursor, direction):
        return self.last_run


@unittest.skipIf(Flask is None, 'flask не установлен')
class LeadRoutesTests(unittest.TestCase):
    def setUp(self):
        self.queries = _Queries()
        self.lead_queries = _LeadQueries()
        for target, double in (('queries', self.queries), ('lead_queries', self.lead_queries)):
            patcher = mock.patch.object(cdr_routes, target, double)
            patcher.start()
            self.addCleanup(patcher.stop)
        schema_patch = mock.patch.object(cdr_routes.schema, 'schema_is_ready', return_value=True)
        schema_patch.start()
        self.addCleanup(schema_patch.stop)

        app = Flask(__name__)
        app.register_blueprint(cdr_routes.build_cdr_blueprint(
            db=_FakeDb(),
            require_api_key=lambda fn: fn,
            build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=lambda: (1, None, None),
        ))
        app.config['TESTING'] = True
        self.client = app.test_client()

    def get(self, path, **params):
        params.setdefault('source', 'amo')
        params.setdefault('date_from', '2026-09-10')
        params.setdefault('date_to', '2026-09-12')
        return self.client.get('/api/cdr/' + path, query_string=params)

    # ── /leads ───────────────────────────────────────────────────────────────

    def test_записи_получают_свои_звонки_включая_допуск(self):
        body = self.get('leads').get_json()
        self.assertEqual(body['total'], 2)
        first = next(row for row in body['leads'] if row['key'] == '101')
        self.assertEqual(first['agg']['touches'], 2)
        self.assertEqual(first['agg']['talks'], 1)
        self.assertEqual(first['agg']['before_lead'], 1)
        self.assertEqual(first['owner_called'], 'да')
        self.assertEqual(first['touches'][0]['operator'], 'Жупан Аружан')
        self.assertEqual(first['touches'][1]['operator'], 'Зинеден Аружан')
        self.assertEqual(first['touches'][0]['started_at'], '2026-09-10 09:59:50')
        second = next(row for row in body['leads'] if row['key'] == '102')
        self.assertEqual(second['agg']['touches'], 0)
        self.assertEqual(second['owner_called'], '')

    def test_сводка_и_окно_звонков_на_сутки_дольше(self):
        body = self.get('leads').get_json()
        self.assertEqual(body['summary']['leads'], 2)
        self.assertEqual(body['summary']['with_touches'], 1)
        self.assertEqual(body['summary']['touches'], 2)
        self.assertEqual(body['summary']['blocks'], 2)
        self.assertTrue(body['summary']['full_url'])
        self.assertEqual(body['period']['touches_to'], '2026-09-13')
        self.assertEqual(body['source']['label'], 'Основа')
        self.assertEqual([s['key'] for s in body['sources']], ['amo', 'crm_paid_hire', 'crm_stream'])

    def test_окно_звонков_не_уходит_в_будущее(self):
        body = self.get('leads', date_to='2026-09-16').get_json()
        self.assertEqual(body['period']['touches_to'], '2026-09-16')

    def test_фильтр_без_звонков_после_привязки(self):
        body = self.get('leads', presence='without').get_json()
        self.assertEqual([row['key'] for row in body['leads']], ['102'])
        self.assertEqual(body['total'], 1)
        body = self.get('leads', presence='with').get_json()
        self.assertEqual([row['key'] for row in body['leads']], ['101'])

    def test_только_звонки_ответственного_не_меняют_ответ_звонил_ли_он(self):
        body = self.get('leads', own_only=1).get_json()
        first = next(row for row in body['leads'] if row['key'] == '101')
        self.assertEqual([t['ext'] for t in first['touches']], ['6474'])
        self.assertEqual(first['owner_called'], 'да')
        self.assertIn('только звонки ответственного', body['filters_note'])

    def test_страницы_режутся_после_фильтров(self):
        body = self.get('leads', page_size=1, page=2).get_json()
        self.assertEqual(body['total'], 2)
        self.assertEqual(len(body['leads']), 1)

    def test_недостающие_сутки_звонков_ставятся_мосту(self):
        self.get('leads')
        self.assertEqual(len(self.queries.enqueued), 4)      # 10, 11, 12, 13 сентября
        self.queries.enqueued.clear()
        self.get('leads', sync=0)
        self.assertEqual(self.queries.enqueued, [])

    def test_покрытие_записей_видно_в_ответе(self):
        self.lead_queries.without_phone = 5
        body = self.get('leads').get_json()
        self.assertEqual(body['leads_coverage']['missing_days'], ['2026-09-12'])
        self.assertEqual(body['leads_coverage']['without_phone'], 5)

    def test_неизвестный_источник_это_400(self):
        response = self.get('leads', source='oktell')
        self.assertEqual(response.status_code, 400)
        self.assertIn('источник', response.get_json()['error'])

    def test_оператору_закрыто(self):
        self.queries.ctx = OPERATOR
        self.assertEqual(self.get('leads').status_code, 403)

    # ── /leads/export ────────────────────────────────────────────────────────

    def test_выгрузка_отдаёт_книгу_с_листами(self):
        response = self.get('leads/export')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(__import__('io').BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['Контекст', 'Лиды', 'Касания', 'Операторы', 'Сводка'])
        header = [cell.value for cell in book['Лиды'][2]]
        self.assertIn('Ссылка на запись', header)
        self.assertEqual(book['Касания'].max_row, 3)      # шапка + два касания

    def test_json_двойник_начинается_с_контекста(self):
        response = self.get('leads/export', format='json')
        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.data.decode('utf-8'))
        self.assertEqual(list(payload)[0], 'КОНТЕКСТ_ДЛЯ_ИИ')
        self.assertEqual(len(payload['записи']), 2)

    # ── /leads/sync ──────────────────────────────────────────────────────────

    def test_догрузка_идёт_в_фоне_и_отвечает_сразу(self):
        with mock.patch.object(cdr_routes._LEAD_SYNC_POOL, 'submit') as submit:
            response = self.client.post('/api/cdr/leads/sync', query_string={
                'source': 'amo', 'date_from': '2026-09-10', 'date_to': '2026-09-12'})
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.get_json()['status'], 'started')
        self.assertEqual(submit.call_count, 1)

    def test_догрузка_не_запускается_поверх_идущей(self):
        self.lead_queries.last_run = {'status': 'running', 'started_at': '2026-09-16 11:55:00'}
        with mock.patch.object(cdr_routes._LEAD_SYNC_POOL, 'submit') as submit:
            response = self.client.post('/api/cdr/leads/sync', query_string={
                'source': 'amo', 'date_from': '2026-09-10', 'date_to': '2026-09-12'})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(submit.call_count, 0)

    def test_брошенный_прогон_не_блокирует_новый(self):
        self.lead_queries.last_run = {'status': 'running', 'started_at': '2026-09-16 09:00:00'}
        with mock.patch.object(cdr_routes._LEAD_SYNC_POOL, 'submit'):
            response = self.client.post('/api/cdr/leads/sync', query_string={
                'source': 'amo', 'date_from': '2026-09-10', 'date_to': '2026-09-12'})
        self.assertEqual(response.status_code, 202)

    def test_догрузка_длиннее_месяца_отклоняется(self):
        response = self.client.post('/api/cdr/leads/sync', query_string={
            'source': 'amo', 'date_from': '2026-07-01', 'date_to': '2026-09-12'})
        self.assertEqual(response.status_code, 400)

    def test_догрузка_закрыта_без_права(self):
        self.queries.ctx = OPERATOR
        response = self.client.post('/api/cdr/leads/sync', query_string={
            'source': 'amo', 'date_from': '2026-09-10', 'date_to': '2026-09-12'})
        self.assertEqual(response.status_code, 403)


if __name__ == '__main__':
    unittest.main()
