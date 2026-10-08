# -*- coding: utf-8 -*-
"""Раздел «Списки Байги» (#356): права, разбор файла, поиск, выгрузка, ручки.

Что сторожится:
  * периметр (решение владельца 05.10.2026): супер-админ и глава «Маркетинга»
    ведут раздел; главы и супервайзеры ОП и СЗоВ читают; операторы ОП и СЗоВ и
    сотрудники «Маркетинга» читают после QR; остальным закрыто; меню во фронте
    отвечает так же, как сервер;
  * выдачи из раздела (кнопка «Доступ», решение владельца 07.10.2026): человеку,
    группе, отделу — только добавляют к кругу; раздают супер-админ и названные
    поимённо; рядовой и по выдаче входит через QR;
  * разбор файла по п. 3 и п. 7 постановки: колонки по названию, период из
    имени (оба вида) или из «Даты», ошибки блокируют, предупреждения — нет,
    у каждой — лист и строка Excel;
  * поиск без учёта казахских букв (Жусип = Жүсіп) и раскладки в номере ВУ,
    пользовательский текст в SQL не попадает;
  * выгрузка недели без фильтров загружается обратно и совпадает построчно;
  * загрузка недели атомарна: занятая неделя без подтверждения — отказ без
    единой записи, недописанные строки — откат;
  * каждая выгрузка и скачивание исходника — в журнал;
  * раздел подключён: схема, Blueprint, пункт меню, гард видимости.

Все ФИО, номера ВУ и ID здесь ВЫДУМАНЫ: настоящий файл недели — персональные
данные водителей, и в публичный репозиторий он не идёт ни целиком, ни строкой.

SQL здесь не проверяется — его проверяет прогон на стенде с базой.
"""

import json
import re
import shutil
import subprocess
import sys
import unittest
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from functools import wraps
from io import BytesIO
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    from flask import Flask
except ImportError:  # pragma: no cover
    Flask = None

from openpyxl import Workbook, load_workbook  # noqa: E402

from baiga import access, filters, parse, queries, report, routes, schema  # noqa: E402
from parcels import access as parcels_access  # noqa: E402
from tests import source_cache  # noqa: E402

APP_JSX = ROOT / 'src' / 'App.jsx'
META_JS = ROOT / 'src' / 'components' / 'baiga' / 'baigaMeta.js'
DATABASE_PY = ROOT / 'database.py'
BOT_PY = ROOT / 'bot_schedule2.py'

NAME = 'Список байги - 21.09.2026 - 27.09.2026.xlsx'
HEADER = ['Позиция', 'Дата', 'Неделя', 'Водитель', 'Наименование приза', 'Сумма', 'Поездок',
          'Город', 'Таксопарк', 'Номер ВУ', 'ID водителя']


def _read(path):
    return path.read_text(encoding='utf-8-sig')


def person(role='operator', department_code='marketing', headed_codes=(), user_id=10, grants=(), circle=None):
    """grants — уровни выдач, под которые человек подпадает (кнопка «Доступ»);
    circle — правки строк «открыт по умолчанию»: {строка: уровень}."""
    return {
        'user_id': user_id, 'name': 'Сотрудник %d' % user_id, 'role': role,
        'department_id': 909, 'department_code': department_code, 'city': 'Алматы',
        'headed_department_ids': [909] if headed_codes else [],
        'headed_department_codes': list(headed_codes),
        'grant_levels': list(grants),
        'circle_levels': dict(circle or {}),
    }


def fake_id(number):
    return '%032x' % (0xbaa9a000 + number)


def line(position, name=None, prize='Не указано', amount=None, trips=40, city='Алматы',
         park='Тестовый парк (Алматы)', license_value=None, driver=None, date_text='21.09.26', week=39):
    return [position, date_text, week, name or 'Тестов%d Т. Т.' % position, prize,
            amount if amount is not None else 300000 - position * 1000, trips, city, park,
            license_value if license_value is not None else 'ZZ%06d' % position,
            driver if driver is not None else fake_id(position)]


def book(sheets):
    """{имя листа: [строки]} → байты xlsx. Первая строка листа — как дали."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def standard_book():
    return book({
        'Алматы-Каскелен': [HEADER, line(1, prize='200 000 тенге'), line(2, prize='5000 тенге'),
                            line(3, city='Шымкент')],
        'МОТО БАЙГА': [HEADER, line(1, name='Жүсіпов А. Б.', prize='10 000 тенге', driver=fake_id(101),
                                    license_value='ZZ000101')],
    })


# ─────────────────────────────────────────────────────────────────────────────
# Права
# ─────────────────────────────────────────────────────────────────────────────

FULL = {'can_open': True, 'can_export': True, 'can_manage': True, 'can_manage_access': False}
EXPORT = {'can_open': True, 'can_export': True, 'can_manage': False, 'can_manage_access': False}
READ = {'can_open': True, 'can_export': False, 'can_manage': False, 'can_manage_access': False}
CLOSED = {'can_open': False, 'can_export': False, 'can_manage': False, 'can_manage_access': False}
# Супер-админ ещё и раздаёт доступ — единственный, кому это положено должностью.
OWNER = dict(FULL, can_manage_access=True)

# Аналитик из именного списка (решение владельца 06.10.2026) — только id.
ANALYST_ID = 540

# Сетка для сверок «все сочетания»: роли портала, отделы раздела и чужие,
# главенство — своё, чужое и двойное.
GRID_ROLES = ('super_admin', 'admin', 'sv', 'supervisor', 'trainer', 'operator', 'trainee',
              'marketing_manager', 'accounting_manager')
GRID_DEPARTMENTS = ('marketing', 'op', 'szov', 'tez', 'front_office', None)
GRID_HEADS = ((), ('marketing',), ('op',), ('szov',), ('tez',), ('op', 'szov'), ('tez', 'marketing'),
              ('tez', 'szov'))


def grid():
    return [(role, code, heads) for role in GRID_ROLES for code in GRID_DEPARTMENTS for heads in GRID_HEADS]


class SwitchTests(unittest.TestCase):
    """Выключатель «только супер-админ»: снят 05.10.2026, но обязан работать."""

    def test_switch_is_off_and_closes_everyone_else_when_on(self):
        self.assertFalse(access.PILOT_SUPER_ADMIN_ONLY)
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True):
            self.assertEqual(access.capabilities(person(role='super_admin', department_code=None)), OWNER)
            for kwargs in ({'department_code': 'szov'}, {'department_code': 'op'}, {'role': 'marketing_manager'},
                           {'role': 'sv', 'department_code': 'op'},
                           {'role': 'admin', 'department_code': 'szov', 'headed_codes': ('szov',)},
                           {'role': 'admin', 'department_code': 'marketing', 'headed_codes': ('marketing',)}):
                self.assertEqual(access.capabilities(person(**kwargs)), CLOSED, kwargs)

    def test_switch_and_analysts_are_the_same_on_both_sides(self):
        app = _read(APP_JSX)
        self.assertIn('const BAIGA_PILOT_SUPER_ADMIN_ONLY = %s;' % str(access.PILOT_SUPER_ADMIN_ONLY).lower(), app)
        ids = ', '.join(str(user_id) for user_id in sorted(access.ANALYST_USER_IDS))
        self.assertIn('const BAIGA_ANALYST_USER_IDS = new Set([%s]);' % ids, app)


class AccessTests(unittest.TestCase):
    """Круг доступа — решение владельца 05.10.2026."""

    def setUp(self):
        patcher = mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False)
        patcher.start()
        self.addCleanup(patcher.stop)

    def caps(self, **kwargs):
        return access.capabilities(person(**kwargs))

    def test_super_admin_does_everything(self):
        self.assertEqual(self.caps(role='super_admin', department_code=None), OWNER)

    def test_marketing_head_has_full_access_whatever_the_base_role(self):
        for role in ('admin', 'sv', 'operator', 'marketing_manager'):
            for own in ('marketing', None):
                self.assertEqual(self.caps(role=role, department_code=own, headed_codes=('marketing',)), FULL,
                                 (role, own))

    def test_op_and_szov_heads_read_without_export(self):
        for code in ('op', 'szov'):
            for role in ('admin', 'sv', 'operator'):
                ctx = person(role=role, department_code=code, headed_codes=(code,))
                self.assertEqual(access.capabilities(ctx), READ, (code, role))
                self.assertFalse(access.requires_sensitive_qr(ctx), (code, role))
        # Глава узнаётся по главенству, а не по своему отделу в карточке.
        self.assertEqual(self.caps(role='admin', department_code=None, headed_codes=('szov',)), READ)

    def test_supervisors_of_op_and_szov_read_without_qr(self):
        for code in ('op', 'szov'):
            for role in ('sv', 'supervisor'):
                ctx = person(role=role, department_code=code)
                self.assertEqual(access.capabilities(ctx), READ, (code, role))
                self.assertFalse(access.requires_sensitive_qr(ctx), (code, role))
        for code in ('marketing', 'tez', 'front_office', None):
            self.assertEqual(self.caps(role='sv', department_code=code), CLOSED, code)

    def test_operators_and_marketing_staff_read_behind_qr(self):
        for kwargs in ({'department_code': 'op'}, {'department_code': 'szov'},
                       {'role': 'marketing_manager', 'department_code': 'marketing'},
                       {'department_code': 'marketing'}):
            ctx = person(**kwargs)
            self.assertEqual(access.capabilities(ctx), READ, kwargs)
            self.assertTrue(access.requires_sensitive_qr(ctx), kwargs)

    def test_everyone_else_is_closed(self):
        for kwargs in (
                {'role': 'trainer', 'department_code': 'szov'}, {'role': 'trainer', 'department_code': 'op'},
                {'role': 'trainer', 'department_code': 'marketing'},
                # Стажёра общий QR-замок не спрашивает — значит, и раздел ему закрыт.
                {'role': 'trainee', 'department_code': 'szov'}, {'role': 'trainee', 'department_code': 'op'},
                # Админ, не возглавляющий отдел, в круг не назван — где бы ни числился.
                {'role': 'admin', 'department_code': None}, {'role': 'admin', 'department_code': 'szov'},
                {'role': 'admin', 'department_code': 'op'}, {'role': 'admin', 'department_code': 'marketing'},
                # Главы прочих отделов — тоже.
                {'role': 'admin', 'department_code': 'tez', 'headed_codes': ('tez',)},
                {'role': 'admin', 'department_code': 'front_office', 'headed_codes': ('front_office',)},
                # Глава чужого отдела, числящийся оператором ОП: замок его не
                # спрашивает (он глава), а без замка рядового в раздел не пускают.
                {'role': 'operator', 'department_code': 'op', 'headed_codes': ('tez',)},
                {'department_code': 'tez'}, {'department_code': 'front_office'}, {'department_code': None},
                {'role': 'sv', 'department_code': 'tez'}):
            self.assertEqual(self.caps(**kwargs), CLOSED, kwargs)

    def test_department_codes_are_compared_without_case_and_spaces(self):
        self.assertEqual(self.caps(department_code=' OP '), READ)
        self.assertEqual(self.caps(role='sv', department_code='SZOV'), READ)
        self.assertEqual(self.caps(role='admin', department_code=None, headed_codes=(' Marketing ',)), FULL)

    def test_analyst_by_id_manages(self):
        with mock.patch.object(access, 'ANALYST_USER_IDS', frozenset({77})):
            self.assertEqual(self.caps(department_code='tez', user_id=77), FULL)
            self.assertEqual(self.caps(department_code='tez', user_id=78), CLOSED)

    def test_named_analyst_has_the_whole_section_without_qr(self):
        """Решение владельца 06.10.2026 про сотрудника отдела аналитики:
        «без qr, полный доступ к разделу». Числится он оператором, и общий
        замок спросил бы у него код."""
        self.assertEqual(access.ANALYST_USER_IDS, frozenset({ANALYST_ID}))
        analyst = person(role='operator', department_code='analytik', user_id=ANALYST_ID)
        self.assertEqual(access.capabilities(analyst), FULL)
        self.assertFalse(access.requires_sensitive_qr(analyst))
        # Выдано человеку, а не отделу: сосед по отделу и его глава раздела не видят.
        for kwargs in ({'user_id': ANALYST_ID + 1}, {'role': 'trainee', 'user_id': ANALYST_ID + 1},
                       {'role': 'admin', 'headed_codes': ('analytik',), 'user_id': 229}):
            self.assertEqual(self.caps(department_code='analytik', **kwargs), CLOSED, kwargs)

    def test_analyst_is_never_asked_for_qr_whatever_the_role_and_department(self):
        """Замок снят с человека из списка, а не с его должности: на любом
        сочетании роли, отдела и главенства он входит без кода и ведёт раздел,
        а тот же профиль с чужим id живёт по общему правилу."""
        with mock.patch.object(access, 'ANALYST_USER_IDS', frozenset({77})):
            asked_without_the_list = 0
            for role, code, heads in grid():
                named = person(role=role, department_code=code, headed_codes=heads, user_id=77)
                self.assertEqual(access.capabilities(named), OWNER if role == 'super_admin' else FULL,
                                 (role, code, heads))
                self.assertFalse(access.requires_sensitive_qr(named), (role, code, heads))
                other = person(role=role, department_code=code, headed_codes=heads, user_id=78)
                # Общее правило — замок «Посылок»; раздел шире него ровно на
                # стажёра, которому раздел можно выдать (GrantAccessTests).
                expected = parcels_access.requires_sensitive_qr(other) or (role == 'trainee' and not heads)
                self.assertEqual(access.requires_sensitive_qr(other), expected, (role, code, heads))
                asked_without_the_list += access.requires_sensitive_qr(other)
            # Сетка не вырождена: среди профилей есть те, кого замок спрашивает.
            self.assertGreater(asked_without_the_list, 10)

    def test_analyst_id_is_read_as_a_number(self):
        with mock.patch.object(access, 'ANALYST_USER_IDS', frozenset({77})):
            for user_id, named in (('77', True), (77.0, True), (None, False), ('', False), ('77a', False)):
                ctx = dict(person(department_code='tez'), user_id=user_id)
                self.assertEqual(access.can_manage(ctx), named, repr(user_id))
                self.assertEqual(access.requires_sensitive_qr(ctx), not named, repr(user_id))

    def test_switch_closes_the_analyst_too(self):
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True):
            analyst = person(role='operator', department_code='analytik', user_id=ANALYST_ID)
            self.assertEqual(access.capabilities(analyst), CLOSED)

    def test_nobody_but_the_named_circle_enters_without_qr(self):
        """Главное свойство раздела: мимо QR-замка входят только супер-админ,
        главы трёх отделов и супервайзеры ОП и СЗоВ — на всех сочетаниях роли,
        отдела и главенства."""
        entered = 0
        for role, code, heads in grid():
            ctx = person(role=role, department_code=code, headed_codes=heads)
            if not access.can_open_section(ctx):
                continue
            entered += 1
            if access.requires_sensitive_qr(ctx):
                # Рядовой: только чтение и только отделы раздела.
                self.assertEqual(access.capabilities(ctx), READ, (role, code, heads))
                self.assertIn(code, access.SECTION_DEPARTMENT_CODES, (role, code, heads))
                continue
            named = (role == 'super_admin'
                     or bool(set(heads) & set(access.SECTION_DEPARTMENT_CODES))
                     or (role in ('sv', 'supervisor') and code in access.READ_DEPARTMENT_CODES))
            self.assertTrue(named, (role, code, heads))
        self.assertGreater(entered, 50)

    def test_only_full_access_exports_and_manages(self):
        for role, code, heads in grid():
            caps = access.capabilities(person(role=role, department_code=code, headed_codes=heads))
            full = role == 'super_admin' or 'marketing' in heads
            self.assertEqual((caps['can_export'], caps['can_manage']), (full, full), (role, code, heads))


class GrantAccessTests(unittest.TestCase):
    """Выдачи из раздела — кнопка «Доступ» (решение владельца 07.10.2026):
    человеку, группе, отделу. В контексте человека они лежат уровнями
    (grant_levels) — кому именно выдано, решает запрос (queries.load_access_context)."""

    def setUp(self):
        patcher = mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_levels_are_a_ladder(self):
        self.assertEqual(access.LEVELS, ('read', 'export', 'full'))
        self.assertEqual(access.SUBJECT_TYPES, ('user', 'group', 'department'))
        for level, caps in (('read', READ), ('export', EXPORT), ('full', FULL)):
            self.assertEqual(access.capabilities(person(department_code='tez', grants=(level,))), caps, level)

    def test_grant_opens_the_section_whatever_the_role_and_department(self):
        for role, code, heads in grid():
            ctx = person(role=role, department_code=code, headed_codes=heads, grants=('read',))
            self.assertTrue(access.can_open_section(ctx), (role, code, heads))

    def test_grant_only_adds_to_the_circle(self):
        """Выдача уровень, положенный по кругу, не понижает и круг не сужает."""
        for role, code, heads in grid():
            base = person(role=role, department_code=code, headed_codes=heads)
            by_circle = access.level_of(base)
            for level in access.LEVELS:
                granted = access.level_of(dict(base, grant_levels=[level]))
                self.assertEqual(granted, access.strongest([by_circle, level]), (role, code, heads, level))

    def test_the_strongest_grant_wins_and_junk_opens_nothing(self):
        self.assertEqual(access.capabilities(person(department_code='tez', grants=('read', 'full', 'export'))), FULL)
        # Опечатка в базе не открывает ничего — ни сама, ни рядом с настоящим уровнем.
        for junk in (('owner',), ('READ',), ('',), (None,), ('read ',)):
            self.assertEqual(access.capabilities(person(department_code='tez', grants=junk)), CLOSED, junk)
        self.assertEqual(access.capabilities(person(department_code='tez', grants=('owner', 'read'))), READ)
        # Контекст без поля (старый вызывающий) — выдач нет.
        bare = person(department_code='tez')
        del bare['grant_levels']
        self.assertEqual(access.capabilities(bare), CLOSED)

    def test_granted_rank_and_file_enter_through_qr(self):
        """Рядовой и по выдаче входит через замок — стажёр тоже: по кругу ему
        закрыто, а выдать раздел можно, и без замка он читал бы то, что оператор
        рядом открывает кодом."""
        for role in ('operator', 'trainee', 'accounting_manager', 'marketing_manager'):
            for code in GRID_DEPARTMENTS:
                ctx = person(role=role, department_code=code, grants=('full',))
                self.assertTrue(access.requires_sensitive_qr(ctx), (role, code))
                # Глава отдела — уже не рядовой, какой бы ни была должность в карточке.
                head = person(role=role, department_code=code, headed_codes=('tez',), grants=('full',))
                self.assertFalse(access.requires_sensitive_qr(head), (role, code))
        for role in ('sv', 'supervisor', 'trainer', 'admin', 'super_admin'):
            for code in GRID_DEPARTMENTS:
                self.assertFalse(access.requires_sensitive_qr(person(role=role, department_code=code,
                                                                     grants=('read',))), (role, code))

    def test_lock_asks_exactly_those_the_portal_gives_a_code_to(self):
        """Набор раздела — тот же, по которому портал выдаёт QR: спросить код у
        должности, которой его не выдают, значит запереть выданный раздел."""
        from driver_chats.access import QR_GATED_ROLES as wide
        from wiki.access import QR_GATED_ROLES as common
        self.assertEqual(access.QR_ASKED_ROLES, frozenset(common) | frozenset(wide))
        self.assertIn('SENSITIVE_QR_GATED_ROLES = frozenset(_WIKI_QR_GATED_ROLES) | '
                      'frozenset(_DRIVER_CHATS_QR_GATED_ROLES)\n', _read(BOT_PY))
        # Кадровик: кода ему портал не выдаёт (решение владельца 22.09.2026) —
        # раздел по выдаче открыт ему без замка, а по кругу не открыт вовсе.
        for code in GRID_DEPARTMENTS:
            self.assertFalse(access.requires_sensitive_qr(person(role='hr_manager', department_code=code,
                                                                 grants=('read',))), code)
            self.assertEqual(access.capabilities(person(role='hr_manager', department_code=code)), CLOSED, code)

    def test_every_role_of_the_portal_is_decided_for_the_lock(self):
        """Раздел выдают любому отделу, а значит — любой должности. Новая
        должность в портале обязана получить решение: спрашивать её о QR или
        нет. Без строки здесь она по выдаче отделу читала бы ФИО и номера ВУ
        без подтверждения — молча."""
        check = re.search(r"role VARCHAR\(32\) NOT NULL CHECK\(role IN \(([^)]*)\)\)", _read(DATABASE_PY))
        roles = set(re.findall(r"'([a-z_]+)'", check.group(1)))
        # Кого замок не спрашивает: старшие (подтверждают сами или им некому) и
        # кадровик, которому портал код не выдаёт (решение владельца 22.09.2026).
        not_asked = {'super_admin', 'admin', 'sv', 'supervisor', 'trainer', 'hr_manager'}
        self.assertEqual(roles, set(access.QR_ASKED_ROLES) | not_asked)
        self.assertFalse(set(access.QR_ASKED_ROLES) & not_asked)
        for role in roles:
            ctx = person(role=role, department_code='tez', grants=('read',))
            self.assertEqual(access.requires_sensitive_qr(ctx), role not in not_asked, role)

    def test_role_is_read_as_written_in_the_card(self):
        for role in (' Operator ', 'TRAINEE', 'Marketing_Manager'):
            self.assertTrue(access.requires_sensitive_qr(person(role=role, department_code='tez', grants=('read',))),
                            role)

    def test_switch_closes_the_granted_and_the_named(self):
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True), \
                mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({77})):
            self.assertEqual(access.capabilities(person(department_code='tez', grants=('full',))), CLOSED)
            self.assertEqual(access.capabilities(person(department_code='tez', user_id=77)), CLOSED)
            self.assertEqual(access.capabilities(person(role='super_admin', department_code=None)), OWNER)

    def test_access_is_handed_out_by_the_super_admin_and_the_named_only(self):
        """Кнопка «Доступ» — «у суперадминов» и названных поимённо. Ни полный
        доступ к разделу, ни главенство, ни выдача её не дают."""
        # 415 — глава «Маркетинга», постановщик раздела.
        self.assertEqual(access.ACCESS_MANAGER_USER_IDS, frozenset({415}))
        with mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({77})):
            for role, code, heads in grid():
                named = person(role=role, department_code=code, headed_codes=heads, user_id=77)
                self.assertTrue(access.can_manage_access(named), (role, code, heads))
                self.assertTrue(access.can_open_section(named), (role, code, heads))
                for grants in ((), ('full',)):
                    other = person(role=role, department_code=code, headed_codes=heads, user_id=78, grants=grants)
                    self.assertEqual(access.can_manage_access(other), role == 'super_admin',
                                     (role, code, heads, grants))

    def test_handing_out_access_is_not_full_access(self):
        """Названному раздел открыт на чтение — кнопка живёт внутри раздела, —
        а вести его он может только своим уровнем: по кругу или по выдаче."""
        with mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({77})):
            outsider = person(role='sv', department_code='tez', user_id=77)
            self.assertEqual(access.capabilities(outsider), dict(READ, can_manage_access=True))
            self.assertEqual(access.capabilities(dict(outsider, grant_levels=['full'])),
                             dict(FULL, can_manage_access=True))
            head = person(role='admin', department_code='marketing', headed_codes=('marketing',), user_id=77)
            self.assertEqual(access.capabilities(head), OWNER)
            # Замок — по должности, как у всех: раздающего-оператора он спрашивает.
            self.assertTrue(access.requires_sensitive_qr(person(department_code='tez', user_id=77)))
            self.assertFalse(access.requires_sensitive_qr(outsider))

    def test_manager_id_is_read_as_a_number(self):
        with mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({77})):
            for user_id, named in (('77', True), (77.0, True), (None, False), ('', False), ('77a', False)):
                ctx = dict(person(department_code='tez'), user_id=user_id)
                self.assertEqual(access.can_manage_access(ctx), named, repr(user_id))


# Кто подпадает под строку круга — записано здесь ещё раз, независимо от
# access._circle_slots: сторож обязан считать сам, а не спрашивать у проверяемого.
_RANK_AND_FILE = ('operator', 'marketing_manager', 'accounting_manager')


def in_circle_row(slot, role, code, heads):
    kind, key = slot.split(':')
    if kind == 'head':
        return key in heads
    if kind == 'sv':
        return role in ('sv', 'supervisor') and code == key
    if kind == 'staff':
        return role in _RANK_AND_FILE and not heads and code == key
    raise AssertionError(slot)


DEPARTMENT_ROWS = ('head:marketing', 'staff:marketing', 'head:op', 'sv:op', 'staff:op',
                   'head:szov', 'sv:szov', 'staff:szov')


class CircleEditTests(unittest.TestCase):
    """Круг раздела — «открыт по умолчанию» — правят из листа «Доступ»
    (решение владельца 08.10.2026: «сделай так чтобы можно было его
    редактировать»). Круг разложен на строки «должность в отделе»; правится
    уровень строки, а не её состав. Правки приходят в контексте: circle_levels."""

    def setUp(self):
        patcher = mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_defaults_are_the_circle_named_by_the_owner(self):
        """Без единой правки круг тот же, что владелец назвал 05–06.10.2026."""
        self.assertEqual(access.circle_defaults(), {
            'head:marketing': 'full', 'staff:marketing': 'read',
            'head:op': 'read', 'sv:op': 'read', 'staff:op': 'read',
            'head:szov': 'read', 'sv:szov': 'read', 'staff:szov': 'read',
            'named:%d' % ANALYST_ID: 'full',
        })
        self.assertEqual(access.CIRCLE_LEVELS, ('none',) + access.LEVELS)
        opened = 0
        for role, code, heads in grid():
            if role == 'super_admin' or 'marketing' in heads:
                expected = 'full'
            elif (set(heads) & {'op', 'szov'}
                  or (role in ('sv', 'supervisor') and code in ('op', 'szov'))
                  or (role in _RANK_AND_FILE and not heads and code in ('marketing', 'op', 'szov'))):
                expected = 'read'
            else:
                expected = None
            ctx = person(role=role, department_code=code, headed_codes=heads)
            self.assertEqual(access.level_of(ctx), expected, (role, code, heads))
            # Контекст без поля правок (старый вызывающий) — круг по умолчанию.
            del ctx['circle_levels']
            self.assertEqual(access.level_of(ctx), expected, (role, code, heads))
            opened += expected is not None
        self.assertGreater(opened, 50)

    def test_edit_reaches_the_people_of_the_row_and_nobody_else(self):
        """На всей сетке: правка строки меняет уровень тех, кто под неё
        подпадает, а у остальных не меняет ничего."""
        defaults = access.circle_defaults()
        for slot in DEPARTMENT_ROWS:
            for level in access.CIRCLE_LEVELS:
                touched = 0
                for role, code, heads in grid():
                    before = access.level_of(person(role=role, department_code=code, headed_codes=heads))
                    after = access.level_of(person(role=role, department_code=code, headed_codes=heads,
                                                   circle={slot: level}))
                    if role == 'super_admin':
                        self.assertEqual(after, 'full', (slot, level))
                        continue
                    mine = [row for row in DEPARTMENT_ROWS if in_circle_row(row, role, code, heads)]
                    if slot not in mine:
                        self.assertEqual(after, before, (slot, level, role, code, heads))
                        continue
                    touched += 1
                    # Старший из строк человека: правленая — по правке, прочие — по умолчанию.
                    levels = [level if row == slot else defaults[row] for row in mine]
                    self.assertEqual(after, access.strongest(levels), (slot, level, role, code, heads))
                self.assertGreater(touched, 0, slot)

    def test_each_row_by_name(self):
        """Поимённо, без формул: кого какая строка открывает и закрывает."""
        people = {
            'head:marketing': person(role='admin', department_code='marketing', headed_codes=('marketing',)),
            'staff:marketing': person(role='marketing_manager', department_code='marketing'),
            'head:op': person(role='admin', department_code='op', headed_codes=('op',)),
            'sv:op': person(role='sv', department_code='op'),
            'staff:op': person(role='operator', department_code='op'),
            'head:szov': person(role='admin', department_code=None, headed_codes=('szov',)),
            'sv:szov': person(role='supervisor', department_code='szov'),
            'staff:szov': person(role='operator', department_code='szov'),
            'named:%d' % ANALYST_ID: person(role='operator', department_code='analytik', user_id=ANALYST_ID),
        }
        self.assertEqual(set(people), set(access.circle_defaults()))
        wanted = {'none': CLOSED, 'read': READ, 'export': EXPORT, 'full': FULL}
        for slot, ctx in people.items():
            for level, caps in wanted.items():
                self.assertEqual(access.capabilities(dict(ctx, circle_levels={slot: level})), caps, (slot, level))
                # Соседние строки от правки не меняются.
                for other, neighbour in people.items():
                    if other != slot:
                        self.assertEqual(access.level_of(dict(neighbour, circle_levels={slot: level})),
                                         access.circle_defaults()[other], (slot, level, other))

    def test_the_strongest_row_of_a_person_wins(self):
        both = person(role='sv', department_code='szov', headed_codes=('op',))     # глава ОП и СВ СЗоВ
        self.assertEqual(access.level_of(dict(both, circle_levels={'head:op': 'none'})), 'read')
        self.assertEqual(access.level_of(dict(both, circle_levels={'head:op': 'none', 'sv:szov': 'full'})), 'full')
        self.assertEqual(access.level_of(dict(both, circle_levels={'head:op': 'export'})), 'export')
        self.assertIsNone(access.level_of(dict(both, circle_levels={'head:op': 'none', 'sv:szov': 'none'})))

    def test_closed_row_and_grants_do_not_cancel_each_other(self):
        """Строку круга закрыли — выданное отдельной выдачей остаётся; выдача
        одному закрытую строку остальным не открывает."""
        closed = {'staff:op': 'none'}
        self.assertEqual(access.capabilities(person(department_code='op', circle=closed)), CLOSED)
        self.assertEqual(access.capabilities(person(department_code='op', circle=closed, grants=('export',))), EXPORT)
        self.assertEqual(access.capabilities(person(department_code='op', circle={'staff:op': 'full'},
                                                    grants=('read',))), FULL)

    def test_junk_in_the_table_closes_the_row_instead_of_resetting_it(self):
        for junk in ('owner', '', None, 'READ', 'read ', 0, True):
            ctx = person(department_code='op', circle={'staff:op': junk})
            self.assertEqual(access.capabilities(ctx), CLOSED, repr(junk))
        # Правки нет вовсе (ключа нет) — значение по умолчанию.
        self.assertEqual(access.capabilities(person(department_code='op', circle={'staff:szov': 'none'})), READ)

    def test_rows_outside_the_circle_open_nothing(self):
        """Правится уровень строки, а не состав круга: строка, которой в круге
        нет (чужой отдел, должность без строки), никому ничего не открывает."""
        edits = {'staff:tez': 'full', 'head:tez': 'full', 'sv:marketing': 'full', 'staff:front_office': 'full',
                 'trainer:op': 'full', 'named:78': 'full', 'super_admin': 'none'}
        for kwargs in ({'department_code': 'tez'}, {'role': 'admin', 'department_code': 'tez', 'headed_codes': ('tez',)},
                       {'role': 'sv', 'department_code': 'marketing'}, {'department_code': 'front_office'},
                       {'role': 'trainer', 'department_code': 'op'}, {'department_code': 'analytik', 'user_id': 78},
                       {'role': 'trainee', 'department_code': 'szov'}):
            self.assertEqual(access.capabilities(person(circle=edits, **kwargs)), CLOSED, kwargs)
        for slot in edits:
            self.assertNotIn(slot, access.circle_defaults())

    def test_super_admin_cannot_be_locked_out(self):
        everything_closed = dict.fromkeys(list(access.circle_defaults()) + ['super_admin'], 'none')
        self.assertEqual(access.capabilities(person(role='super_admin', department_code=None,
                                                    circle=everything_closed)), OWNER)
        self.assertNotIn('super_admin', access.circle_defaults())
        # Остальным при этом закрыто всё — на всей сетке.
        for role, code, heads in grid():
            if role != 'super_admin':
                ctx = person(role=role, department_code=code, headed_codes=heads, circle=everything_closed)
                self.assertEqual(access.capabilities(ctx), CLOSED, (role, code, heads))

    def test_named_row_follows_the_named_list(self):
        with mock.patch.object(access, 'ANALYST_USER_IDS', frozenset({77})):
            self.assertIn('named:77', access.circle_defaults())
            self.assertNotIn('named:%d' % ANALYST_ID, access.circle_defaults())
            named = person(department_code='tez', user_id=77)
            self.assertEqual(access.capabilities(named), FULL)
            self.assertEqual(access.capabilities(dict(named, circle_levels={'named:77': 'read'})), READ)
            self.assertEqual(access.capabilities(dict(named, circle_levels={'named:77': 'none'})), CLOSED)
            # Без замка он и тогда, когда раздел достался ему иначе — выдачей.
            granted = dict(named, circle_levels={'named:77': 'none'}, grant_levels=['read'])
            self.assertEqual(access.capabilities(granted), READ)
            self.assertFalse(access.requires_sensitive_qr(granted))
            # Чужая именная строка его не касается.
            self.assertEqual(access.capabilities(dict(named, circle_levels={'named:78': 'none'})), FULL)

    def test_the_one_who_hands_out_access_keeps_the_section(self):
        """Раздающий не запирает сам себя: закрыл свою строку — чтение и кнопка
        «Доступ» остаются, вернуть уровень он может там же."""
        everything_closed = dict.fromkeys(access.circle_defaults(), 'none')
        with mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({77})):
            head = person(role='admin', department_code='marketing', headed_codes=('marketing',), user_id=77,
                          circle=everything_closed)
            self.assertEqual(access.capabilities(head), dict(READ, can_manage_access=True))

    def test_lock_does_not_depend_on_the_circle(self):
        edits = {'staff:op': 'full', 'head:op': 'none', 'sv:szov': 'export'}
        for role, code, heads in grid():
            plain = person(role=role, department_code=code, headed_codes=heads)
            self.assertEqual(access.requires_sensitive_qr(dict(plain, circle_levels=edits)),
                             access.requires_sensitive_qr(plain), (role, code, heads))

    def test_switch_is_stronger_than_the_circle(self):
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True):
            self.assertEqual(access.capabilities(person(department_code='op', circle={'staff:op': 'full'})), CLOSED)

    def test_rows_for_the_sheet_describe_the_real_rule(self):
        """Строки «открыт по умолчанию» в листе — не пересказ: порядок и состав
        берутся из того же круга, уровень — как его посчитает правило."""
        rows = access.circle()
        self.assertEqual([row['slot'] for row in rows],
                         ['super_admin'] + list(DEPARTMENT_ROWS) + ['named:%d' % ANALYST_ID])
        self.assertEqual(rows[0], {'slot': 'super_admin', 'kind': 'super_admin', 'level': 'full', 'qr': False,
                                   'locked': True})
        # Правятся все строки, кроме строки супер-админа, — ровно те, что знает ручка.
        self.assertEqual({row['slot'] for row in rows if not row['locked']}, set(access.circle_defaults()))
        self.assertEqual([row['locked'] for row in rows], [True] + [False] * (len(rows) - 1))
        by_slot = {row['slot']: row for row in rows}
        self.assertEqual(by_slot['staff:op'], {'slot': 'staff:op', 'kind': 'staff', 'code': 'op', 'level': 'read',
                                               'qr': True, 'locked': False})
        self.assertEqual(by_slot['named:%d' % ANALYST_ID]['user_id'], ANALYST_ID)
        # «После QR» — у строк рядовых, и замок правила с этим согласен.
        for row in rows[1:]:
            self.assertEqual(row['qr'], row['kind'] == 'staff', row['slot'])
        self.assertTrue(access.requires_sensitive_qr(person(department_code='op')))
        self.assertFalse(access.requires_sensitive_qr(person(role='sv', department_code='op')))
        # Уровень строки — с правками: закрытая и испорченная показаны как «нет».
        edited = {row['slot']: row['level'] for row in access.circle({'staff:op': 'none', 'head:szov': 'export',
                                                                      'sv:op': 'owner'})}
        self.assertEqual((edited['staff:op'], edited['head:szov'], edited['sv:op'], edited['staff:szov']),
                         ('none', 'export', 'none', 'read'))
        for slot, level in edited.items():
            if slot in access.circle_defaults() and slot.split(':')[0] != 'named':
                kind, code = slot.split(':')
                sample = {'head': person(role='admin', department_code=code, headed_codes=(code,)),
                          'sv': person(role='sv', department_code=code),
                          'staff': person(role='operator', department_code=code)}[kind]
                got = access.level_of(dict(sample, circle_levels={'staff:op': 'none', 'head:szov': 'export',
                                                                  'sv:op': 'owner'}))
                self.assertEqual(got or 'none', level, slot)


# ─────────────────────────────────────────────────────────────────────────────
# Разбор файла
# ─────────────────────────────────────────────────────────────────────────────

class ParseHappyPathTests(unittest.TestCase):

    def setUp(self):
        self.result = parse.parse_workbook(standard_book(), NAME)

    def test_week_is_read_whole(self):
        result = self.result
        self.assertTrue(result['ok'])
        self.assertEqual(result['errors'], [])
        self.assertEqual(result['period'], {'start': date(2026, 9, 21), 'end': date(2026, 9, 27),
                                            'source': 'name', 'week': 39})
        self.assertEqual([(sheet['name'], sheet['rows']) for sheet in result['sheets']],
                         [('Алматы-Каскелен', 3), ('МОТО БАЙГА', 1)])
        self.assertEqual(result['stats'], {'rows': 4, 'drivers': 4, 'prize_rows': 3, 'prize_total': 215000})

    def test_sheet_name_becomes_zachet_and_row_keeps_excel_number(self):
        first = self.result['rows'][0]
        self.assertEqual((first['zachet'], first['row_number'], first['sheet_order']), ('Алматы-Каскелен', 2, 0))
        moto = self.result['rows'][-1]
        self.assertEqual((moto['zachet'], moto['sheet_order']), ('МОТО БАЙГА', 1))

    def test_date_stays_text_as_in_the_file(self):
        row = self.result['rows'][0]
        self.assertEqual((row['date_text'], row['row_date'], row['week_number']), ('21.09.26', date(2026, 9, 21), 39))

    def test_prize_number_is_extracted(self):
        prizes = [(row['prize_name'], row['has_prize'], row['prize_amount']) for row in self.result['rows']]
        self.assertEqual(prizes[0], ('200 000 тенге', True, 200000))
        self.assertEqual(prizes[2], ('Не указано', False, None))

    def test_kazakh_letters_are_folded_for_search(self):
        moto = self.result['rows'][-1]
        self.assertEqual(moto['driver_name'], 'Жүсіпов А. Б.')
        self.assertIn('жусипов', moto['search_text'])
        self.assertIn(fake_id(101), moto['search_text'])

    def test_rows_are_not_in_the_public_summary(self):
        summary = parse.public_summary(self.result)
        self.assertNotIn('rows', summary)
        self.assertEqual(summary['period']['start'], '2026-09-21')


class ParseFormatTests(unittest.TestCase):

    def test_columns_are_found_by_name_not_place(self):
        order = [5, 0, 10, 3, 1, 2, 4, 6, 7, 8, 9]
        header = [HEADER[i] for i in order] + [None, None]
        row = line(1)
        content = book({'Астана': [['Итоги Байги'], header, [row[i] for i in order] + [None, None]]})
        result = parse.parse_workbook(content, NAME)
        self.assertTrue(result['ok'], result['errors'])
        self.assertEqual(result['rows'][0]['row_number'], 3)
        self.assertEqual(result['rows'][0]['driver_id'], fake_id(1))

    def test_license_written_as_a_number_is_read_as_text(self):
        content = book({'Астана': [HEADER, line(1, license_value=9912345678)]})
        row = parse.parse_workbook(content, NAME)['rows'][0]
        self.assertEqual((row['license'], row['license_key']), ('9912345678', '9912345678'))

    def test_period_from_underscored_name(self):
        self.assertEqual(parse.period_from_name('Список_байги_-_14_09_2026_-_20_09_2026.xlsx'),
                         (date(2026, 9, 14), date(2026, 9, 20)))
        self.assertEqual(parse.period_from_name('-_21.09.2026_-_27.09.2026.xlsx'),
                         (date(2026, 9, 21), date(2026, 9, 27)))
        self.assertIsNone(parse.period_from_name('байга итог.xlsx'))
        self.assertIsNone(parse.period_from_name('xlsx'))

    def test_range_notation_is_not_read_as_year_2027(self):
        """«21.09-27.09.2026»: кусок «21.09-27» раньше читался как 21.09.2027."""
        self.assertEqual(parse.dates_from_name('Список байги 21.09-27.09.2026.xlsx'),
                         [date(2026, 9, 21), date(2026, 9, 27)])
        self.assertEqual(parse.dates_from_name('Байга 1.09-7.09.2026.xlsx'), [date(2026, 9, 1), date(2026, 9, 7)])

    def test_period_must_match_the_date_column(self):
        week = [date(2026, 9, 21)] * 5
        period, error = parse.resolve_period('Список байги - 21.09.2026 - 27.09.2026.xlsx', week)
        self.assertEqual((period['start'], period['end'], period['source'], error),
                         (date(2026, 9, 21), date(2026, 9, 27), 'name', None))
        _, error = parse.resolve_period('Список байги - 14.09.2026 - 20.09.2026.xlsx', week)
        self.assertIn('не совпадает с колонкой «Дата»', error)
        _, error = parse.resolve_period('Список байги - 20.09.2026 - 27.09.2026.xlsx', week)
        self.assertIn('не из 7 дней', error)

    def test_single_date_in_name_is_start_or_end_by_the_data(self):
        week = [date(2026, 9, 21)] * 3
        start, _ = parse.resolve_period('Список байги - 21.09.2026.xlsx', week)
        end, _ = parse.resolve_period('Список байги - 27.09.2026.xlsx', week)
        self.assertEqual((start['start'], start['end']), (date(2026, 9, 21), date(2026, 9, 27)))
        self.assertEqual((end['start'], end['end']), (date(2026, 9, 21), date(2026, 9, 27)))
        _, error = parse.resolve_period('Список байги - 02.10.2026.xlsx', week)
        self.assertIn('не совпадает', error)

    def test_period_from_data_takes_the_common_date_not_a_typo(self):
        period, error = parse.resolve_period('байга итог.xlsx', [date(2026, 9, 21)] * 5 + [date(2025, 9, 21)])
        self.assertEqual((period['start'], period['source'], error), (date(2026, 9, 21), 'data', None))

    def test_wrong_period_in_name_blocks_the_upload(self):
        result = parse.parse_workbook(standard_book(), 'Список байги 21.09.2027 - 27.09.2027.xlsx')
        self.assertFalse(result['ok'])
        self.assertIn('не совпадает', result['errors'][0]['message'])

    def test_period_from_dates_when_name_has_none(self):
        result = parse.parse_workbook(standard_book(), 'байга итог.xlsx')
        self.assertTrue(result['ok'])
        self.assertEqual((result['period']['start'], result['period']['end'], result['period']['source']),
                         (date(2026, 9, 21), date(2026, 9, 27), 'data'))
        self.assertEqual([group['kind'] for group in result['warnings']], ['period_from_data'])

    def test_not_an_xlsx_is_refused_with_a_reason(self):
        cases = [(b'', 'BAIGA_FILE_EMPTY'), (b'\xd0\xcf\x11\xe0' + b'0' * 20, 'BAIGA_FILE_XLS'),
                 (b'just text', 'BAIGA_FILE_NOT_XLSX'), (b'PK\x03\x04broken', 'BAIGA_FILE_BROKEN')]
        for content, code in cases:
            with self.assertRaises(parse.ParseError) as caught:
                parse.parse_workbook(content, NAME)
            self.assertEqual(caught.exception.code, code)
        with mock.patch.object(parse, 'MAX_FILE_BYTES', 10):
            with self.assertRaises(parse.ParseError) as caught:
                parse.parse_workbook(standard_book(), NAME)
            self.assertEqual(caught.exception.code, 'BAIGA_FILE_TOO_BIG')


class ParseErrorTests(unittest.TestCase):

    def test_numbers_beyond_the_columns_are_errors_not_a_500_later(self):
        row = line(1)
        row[6] = 99999999999
        errors = self.errors([HEADER, row])
        self.assertEqual((errors[0]['message'], errors[0]['column']), ('Слишком большое число', 'Поездок'))

    def test_canonical_title_beats_an_alias_and_duplicates_block(self):
        header = ['ID'] + HEADER
        data = ['777'] + line(1)
        result = parse.parse_workbook(book({'Астана': [header, data]}), NAME)
        self.assertTrue(result['ok'], result['errors'])
        self.assertEqual(result['rows'][0]['driver_id'], fake_id(1))
        errors = self.errors([HEADER + ['Сумма'], line(1) + [5]])
        self.assertEqual(errors[0]['message'], 'Колонка встречается дважды: «Сумма»')

    def test_sheet_with_data_but_without_header_blocks(self):
        result = parse.parse_workbook(book({'Астана': [HEADER, line(1)], 'МОТО БАЙГА': [line(2), line(3)]}), NAME)
        self.assertFalse(result['ok'])
        self.assertEqual(result['errors'][0]['sheet'], 'МОТО БАЙГА')

    def test_broken_sheet_is_a_file_error_not_a_500(self):
        stream = BytesIO(standard_book())
        source = zipfile.ZipFile(stream)
        broken = BytesIO()
        with zipfile.ZipFile(broken, 'w') as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    data = data[:len(data) // 2]
                target.writestr(item, data)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_workbook(broken.getvalue(), NAME)
        self.assertEqual(caught.exception.code, 'BAIGA_FILE_BROKEN')

    def test_password_protected_file_says_so(self):
        content = b'\xd0\xcf\x11\xe0' + b'\x00' * 64 + 'EncryptedPackage'.encode('utf-16-le')
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_workbook(content, NAME)
        self.assertEqual(caught.exception.code, 'BAIGA_FILE_ENCRYPTED')

    def errors(self, rows, title='Астана'):
        result = parse.parse_workbook(book({title: rows}), NAME)
        self.assertFalse(result['ok'])
        return result['errors']

    def test_missing_column_names_the_sheet(self):
        errors = self.errors([HEADER[:-1], line(1)[:-1]])
        self.assertEqual(errors, [{'message': 'Нет колонок: «ID водителя»', 'sheet': 'Астана', 'row': 1}])

    def test_empty_required_cell_names_sheet_row_and_column(self):
        row = line(1)
        row[3] = None
        errors = self.errors([HEADER, line(1, driver=fake_id(9)), row])
        self.assertEqual(errors, [{'message': 'Пустая ячейка', 'sheet': 'Астана', 'row': 3, 'column': 'Водитель'}])

    def test_not_a_number_and_bad_date_block(self):
        bad = line(2, date_text='31.02.26')
        bad[5] = '12а'
        errors = self.errors([HEADER, line(1), bad])
        self.assertEqual([(error['row'], error['column']) for error in errors], [(3, 'Сумма'), (3, 'Дата')])

    def test_whole_numbers_with_spaces_are_numbers(self):
        row = line(1)
        row[5] = '1 234 567'
        result = parse.parse_workbook(book({'Астана': [HEADER, row]}), NAME)
        self.assertTrue(result['ok'], result['errors'])
        self.assertEqual(result['rows'][0]['amount'], 1234567)
        for text in ('12.5', '-3'):
            row[6] = text
            result = parse.parse_workbook(book({'Астана': [HEADER, row]}), NAME)
            self.assertFalse(result['ok'], text)

    def test_file_without_any_table_is_an_error(self):
        result = parse.parse_workbook(book({'Лист1': [['просто текст']]}), NAME)
        self.assertFalse(result['ok'])
        self.assertIn('нет таблицы Байги', result['errors'][0]['message'])


class ParseWarningTests(unittest.TestCase):

    def kinds(self, sheets, name=NAME):
        result = parse.parse_workbook(book(sheets), name)
        self.assertTrue(result['ok'], result['errors'])
        return {group['kind']: group for group in result['warnings']}

    def test_warnings_do_not_block_and_name_the_row(self):
        warnings = self.kinds({
            'Астана': [HEADER, line(1, license_value='ZZ0000001'), line(2, driver='abc'),
                       line(3, prize='Смартфон'), line(4, date_text='14.09.26'),
                       line(5, driver=fake_id(1), license_value='ZZ000005')],
            'Пустой': [['заметки аналитика'], [], ['ещё заметка']],
        })
        self.assertEqual(set(warnings), {'license', 'driver_id', 'prize', 'date_period', 'duplicate',
                                         'sheet_skipped'})
        self.assertEqual(warnings['duplicate']['items'][0],
                         {'message': 'Уже есть: лист «Астана», строка 2', 'sheet': 'Астана', 'row': 6,
                          'column': 'ID водителя'})
        self.assertEqual(warnings['license']['items'][0]['row'], 2)
        self.assertEqual(warnings['sheet_skipped']['items'][0]['sheet'], 'Пустой')

    def test_word_prize_counts_as_prize_without_amount(self):
        has, amount, recognised = parse.prize_amount('Смартфон')
        self.assertEqual((has, amount, recognised), (True, None, False))
        self.assertEqual(parse.prize_amount('Не указано'), (False, None, True))
        self.assertEqual(parse.prize_amount('10 000 тенге'), (True, 10000, True))
        self.assertEqual(parse.prize_amount('2000\u00a0тенге'), (True, 2000, True))

    def test_cyrillic_letters_in_license_are_the_same_number(self):
        self.assertEqual(parse.license_key('АВ 123-456'), 'AB123456')
        self.assertTrue(parse.is_standard_license(parse.license_key('ав123456')))


# ─────────────────────────────────────────────────────────────────────────────
# Поиск
# ─────────────────────────────────────────────────────────────────────────────

class FilterTests(unittest.TestCase):

    def test_empty_filters_select_everything(self):
        where, params = filters.where(filters.normalize({}))
        self.assertEqual((where, params), ('TRUE', {}))

    def test_user_text_never_reaches_sql(self):
        chosen = filters.normalize({'q': "Жүсіп'; DROP TABLE x; --", 'city': "Шымкент'--",
                                    'zachet': 'Астана', 'park': 'Парк', 'prize': '5000 тенге'})
        where, params = filters.where(chosen)
        self.assertNotIn('DROP', where)
        self.assertNotIn('Шымкент', where)
        self.assertNotIn('%%', where)
        self.assertEqual(params['city'], "Шымкент'--")
        self.assertEqual(params['w0'], '%жусип\';%')

    def test_kazakh_letters_and_case_do_not_matter(self):
        _, params = filters.where(filters.normalize({'q': 'ЖҮСІП'}))
        self.assertEqual(params['w0'], '%жусип%')

    def test_license_typed_in_russian_layout_is_also_searched_in_latin(self):
        where, params = filters.where(filters.normalize({'q': 'ав123'}))
        self.assertEqual((params['w0'], params['w0l']), ('%ав123%', '%ab123%'))
        self.assertIn(' OR ', where)

    def test_like_wildcards_are_escaped(self):
        _, params = filters.where(filters.normalize({'q': '100%_'}))
        self.assertEqual(params['w0'], '%100\\%\\_%')

    def test_ranges_and_prize_modes(self):
        chosen = filters.normalize({'amount_min': '500 000', 'amount_max': 100, 'position_max': '10',
                                    'prize': 'any', 'trips_min': 'много'})
        where, params = filters.where(chosen)
        self.assertEqual((params['amount_min'], params['amount_max'], params['position_max']), (100, 500000, 10))
        self.assertNotIn('trips_min', params)
        self.assertIn('r.has_prize', where)
        where, _ = filters.where(filters.normalize({'prize': 'none'}))
        self.assertEqual(where, 'NOT r.has_prize')

    def test_period_is_a_date(self):
        _, params = filters.where(filters.normalize({'period': '2026-09-21'}))
        self.assertEqual(params['period'], date(2026, 9, 21))
        self.assertNotIn('period', filters.where(filters.normalize({'period': 'вчера'}))[1])

    def test_pasted_list_is_split_like_a_human_means_it(self):
        text = '%s, ZZ000002\nAB 123456\n\nZZ000003 ZZ000004;%s' % (fake_id(1), fake_id(1).upper())
        self.assertEqual(filters.split_tokens(text),
                         [fake_id(1), 'ZZ000002', 'AB 123456', 'ZZ000003', 'ZZ000004', fake_id(1).upper()])

    def test_pasted_list_accepts_fleet_links_and_cyrillic_licenses(self):
        link = 'https://fleet.yandex.kz/contractors?park_id=%s&contractor_id=%s' % (fake_id(900), fake_id(5))
        drivers, licenses, key_of = filters.classify_tokens([fake_id(1).upper(), link, 'ав123456', '9912345678'])
        self.assertIn(fake_id(1), drivers)
        self.assertIn(fake_id(5), drivers)
        self.assertNotIn(fake_id(900), drivers)
        self.assertEqual(licenses, ['9912345678', 'AB123456'])
        self.assertEqual(key_of[link], [('driver', fake_id(5))])

    def test_id_only_list_does_not_match_empty_licences(self):
        """Заглушка [''] находила каждую строку с пустым номером ВУ — чужих водителей."""
        where, params = filters.where(filters.normalize({'list': fake_id(1)}))
        self.assertEqual(where, '(r.driver_key = ANY(%(list_ids)s))')
        self.assertEqual(params, {'list_ids': [fake_id(1)]})
        self.assertNotIn('', params['list_ids'])

    def test_non_standard_id_is_found_as_an_id(self):
        drivers, licenses, _ = filters.classify_tokens(['7d1f0c2e9a'])
        self.assertIn('7d1f0c2e9a', drivers)
        self.assertEqual(filters.not_found(['7d1f0c2e9a'], {'7d1f0c2e9a'}, set()), [])

    def test_driver_card_looks_up_the_exact_id(self):
        where, params = filters.where(filters.normalize({'driver': 'ABC12'}))
        self.assertEqual((where, params), ('r.driver_key = %(driver)s', {'driver': 'abc12'}))

    def test_zachet_sort_is_one_order_across_weeks(self):
        sql, key, _ = filters.order_by('zachet', 'asc')
        self.assertTrue(sql.startswith('MIN(r.sheet_order) OVER (PARTITION BY r.zachet) ASC NULLS LAST, '
                                       'r.zachet ASC NULLS LAST'))
        self.assertEqual(key, 'zachet')

    def test_not_found_is_reported_as_pasted(self):
        tokens = [fake_id(1), 'ав123456', 'ZZ000009']
        self.assertEqual(filters.not_found(tokens, {fake_id(1)}, {'AB123456'}), ['ZZ000009'])

    def test_sort_is_whitelisted(self):
        sql, key, direction = filters.order_by('amount', 'desc')
        self.assertTrue(sql.startswith('r.amount DESC NULLS LAST'))
        self.assertEqual((key, direction), ('amount', 'desc'))
        sql, key, direction = filters.order_by('r.id; DROP', 'asc')
        self.assertTrue(sql.startswith('r.period_start DESC'))
        self.assertEqual((key, direction), ('week', 'desc'))

    def test_page_sizes_are_fixed(self):
        self.assertEqual(filters.page_args({'size': 500, 'page': 3}), (3, 500))
        self.assertEqual(filters.page_args({'size': 7, 'page': -1}), (1, 50))

    def test_journal_keeps_filters_without_the_list_itself(self):
        chosen = filters.normalize({'city': 'Шымкент', 'list': 'ZZ000001\nZZ000002', 'period': '2026-09-21'})
        self.assertEqual(filters.public(chosen),
                         {'city': 'Шымкент', 'period': '2026-09-21', 'list': {'count': 2}})


# ─────────────────────────────────────────────────────────────────────────────
# Выгрузка
# ─────────────────────────────────────────────────────────────────────────────

class ReportTests(unittest.TestCase):

    def setUp(self):
        self.content = standard_book()
        self.parsed = parse.parse_workbook(self.content, NAME)
        self.rows = sorted(self.parsed['rows'], key=lambda row: (row['sheet_order'], row['row_number']))

    def test_week_without_filters_loads_back_row_for_row(self):
        out = report.build(self.rows, report.MODE_SHEETS)
        name = report.file_name(date(2026, 9, 21), date(2026, 9, 27))
        self.assertEqual(name, NAME)
        back = parse.parse_workbook(out, name)
        self.assertTrue(back['ok'], back['errors'])
        keys = [key for key in self.rows[0] if key != 'row_date']
        self.assertEqual([[row[key] for key in keys] for row in back['rows']],
                         [[row[key] for key in keys] for row in self.rows])
        source = load_workbook(BytesIO(self.content))
        exported = load_workbook(BytesIO(out))
        self.assertEqual(source.sheetnames, exported.sheetnames)
        for left, right in zip(source.worksheets, exported.worksheets):
            self.assertEqual([[str(value) for value in row] for row in left.iter_rows(values_only=True)],
                             [[str(value) for value in row] for row in right.iter_rows(values_only=True)])

    def test_ids_and_licenses_are_text_header_frozen_with_filter(self):
        worksheet = load_workbook(BytesIO(report.build(self.rows))).worksheets[0]
        self.assertEqual(worksheet.freeze_panes, 'A2')
        self.assertEqual(worksheet.auto_filter.ref, 'A1:K4')
        self.assertEqual((worksheet['K2'].number_format, worksheet['J2'].number_format), ('@', '@'))
        self.assertIsInstance(worksheet['F2'].value, int)

    def test_green_corners_are_suppressed_on_every_sheet(self):
        archive = zipfile.ZipFile(BytesIO(report.build(self.rows)))
        sheets = sorted(name for name in archive.namelist() if name.startswith('xl/worksheets/sheet'))
        self.assertEqual(len(sheets), 2)
        for name in sheets:
            xml = archive.read(name).decode('utf-8')
            self.assertIn('numberStoredAsText="1" twoDigitTextYear="1"', xml)
            self.assertLess(xml.find('<pageMargins'), xml.find('<ignoredErrors'))

    def test_single_sheet_keeps_file_columns_and_adds_zachet_last(self):
        worksheet = load_workbook(BytesIO(report.build(self.rows, report.MODE_SINGLE))).worksheets[0]
        self.assertEqual(worksheet.title, report.SINGLE_SHEET_TITLE)
        self.assertEqual([cell.value for cell in worksheet[1]], HEADER + ['Зачёт'])
        self.assertEqual(worksheet.max_row, 5)

    def test_text_starting_with_equals_is_not_a_formula(self):
        # Исходник — со СТРОКОВЫМИ ячейками: append() у openpyxl сам сделал бы
        # из «=1+2» формулу без значения, а в файле аналитика это текст.
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Астана'
        for number, values in enumerate([HEADER, line(1, name='=1+2', prize='=HYPERLINK("x")', park='=A1')], 1):
            for column, value in enumerate(values, 1):
                cell = worksheet.cell(row=number, column=column, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'
        stream = BytesIO()
        workbook.save(stream)
        rows = parse.parse_workbook(stream.getvalue(), NAME)['rows']
        self.assertEqual(rows[0]['driver_name'], '=1+2')
        out = report.build(rows)
        archive = zipfile.ZipFile(BytesIO(out))
        self.assertNotIn('<f>', archive.read('xl/worksheets/sheet1.xml').decode('utf-8'))
        back = parse.parse_workbook(out, NAME)['rows'][0]
        self.assertEqual((back['driver_name'], back['prize_name'], back['park']), ('=1+2', '=HYPERLINK("x")', '=A1'))

    def test_sheet_titles_ignore_case_and_stay_within_31(self):
        rows = parse.parse_workbook(book({'МОТО БАЙГА': [HEADER, line(1)]}), NAME)['rows']
        twin = dict(rows[0], zachet='Мото Байга', driver_key='other')
        self.assertEqual(len(report.group_by_zachet([rows[0], twin])), 1)
        titles = report._sheet_titles(['А' * 40, 'а' * 40])
        self.assertEqual(len(set(t.casefold() for t in titles)), 2)
        self.assertTrue(all(len(title) <= 31 for title in titles))

    def test_file_name_for_many_weeks(self):
        self.assertEqual(report.file_name(), 'Список байги - все недели.xlsx')


# ─────────────────────────────────────────────────────────────────────────────
# Схема
# ─────────────────────────────────────────────────────────────────────────────

class _Cursor:
    def __init__(self):
        self.statements = []
        self.params = []

    def execute(self, sql, params=None):
        self.statements.append(' '.join(sql.split()))
        self.params.append(params)


class SchemaTests(unittest.TestCase):

    def test_tables_go_first_and_everything_is_idempotent(self):
        cursor = _Cursor()
        schema.init_baiga_schema(cursor)
        kinds = ['table' if 'CREATE TABLE' in sql else 'other' for sql in cursor.statements]
        self.assertEqual(kinds[:7], ['table'] * 7)
        self.assertNotIn('table', kinds[7:])
        ddl = [sql for sql in cursor.statements if 'SAVEPOINT' not in sql]
        self.assertTrue(all('IF NOT EXISTS' in sql for sql in ddl))

    def test_one_active_upload_per_week(self):
        ddl = ' '.join(' '.join(statement.split()) for statement in schema._STATEMENTS)
        self.assertIn("CREATE UNIQUE INDEX IF NOT EXISTS uq_baiga_uploads_active_week "
                      "ON baiga_uploads(campaign, period_start) WHERE status = 'active'", ddl)

    def test_one_grant_per_subject(self):
        ddl = ' '.join(' '.join(statement.split()) for statement in schema._STATEMENTS)
        # На этом ключе стоит и «повторная выдача меняет уровень» (ON CONFLICT),
        # и поиск выдач человека на каждом запросе раздела.
        self.assertIn('CREATE UNIQUE INDEX IF NOT EXISTS uq_baiga_access_grants_subject '
                      'ON baiga_access_grants(subject_type, subject_id)', ddl)
        self.assertIn('CREATE TABLE IF NOT EXISTS baiga_access_log', ddl)

    def test_grant_tables_are_asked_once_per_process_after_they_appear(self):
        """Проверка доступа идёт на каждом запросе раздела: «таблицы есть»
        запоминается на процесс, «нет» — нет (схема могла лечь следующим стартом)."""
        class Answering(_Cursor):
            def __init__(self, answers):
                super().__init__()
                self.answers = list(answers)

            def fetchone(self):
                return (self.answers.pop(0),)

        with mock.patch.dict(schema._grants_seen, {'ready': False}):
            cursor = Answering([False, False, True])
            self.assertEqual([schema.grants_ready(cursor) for _ in range(5)], [False, False, True, True, True])
            self.assertEqual(len(cursor.statements), 3)
            self.assertEqual(cursor.params, [{'table': 'public.baiga_access_log'}] * 3)
        # Таблица правок круга моложе таблиц выдач — и спрашивается отдельно:
        # не легла она, выдачи работают как работали.
        with mock.patch.dict(schema._circle_seen, {'ready': False}):
            cursor = Answering([False, True])
            self.assertEqual([schema.circle_ready(cursor) for _ in range(4)], [False, True, True, True])
            self.assertEqual(cursor.params, [{'table': 'public.baiga_access_circle'}] * 2)
            self.assertIn('SELECT to_regclass(%(table)s) IS NOT NULL', cursor.statements[0])
        # «Есть» одной таблицы за другую не отвечает. Выдачи уже видели, а таблица
        # правок круга не легла: общая память ответила бы «есть», контекст пошёл бы
        # в таблицу, которой нет, и раздел падал бы на каждом запросе.
        with mock.patch.dict(schema._grants_seen, {'ready': True}), \
                mock.patch.dict(schema._circle_seen, {'ready': False}):
            cursor = Answering([False])
            self.assertEqual((schema.grants_ready(cursor), schema.circle_ready(cursor)), (True, False))
            self.assertEqual(cursor.params, [{'table': 'public.baiga_access_circle'}])
            self.assertEqual((schema._grants_seen['ready'], schema._circle_seen['ready']), (True, False))
        with mock.patch.dict(schema._grants_seen, {'ready': False}), \
                mock.patch.dict(schema._circle_seen, {'ready': True}):
            cursor = Answering([False])
            self.assertEqual((schema.circle_ready(cursor), schema.grants_ready(cursor)), (True, False))
            self.assertEqual(cursor.params, [{'table': 'public.baiga_access_log'}])

    def test_circle_edits_are_one_row_per_slot(self):
        ddl = ' '.join(' '.join(statement.split()) for statement in schema._STATEMENTS)
        self.assertIn('CREATE TABLE IF NOT EXISTS baiga_access_circle ( slot VARCHAR(48) PRIMARY KEY, '
                      'level VARCHAR(16) NOT NULL,', ddl)

    def test_trigram_index_cannot_break_the_section(self):
        class Failing(_Cursor):
            def execute(self, sql, params=None):
                super().execute(sql, params)
                if 'pg_trgm' in sql and 'EXTENSION' in sql:
                    raise RuntimeError('нет расширения')

        cursor = Failing()
        schema.init_baiga_schema(cursor)
        self.assertEqual(cursor.statements[-1], 'ROLLBACK TO SAVEPOINT baiga_trgm')


class ContextTests(unittest.TestCase):
    """Контекст доступа: профиль, главенство и уровни выдач одним запросом.
    Сам SQL исполняется на стенде и в tests/test_baiga_access_postgres.py."""

    class Cursor:
        def __init__(self, row):
            self.row = row
            self.calls = []

        def execute(self, sql, params=None):
            self.calls.append((' '.join(sql.split()), params))

        def fetchone(self):
            return self.row

    def load(self, row, ready=True, user_id=31, circle=True):
        cursor = self.Cursor(row)
        with mock.patch.object(schema, 'grants_ready', lambda cursor: ready), \
                mock.patch.object(schema, 'circle_ready', lambda cursor: circle):
            return queries.load_access_context(cursor, user_id), cursor

    def test_role_stays_as_written_and_grant_levels_come_along(self):
        ctx, cursor = self.load(('Кадрова К. К.', ' HR_Manager ', 168, 'hr', None, [7], ['HR'], ['read', 'full'],
                                 {'staff:op': 'none'}), user_id='31')
        self.assertEqual(ctx, {
            'user_id': 31, 'name': 'Кадрова К. К.', 'role': 'hr_manager', 'department_id': 168,
            'department_code': 'hr', 'city': None, 'headed_department_ids': [7],
            'headed_department_codes': ['HR'], 'grant_levels': ['read', 'full'],
            'circle_levels': {'staff:op': 'none'},
        })
        # Замку нужна должность из карточки: «Посылки» свели бы кадровика к
        # оператору и спросили бы код, которого портал ему не выдаёт.
        self.assertEqual(parcels_access.normalize_role(ctx['role']), 'operator')
        self.assertFalse(access.requires_sensitive_qr(dict(ctx, headed_department_ids=[])))
        self.assertEqual(len(cursor.calls), 1)
        sql, params = cursor.calls[0]
        self.assertEqual(params, {'user_id': 31})
        # Три вида выдач: самому человеку, его группе, его отделу (и отделу,
        # который он возглавляет, числясь в другом).
        self.assertIn("(a.subject_type = 'user' AND a.subject_id = %(user_id)s)", sql)
        self.assertIn("(a.subject_type = 'group' AND a.subject_id IN (SELECT group_id FROM my_groups))", sql)
        self.assertIn("(a.subject_type = 'department' AND a.subject_id IN (SELECT id FROM my_departments))", sql)
        # Отдел — свой ДЕЙСТВУЮЩИЙ и возглавляемые (тоже действующие): закрытому
        # отделу выдача ничего не открывает, как архивной группе.
        self.assertIn("my_departments AS ( SELECT d.id FROM departments d WHERE d.is_active "
                      "AND d.id = (SELECT department_id FROM me) UNION SELECT id FROM headed )", sql)
        self.assertIn("headed AS ( SELECT d.id, d.code FROM departments d "
                      "WHERE d.head_user_id = %(user_id)s AND d.is_active )", sql)
        # Группы — ЭТОГО человека, уровни — из ЕГО выдач: без этих двух строк
        # одна выдача открывала бы раздел всему порталу.
        self.assertIn("WHERE m.user_id = %(user_id)s AND m.start_date <= CURRENT_DATE", sql)
        # Последние две колонки: уровни выдач ЭТОГО человека и правки круга (они одни на всех).
        self.assertTrue(sql.endswith(
            "COALESCE((SELECT array_agg(DISTINCT level) FROM my_grants), '{}'), "
            "COALESCE((SELECT json_object_agg(slot, level) FROM baiga_access_circle), '{}'::json)"), sql[-200:])
        self.assertIn("my_grants AS ( SELECT a.level FROM baiga_access_grants a WHERE (a.subject_type = 'user'", sql)

    def test_group_is_the_same_thing_as_in_the_wiki(self):
        """Состав группы — действующее членство оператора или супервайзера в
        действующей группе: то же определение, что у вики (my_groups)."""
        sql = ' '.join(queries._CONTEXT_WITH_GRANTS_SQL.split())
        wiki = ' '.join(_read(ROOT / 'wiki' / 'queries.py').split())
        for table, column in (('group_operator_memberships', 'operator_id'),
                              ('group_supervisor_memberships', 'supervisor_id')):
            self.assertIn('SELECT group_id, %s' % column, sql)
            self.assertIn('FROM %s' % table, sql)
            self.assertIn('FROM %s' % table, wiki)
        self.assertIn("JOIN groups g ON g.id = m.group_id AND g.status = 'active'", sql)
        self.assertIn('m.start_date <= CURRENT_DATE AND (m.end_date IS NULL OR m.end_date >= CURRENT_DATE)', sql)
        self.assertIn("JOIN groups g ON g.id = gom.group_id AND g.status = 'active'", wiki)
        self.assertIn('gom.start_date <= CURRENT_DATE AND (gom.end_date IS NULL OR gom.end_date >= CURRENT_DATE)', wiki)

    def test_without_grant_tables_the_context_is_the_circle_alone(self):
        ctx, cursor = self.load(('Оператов О. О.', 'operator', 191, 'op', None, [], [], [], {}), ready=False)
        self.assertEqual((ctx['grant_levels'], ctx['circle_levels']), ([], {}))
        self.assertEqual(access.capabilities(ctx), READ)
        sql = cursor.calls[0][0]
        self.assertNotIn('baiga_access_grants', sql)
        self.assertNotIn('baiga_access_circle', sql)
        self.assertNotIn('memberships', sql)
        self.assertTrue(sql.endswith("'{}'::varchar[], '{}'::json"), sql[-60:])

    def test_without_the_circle_table_grants_work_and_the_circle_is_the_default(self):
        """Таблица правок круга не легла, а таблицы выдач на месте: выдачи
        действуют, круг — по умолчанию; раздел не отвечает ошибкой каждому."""
        ctx, cursor = self.load(('Оператов О. О.', 'sv', 70, 'tez', None, [], [], ['export'], {}), circle=False)
        self.assertEqual((ctx['grant_levels'], ctx['circle_levels']), (['export'], {}))
        self.assertEqual(access.capabilities(ctx), EXPORT)
        sql = cursor.calls[0][0]
        self.assertIn('FROM baiga_access_grants a', sql)
        self.assertNotIn('baiga_access_circle', sql)
        self.assertTrue(sql.endswith("FROM my_grants), '{}'), '{}'::json"), sql[-60:])

    def test_unknown_user_has_no_context(self):
        self.assertIsNone(self.load(None)[0])
        self.assertIsNone(self.load((None, None, None, None, None, [], [], [], {}))[0])

    def test_staff_means_everyone_who_can_still_log_in(self):
        """Число людей под выдачей и список «кому выдать» считают тех, кто в
        портал входит: уволенный — нет, человек в отпуске — да."""
        self.assertEqual(queries._STAFF, "status NOT IN ('fired', 'dismissal')")
        login = _read(BOT_PY).split("user_status = str(user_profile[11] or '').strip().lower()")[1][:120]
        self.assertIn("if user_status in ('fired', 'dismissal'):", login)



class GrantQueryTests(unittest.TestCase):
    """Запросы выдач — текстом и параметрами, с курсором-записывателем. Что они
    ЗНАЧАТ на настоящей базе, исполняет tests/test_baiga_access_postgres.py
    (нужен локальный Postgres); здесь — сторож, который работает везде: условия
    живости адресата, блокировки и то, что пишется при смене уровня и снятии."""

    ACTOR = {'user_id': 415, 'name': 'Раздающий Р. Р.'}

    class Cursor:
        def __init__(self, rows=()):
            self.rows = list(rows)
            self.calls = []

        def execute(self, sql, params=None):
            self.calls.append((' '.join(sql.split()), params))

        def fetchall(self):
            return list(self.rows)

        def fetchone(self):
            return self.rows[0] if self.rows else None

    def batches(self):
        """Подмена пачечной записи: что и в какую таблицу ушло."""
        written = []

        def record(cursor, sql, values, **kwargs):
            written.append((' '.join(sql.split()), list(values)))

        patcher = mock.patch.object(queries, 'execute_values', record)
        patcher.start()
        self.addCleanup(patcher.stop)
        return written

    def test_only_living_subjects_are_found(self):
        cursor = self.Cursor([('user', 31, 'Оператов О. О.'), ('group', 5, 'Регионы')])
        found = queries.find_subjects(cursor, [('user', 31), ('group', 5), ('department', 44), ('user', '32')])
        self.assertEqual(found, {('user', 31): 'Оператов О. О.', ('group', 5): 'Регионы'})
        sql, params = cursor.calls[0]
        self.assertEqual(params, {'user': [31, 32], 'group': [5], 'department': [44]})
        # Уволенному, архивной группе и закрытому отделу не выдают.
        self.assertIn("WHERE u.id = ANY(%(user)s::int[]) AND u.status NOT IN ('fired', 'dismissal')", sql)
        self.assertIn("WHERE g.id = ANY(%(group)s::int[]) AND g.status = 'active'", sql)
        self.assertIn("WHERE d.id = ANY(%(department)s::int[]) AND d.is_active", sql)

    def test_catalog_offers_only_the_living(self):
        cursor = self.Cursor([('department', 44, 'Бухгалтерия', None, None, 1),
                              ('group', 5, 'Регионы', 'СЗоВ', None, 8),
                              ('user', 31, 'Оператов О. О.', 'Тез КЦ', 'operator', None)])
        catalog = queries.access_catalog(cursor)
        self.assertEqual(catalog, {
            'department': [{'id': 44, 'name': 'Бухгалтерия', 'detail': None, 'role': None, 'people': 1}],
            'group': [{'id': 5, 'name': 'Регионы', 'detail': 'СЗоВ', 'role': None, 'people': 8}],
            'user': [{'id': 31, 'name': 'Оператов О. О.', 'detail': 'Тез КЦ', 'role': 'operator', 'people': None}],
        })
        sql = cursor.calls[0][0]
        self.assertIn("staff AS ( SELECT id, name, role, department_id FROM users "
                      "WHERE status NOT IN ('fired', 'dismissal') )", sql)
        self.assertIn("FROM departments d WHERE d.is_active UNION ALL", sql)
        self.assertIn("WHERE g.status = 'active' UNION ALL", sql)
        self.assertIn("FROM staff s LEFT JOIN departments sd ON sd.id = s.department_id ORDER BY 1, 3, 2", sql)
        # Людей в группе считают по действующему членству тех, кто в штате.
        self.assertIn("JOIN staff s ON s.id = m.user_id WHERE m.start_date <= CURRENT_DATE "
                      "AND (m.end_date IS NULL OR m.end_date >= CURRENT_DATE) GROUP BY 1", sql)

    def test_sheet_marks_grants_that_open_nothing(self):
        cursor = self.Cursor()
        queries.list_grants(cursor)
        sql = cursor.calls[0][0]
        self.assertIn("CASE a.subject_type WHEN 'user' THEN COALESCE(u.status NOT IN ('fired', 'dismissal'), FALSE) "
                      "WHEN 'group' THEN COALESCE(g.status = 'active', FALSE) "
                      "ELSE COALESCE(d.is_active, FALSE) END", sql)
        self.assertIn("LEFT JOIN users u ON a.subject_type = 'user' AND u.id = a.subject_id", sql)
        self.assertIn("LEFT JOIN groups g ON a.subject_type = 'group' AND g.id = a.subject_id", sql)
        self.assertIn("LEFT JOIN departments d ON a.subject_type = 'department' AND d.id = a.subject_id", sql)

    def test_already_granted_levels_are_asked_in_one_query(self):
        cursor = self.Cursor([('user', 31, 'read')])
        self.assertEqual(queries.granted_levels(cursor, [('user', 31), ('group', '5')]), {('user', 31): 'read'})
        sql, params = cursor.calls[0]
        self.assertIn('WHERE (subject_type, subject_id) IN %(pairs)s', sql)
        self.assertEqual(params, {'pairs': (('user', 31), ('group', 5))})
        # Пустой список — без запроса: «IN ()» базу не спрашивают.
        empty = self.Cursor()
        self.assertEqual(queries.granted_levels(empty, []), {})
        self.assertEqual(empty.calls, [])

    def test_repeat_grant_rewrites_the_level_and_who_granted(self):
        written = self.batches()
        with mock.patch.object(queries, 'granted_levels', lambda cursor, subjects: {('user', 31): 'read'}):
            result = queries.write_grants(self.Cursor(), [('user', 31, 'Оператов О. О.'), ('group', 5, 'Регионы')],
                                          'full', self.ACTOR)
        self.assertEqual(result, (1, 1))
        (grants_sql, grants), (log_sql, log) = written
        self.assertIn('INSERT INTO baiga_access_grants (subject_type, subject_id, level, granted_by, granted_by_name) '
                      'VALUES %s ON CONFLICT (subject_type, subject_id) DO UPDATE SET level = EXCLUDED.level, '
                      'granted_by = EXCLUDED.granted_by, granted_by_name = EXCLUDED.granted_by_name, '
                      "granted_at = (CURRENT_TIMESTAMP AT TIME ZONE 'Asia/Almaty')", grants_sql)
        # Сначала новые, потом сменившие уровень — и у всех выдал тот, кто сохранил.
        self.assertEqual(grants, [('group', 5, 'full', 415, 'Раздающий Р. Р.'),
                                  ('user', 31, 'full', 415, 'Раздающий Р. Р.')])
        self.assertIn('INSERT INTO baiga_access_log (action, subject_type, subject_id, subject_label, level_before, '
                      'level_after, actor_user_id, actor_name) VALUES %s', log_sql)
        self.assertEqual(log, [('grant', 'group', 5, 'Регионы', None, 'full', 415, 'Раздающий Р. Р.'),
                               ('change', 'user', 31, 'Оператов О. О.', 'read', 'full', 415, 'Раздающий Р. Р.')])

    def test_same_level_again_writes_nothing(self):
        written = self.batches()
        with mock.patch.object(queries, 'granted_levels', lambda cursor, subjects: {('user', 31): 'full'}):
            self.assertEqual(queries.write_grants(self.Cursor(), [('user', 31, 'Оператов О. О.')], 'full', self.ACTOR),
                             (0, 0))
        self.assertEqual(written, [])

    def test_level_change_updates_this_grant_and_leaves_a_trace(self):
        written = self.batches()
        cursor = self.Cursor()
        grant = {'id': 7, 'subject_type': 'group', 'subject_id': 5, 'level': 'read', 'label': 'Регионы'}
        queries.set_grant_level(cursor, grant, 'export', self.ACTOR)
        sql, params = cursor.calls[0]
        self.assertIn('UPDATE baiga_access_grants SET level = %(level)s, granted_by = %(actor_id)s, '
                      'granted_by_name = %(actor_name)s, granted_at =', sql)
        self.assertTrue(sql.endswith('WHERE id = %(id)s'), sql)
        # В базу уходит НОВЫЙ уровень и номер этой выдачи.
        self.assertEqual(params, {'level': 'export', 'actor_id': 415, 'actor_name': 'Раздающий Р. Р.', 'id': 7})
        self.assertEqual(written[0][1], [('change', 'group', 5, 'Регионы', 'read', 'export', 415, 'Раздающий Р. Р.')])

    def test_revoke_deletes_this_grant_and_leaves_a_trace(self):
        written = self.batches()
        cursor = self.Cursor()
        grant = {'id': 7, 'subject_type': 'group', 'subject_id': 5, 'level': 'full', 'label': None}
        queries.delete_grant(cursor, grant, self.ACTOR)
        self.assertEqual(cursor.calls, [('DELETE FROM baiga_access_grants WHERE id = %(id)s', {'id': 7})])
        # Снятие — тоже в журнал, даже когда адресата уже нет и подписать его нечем.
        self.assertEqual(written[0][1], [('revoke', 'group', 5, None, 'full', None, 415, 'Раздающий Р. Р.')])

    def test_grant_is_read_under_a_row_lock_when_asked(self):
        row = (7, 'group', 5, 'read', 'Регионы')
        locked, plain = self.Cursor([row]), self.Cursor([row])
        self.assertEqual(queries.get_grant(locked, '7', lock=True),
                         {'id': 7, 'subject_type': 'group', 'subject_id': 5, 'level': 'read', 'label': 'Регионы'})
        queries.get_grant(plain, 7)
        self.assertTrue(locked.calls[0][0].endswith('WHERE a.id = %(id)s FOR UPDATE OF a'), locked.calls[0][0])
        self.assertTrue(plain.calls[0][0].endswith('WHERE a.id = %(id)s'), plain.calls[0][0])
        self.assertEqual(locked.calls[0][1], {'id': 7})
        self.assertIsNone(queries.get_grant(self.Cursor(), 7))

    def test_access_edits_take_one_lock_for_the_whole_section(self):
        cursor = self.Cursor()
        queries.lock_access(cursor)
        self.assertEqual(cursor.calls, [("SELECT pg_advisory_xact_lock(hashtext('baiga:access'))", None)])

    def test_circle_edits_are_read_whole(self):
        stamp = datetime(2026, 10, 8, 12, 0)
        cursor = self.Cursor([('staff:op', 'none', 'Раздающий Р. Р.', stamp), ('head:szov', 'export', None, stamp)])
        self.assertEqual(queries.circle_edits(cursor), {
            'staff:op': {'level': 'none', 'updated_by_name': 'Раздающий Р. Р.', 'updated_at': stamp},
            'head:szov': {'level': 'export', 'updated_by_name': None, 'updated_at': stamp}})
        self.assertEqual(cursor.calls, [('SELECT slot, level, updated_by_name, updated_at FROM baiga_access_circle',
                                         None)])
        self.assertEqual(queries.circle_edits(self.Cursor()), {})

    def test_circle_edit_is_one_row_per_slot_and_leaves_a_trace(self):
        written = self.batches()
        cursor = self.Cursor()
        queries.set_circle_level(cursor, 'staff:op', 'read', 'none', self.ACTOR)
        (sql, params), = cursor.calls
        # Строка на правленую строку круга: повторная правка переписывает уровень,
        # автора и время, а не кладёт вторую строку.
        self.assertEqual(sql, 'INSERT INTO baiga_access_circle (slot, level, updated_by, updated_by_name) '
                              'VALUES (%(slot)s, %(level)s, %(actor_id)s, %(actor_name)s) '
                              'ON CONFLICT (slot) DO UPDATE SET level = EXCLUDED.level, '
                              'updated_by = EXCLUDED.updated_by, updated_by_name = EXCLUDED.updated_by_name, '
                              "updated_at = (CURRENT_TIMESTAMP AT TIME ZONE 'Asia/Almaty')")
        # В базу уходит НОВЫЙ уровень; прежний — только в журнал.
        self.assertEqual(params, {'slot': 'staff:op', 'level': 'none', 'actor_id': 415,
                                  'actor_name': 'Раздающий Р. Р.'})
        (log_sql, log), = written
        self.assertIn('INSERT INTO baiga_access_log', log_sql)
        self.assertEqual(log, [('circle', 'circle', 0, 'staff:op', 'read', 'none', 415, 'Раздающий Р. Р.')])


# ─────────────────────────────────────────────────────────────────────────────
# Ручки — на подменённом слое запросов
# ─────────────────────────────────────────────────────────────────────────────

class Store:
    """Память вместо базы: ровно те функции queries, что зовут роуты."""

    def __init__(self):
        self.uploads = {}
        self.files = {}
        self.rows = []
        self.exports = []
        self.writes = []
        self.short_write = False
        self.export_result = None
        # Доступ: выдачи, их журнал и «живые» адресаты (вид, id) → подпись.
        self.grants = {}
        self.access_log = []
        self.access_locks = 0
        # Шаги правки доступа по порядку: блокировка обязана идти первой.
        self.access_steps = []
        self.subjects = {('user', 31): 'Оператов О. О.', ('user', 32): 'Тренеров Т. Т.',
                         ('group', 5): 'Регионы', ('department', 44): 'Бухгалтерия'}
        # Правки круга — таблица baiga_access_circle: {строка: уровень} и кто правил.
        self.circle = {}
        self.circle_by = {}
        # Таблицы правок круга нет (миграция не легла): любой запрос к ней падает.
        self.circle_missing = False
        # Названные поимённо, которых в штате уже нет.
        self.gone_people = set()

    # экран и поиск
    def list_weeks(self, cursor, campaign):
        return [dict(upload) for upload in self.uploads.values() if upload['status'] == 'active']

    def filter_options(self, cursor):
        return {'zachets': sorted({row['zachet'] for row in self.rows}), 'cities': [], 'parks': [], 'prizes': []}

    def totals(self, cursor, where_sql, params):
        return {'rows': len(self.rows), 'drivers': len({row['driver_key'] for row in self.rows}),
                'prize_rows': 0, 'prize_total': 0}

    @staticmethod
    def _screen_row(row):
        # Тот же список полей, что у настоящей выборки (queries.ROW_FIELDS).
        return {name: row.get(name) for name in queries.ROW_FIELDS}

    def search(self, cursor, where_sql, params, order_sql, limit, offset):
        self.last_search = (where_sql, params, order_sql, limit, offset)
        return [self._screen_row(row) for row in self.rows[offset:offset + limit]]

    def zachet_summary(self, cursor, where_sql, params):
        names = []
        for row in self.rows:
            if row['zachet'] not in names:
                names.append(row['zachet'])
        return [{'name': name, 'rows': sum(1 for row in self.rows if row['zachet'] == name),
                 'prize_rows': 0, 'prize_total': 0} for name in names]

    def find_overlapping_uploads(self, cursor, campaign, start, end):
        return [dict(upload) for upload in self.uploads.values()
                if upload['status'] == 'active' and upload['period_start'] != start
                and upload['period_start'] <= end and upload['period_end'] >= start]

    def found_keys(self, cursor, where_sql, params):
        return {row['driver_key'] for row in self.rows}, {row['license_key'] for row in self.rows}

    def export_rows(self, cursor, where_sql, params, order_sql, limit):
        self.last_export = (where_sql, params, order_sql, limit)
        rows = self.export_result if self.export_result is not None else self.rows
        return [self._screen_row(row) for row in rows][:limit + 1]

    # загрузка
    def lock_week(self, cursor, campaign, period_start):
        self.writes.append(('lock', period_start))

    def find_active_upload(self, cursor, campaign, period_start):
        for upload in self.uploads.values():
            if upload['status'] == 'active' and upload['period_start'] == period_start:
                return dict(upload, file_sha256=upload.get('sha'))
        return None

    def insert_upload(self, cursor, result, actor):
        upload_id = max(self.uploads, default=0) + 1
        self.uploads[upload_id] = {
            'id': upload_id, 'status': 'active', 'period_start': result['period']['start'],
            'period_end': result['period']['end'], 'rows_count': result['stats']['rows'],
            'file_name': result['file_name'], 'uploaded_by_name': actor['name'],
            'uploaded_at': datetime(2026, 10, 2, 12, 0), 'sha': result['sha256'],
        }
        self.writes.append(('insert_upload', upload_id))
        return upload_id

    def insert_file(self, cursor, upload_id, content):
        self.files[upload_id] = content
        self.writes.append(('insert_file', upload_id))

    def insert_rows(self, cursor, upload_id, period_start, rows):
        self.writes.append(('insert_rows', upload_id))
        for row in rows:
            self.rows.append(dict(row, upload_id=upload_id, period_start=period_start, id=len(self.rows) + 1))
        return len(rows) - 1 if self.short_write else len(rows)

    def close_upload(self, cursor, upload_id, status, actor, replaced_by=None):
        self.uploads[upload_id]['status'] = status
        self.rows = [row for row in self.rows if row.get('upload_id') != upload_id]
        self.files.pop(upload_id, None)
        self.writes.append(('close', upload_id, status))

    def mark_replaced_by(self, cursor, upload_id, replaced_by):
        self.uploads[upload_id]['replaced_by'] = replaced_by

    def get_upload(self, cursor, upload_id, lock=False):
        upload = self.uploads.get(int(upload_id))
        return dict(upload) if upload else None

    def read_file(self, cursor, upload_id):
        return self.files.get(int(upload_id))

    def list_uploads(self, cursor, limit=200):
        return [dict(upload) for upload in self.uploads.values()]

    def list_exports(self, cursor, limit=200):
        return list(self.exports)

    def log_export(self, cursor, **kwargs):
        self.exports.append(kwargs)

    # доступ
    def list_grants(self, cursor):
        return [dict(grant, label=self.subjects.get((grant['subject_type'], grant['subject_id'])),
                     detail=None, role=None, people=None, active=True)
                for grant in sorted(self.grants.values(), key=lambda grant: grant['id'])]

    def access_catalog(self, cursor):
        catalog = {'department': [], 'group': [], 'user': []}
        for (kind, ident), name in sorted(self.subjects.items()):
            catalog[kind].append({'id': ident, 'name': name, 'detail': None, 'role': None, 'people': None})
        return catalog

    def circle_names(self, cursor, codes, user_ids):
        return ({code: 'Отдел %s' % code for code in codes},
                {user_id: 'Названный %d' % user_id for user_id in user_ids if user_id not in self.gone_people})

    def _circle_table(self):
        if self.circle_missing:
            raise RuntimeError('relation "baiga_access_circle" does not exist')

    def circle_edits(self, cursor):
        self._circle_table()
        self.access_steps.append('circle:read')
        return {slot: {'level': level, 'updated_by_name': self.circle_by.get(slot),
                       'updated_at': datetime(2026, 10, 8, 12, 0) if slot in self.circle_by else None}
                for slot, level in self.circle.items()}

    def set_circle_level(self, cursor, slot, before, level, actor):
        self._circle_table()
        self.access_steps.append('circle:set')
        self.circle[slot] = level
        self.circle_by[slot] = actor.get('name')
        self.access_log.append({'action': 'circle', 'subject_type': 'circle', 'subject_id': 0, 'label': slot,
                                'before': before, 'after': level,
                                'actor_id': actor['user_id'], 'actor_name': actor.get('name')})

    def find_subjects(self, cursor, subjects):
        self.access_steps.append('find')
        return {subject: self.subjects[subject] for subject in subjects if subject in self.subjects}

    def lock_access(self, cursor):
        self.access_locks += 1
        self.access_steps.append('lock')

    def granted_levels(self, cursor, subjects):
        by_subject = {(grant['subject_type'], grant['subject_id']): grant['level'] for grant in self.grants.values()}
        return {subject: by_subject[subject] for subject in subjects if subject in by_subject}

    def execute_values(self, cursor, sql, values, **kwargs):
        """Пачечная запись — как её зовёт настоящий queries.write_grants."""
        if 'baiga_access_grants' in sql:
            self.access_steps.append('write')
            for kind, ident, level, actor_id, actor_name in values:
                found = next((grant for grant in self.grants.values()
                              if (grant['subject_type'], grant['subject_id']) == (kind, ident)), None)
                if found is None:
                    found = {'id': max(self.grants, default=0) + 1, 'subject_type': kind, 'subject_id': ident}
                    self.grants[found['id']] = found
                found.update(level=level, granted_by_name=actor_name, granted_at=datetime(2026, 10, 7, 15, 0))
        elif 'baiga_access_log' in sql:
            fields = ('action', 'subject_type', 'subject_id', 'label', 'before', 'after', 'actor_id', 'actor_name')
            self.access_log.extend(dict(zip(fields, row)) for row in values)
        else:  # pragma: no cover
            raise AssertionError('неожиданная пачечная запись: %s' % sql)

    def get_grant(self, cursor, grant_id, lock=False):
        self.access_steps.append('get:locked' if lock else 'get')
        grant = self.grants.get(int(grant_id))
        if not grant:
            return None
        return dict(grant, label=self.subjects.get((grant['subject_type'], grant['subject_id'])))

    def set_grant_level(self, cursor, grant, level, actor):
        self.access_steps.append('set')
        self.access_log.append({'action': 'change', 'subject_type': grant['subject_type'],
                                'subject_id': grant['subject_id'], 'label': grant.get('label'),
                                'before': grant['level'], 'after': level,
                                'actor_id': actor['user_id'], 'actor_name': actor.get('name')})
        self.grants[grant['id']].update(level=level, granted_by_name=actor.get('name'))

    def delete_grant(self, cursor, grant, actor):
        self.access_steps.append('delete')
        self.access_log.append({'action': 'revoke', 'subject_type': grant['subject_type'],
                                'subject_id': grant['subject_id'], 'label': grant.get('label'),
                                'before': grant['level'], 'after': None,
                                'actor_id': actor['user_id'], 'actor_name': actor.get('name')})
        del self.grants[grant['id']]


class FakeDb:
    def __init__(self):
        self.commits = 0
        self.rollbacks = 0

    @contextmanager
    def _get_cursor(self):
        try:
            yield object()
        except Exception:
            self.rollbacks += 1
            raise
        self.commits += 1


@unittest.skipIf(Flask is None, 'Flask не установлен')
class _Base(unittest.TestCase):

    def setUp(self):
        self.store = Store()
        self.db = FakeDb()
        self.viewer = person(role='super_admin', department_code=None)
        self.qr_granted = False
        self.qr_error = None
        self.authorized = False

        def load_access_context(cursor, user_id):
            # Правки круга — из «базы», как у настоящей выборки контекста: ручка
            # правит таблицу, а следующий запрос любого человека это уже видит.
            return dict(self.viewer, circle_levels=dict(self.store.circle))

        patches = {'load_access_context': load_access_context}
        for name in ('list_weeks', 'filter_options', 'totals', 'search', 'found_keys', 'export_rows',
                     'zachet_summary', 'find_overlapping_uploads',
                     'lock_week', 'find_active_upload', 'insert_upload', 'insert_file', 'insert_rows',
                     'close_upload', 'mark_replaced_by', 'get_upload', 'read_file', 'list_uploads',
                     'list_exports', 'log_export',
                     # Доступ. write_grants и log_access — настоящие: подменена только
                     # пачечная запись под ними (execute_values).
                     'list_grants', 'access_catalog', 'circle_names', 'find_subjects', 'lock_access',
                     'granted_levels', 'execute_values', 'get_grant', 'set_grant_level', 'delete_grant',
                     'circle_edits', 'set_circle_level'):
            patches[name] = getattr(self.store, name)
        patcher = mock.patch.multiple(queries, **patches)
        patcher.start()
        self.addCleanup(patcher.stop)
        ready = mock.patch.object(schema, 'schema_is_ready', lambda cursor: True)
        ready.start()
        self.addCleanup(ready.stop)
        self.grants_ready = True
        grants = mock.patch.object(schema, 'grants_ready', lambda cursor: self.grants_ready)
        grants.start()
        self.addCleanup(grants.stop)
        self.circle_ready = True
        circle = mock.patch.object(schema, 'circle_ready', lambda cursor: self.circle_ready)
        circle.start()
        self.addCleanup(circle.stop)
        # Ручки проверяются со снятым выключателем; сам выключатель — SwitchTests.
        pilot = mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False)
        pilot.start()
        self.addCleanup(pilot.stop)

        # Подмены строгие, как настоящие слои: без декоратора авторизации
        # запрашивающего нет, а QR подтверждён КОНКРЕТНОМУ человеку. С тождеством
        # вместо них снятый декоратор и вопрос «подтверждён ли None» проходили.
        def require_api_key(handler):
            @wraps(handler)
            def wrapper(*args, **kwargs):
                self.authorized = True
                try:
                    return handler(*args, **kwargs)
                finally:
                    self.authorized = False
            return wrapper

        def resolve_requester():
            if not self.authorized:
                return None, None, ('Unauthorized', 401)
            return self.viewer['user_id'], None, None

        def sensitive_access_granted(user_id):
            if self.qr_error:
                raise self.qr_error
            return self.qr_granted and user_id == self.viewer['user_id']

        app = Flask(__name__)
        app.register_blueprint(routes.build_baiga_blueprint(
            db=self.db,
            require_api_key=require_api_key,
            build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=resolve_requester,
            sensitive_access_granted=sensitive_access_granted,
        ))
        app.config['TESTING'] = True
        self.client = app.test_client()

    def upload(self, content=None, name=NAME, replace=False, path='/api/baiga/uploads'):
        data = {'file': (BytesIO(content if content is not None else standard_book()), 'week.xlsx'),
                'file_name': name}
        if replace:
            data['replace'] = '1'
        return self.client.post(path, data=data, content_type='multipart/form-data')

    def writes(self):
        return [write for write in self.store.writes if write[0] != 'lock']


class GateTests(_Base):

    # Все ручки раздела — чтобы гейт проверялся на каждой, а не на «экране».
    def every_route(self):
        return [
            self.client.get('/api/baiga'),
            self.client.post('/api/baiga/rows', json={}),
            self.client.post('/api/baiga/export', json={}),
            self.upload(path='/api/baiga/uploads/preview'),
            self.upload(),
            self.client.delete('/api/baiga/uploads/1'),
            self.client.get('/api/baiga/uploads/1/file'),
            self.client.get('/api/baiga/journal'),
            self.client.get('/api/baiga/access'),
            self.client.post('/api/baiga/access/grants',
                             json={'subjects': [{'type': 'user', 'id': 31}], 'level': 'full'}),
            self.client.patch('/api/baiga/access/grants/1', json={'level': 'full'}),
            self.client.delete('/api/baiga/access/grants/1'),
            self.client.patch('/api/baiga/access/circle', json={'slot': 'staff:op', 'level': 'full'}),
        ]

    def codes(self, responses):
        return [(response.status_code, (response.get_json() or {}).get('code')) for response in responses]

    def test_rank_and_file_need_qr_before_anything(self):
        for kwargs in ({'department_code': 'szov'}, {'department_code': 'op'},
                       {'role': 'marketing_manager', 'department_code': 'marketing'}):
            self.viewer = person(**kwargs)
            self.qr_granted = False
            self.assertEqual(set(self.codes(self.every_route())), {(403, 'SENSITIVE_ACCESS_REQUIRED')}, kwargs)
            self.qr_granted = True
            data = self.client.get('/api/baiga').get_json()
            self.assertEqual(data['capabilities'], READ, kwargs)
            self.assertEqual(self.client.post('/api/baiga/rows', json={}).status_code, 200, kwargs)
        self.assertEqual(self.writes(), [])

    def test_broken_qr_check_fails_closed(self):
        """Проверка QR упала (база, сеть) — рядовой получает отказ, а не данные."""
        self.upload()
        self.viewer = person(department_code='op')
        self.qr_granted = True
        self.qr_error = RuntimeError('база недоступна')
        for response in self.every_route():
            self.assertEqual(response.status_code, 500)
            self.assertNotIn('rows', response.get_json() or {})
        self.assertEqual(self.store.exports, [])

    def test_qr_is_asked_about_the_person_who_came(self):
        """Гейт спрашивает подтверждение по id пришедшего. Подмена отвечает «да»
        только на него: вопрос про None или про id отдела оставил бы оператора
        с подтверждённым QR за замком навсегда."""
        self.viewer = person(department_code='op', user_id=31)
        self.qr_granted = True
        self.assertEqual(self.client.get('/api/baiga').status_code, 200)
        self.assertEqual(self.client.post('/api/baiga/rows', json={}).status_code, 200)

    def test_without_the_auth_layer_nobody_gets_in(self):
        """Запрашивающего даёт только слой авторизации: ручка, оставшаяся без
        него, не должна отвечать данными."""
        self.assertEqual(self.client.get('/api/baiga').status_code, 200)
        bare = Flask('bare')
        bare.register_blueprint(routes.build_baiga_blueprint(
            db=self.db, require_api_key=lambda handler: handler,
            build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=lambda: (None, None, ('Unauthorized', 401)),
            sensitive_access_granted=lambda user_id: True))
        client = bare.test_client()
        self.assertEqual(client.get('/api/baiga').status_code, 401)
        self.assertEqual(client.post('/api/baiga/rows', json={}).status_code, 401)

    def test_closed_section_says_so_before_qr(self):
        for kwargs in ({'department_code': 'tez'}, {'role': 'trainee', 'department_code': 'szov'},
                       {'role': 'trainer', 'department_code': 'op'}, {'role': 'admin', 'department_code': None},
                       {'role': 'admin', 'department_code': 'tez', 'headed_codes': ('tez',)}):
            self.viewer = person(**kwargs)
            # Подтверждённый QR раздел не открывает: он про данные, а не про круг.
            for granted in (False, True):
                self.qr_granted = granted
                self.assertEqual(set(self.codes(self.every_route())), {(403, 'BAIGA_SECTION_CLOSED')},
                                 (kwargs, granted))
        self.assertEqual(self.writes(), [])

    def test_readers_search_but_cannot_export_or_upload(self):
        self.upload()
        before = list(self.store.writes)
        self.qr_granted = True
        for kwargs in ({'role': 'sv', 'department_code': 'szov'}, {'role': 'sv', 'department_code': 'op'},
                       {'role': 'admin', 'department_code': 'op', 'headed_codes': ('op',)},
                       {'role': 'admin', 'department_code': 'szov', 'headed_codes': ('szov',)},
                       {'department_code': 'op'}, {'role': 'marketing_manager', 'department_code': 'marketing'}):
            self.viewer = person(**kwargs)
            self.assertEqual(self.codes(self.every_route()), [
                (200, None), (200, None), (403, 'BAIGA_EXPORT_FORBIDDEN'),
                (403, 'BAIGA_MANAGE_FORBIDDEN'), (403, 'BAIGA_MANAGE_FORBIDDEN'),
                (403, 'BAIGA_MANAGE_FORBIDDEN'), (403, 'BAIGA_MANAGE_FORBIDDEN'),
                (403, 'BAIGA_MANAGE_FORBIDDEN'),
                (403, 'BAIGA_ACCESS_FORBIDDEN'), (403, 'BAIGA_ACCESS_FORBIDDEN'),
                (403, 'BAIGA_ACCESS_FORBIDDEN'), (403, 'BAIGA_ACCESS_FORBIDDEN'),
                (403, 'BAIGA_ACCESS_FORBIDDEN'),
            ], kwargs)
        # Ни одной записи и ни одной строки журнала выгрузок от читающих.
        self.assertEqual(self.store.writes, before)
        self.assertEqual(self.store.exports, [])
        self.assertEqual((self.store.grants, self.store.access_log, self.store.access_locks), ({}, [], 0))
        # И круг они себе не подняли: ручка правки круга стоит за тем же гейтом.
        self.assertEqual(self.store.circle, {})
        self.assertEqual(self.store.uploads[1]['status'], 'active')

    def test_heads_and_supervisors_are_not_asked_for_qr(self):
        self.qr_granted = False
        for kwargs in ({'role': 'sv', 'department_code': 'szov'}, {'role': 'sv', 'department_code': 'op'},
                       {'role': 'admin', 'department_code': 'op', 'headed_codes': ('op',)},
                       {'role': 'admin', 'department_code': 'marketing', 'headed_codes': ('marketing',)}):
            self.viewer = person(**kwargs)
            self.assertEqual(self.client.get('/api/baiga').status_code, 200, kwargs)

    def test_marketing_head_exports_uploads_and_reads_the_journal(self):
        self.viewer = person(role='admin', department_code='marketing', headed_codes=('marketing',))
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], FULL)
        self.assertEqual(self.upload().status_code, 201)
        for row in self.store.rows:
            row['period_start'] = date(2026, 9, 21)
        self.assertEqual(self.client.post('/api/baiga/export', json={}).status_code, 200)
        self.assertEqual(self.client.get('/api/baiga/uploads/1/file').status_code, 200)
        self.assertEqual(len(self.client.get('/api/baiga/journal').get_json()['uploads']), 1)
        self.assertEqual(self.client.delete('/api/baiga/uploads/1').status_code, 200)

    def test_named_analyst_does_everything_and_the_qr_key_is_not_even_asked(self):
        """Аналитик — оператор по должности, но раздел выдан ему целиком и без
        замка. Проверка QR здесь сломана намеренно: тронь её гейт — ручки
        ответили бы 500, а не данными (ровно так падает рядовой в
        test_broken_qr_check_fails_closed)."""
        self.viewer = person(role='operator', department_code='analytik', user_id=ANALYST_ID)
        self.qr_granted = False
        self.qr_error = RuntimeError('QR у аналитика не спрашивают')
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], FULL)
        self.assertEqual(self.client.post('/api/baiga/rows', json={}).status_code, 200)
        self.assertEqual(self.upload(path='/api/baiga/uploads/preview').status_code, 200)
        self.assertEqual(self.upload().status_code, 201)
        for row in self.store.rows:
            row['period_start'] = date(2026, 9, 21)
        self.assertEqual(self.client.post('/api/baiga/export', json={}).status_code, 200)
        self.assertEqual(self.client.get('/api/baiga/uploads/1/file').status_code, 200)
        self.assertEqual(len(self.client.get('/api/baiga/journal').get_json()['uploads']), 1)
        self.assertEqual(self.upload(replace=True).status_code, 201)
        self.assertEqual(self.client.delete('/api/baiga/uploads/2').status_code, 200)
        # В журнале выгрузок — он сам, а не безымянная запись.
        self.assertEqual({entry['actor']['user_id'] for entry in self.store.exports}, {ANALYST_ID})

    def test_analysts_neighbour_in_the_department_stays_out(self):
        """Допуск именной: коллега по отделу аналитики с подтверждённым QR
        получает «раздел закрыт» на каждой ручке."""
        self.viewer = person(role='operator', department_code='analytik', user_id=ANALYST_ID + 1)
        self.qr_granted = True
        self.assertEqual(set(self.codes(self.every_route())), {(403, 'BAIGA_SECTION_CLOSED')})
        self.assertEqual(self.writes(), [])


class ScreenAndSearchTests(_Base):

    def test_screen_is_one_request_with_limits(self):
        data = self.client.get('/api/baiga').get_json()
        self.assertTrue(data['schema_ready'])
        self.assertEqual(data['capabilities'], OWNER)
        self.assertEqual(data['limits']['page_sizes'], [50, 100, 500])

    def test_search_pages_and_reports_list_misses(self):
        self.upload()
        response = self.client.post('/api/baiga/rows', json={
            'filters': {'list': '%s\nZZ999999' % fake_id(1)}, 'page': 1, 'size': 50, 'sort': 'amount', 'dir': 'desc'})
        data = response.get_json()
        self.assertEqual(data['totals']['rows'], 4)
        self.assertEqual(data['list'], {'count': 2, 'not_found': ['ZZ999999'], 'not_found_total': 1})
        self.assertEqual((data['sort'], data['dir'], data['size']), ('amount', 'desc', 50))
        self.assertNotIn('search_text', data['rows'][0])
        self.assertEqual([item['name'] for item in data['zachets']], ['Алматы-Каскелен', 'МОТО БАЙГА'])

    def test_dates_are_iso_not_gmt(self):
        self.upload()
        data = self.client.get('/api/baiga').get_json()
        self.assertEqual(data['weeks'][0]['period_start'], '2026-09-21')
        self.assertEqual(data['weeks'][0]['uploaded_at'], '2026-10-02T12:00:00')


class UploadTests(_Base):

    def test_preview_writes_nothing_and_shows_the_week(self):
        response = self.upload(path='/api/baiga/uploads/preview')
        data = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(data['ok'])
        self.assertNotIn('rows', data)
        self.assertEqual(data['period']['start'], '2026-09-21')
        self.assertIsNone(data['existing'])
        self.assertEqual(self.writes(), [])

    def test_upload_writes_week_file_and_rows(self):
        response = self.upload()
        self.assertEqual(response.status_code, 201)
        self.assertEqual(self.writes(), [('insert_upload', 1), ('insert_file', 1), ('insert_rows', 1)])
        self.assertEqual(len(self.store.rows), 4)
        self.assertEqual(self.store.uploads[1]['file_name'], NAME)

    def test_same_week_again_needs_confirmation_and_writes_nothing(self):
        self.upload()
        before = list(self.store.writes)
        response = self.upload()
        self.assertEqual((response.status_code, response.get_json()['code']), (409, 'BAIGA_WEEK_EXISTS'))
        self.assertTrue(response.get_json()['existing']['same_file'])
        self.assertEqual([write for write in self.store.writes if write not in before and write[0] != 'lock'], [])

    def test_confirmed_replacement_replaces_not_duplicates(self):
        self.upload()
        response = self.upload(replace=True)
        self.assertEqual(response.status_code, 201)
        self.assertTrue(response.get_json()['replaced'])
        self.assertEqual(self.store.uploads[1]['status'], 'replaced')
        self.assertEqual(self.store.uploads[1]['replaced_by'], 2)
        self.assertEqual(len(self.store.rows), 4)
        self.assertNotIn(1, self.store.files)
        # Прежняя неделя закрыта ДО записи новой: иначе две активные загрузки.
        closes = [i for i, write in enumerate(self.store.writes) if write[0] == 'close']
        inserts = [i for i, write in enumerate(self.store.writes) if write == ('insert_upload', 2)]
        self.assertLess(closes[0], inserts[0])

    def test_file_with_errors_is_refused_without_writes(self):
        bad = book({'Астана': [HEADER, line(1, date_text='вчера')]})
        response = self.upload(bad)
        data = response.get_json()
        self.assertEqual((response.status_code, data['code']), (422, 'BAIGA_FILE_HAS_ERRORS'))
        self.assertEqual(data['preview']['errors'][0]['column'], 'Дата')
        self.assertEqual(self.writes(), [])

    def test_overlapping_week_is_refused_in_preview_and_upload(self):
        self.store.uploads[9] = {'id': 9, 'status': 'active', 'period_start': date(2026, 9, 22),
                                 'period_end': date(2026, 9, 28), 'rows_count': 5, 'file_name': 'x',
                                 'uploaded_by_name': 'x', 'uploaded_at': datetime(2026, 10, 1, 9, 0)}
        preview = self.upload(path='/api/baiga/uploads/preview').get_json()
        self.assertFalse(preview['ok'])
        self.assertIn('пересекается', preview['errors'][0]['message'])
        response = self.upload()
        self.assertEqual((response.status_code, response.get_json()['code']), (409, 'BAIGA_WEEK_OVERLAP'))
        self.assertEqual(self.writes(), [])

    def test_short_write_rolls_back(self):
        self.store.short_write = True
        response = self.upload()
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.db.rollbacks, 1)

    def test_not_excel_and_missing_file(self):
        response = self.upload(b'hello')
        self.assertEqual((response.status_code, response.get_json()['code']), (422, 'BAIGA_FILE_NOT_XLSX'))
        response = self.client.post('/api/baiga/uploads', data={}, content_type='multipart/form-data')
        self.assertEqual(response.get_json()['code'], 'BAIGA_FILE_REQUIRED')

    def test_cyrillic_file_name_is_kept_as_is(self):
        self.upload(name='C:\\Users\\аналитик\\' + NAME)
        self.assertEqual(self.store.uploads[1]['file_name'], NAME)


class DeleteSourceJournalTests(_Base):

    def setUp(self):
        super().setUp()
        self.upload()

    def test_delete_week(self):
        response = self.client.delete('/api/baiga/uploads/1')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.store.uploads[1]['status'], 'deleted')
        self.assertEqual(self.store.rows, [])
        again = self.client.delete('/api/baiga/uploads/1')
        self.assertEqual(again.get_json()['code'], 'BAIGA_UPLOAD_CLOSED')

    def test_source_download_is_journaled(self):
        response = self.client.get('/api/baiga/uploads/1/file')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, self.store.files[1])
        self.assertIn("filename*=UTF-8''", response.headers['Content-Disposition'])
        self.assertIn('Content-Disposition', response.headers['Access-Control-Expose-Headers'])
        self.assertEqual(self.store.exports[-1]['kind'], 'source')
        self.client.delete('/api/baiga/uploads/1')
        self.assertEqual(self.client.get('/api/baiga/uploads/1/file').get_json()['code'], 'BAIGA_FILE_GONE')

    def test_export_is_journaled_with_filters(self):
        response = self.client.post('/api/baiga/export', json={'filters': {'city': 'Шымкент'}, 'mode': 'single'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['X-Rows'], '4')
        self.assertIn('Content-Disposition', response.headers['Access-Control-Expose-Headers'])
        logged = self.store.exports[-1]
        self.assertEqual((logged['kind'], logged['mode'], logged['filters'], logged['rows_count']),
                         ('export', 'single', {'city': 'Шымкент'}, 4))

    def test_export_by_sheets_goes_in_file_order(self):
        self.client.post('/api/baiga/export', json={'mode': 'sheets', 'sort': 'amount'})
        self.assertEqual(self.store.last_export[2], 'r.period_start, r.sheet_order, r.row_number, r.id')

    def test_export_limit_and_empty_selection(self):
        with mock.patch.object(report, 'EXPORT_LIMIT', 3):
            self.store.last_export = None
            response = self.client.post('/api/baiga/export', json={})
            self.assertEqual(response.get_json()['code'], 'BAIGA_EXPORT_TOO_BIG')
            # Отказ — по счёту, без выборки строк в память.
            self.assertIsNone(self.store.last_export)
        self.store.rows = []
        self.assertEqual(self.client.post('/api/baiga/export', json={}).get_json()['code'], 'BAIGA_EXPORT_EMPTY')
        self.assertEqual([entry for entry in self.store.exports if entry['kind'] == 'export'], [])

    def test_journal_lists_uploads(self):
        data = self.client.get('/api/baiga/journal').get_json()
        self.assertEqual(len(data['uploads']), 1)


# ─────────────────────────────────────────────────────────────────────────────
# Проводка
# ─────────────────────────────────────────────────────────────────────────────

class AccessSheetTests(_Base):
    """Лист «Доступ»: ручки /api/baiga/access*. Смотрит супер-админ, пока тест
    не назначит другого."""

    ACCESS_ROUTES = 5

    def grant(self, subjects, level='read'):
        return self.client.post('/api/baiga/access/grants', json={
            'subjects': [{'type': kind, 'id': ident} for kind, ident in subjects], 'level': level})

    def access_routes(self):
        return [
            self.client.get('/api/baiga/access'),
            self.grant([('user', 31)], 'full'),
            self.client.patch('/api/baiga/access/grants/1', json={'level': 'full'}),
            self.client.delete('/api/baiga/access/grants/1'),
            self.client.patch('/api/baiga/access/circle', json={'slot': 'staff:op', 'level': 'full'}),
        ]

    def state(self):
        return (json.dumps(self.store.grants, default=str, sort_keys=True), len(self.store.access_log),
                dict(self.store.circle))

    def test_sheet_is_one_request(self):
        self.assertEqual(self.grant([('group', 5)]).status_code, 200)
        data = self.client.get('/api/baiga/access').get_json()
        self.assertEqual(set(data), {'grants', 'circle', 'catalog', 'max_subjects'})
        self.assertEqual([(grant['subject_type'], grant['label'], grant['level']) for grant in data['grants']],
                         [('group', 'Регионы', 'read')])
        self.assertEqual(data['grants'][0]['granted_at'], '2026-10-07T15:00:00')
        self.assertEqual(data['max_subjects'], access.MAX_GRANT_SUBJECTS)
        self.assertEqual(set(data['catalog']), {'department', 'group', 'user'})
        self.assertEqual([item['name'] for item in data['catalog']['user']], ['Оператов О. О.', 'Тренеров Т. Т.'])
        # Круг — строками «должность в отделе», в порядке правила, с названиями
        # отделов и именами названных; id людей наружу не уходят.
        circle = {row['slot']: row for row in data['circle']}
        self.assertEqual([row['slot'] for row in data['circle']], [row['slot'] for row in access.circle()])
        self.assertEqual(circle['super_admin'], {'slot': 'super_admin', 'kind': 'super_admin', 'level': 'full',
                                                 'qr': False, 'locked': True, 'code': None, 'department': None,
                                                 'person': None, 'updated_by_name': None, 'updated_at': None})
        self.assertEqual(circle['head:marketing'], {'slot': 'head:marketing', 'kind': 'head', 'level': 'full',
                                                    'qr': False, 'locked': False, 'code': 'marketing',
                                                    'department': 'Отдел marketing', 'person': None,
                                                    'updated_by_name': None, 'updated_at': None})
        self.assertEqual(circle['staff:szov'], {'slot': 'staff:szov', 'kind': 'staff', 'level': 'read',
                                                'qr': True, 'locked': False, 'code': 'szov',
                                                'department': 'Отдел szov', 'person': None,
                                                'updated_by_name': None, 'updated_at': None})
        named = circle['named:%d' % ANALYST_ID]
        self.assertEqual((named['kind'], named['level'], named['person'], named['department'], named['locked']),
                         ('named', 'full', 'Названный %d' % ANALYST_ID, None, False))
        self.assertNotIn('user_id', json.dumps(data['circle']))
        # Названного в штате уже нет — строки нет: подписать её нечем.
        self.store.gone_people.add(ANALYST_ID)
        slots = [row['slot'] for row in self.client.get('/api/baiga/access').get_json()['circle']]
        self.assertEqual(slots, [row['slot'] for row in access.circle() if row['kind'] != 'named'])

    def test_grant_to_many_at_once(self):
        response = self.grant([('user', 31), ('group', 5), ('department', 44)], 'export')
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual((data['granted'], data['changed']), (3, 0))
        # Ответ несёт свежий список — второго запроса «перечитать» листу не нужно.
        self.assertEqual({(grant['subject_type'], grant['subject_id'], grant['level']) for grant in data['grants']},
                         {('user', 31, 'export'), ('group', 5, 'export'), ('department', 44, 'export')})
        self.assertEqual([(entry['action'], entry['subject_type'], entry['label'], entry['before'], entry['after'])
                          for entry in self.store.access_log],
                         [('grant', 'user', 'Оператов О. О.', None, 'export'),
                          ('grant', 'group', 'Регионы', None, 'export'),
                          ('grant', 'department', 'Бухгалтерия', None, 'export')])
        # В журнале — тот, кто выдал, а не безымянная запись.
        self.assertEqual({(entry['actor_id'], entry['actor_name']) for entry in self.store.access_log},
                         {(self.viewer['user_id'], self.viewer['name'])})
        self.assertEqual({grant['granted_by_name'] for grant in self.store.grants.values()}, {self.viewer['name']})
        self.assertEqual(self.store.access_locks, 1)
        # Блокировка — до чтения адресатов: два раздающих не читают одно состояние.
        self.assertEqual(self.store.access_steps, ['lock', 'find', 'write'])

    def test_level_change_and_revoke_lock_before_they_read(self):
        self.grant([('group', 5)], 'read')
        del self.store.access_steps[:]
        self.client.patch('/api/baiga/access/grants/1', json={'level': 'full'})
        self.assertEqual(self.store.access_steps, ['lock', 'get:locked', 'set'])
        del self.store.access_steps[:]
        self.client.delete('/api/baiga/access/grants/1')
        self.assertEqual(self.store.access_steps, ['lock', 'get:locked', 'delete'])
        # Отказ «выдачи уже нет» — тоже под блокировкой и без записи.
        del self.store.access_steps[:]
        self.client.delete('/api/baiga/access/grants/1')
        self.assertEqual(self.store.access_steps, ['lock', 'get:locked'])

    def test_one_missing_subject_refuses_the_whole_grant(self):
        """Всё или ничего: половина выданного списка читалась бы как «не
        удалось» при наполовину открытом разделе."""
        self.grant([('user', 31)])
        before = self.state()
        for subjects in ([('user', 32), ('group', 999)], [('department', 31)], [('user', 5)]):
            response = self.grant(subjects, 'full')
            self.assertEqual((response.status_code, response.get_json()['code']),
                             (422, 'BAIGA_ACCESS_SUBJECT_GONE'), subjects)
            self.assertEqual(self.state(), before, subjects)

    def test_repeat_grant_changes_the_level_and_the_same_level_is_left_alone(self):
        self.grant([('user', 31)], 'read')
        data = self.grant([('user', 31), ('user', 32)], 'full').get_json()
        self.assertEqual((data['granted'], data['changed']), (1, 1))
        self.assertEqual(len(self.store.grants), 2)
        self.assertEqual([(entry['action'], entry['subject_id'], entry['before'], entry['after'])
                          for entry in self.store.access_log[1:]],
                         [('grant', 32, None, 'full'), ('change', 31, 'read', 'full')])
        # Тот же уровень ещё раз: ни записи, ни строки журнала, «выдал» не переписан.
        self.viewer = person(role='super_admin', department_code=None, user_id=11)
        before = self.state()
        data = self.grant([('user', 31), ('user', 32)], 'full').get_json()
        self.assertEqual((data['granted'], data['changed']), (0, 0))
        self.assertEqual(self.state(), before)
        self.assertEqual(len(data['grants']), 2)

    def test_same_subject_twice_in_a_request_counts_once(self):
        data = self.grant([('user', 31), ('user', 31)], 'read').get_json()
        self.assertEqual((data['granted'], len(self.store.grants), len(self.store.access_log)), (1, 1, 1))

    def test_bad_requests_are_refused_before_any_write(self):
        self.grant([('user', 31)])
        before, locks = self.state(), self.store.access_locks
        many = [{'type': 'user', 'id': ident} for ident in range(1, access.MAX_GRANT_SUBJECTS + 2)]
        cases = [
            ({'subjects': [{'type': 'user', 'id': 32}]}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'subjects': [{'type': 'user', 'id': 32}], 'level': 'owner'}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'level': 'read'}, 'BAIGA_ACCESS_SUBJECT_REQUIRED'),
            ({'subjects': [], 'level': 'read'}, 'BAIGA_ACCESS_SUBJECT_REQUIRED'),
            ({'subjects': {'type': 'user', 'id': 32}, 'level': 'read'}, 'BAIGA_ACCESS_SUBJECT_REQUIRED'),
            ({'subjects': [{'type': 'direction', 'id': 32}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'otp_role', 'id': 32}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'user', 'id': '32'}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'user', 'id': 0}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'user', 'id': -3}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'user', 'id': True}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'user', 'id': 3.5}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [{'type': 'user'}], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': ['user:32'], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': [None], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SUBJECT'),
            # Хороший адресат рядом с плохим не спасает запрос и не пишется сам.
            ({'subjects': [{'type': 'user', 'id': 32}, {'type': 'user', 'id': 'x'}], 'level': 'read'},
             'BAIGA_ACCESS_BAD_SUBJECT'),
            ({'subjects': many, 'level': 'read'}, 'BAIGA_ACCESS_TOO_MANY'),
        ]
        for body, code in cases:
            response = self.client.post('/api/baiga/access/grants', json=body)
            self.assertEqual((response.status_code, response.get_json()['code']), (400, code), body)
            self.assertEqual((self.state(), self.store.access_locks), (before, locks), body)
        # Ровно предел — проходит разбор (и упирается уже в «нет таких людей»).
        response = self.client.post('/api/baiga/access/grants', json={'subjects': many[:-1], 'level': 'read'})
        self.assertEqual(response.get_json()['code'], 'BAIGA_ACCESS_SUBJECT_GONE')

    def test_level_is_changed_and_grant_is_revoked_by_id(self):
        self.grant([('group', 5)], 'read')
        response = self.client.patch('/api/baiga/access/grants/1', json={'level': 'full'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([grant['level'] for grant in response.get_json()['grants']], ['full'])
        self.assertEqual([(entry['action'], entry['label'], entry['before'], entry['after'])
                          for entry in self.store.access_log[1:]], [('change', 'Регионы', 'read', 'full')])
        # Тот же уровень — без записи; незнакомый — отказ; чужого id — нет.
        before = self.state()
        self.assertEqual(self.client.patch('/api/baiga/access/grants/1', json={'level': 'full'}).status_code, 200)
        for body in ({'level': 'owner'}, {}, {'level': None}):
            response = self.client.patch('/api/baiga/access/grants/1', json=body)
            self.assertEqual((response.status_code, response.get_json()['code']), (400, 'BAIGA_ACCESS_BAD_LEVEL'))
        response = self.client.patch('/api/baiga/access/grants/9', json={'level': 'read'})
        self.assertEqual((response.status_code, response.get_json()['code']), (404, 'BAIGA_ACCESS_GRANT_NOT_FOUND'))
        self.assertEqual(self.state(), before)

        response = self.client.delete('/api/baiga/access/grants/1')
        self.assertEqual((response.status_code, response.get_json()['grants']), (200, []))
        self.assertEqual(self.store.grants, {})
        self.assertEqual([(entry['action'], entry['label'], entry['before'], entry['after'])
                          for entry in self.store.access_log[2:]], [('revoke', 'Регионы', 'full', None)])
        again = self.client.delete('/api/baiga/access/grants/1')
        self.assertEqual((again.status_code, again.get_json()['code']), (404, 'BAIGA_ACCESS_GRANT_NOT_FOUND'))
        self.assertEqual(len(self.store.access_log), 3)

    def test_full_access_to_the_section_does_not_hand_out_access(self):
        """Глава «Маркетинга» ведёт раздел целиком, но раздаёт его только
        названный поимённо: тот же профиль с чужим id получает отказ."""
        self.grant([('group', 5)])
        before = self.state()
        for viewer in (person(role='admin', department_code='marketing', headed_codes=('marketing',)),
                       person(department_code='analytik', user_id=ANALYST_ID),
                       person(role='sv', department_code='tez', grants=('full',)),
                       person(role='admin', department_code='szov', headed_codes=('szov',))):
            self.viewer = viewer
            self.assertTrue(access.can_open_section(viewer))
            self.assertFalse(self.client.get('/api/baiga').get_json()['capabilities']['can_manage_access'])
            codes = [(response.status_code, response.get_json()['code']) for response in self.access_routes()]
            self.assertEqual(codes, [(403, 'BAIGA_ACCESS_FORBIDDEN')] * self.ACCESS_ROUTES, viewer)
            self.assertEqual(self.state(), before, viewer)

    def test_named_manager_hands_out_access(self):
        self.viewer = person(role='admin', department_code='marketing', headed_codes=('marketing',), user_id=415)
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], OWNER)
        self.assertEqual(self.grant([('department', 44)], 'read').status_code, 200)
        self.assertEqual(self.store.access_log[0]['actor_id'], 415)
        self.assertEqual(self.client.get('/api/baiga/access').status_code, 200)

    def test_named_manager_outside_the_circle_opens_the_section_and_stays_behind_qr(self):
        with mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({77})):
            self.viewer = person(role='operator', department_code='tez', user_id=77)
            self.qr_granted = False
            codes = {(response.status_code, response.get_json()['code']) for response in self.access_routes()}
            self.assertEqual(codes, {(403, 'SENSITIVE_ACCESS_REQUIRED')})
            self.assertEqual(self.store.grants, {})
            self.qr_granted = True
            self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'],
                             dict(READ, can_manage_access=True))
            self.assertEqual(self.grant([('user', 31)], 'export').status_code, 200)
            # Свой уровень раздающий поднимает той же выдачей — себе, как любому.
            self.store.subjects[('user', 77)] = self.viewer['name']
            self.assertEqual(self.grant([('user', 77)], 'full').status_code, 200)

    def test_switch_closes_the_sheet_for_the_named(self):
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True):
            self.viewer = person(role='admin', department_code='marketing', headed_codes=('marketing',), user_id=415)
            codes = {(response.status_code, response.get_json()['code']) for response in self.access_routes()}
            self.assertEqual(codes, {(403, 'BAIGA_SECTION_CLOSED')})
            self.viewer = person(role='super_admin', department_code=None)
            self.assertEqual(self.client.get('/api/baiga/access').status_code, 200)

    def test_not_deployed_tables_answer_plainly_and_the_section_lives(self):
        """Миграция выдач не легла: лист честно говорит об этом, а раздел
        работает по кругу — как до кнопки «Доступ»."""
        self.grants_ready = False
        codes = [(response.status_code, response.get_json()['code']) for response in self.access_routes()]
        self.assertEqual(codes, [(409, 'BAIGA_ACCESS_NOT_READY')] * self.ACCESS_ROUTES)
        self.assertEqual((self.store.grants, self.store.access_log, self.store.access_locks), ({}, [], 0))
        self.assertEqual(self.store.circle, {})
        self.assertEqual(self.client.get('/api/baiga').status_code, 200)
        self.assertEqual(self.client.post('/api/baiga/rows', json={}).status_code, 200)

    def test_granted_person_gets_exactly_the_granted_level(self):
        self.upload()
        for row in self.store.rows:
            row['period_start'] = date(2026, 9, 21)
        expected = {
            'read': (READ, 403, 403),
            'export': (EXPORT, 200, 403),
            'full': (FULL, 200, 200),
        }
        for level, (caps, export_status, journal_status) in expected.items():
            self.viewer = person(role='sv', department_code='tez', grants=(level,))
            self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], caps, level)
            self.assertEqual(self.client.post('/api/baiga/rows', json={}).status_code, 200, level)
            self.assertEqual(self.client.post('/api/baiga/export', json={}).status_code, export_status, level)
            self.assertEqual(self.client.get('/api/baiga/journal').status_code, journal_status, level)
            self.assertEqual(self.client.get('/api/baiga/access').status_code, 403, level)

    def test_granted_rank_and_file_wait_for_qr_on_every_route(self):
        for kwargs in ({'role': 'trainee', 'department_code': 'front_office'},
                       {'role': 'operator', 'department_code': 'tez'},
                       {'role': 'accounting_manager', 'department_code': 'accounting'}):
            self.viewer = person(grants=('full',), **kwargs)
            self.qr_granted = False
            codes = {(response.status_code, response.get_json()['code'])
                     for response in (self.client.get('/api/baiga'), self.client.post('/api/baiga/rows', json={}),
                                      self.client.post('/api/baiga/export', json={}),
                                      self.client.get('/api/baiga/journal'), self.client.get('/api/baiga/access'))}
            self.assertEqual(codes, {(403, 'SENSITIVE_ACCESS_REQUIRED')}, kwargs)
            self.qr_granted = True
            self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], FULL, kwargs)
        # Кадровику кода не выдают — выданный раздел открыт ему без замка.
        self.viewer = person(role='hr_manager', department_code='hr', grants=('read',))
        self.qr_granted = False
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], READ)


class CircleRouteTests(_Base):
    """PATCH /api/baiga/access/circle — уровень строки «открыт по умолчанию».
    Правит супер-админ, пока тест не назначит другого; «база» правок —
    store.circle, и контекст каждого следующего запроса читает её же."""

    OP = {'department_code': 'op'}
    SV_OP = {'role': 'sv', 'department_code': 'op'}

    def edit(self, slot, level):
        return self.client.patch('/api/baiga/access/circle', json={'slot': slot, 'level': level})

    def as_person(self, **kwargs):
        """Что раздел отвечает человеку: (код экрана, права или код отказа)."""
        admin, self.viewer = self.viewer, person(**kwargs)
        try:
            response = self.client.get('/api/baiga')
            data = response.get_json()
            return response.status_code, data.get('capabilities') or data.get('code')
        finally:
            self.viewer = admin

    def state(self):
        return dict(self.store.circle), len(self.store.access_log), self.store.access_locks

    def test_level_of_a_row_is_changed_and_the_people_of_the_row_get_it(self):
        self.qr_granted = True
        self.assertEqual(self.as_person(**self.OP), (200, READ))
        response = self.edit('staff:op', 'export')
        self.assertEqual(response.status_code, 200)
        # Блокировка — до чтения таблицы: два раздающих не читают одно состояние.
        # После записи таблица перечитана: время правки ставит база.
        self.assertEqual(self.store.access_steps, ['lock', 'circle:read', 'circle:set', 'circle:read'])
        data = response.get_json()
        # Ответ — свежие строки круга, те же, что отдаст лист: второго запроса нет.
        self.assertEqual(set(data), {'circle'})
        self.assertEqual(data['circle'], self.client.get('/api/baiga/access').get_json()['circle'])
        levels = {row['slot']: row['level'] for row in data['circle']}
        self.assertEqual(levels['staff:op'], 'export')
        self.assertEqual({slot: level for slot, level in levels.items() if slot != 'staff:op'},
                         {row['slot']: row['level'] for row in access.circle() if row['slot'] != 'staff:op'})
        self.assertEqual(self.store.circle, {'staff:op': 'export'})
        # У правленой строки — кто и когда её правил; у нетронутых этого нет.
        edited = {row['slot']: (row['updated_by_name'], row['updated_at']) for row in data['circle']}
        self.assertEqual(edited.pop('staff:op'), (self.viewer['name'], '2026-10-08T12:00:00'))
        self.assertEqual(set(edited.values()), {(None, None)})
        # Оператор ОП выгружает; его супервайзер и оператор СЗоВ — как были.
        self.assertEqual(self.as_person(**self.OP), (200, EXPORT))
        self.assertEqual(self.as_person(**self.SV_OP), (200, READ))
        self.assertEqual(self.as_person(department_code='szov'), (200, READ))
        # След: что было (значение по умолчанию), что стало и кто правил.
        self.assertEqual(self.store.access_log, [{
            'action': 'circle', 'subject_type': 'circle', 'subject_id': 0, 'label': 'staff:op',
            'before': 'read', 'after': 'export',
            'actor_id': self.viewer['user_id'], 'actor_name': self.viewer['name']}])

    def test_closed_row_closes_the_section_for_its_people_only(self):
        self.qr_granted = True
        self.assertEqual(self.edit('staff:op', 'none').status_code, 200)
        self.assertEqual(self.as_person(**self.OP), (403, 'BAIGA_SECTION_CLOSED'))
        self.assertEqual(self.as_person(**self.SV_OP), (200, READ))
        self.assertEqual(self.as_person(department_code='szov'), (200, READ))
        self.assertEqual(self.as_person(role='admin', department_code='op', headed_codes=('op',)), (200, READ))
        # Закрытая строка в листе — «нет», и следом её открывают обратно.
        shown = {row['slot']: row['level'] for row in self.client.get('/api/baiga/access').get_json()['circle']}
        self.assertEqual((shown['staff:op'], shown['sv:op']), ('none', 'read'))
        self.assertEqual(self.edit('staff:op', 'read').status_code, 200)
        self.assertEqual(self.as_person(**self.OP), (200, READ))
        self.assertEqual([(entry['before'], entry['after']) for entry in self.store.access_log],
                         [('read', 'none'), ('none', 'read')])

    def test_every_row_and_every_level_goes_through(self):
        """Каждая строка × каждый уровень: ответ, таблица и строка листа."""
        defaults = access.circle_defaults()
        for slot in defaults:
            for level in access.CIRCLE_LEVELS:
                response = self.edit(slot, level)
                self.assertEqual(response.status_code, 200, (slot, level))
                shown = {row['slot']: row['level'] for row in response.get_json()['circle']}
                self.assertEqual(shown[slot], level, (slot, level))
            # Вернули как было — и раздел отвечает как до правок.
            self.assertEqual(self.edit(slot, defaults[slot]).status_code, 200)
        self.assertEqual(self.store.circle, defaults)
        self.assertEqual([row['level'] for row in self.client.get('/api/baiga/access').get_json()['circle']],
                         [row['level'] for row in access.circle()])

    def test_same_level_writes_nothing(self):
        # Уровень по умолчанию, пока строки в таблице нет, — тоже «тот же».
        before = self.state()
        self.assertEqual(self.edit('staff:op', 'read').status_code, 200)
        self.assertEqual(self.edit('head:marketing', 'full').status_code, 200)
        self.assertEqual((self.store.circle, self.store.access_log), ({}, []))
        self.assertEqual(self.store.access_steps, ['lock', 'circle:read'] * 2)
        self.assertEqual(self.state()[:2], before[:2])
        self.edit('staff:op', 'full')
        del self.store.access_steps[:]
        before = self.state()
        self.viewer = person(role='super_admin', department_code=None, user_id=11)
        self.assertEqual(self.edit('staff:op', 'full').status_code, 200)
        self.assertEqual(self.state()[:2], before[:2])
        self.assertEqual(self.store.access_steps, ['lock', 'circle:read'])

    def test_bad_requests_are_refused_before_any_write(self):
        self.edit('sv:szov', 'export')
        before = self.state()
        cases = [
            # Строки супер-админа среди правимых нет: запереть раздел от всех нельзя.
            ({'slot': 'super_admin', 'level': 'none'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'super_admin', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            # Правится уровень строки, а не состав круга: чужого отдела, должности
            # без строки и чужого имени в нём нет.
            ({'slot': 'staff:tez', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'head:tez', 'level': 'full'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'sv:marketing', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'trainer:op', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'named:%d' % (ANALYST_ID + 1), 'level': 'full'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'named:415', 'level': 'full'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'STAFF:OP', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'staff', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': '', 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': None, 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': ['staff:op'], 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 7, 'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'level': 'read'}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({}, 'BAIGA_ACCESS_BAD_SLOT'),
            ({'slot': 'staff:op'}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': ''}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': None}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': 'owner'}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': 'нет'}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': ['read']}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': 0}, 'BAIGA_ACCESS_BAD_LEVEL'),
            ({'slot': 'staff:op', 'level': False}, 'BAIGA_ACCESS_BAD_LEVEL'),
        ]
        for body, code in cases:
            response = self.client.patch('/api/baiga/access/circle', json=body)
            self.assertEqual((response.status_code, response.get_json()['code']), (400, code), body)
            self.assertEqual(self.state(), before, body)
        # Не JSON вовсе — тот же отказ, а не 500.
        response = self.client.patch('/api/baiga/access/circle', data='slot=staff:op', content_type='text/plain')
        self.assertEqual((response.status_code, response.get_json()['code']), (400, 'BAIGA_ACCESS_BAD_SLOT'))
        self.assertEqual(self.state(), before)

    def test_junk_in_the_table_is_shown_closed_and_is_overwritten(self):
        self.qr_granted = True
        self.store.circle['staff:op'] = 'owner'
        shown = {row['slot']: row['level'] for row in self.client.get('/api/baiga/access').get_json()['circle']}
        self.assertEqual(shown['staff:op'], 'none')
        self.assertEqual(self.as_person(**self.OP), (403, 'BAIGA_SECTION_CLOSED'))
        # «Нет» поверх мусора — не «тот же уровень»: в таблицу ложится настоящее значение.
        self.assertEqual(self.edit('staff:op', 'none').status_code, 200)
        self.assertEqual(self.store.circle, {'staff:op': 'none'})
        self.assertEqual([(entry['before'], entry['after']) for entry in self.store.access_log], [('owner', 'none')])

    def test_without_the_table_the_edit_says_so_and_the_sheet_shows_the_defaults(self):
        """Таблица правок моложе выдач: не легла она — лист открывается с кругом
        по умолчанию, выдачи работают, а правка честно отвечает «ещё нет»."""
        self.circle_ready = False
        self.store.circle_missing = True
        before = self.state()
        response = self.edit('staff:op', 'none')
        self.assertEqual((response.status_code, response.get_json()['code']), (409, 'BAIGA_ACCESS_NOT_READY'))
        self.assertEqual(self.state(), before)
        sheet = self.client.get('/api/baiga/access')
        self.assertEqual(sheet.status_code, 200)
        self.assertEqual([row['level'] for row in sheet.get_json()['circle']],
                         [row['level'] for row in access.circle()])
        granted = self.client.post('/api/baiga/access/grants',
                                   json={'subjects': [{'type': 'user', 'id': 31}], 'level': 'read'})
        self.assertEqual(granted.status_code, 200)
        # И без таблиц выдач — тот же ответ, а не правка круга вслепую.
        self.store.circle_missing = False
        self.circle_ready, self.grants_ready = True, False
        response = self.edit('staff:op', 'none')
        self.assertEqual((response.status_code, response.get_json()['code']), (409, 'BAIGA_ACCESS_NOT_READY'))
        self.assertEqual(self.store.circle, {})

    def test_only_those_who_hand_out_access_edit_the_circle(self):
        """Полный доступ к разделу круг не правит: глава «Маркетинга» с чужим id,
        аналитик, человек с выданным «полным» — отказ и ни одной записи."""
        for viewer in (person(role='admin', department_code='marketing', headed_codes=('marketing',)),
                       person(department_code='analytik', user_id=ANALYST_ID),
                       person(role='sv', department_code='tez', grants=('full',)),
                       person(role='admin', department_code='szov', headed_codes=('szov',))):
            self.viewer = viewer
            for slot, level in (('staff:op', 'full'), ('head:szov', 'full'), ('head:marketing', 'none')):
                response = self.edit(slot, level)
                self.assertEqual((response.status_code, response.get_json()['code']),
                                 (403, 'BAIGA_ACCESS_FORBIDDEN'), (viewer, slot))
        self.assertEqual(self.state(), ({}, 0, 0))

    def test_named_manager_edits_the_circle_and_cannot_lock_himself_out(self):
        manager = person(role='admin', department_code='marketing', headed_codes=('marketing',), user_id=415)
        self.viewer = manager
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], OWNER)
        self.assertEqual(self.edit('sv:szov', 'export').status_code, 200)
        self.assertEqual(self.store.access_log[0]['actor_id'], 415)
        # Свою строку он закрыл: вести раздел перестал, но читает и раздаёт —
        # и там же возвращает уровень обратно.
        self.assertEqual(self.edit('head:marketing', 'none').status_code, 200)
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], dict(READ, can_manage_access=True))
        self.assertEqual(self.client.get('/api/baiga/journal').get_json()['code'], 'BAIGA_MANAGE_FORBIDDEN')
        self.assertEqual(self.client.get('/api/baiga/access').status_code, 200)
        self.assertEqual(self.edit('head:marketing', 'full').status_code, 200)
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], OWNER)

    def test_super_admin_stays_in_with_everything_closed(self):
        for slot in access.circle_defaults():
            self.assertEqual(self.edit(slot, 'none').status_code, 200, slot)
        self.assertEqual(self.client.get('/api/baiga').get_json()['capabilities'], OWNER)
        shown = {row['slot']: row['level'] for row in self.client.get('/api/baiga/access').get_json()['circle']}
        self.assertEqual(shown.pop('super_admin'), 'full')
        self.assertEqual(set(shown.values()), {'none'})
        self.qr_granted = True
        for kwargs in (self.OP, self.SV_OP, {'role': 'marketing_manager', 'department_code': 'marketing'},
                       {'role': 'admin', 'department_code': 'marketing', 'headed_codes': ('marketing',)},
                       {'role': 'admin', 'department_code': 'szov', 'headed_codes': ('szov',)},
                       {'department_code': 'analytik', 'user_id': ANALYST_ID}):
            self.assertEqual(self.as_person(**kwargs), (403, 'BAIGA_SECTION_CLOSED'), kwargs)

    def test_named_row_is_edited_like_the_rest_and_keeps_no_lock(self):
        analyst = {'role': 'operator', 'department_code': 'analytik', 'user_id': ANALYST_ID}
        slot = 'named:%d' % ANALYST_ID
        self.qr_granted = False
        self.assertEqual(self.as_person(**analyst), (200, FULL))
        self.assertEqual(self.edit(slot, 'read').status_code, 200)
        self.assertEqual(self.as_person(**analyst), (200, READ))
        self.assertEqual(self.edit(slot, 'none').status_code, 200)
        self.assertEqual(self.as_person(**analyst), (403, 'BAIGA_SECTION_CLOSED'))
        # Сосед по отделу от правок именной строки ничего не получил.
        self.assertEqual(self.as_person(role='operator', department_code='analytik', user_id=ANALYST_ID + 1),
                         (403, 'BAIGA_SECTION_CLOSED'))

    def test_closed_row_keeps_what_was_granted_separately(self):
        """Строку закрыли — выдача человеку остаётся: круг и выдачи складываются,
        а не отменяют друг друга."""
        self.qr_granted = True
        self.edit('staff:op', 'none')
        self.assertEqual(self.as_person(grants=('export',), **self.OP), (200, EXPORT))
        self.assertEqual(self.as_person(**self.OP), (403, 'BAIGA_SECTION_CLOSED'))

    def test_rank_and_file_of_an_opened_row_still_wait_for_qr(self):
        """Строку подняли до «полного» — замок остаётся: оператор ведёт раздел
        только после QR-подтверждения."""
        self.edit('staff:op', 'full')
        self.qr_granted = False
        self.assertEqual(self.as_person(**self.OP), (403, 'SENSITIVE_ACCESS_REQUIRED'))
        self.qr_granted = True
        self.assertEqual(self.as_person(**self.OP), (200, FULL))
        # И кнопки «Доступ» полный уровень не даёт.
        self.viewer = person(**self.OP)
        response = self.edit('staff:op', 'none')
        self.assertEqual((response.status_code, response.get_json()['code']), (403, 'BAIGA_ACCESS_FORBIDDEN'))
        self.assertEqual(self.store.circle, {'staff:op': 'full'})


class WiringTests(unittest.TestCase):

    def setUp(self):
        self.app = _read(APP_JSX)

    def test_schema_is_deployed(self):
        source = _read(DATABASE_PY)
        thermo = source.index('            self._init_thermoboxes_schema_tx(cursor)\n')
        baiga = source.index('            self._init_baiga_schema_tx(cursor)\n')
        self.assertLess(thermo, baiga)
        self.assertIn('def _init_baiga_schema_tx(self, cursor):', source)
        self.assertIn('SAVEPOINT baiga_schema', source)

    def test_blueprint_is_registered_with_the_qr_key(self):
        source = _read(BOT_PY)
        block = source.split('from baiga.routes import build_baiga_blueprint')[1].split('except Exception')[0]
        self.assertIn('app.register_blueprint(build_baiga_blueprint(', block)
        # Все три: без авторизации раздел не знает, кто пришёл, а с чужим
        # resolve_requester отвечал бы не тому человеку.
        self.assertIn('        require_api_key=require_api_key,\n', block)
        self.assertIn('        resolve_requester=_resolve_requester,\n', block)
        self.assertIn('sensitive_access_granted=_sensitive_access_granted_for_user', block)
        self.assertIn('            @require_api_key\n', _read(ROOT / 'baiga' / 'routes.py'))

    def test_lazy_view_behind_the_qr_gate(self):
        self.assertIn("const BaigaView = lazyWithRetry(() => import('./components/baiga/BaigaView'));", self.app)
        # Условие целиком и порядок веток: с подстрокой проходила и инверсия
        # замка («!sensitiveSectionsLocked ?»), при которой раздел не открылся бы
        # ни у кого из круга.
        start = '{view === "baiga" && canAccessBaigaSection && (baigaLocked ? ('
        self.assertEqual(self.app.count(start), 1)
        self.assertEqual(self.app.count('{view === "baiga" && '), 1)
        block = self.app.split(start)[1].split('</Suspense>')[0]
        gate, otherwise, section = (block.index('<SensitiveSectionGate'), block.index(') : ('),
                                    block.index('<BaigaView'))
        self.assertLess(gate, otherwise)
        self.assertLess(otherwise, section)
        # Замок у раздела свой: общий спросил бы код и у аналитика из списка.
        self.assertIn('const baigaLocked = baigaQrRequiredFor(user) && !sensitiveAccess.granted;', self.app)
        self.assertIn('const baigaChecking = baigaLocked && !sensitiveAccess.checked;', self.app)
        self.assertIn('checking={baigaChecking}', block[:otherwise])
        self.assertNotIn('sensitiveSections', block[:section])
        # Доступ к экрану считает предикат раздела, а не роль напрямую.
        self.assertIn('const canAccessBaigaSection = canAccessBaigaSectionForUser(user);', self.app)

    def test_menu_item_is_declared_once_in_the_common_part(self):
        self.assertEqual(self.app.count("handleSidebarViewNavigation(e, 'baiga')"), 1)
        item = self.app.split("handleSidebarViewNavigation(e, 'baiga')")[1][:700]
        self.assertIn('Списки Байги', item)
        # Пункт показывает ровно условие доступа — ни шире, ни уже.
        wrapper = self.app.split("handleSidebarViewNavigation(e, 'baiga')")[0][-420:]
        self.assertIn('{canAccessBaigaSection && (\n'
                      '                                    <SidebarDeptScope section="baiga" '
                      'activeCode={activeDeptCode}>\n', wrapper)
        self.assertLess(self.app.index("handleSidebarViewNavigation(e, 'thermoboxes')"),
                        self.app.index("handleSidebarViewNavigation(e, 'baiga')"))

    def test_screen_keeps_no_copy_of_the_circle(self):
        """Круг раздела правят из листа «Доступ», поэтому экран его не считает:
        второй копии правила (отделы, должности) в App.jsx быть не должно —
        она показывала бы пункт тому, кому раздел уже закрыли."""
        for name in ('BAIGA_MANAGE_DEPARTMENT_CODE', 'BAIGA_READ_DEPARTMENT_CODES', 'BAIGA_SECTION_DEPARTMENT_CODES'):
            self.assertNotIn(name, self.app)
        body = self.app.split('const canAccessBaigaSectionForUser = (userLike) => {')[1].split('\n};')[0]
        code = [line.strip() for line in body.splitlines() if line.strip() and not line.strip().startswith('//')]
        self.assertEqual(code, [
            "if (normalizeRole(userLike?.role) === 'super_admin') return true;",
            'if (BAIGA_PILOT_SUPER_ADMIN_ONLY) return false;',
            'return userLike?.baiga_access === true;',
        ])
        # Круг по умолчанию задан на сервере — и только там.
        self.assertEqual(access.SECTION_DEPARTMENT_CODES, ('marketing', 'op', 'szov'))
        self.assertEqual([code for code, _kinds in access.CIRCLE], list(access.SECTION_DEPARTMENT_CODES))

    def _front_answers(self, users, analysts=None, predicate='canAccessBaigaSectionForUser', pilot=None):
        """Настоящий предикат из App.jsx, выполненный node, — поведение, а не текст.

        analysts — подставить свой список аналитиков вместо боевого.
        predicate — что спросить: пункт меню или замок раздела.
        pilot — подставить своё значение выключателя «только супер-админ»."""
        node = shutil.which('node')
        if not node:
            self.skipTest('node недоступен')
        app = self.app
        start = app.index('const SENSITIVE_QR_GATED_ROLES = ')
        qr_lock = app[start:app.index('\n);\n', app.index('const sensitiveSectionQrRequiredFor = ')) + 3]
        # Замок раздела спрашивает всех, кому портал выдаёт код: общий замок плюс
        # замок «Чатов водителей» (стажёр) — sensitiveQrAvailableFor.
        start = app.index('const driverChatsQrRequiredFor = (userLike) => {')
        wide_lock = app[start:app.index('\n);\n', app.index('const sensitiveQrAvailableFor = ')) + 3]
        start = app.index('const BAIGA_ANALYST_USER_IDS = ')
        section = app[start:app.index('\n);\n', app.index('const baigaQrRequiredFor = ')) + 3]
        self.assertIn('const canAccessBaigaSectionForUser = ', section)
        if analysts is not None:
            section, replaced = re.subn(r'const BAIGA_ANALYST_USER_IDS = new Set\(\[[^\]]*\]\);',
                                        'const BAIGA_ANALYST_USER_IDS = new Set(%s);' % json.dumps(list(analysts)),
                                        section)
            self.assertEqual(replaced, 1)
        if pilot is not None:
            section, replaced = re.subn(r'const BAIGA_PILOT_SUPER_ADMIN_ONLY = (?:true|false);',
                                        'const BAIGA_PILOT_SUPER_ADMIN_ONLY = %s;' % str(bool(pilot)).lower(),
                                        section)
            self.assertEqual(replaced, 1)
        script = '\n'.join((
            "import { normalizeRole, isDepartmentHead, isSupervisorRole } from '%s';"
            % (ROOT / 'src' / 'utils' / 'roles.js').as_uri(),
            qr_lock, wide_lock, section,
            'const users = %s;' % json.dumps(users),
            'process.stdout.write(JSON.stringify(users.map((u) => %s(u))));' % predicate,
        ))
        out = subprocess.run([node, '--input-type=module', '-e', script], capture_output=True, check=True)
        answers = json.loads(out.stdout.decode('utf-8'))
        self.assertEqual(len(answers), len(users))
        return answers

    @staticmethod
    def _front_user(role, code, heads, user_id=10, granted=None):
        user = {'id': user_id, 'role': role, 'department_code': code,
                'headed_department_id': 1 if heads else None, 'headed_department_codes': list(heads)}
        if granted is not None:
            # Флаг профиля: ответ сервера «раздел открыт» (bot_schedule2.py).
            user['baiga_access'] = granted
        return user

    def _profile_user(self, role, code, heads, user_id=10, grants=(), circle=None):
        """Человек, каким его видит экран: профиль с флагом, который посчитал
        сервер (как _baiga_section_open_for — access.can_open_section)."""
        ctx = person(role=role, department_code=code, headed_codes=heads, user_id=user_id, grants=grants,
                     circle=circle)
        return self._front_user(role, code, heads, user_id=user_id, granted=access.can_open_section(ctx)), ctx

    def test_menu_item_follows_the_answer_of_the_server(self):
        """Пункт меню — это ответ сервера, пришедший флагом профиля: на всей
        сетке, с кругом по умолчанию. Иначе человек видит пункт, а раздел
        отвечает отказом, — или доступ есть, а пункта нет."""
        cases = grid()
        pairs = [self._profile_user(*case) for case in cases]
        front = self._front_answers([user for user, _ctx in pairs])
        shown = 0
        for case, (_user, ctx), answer in zip(cases, pairs, front):
            self.assertEqual(answer, access.can_open_section(ctx), case)
            shown += answer
        # Сетка не вырождена: открыто заметной части и далеко не всем.
        self.assertGreater(shown, 50)
        self.assertLess(shown, len(cases) - 50)

    def test_without_the_flag_the_screen_opens_nothing_by_itself(self):
        """Должность и отдел на экране ничего не доказывают: без ответа сервера
        пункта нет ни у кого, кроме супер-админа, — даже у тех, кто в круге по
        умолчанию. И «почти true» (строка из старого кэша) ответом не считается."""
        cases = grid()
        bare = self._front_answers([self._front_user(*case) for case in cases])
        denied = self._front_answers([self._front_user(*case, granted=False) for case in cases])
        for (role, code, heads), without, refused in zip(cases, bare, denied):
            self.assertEqual((without, refused), (role == 'super_admin',) * 2, (role, code, heads))
        odd = [dict(self._front_user('operator', 'op', ()), baiga_access=value)
               for value in ('true', 1, 'yes', [], {}, None)]
        self.assertEqual(self._front_answers(odd), [False] * len(odd))

    def test_edited_circle_reaches_the_menu_through_the_flag(self):
        """Строку круга закрыли или открыли в листе «Доступ» — пункт меню следует
        за сервером у каждого, а не за должностью и отделом."""
        edits = {'staff:op': 'none', 'head:szov': 'none', 'sv:szov': 'full'}
        cases = grid()
        pairs = [self._profile_user(*case, circle=edits) for case in cases]
        front = self._front_answers([user for user, _ctx in pairs])
        changed = 0
        for case, (_user, ctx), answer in zip(cases, pairs, front):
            self.assertEqual(answer, access.can_open_section(ctx), case)
            changed += answer != access.can_open_section(person(role=case[0], department_code=case[1],
                                                               headed_codes=case[2]))
        # Правка действительно кого-то закрыла — и экран это показал.
        self.assertGreater(changed, 5)
        closed = self._profile_user('operator', 'op', (), circle=edits)[0]
        self.assertEqual(self._front_answers([closed]), [False])

    def test_named_analyst_sees_the_menu_item_and_no_lock(self):
        """Боевой список, без подстановок: сотрудник отдела аналитики видит
        пункт и входит без замка, его сосед по отделу — нет."""
        users = [self._profile_user('operator', 'analytik', (), user_id=user_id)[0]
                 for user_id in (ANALYST_ID, ANALYST_ID + 1)]
        self.assertEqual(self._front_answers(users), [True, False])
        self.assertEqual(self._front_answers(users, predicate='baigaQrRequiredFor'), [False, True])
        # Id в профиле бывает и строкой — замок от этого не меняется.
        as_text = [dict(users[0], id=str(ANALYST_ID))]
        self.assertEqual(self._front_answers(as_text, predicate='baigaQrRequiredFor'), [False])
        # Его строку в круге закрыли — пункта нет и у него.
        closed = self._profile_user('operator', 'analytik', (), user_id=ANALYST_ID,
                                    circle={'named:%d' % ANALYST_ID: 'none'})[0]
        self.assertEqual(self._front_answers([closed]), [False])

    def test_switch_closes_everyone_on_the_screen_like_on_the_server(self):
        """Выключатель «только супер-админ» стоит раньше ответа сервера: включили
        его на экране — пункта нет ни у кого, даже с флагом. На сервере он же
        закрывает и круг, и выдачи, и названных поимённо."""
        users = [self._front_user('operator', 'analytik', (), user_id=ANALYST_ID, granted=True),
                 self._front_user('super_admin', None, (), granted=True),
                 self._front_user('admin', 'marketing', ('marketing',), granted=True),
                 self._front_user('operator', 'tez', (), granted=True)]
        self.assertEqual(self._front_answers(users, pilot=True), [False, True, False, False])
        self.assertEqual(self._front_answers(users, pilot=False), [True, True, True, True])
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True):
            self.assertEqual(
                [access.can_open_section(person(role='operator', department_code='analytik', user_id=ANALYST_ID)),
                 access.can_open_section(person(role='super_admin', department_code=None)),
                 access.can_open_section(person(role='admin', department_code='marketing',
                                                headed_codes=('marketing',))),
                 access.can_open_section(person(department_code='tez', grants=('full',)))],
                [False, True, False, False])

    def test_lock_on_the_screen_answers_like_the_server(self):
        """Замок экрана и замок сервера обязаны совпадать у каждого, кому раздел
        открыт: иначе экран рисует раздел, а ручки отвечают «нужен QR», — или
        человек видит замок там, где сервер пускает. Сверяем на всех сочетаниях
        роли, отдела и главенства — по кругу и по выдаче, для человека из
        именного списка и для чужого id; с кадровиком (кода ему не выдают) и
        стажёром (по выдаче его спрашивают)."""
        cases = grid() + [('hr_manager', code, heads) for code in GRID_DEPARTMENTS + ('hr',)
                          for heads in ((), ('hr',))]
        compared = asked = 0
        with mock.patch.object(access, 'ANALYST_USER_IDS', frozenset({77})):
            for user_id in (77, 78):
                for grants in ((), ('read',)):
                    # По одному набору за запуск node: сетка целиком не влезает в
                    # предел длины одного аргумента командной строки (Linux, 128 КиБ).
                    pairs = [self._profile_user(*case, user_id=user_id, grants=grants) for case in cases]
                    users = [user for user, _ctx in pairs]
                    opened = self._front_answers(users, analysts=(77,))
                    locked = self._front_answers(users, analysts=(77,), predicate='baigaQrRequiredFor')
                    for case, (_user, ctx), is_open, is_locked in zip(cases, pairs, opened, locked):
                        self.assertEqual(is_open, access.can_open_section(ctx), (case, user_id, grants))
                        if not is_open:
                            continue
                        self.assertEqual(is_locked, access.requires_sensitive_qr(ctx), (case, user_id, grants))
                        compared += 1
                        asked += is_locked
        trainees = [self._front_user('trainee', code, (), granted=True) for code in GRID_DEPARTMENTS]
        self.assertEqual(self._front_answers(trainees, predicate='baigaQrRequiredFor'), [True] * len(trainees))
        # Сетка не вырождена: есть и те, кого замок спрашивает, и те, кого нет.
        self.assertGreater(asked, 30)
        self.assertGreater(compared - asked, 400)

    def test_menu_never_leads_past_the_qr_lock(self):
        """Должность, которой портал QR не выдаёт (кадровик), в круг не входит,
        где бы ни числилась: пункта у неё нет, а по выдаче раздел открыт ей без
        замка — запереть её было бы навсегда."""
        for code in access.SECTION_DEPARTMENT_CODES:
            user, ctx = self._profile_user('hr_manager', code, ())
            self.assertFalse(access.can_open_section(ctx), code)
            self.assertEqual(self._front_answers([user]), [False], code)
            user, ctx = self._profile_user('hr_manager', code, (), grants=('read',))
            self.assertEqual((self._front_answers([user]), self._front_answers([user], predicate='baigaQrRequiredFor')),
                             ([True], [False]), code)
            self.assertFalse(access.requires_sensitive_qr(ctx), code)

    def _profile_flag(self, ctx=None, error=None):
        """Настоящая _baiga_section_open_for из монолита — без его импорта."""
        import logging

        class Db:
            @contextmanager
            def _get_cursor(self):
                yield object()

        def load(cursor, user_id):
            if error:
                raise error
            return dict(ctx, user_id=user_id) if ctx is not None else None

        namespace = {'db': Db(), 'logging': logging}
        node = source_cache.function_copy(BOT_PY, '_baiga_section_open_for')
        module = __import__('ast').Module(body=[node], type_ignores=[])
        exec(compile(module, '<bot:_baiga_section_open_for>', 'exec'), namespace)
        for patcher in (mock.patch.object(queries, 'load_access_context', load),
                        mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', False),
                        mock.patch.object(logging, 'exception', lambda *args, **kwargs: None)):
            patcher.start()
            self.addCleanup(patcher.stop)
        return namespace['_baiga_section_open_for']

    def test_profile_flag_is_the_answer_of_the_section_rule(self):
        """Флаг считает то же правило, что гейт ручек: выдача открывает, круг
        открывает, остальным — false. Сбой базы не роняет вход в портал."""
        cases = [(person(role='sv', department_code='tez', grants=('read',)), True),
                 (person(role='sv', department_code='tez'), False),
                 (person(department_code='op'), True),
                 (person(role='trainer', department_code='szov'), False),
                 (None, False)]
        for ctx, expected in cases:
            self.assertIs(self._profile_flag(ctx)(31), expected, ctx)
        self.assertIs(self._profile_flag(person(grants=('read',)))(None), False)
        self.assertIs(self._profile_flag(person(grants=('read',)))(0), False)
        self.assertIs(self._profile_flag(error=RuntimeError('база недоступна'))(31), False)
        with mock.patch.object(access, 'ACCESS_MANAGER_USER_IDS', frozenset({31})):
            self.assertIs(self._profile_flag(person(role='sv', department_code='tez'))(31), True)

    def test_profile_flag_survives_a_failed_context_by_the_default_circle(self):
        """Контекст раздела не прочитался (база занята, идёт выкладка) — флаг
        считается по кругу по умолчанию из профиля: экран своего круга больше
        не держит, и без этого пункт меню пропал бы у всего круга разом."""
        broken = self._profile_flag(error=RuntimeError('база недоступна'))

        def profile(role, code, heads=()):
            return {'role': role, 'department_code': code,
                    'headed_department_ids': [909] if heads else [], 'headed_department_codes': list(heads)}

        for case, expected in ((profile('operator', 'op'), True), (profile('sv', 'szov'), True),
                               (profile('admin', 'marketing', ('marketing',)), True),
                               (profile('marketing_manager', 'marketing'), True),
                               (profile('operator', 'tez'), False), (profile('trainer', 'op'), False),
                               (profile('trainee', 'szov'), False), (profile('admin', None), False),
                               (profile(None, None), False), ({}, False), (None, False)):
            self.assertIs(broken(31, case), expected, case)
        self.assertIs(broken(ANALYST_ID, profile('operator', 'analytik')), True)
        self.assertIs(broken(1, profile('super_admin', None)), True)
        # На всей сетке запасной расчёт — ровно круг по умолчанию, без выдач и правок.
        for role, code, heads in grid():
            self.assertIs(broken(31, profile(role, code, heads)),
                          access.can_open_section(person(role=role, department_code=code, headed_codes=heads,
                                                         user_id=31)), (role, code, heads))
        # Выключатель закрывает и запасной расчёт.
        with mock.patch.object(access, 'PILOT_SUPER_ADMIN_ONLY', True):
            self.assertIs(broken(31, profile('operator', 'op')), False)
            self.assertIs(broken(1, profile('super_admin', None)), True)
        # Контекст прочитался — отвечает он, а не профиль: закрытая строка круга
        # пункт меню прячет, выдача — показывает.
        closed = self._profile_flag(person(department_code='op', circle={'staff:op': 'none'}))
        self.assertIs(closed(31, profile('operator', 'op')), False)
        granted = self._profile_flag(person(role='sv', department_code='tez', grants=('read',)))
        self.assertIs(granted(31, profile('sv', 'tez')), True)
        # Человека в базе нет — профиль его не «открывает».
        self.assertIs(self._profile_flag(None)(31, profile('operator', 'op')), False)

    def test_profile_carries_the_flag(self):
        payload = source_cache.function_node(BOT_PY, '_get_user_payload')
        lines = source_cache.read(BOT_PY).splitlines()
        body = '\n'.join(lines[payload.lineno - 1:payload.end_lineno])
        # Запасному расчёту профиль отдаёт и id возглавляемых отделов: по ним
        # правило отличает главу от рядового (is_department_head).
        self.assertIn("    baiga_access = _baiga_section_open_for(user_id, {\n"
                      "        'role': role, 'department_code': department_code,\n"
                      "        'headed_department_ids': headed_department_ids, "
                      "'headed_department_codes': headed_department_codes,\n"
                      "    }) if user_id is not None else False\n", body)
        self.assertIn('        "baiga_access": baiga_access,\n', body)
        # Сам предикат читает флаг ПОСЛЕ выключателя — иначе выключатель не закрыл бы выданных.
        predicate = self.app.split('const canAccessBaigaSectionForUser = (userLike) => {')[1].split('\n};')[0]
        self.assertLess(predicate.index('if (BAIGA_PILOT_SUPER_ADMIN_ONLY) return false;'),
                        predicate.index('return userLike?.baiga_access === true;'))

    def test_access_button_opens_the_sheet(self):
        view = _read(ROOT / 'src' / 'components' / 'baiga' / 'BaigaView.jsx')
        self.assertIn("import BaigaAccessSheet from './BaigaAccessSheet';", view)
        button = view.split('{capabilities.can_manage_access && (')[1].split('</button>')[0]
        self.assertIn('<button type="button" className={`${iosBtnSecondary} shrink-0`}\n'
                      '                                onClick={() => setAccessOpen(true)} '
                      'aria-label="Доступ к разделу">', button)
        self.assertIn('<span className="hidden sm:inline">Доступ</span>', button)
        sheet = view.split('{capabilities.can_manage_access && (')[2].split('/>')[0]
        self.assertIn('<BaigaAccessSheet', sheet)
        self.assertIn('open={accessOpen}', sheet)
        source = _read(ROOT / 'src' / 'components' / 'baiga' / 'BaigaAccessSheet.jsx')
        # Пять ручек листа — те же адреса, что у сервера.
        for call in ("axios.get(`${apiBaseUrl}/api/baiga/access`",
                     "axios.post(`${apiBaseUrl}/api/baiga/access/grants`",
                     "axios.patch(`${apiBaseUrl}/api/baiga/access/grants/${draft.grant.id}`",
                     "axios.delete(`${apiBaseUrl}/api/baiga/access/grants/${draft.grant.id}`",
                     "axios.patch(`${apiBaseUrl}/api/baiga/access/circle`"):
            self.assertEqual(source.count(call), 1, call)
        rules = {(rule.rule, tuple(sorted(rule.methods - {'HEAD', 'OPTIONS'})))
                 for rule in self._url_map() if '/access' in rule.rule}
        self.assertEqual(rules, {('/api/baiga/access', ('GET',)), ('/api/baiga/access/grants', ('POST',)),
                                 ('/api/baiga/access/grants/<int:grant_id>', ('PATCH',)),
                                 ('/api/baiga/access/grants/<int:grant_id>', ('DELETE',)),
                                 ('/api/baiga/access/circle', ('PATCH',))})

    @staticmethod
    def _url_map():
        if Flask is None:
            raise unittest.SkipTest('Flask не установлен')
        app = Flask('map')
        app.register_blueprint(routes.build_baiga_blueprint(
            db=None, require_api_key=lambda handler: handler, build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=lambda: (None, None, ('Unauthorized', 401)),
            sensitive_access_granted=lambda user_id: False))
        return list(app.url_map.iter_rules())

    def test_qr_status_is_asked_when_the_section_opens(self):
        """Без этого оператор, не заходивший до «Списков Байги» в другой закрытый
        раздел, остаётся на «Проверяем доступ…» без кнопки QR."""
        self.assertIn("|| view === 'baiga' || view === 'driver_mailings') {\n"
                      "                    fetchSensitiveAccessStatus();", self.app)

    def test_direct_link_without_access_leads_away_not_to_a_blank_screen(self):
        """У СЗоВ нет allowlist'а разделов, и гард отдела закрытого стажёра не
        уведёт: без своей ветки ссылка ?view=baiga оставляла пустой экран."""
        effect = self.app.split("if (view === 'baiga' && !canAccessBaigaSection) {")[1]
        branch, rest = effect.split('\n                }\n', 1)
        self.assertIn("redirectToView('hours');", branch)
        self.assertIn("redirectToView('sv_list');", branch)
        deps = rest.split('}, [', 1)[1].split(']);', 1)[0]
        self.assertIn('canAccessBaigaSection', deps)

    def test_subtitle_promises_export_only_to_those_who_have_it(self):
        view = _read(ROOT / 'src' / 'components' / 'baiga' / 'BaigaView.jsx')
        self.assertIn('subtitle={capabilities.can_export', view)
        self.assertNotIn('аналитик загрузит', view)

    def test_buttons_follow_the_capabilities(self):
        """Выгрузка — по can_export, вкладка журнала, загрузка и её окно — по
        can_manage; сам журнал открывается только с can_manage."""
        view = _read(ROOT / 'src' / 'components' / 'baiga' / 'BaigaView.jsx')
        self.assertIn("const capabilities = screen?.capabilities || {};", view)
        self.assertEqual(view.count("{capabilities.can_export && tab === 'search' && ("), 1)
        self.assertEqual(view.count('{capabilities.can_manage && ('), 3)
        self.assertIn("if (tab === 'journal' && capabilities.can_manage) {", view)
        # «Доступ» — кнопка и её окно — по своему праву: полный доступ к разделу
        # его не даёт.
        self.assertEqual(view.count('{capabilities.can_manage_access && ('), 2)
        self.assertEqual(len(re.findall(r'capabilities\.can_\w+', view)),
                         len(re.findall(r'capabilities\.can_(?:export|manage|manage_access)\b', view)))

    def test_guards_and_registries(self):
        self.assertIn("if (view === 'baiga' && canAccessBaigaSection) return;", self.app)
        self.assertIn("    baiga: ['marketing', 'op', 'szov'],", self.app)
        self.assertIn("    baiga: 'Baiga lists',", self.app)
        self.assertIn("canAccessBaigaSection && deptAllowsInner('baiga'),", self.app)
        # Тренеру раздел по кругу закрыт, но выдать его можно и ему: без строки
        # в списке гард вида уводил бы тренера с выданным разделом в «Опросы».
        # Тренера без выдачи уводит ветка view === 'baiga' (тест ниже).
        trainer = self.app.split('const TRAINER_ALLOWED_VIEWS = Object.freeze([')[1].split(']);')[0]
        # Элемент списка, а не слово в комментарии рядом с ним.
        self.assertRegex(trainer, r"(?m)^    'baiga',$")
        closed = self.app.split("if (view === 'baiga' && !canAccessBaigaSection) {")[1]
        closed = closed.split('\n                }\n')[0]
        self.assertIn("else if (isPlainTrainer) redirectToView('surveys');", closed)
        deps = self.app.split("if (view === 'baiga' && canAccessBaigaSection) return;")[1]
        deps = deps.split('}, [')[1].split(']);')[0]
        self.assertIn('canAccessBaigaSection', deps)

    def test_section_params_leave_the_address_with_the_section(self):
        for fn in ('const buildAppViewUrl = (nextView) => {', 'const syncAppViewWithUrl = (nextView) => {'):
            body = self.app.split(fn)[1].split('\n};')[0]
            self.assertIn("if (nextView !== 'baiga') stripBaigaParams(url);", body)

    def test_mirrors_of_server_constants(self):
        meta = _read(META_JS)
        self.assertIn('export const PAGE_SIZES = [%s];' % ', '.join(map(str, filters.PAGE_SIZES)), meta)
        self.assertIn("export const PRIZE_ANY = '%s';" % filters.PRIZE_ANY, meta)
        self.assertIn("export const PRIZE_NONE = '%s';" % filters.PRIZE_NONE, meta)
        for key in filters.SORT_COLUMNS:
            self.assertRegex(meta, r"\b%s: '" % key)
        for key in filters.RANGES:
            self.assertIn("'%s'" % key, meta)

    def test_menu_icons_are_components_not_elements(self):
        folder = ROOT / 'src' / 'components' / 'baiga'
        for path in folder.glob('*.jsx'):
            source = _read(path)
            for menu in re.findall(r'(?s)const \w*[mM]enuItems = \[(.*?)\];', source):
                self.assertNotRegex(menu, r'icon:\s*<', path.name)


if __name__ == '__main__':
    unittest.main()
