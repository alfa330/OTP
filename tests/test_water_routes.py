# -*- coding: utf-8 -*-
"""Роуты раздела «Учёт воды» — на подменённом слое запросов, без базы.

Что сторожится:
  * гейты: чужой отдел, QR, «колл-центр не выдаёт», «учёт ведёт руководитель»;
  * выдача перепроверяет право сама и при отказе НЕ пишет ничего — ни строки
    журнала, ни списания (db._get_cursor коммитит и то, что записано до 4xx);
  * вид выдачи, сменившийся за время открытого экрана, не выдаётся молча;
  * «Требуется закупка» уходит в Telegram один раз — на пересечении порога;
  * условия программы сохраняются целиком или не сохраняются вовсе;
  * выгрузка требует период и отдаёт журнал с телефоном и ID текстом.

SQL здесь не проверяется — его проверяет прогон на стенде с базой; зато
реальные `_office_row` и `rules` участвуют как есть.
"""

import sys
import unittest
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    from flask import Flask
except ImportError:  # pragma: no cover
    Flask = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

from water import queries as water_queries  # noqa: E402
from water import routes as water_routes  # noqa: E402
from water import schema as water_schema  # noqa: E402

TODAY = date(2026, 9, 30)
NOW = datetime(2026, 9, 30, 12, 0)
ACCOUNT = 'c' * 32
ISSUER = 424  # офисник из списка владельца (water.access.ISSUER_USER_IDS)


def person(role='operator', department_code='front_office', headed_codes=(), user_id=ISSUER):
    return {
        'user_id': user_id, 'name': 'Сотрудник %d' % user_id, 'role': role,
        'department_id': 909, 'department_code': department_code, 'city': 'Алматы',
        'headed_department_ids': [909] if headed_codes else [],
        'headed_department_codes': list(headed_codes),
    }


def crm_driver(total=0, week=0, tariffs=('econom', 'business'), account=ACCOUNT, iin='900101300123'):
    return {
        'account_id': account, 'iin': iin, 'name': 'Тестов Тест', 'phone': '+77000000001',
        'park': 'Тестовый парк', 'park_id': 'p' * 32, 'car': 'Camry · 001AAA01',
        'tariffs': list(tariffs), 'work_status': 'working', 'is_blocked': False, 'fired': False,
        'registered_at': '2026-09-01', 'orders': {'total': total, 'week': week},
        'photo_control': {'status': None, 'checked_at': None},
        'info': {'driver': {'iin': iin}},
    }


class Store:
    """Память вместо базы: ровно те функции queries, что зовут роуты."""

    def __init__(self):
        self.settings = {
            'min_trips': 20, 'cooldown_days': 7, 'welcome_blocks': 1, 'activity_blocks': 1,
            'tariffs': ['business', 'ultimate'], 'low_threshold': 20, 'buy_threshold': 10,
            'notify_user_ids': [], 'updated_by_name': None, 'updated_at': None,
        }
        self.offices = {}
        self.issues = []
        self.movements = []
        self.locks = []
        self.recipients = [{'user_id': 5, 'name': 'Руководитель', 'chat_id': 555}]
        self.candidates = [{'id': 5, 'name': 'Руководитель', 'city': None, 'has_telegram': True}]
        self.directory = [
            {'id': 43, 'city': 'Алматы', 'name': 'Офис для подключения тарифа «Бизнес»', 'address': 'Байзакова 78А'},
            {'id': 47, 'city': 'Алматы', 'name': 'Офис Алматы №1', 'address': 'Жамбыла 172В'},
            {'id': 45, 'city': 'Алматы', 'name': 'Офис для подключения тарифа «Wolt»', 'address': 'Толе би 4'},
            {'id': 44, 'city': 'Астана', 'name': 'Офис для подключения тарифа «Бизнес»', 'address': 'Сарыарка 31'},
            {'id': 49, 'city': 'Астана', 'name': 'Офис Астана', 'address': 'Сарыарка 31'},
            {'id': 65, 'city': 'Шымкент', 'name': 'Офис Шымкент', 'address': 'Республики 17'},
        ]
        self.fail_welcome_unique = False

    # офисы
    def add(self, stock=50, active=True, own_low=None, own_buy=None, wiki_id=47):
        office_id = len(self.offices) + 1
        wiki = next(item for item in self.directory if item['id'] == wiki_id)
        self.offices[office_id] = {
            'id': office_id, 'office_id': wiki_id, 'city': wiki['city'], 'name': wiki['name'],
            'address': wiki['address'], 'stock': stock, 'low_threshold': own_low,
            'buy_threshold': own_buy, 'is_active': active, 'created_at': NOW, 'updated_at': NOW,
        }
        return office_id

    def _row(self, raw, settings):
        return water_queries._office_row(
            tuple(raw[name] for name in water_queries._OFFICE_COLUMNS), settings)

    def get_settings(self, cursor):
        return dict(self.settings)

    def update_settings(self, cursor, fields, actor):
        self.settings.update(fields)
        return dict(self.settings)

    def list_offices(self, cursor, settings, include_inactive=False):
        return [self._row(raw, settings) for raw in self.offices.values()
                if include_inactive or raw['is_active']]

    def read_office(self, cursor, water_office_id, settings, for_update=False):
        raw = self.offices.get(int(water_office_id))
        return self._row(raw, settings) if raw else None

    def taken_office_ids(self, cursor):
        return {raw['office_id'] for raw in self.offices.values()}

    def directory_rows(self, cursor):
        return list(self.directory)

    def directory_office(self, cursor, office_id):
        return next((item for item in self.directory if item['id'] == office_id), None)

    def add_office(self, cursor, wiki_office, stock, low_threshold, buy_threshold, settings, actor):
        office_id = self.add(stock=stock, own_low=low_threshold, own_buy=buy_threshold,
                             wiki_id=wiki_office['id'])
        return self.read_office(cursor, office_id, settings)

    def update_office(self, cursor, water_office_id, fields, settings):
        self.offices[water_office_id].update(fields)
        return self.read_office(cursor, water_office_id, settings)

    def apply_movement(self, cursor, office, kind, delta, comment, actor):
        raw = self.offices[office['id']]
        raw['stock'] += delta
        movement = {'kind': kind, 'delta': delta, 'stock_after': raw['stock'], 'comment': comment}
        self.movements.append(movement)
        return movement

    def list_movements(self, cursor, water_office_id, limit=30):
        return list(self.movements)

    # выдачи
    def lock_driver(self, cursor, account_id, iin=None):
        self.locks.append((account_id, iin))

    def person_history(self, cursor, account_id, iin=None, limit=20):
        rows = [item for item in self.issues
                if not item.get('canceled_at')
                and (item['driver_account_id'] == account_id or (iin and item.get('driver_iin') == iin))]
        return sorted(rows, key=lambda item: item['created_at'], reverse=True)

    def issue_for_cancel(self, cursor, issue_id, for_update=False):
        item = next((row for row in self.issues if row['id'] == int(issue_id)), None)
        if not item:
            return None
        return {'id': item['id'], 'water_office_id': item['water_office_id'],
                'driver_account_id': item['driver_account_id'], 'driver_iin': item.get('driver_iin'),
                'blocks': item['blocks'], 'canceled_at': item.get('canceled_at')}

    def cancel_issue(self, cursor, issue, reason, actor):
        item = next(row for row in self.issues if row['id'] == issue['id'])
        item.update(canceled_at=NOW, canceled_by_name=actor.get('name'), cancel_reason=reason)
        self.offices[item['water_office_id']]['stock'] += item['blocks']
        return dict(item, created_at=item['created_at'].isoformat(), canceled_at=NOW.isoformat())

    def create_issue(self, cursor, office, driver, verdict, blocks, actor):
        if self.fail_welcome_unique and verdict['kind'] == 'welcome':
            from psycopg2 import errors
            raise errors.UniqueViolation('uq_water_welcome_account')
        raw = self.offices[office['id']]
        raw['stock'] -= blocks
        item = {
            'id': len(self.issues) + 1, 'water_office_id': office['id'], 'city': office['city'],
            'office_name': office['name'], 'kind': verdict['kind'], 'blocks': blocks,
            'stock_after': raw['stock'], 'driver_account_id': driver['account_id'],
            'driver_iin': driver.get('iin'), 'driver_name': driver.get('name'),
            'driver_phone': driver.get('phone'), 'driver_park': driver.get('park'),
            'driver_park_id': driver.get('park_id'), 'driver_tariffs': driver.get('tariffs'),
            'orders_total': driver['orders']['total'], 'orders_counted': verdict['trips'],
            'orders_basis': verdict['trips_basis'], 'issued_by': actor['user_id'],
            'fk_source': verdict.get('fk') if verdict['kind'] == 'welcome' else None,
            'issued_by_name': actor.get('name'), 'created_at': NOW,
        }
        self.issues.append(item)
        return dict(item, created_at=NOW.isoformat())

    def list_issues(self, cursor, filters, limit=50, offset=0):
        return list(self.issues), len(self.issues)

    def issues_for_export(self, cursor, filters, limit):
        return [dict(item, created_at=item['created_at'].isoformat()) for item in self.issues], len(self.issues)

    def filter_values(self, cursor):
        return {'staff': [], 'parks': []}

    def dashboard(self, cursor, settings, date_from, date_to, today):
        return []

    def notify_candidates(self, cursor):
        return list(self.candidates)

    def front_office_head_ids(self, cursor):
        return []

    def notify_recipients(self, cursor, settings):
        return list(self.recipients)


class FakeDb:
    def __init__(self):
        self.commits = 0
        self.rollbacks = 0

    @contextmanager
    def _get_cursor(self):
        try:
            yield object()
        except Exception:
            self.rollbacks += 1
            raise
        self.commits += 1


@unittest.skipIf(Flask is None, 'Flask не установлен')
class _Base(unittest.TestCase):
    def setUp(self):
        self.store = Store()
        self.db = FakeDb()
        self.viewer = person()
        self.qr_ok = True
        self.sent = []
        self.crm = crm_driver()
        store = self.store
        patches = {
            'load_access_context': lambda cursor, user_id: dict(self.viewer),
            'today_almaty': lambda: TODAY,
            'now_almaty': lambda: NOW,
            'directory': store.directory_rows,
        }
        for name in ('get_settings', 'update_settings', 'list_offices', 'read_office',
                     'taken_office_ids', 'directory_office', 'add_office', 'update_office',
                     'apply_movement', 'list_movements', 'lock_driver', 'person_history',
                     'create_issue', 'list_issues', 'issues_for_export', 'filter_values',
                     'issue_for_cancel', 'cancel_issue',
                     'dashboard', 'notify_candidates', 'front_office_head_ids',
                     'notify_recipients'):
            patches[name] = getattr(store, name)
        patcher = mock.patch.multiple(water_queries, **patches)
        patcher.start()
        self.addCleanup(patcher.stop)
        ready = mock.patch.object(water_schema, 'schema_is_ready', lambda cursor: True)
        ready.start()
        self.addCleanup(ready.stop)

        self.lookup_override = None

        def lookup(raw):
            if self.lookup_override:
                return self.lookup_override(raw)
            if not raw:
                raise water_routes.water_driver.DriverLookupError('Вставьте ссылку', code='bad_account_id', status=400)
            return dict(self.crm, orders=dict(self.crm['orders']))

        def send(chat_id, text, parse_mode=None, reply_markup=None):
            self.sent.append((chat_id, text, reply_markup))

            class _Ok:
                status_code = 200
            return _Ok()

        app = Flask(__name__)
        app.register_blueprint(water_routes.build_water_blueprint(
            db=self.db,
            require_api_key=lambda f: f,
            build_cors_preflight_response=lambda: ('', 204),
            resolve_requester=lambda: (self.viewer['user_id'], None, None),
            sensitive_access_granted=lambda _uid: self.qr_ok,
            send_telegram=send,
            web_app_base_url='https://example.invalid/OTP',
            driver_lookup=lookup,
        ))
        app.config['TESTING'] = True
        self.client = app.test_client()

    def issue(self, **body):
        body.setdefault('link', ACCOUNT)
        # ФК нового водителя CRM у тестового водителя не знает — сотрудник
        # подтверждает его отметкой. Отказ без отметки — FkTests.
        body.setdefault('fk_confirmed', True)
        return self.client.post('/api/water/issues', json=body)


class GateTests(_Base):
    def test_other_department_is_closed(self):
        self.viewer = person(department_code='marketing')
        response = self.client.get('/api/water/ping')
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'WATER_SECTION_CLOSED')

    def test_operator_without_qr_is_stopped(self):
        self.qr_ok = False
        response = self.client.get('/api/water/offices')
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'SENSITIVE_ACCESS_REQUIRED')

    def test_call_centre_checks_but_does_not_issue(self):
        self.viewer = person(department_code='szov')
        office_id = self.store.add()
        check = self.client.post('/api/water/check', json={'link': ACCOUNT})
        self.assertEqual(check.status_code, 200)
        self.assertTrue(check.get_json()['verdict']['allowed'])
        response = self.issue(water_office_id=office_id)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'WATER_ISSUE_FORBIDDEN')
        self.assertEqual(self.store.issues, [])

    def test_sales_check_but_do_not_issue(self):
        """ОП — права колл-центра СЗоВ (02.10.2026): проверка есть, выдачи нет."""
        self.viewer = person(department_code='op')
        office_id = self.store.add()
        check = self.client.post('/api/water/check', json={'link': ACCOUNT})
        self.assertEqual(check.status_code, 200)
        response = self.issue(water_office_id=office_id)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'WATER_ISSUE_FORBIDDEN')
        self.assertEqual(self.store.issues, [])

    def test_front_office_operator_does_not_manage(self):
        office_id = self.store.add()
        for method, url, body in (
                ('put', '/api/water/settings', {'min_trips': 5}),
                ('post', '/api/water/offices', {'office_id': 65, 'stock': 1}),
                ('post', '/api/water/offices/%d/intake' % office_id, {'blocks': 5}),
                ('post', '/api/water/offices/%d/recount' % office_id, {'stock': 5, 'comment': 'x'}),
                ('patch', '/api/water/offices/%d' % office_id, {'low_threshold': 5})):
            response = getattr(self.client, method)(url, json=body)
            self.assertEqual(response.status_code, 403, url)
            self.assertEqual(response.get_json()['code'], 'WATER_MANAGE_FORBIDDEN')
        self.assertEqual(self.store.settings['min_trips'], 20)
        self.assertEqual(self.store.offices[office_id]['stock'], 50)

    def test_ping_carries_capabilities_and_public_settings(self):
        body = self.client.get('/api/water/ping').get_json()
        self.assertTrue(body['capabilities']['can_issue'])
        self.assertFalse(body['capabilities']['can_view_journal'])
        self.assertNotIn('notify_user_ids', body['settings'])


class CheckTests(_Base):
    def test_check_answers_without_iin_and_crm_snapshot(self):
        office_id = self.store.add()
        body = self.client.post('/api/water/check',
                                json={'link': ACCOUNT, 'water_office_id': office_id}).get_json()
        self.assertNotIn('iin', body['driver'])
        self.assertNotIn('info', body['driver'])
        self.assertEqual(body['verdict']['kind'], 'welcome')
        self.assertEqual(body['verdict']['stock'], 50)

    def test_crm_outage_does_not_suggest_manual_entry(self):
        """Текст клиента CRM писался для «Посылок» («…или заполните вручную»);
        у воды ручного пути нет — право решает CRM."""
        def down(raw):
            raise water_routes.water_driver.DriverLookupError(
                'CRM не ответила — попробуйте ещё раз или заполните вручную',
                code='crm_unavailable', status=502)
        self.lookup_override = down
        response = self.client.post('/api/water/check', json={'link': ACCOUNT})
        self.assertEqual(response.status_code, 502)
        self.assertNotIn('вручную', response.get_json()['error'])

    def test_crm_error_is_passed_through_with_its_status(self):
        response = self.client.post('/api/water/check', json={'link': ''})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()['code'], 'bad_account_id')


class IssueTests(_Base):
    def test_welcome_issue_writes_the_journal_and_takes_the_stock(self):
        office_id = self.store.add(stock=50)
        response = self.issue(water_office_id=office_id, kind='welcome')
        self.assertEqual(response.status_code, 201, response.get_json())
        body = response.get_json()
        self.assertEqual(body['issue']['kind'], 'welcome')
        self.assertEqual(body['office']['stock'], 49)
        self.assertEqual(self.store.offices[office_id]['stock'], 49)
        self.assertEqual(len(self.store.issues), 1)
        self.assertEqual(self.store.locks, [(ACCOUNT, '900101300123')])
        # Вердикт после выдачи — уже без права: второй приветственный не положен.
        self.assertFalse(body['verdict']['allowed'])

    def test_refusal_writes_nothing(self):
        office_id = self.store.add(stock=50)
        self.crm = crm_driver(total=500, week=3)
        response = self.issue(water_office_id=office_id)
        self.assertEqual(response.status_code, 409)
        body = response.get_json()
        self.assertEqual(body['code'], 'WATER_NOT_ELIGIBLE')
        self.assertIn('менее 20 поездок', body['error'])
        self.assertEqual(self.store.issues, [])
        self.assertEqual(self.store.offices[office_id]['stock'], 50)

    def test_second_welcome_is_refused(self):
        office_id = self.store.add(stock=50)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 201)
        response = self.issue(water_office_id=office_id)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(len(self.store.issues), 1)

    def test_same_person_in_another_park_gets_no_second_welcome(self):
        office_id = self.store.add(stock=50)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 201)
        self.crm = crm_driver(account='d' * 32)  # тот же ИИН, другой аккаунт
        response = self.issue(water_office_id=office_id, link='d' * 32)
        self.assertEqual(response.status_code, 409)

    def test_kind_changed_while_the_screen_was_open(self):
        """Человек нажал «Выдать приветственный», а водитель уже выполнил заказ."""
        office_id = self.store.add(stock=50)
        self.crm = crm_driver(total=30, week=30)
        response = self.issue(water_office_id=office_id, kind='welcome')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['code'], 'WATER_VERDICT_CHANGED')
        self.assertEqual(self.store.issues, [])

    def test_more_blocks_than_allowed_is_refused_before_writing(self):
        office_id = self.store.add(stock=50)
        response = self.issue(water_office_id=office_id, blocks=2)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()['code'], 'WATER_BLOCKS_TOO_MANY')
        self.assertEqual(self.store.issues, [])

    def test_stock_that_ran_out_meanwhile_says_so_and_sends_a_fresh_verdict(self):
        """Настройки дают 2 блока, на экране было 3, сосед успел выдать —
        осталась одна. Ответ называет полку, а не правило, и несёт вердикт."""
        self.store.settings['activity_blocks'] = 2
        office_id = self.store.add(stock=1)
        self.crm = crm_driver(total=500, week=30)
        response = self.issue(water_office_id=office_id, blocks=2)
        self.assertEqual(response.status_code, 400)
        body = response.get_json()
        self.assertEqual(body['error'], 'В офисе осталось 1 блок')
        self.assertEqual(body['verdict']['blocks_max'], 1)
        self.assertEqual(self.store.issues, [])

    def test_setting_ceiling_reads_in_genitive(self):
        office_id = self.store.add(stock=50)
        response = self.issue(water_office_id=office_id, blocks=2)
        self.assertEqual(response.get_json()['error'], 'За одну выдачу — не больше 1 блока')

    def test_crm_acronym_keeps_its_case_in_the_refusal(self):
        office_id = self.store.add(stock=50)
        self.crm = dict(crm_driver(), orders={'total': None, 'week': None})
        body = self.issue(water_office_id=office_id).get_json()
        self.assertIn('Выдать воду нельзя: CRM не отдала', body['error'])

    def test_bad_blocks_value(self):
        office_id = self.store.add()
        for value in (0, -1, 'два', 1.5):
            response = self.issue(water_office_id=office_id, blocks=value)
            self.assertEqual(response.status_code, 400, value)
        self.assertEqual(self.store.issues, [])

    def test_empty_office_is_refused(self):
        office_id = self.store.add(stock=0)
        response = self.issue(water_office_id=office_id)
        self.assertEqual(response.status_code, 409)
        self.assertIn('в офисе нет воды', response.get_json()['error'])

    def test_inactive_office_is_not_found(self):
        office_id = self.store.add(active=False)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 404)

    def test_office_is_required(self):
        self.assertEqual(self.issue().status_code, 400)

    def test_race_on_welcome_is_a_clean_409(self):
        self.store.fail_welcome_unique = True
        office_id = self.store.add()
        response = self.issue(water_office_id=office_id)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['code'], 'WATER_WELCOME_TAKEN')
        self.assertEqual(self.db.rollbacks, 1)


class FkTests(_Base):
    """«У новых водителей обязательно должен быть пройден ФК» (30.09.2026)."""

    def test_welcome_without_confirmation_is_refused_and_writes_nothing(self):
        office_id = self.store.add(stock=50)
        response = self.issue(water_office_id=office_id, fk_confirmed=False)
        self.assertEqual(response.status_code, 409)
        self.assertIn('подтвердите, что водитель прошёл ФК', response.get_json()['error'])
        self.assertEqual(self.store.issues, [])
        self.assertEqual(self.store.offices[office_id]['stock'], 50)

    def test_confirmation_is_recorded(self):
        office_id = self.store.add(stock=50)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 201)
        self.assertEqual(self.store.issues[0]['fk_source'], 'confirmed')

    def test_passed_in_crm_needs_no_confirmation(self):
        office_id = self.store.add(stock=50)
        self.crm = dict(crm_driver(), photo_control={'status': 'passed', 'checked_at': None})
        self.assertEqual(self.issue(water_office_id=office_id, fk_confirmed=False).status_code, 201)
        self.assertEqual(self.store.issues[0]['fk_source'], 'passed')

    def test_failed_in_crm_cannot_be_overridden_by_the_toggle(self):
        office_id = self.store.add(stock=50)
        self.crm = dict(crm_driver(), photo_control={'status': 'failed', 'checked_at': None})
        response = self.issue(water_office_id=office_id, fk_confirmed=True)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.store.issues, [])

    def test_only_a_real_true_counts_as_confirmation(self):
        office_id = self.store.add(stock=50)
        self.assertEqual(self.issue(water_office_id=office_id, fk_confirmed='yes').status_code, 409)


class CancelTests(_Base):
    """Отмена ошибочной выдачи (01.10.2026): руководитель, с причиной, блоки
    возвращаются, водителю выдача больше не засчитывается."""

    def setUp(self):
        super().setUp()
        self.office_id = self.store.add(stock=10)
        self.assertEqual(self.issue(water_office_id=self.office_id).status_code, 201)
        self.issue_id = self.store.issues[0]['id']
        self.viewer = person(role='admin', headed_codes=('front_office',))

    def cancel(self, **body):
        return self.client.post('/api/water/issues/%d/cancel' % self.issue_id, json=body)

    def test_front_office_cannot_cancel(self):
        self.viewer = person()
        response = self.cancel(reason='ошиблись')
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'WATER_MANAGE_FORBIDDEN')
        self.assertIsNone(self.store.issues[0].get('canceled_at'))

    def test_reason_is_required_and_nothing_changes_without_it(self):
        for body in ({}, {'reason': '   '}):
            response = self.cancel(**body)
            self.assertEqual(response.get_json()['code'], 'WATER_CANCEL_REASON_REQUIRED')
        self.assertIsNone(self.store.issues[0].get('canceled_at'))
        self.assertEqual(self.store.offices[self.office_id]['stock'], 9)

    def test_cancel_returns_the_blocks_and_marks_the_row(self):
        response = self.cancel(reason='ошиблись водителем')
        self.assertEqual(response.status_code, 200, response.get_json())
        body = response.get_json()
        self.assertEqual(body['office']['stock'], 10)
        self.assertEqual(body['issue']['cancel_reason'], 'ошиблись водителем')
        self.assertTrue(body['issue']['canceled_at'])
        self.assertEqual(self.store.locks[-1], (ACCOUNT, '900101300123'))

    def test_second_cancel_is_refused(self):
        self.assertEqual(self.cancel(reason='ошиблись').status_code, 200)
        response = self.cancel(reason='ещё раз')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.store.offices[self.office_id]['stock'], 10)

    def test_unknown_issue_is_404(self):
        self.issue_id = 999
        self.assertEqual(self.cancel(reason='x').status_code, 404)

    def test_canceled_welcome_can_be_issued_again(self):
        """Ради этого отмена и нужна: ошибочная выдача не должна закрывать
        водителю приветственный блок."""
        self.cancel(reason='не тот водитель')
        self.viewer = person()
        response = self.issue(water_office_id=self.office_id)
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(response.get_json()['issue']['kind'], 'welcome')


def manager():
    """Руководитель регионов / фронт-офисов — глава отдела «Фронт офисы»."""
    return person(role='admin', headed_codes=('front_office',))


class PerimeterTests(_Base):
    """Кому открыт раздел и журнал (решение владельца 01.10.2026)."""

    JOURNAL = ('/api/water/issues', '/api/water/filters',
               '/api/water/issues/export?date_from=2026-09-01&date_to=2026-09-30')

    def test_office_staff_outside_the_list_are_closed(self):
        self.viewer = person(user_id=10)
        response = self.client.get('/api/water/ping')
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()['code'], 'WATER_SECTION_CLOSED')

    def test_listed_office_staff_issue_under_their_own_name(self):
        office_id = self.store.add(stock=5)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 201)
        self.assertEqual(self.store.issues[0]['issued_by'], ISSUER)

    def test_tez_is_closed_even_for_its_head(self):
        for viewer in (person(department_code='tez'),
                       person(role='admin', department_code='tez', headed_codes=('tez',))):
            self.viewer = viewer
            response = self.client.get('/api/water/ping')
            self.assertEqual(response.status_code, 403, viewer)
            self.assertEqual(response.get_json()['code'], 'WATER_SECTION_CLOSED')

    def test_trainer_of_call_centre_gets_in_like_operators(self):
        self.viewer = person(role='trainer', department_code='szov')
        body = self.client.get('/api/water/ping').get_json()
        self.assertTrue(body['capabilities']['can_open'])
        self.assertFalse(body['capabilities']['can_issue'])
        self.assertEqual(self.client.post('/api/water/check', json={'link': ACCOUNT}).status_code, 200)

    def test_journal_is_closed_to_issuers_and_the_call_centre(self):
        for viewer in (person(), person(role='sv', department_code='szov'),
                       person(role='admin', department_code='szov', headed_codes=('szov',)),
                       person(role='admin', department_code='op', headed_codes=('op',))):
            self.viewer = viewer
            for url in self.JOURNAL:
                response = self.client.get(url)
                self.assertEqual(response.status_code, 403, (viewer['role'], url))
                self.assertEqual(response.get_json()['code'], 'WATER_JOURNAL_FORBIDDEN')

    def test_journal_is_open_to_the_manager_and_super_admin(self):
        for viewer in (manager(), person(role='super_admin', department_code='szov')):
            self.viewer = viewer
            for url in self.JOURNAL[:2]:
                self.assertEqual(self.client.get(url).status_code, 200, (viewer['role'], url))


class BuyAlertTests(_Base):
    def test_crossing_the_threshold_sends_one_telegram(self):
        office_id = self.store.add(stock=11)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 201)
        self.assertEqual(len(self.sent), 1)
        chat_id, text, markup = self.sent[0]
        self.assertEqual(chat_id, 555)
        self.assertIn('Требуется закупка воды', text)
        self.assertIn('Осталось 10 блоков', text)
        self.assertEqual(markup['inline_keyboard'][0][0]['url'],
                         'https://example.invalid/OTP/?view=water')

    def test_below_the_threshold_it_stays_quiet(self):
        office_id = self.store.add(stock=9)
        self.assertEqual(self.issue(water_office_id=office_id).status_code, 201)
        self.assertEqual(self.sent, [])

    def test_recount_down_through_the_threshold_alerts_too(self):
        self.viewer = person(role='admin', headed_codes=('front_office',))
        office_id = self.store.add(stock=30)
        response = self.client.post('/api/water/offices/%d/recount' % office_id,
                                    json={'stock': 4, 'comment': 'Пересчитали полку'})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(len(self.sent), 1)

    def test_intake_never_alerts(self):
        self.viewer = person(role='admin', headed_codes=('front_office',))
        office_id = self.store.add(stock=3)
        self.client.post('/api/water/offices/%d/intake' % office_id, json={'blocks': 2})
        self.assertEqual(self.sent, [])


class ManageTests(_Base):
    def setUp(self):
        super().setUp()
        self.viewer = person(role='admin', headed_codes=('front_office',))

    def test_add_office_with_start_stock(self):
        response = self.client.post('/api/water/offices', json={'office_id': 49, 'stock': 40})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.get_json()['office']['stock'], 40)

    def test_offices_outside_the_program_are_refused(self):
        """Воду выдают только в Алматы и Астане и не в офисах Wolt (01.10.2026),
        а с 05.10.2026 и не в офисе «Бизнес» Астаны: сервер держит это сам, а
        не только список на экране."""
        for wiki_id in (65, 45, 44):
            response = self.client.post('/api/water/offices', json={'office_id': wiki_id, 'stock': 5})
            self.assertEqual(response.status_code, 400, wiki_id)
            self.assertEqual(response.get_json()['code'], 'WATER_OFFICE_OUTSIDE_PROGRAM')
            # Текст один на все причины отказа, поэтому называет каждую.
            self.assertEqual(response.get_json()['error'],
                             'Учёт воды ведётся только в офисах Алматы и Астаны, кроме офисов Wolt '
                             'и офиса «Бизнес» в Астане')
        self.assertEqual(self.store.offices, {})

    def test_almaty_business_office_is_still_added(self):
        """Убран только офис «Бизнес» Астаны — алматинский с тем же названием
        заводится, как раньше."""
        response = self.client.post('/api/water/offices', json={'office_id': 43, 'stock': 5})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.get_json()['office']['city'], 'Алматы')

    def test_add_office_twice_is_refused(self):
        self.store.add(wiki_id=47)
        response = self.client.post('/api/water/offices', json={'office_id': 47, 'stock': 1})
        self.assertEqual(response.status_code, 409)

    def test_unknown_office_is_refused(self):
        response = self.client.post('/api/water/offices', json={'office_id': 999, 'stock': 1})
        self.assertEqual(response.status_code, 404)

    def test_directory_offers_only_almaty_and_astana_without_wolt(self):
        """И без офиса «Бизнес» Астаны (44) — алматинский «Бизнес» (43) в списке."""
        body = self.client.get('/api/water/offices').get_json()
        self.assertEqual([item['id'] for item in body['directory']], [43, 47, 49])

    def test_directory_hides_offices_already_in_the_program(self):
        self.store.add(wiki_id=47)
        body = self.client.get('/api/water/offices').get_json()
        self.assertEqual([item['id'] for item in body['directory']], [43, 49])

    def test_intake_adds(self):
        office_id = self.store.add(stock=5)
        response = self.client.post('/api/water/offices/%d/intake' % office_id,
                                    json={'blocks': 20, 'comment': 'накладная 12'})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.get_json()['office']['stock'], 25)

    def test_intake_needs_a_positive_number(self):
        office_id = self.store.add(stock=5)
        for value in (0, -3, '', None, 'много'):
            response = self.client.post('/api/water/offices/%d/intake' % office_id,
                                        json={'blocks': value})
            self.assertEqual(response.status_code, 400, value)
        self.assertEqual(self.store.offices[office_id]['stock'], 5)

    def test_recount_needs_a_reason_and_a_change(self):
        office_id = self.store.add(stock=5)
        no_reason = self.client.post('/api/water/offices/%d/recount' % office_id, json={'stock': 3})
        self.assertEqual(no_reason.get_json()['code'], 'WATER_COMMENT_REQUIRED')
        same = self.client.post('/api/water/offices/%d/recount' % office_id,
                                json={'stock': 5, 'comment': 'проверили'})
        self.assertEqual(same.get_json()['code'], 'WATER_RECOUNT_SAME')
        self.assertEqual(self.store.movements, [])

    def test_office_threshold_override_and_reset(self):
        office_id = self.store.add(stock=15)
        response = self.client.patch('/api/water/offices/%d' % office_id,
                                     json={'low_threshold': 16, 'buy_threshold': 15})
        self.assertEqual(response.get_json()['office']['status'], 'buy')
        reset = self.client.patch('/api/water/offices/%d' % office_id,
                                  json={'low_threshold': None, 'buy_threshold': None})
        office = reset.get_json()['office']
        self.assertEqual((office['low_threshold'], office['buy_threshold']), (20, 10))
        self.assertEqual(office['status'], 'low')

    def test_office_buy_above_low_is_refused(self):
        office_id = self.store.add()
        response = self.client.patch('/api/water/offices/%d' % office_id, json={'buy_threshold': 25})
        self.assertEqual(response.get_json()['code'], 'WATER_THRESHOLDS_ORDER')
        self.assertIsNone(self.store.offices[office_id]['buy_threshold'])

    def test_settings_are_saved_whole_or_not_at_all(self):
        response = self.client.put('/api/water/settings',
                                   json={'min_trips': 15, 'tariffs': []})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.store.settings['min_trips'], 20)

    def test_settings_thresholds_order(self):
        response = self.client.put('/api/water/settings',
                                   json={'low_threshold': 5, 'buy_threshold': 6})
        self.assertEqual(response.get_json()['code'], 'WATER_THRESHOLDS_ORDER')

    def test_settings_save(self):
        response = self.client.put('/api/water/settings', json={
            'min_trips': 25, 'cooldown_days': 6, 'tariffs': ['business', 'ultimate', 'maybach'],
            'low_threshold': 30, 'buy_threshold': 12, 'notify_user_ids': [5]})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(self.store.settings['tariffs'], ['business', 'ultimate', 'maybach'])
        self.assertEqual(self.store.settings['notify_user_ids'], [5])

    def test_recipient_who_left_does_not_lock_the_settings(self):
        """Выбранного уволили — он пропал из кандидатов, но сохранение прочих
        условий программы проходит, а на экран он не приезжает вовсе."""
        self.store.settings['notify_user_ids'] = [5, 57]
        body = self.client.get('/api/water/settings').get_json()
        self.assertEqual(body['settings']['notify_user_ids'], [5])
        response = self.client.put('/api/water/settings', json={'min_trips': 25, 'notify_user_ids': [5, 57]})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(self.store.settings['min_trips'], 25)

    def test_deactivating_an_office_ignores_untouched_thresholds(self):
        office_id = self.store.add()
        self.store.offices[office_id]['buy_threshold'] = 25  # пара уже перевёрнута
        response = self.client.patch('/api/water/offices/%d' % office_id, json={'is_active': False})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertFalse(self.store.offices[office_id]['is_active'])

    def test_global_thresholds_cannot_flip_an_office_pair(self):
        office_id = self.store.add()
        self.store.offices[office_id]['buy_threshold'] = 15  # свой, low — общий 20
        response = self.client.put('/api/water/settings', json={'low_threshold': 12, 'buy_threshold': 10})
        self.assertEqual(response.get_json()['code'], 'WATER_THRESHOLDS_ORDER')
        self.assertIn('Офис Алматы №1', response.get_json()['error'])
        self.assertEqual(self.store.settings['low_threshold'], 20)

    def test_notify_only_known_people(self):
        response = self.client.put('/api/water/settings', json={'notify_user_ids': [5, 77]})
        self.assertEqual(response.get_json()['code'], 'WATER_NOTIFY_UNKNOWN')

    def test_bad_tariff_code(self):
        response = self.client.put('/api/water/settings', json={'tariffs': ['<b>']})
        self.assertEqual(response.status_code, 400)

    def test_settings_show_recipients_only_to_the_manager(self):
        body = self.client.get('/api/water/settings').get_json()
        self.assertIn('candidates', body)
        self.viewer = person()
        body = self.client.get('/api/water/settings').get_json()
        self.assertNotIn('candidates', body)
        self.assertNotIn('notify_user_ids', body['settings'])


@unittest.skipIf(load_workbook is None, 'openpyxl не установлен')
class ExportTests(_Base):
    def setUp(self):
        super().setUp()
        self.viewer = manager()  # журнал и выгрузка — руководителю

    def test_period_is_required(self):
        response = self.client.get('/api/water/issues/export')
        self.assertEqual(response.get_json()['code'], 'WATER_PERIOD_REQUIRED')

    def test_period_longer_than_a_year_is_refused(self):
        response = self.client.get('/api/water/issues/export?date_from=2025-01-01&date_to=2026-09-30')
        self.assertEqual(response.get_json()['code'], 'WATER_PERIOD_TOO_LONG')

    def test_workbook_carries_the_journal(self):
        office_id = self.store.add()
        self.issue(water_office_id=office_id)
        response = self.client.get('/api/water/issues/export?date_from=2026-09-01&date_to=2026-09-30')
        self.assertEqual(response.status_code, 200)
        from urllib.parse import unquote
        self.assertIn('Выдачи воды 01.09.2026 — 30.09.2026.xlsx',
                      unquote(response.headers['Content-Disposition']))
        book = load_workbook(BytesIO(response.data))
        sheet = book['Выдачи']
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(header[:5], ['Дата и время', 'Город', 'Офис', 'Сотрудник', 'Водитель'])
        row = [cell.value for cell in sheet[2]]
        self.assertEqual(row[header.index('Вид выдачи')], 'Приветственная')
        self.assertEqual(row[header.index('Телефон')], '+77000000001')
        self.assertEqual(row[header.index('ID водителя')], ACCOUNT)
        self.assertEqual(row[header.index('Тариф')], 'Эконом, Business')
        self.assertEqual(sheet.cell(row=2, column=header.index('Телефон') + 1).number_format, '@')
        context = [[cell.value for cell in line] for line in book['Контекст'].iter_rows()]
        self.assertIn(['Период', '01.09.2026 — 30.09.2026'], context)

    def test_formula_like_text_stays_text(self):
        """ФИО и парк приходят из CRM, их вводят чужие люди: «=HYPERLINK(…)»
        обязан лечь в файл строкой, а не живой формулой."""
        office_id = self.store.add()
        self.crm = dict(crm_driver(), name='=1+1', park='=HYPERLINK("http://x","p")')
        self.issue(water_office_id=office_id)
        response = self.client.get('/api/water/issues/export?date_from=2026-09-01&date_to=2026-09-30&q=%3Dcmd')
        book = load_workbook(BytesIO(response.data))
        sheet = book['Выдачи']
        header = [cell.value for cell in sheet[1]]
        for title, expected in (('Водитель', '=1+1'), ('Парк', '=HYPERLINK("http://x","p")')):
            cell = sheet.cell(row=2, column=header.index(title) + 1)
            self.assertEqual(cell.data_type, 's', title)
            self.assertEqual(cell.value, expected)
        context = {row[0].value: row[1] for row in book['Контекст'].iter_rows(min_row=2) if row[0].value}
        self.assertEqual(context['Отбор'].data_type, 's')

    def test_canceled_issue_stays_in_the_file_but_not_in_the_total(self):
        office_id = self.store.add(stock=10)
        self.issue(water_office_id=office_id)
        self.store.issues[0].update(canceled_at=NOW, canceled_by_name='Руководитель', cancel_reason='ошиблись')
        response = self.client.get('/api/water/issues/export?date_from=2026-09-01&date_to=2026-09-30')
        book = load_workbook(BytesIO(response.data))
        sheet = book['Выдачи']
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(row=2, column=header.index('Отменена') + 1).value,
                         '30.09.2026 12:00 · Руководитель · ошиблись')
        context = {row[0].value: row[1].value for row in book['Контекст'].iter_rows(min_row=2) if row[0].value}
        self.assertEqual(context['Блоков выдано'], 0)
        self.assertEqual(context['Отменено выдач (в «Блоков выдано» не входят)'], 1)

    def test_kind_filter_is_validated(self):
        response = self.client.get('/api/water/issues?kind=free')
        self.assertEqual(response.status_code, 400)


if __name__ == '__main__':
    unittest.main()
