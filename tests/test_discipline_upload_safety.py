# -*- coding: utf-8 -*-
"""«Дисциплина»: загрузки без антивируса (решение владельца 23.10.2026).

Три дыры, которые закрыл владелец 30.10.2026 («нужно обезопасить»):

* имя файла выдавало один тип за другой — невидимые и «переворачивающие»
  символы (U+202E разворачивает строку) и полосы пробелов прятали настоящее
  расширение;
* DOCX мог заставить Word при открытии загрузить что-то из сети (внешний
  шаблон, внешний объект) — такой документ теперь не принимается;
* тяжёлый шаблон акта раздувал память сервера при каждом черновике, а
  вложенные таблицы с объединёнными ячейками держали процессор минутами.

Документы здесь только безобидные (обычный DOCX, раздутый текстом шаблон,
битый архив); решения о связях и типах частей проверяются на строках.
Базы не нужно.
"""
import io
import time
import zipfile

import pytest

docx = pytest.importorskip('docx')
from docx import Document  # noqa: E402

from discipline import act_docx, files, report  # noqa: E402

RLO = chr(0x202E)          # разворачивает строку справа налево
ZWSP = chr(0x200B)         # пробел нулевой ширины
BOM = chr(0xFEFF)
ISOLATE, POP_ISOLATE = chr(0x2067), chr(0x2069)
NBSP, NARROW_NBSP = chr(0xA0), chr(0x202F)
HANGUL_FILLER = chr(0x3164)
PDF = b'%PDF-1.4\n' + b'0' * 200
W_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"')


def _docx(body_xml=None, extra=None):
    """Обычный документ python-docx; body_xml — своё тело, extra — лишние части."""
    buffer = io.BytesIO()
    Document().save(buffer)
    if body_xml is None and not extra:
        return buffer.getvalue()
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if body_xml is not None and info.filename == 'word/document.xml':
                data = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        '<w:document %s><w:body>%s<w:sectPr/></w:body></w:document>' % (W_NS, body_xml)).encode()
            target.writestr(info.filename, data)
        for name, data in (extra or {}).items():
            target.writestr(name, data)
    return out.getvalue()


def _refusal(data, **kwargs):
    with pytest.raises(files.FileError) as caught:
        files.check_docx(data, name='документ.docx', **kwargs)
    assert caught.value.code == 'DISCIPLINE_DOCX_REFUSED'
    return caught.value.message


# ── Имя файла ───────────────────────────────────────────────────────────────

def test_name_keeps_the_real_extension_visible():
    # «объяснительная‮fdp.doc» на экране читалось как «объяснительнаяcod.pdf».
    assert files.clean_name('объяснительная' + RLO + 'fdp.doc') == 'объяснительнаяfdp.doc'
    assert files.clean_name('акт' + ZWSP + '.pdf' + BOM) == 'акт.pdf'
    assert files.clean_name(ISOLATE + 'скан' + POP_ISOLATE + '.jpg') == 'скан.jpg'
    # Полоса пробелов уводила «.doc» за край строки.
    assert files.clean_name('акт.pdf' + NBSP * 40 + NARROW_NBSP + '\t.doc') == 'акт.pdf .doc'
    assert files.clean_name('акт' + HANGUL_FILLER * 5 + '.docx') == 'акт.docx'


def test_name_is_normalized_and_cyrillic_survives():
    # macOS присылает «й» двумя символами — в базе и на экране он один.
    assert files.clean_name('Объяснительная Сейтказиной.docx'.replace('й', 'и' + chr(0x306))) \
        == 'Объяснительная Сейтказиной.docx'
    assert files.safe_name('C:\\Users\\Оператор\\Рабочий стол\\акт: «копия»?.pdf') == 'акт «копия».pdf'
    assert files.safe_name(RLO + ZWSP) == 'file'


def test_long_name_is_cut_by_its_stem():
    name = files.safe_name('а' * 300 + '.docx')
    assert len(name) == files.NAME_MAX and name.endswith('.docx')
    assert files.safe_name('б' * 300) == 'б' * files.NAME_MAX


def test_upload_check_returns_the_clean_name():
    name, ext, content_type = files.check('explanation', filename='скан' + RLO + 'gpj.pdf', data=PDF)
    assert (name, ext, content_type) == ('сканgpj.pdf', 'pdf', 'application/pdf')


def test_journal_shows_file_names_the_way_the_file_list_does():
    assert report.changes_text({'file': 'объяснительная' + RLO + 'fdp.doc'}) == 'Файл: объяснительнаяfdp.doc'
    assert RLO not in report.changes_text({'files': ['акт' + RLO + 'fdp.doc', 'фото.png']})


# ── DOCX: связи пакета и типы частей — на строках ──────────────────────────

@pytest.mark.parametrize('target, remote', [
    ('https://example.com/a.dotx', True),
    ('HTTP://example.com/a.png', True),
    ('ftp://example.com/a', True),
    ('\\\\server\\share\\t.dotx', True),
    ('//server/share/t.dotx', True),
    ('file://server/share/t.dotx', True),
    ('file:///%5C%5Cserver%5Cshare%5Ct.dotx', True),
    ('mailto:hr@example.com', True),
    # Табуляция, перевод строки и пробел внутри адреса читатель ссылки выбрасывает.
    ('ht\ttp://example.com/t.dotx', True),
    ('http\n://example.com/t.dotx', True),
    (' https://example.com/t.dotx', True),
    ('file:\t//server/share/t.dotx', True),
    ('file:///C:/Users/a/AppData/Roaming/Microsoft/Templates/Normal.dotm', False),
    # Word для Mac: свой диск, без хоста — не сеть.
    ('file:///Users/a/Library/Group Containers/UBF8T346G9.Office/User Content/Templates/Акт.dotx', False),
    ('Macintosh HD:Users:a:Library:Normal.dotm', False),
    ('file://localhost/C:/Users/a/t.dotx', False),
    ('file:////server/share/t.dotx', True),
    ('file:///C:\\Users\\a\\Documents\\Акт.dotx', False),
    ('C:\\Users\\a\\Шаблон.dotx', False),
    ('Normal.dotm', False),
    ('media/image1.png', False),
    ('../customXml/item1.xml', False),
    ('', False),
    (None, False),
])
def test_remote_target(target, remote):
    assert files.remote_target(target) is remote


REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'


@pytest.mark.parametrize('rel_type, target, mode, refused', [
    (REL + 'image', 'media/image1.png', None, None),
    (REL + 'styles', 'styles.xml', None, None),
    (REL + 'hyperlink', 'https://example.com/правила', 'External', None),
    ('http://purl.oclc.org/ooxml/officeDocument/relationships/hyperlink', 'https://example.com', 'External', None),
    (REL + 'attachedTemplate', 'file:///C:/Users/a/Templates/Акт.dotx', 'External', None),
    (REL + 'attachedTemplate', 'https://example.com/t.dotx', 'External', 'шаблон из сети'),
    (REL + 'attachedTemplate', '\\\\server\\share\\t.dotx', 'external', 'шаблон из сети'),
    (REL + 'image', 'https://example.com/a.png', 'External', 'внешний объект'),
    (REL + 'image', 'C:/Users/a/a.png', 'External', 'внешний объект'),
    # Ссылка в сеть без пометки «внешняя» — всё равно внешняя.
    (REL + 'image', 'https://example.com/a.png', None, 'внешний объект'),
    (REL + 'attachedTemplate', 'ht\ttp://example.com/t.dotx', 'External', 'шаблон из сети'),
    # Кусок другого формата, который Word вливает при открытии, — даже внутренний.
    (REL + 'aFChunk', 'afchunk.htm', None, 'подмешивает'),
])
def test_relationship_problem(rel_type, target, mode, refused):
    problem = files.relationship_problem(rel_type, target, mode)
    if refused is None:
        assert problem is None
    else:
        assert refused in problem and 'PDF' in problem


def test_content_types_problem():
    main = files._WORD_MAIN
    assert files.content_types_problem([main, 'application/xml', 'image/png']) is None
    assert files.content_types_problem([main.upper()]) is None
    assert 'не документ Word' in files.content_types_problem(['application/xml'])
    assert 'не документ Word' in files.content_types_problem([
        'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml'])
    for kind in ('application/vnd.ms-word.document.macroEnabled.main+xml', 'application/vnd.ms-office.vbaProject',
                 'application/vnd.ms-office.activeX+xml'):
        assert 'макросы' in files.content_types_problem([main, kind])


# ── DOCX: пакет целиком — безобидные файлы ─────────────────────────────────

def test_ordinary_document_passes_as_any_kind():
    data = _docx()
    files.check_docx(data, name='объяснительная.docx')
    files.check_docx(data, name='шаблон.docx', template=True)
    name, ext, _type = files.check('explanation', filename='объяснительная.docx', data=data)
    assert (name, ext) == ('объяснительная.docx', 'docx')


def test_not_a_word_package_is_refused():
    assert 'повреждён' in _refusal(b'PK\x03\x04' + b'0' * 50)
    other = io.BytesIO()
    with zipfile.ZipFile(other, 'w') as archive:
        archive.writestr('notes.txt', 'просто архив')
    assert 'повреждён' in _refusal(other.getvalue())
    without_word = io.BytesIO()
    with zipfile.ZipFile(without_word, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
                                                'content-types"><Default Extension="xml" '
                                                'ContentType="application/xml"/></Types>')
    assert 'не документ Word' in _refusal(without_word.getvalue())


def test_duplicate_part_names_are_refused():
    # Имена частей пакета регистр не различают: проверке — одна часть, Word — другая.
    data = _docx(extra={'word/Document.xml': b'<x/>'})
    assert 'повреждён' in _refusal(data)


def test_unpacked_size_is_bounded():
    data = _docx(extra={'word/media/fill.bin': b'\0' * (files.TEMPLATE_MAX_UNPACKED + 1)})
    assert len(data) < 1024 * 1024
    files.check_docx(data, name='объяснительная.docx')
    assert 'в распакованном виде больше' in _refusal(data, template=True)


def _renamed_main_part(body_xml):
    """Тот же безобидный документ, но основная часть лежит в word/document.dat:
    python-docx находит её по связи и типу из [Content_Types].xml, не по имени."""
    source = zipfile.ZipFile(io.BytesIO(_docx(body_xml)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            name = info.filename.replace('word/document.xml', 'word/document.dat').replace(
                'word/_rels/document.xml.rels', 'word/_rels/document.dat.rels')
            if info.filename in ('[Content_Types].xml', '_rels/.rels'):
                data = data.replace(b'document.xml', b'document.dat')
            target.writestr(name, data)
    return out.getvalue()


def test_template_limits_follow_the_part_type_not_its_name():
    """Пределы шаблона обходились переименованием части: считались только
    *.xml, а разбирает python-docx по типу (разбор, 30.10.2026)."""
    data = _renamed_main_part('<w:p><w:r><w:t>{{ФИО}} {{НОМЕР}}</w:t></w:r></w:p>' * 90_000)
    assert 'word/document.xml' not in zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert 'больше 4 МБ' in _refusal(data, template=True)
    small = _renamed_main_part('<w:p><w:r><w:t>{{ФИО}}</w:t></w:r></w:p>')
    files.check_docx(small, name='шаблон.docx', template=True)
    assert [p.text for p in Document(io.BytesIO(small)).paragraphs] == ['{{ФИО}}']


def test_doc_must_be_a_real_word_97_file():
    """RTF, HTML и текст с расширением .doc Word открывает по содержимому, и
    проверка DOCX обходилась бы «Сохранить как .doc» (разбор, 30.10.2026)."""
    for data in (b'{\\rtf1\\ansi hello}', b'<html><body>hello</body></html>', 'просто текст'.encode()):
        with pytest.raises(files.FileError) as caught:
            files.check('explanation', filename='объяснительная.doc', data=data)
        assert 'не документ Word 97–2003' in caught.value.message
    word97 = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\0' * 512
    assert files.check('explanation', filename='объяснительная.doc', data=word97)[1:] == ('doc', 'application/msword')


def test_template_xml_and_elements_are_bounded():
    # Архив на десятки КБ, XML на мегабайты: так шаблон раздувал память сервера.
    heavy_text = _docx('<w:p><w:r><w:t>{{ФИО}} {{НОМЕР}}</w:t></w:r></w:p>' * 90_000)
    assert len(heavy_text) < 200 * 1024
    files.check_docx(heavy_text, name='объяснительная.docx')
    assert 'больше 4 МБ' in _refusal(heavy_text, template=True)
    many_elements = _docx('<w:p/>' * (files.TEMPLATE_MAX_ELEMENTS + 1))
    assert 'больше 100000 элементов' in _refusal(many_elements, template=True)
    files.check_docx(_docx('<w:p/>' * 5_000), name='шаблон.docx', template=True)


# ── Шаблон акта: обход за один проход и проверка перед разбором ────────────

def _nested_table(depth, inner):
    grid = '<w:tblGrid>' + '<w:gridCol w:w="100"/>' * 20 + '</w:tblGrid>'
    body = inner
    for _level in range(depth):
        body = ('<w:tbl><w:tblPr/>%s<w:tr><w:tc><w:tcPr><w:gridSpan w:val="20"/></w:tcPr>%s<w:p/></w:tc></w:tr>'
                '</w:tbl>' % (grid, body))
    return body


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _texts(data):
    """Текст каждого абзаца тела — с таблицами и полями содержимого."""
    body = Document(io.BytesIO(data)).element.body
    return [''.join(t.text or '' for t in p.iter(W + 't')) for p in body.iter(W + 'p')]


def test_nested_merged_tables_are_walked_once():
    # Ячейка на 20 колонок сетки приходила из row.cells 20 раз: глубина 5 —
    # 3,2 млн проходов и 74 с на сервере (замер 30.10.2026). Теперь — один.
    template = _docx(_nested_table(7, '<w:p><w:r><w:t>{{ФИО}}</w:t></w:r></w:p>') + '<w:p/>')
    started = time.monotonic()
    result = act_docx.fill_template(template, {'{{ФИО}}': 'Иванова Анна'})
    assert time.monotonic() - started < 2
    assert 'Иванова Анна' in _texts(result)


def test_placeholders_in_content_controls_and_text_runs_are_filled():
    body = ('<w:sdt><w:sdtContent><w:p><w:r><w:t>{{НОМЕР_АКТА}}</w:t></w:r></w:p></w:sdtContent></w:sdt>'
            '<w:p><w:r><w:t>{{ФИ</w:t></w:r><w:r><w:t>О}}</w:t></w:r></w:p>')
    texts = _texts(act_docx.fill_template(_docx(body), {'{{НОМЕР_АКТА}}': 'А-1', '{{ФИО}}': 'Иванова Анна'}))
    assert texts[:2] == ['А-1', 'Иванова Анна']


def test_heavy_template_is_refused_before_parsing():
    # Черновик скачивает СВ: ему не совет загрузчику, а кто и что должен сделать.
    heavy = _docx('<w:p><w:r><w:t>{{ФИО}}</w:t></w:r></w:p>' * 150_000)
    with pytest.raises(files.FileError) as caught:
        act_docx.fill_template(heavy, {'{{ФИО}}': 'Иванова Анна'})
    assert (caught.value.code, caught.value.status) == ('DISCIPLINE_TEMPLATE_REFUSED', 409)
    assert 'руководителю нужно загрузить его заново' in caught.value.message


def test_text_box_next_to_a_placeholder_survives_and_the_rest_is_filled():
    """Абзац «АКТ № …» с надписью «УТВЕРЖДАЮ …»: живой обход дерева обрывался
    на нём, и метки ниже оставались незаполненными; присваивание текста
    прогону стирало саму надпись (разбор, 30.10.2026)."""
    body = ('<w:p><w:r><w:t>АКТ № {{НОМЕР_АКТА}}</w:t></w:r>'
            '<w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
            '<w:p><w:r><w:t>УТВЕРЖДАЮ {{СОСТАВИЛ}}</w:t></w:r></w:p>'
            '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
            '<w:p><w:r><w:t>Сотрудник: {{ФИО}}</w:t></w:r></w:p>'
            '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>{{РЕШЕНИЕ}}</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    result = act_docx.fill_template(_docx(body), {'{{НОМЕР_АКТА}}': 'А-7', '{{СОСТАВИЛ}}': 'Сидорова Жанна',
                                                  '{{ФИО}}': 'Иванова Анна', '{{РЕШЕНИЕ}}': 'Выговор'})
    xml = zipfile.ZipFile(io.BytesIO(result)).read('word/document.xml').decode('utf-8')
    assert '{{' not in xml
    for text in ('АКТ № А-7', 'УТВЕРЖДАЮ Сидорова Жанна', 'Сотрудник: Иванова Анна', 'Выговор'):
        assert text in xml
    assert 'v:textbox' in xml


def test_picture_before_the_placeholder_keeps_its_place():
    """Логотип в начале строки «АКТ № …» — отдельный прогон: текст ложится в
    прогон, где он был, и картинка не уезжает за него."""
    body = ('<w:p><w:r><w:pict><v:shape><v:imagedata/></v:shape></w:pict></w:r>'
            '<w:r><w:t>АКТ № {{НОМЕР</w:t></w:r><w:r><w:t>_АКТА}}</w:t></w:r></w:p>')
    result = act_docx.fill_template(_docx(body), {'{{НОМЕР_АКТА}}': 'А-7'})
    xml = zipfile.ZipFile(io.BytesIO(result)).read('word/document.xml').decode('utf-8')
    assert '{{' not in xml and 'АКТ № А-7' in xml
    assert xml.index('v:imagedata') < xml.index('АКТ № А-7')


def test_svg_picture_is_not_template_xml():
    """SVG — тоже XML, но python-docx хранит картинку байтами и не разбирает:
    в предел XML шаблона она не входит (второй разбор, 30.10.2026)."""
    types = ({'svg': 'image/svg+xml', 'xml': 'application/xml'},
             {'/word/document.dat': files._WORD_MAIN})
    assert files._is_xml_part('word/media/image1.svg', types) is False
    assert files._is_xml_part('word/styles.xml', types) is True
    assert files._is_xml_part('word/document.dat', types) is True
    assert files._is_xml_part('word/media/image2.png', types) is False


def test_draft_is_logged_only_after_it_is_built():
    """Отказ черновика по шаблону не должен оставлять в журнале «Выгружен
    черновик акта»: запись — после сборки."""
    import inspect
    from discipline import routes
    source = inspect.getsource(routes)
    draft = source[source.index("def discipline_act_draft"):source.index("def discipline_explanation_blank")]
    assert draft.index("act_docx.build(") < draft.index("'act_docx', 'act'")
