# -*- coding: utf-8 -*-
"""Выгрузка журнала ознакомления в Excel (ТЗ #300, п.14).

Отдельный модуль, а не кусок routes.py: сборка книги — это чистая работа над
готовыми строками, без базы и без flask, и проверять её удобно без сервера.

Оформление шапки взято у выгрузки задач (bot_schedule2.py:
_task_export_fill_sheet), цвета ответов — у выгрузки «Опросов»: человек
открывает эти файлы в одном Excel, и разные шапки и разные «зелёный — верно»
читались бы как разные продукты.

Листы повторяют экран «Результаты»: «Ознакомление» — список сотрудников,
«Попытки» — разбор каждой попытки, «Вопросы» — где ошибаются. Двух последних
нет у новости без теста.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import access as news_access

# (ключ, заголовок, ширина). Порядок — порядок чтения: сначала кто, потом что
# с ним произошло. Первые восемь — «минимальные поля выгрузки» из ТЗ дословно,
# остальные добавлены теми же данными, что уже есть в журнале на экране.
COLUMNS = (
    ('name', 'ФИО', 32),
    ('user_id', 'ID сотрудника', 14),
    ('department', 'Подразделение', 30),
    ('role', 'Должность', 18),
    ('published_at', 'Дата публикации', 18),
    ('shown_at', 'Первый показ', 18),
    ('confirmed_at', 'Дата ознакомления', 18),
    ('score', 'Результат теста', 16),
    ('attempts', 'Попыток', 10),
    ('trainer', 'Тренажёр', 12),
    ('wave', 'Волна', 9),
    ('wave_planned_at', 'Волна по плану', 18),
    ('wave_activated_at', 'Волна фактически', 18),
    # Время в окне Oktell (23.09.2026). У объявлений в портал замера нет, и
    # колонки пропадают сами — как пустые (см. build).
    ('read_spent', 'На чтении', 11),
    ('quiz_spent', 'На тесте', 11),
    ('overtime', 'Сверх отведённого', 22),
    ('status', 'Статус', 34),
)

# Лист «Попытки»: одна строка — одна попытка, дальше по колонке на вопрос.
ATTEMPT_COLUMNS = (
    ('name', 'ФИО', 32),
    ('user_id', 'ID сотрудника', 14),
    ('department', 'Подразделение', 30),
    ('attempt_no', 'Попытка', 10),
    ('created_at', 'Время', 18),
    ('score', 'Верных ответов', 16),
    ('needed', 'Нужно для зачёта', 12),
    ('result', 'Итог', 15),
)

# Лист «Вопросы»: одна строка — один вопрос. Не строка на вариант, как в
# «Опросах»: там ответы бывают рейтингом и текстом, здесь вариант всегда один
# верный, и руководителю нужен ответ «какой вопрос заваливают и что выбирают
# вместо верного» — он читается в одной строке.
QUESTION_COLUMNS = (
    ('number', '№', 6),
    ('prompt', 'Вопрос', 44),
    ('correct', 'Верный ответ', 28),
    ('answered', 'Ответили', 11),
    ('wrong', 'Ошиблись', 11),
    ('wrong_share', 'Ошиблись, %', 12),
    ('top_wrong', 'Чаще всего выбирали неверный', 30),
    ('distribution', 'Все ответы', 46),
)

ATTEMPTS_SHEET = 'Попытки'
# По первой попытке — как вкладка «Вопросы» на экране (queries.question_stats).
# Сказано в имени листа: иначе «ошиблись 5 из 80» при двадцати попытках у
# одного человека читалось бы как неправда.
QUESTIONS_SHEET = 'Вопросы (1-я попытка)'

# СТРОЧНЫМИ. «DD.MM.YYYY» Excel понимает, а Numbers и Quick Look — то, чем
# файл открывается на iPhone и Mac, — читают «D» как день ГОДА: 21 сентября
# превращалось в «264.09.2026». Поймано на боевой выгрузке.
DATE_FORMAT = 'dd.mm.yyyy hh:mm'

_DATE_KEYS = ('published_at', 'shown_at', 'confirmed_at', 'created_at',
              'wave_planned_at', 'wave_activated_at')
_NUMBER_KEYS = ('user_id', 'attempts', 'wave', 'attempt_no', 'needed',
                'number', 'answered', 'wrong')

# Цвета ответа — те же, что в выгрузке «Опросов» (bot_schedule2.py:
# answer_correct_fill и соседи).
_RIGHT = (PatternFill(fill_type='solid', fgColor='DCFCE7'), Font(color='166534'))
_WRONG = (PatternFill(fill_type='solid', fgColor='FEF2F2'), Font(color='991B1B'))
_EMPTY = (PatternFill(fill_type='solid', fgColor='F3F4F6'), Font(color='6B7280'))

# Перерасход времени — красным, как в журнале на экране (решение владельца
# 25.09.2026: «сохранять таких и подмечать красным»). Тот же красный, что у
# неверного ответа: файл открывают рядом с «Опросами». Красим и ФИО — строку
# превысившего находят глазами, не листая до колонки «Сверх отведённого».
_OVER = _WRONG

_QUESTION_TITLE_LIMIT = 100


def _duration(seconds):
    """«4:12», «1:02:05». Текстом, а не долей суток: формат длительности
    Numbers и Quick Look читают по-своему (см. DATE_FORMAT выше), а время в
    окне читают глазами, не складывают."""
    if seconds is None:
        return ''
    total = max(0, int(seconds))
    hours, rest = divmod(total, 3600)
    minutes, secs = divmod(rest, 60)
    return ('%d:%02d:%02d' % (hours, minutes, secs) if hours
            else '%d:%02d' % (minutes, secs))


def _overtime(row):
    """«чтение +1:12 · тест +0:30» — на сколько превышено отведённое время."""
    parts = []
    for key, label in (('read_over_seconds', 'чтение'), ('quiz_over_seconds', 'тест')):
        over = row.get(key)
        if over:
            parts.append('%s +%s' % (label, _duration(over)))
    return ' · '.join(parts)


def _moment(value):
    """ISO-строка журнала → datetime для ячейки. Не разобрали — пусто.

    Время в базе наивное (часы Алматы), таким и уходит в файл: перевод в UTC
    сдвинул бы «подтверждено в 09:14» на пять часов у всех сразу.
    """
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def _score(correct, total):
    # «2 из 3» — так же, как на экране. Числом долю не пишем: в Excel её
    # немедленно сложат с другими долями и получат бессмыслицу.
    return '%s из %s' % (correct or 0, total) if total else ''


def _value(key, row, post):
    if key == 'user_id':
        return row.get('user_id')
    if key == 'name':
        return row.get('name') or '—'
    if key == 'department':
        return row.get('department_name') or ''
    if key == 'role':
        return news_access.ROLE_TITLES.get(
            str(row.get('role') or '').lower(), row.get('role') or '')
    if key == 'published_at':
        return _moment(post.get('published_at'))
    if key in _DATE_KEYS:
        return _moment(row.get(key))
    if key == 'score':
        return _score(row.get('last_correct'), row.get('last_total'))
    if key == 'attempts':
        return row.get('attempts') or 0
    if key == 'trainer':
        return 'пройден' if row.get('trainer_passed_at') else ''
    if key == 'wave':
        return row.get('wave_no') or ''
    if key == 'read_spent':
        return _duration(row.get('read_spent_seconds'))
    if key == 'quiz_spent':
        return _duration(row.get('quiz_spent_seconds'))
    if key == 'overtime':
        return _overtime(row)
    if key == 'status':
        # Подпись берётся у сервера, а не пишется здесь второй раз: файл и
        # экран обязаны называть одно состояние одним словом.
        label = news_access.PERSON_STATUS_LABELS.get(row.get('status'), '')
        return label + ('' if row.get('in_audience', True) else ' (уже не в адресатах)')
    return ''


def _sheet(book, title, headers, rows, *, first=False):
    """Лист с оформленной шапкой. headers: [(ключ, заголовок, ширина)],
    rows: [[значение, …]] в том же порядке. Возвращает лист."""
    sheet = book.active if first else book.create_sheet()
    sheet.title = title

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='1F2937')
    for column, (_key, text, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width

    for index, values in enumerate(rows, start=2):
        for column, ((key, _text, _width), value) in enumerate(zip(headers, values), start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            if key in _DATE_KEYS:
                cell.number_format = DATE_FORMAT
            elif key in _NUMBER_KEYS:
                cell.number_format = '0'

    # Шапка закреплена и с фильтром: журнал на двести человек иначе не читается.
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)),
                                         max(len(rows) + 1, 2))
    return sheet


def _question_title(question):
    # «Q1: …» — как колонки ответов в выгрузке «Опросов». Длинный вопрос
    # режем: шапка на полэкрана прячет сами ответы.
    text = ' '.join(str(question.get('prompt') or '').split())
    if len(text) > _QUESTION_TITLE_LIMIT:
        text = text[:_QUESTION_TITLE_LIMIT - 1].rstrip() + '…'
    return 'Q%s: %s' % (question.get('number'), text)


def _attempts_sheet(book, rows, questions, attempts):
    """Каждая попытка каждого человека с ответами — «Ответы» опросов, только
    попыток у человека бывает двадцать, и каждая здесь своей строкой."""
    known = {row.get('user_id'): row for row in rows}
    ordered = sorted(attempts, key=lambda item: (
        str((known.get(item['user_id']) or item).get('name') or '').casefold(),
        item['user_id'], item['attempt_no']))

    headers = list(ATTEMPT_COLUMNS) + [
        ('q%s' % question['id'], _question_title(question), 28) for question in questions]
    table, marks = [], []
    for item in ordered:
        person = known.get(item['user_id']) or {}
        line = [
            person.get('name') or item.get('name') or '—',
            item['user_id'],
            person.get('department_name') or '',
            item['attempt_no'],
            _moment(item.get('created_at')),
            _score(item['correct'], item['total']),
            item['needed'],
            'Засчитана' if item['passed'] else 'Не засчитана',
        ]
        tones = []
        for question in questions:
            chosen = item['answers'].get(str(question['id']))
            labels = [option['label'] for option in question['options']]
            try:
                index = int(chosen)
            except (TypeError, ValueError):
                index = None
            if index is None or not 0 <= index < len(labels):
                line.append('—')
                tones.append(_EMPTY)
            else:
                line.append(labels[index])
                tones.append(_RIGHT if index == question['correct'] else _WRONG)
        table.append(line)
        marks.append((item['passed'], tones))

    sheet = _sheet(book, ATTEMPTS_SHEET, headers, table)
    first_answer = len(ATTEMPT_COLUMNS) + 1
    for index, (passed, tones) in enumerate(marks, start=2):
        verdict = sheet.cell(row=index, column=len(ATTEMPT_COLUMNS))
        verdict.fill, verdict.font = _RIGHT if passed else _WRONG
        for offset, (fill, font) in enumerate(tones):
            cell = sheet.cell(row=index, column=first_answer + offset)
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    return sheet


def _questions_sheet(book, questions):
    """Где ошибаются (ТЗ #300, п.12) — вкладка «Вопросы» экрана одной строкой
    на вопрос."""
    table = []
    for question in questions:
        options = question['options']
        right = next((option for option in options if option['is_correct']), None)
        wrong = [option for option in options if not option['is_correct'] and option['count']]
        top = max(wrong, key=lambda option: option['count']) if wrong else None
        table.append([
            question['number'],
            question['prompt'],
            right['label'] if right else '',
            question['answered'],
            question['wrong_people'],
            # Долей, а не «38»: с форматом «0%» Excel и покажет 38%, и
            # отсортирует как число.
            (question['wrong_people'] / question['answered']) if question['answered'] else 0,
            ('%s — %s' % (top['label'], top['count'])) if top else '',
            '\n'.join('%s — %s (%s%%)%s' % (
                option['label'], option['count'],
                ('%g' % option['percent']).replace('.', ','),
                ' · верный' if option['is_correct'] else '') for option in options),
        ])
    sheet = _sheet(book, QUESTIONS_SHEET, QUESTION_COLUMNS, table)
    for index in range(2, len(table) + 2):
        sheet.cell(row=index, column=6).number_format = '0%'
        for column in (2, 3, 7, 8):
            sheet.cell(row=index, column=column).alignment = Alignment(
                wrap_text=True, vertical='top')
        sheet.cell(row=index, column=3).fill, sheet.cell(row=index, column=3).font = _RIGHT
    return sheet


def build(post, rows, questions=(), attempts=()):
    """Книга Excel по журналу новости. Возвращает BytesIO, готовый к отдаче.

    questions — queries.question_stats, attempts — queries.post_attempts. Оба
    пусты у новости без теста, и тогда в книге один лист.
    """
    values = {key: [_value(key, row, post) for row in rows]
              for key, _title, _width in COLUMNS}
    # Колонку, пустую во ВСЕХ строках, не выводим: заголовок без единого
    # значения — шум, из-за которого лист приходится листать вбок. У новости
    # без теста, тренажёра и волн так уходит половина таблицы.
    headers = [column for column in COLUMNS
               if not rows or any(value not in (None, '', 0) for value in values[column[0]])
               or column[0] in ('user_id', 'attempts')]

    book = Workbook()
    sheet = _sheet(book, 'Ознакомление', headers,
                   [[values[key][index] for key, _t, _w in headers]
                    for index in range(len(rows))],
                   first=True)
    keys = [key for key, _t, _w in headers]
    if 'overtime' in keys:
        marked = [keys.index(key) + 1 for key in ('overtime', 'name') if key in keys]
        for index, value in enumerate(values['overtime'], start=2):
            if value:
                for column in marked:
                    cell = sheet.cell(row=index, column=column)
                    cell.fill, cell.font = _OVER
    if questions and attempts:
        _attempts_sheet(book, rows, questions, attempts)
    if questions:
        _questions_sheet(book, questions)

    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream
