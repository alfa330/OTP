"""Выгрузка таблицы термокоробов в xlsx — «Экспорт в Excel» из постановки #363.

Колонки и подписи — как в листе «Условия выдачи коробов», из которого таблица
перенесена: файл должен читаться тем, кто вёл Google-таблицу, без перевода.

Остатки пишутся числом, а не текстом: зелёного уголка «число сохранено как
текст» у них не будет, и сумма по колонке в Excel считается сама.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import rules

# (заголовок, ширина, функция значения)
_COLUMNS = (
    ('Город', 16, lambda row: row.get('city') or ''),
    ('Адрес', 38, lambda row: row.get('address') or row.get('name') or ''),
    ('Бесплатные термокороба', 14, lambda row: row.get('free_boxes') or 0),
    ('Термопакет', 12, lambda row: row.get('thermo_bags') or 0),
    ('Б/У короб', 10, lambda row: row.get('used_boxes') or 0),
    ('Тариф', 26, lambda row: rules.tariff_label(row.get('tariff'))),
    ('Заказы', 14, lambda row: rules.orders_label(row.get('min_orders') or 0)),
    ('Срок', 16, lambda row: rules.period_label(row.get('period_days') or 7)),
    ('Отдельное условие', 22, lambda row: row.get('special_condition') or ''),
    ('Депозит', 15, lambda row: rules.deposit_label(row.get('deposit_tenge') or 0)),
    ('Обновлено', 18, lambda row: _stamp(row.get('updated_at'))),
    ('Кто обновил', 24, lambda row: row.get('updated_by_name') or ''),
)

_NUMERIC = {'Бесплатные термокороба', 'Термопакет', 'Б/У короб'}


def _stamp(value):
    """«2026-10-01T14:20:05» → «01.10.2026 14:20» — часы Алматы, как в базе."""
    text = str(value or '')
    if len(text) < 16:
        return text
    return '%s.%s.%s %s' % (text[8:10], text[5:7], text[0:4], text[11:16])


def build(rows):
    """Книга с одним листом. Возвращает байты xlsx."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Условия выдачи коробов'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='334155')
    for index, (title, width, _) in enumerate(_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = 'A2'

    for row_index, row in enumerate(rows, start=2):
        for column_index, (title, _, value_of) in enumerate(_COLUMNS, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value_of(row))
            if title in _NUMERIC:
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(vertical='top', wrap_text=True)

    sheet.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(_COLUMNS)), max(1, len(rows) + 1))

    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def file_name(day):
    """«Термокороба 01.10.2026.xlsx». Пара — exportFileName в thermoboxMeta.js."""
    return 'Термокороба %s.xlsx' % day.strftime('%d.%m.%Y')
