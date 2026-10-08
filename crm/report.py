# -*- coding: utf-8 -*-
"""Выгрузка «Обращений» в xlsx за выбранный период (владелец, 30.09.2026:
«выгрузка в формате Excel по выбранному периоду по всем обращениям»).

Два листа: «Контекст» и «Обращения». Контекст — первым, как у остальных
выгрузок портала («Жалобы», «Посылки», «Касания»): обращения живые, ответы
дописываются днями, и через месяц по файлу иначе не понять, на какую дату он
собран и что в него вошло.

Соглашения те же, что в complaints/report.py:

* ИИН и телефон — ТЕКСТОМ (числом ИИН уезжает в экспоненту, телефон теряет
  ведущие нули), зелёный уголок гасится тегом <ignoredErrors>; функция
  приходит аргументом, потому что живёт в монолите;
* даты — настоящими датами, чтобы сортировались и шли в сводную;
* книга потоковая (write_only): ширины и закрепление задаются до первой строки.

Жалобы в эту выгрузку не входят: у них свой круг доступа (сотрудник, на
которого жалоба, не видит её никогда) и своя выгрузка в разделе «Жалобы».
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import scenarios, schema

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
LABEL_FONT = Font(bold=True)
WRAP = Alignment(vertical='top', wrap_text=True)
DATE_FORMAT = 'DD.MM.YYYY HH:MM'

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Подписи статусов — те же, что в ленте раздела (CrmTicketsView STATUS_META).
STATUS_TITLES = {
    'open': 'Отправлено',
    'in_progress': 'В работе',
    'answered': 'Есть ответ',
    'resolved': 'Решено',
    'cancelled': 'Отменено',
}

DELIVERY_TITLES = {
    'sent': 'Доставлено',
    'failed': 'Не доставлено',
    'pending': 'Ждёт отправки',
}

# Город и парк у тематик спрашиваются разными вопросами (у посылок —
# «город посылки», у офисов — «город офиса», у сотрудничества — свои поля), а
# в выгрузке это одна колонка: по ней фильтруют.
CITY_KEYS = ('city', 'parcel_city', 'office_city', 'coop_city')
PARK_KEYS = ('park', 'coop_park')

COLUMNS = [
    # (заголовок, ключ, ширина)
    ('№', 'id', 7),
    ('Создано', 'created_at', 17),
    ('Тематика', 'queue_title', 20),
    ('Тема', 'topic', 30),
    ('Заголовок', 'subject', 44),
    ('Статус', 'status_title', 13),
    ('Кто завёл', 'created_by_name', 26),
    ('Отдел', 'department_name', 18),
    ('Клиент', 'client_name', 26),
    ('Телефон', 'client_phone', 16),
    ('ИИН', 'iin', 15),
    ('Город', 'city', 14),
    ('Таксопарк', 'park', 18),
    ('Группа', 'group_title', 26),
    ('Доставка', 'delivery_title', 14),
    ('Первый ответ', 'first_reply_at', 17),
    ('До первого ответа, мин', 'reply_minutes', 12),
    ('Закрыто', 'resolved_at', 17),
    ('Кто закрыл', 'resolved_by_name', 24),
    ('Метки', 'flags_title', 22),
    ('Текст обращения', 'body', 70),
]

# Проверка супервайзером до группы (schema.REVIEW_*, задача #297). Колонки
# появляются только когда в периоде есть хоть одно такое обращение: проходят
# её единицы, и четыре пустые колонки в каждой остальной выгрузке — шум.
REVIEW_TITLES = {
    schema.REVIEW_PENDING: 'Ждёт проверки',
    schema.REVIEW_SENT: 'Отправлено в группу',
    schema.REVIEW_RESOLVED: 'Решено супервайзером',
}

REVIEW_COLUMNS = [
    ('Проверка супервайзера', 'review_title', 22),
    ('Кто проверил', 'review_by_name', 24),
    ('Когда проверил', 'review_at', 17),
    ('Итог проверки', 'review_note', 50),
]

# Обращение, которое в группу не уходило: ждёт проверки или решено самим
# супервайзером. У него нет ни группы, ни доставки — и писать их в файл нельзя.
_NEVER_SENT = (schema.REVIEW_PENDING, schema.REVIEW_RESOLVED)

TEXT_COLUMNS = ('client_phone', 'iin')
DATE_KEYS = ('created_at', 'first_reply_at', 'resolved_at', 'review_at')
WRAP_KEYS = ('body', 'review_note')


def columns_for(items):
    """Колонки листа для этого набора обращений: колонки проверки встают перед
    «Метками», если проверку проходило хоть одно."""
    if not any((item or {}).get('review_state') for item in items or ()):
        return COLUMNS
    at = [key for _title, key, _width in COLUMNS].index('flags_title')
    return COLUMNS[:at] + REVIEW_COLUMNS + COLUMNS[at:]


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def _as_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str) and value:
        try:
            return datetime.fromisoformat(value.replace('Z', ''))
        except ValueError:
            return value
    return None


def _answer(answers, keys):
    """Первый непустой ответ из списка ключей. Ответ «да/нет с уточнением»
    хранится словарём — такой в колонку города или парка не годится."""
    for key in keys:
        value = (answers or {}).get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ''


def row_values(item, columns=COLUMNS):
    """Обращение → значения строки листа. Отдельной функцией, чтобы тест
    проверял содержимое, не разбирая готовый xlsx."""
    answers = item.get('answers') or {}
    scenario = scenarios.get(item.get('scenario_key')) or {}
    created = _as_datetime(item.get('created_at'))
    replied = _as_datetime(item.get('first_reply_at'))
    minutes = None
    if isinstance(created, datetime) and isinstance(replied, datetime) and replied >= created:
        minutes = int((replied - created).total_seconds() // 60)
    review = item.get('review_state')
    never_sent = review in _NEVER_SENT
    values = dict(item)
    values.update({
        'topic': scenario.get('title') or item.get('topic_title') or '',
        # «Отправлено» про обращение, которое ждёт супервайзера, — неправда.
        'status_title': ('На проверке' if review == schema.REVIEW_PENDING
                         else STATUS_TITLES.get(item.get('status'), item.get('status') or '')),
        'iin': _answer(answers, ('iin',)),
        'city': _answer(answers, CITY_KEYS),
        'park': _answer(answers, PARK_KEYS),
        # Куда обращение ушло на самом деле: у уведённой темы это не чат
        # тематики (crm_tickets.tg_chat_title — снимок при создании). Не
        # уходившему в группу чат тематики не подставляем: его там не было.
        'group_title': '' if never_sent else (
            item.get('tg_chat_title') or item.get('queue_chat_title') or ''),
        'delivery_title': '' if never_sent else DELIVERY_TITLES.get(
            item.get('delivery_status'), ''),
        'reply_minutes': minutes,
        'flags_title': ', '.join(scenarios.FLAG_LABELS.get(flag, flag)
                                 for flag in (item.get('flags') or [])),
        'review_title': REVIEW_TITLES.get(review, ''),
    })
    return [values.get(key) for _title, key, _width in columns]


def _header(sheet, titles):
    cells = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGN
        cells.append(cell)
    sheet.append(cells)


def _ru(value):
    try:
        return date.fromisoformat(str(value)).strftime('%d.%m.%Y')
    except ValueError:
        return str(value or '')


def build_workbook(items, *, date_from, date_to, generated_by, generated_at=None,
                   truncated=False, text_warning_patch=None):
    """Книга целиком. Возвращает BytesIO, готовый к send_file."""
    generated_at = generated_at or datetime.now()
    book = Workbook(write_only=True)

    context = book.create_sheet('Контекст')
    context.column_dimensions['A'].width = 30
    context.column_dimensions['B'].width = 90
    title = WriteOnlyCell(context, value='Обращения — выгрузка')
    title.font = TITLE_FONT
    context.append([title])
    rows = [
        ('Собрано', generated_at.strftime('%d.%m.%Y %H:%M')),
        ('Кто выгрузил', generated_by or ''),
        ('Период', '%s — %s (по дате создания обращения)' % (_ru(date_from), _ru(date_to))),
        ('Строк', len(items)),
    ]
    if truncated:
        rows.append(('Внимание', 'В выгрузку вошли не все обращения периода — сократите период.'))
    columns = columns_for(items)
    reviewed = columns is not COLUMNS
    rows.append(('Что в листе «Обращения»',
                 'Каждая строка — одно обращение. «Группа» — куда оно ушло на самом деле. '
                 'Жалобы сюда не входят: их выгрузка — в разделе «Жалобы».'
                 + (' «Проверка супервайзера» — у обращений, которые сначала проверяет '
                    'супервайзер: он решает их сам или отправляет в группу; у остальных '
                    'колонка пустая.' if reviewed else '')))
    for label, value in rows:
        head = WriteOnlyCell(context, value=label)
        head.font = LABEL_FONT
        body = WriteOnlyCell(context, value=_clean(value))
        body.alignment = WRAP
        context.append([head, body])

    sheet = book.create_sheet('Обращения')
    for index, (_title, _key, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'B2'
    _header(sheet, [title for title, _key, _width in columns])
    for item in items:
        cells = []
        for (_title, key, _width), value in zip(columns, row_values(item, columns)):
            if key in DATE_KEYS:
                value = _as_datetime(value)
            cell = WriteOnlyCell(sheet, value=_clean(value if value is not None else ''))
            if key in DATE_KEYS and isinstance(value, datetime):
                cell.number_format = DATE_FORMAT
            if key in TEXT_COLUMNS:
                cell.number_format = '@'
            if key in WRAP_KEYS:
                cell.alignment = WRAP
            cells.append(cell)
        sheet.append(cells)
    if items:
        sheet.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(columns)), len(items) + 1)

    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    if text_warning_patch and items:
        keys = [key for _t, key, _w in columns]
        first = get_column_letter(1 + keys.index('client_phone'))
        last = get_column_letter(1 + keys.index('iin'))
        try:
            stream = text_warning_patch(stream, '%s2:%s%d' % (first, last, len(items) + 1),
                                        sheet_path='xl/worksheets/sheet2.xml')
            stream.seek(0)
        except Exception:  # noqa: BLE001 — уголок в Excel не повод терять выгрузку
            stream.seek(0)
    return stream


def filename(date_from, date_to):
    return 'Обращения %s–%s.xlsx' % (_ru(date_from), _ru(date_to))
