# -*- coding: utf-8 -*-
"""Детализация звонков раздела «Удаленный КЦ» в xlsx (запрос владельца 15.10.2026: «выгрузить
детализацию за выбранный период — что-то наподобие журнала звонков»).

Журнал раздела — это водители, строка на человека; детализация — звонки: каждая попытка
оператора за выбранные дни строкой, как в истории карточки водителя. Три листа:

    «Звонки»     когда, кто, с какой линии, кому, чем кончилось у Binotel, сколько говорили,
                 какой итог поставил оператор и что написал — по времени звонка;
    «Операторы»  итоги каждого за период по тем же правилам, что сводка «Операторы» раздела
                 (там — один день), и строка «Итого»;
    «Контекст»   отдел, период, кто и когда собрал, как читать колонки.

Исход и его подписи — те же, что в журнале и карточке (service.attempt_outcome; RESULT_LABELS —
двойник RESULT_LABEL в DialListJournal.jsx, приписки — attemptHints.js): файл и экран не должны
называть один звонок по-разному. Итоги на листе «Операторы» считаются по тем же строкам, что
на листе «Звонки», — фильтр по оператору в Excel даёт ровно их. Дозвон засчитывается от
DIAL_MIN_SEC секунд разговора (владелец, 29.10.2026: «с засчитанного кол-во дозвонов убрать звонки
меньше 10 секунд»): короткий разговор остаётся «Дозвонились» по Binotel, как в журнале, но в
«Подробностях» помечен, а в «Дозвонились» листа «Операторы» не входит.

Номер водителя — только маской (service.mask_phone), как в журнале: полный номер из раздела не
выходит ни в одну ручку, а файл — тоже ручка. ИИН и ID водителя — полностью (владелец, 29.10.2026:
«в детализацию и в отчеты добавить id и иин водителя»): по ним водителя находят в Sapar, в Fleet и в
выгрузках iTaxi. ID — Yandex Driver Id из Sapar (driver_ids.py); аккаунтов несколько — все, через
запятую; Sapar водителя не знает — пусто. Звонки на номера «Реестра тестовых номеров» не попадают —
как и в сводку «Операторы».

Отчёт главы по своей группе (владелец, 15.10.2026: «глава Регионов мог выгружать отчеты только по
своей группе»; service.DIAL_LIST_GROUP_REPORT_HEAD_CODES) — тот же файл, но только звонки тех, кто в
день звонка состоял в её группах (GROUP_FILTER_SQL), и строка «Группа» в «Контексте».

Соглашения — как у выгрузок «Касаний» и «Учёта воды» (cdr/report.py, water/report.py): дата и
время — настоящей датой, разговор — длительностью (читается «3:25» и складывается), линия —
текстом, зелёный уголок «Число сохранено как текст» гасит функция монолита (приходит
аргументом). Книга пишется потоково (write_only): ширины и закрепление — до первой строки,
стили — модульные константы; строки идут из базы порциями одного SELECT (один снимок, без
копии месяца в памяти).
"""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from test_numbers import keys as test_keys

from . import driver_ids, service, signing

# Сколько дней за раз. Месяц — естественная единица отчёта (зарплата, итоги месяца), и в 31 день
# помещается любой календарный месяц. Двойник — EXPORT_MAX_DAYS в callsExport.js.
MAX_DAYS = 31
FETCH_CHUNK = 2000

# Подписи исхода — дословно RESULT_LABEL журнала (DialListJournal.jsx); '' — исход ещё едет от АТС.
RESULT_LABELS = {
    "answered": "Дозвонились",
    "busy": "Занято",
    "no_answer": "Не поднял трубку",
    "other": "Не состоялся",
    "failed": "Ошибка АТС",
    "cancelled": "Отменён до ответа",
    "": "Ждём исход от АТС",
}

# Приписки автосброса — те же слова, что в истории карточки (attemptHints.js).
RING_TIMEOUT_NOTE = "сброшен автоматически: водитель не взял трубку"
ANNOUNCEMENT_NOTE = "сброшен автоматически: сообщение сети «абонент недоступен»"
CANCELLED_NOTE = "попытка не засчитана"
# Дозвон засчитывается от этой длительности разговора (владелец, 29.10.2026: «с засчитанного кол-во
# дозвонов убрать звонки меньше 10 секунд»; service.ATTEMPT_DIAL_COUNTED_SQL — двойник в SQL).
DIAL_MIN_SEC = signing.SUCCESS_MIN_BILLSEC
SHORT_TALK_NOTE = f"разговор короче {DIAL_MIN_SEC} с — в дозвоны не засчитан"

# Почему звонок не начался (service.call_error_kind). Сам текст Binotel в файл не идёт: он
# английский, а у попыток до 13.10.2026 в нём бывает номер водителя.
FAILURE_NOTES = {
    service.CALL_ERROR_LINE: "линия оператора не на связи",
    service.CALL_ERROR_NETWORK: "нет связи с АТС",
    service.CALL_ERROR_PBX_BUSY: "АТС не приняла звонок",
    service.CALL_ERROR_RATE: "слишком частые запросы к АТС",
    service.CALL_ERROR_PBX: "АТС отказала",
}
# Код ответа АТС полезен только там, где разбираться руководителю (отказ АТС, сбой 105).
FAILURE_CODE_KINDS = (service.CALL_ERROR_PBX, service.CALL_ERROR_PBX_BUSY)

SHEET_CALLS = "Звонки"
SHEET_OPERATORS = "Операторы"
SHEET_CONTEXT = "Контекст"

# key, заголовок, ширина. Порядок — как читают: когда и кто, кому, чем кончилось, что сказал
# оператор; служебное — в конце.
COLUMNS = (
    ("at", "Дата и время", 19),
    ("operator", "Оператор", 28),
    ("line", "Линия", 9),
    ("driver", "Водитель", 30),
    ("driver_ids", "ID водителя", 38),
    ("iin", "ИИН", 15),
    ("phone", "Номер", 17),
    ("result", "Результат", 19),
    ("talk_sec", "Разговор", 10),
    ("outcome", "Итог оператора", 30),
    ("comment", "Комментарий оператора", 44),
    ("success", "Успешка", 10),
    ("note", "Подробности", 46),
)
# Цифрами, но текстом: линия так хранится и так её показывает раздел, у ИИН числом пропал бы ведущий
# ноль, ID — код, а не число.
TEXT_COLUMNS = ("line", "driver_ids", "iin")

OPERATOR_COLUMNS = (
    ("Оператор", 30),
    ("Звонков", 10),
    (f"Дозвонились (от {DIAL_MIN_SEC} с)", 13),
    ("Дозвон", 9),
    ("Разговор", 12),
    ("Средний разговор", 12),
    ("Успешки", 10),
    ("Отменено до ответа", 13),
    ("Ошибки АТС", 11),
)

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
SECTION_FONT = Font(bold=True, size=11)
TOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
TOTAL_FONT = Font(bold=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
WRAP = Alignment(wrap_text=True)

DATETIME_FORMAT = "DD.MM.YYYY HH:MM:SS"
# Разговор — долей суток с форматом длительности: Excel показывает «3:25» и честно складывает.
TALK_FORMAT = "[m]:ss"
TALK_TOTAL_FORMAT = "[h]:mm:ss"
PERCENT_FORMAT = "0%"

# Звонки отдела за дни [date_from; date_to] по Алматы. Отдел — базы, которую обзванивали, а не
# человека (как сводка «Операторы»): на линии отдела сидят и сотрудники других отделов.
_CALLS_SELECT = """
    SELECT t.requested_at, t.operator_id, COALESCE(u.name, ''), t.internal_number, t.state,
           t.disposition, t.billsec, t.cancelled, t.auto_hangup, t.talk_detected, t.ring_timeout_at,
           t.api_error, l.full_name, l.phone_norm, o.name, os.name, COALESCE(t.operator_comment, ''),
           COALESCE(l.success_attempt_id = t.id, FALSE), l.iin, COALESCE(l.yandex_driver_ids, '')
    FROM dial_list_attempts t
    JOIN dial_list_assignments a ON a.id = t.assignment_id
    JOIN dial_list_leads l ON l.id = a.lead_id
    LEFT JOIN users u ON u.id = t.operator_id
    LEFT JOIN dial_list_outcomes o ON o.id = t.outcome_id
    LEFT JOIN dial_list_outcome_subtypes os ON os.id = t.outcome_subtype_id
    WHERE l.department_id = %(department_id)s
      AND t.requested_at >= (%(date_from)s::date)::timestamp AT TIME ZONE 'Asia/Almaty'
      AND t.requested_at < (%(date_to)s::date + 1)::timestamp AT TIME ZONE 'Asia/Almaty'
      AND """ + test_keys.DIAL_LIST_LEAD_NOT_TEST_SQL
# Отчёт главы по своей группе (service.DIAL_LIST_GROUP_REPORT_HEAD_CODES): звонок — того, кто В ДЕНЬ
# ЗВОНКА состоял в группе. Членство датированное: перешёл в группу 14-го — его звонки 13-го в отчёт
# не идут, ушёл из неё — уходят и звонки после ухода.
GROUP_FILTER_SQL = """
      AND EXISTS (SELECT 1 FROM group_operator_memberships gm
                  WHERE gm.operator_id = t.operator_id AND gm.group_id = ANY(%(group_ids)s)
                    AND gm.start_date <= (t.requested_at AT TIME ZONE 'Asia/Almaty')::date
                    AND (gm.end_date IS NULL OR gm.end_date >= (t.requested_at AT TIME ZONE 'Asia/Almaty')::date))"""
_CALLS_ORDER = """
    ORDER BY t.requested_at, t.id
"""


def calls_sql(group_ids=None):
    """Запрос звонков: весь отдел (group_ids None) или только люди групп."""
    return _CALLS_SELECT + (GROUP_FILTER_SQL if group_ids is not None else "") + _CALLS_ORDER


CALLS_SQL = calls_sql()


# ── период ──────────────────────────────────────────────────────────────────

def parse_days(date_from, date_to, today=None, max_days=MAX_DAYS):
    """Границы периода (date, date) из 'YYYY-MM-DD'. ValueError — словами для руководителя.

    Одна граница — однодневный период; перепутанные местами — меняем; хвост в будущем отрезаем
    (звонков там нет, а «по 31.10» в файле читалось бы как «весь октябрь проверен»). today —
    сегодня по Алматы, если не передан."""
    today = today or service._today()

    def _day(value):
        text = str(value or "").strip()
        if not text:
            return None
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            raise ValueError("Даты периода — в виде ГГГГ-ММ-ДД")

    start, end = _day(date_from), _day(date_to)
    if start is None and end is None:
        raise ValueError("Выберите период выгрузки")
    start, end = start or end, end or start
    if end < start:
        start, end = end, start
    if start > today:
        raise ValueError("Период ещё не наступил — звонков в нём нет")
    end = min(end, today)
    if (end - start).days + 1 > max_days:
        # «не больше 31 дня», «не больше 7 дней»: после «больше» — родительный падеж.
        word = "дня" if max_days % 10 == 1 and max_days % 100 != 11 else "дней"
        raise ValueError(f"Не больше {max_days} {word} за раз — выберите период короче")
    return start, end


def period_label(day_from, day_to):
    if day_from == day_to:
        return day_from.strftime("%d.%m.%Y")
    return f"{day_from:%d.%m.%Y} — {day_to:%d.%m.%Y}"


def report_filename(day_from, day_to):
    """«Детализация звонков 01.10.2026 — 15.10.2026.xlsx». Имя файла у скачанного ставит сайт
    (Content-Disposition через CORS ему не виден) — то же правило в callsExport.js."""
    return f"Детализация звонков {period_label(day_from, day_to)}.xlsx"


# ── строка звонка ───────────────────────────────────────────────────────────

def failure_note(api_error):
    """Почему звонок не начался — словами, без текста Binotel."""
    text = str(api_error or "")
    if text == service.BINOTEL_NOT_CONFIGURED:
        return text
    kind = service.call_error_kind(None, text)
    code = service.binotel_error_code(text)
    note = FAILURE_NOTES.get(kind, FAILURE_NOTES[service.CALL_ERROR_PBX])
    return f"{note} (код {code})" if code is not None and kind in FAILURE_CODE_KINDS else note


def counted_dial(result, billsec):
    """Засчитанный дозвон: разговор не короче DIAL_MIN_SEC (двойник service.ATTEMPT_DIAL_COUNTED_SQL)."""
    return result == "answered" and int(billsec or 0) >= DIAL_MIN_SEC


def call_note(result, auto_hangup="", ring_timeout_at=None, api_error="", billsec=None):
    """«Подробности» звонка — то, что карточка водителя пишет под попыткой; '' — нечего добавить.
    Приписку автосброса получает только недозвон: у разговора и у отмены она была бы неправдой
    (как autoHangupHint в attemptHints.js); сообщение сети сильнее гудков. Разговор короче
    DIAL_MIN_SEC (billsec передан) — пометка, что в дозвоны он не засчитан."""
    if result == "failed":
        return failure_note(api_error)
    if result == "cancelled":
        return CANCELLED_NOTE
    if result == "answered":
        return SHORT_TALK_NOTE if billsec is not None and not counted_dial(result, billsec) else ""
    if result == "":
        return ""
    if auto_hangup == service.AUTO_HANGUP_ANNOUNCEMENT:
        return ANNOUNCEMENT_NOTE
    if ring_timeout_at is not None or auto_hangup == service.AUTO_HANGUP_RING_TIMEOUT:
        return RING_TIMEOUT_NOTE
    return ""


def _local(value):
    """Момент из базы → время Алматы без пояса (Excel поясов не знает)."""
    if not isinstance(value, datetime):
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(service.PERIOD_TZ).replace(tzinfo=None)


def call_from_row(r):
    """Строка CALLS_SQL → звонок словарём (одно место, где колонки запроса знают по номерам)."""
    (requested_at, operator_id, operator_name, internal_number, state, disposition, billsec,
     cancelled, auto_hangup, talk_detected, ring_timeout_at, api_error, full_name, phone_norm,
     outcome_name, subtype_name, comment, is_success, iin, yandex_ids) = r
    auto_hangup = auto_hangup or ""
    result = service.attempt_outcome(state or "", disposition or "", billsec, auto_hangup, talk_detected,
                                     cancelled=bool(cancelled))
    talk_sec = int(billsec or 0) if result == "answered" else 0
    return {
        "at": _local(requested_at),
        "operator_id": operator_id,
        "operator": operator_name or "",
        "line": internal_number or "",
        "driver": full_name or "",
        "driver_ids": yandex_ids or "",
        "iin": iin or "",
        "phone": service.mask_phone(phone_norm),
        "result": result,
        # Разговор — только у дозвона: у сообщения сети Binotel пишет длительность автоответчика.
        "talk_sec": talk_sec,
        # В «Дозвонились» листа «Операторы» — только разговор не короче DIAL_MIN_SEC.
        "counted": counted_dial(result, talk_sec),
        # Итог с типом — одной строкой, как бейдж в журнале: «Отказ · Дорого».
        "outcome": " · ".join(part for part in (outcome_name, subtype_name) if part),
        "comment": comment or "",
        "success": bool(is_success),
        "note": call_note(result, auto_hangup, ring_timeout_at, api_error, billsec=talk_sec),
    }


def iter_calls(cursor, department_id, day_from, day_to, group_ids=None, chunk=FETCH_CHUNK):
    """Звонки периода по времени — порциями ОДНОГО запроса (один снимок базы: страницы
    LIMIT/OFFSET потеряли бы или задвоили звонки, которые закрываются прямо сейчас).
    group_ids — только люди этих групп (пустой список — ничего: отказ, а не «весь отдел»)."""
    params = {"department_id": int(department_id), "date_from": day_from.isoformat(),
              "date_to": day_to.isoformat()}
    if group_ids is not None:
        params["group_ids"] = [int(g) for g in group_ids]
    cursor.execute(calls_sql(group_ids), params)
    while True:
        rows = cursor.fetchmany(int(chunk))
        if not rows:
            return
        for row in rows:
            yield call_from_row(row)


def department_name(cursor, department_id):
    cursor.execute("SELECT name FROM departments WHERE id = %s", (int(department_id),))
    row = cursor.fetchone()
    return (row[0] or "") if row else ""


# ── книга ───────────────────────────────────────────────────────────────────

def _clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _text(sheet, value, *, keep_format=False):
    """Ячейка, которая ОСТАЁТСЯ строкой: комментарий оператора с «=» в начале иначе стал бы
    формулой (тот же приём, что water/report.py::_text). Формат «текст» — только непустой."""
    text = _clean(value)
    if text in (None, ""):
        return None
    cell = WriteOnlyCell(sheet, value=text)
    cell.data_type = "s"
    if keep_format:
        cell.number_format = "@"
    return cell


def _duration(sheet, seconds, number_format=TALK_FORMAT, style=None):
    """Секунды → длительность Excel (доля суток). Ноль — пустая клетка: «0:00» читалось бы как
    «разговор был и длился ноль»."""
    seconds = int(seconds or 0)
    if seconds <= 0 and style is None:
        return None
    cell = WriteOnlyCell(sheet, value=seconds / 86400.0 if seconds > 0 else None)
    cell.number_format = number_format
    if style:
        cell.font, cell.fill = style
    return cell


def _header(sheet, titles, widths):
    """Шапка, ширины и закрепление — ДО первой строки (ловушка потокового режима, cdr/report.py)."""
    sheet.freeze_panes = "A2"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    row = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGN
        row.append(cell)
    sheet.append(row)


def _ids_cell(sheet, value):
    """ID водителя. Аккаунтов несколько — перенос у запятой: каждый ID своей строкой в ячейке, второй
    не прячется за краем колонки."""
    cell = _text(sheet, value, keep_format=True)
    if cell is not None and driver_ids.SEPARATOR in value:
        cell.alignment = WRAP
    return cell


def _call_cells(sheet, call):
    at = WriteOnlyCell(sheet, value=call["at"])
    at.number_format = DATETIME_FORMAT
    return [
        at,
        _text(sheet, call["operator"]),
        _text(sheet, call["line"], keep_format=True),
        _text(sheet, call["driver"]),
        _ids_cell(sheet, call["driver_ids"]),
        _text(sheet, call["iin"], keep_format=True),
        _text(sheet, call["phone"]),
        _text(sheet, RESULT_LABELS.get(call["result"], call["result"])),
        _duration(sheet, call["talk_sec"]),
        _text(sheet, call["outcome"]),
        _text(sheet, call["comment"]),
        _text(sheet, "да" if call["success"] else ""),
        _text(sheet, call["note"]),
    ]


# answered — засчитанные дозвоны (counted), counted_talk_sec — их разговор: из него средний разговор.
# talk_sec — весь разговор, и короткий тоже: это время, которое оператор говорил.
TOTAL_KEYS = ("calls", "answered", "talk_sec", "counted_talk_sec", "successes", "cancelled", "failed")


def _empty_totals(name):
    return {"name": name, **dict.fromkeys(TOTAL_KEYS, 0)}


def _count(totals, call):
    totals["calls"] += 1
    totals["answered"] += int(call["counted"])
    totals["talk_sec"] += call["talk_sec"]
    totals["counted_talk_sec"] += call["talk_sec"] if call["counted"] else 0
    totals["successes"] += int(call["success"])
    totals["cancelled"] += int(call["result"] == "cancelled")
    totals["failed"] += int(call["result"] == "failed")


def _operator_cells(sheet, totals, total_row=False):
    style = (TOTAL_FONT, TOTAL_FILL) if total_row else None

    def plain(value, number_format=None):
        cell = WriteOnlyCell(sheet, value=value)
        if number_format:
            cell.number_format = number_format
        if style:
            cell.font, cell.fill = style
        return cell

    name = plain(_clean(totals["name"]) or "—")
    name.data_type = "s"
    calls, answered, talk = totals["calls"], totals["answered"], totals["talk_sec"]
    counted_talk = totals["counted_talk_sec"]
    return [
        name,
        plain(calls),
        plain(answered),
        # Доля — от всех звонков, как «Дозвонились · N %» на вкладке «Операторы».
        plain(answered / calls if calls else None, PERCENT_FORMAT),
        _duration(sheet, talk, TALK_TOTAL_FORMAT, style),
        # Средний — по засчитанным дозвонам: короткий разговор не в счёт и здесь.
        _duration(sheet, round(counted_talk / answered) if answered else 0, TALK_FORMAT, style),
        plain(totals["successes"]),
        plain(totals["cancelled"]),
        plain(totals["failed"]),
    ]


def _fill_context(sheet, *, department, day_from, day_to, generated_at, generated_by, written, groups=None):
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 96
    lines = [
        ("Удаленный КЦ — детализация звонков", None),
        ("", None),
        ("Отдел", department or "—"),
        # Отчёт по своей группе — видно, что в файле не весь отдел.
        *([("Группа", ", ".join(groups) + " — звонки тех, кто в день звонка состоял в группе")]
          if groups is not None else []),
        ("Период", period_label(day_from, day_to)),
        ("Собрано", generated_at.strftime("%d.%m.%Y %H:%M") + " (Алматы)"),
        ("Собрал", generated_by or "—"),
        ("Звонков", written),
        ("", None),
        ("Как читать", None),
        ("Дата и время", "Когда оператор нажал «Позвонить». Время — Алматы."),
        ("Результат", "Чем кончился звонок по данным Binotel. «Дозвонились» — водитель ответил и был "
                      "разговор; сообщение сети «абонент недоступен» и ответ без единой реплики — недозвон. "
                      f"Разговор короче {DIAL_MIN_SEC} с в дозвоны не засчитывается — так и помечен в «Подробностях»."),
        ("Отменён до ответа", "Оператор сам завершил звонок до ответа водителя — попытка не засчитана."),
        ("Ошибка АТС", "Звонок не начался; почему — в «Подробностях»."),
        ("Успешка", "Водитель подписал документы после этого разговора — успешка засчитана оператору."),
        ("ID водителя", "Yandex Driver Id из Sapar — тот же, что в выгрузке услуг iTaxi. Аккаунтов у водителя "
                        "несколько — все, через запятую. Пусто — Sapar водителя не знает."),
        ("ИИН", "Полностью, из базы водителей раздела."),
        ("Номер", "Только последние четыре цифры, как в журнале раздела."),
        ("Операторы", f"Итоги по тем же звонкам, что на листе «Звонки». «Дозвонились» — разговоры от {DIAL_MIN_SEC} с, "
                      "«Дозвон» — их доля от всех звонков, как на вкладке «Операторы»; «Средний разговор» — по "
                      "ним же, «Разговор» — весь, и короткий тоже."),
        ("Тестовые номера", "Звонки на номера из «Реестра тестовых номеров» в файл не входят."),
    ]
    for index, (label, value) in enumerate(lines, start=1):
        left = WriteOnlyCell(sheet, value=label)
        if value is None:
            if label:
                left.font = TITLE_FONT if index == 1 else SECTION_FONT
            sheet.append([left])
            continue
        right = (_text(sheet, value) if isinstance(value, str) else None) or WriteOnlyCell(sheet, value=value)
        right.alignment = WRAP_TOP
        sheet.append([left, right])


def build_workbook(calls, *, department, day_from, day_to, generated_by="", generated_at=None,
                   text_warning_patch=None, groups=None):
    """calls — итерируемое звонков (call_from_row; можно генератор). → (BytesIO, звонков).
    groups — названия групп, если отчёт по группе (строка «Группа» в «Контексте»)."""
    generated_at = generated_at or datetime.now(service.PERIOD_TZ)
    workbook = Workbook(write_only=True)
    # Порядок создания = порядок листов: «Звонки» — sheet1.xml (по нему гасится уголок).
    calls_sheet = workbook.create_sheet(SHEET_CALLS)
    operators_sheet = workbook.create_sheet(SHEET_OPERATORS)
    context_sheet = workbook.create_sheet(SHEET_CONTEXT)

    _header(calls_sheet, [title for _key, title, _width in COLUMNS], [width for _k, _t, width in COLUMNS])
    totals = {}
    written = 0
    for call in calls:
        calls_sheet.append(_call_cells(calls_sheet, call))
        written += 1
        person = totals.get(call["operator_id"])
        if person is None:
            person = totals[call["operator_id"]] = _empty_totals(call["operator"])
        _count(person, call)
    if written:
        calls_sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{written + 1}"

    _header(operators_sheet, [title for title, _width in OPERATOR_COLUMNS],
            [width for _title, width in OPERATOR_COLUMNS])
    overall = _empty_totals("Итого")
    for person in sorted(totals.values(), key=lambda t: (-t["calls"], t["name"].lower())):
        operators_sheet.append(_operator_cells(operators_sheet, person))
        for key in TOTAL_KEYS:
            overall[key] += person[key]
    if len(totals) > 1:
        operators_sheet.append(_operator_cells(operators_sheet, overall, total_row=True))

    _fill_context(context_sheet, department=department, day_from=day_from, day_to=day_to,
                  generated_at=generated_at, generated_by=generated_by, written=written, groups=groups)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    if callable(text_warning_patch) and written:
        sqref = " ".join("{0}2:{0}{1}".format(get_column_letter(_column_index(key)), written + 1)
                         for key in TEXT_COLUMNS)
        try:
            stream = text_warning_patch(stream, sqref, sheet_path="xl/worksheets/sheet1.xml")
        except Exception:  # noqa: BLE001 — значок в углу ячейки не повод не отдать файл
            stream.seek(0)
    return stream, written


def _column_index(key):
    return next(index for index, (name, _title, _width) in enumerate(COLUMNS, start=1) if name == key)
