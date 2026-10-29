# -*- coding: utf-8 -*-
"""Выгрузка рассылки автоинформатора в xlsx: водители, все звонки, контекст.

Номер — маской, как в журнале раздела: полный номер из раздела не выходит ни в одну ручку,
а файл — тоже ручка. ИИН и ID водителя — полностью, из базы «Удаленного КЦ» по номеру (владелец,
29.10.2026: «в детализацию и в отчеты добавить id и иин водителя»), как в «Детализации»
(dial_list/calls_export.py). Время — по Алматы настоящей датой, секунды — числом (складываются)."""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from dial_list.driver_ids import SEPARATOR as IDS_SEPARATOR

from . import rules

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
WRAP = Alignment(wrap_text=True, vertical='top')
IDS_WRAP = Alignment(wrap_text=True)
DATETIME_FORMAT = 'DD.MM.YYYY HH:MM:SS'
# ID водителя и ИИН — третья и четвёртая колонки обоих листов: ИИН цифрами, но текстом (числом пропал бы
# ведущий ноль), зелёный уголок «Число сохранено как текст» у них гасит функция монолита.
TEXT_COLUMNS = ('C', 'D')

STATE_LABELS = {
    rules.STATE_DONE: 'Дозвонились',
    rules.STATE_GIVEN_UP: 'Попытки кончились',
    rules.STATE_PENDING: 'В очереди',
    rules.STATE_CALLING: 'Звоним',
    rules.STATE_STOPPED: 'Рассылку завершили раньше',
}
HANGUP_LABELS = {rules.HANGUP_DRIVER: 'водитель', rules.HANGUP_ROBOT: 'робот (запись кончилась)',
                 rules.HANGUP_NETWORK: 'сеть', '': ''}


def _local(value):
    if not value:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(rules.ALMATY).replace(tzinfo=None)


def _seconds(ms):
    return round(int(ms) / 1000, 1) if ms is not None else None


def _text(value):
    return ILLEGAL_CHARACTERS_RE.sub('', str(value if value is not None else ''))


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = 'A2'
    row = []
    for header in headers:
        cell = WriteOnlyCell(ws, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        row.append(cell)
    ws.append(row)
    return ws


def _datetime_cell(ws, value):
    cell = WriteOnlyCell(ws, value=_local(value))
    cell.number_format = DATETIME_FORMAT
    return cell


def _identity_cells(ws, row):
    """ID водителя и ИИН — строкой, с форматом «текст». Аккаунтов несколько — каждый ID своей строкой
    в ячейке (перенос у запятой)."""
    cells = []
    for value in (row.get("driver_ids") or '', row.get("iin") or ''):
        cell = WriteOnlyCell(ws, value=_text(value) or None)
        if value:
            cell.data_type, cell.number_format = 's', '@'
        if IDS_SEPARATOR in value:
            cell.alignment = IDS_WRAP
        cells.append(cell)
    return cells


def build(detail, targets, calls, text_warning_patch=None):
    """text_warning_patch — функция монолита (_excel_suppress_number_as_text_warning), гасит зелёный
    уголок у ИИН; нет её — файл тот же, только с уголком."""
    wb = Workbook(write_only=True)
    parts = detail.get("parts") or []

    ws = _sheet(wb, 'Водители', ['Строка файла', 'ФИО', 'ID водителя', 'ИИН', 'Номер', 'Итог', 'Попыток',
                                  'Последний звонок', 'Слушал, с', 'Где положил трубку', 'Подписал (Sapar)',
                                  'Подпись после звонка'],
                [12, 32, 38, 15, 18, 26, 10, 20, 11, 30, 20, 20])
    for t in targets:
        signed_at = t.get("signed_at")
        ws.append([t["row_number"], _text(t["full_name"]), *_identity_cells(ws, t), t["phone"],
                   STATE_LABELS.get(t["state"], t["state"]), t["attempts"], _datetime_cell(ws, t["last_call_at"]),
                   _seconds(t["listened_ms"]) if t["state"] == rules.STATE_DONE else None,
                   rules.stage_label(t["stage"]) if t["state"] == rules.STATE_DONE else '',
                   _datetime_cell(ws, signed_at) if signed_at else None,
                   ('да' if t.get("signed_after_call") else 'подписал до звонка') if signed_at else ''])

    ws = _sheet(wb, 'Звонки', ['Время звонка', 'ФИО', 'ID водителя', 'ИИН', 'Номер', 'Попытка', 'Линия', 'Итог',
                                'Звонило до ответа, с', 'Слушал, с', 'Где положил трубку', 'Кто положил',
                                'Нажатия'],
                [20, 32, 38, 15, 18, 9, 8, 18, 14, 11, 30, 24, 10])
    for c in calls:
        answered = c["result"] == rules.RESULT_ANSWERED
        ws.append([_datetime_cell(ws, c["issued_at"]), _text(c["full_name"]), *_identity_cells(ws, c), c["phone"],
                   c["attempt_no"], c["line"], rules.RESULT_LABELS.get(c["result"], c["result"]), _seconds(c["ring_ms"]),
                   _seconds(c["listened_ms"]) if answered else None,
                   rules.stage_label(c["stage"]) if answered else '',
                   HANGUP_LABELS.get(c["hangup_by"], c["hangup_by"]), c["dtmf"]])

    analytics = detail.get("analytics") or {}
    stages = analytics.get("stages") or {}
    ws = _sheet(wb, 'Контекст', ['Что', 'Значение'], [34, 80])
    part_text = '; '.join(
        f"{rules.LANG_LABELS.get(p['lang'], p['lang'])}: {p['start_ms'] / 1000:.1f}–{p['end_ms'] / 1000:.1f} с"
        for p in parts)
    rows = [
        ('Рассылка', detail.get("name")),
        ('Создана', detail.get("created_at")),
        ('Запись', f"{(detail.get('audio_ms') or 0) / 1000:.1f} с; {part_text}"),
        ('Окно звонков (Алматы)', f"{detail.get('window_start')}–{detail.get('window_end')}"),
        ('Темп', f"не больше {detail.get('max_per_minute')} звонков в минуту"),
        ('Попыток на водителя', detail.get("max_attempts")),
        ('Повтор недозвона через', f"{detail.get('retry_after_min')} мин"),
        ('Водителей в списке', detail.get("targets_total")),
        ('Дозвонились', analytics.get("answered")),
        ('Положили на казахской части', stages.get('kz')),
        ('Положили на русской части', stages.get('ru')),
        ('Дослушали до конца', stages.get(rules.STAGE_FULL)),
    ]
    signing = analytics.get("signing")
    if signing:
        rows += [
            ('Подписали после звонка робота', signing.get("after_call")),
            ('…из них без разговора с оператором', signing.get("after_call_self")),
            ('Подписали ещё до звонка', signing.get("before_call")),
            ('Водителей найдено в базе «Обзвона»', signing.get("matched")),
            ('Подписание проверено (Sapar)',
             _local(datetime.fromisoformat(signing["checked_at"])) if signing.get("checked_at") else '—'),
        ]
    rows += [
        ('Выгружено', datetime.now(rules.ALMATY).strftime('%d.%m.%Y %H:%M')),
        ('Как читать', '«Слушал» — секунды от ответа водителя до отбоя; запись играет с момента ответа. '
                       '«Где положил трубку» — часть записи, которая звучала в момент отбоя. '
                       '«Линия занята» — Binotel не дал линию (у линии два звонка одновременно); '
                       '«Сеть не пропустила» — вызов не дошёл до водителя (сеть или оператор связи); '
                       'попыткой водителя ни то, ни другое не считается. Номер — маской. ИИН и ID водителя — '
                       'из базы «Удаленного КЦ» по номеру; ID — Yandex Driver Id из Sapar, аккаунтов несколько — '
                       'через запятую; пусто — водителя нет в базе или Sapar его не знает.'),
    ]
    for key, value in rows:
        ws.append([key, _text(value)])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    if callable(text_warning_patch):
        # «Водители» — sheet1.xml, «Звонки» — sheet2.xml (порядок создания листов).
        for sheet_path, written in (('xl/worksheets/sheet1.xml', len(targets)),
                                    ('xl/worksheets/sheet2.xml', len(calls))):
            if not written:
                continue
            sqref = ' '.join(f'{column}2:{column}{written + 1}' for column in TEXT_COLUMNS)
            try:
                stream = text_warning_patch(stream, sqref, sheet_path=sheet_path)
            except Exception:  # noqa: BLE001 — значок в углу ячейки не повод не отдать файл
                stream.seek(0)
    return stream
