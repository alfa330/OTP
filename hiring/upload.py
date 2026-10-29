# -*- coding: utf-8 -*-
"""Разбор файла базы контактов (ТЗ 5.1): xlsx / csv с полями ФИО, телефон, город, комментарий.

Колонки узнаются по заголовкам (любой регистр, русские и английские слова); без заголовков
берётся порядок ТЗ: ФИО, телефон, город, комментарий. Строка без телефона или с неразборным
номером — ошибка с номером строки, она попадает в отчёт партии. Дубли внутри файла — по
телефону, остаётся первая строка.
"""

import csv
import io

from . import rules

MAX_ROWS = 50000

_HEADERS = {
    'name': ('фио', 'имя', 'ф.и.о', 'ф. и. о.', 'name', 'full name', 'fullname', 'кандидат', 'контакт'),
    'phone': ('телефон', 'номер', 'тел', 'тел.', 'phone', 'mobile', 'моб', 'номер телефона'),
    'city': ('город', 'city', 'регион', 'населённый пункт', 'населенный пункт'),
    'comment': ('комментарий', 'коммент', 'comment', 'примечание', 'заметка', 'note', 'notes'),
}


class UploadError(ValueError):
    pass


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _detect_columns(header):
    """{поле: индекс} по строке заголовков; None — это не заголовки."""
    found = {}
    for index, raw in enumerate(header):
        text = _cell(raw).lower().strip(' :.')
        if not text:
            continue
        for field, names in _HEADERS.items():
            if field in found:
                continue
            if text in names or any(text.startswith(name) for name in names if len(name) > 3):
                found[field] = index
                break
    return found if 'phone' in found else None


def _rows_from_xlsx(raw):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()


def _rows_from_csv(raw):
    text = None
    for encoding in ('utf-8-sig', 'utf-8', 'cp1251'):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise UploadError('Не удалось прочитать файл: неизвестная кодировка')
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ';' if sample.count(';') >= sample.count(',') else ','
    for row in csv.reader(io.StringIO(text), dialect):
        yield row


def parse(raw, file_name):
    """(контакты, ошибки, всего строк). Контакт — {'row', 'full_name', 'phone', 'city', 'comment'};
    ошибка — {'row', 'value', 'error'}."""
    name = str(file_name or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        rows = _rows_from_xlsx(raw)
    elif name.endswith(('.csv', '.txt')):
        rows = _rows_from_csv(raw)
    else:
        raise UploadError('Нужен файл Excel (.xlsx) или CSV')

    contacts, errors, seen = [], [], set()
    columns = None
    total = 0
    for number, row in enumerate(rows, start=1):
        if number > MAX_ROWS + 1:
            raise UploadError('В файле больше %d строк — разделите его' % MAX_ROWS)
        cells = [_cell(value) for value in (row or [])]
        if not any(cells):
            continue
        if columns is None:
            detected = _detect_columns(cells)
            if detected is not None:
                columns = detected
                continue
            columns = {'name': 0, 'phone': 1, 'city': 2, 'comment': 3}
        total += 1

        def pick(field):
            index = columns.get(field)
            return cells[index] if index is not None and index < len(cells) else ''

        phone_raw = pick('phone')
        phone = rules.normalize_phone(phone_raw)
        if not phone:
            errors.append({'row': number, 'value': phone_raw[:40], 'error': 'Пустой или неразборный номер'})
            continue
        if phone in seen:
            errors.append({'row': number, 'value': phone, 'error': 'Повтор номера в файле'})
            continue
        seen.add(phone)
        contacts.append({
            'row': number,
            'full_name': pick('name')[:200],
            'phone': phone,
            'city': pick('city')[:120] or None,
            'comment': pick('comment')[:2000] or None,
        })
    return contacts, errors, total
