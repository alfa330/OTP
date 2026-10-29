# -*- coding: utf-8 -*-
"""Итоговый дашборд воронки и отчёты руководителя (ТЗ 8.1, 8.2), выгрузка в Excel.

Логика расчёта — ровно по ТЗ 8.1:

* когорта — карточки, взятые в работу в отчётном месяце (первый вход в «В работе»);
  их дальнейшее движение считается на дату просмотра;
* карточка достигла этапа, если ХОТЬ РАЗ была в этой колонке (hiring_card_stages),
  а не только находится в ней сейчас; «Дозвон» — хотя бы один звонок с результатом «дозвон»;
* конверсия этапа = достигли ÷ достигли предыдущего; сквозная = достигли ÷ взято в работу;
* кандидаты — в головах; план месяца — в FTE: каждый «Отработал 7 дней» закрывает 1 FTE.

Фильтры (месяц, рекрутеры, источники, «дата этапа») действуют на все блоки разом: когорта
собирается одним запросом, остальное — агрегаты по ней в памяти (сотни–тысячи карточек в
месяц). По каждой цифре отдаётся ключ, по которому ручка вернёт список карточек (ТЗ 8.1:
«клик по любой цифре открывает список»).
"""

import statistics
from collections import defaultdict
from datetime import datetime, timedelta
from io import BytesIO

from . import queries, rules

# Этап «дата этапа» для дополнительного фильтра (ТЗ 8.1): поступление — по received_at,
# остальные — по входу в колонку.
STAGE_DATE_OPTIONS = (
    ('received', 'Поступление'), (rules.IN_WORK, 'Взятие в работу'), (rules.INVITED_TRAINING, 'Приглашен на обучение'),
    (rules.ON_TRAINING, 'Вышел на обучение'), (rules.ON_LINE, 'Вышел на линию'), (rules.HIRED, 'Отработал 7 дней'),
    (rules.REFUSED, 'Отказ'), (rules.DROPPED, 'Отсев'),
)
STAGE_DATE_CODES = tuple(code for code, _ in STAGE_DATE_OPTIONS)

# Пары этапов для отчёта «Время на этапах» (ТЗ 8.2).
TIME_SPANS = (
    ('received', rules.IN_WORK, 'От поступления до взятия в работу'),
    (rules.IN_WORK, 'invited', 'От взятия в работу до приглашения'),
    ('invited', rules.ON_TRAINING, 'От приглашения до выхода на обучение'),
    (rules.ON_TRAINING, rules.ON_LINE, 'От обучения до выхода на линию'),
    (rules.ON_LINE, rules.HIRED, 'От выхода на линию до 7 смен'),
)


def parse_filters(args, *, forced_recruiter_id=None):
    """Фильтры из query-строки. month — 'YYYY-MM' (по умолчанию текущий); period from/to —
    альтернатива месяцу для отчётов 8.2; recruiters / sources — через запятую; stage + stage_from /
    stage_to — «дата этапа»."""
    now = rules.now_almaty()
    month = str(args.get('month') or '').strip()
    period_from = _date_arg(args.get('from'))
    period_to = _date_arg(args.get('to'))
    if period_from and period_to:
        cohort_from, cohort_to = period_from, period_to + timedelta(days=1)
        month = None
    else:
        if not month:
            month = rules.month_of(now)
        try:
            cohort_from, cohort_to = rules.month_bounds(month)
        except (TypeError, ValueError):
            month = rules.month_of(now)
            cohort_from, cohort_to = rules.month_bounds(month)
    recruiters = _ids_arg(args.get('recruiters'))
    if forced_recruiter_id:
        recruiters = [int(forced_recruiter_id)]
    stage = str(args.get('stage') or '').strip()
    stage_from = _date_arg(args.get('stage_from'))
    stage_to = _date_arg(args.get('stage_to'))
    if stage not in STAGE_DATE_CODES or not (stage_from or stage_to):
        stage, stage_from, stage_to = None, None, None
    return {
        'month': month, 'cohort_from': cohort_from, 'cohort_to': cohort_to,
        'recruiters': recruiters, 'sources': _ids_arg(args.get('sources')),
        'vacancy_id': _int(args.get('vacancy_id')), 'city': str(args.get('city') or '').strip() or None,
        'stage': stage, 'stage_from': stage_from, 'stage_to': (stage_to + timedelta(days=1)) if stage_to else None,
    }


def _int(value):
    try:
        return int(value) if value not in (None, '') else None
    except (TypeError, ValueError):
        return None


def _ids_arg(value):
    if isinstance(value, (list, tuple)):
        raw = value
    else:
        raw = str(value or '').split(',')
    return sorted({int(x) for x in raw if str(x).strip().isdigit()})


def _date_arg(value):
    text = str(value or '').strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text[:10])
    except ValueError:
        return None


# ─── Когорта ──────────────────────────────────────────────────────────────────

_COHORT_SQL = """
    WITH first_work AS (
        SELECT card_id, MIN(entered_at) AS taken_first FROM hiring_card_stages WHERE stage = 'in_work' GROUP BY card_id
    )
    SELECT c.id, c.candidate_no, c.full_name, c.kind, c.recruiter_id, r.name AS recruiter_name, c.source_id,
           s.name AS source_name, c.batch_id, b.file_name AS batch_file, c.vacancy_id, v.name AS vacancy_name,
           c.city, c.stage, c.is_closed, c.close_initiator, c.close_reason_text, c.close_stage, c.close_day,
           c.received_at, fw.taken_first, c.hired_at,
           (SELECT m.status FROM hiring_motivation m WHERE m.card_id = c.id) AS motivation_status
      FROM hiring_cards c
      JOIN first_work fw ON fw.card_id = c.id
      LEFT JOIN users r ON r.id = c.recruiter_id
      LEFT JOIN hiring_sources s ON s.id = c.source_id
      LEFT JOIN hiring_batches b ON b.id = c.batch_id
      LEFT JOIN hiring_vacancies v ON v.id = c.vacancy_id
     WHERE fw.taken_first >= %(cohort_from)s AND fw.taken_first < %(cohort_to)s
"""


def _cohort(cursor, filters):
    where, params = [], dict(cohort_from=filters['cohort_from'], cohort_to=filters['cohort_to'])
    if filters.get('recruiters'):
        where.append("c.recruiter_id = ANY(%(recruiters)s)")
        params['recruiters'] = filters['recruiters']
    if filters.get('sources'):
        where.append("c.source_id = ANY(%(sources)s)")
        params['sources'] = filters['sources']
    if filters.get('vacancy_id'):
        where.append("c.vacancy_id = %(vacancy_id)s")
        params['vacancy_id'] = filters['vacancy_id']
    if filters.get('city'):
        where.append("c.city ILIKE %(city)s")
        params['city'] = filters['city']
    if filters.get('stage'):
        if filters['stage'] == 'received':
            column = "c.received_at"
            where.append("(%(stage_from)s::timestamp IS NULL OR " + column + " >= %(stage_from)s)")
            where.append("(%(stage_to)s::timestamp IS NULL OR " + column + " < %(stage_to)s)")
        else:
            where.append("""EXISTS (SELECT 1 FROM hiring_card_stages st WHERE st.card_id = c.id AND st.stage = %(stage)s
                                     AND (%(stage_from)s::timestamp IS NULL OR st.entered_at >= %(stage_from)s)
                                     AND (%(stage_to)s::timestamp IS NULL OR st.entered_at < %(stage_to)s))""")
            params['stage'] = filters['stage']
        params['stage_from'] = filters.get('stage_from')
        params['stage_to'] = filters.get('stage_to')
    sql = _COHORT_SQL + (" AND " + " AND ".join(where) if where else "") + " ORDER BY fw.taken_first, c.id"
    cursor.execute(sql, params)
    columns = [column[0] for column in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


def _reached(cursor, ids):
    """{card_id: {stage: первый вход}} + «дозвон» и «приглашён (любой)»."""
    reached = defaultdict(dict)
    if not ids:
        return reached
    cursor.execute("""
        SELECT card_id, stage, MIN(entered_at) FROM hiring_card_stages WHERE card_id = ANY(%s) GROUP BY card_id, stage
    """, (ids,))
    for card_id, stage, moment in cursor.fetchall():
        reached[int(card_id)][stage] = moment
    cursor.execute("""
        SELECT card_id, MIN(at) FROM hiring_touches WHERE card_id = ANY(%s) AND kind = 'call' AND result = 'dial' GROUP BY card_id
    """, (ids,))
    for card_id, moment in cursor.fetchall():
        reached[int(card_id)]['dialed'] = moment
    for card_id, stages in reached.items():
        moments = [stages[s] for s in (rules.INVITED_OFFICE, rules.INVITED_TRAINING) if s in stages]
        if moments:
            stages['invited'] = min(moments)
    return reached


def _extra_counts(cursor, ids):
    """Касания и возвраты в пул по карточкам когорты (отчёт «Рекрутеры»)."""
    touches, returned = defaultdict(int), defaultdict(int)
    if not ids:
        return touches, returned
    cursor.execute("""
        SELECT card_id, COUNT(*) FROM hiring_touches WHERE card_id = ANY(%s) AND kind IN ('call', 'whatsapp', 'comment')
         GROUP BY card_id
    """, (ids,))
    for card_id, count in cursor.fetchall():
        touches[int(card_id)] = int(count)
    cursor.execute("""
        SELECT card_id, COUNT(*) FROM hiring_assignments WHERE card_id = ANY(%s) AND reason IN ('auto_timeout', 'release', 'recruiter_off')
         GROUP BY card_id
    """, (ids,))
    for card_id, count in cursor.fetchall():
        returned[int(card_id)] = int(count)
    return touches, returned


# ─── Дашборд ──────────────────────────────────────────────────────────────────

class Aggregator:
    """Считает блоки по когорте и запоминает, из каких карточек сложилась каждая цифра."""

    def __init__(self, cards, reached):
        self.cards = cards
        self.reached = reached
        self.ids = {}

    def _remember(self, key, ids):
        self.ids[key] = sorted(ids)
        return len(ids)

    def hit(self, card, stage):
        return stage in self.reached.get(card['id'], {})

    def funnel(self, cards, prefix):
        """Строки воронки: достигли, конверсия от предыдущего, сквозная."""
        rows = []
        base = None
        previous = None
        for stage in rules.FUNNEL_STAGES:
            ids = [card['id'] for card in cards if self.hit(card, stage)]
            count = self._remember('%s:%s' % (prefix, stage), ids)
            if base is None:
                base = count
            rows.append({
                'stage': stage, 'label': rules.FUNNEL_LABELS[stage], 'count': count,
                'from_previous': 100.0 if previous is None and count else rules.through_rate(count, previous),
                'through': rules.through_rate(count, base), 'key': '%s:%s' % (prefix, stage),
            })
            previous = count
        return rows

    def losses(self, cards, prefix):
        out = {}
        for kind in (rules.REFUSED, rules.DROPPED):
            closed = [card for card in cards if card['stage'] == kind and card['is_closed']]
            total = self._remember('%s:%s' % (prefix, kind), [card['id'] for card in closed])
            by_initiator = {}
            for initiator, _label in rules.SUBSTAGES[kind]:
                ids = [card['id'] for card in closed if card['close_initiator'] == initiator]
                by_initiator[initiator] = {'count': self._remember('%s:%s:%s' % (prefix, kind, initiator), ids),
                                           'label': rules.SUBSTAGE_LABELS[kind][initiator],
                                           'key': '%s:%s:%s' % (prefix, kind, initiator)}
            reasons = defaultdict(list)
            for card in closed:
                reasons[(card['close_initiator'], card['close_reason_text'] or '—')].append(card['id'])
            top = sorted(reasons.items(), key=lambda item: (-len(item[1]), item[0][1]))
            reason_rows = []
            for (initiator, text), ids in top:
                key = '%s:%s:reason:%s:%s' % (prefix, kind, initiator, text)
                reason_rows.append({'initiator': initiator, 'initiator_label': rules.SUBSTAGE_LABELS[kind].get(initiator, ''),
                                    'reason': text, 'count': self._remember(key, ids), 'key': key})
            block = {'kind': kind, 'label': rules.STAGE_LABELS[kind], 'count': total, 'key': '%s:%s' % (prefix, kind),
                     'share_of_received': rules.through_rate(total, len(cards)), 'by_initiator': by_initiator,
                     'reasons': reason_rows}
            if kind == rules.DROPPED:
                by_day = defaultdict(list)
                for card in closed:
                    by_day[(card['close_stage'] or '—', card['close_day'] or 0)].append(card['id'])
                block['by_day'] = [{'stage': stage, 'stage_label': rules.DROP_STAGE_LABELS.get(stage, stage), 'day': day,
                                    'count': self._remember('%s:%s:day:%s:%s' % (prefix, kind, stage, day), ids),
                                    'key': '%s:%s:day:%s:%s' % (prefix, kind, stage, day)}
                                   for (stage, day), ids in sorted(by_day.items(), key=lambda item: (item[0][0], item[0][1]))]
            out[kind] = block
        return out

    def group_rows(self, cards, group_key, label_key, prefix, extra=None):
        """Строка на рекрутера / источник / партию: каждый этап, сквозная, отказы, отсев, мотивация."""
        groups = defaultdict(list)
        labels = {}
        for card in cards:
            gid = card.get(group_key)
            groups[gid].append(card)
            labels[gid] = card.get(label_key) or '—'
        rows = []
        for gid, members in groups.items():
            prefix_g = '%s:%s' % (prefix, gid)
            row = {'id': gid, 'label': labels[gid], 'taken': len(members), 'stages': {}}
            for stage in rules.FUNNEL_STAGES:
                ids = [card['id'] for card in members if self.hit(card, stage)]
                row['stages'][stage] = {'count': self._remember('%s:%s' % (prefix_g, stage), ids), 'key': '%s:%s' % (prefix_g, stage)}
            hired = row['stages'][rules.HIRED]['count']
            row['through'] = rules.through_rate(hired, len(members))
            for kind in (rules.REFUSED, rules.DROPPED):
                ids = [card['id'] for card in members if card['stage'] == kind and card['is_closed']]
                row[kind] = {'count': self._remember('%s:%s' % (prefix_g, kind), ids), 'key': '%s:%s' % (prefix_g, kind)}
            ids = [card['id'] for card in members if card.get('motivation_status') in (rules.MOTIVATION_COUNTED, rules.MOTIVATION_PAID)]
            row['motivation'] = {'count': self._remember('%s:motivation' % prefix_g, ids), 'key': '%s:motivation' % prefix_g}
            if extra:
                row.update(extra(members))
            rows.append(row)
        rows.sort(key=lambda item: (-item['taken'], item['label']))
        return rows


def dashboard(cursor, filters, *, with_ids=False):
    cards = _cohort(cursor, filters)
    ids = [card['id'] for card in cards]
    reached = _reached(cursor, ids)
    agg = Aggregator(cards, reached)
    touches, returned = _extra_counts(cursor, ids)

    cursor.execute("SELECT COUNT(*) FROM hiring_cards c WHERE c.received_at >= %s AND c.received_at < %s"
                   + (" AND c.source_id = ANY(%s)" if filters.get('sources') else ""),
                   [filters['cohort_from'], filters['cohort_to']] + ([filters['sources']] if filters.get('sources') else []))
    received_in_period = int(cursor.fetchone()[0])

    funnel = agg.funnel(cards, 'total')
    counts = {row['stage']: row['count'] for row in funnel}
    hired_ids = [card['id'] for card in cards if agg.hit(card, rules.HIRED)]
    motivation_ids = [card['id'] for card in cards if card.get('motivation_status') in (rules.MOTIVATION_COUNTED, rules.MOTIVATION_PAID)]
    agg._remember('total:motivation', motivation_ids)

    plan_fte = 0.0
    if filters.get('month'):
        month_start = filters['cohort_from'].date()
        plans = queries.list_plans(cursor, month_start)
        for plan in plans:
            if filters.get('recruiters') and plan.get('recruiter_id') and plan['recruiter_id'] not in filters['recruiters']:
                continue
            if filters.get('vacancy_id') and plan.get('vacancy_id') and plan['vacancy_id'] != filters['vacancy_id']:
                continue
            plan_fte += float(plan['plan_fte'] or 0)
    cursor.execute("SELECT COUNT(*) FROM hiring_cards c WHERE c.hired_at >= %s AND c.hired_at < %s"
                   + (" AND c.recruiter_id = ANY(%s)" if filters.get('recruiters') else "")
                   + (" AND c.source_id = ANY(%s)" if filters.get('sources') else ""),
                   [filters['cohort_from'], filters['cohort_to']]
                   + ([filters['recruiters']] if filters.get('recruiters') else [])
                   + ([filters['sources']] if filters.get('sources') else []))
    closed_fte = int(cursor.fetchone()[0])

    totals = [
        {'key': 'received', 'label': 'Поступило', 'value': received_in_period, 'hint': 'в пул за период, вне когорты'},
        {'key': 'total:in_work', 'label': 'Взято в работу', 'value': counts[rules.IN_WORK]},
        {'key': 'total:on_training', 'label': 'Вышли на обучение', 'value': counts[rules.ON_TRAINING]},
        {'key': 'total:on_line', 'label': 'Вышли на линию', 'value': counts[rules.ON_LINE]},
        {'key': 'total:hired', 'label': 'Отработали 7 дней', 'value': counts[rules.HIRED]},
        {'key': 'total:hired', 'label': 'Найм закрыт', 'value': len(hired_ids), 'hint': 'в головах'},
        {'key': 'closed_fte', 'label': 'Закрыто FTE плана', 'value': closed_fte, 'plan': plan_fte},
        {'key': 'total:hired', 'label': 'Сквозная конверсия', 'value': rules.through_rate(counts[rules.HIRED], counts[rules.IN_WORK]),
         'unit': '%', 'hint': 'взято в работу → 7 дней'},
        {'key': 'total:motivation', 'label': 'Засчитано мотивации', 'value': len(motivation_ids)},
    ]

    def recruiter_extra(members):
        return {'touches': sum(touches.get(card['id'], 0) for card in members),
                'returned': sum(returned.get(card['id'], 0) for card in members)}

    by_recruiter = agg.group_rows(cards, 'recruiter_id', 'recruiter_name', 'rec', extra=recruiter_extra)
    by_source = agg.group_rows(cards, 'source_id', 'source_name', 'src')
    batches = [card for card in cards if card.get('batch_id')]
    by_batch = agg.group_rows(batches, 'batch_id', 'batch_file', 'batch')
    for row in by_batch:
        source = next((card['source_name'] for card in batches if card['batch_id'] == row['id']), None)
        row['label'] = '%s · %s' % (source or '—', row['label'])

    result = {
        'filters': {'month': filters.get('month'), 'from': filters['cohort_from'].date().isoformat(),
                    'to': (filters['cohort_to'] - timedelta(days=1)).date().isoformat(),
                    'recruiters': filters.get('recruiters'), 'sources': filters.get('sources'), 'stage': filters.get('stage'),
                    'vacancy_id': filters.get('vacancy_id'), 'city': filters.get('city')},
        'cohort_size': len(cards),
        'totals': totals,
        'funnel': funnel,
        'losses': agg.losses(cards, 'total'),
        'by_recruiter': by_recruiter,
        'by_source': by_source,
        'by_batch': by_batch,
        'time_on_stages': time_on_stages(cards, reached),
    }
    if with_ids:
        result['ids'] = agg.ids
    return result


def card_ids_for(cursor, filters, key):
    """Карточки, из которых сложилась цифра key — тот же расчёт, что у дашборда."""
    data = dashboard(cursor, filters, with_ids=True)
    if key == 'received':
        cursor.execute("SELECT id FROM hiring_cards c WHERE c.received_at >= %s AND c.received_at < %s"
                       + (" AND c.source_id = ANY(%s)" if filters.get('sources') else "") + " ORDER BY id DESC LIMIT 2000",
                       [filters['cohort_from'], filters['cohort_to']] + ([filters['sources']] if filters.get('sources') else []))
        return [int(row[0]) for row in cursor.fetchall()]
    if key == 'closed_fte':
        cursor.execute("SELECT id FROM hiring_cards c WHERE c.hired_at >= %s AND c.hired_at < %s"
                       + (" AND c.recruiter_id = ANY(%s)" if filters.get('recruiters') else "")
                       + (" AND c.source_id = ANY(%s)" if filters.get('sources') else "") + " ORDER BY id DESC",
                       [filters['cohort_from'], filters['cohort_to']]
                       + ([filters['recruiters']] if filters.get('recruiters') else [])
                       + ([filters['sources']] if filters.get('sources') else []))
        return [int(row[0]) for row in cursor.fetchall()]
    return list(data['ids'].get(key, []))


def time_on_stages(cards, reached):
    """Среднее и медиана (в днях) между первыми входами в этапы по когорте."""
    rows = []
    for start, end, label in TIME_SPANS:
        spans = []
        for card in cards:
            stages = reached.get(card['id'], {})
            begin = card['received_at'] if start == 'received' else stages.get(start)
            finish = stages.get(end)
            if begin and finish and finish >= begin:
                spans.append((finish - begin).total_seconds() / 86400.0)
        rows.append({'label': label, 'count': len(spans),
                     'avg_days': round(sum(spans) / len(spans), 1) if spans else None,
                     'median_days': round(statistics.median(spans), 1) if spans else None})
    return rows


def motivation_report(cursor, filters):
    """Начисления за месяц по рекрутерам и кандидатам (ТЗ 7.2, 8.2)."""
    recruiter_id = filters['recruiters'][0] if filters.get('recruiters') and len(filters['recruiters']) == 1 else None
    rows = queries.motivation_list(cursor, filters['cohort_from'], filters['cohort_to'], recruiter_id=recruiter_id)
    if filters.get('recruiters') and len(filters['recruiters']) > 1:
        rows = [row for row in rows if row.get('recruiter_id') in filters['recruiters']]
    by_recruiter = defaultdict(lambda: {'counted': 0, 'paid': 0, 'pending': 0, 'cancelled': 0, 'items': []})
    for row in rows:
        bucket = by_recruiter[(row.get('recruiter_id'), row.get('recruiter_name') or '—')]
        bucket[row['status']] = bucket.get(row['status'], 0) + 1
        bucket['items'].append(row)
    return [{'recruiter_id': key[0], 'recruiter_name': key[1], **value}
            for key, value in sorted(by_recruiter.items(), key=lambda item: item[0][1])]


# ─── Excel ────────────────────────────────────────────────────────────────────

def export_dashboard(data, title):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill('solid', fgColor='1F2937')
    header_font = Font(bold=True, color='FFFFFF')
    book = Workbook()

    def sheet(name, headers, rows, widths=None):
        ws = book.create_sheet(title=name)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in rows:
            ws.append(row)
        for index, width in enumerate(widths or (), start=1):
            ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
        ws.freeze_panes = 'A2'
        return ws

    book.remove(book.active)
    filters = data['filters']
    period = filters.get('month') or ('%s — %s' % (filters.get('from'), filters.get('to')))
    sheet('Итоги', ['Показатель', 'Значение', 'Пояснение'],
          [[row['label'], row['value'] if row['value'] is not None else '—',
            (row.get('hint') or '') + ((' · план %.2f FTE' % row['plan']) if row.get('plan') else '')] for row in data['totals']],
          widths=(34, 14, 40))
    ws = book['Итоги']
    ws.insert_rows(1)
    ws['A1'] = '%s · %s' % (title, period)
    ws['A1'].font = Font(bold=True, size=13)
    sheet('Воронка', ['Этап', 'Кол-во', 'Конверсия от предыдущего, %', 'Сквозная конверсия, %'],
          [[row['label'], row['count'], row['from_previous'] if row['from_previous'] is not None else '—',
            row['through'] if row['through'] is not None else '—'] for row in data['funnel']], widths=(36, 10, 26, 24))
    loss_rows = []
    for kind in (rules.REFUSED, rules.DROPPED):
        block = data['losses'][kind]
        loss_rows.append([block['label'], '', block['count'], block['share_of_received'] if block['share_of_received'] is not None else '—'])
        for row in block['reasons']:
            loss_rows.append(['', '%s · %s' % (row['initiator_label'], row['reason']), row['count'], ''])
    sheet('Потери', ['Вид', 'Инициатор · причина', 'Кол-во', 'Доля от взятых, %'], loss_rows, widths=(14, 60, 10, 18))
    stage_headers = [rules.FUNNEL_LABELS[s] for s in rules.FUNNEL_STAGES]
    sheet('Рекрутеры', ['Рекрутер', 'Взято'] + stage_headers + ['Сквозная, %', 'Отказы', 'Отсев', 'Мотивация', 'Касаний', 'Возвращено в пул'],
          [[row['label'], row['taken']] + [row['stages'][s]['count'] for s in rules.FUNNEL_STAGES]
           + [row['through'] if row['through'] is not None else '—', row[rules.REFUSED]['count'], row[rules.DROPPED]['count'],
              row['motivation']['count'], row.get('touches', 0), row.get('returned', 0)] for row in data['by_recruiter']],
          widths=(28, 8) + (14,) * len(stage_headers) + (12, 10, 10, 12, 10, 16))
    sheet('Источники', ['Источник', 'Взято'] + stage_headers + ['Сквозная, %', 'Отказы', 'Отсев'],
          [[row['label'], row['taken']] + [row['stages'][s]['count'] for s in rules.FUNNEL_STAGES]
           + [row['through'] if row['through'] is not None else '—', row[rules.REFUSED]['count'], row[rules.DROPPED]['count']]
           for row in data['by_source']] +
          [['Партия: ' + row['label'], row['taken']] + [row['stages'][s]['count'] for s in rules.FUNNEL_STAGES]
           + [row['through'] if row['through'] is not None else '—', row[rules.REFUSED]['count'], row[rules.DROPPED]['count']]
           for row in data['by_batch']],
          widths=(40, 8) + (14,) * len(stage_headers) + (12, 10, 10))
    sheet('Время на этапах', ['Отрезок', 'Карточек', 'Среднее, дней', 'Медиана, дней'],
          [[row['label'], row['count'], row['avg_days'] if row['avg_days'] is not None else '—',
            row['median_days'] if row['median_days'] is not None else '—'] for row in data['time_on_stages']],
          widths=(44, 10, 14, 14))
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def export_motivation(rows, period_label):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    ws = book.active
    ws.title = 'Мотивация'
    ws.append(['Рекрутер', 'ID кандидата', 'Кандидат', 'Вакансия', 'Выход на линию', 'Смен', 'Статус', 'Дата подтверждения', 'Выплачено'])
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='1F2937')
        cell.font = Font(bold=True, color='FFFFFF')
    for group in rows:
        for item in group['items']:
            ws.append([group['recruiter_name'], item.get('candidate_no'), item.get('full_name'), item.get('vacancy_name') or '',
                       item.get('line_started_at'), item.get('shifts_done'), rules.MOTIVATION_LABELS.get(item['status'], item['status']),
                       item.get('confirmed_at'), item.get('paid_at')])
    ws.insert_rows(1)
    ws['A1'] = 'Мотивация рекрутеров · %s' % period_label
    ws['A1'].font = Font(bold=True, size=13)
    for column, width in zip('ABCDEFGHI', (28, 12, 30, 24, 16, 8, 22, 20, 20)):
        ws.column_dimensions[column].width = width
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
