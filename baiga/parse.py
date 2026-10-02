# -*- coding: utf-8 -*-
"""Разбор и проверка еженедельного файла Байги. Ни базы, ни Flask, ни сети.

Формат (постановка #356, п. 3; сверено с настоящим файлом за 21–27.09.2026):

* несколько листов-зачётов (группы городов + «МОТО БАЙГА»), список меняется —
  имя листа сохраняется полем «Зачёт»;
* на каждом листе одинаковые 11 колонок, названия — в первой строке; колонки
  ищутся по НАЗВАНИЮ, а не по месту; пустые колонки справа не читаются;
* «Дата» — текст «ДД.ММ.ГГ», одна на весь файл: понедельник, начало периода;
* «Неделя» — номер недели (в живом файле — номер по ISO);
* «Наименование приза» — «Не указано» или «200 000 тенге»; число из него —
  сумма приза для фильтра;
* «Номер ВУ» почти всегда «AA123456», но встречаются записанные ЧИСЛОМ
  (10 цифр) — номер читается текстом при любом типе ячейки;
* «ID водителя» — 32 шестнадцатеричных символа (id в CRM yataxi / Флите).

Период берётся из имени файла («Список байги - 21.09.2026 - 27.09.2026.xlsx»,
и тот же вид с подчёркиваниями, в каком файл часто приходит после пересылки).
Нет дат в имени — из колонки «Дата» (с предупреждением): переименованный файл
не должен становиться «неделей без периода».

Что блокирует загрузку и что только предупреждает — п. 7 постановки:

    ошибка          не тот формат; нет нужных колонок; пустая обязательная
                    ячейка; не число в числовой колонке; нераспознанная дата
    предупреждение  нестандартный ID или номер ВУ; водитель повторяется; дата
                    вне периода; приз не распознан; лист без таблицы

Ошибки и предупреждения называют лист и номер строки Excel — аналитик правит
файл, а не гадает, где в нём «строка 412».
"""

import hashlib
import re
from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO

from wiki.text import fold_kazakh

MAX_FILE_BYTES = 20 * 1024 * 1024
# Строк во всём файле. Эталонная неделя — 900; сто тысяч — уже не неделя Байги,
# а чей-то чужой файл, и разбирать его до конца незачем.
MAX_ROWS = 100000
# Сколько ошибок отдать экрану. Не тот файл даёт тысячи одинаковых — первых
# двухсот хватает понять, в чём дело, остальные идут счётчиком.
MAX_ERRORS = 200
# Сколько строк каждого вида предупреждений показать (счётчик — полный).
MAX_WARNING_ITEMS = 50
# Где искать строку названий колонок: по постановке она первая, но заголовок
# «Итоги Байги» над таблицей не должен делать файл нечитаемым.
HEADER_SCAN_ROWS = 10
# Сколько колонок файла надо узнать, чтобы считать строку шапкой таблицы.
HEADER_MIN_KNOWN = 4

CAMPAIGN = 'baiga'

# key, заголовок в файле (он же — в выгрузке). Порядок — порядок колонок файла:
# выгрузка без фильтров обязана совпасть с исходником построчно (п. 8).
COLUMNS = (
    ('position', 'Позиция'),
    ('date', 'Дата'),
    ('week', 'Неделя'),
    ('driver_name', 'Водитель'),
    ('prize_name', 'Наименование приза'),
    ('amount', 'Сумма'),
    ('trips', 'Поездок'),
    ('city', 'Город'),
    ('park', 'Таксопарк'),
    ('license', 'Номер ВУ'),
    ('driver_id', 'ID водителя'),
)
COLUMN_TITLES = dict(COLUMNS)

# Другие написания тех же колонок. Немного и только очевидные: лишний синоним
# опаснее отсутствующего — он молча возьмёт не ту колонку.
_ALIASES = {
    'место': 'position',
    'фио': 'driver_name',
    'фио водителя': 'driver_name',
    'приз': 'prize_name',
    'поездки': 'trips',
    'кол-во поездок': 'trips',
    'количество поездок': 'trips',
    'парк': 'park',
    'ву': 'license',
    'номер водительского удостоверения': 'license',
    'id': 'driver_id',
}

# Пустая ячейка здесь — ошибка: без этих полей строка не строка Байги.
REQUIRED_CELLS = ('position', 'date', 'week', 'driver_name', 'amount', 'trips', 'driver_id')

# Подписи видов предупреждений, в порядке показа.
WARNING_LABELS = {
    'period_from_data': 'Период взят из колонки «Дата»',
    'sheet_skipped': 'Лист пропущен',
    'date_period': 'Дата вне периода',
    'duplicate': 'Водитель повторяется',
    'driver_id': 'Нестандартный ID водителя',
    'license': 'Нестандартный номер ВУ',
    'prize': 'Приз не распознан',
}

# Пределы колонок таблицы baiga_rows. Длиннее — это уже не ФИО и не номер, а
# мусор в ячейке; строку не режем молча, а показываем.
_LIMITS = {'driver_name': 255, 'prize_name': 255, 'city': 120, 'park': 255,
           'license': 64, 'driver_id': 64}

_NO_PRIZE = {'', 'не указано', 'нет', '-', '—', 'без приза'}

_PRIZE_TEXT = re.compile(r'^(\d[\d   ]*)\s*(?:тенге|тг\.?|₸|kzt)?$', re.IGNORECASE)
_INT_TEXT = re.compile(r'^[+-]?\d[\d   ]*(?:[.,]0+)?$')
_DATE_TEXT = re.compile(r'^(\d{1,2})[./-](\d{1,2})[./-](\d{4}|\d{2})$')
# Дата в имени файла: «21.09.2026», «14_09_2026», «21-09-26». Внутри ОДНОЙ даты
# разделитель один и тот же (обратная ссылка \2): иначе в записи диапазона
# «21.09-27.09.2026» кусок «21.09-27» читался как 21.09.2027. Пересылка меняет
# точки на «_», но меняет их все разом.
_NAME_DATE = re.compile(r'(?<!\d)(\d{1,2})([._\-/])(\d{1,2})\2(\d{4}|\d{2})(?![._\-/]?\d)')
# Диапазон с годом только в конце: «21.09-27.09.2026», «21.09 – 27.09.2026».
_NAME_RANGE = re.compile(
    r'(?<!\d)(\d{1,2})([._])(\d{1,2})\s*[-–—]\s*(\d{1,2})\2(\d{1,2})\2(\d{4}|\d{2})(?![._\-/]?\d)')

# Неделя Байги — понедельник…воскресенье, семь дней.
WEEK_DAYS = 7

# Потолки чисел — пределы колонок baiga_rows (INTEGER и BIGINT). Номер телефона,
# вставленный в «Поездок», иначе прошёл бы предпросмотр и уронил загрузку.
_INT32_MAX = 2147483647
_AMOUNT_MAX = 10 ** 15

_STANDARD_DRIVER_ID = re.compile(r'^[0-9a-f]{32}$')
_STANDARD_LICENSE = re.compile(r'^[A-Z]{2}\d{6}$')

# Кириллица, набранная вместо латиницы в номере ВУ (раскладка): «АВ123456».
_CYR_TO_LAT_UPPER = str.maketrans('АВЕКМНОРСТУХ', 'ABEKMHOPCTYX')
_CYR_TO_LAT_LOWER = str.maketrans('авекмнорстух', 'abekmhopctyx')

_SPACES = re.compile(r'[\s  ​]+')
_LICENSE_JUNK = re.compile(r'[\s  ​\-–—_.]+')


class ParseError(Exception):
    """Файл не разобрать вовсе — до строк дело не дошло."""

    def __init__(self, message, code='BAIGA_FILE_INVALID'):
        super().__init__(message)
        self.message = message
        self.code = code


# ─────────────────────────────────────────────────────────────────────────────
# Мелкие разборы — общие с поиском (filters.py) и тестами
# ─────────────────────────────────────────────────────────────────────────────

def normalize_header(value):
    text = _SPACES.sub(' ', str(value or '')).strip().lower().replace('ё', 'е')
    return text.rstrip(':').strip()


_HEADER_KEYS = {normalize_header(title): key for key, title in COLUMNS}
_HEADER_KEYS.update({normalize_header(alias): key for alias, key in _ALIASES.items()})


def header_key(value):
    return _HEADER_KEYS.get(normalize_header(value))


def clean_text(value):
    """Текст ячейки: число без «.0», пробелы схлопнуты, края обрезаны."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime('%d.%m.%y')
    return _SPACES.sub(' ', str(value)).strip()


def to_int(value):
    """Целое из ячейки или None, если там не целое число."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    if isinstance(value, str):
        text = value.strip()
        if _INT_TEXT.match(text):
            digits = re.sub(r'[   ]', '', text)
            return int(re.sub(r'[.,]0+$', '', digits))
    return None


def parse_date(value):
    """Дата из ячейки: настоящая дата Excel или текст «ДД.ММ.ГГ(ГГ)»."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        match = _DATE_TEXT.match(value.strip())
        if match:
            day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
            if year < 100:
                year += 2000
            try:
                return date(year, month, day)
            except ValueError:
                return None
    return None


def _make_date(day, month, year):
    year = int(year)
    if year < 100:
        year += 2000
    try:
        return date(year, int(month), int(day))
    except ValueError:
        return None


def dates_from_name(file_name):
    """Даты из имени файла по порядку: диапазон «21.09-27.09.2026» — две даты,
    «… - 21.09.2026 - 27.09.2026» — две, «… 21.09.2026» — одна."""
    name = str(file_name or '')
    match = _NAME_RANGE.search(name)
    if match:
        start = _make_date(match.group(1), match.group(3), match.group(6))
        end = _make_date(match.group(4), match.group(5), match.group(6))
        if start and end:
            return [start, end]
    found = []
    for match in _NAME_DATE.finditer(name):
        value = _make_date(match.group(1), match.group(3), match.group(4))
        if value:
            found.append(value)
    return found


def period_from_name(file_name):
    """(начало, конец), если имя называет неделю целиком (две даты), иначе None."""
    found = dates_from_name(file_name)
    return (found[0], found[1]) if len(found) >= 2 else None


def _fmt(value):
    return value.strftime('%d.%m.%Y')


def resolve_period(file_name, row_dates):
    """Период недели: ({start, end, source} или None, ошибка или None).

    Период — ключ недели: по нему повторная загрузка заменяет неделю, а не
    ложится рядом. Поэтому имя файла сверяется с колонкой «Дата» (в файле она
    одна на все строки — понедельник недели), а расхождение БЛОКИРУЕТ загрузку:
    неделя под чужой датой дала бы дубль водителей или заменила бы чужую неделю.

      * две даты в имени — неделя из них: ровно 7 дней, и «Дата» строк в ней;
      * одна дата — это начало или конец недели, решает «Дата» строк;
      * дат в имени нет — неделя от «Даты» строк (самой частой, а не самой
        ранней: одна опечатка «21.09.25» не должна уносить неделю на год назад).
    """
    data_start = Counter(row_dates).most_common(1)[0][0] if row_dates else None
    named = dates_from_name(file_name)
    if len(named) >= 2:
        start, end = named[0], named[1]
        if (end - start).days != WEEK_DAYS - 1:
            return None, ('В имени файла период не из 7 дней: %s – %s. Неделя Байги — с понедельника '
                          'по воскресенье' % (_fmt(start), _fmt(end)))
        if data_start and not start <= data_start <= end:
            return None, ('Период в имени файла (%s – %s) не совпадает с колонкой «Дата» (%s). '
                          'Проверьте имя файла' % (_fmt(start), _fmt(end), _fmt(data_start)))
        return {'start': start, 'end': end, 'source': 'name'}, None
    if not data_start:
        return None, 'Не удалось определить период: в колонке «Дата» нет дат'
    if len(named) == 1:
        single = named[0]
        if single == data_start:
            return {'start': single, 'end': single + timedelta(days=WEEK_DAYS - 1), 'source': 'name'}, None
        if single - timedelta(days=WEEK_DAYS - 1) == data_start:
            return {'start': data_start, 'end': single, 'source': 'name'}, None
        return None, ('Дата в имени файла (%s) не совпадает с колонкой «Дата» (%s). Проверьте имя файла'
                      % (_fmt(single), _fmt(data_start)))
    return {'start': data_start, 'end': data_start + timedelta(days=WEEK_DAYS - 1), 'source': 'data'}, None


def prize_amount(text):
    """(есть ли приз, сумма приза или None, распознан ли текст).

    «Не указано» — приза нет. «200 000 тенге» — приз на 200 000. Приз словами
    («Смартфон») — приз есть, суммы нет: в «только с призом» он попадает, в
    сумму призов — нет, и об этом предупреждение.
    """
    value = clean_text(text)
    if value.lower() in _NO_PRIZE:
        return False, None, True
    match = _PRIZE_TEXT.match(value)
    if match:
        amount = int(re.sub(r'\D', '', match.group(1)))
        if amount <= _INT32_MAX:
            return True, amount, True
    return True, None, False


def license_key(value):
    """Номер ВУ для сравнения: без пробелов и дефисов, заглавными, латиницей."""
    text = _LICENSE_JUNK.sub('', clean_text(value)).upper()
    return text.translate(_CYR_TO_LAT_UPPER)[:64]


def driver_key(value):
    return clean_text(value).lower()[:64]


def fold_text(value):
    """Свёртка для поиска: казахские буквы → русские двойники, ё → е, нижний
    регистр. Та же функция, что у вики, — Жусип = Жүсіп."""
    return fold_kazakh(clean_text(value)).lower()


def latin_variant(value):
    """Тот же текст, где кириллица-двойник заменена латиницей («ав12» → «ab12»)."""
    return str(value or '').translate(_CYR_TO_LAT_LOWER)


def search_text(name, license_value, driver_value):
    return ' '.join(part for part in (
        fold_text(name), license_key(license_value).lower(), driver_key(driver_value),
    ) if part)


def file_sha256(content):
    return hashlib.sha256(content or b'').hexdigest()


def is_standard_license(key):
    return bool(_STANDARD_LICENSE.match(key or ''))


def is_standard_driver_id(key):
    return bool(_STANDARD_DRIVER_ID.match(key or ''))


# ─────────────────────────────────────────────────────────────────────────────
# Учёт ошибок и предупреждений
# ─────────────────────────────────────────────────────────────────────────────

class _Issues:
    def __init__(self):
        self.errors = []
        self.errors_total = 0
        self.warnings = {kind: {'kind': kind, 'label': label, 'count': 0, 'items': []}
                         for kind, label in WARNING_LABELS.items()}

    @staticmethod
    def _item(message, sheet, row, column):
        item = {'message': message}
        if sheet is not None:
            item['sheet'] = sheet
        if row is not None:
            item['row'] = row
        if column is not None:
            item['column'] = column
        return item

    def error(self, message, sheet=None, row=None, column=None):
        self.errors_total += 1
        if len(self.errors) < MAX_ERRORS:
            self.errors.append(self._item(message, sheet, row, column))

    def warn(self, kind, message, sheet=None, row=None, column=None):
        bucket = self.warnings[kind]
        bucket['count'] += 1
        if len(bucket['items']) < MAX_WARNING_ITEMS:
            bucket['items'].append(self._item(message, sheet, row, column))

    def warning_groups(self):
        return [bucket for bucket in self.warnings.values() if bucket['count']]

    def warnings_total(self):
        return sum(bucket['count'] for bucket in self.warnings.values())


# ─────────────────────────────────────────────────────────────────────────────
# Книга
# ─────────────────────────────────────────────────────────────────────────────

def _open_workbook(content):
    """Книга в потоковом режиме или ParseError с понятной причиной."""
    if not content:
        raise ParseError('Файл пустой', 'BAIGA_FILE_EMPTY')
    if len(content) > MAX_FILE_BYTES:
        raise ParseError('Файл больше %d МБ' % (MAX_FILE_BYTES // (1024 * 1024)),
                         'BAIGA_FILE_TOO_BIG')
    if content[:4] == b'\xd0\xcf\x11\xe0':
        # Тот же контейнер OLE у .xlsx, сохранённого с паролем на открытие, —
        # естественный выбор для файла с персональными данными. Совет «сохраните
        # как .xlsx» ему не поможет: пароль останется.
        if 'EncryptedPackage'.encode('utf-16-le') in content:
            raise ParseError('Файл защищён паролем — снимите пароль в Excel и загрузите файл снова',
                             'BAIGA_FILE_ENCRYPTED')
        raise ParseError('Это файл старого формата .xls — откройте его в Excel и сохраните '
                         'как «Книга Excel (.xlsx)»', 'BAIGA_FILE_XLS')
    if content[:4] != b'PK\x03\x04':
        raise ParseError('Это не файл Excel (.xlsx)', 'BAIGA_FILE_NOT_XLSX')
    from openpyxl import load_workbook
    try:
        # read_only — поток, а не дерево: файл на 20 МБ в полном режиме съел бы
        # память инстанса. data_only — значения формул, а не их текст.
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl бросает что угодно
        raise ParseError('Файл не открылся как книга Excel (.xlsx)', 'BAIGA_FILE_BROKEN') from exc


_CANONICAL_KEYS = {normalize_header(title): key for key, title in COLUMNS}
_ALIAS_KEYS = {normalize_header(alias): key for alias, key in _ALIASES.items()}

# Сколько непустых ячеек в строке, чтобы считать её строкой данных, а не заметкой.
_DATA_ROW_MIN_CELLS = 3


def map_header(values):
    """Строка шапки → ({key: индекс колонки}, [повторённые названия]).

    Два прохода: сначала точные названия колонок файла, и только потом синонимы
    для того, чего не нашлось. Иначе лишняя колонка «ID» (номер строки из
    выгрузки CRM) слева от «ID водителя» молча забрала бы себе ID водителя. Две
    колонки с одним точным названием — ошибка: какую из них брать, файл не
    говорит, а молча взять левую значит загрузить не те числа.
    """
    columns, duplicates = {}, []
    for index, value in enumerate(values or ()):
        key = _CANONICAL_KEYS.get(normalize_header(value))
        if not key:
            continue
        if key in columns:
            if COLUMN_TITLES[key] not in duplicates:
                duplicates.append(COLUMN_TITLES[key])
            continue
        columns[key] = index
    for index, value in enumerate(values or ()):
        key = _ALIAS_KEYS.get(normalize_header(value))
        if key and key not in columns:
            columns[key] = index
    return columns, duplicates


def _find_header(rows_iter):
    """(номер строки шапки, {key: индекс}, [повторы], есть ли на листе данные).

    Шапка ищется в первых строках листа: заголовок «Итоги Байги» над таблицей не
    должен делать файл нечитаемым. Строки, прочитанные в поисках шапки, назад
    не нужны — поток не мотается.
    """
    scanned = 0
    data_like = False
    for number, values in rows_iter:
        scanned += 1
        columns, duplicates = map_header(values)
        if len(columns) >= HEADER_MIN_KNOWN:
            return number, columns, duplicates, True
        if sum(1 for value in (values or ()) if not _empty(value)) >= _DATA_ROW_MIN_CELLS:
            data_like = True
        if scanned >= HEADER_SCAN_ROWS:
            break
    return None, {}, [], data_like


def _cell(values, index):
    if index is None or values is None or index >= len(values):
        return None
    return values[index]


def _empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _row_from_cells(cells, *, sheet, sheet_order, number, issues):
    """Одна строка листа → словарь строки таблицы или None (если ошибки)."""
    ok = True
    for key in REQUIRED_CELLS:
        if _empty(cells[key]):
            issues.error('Пустая ячейка', sheet, number, COLUMN_TITLES[key])
            ok = False

    numbers = {}
    for key in ('position', 'week', 'amount', 'trips'):
        if _empty(cells[key]):
            continue
        value = to_int(cells[key])
        if value is None:
            issues.error('Не целое число: «%s»' % clean_text(cells[key])[:40], sheet, number,
                         COLUMN_TITLES[key])
            ok = False
            continue
        numbers[key] = value
    if 'position' in numbers and numbers['position'] < 1:
        issues.error('Позиция меньше единицы', sheet, number, COLUMN_TITLES['position'])
        ok = False
    if 'week' in numbers and not 1 <= numbers['week'] <= 53:
        issues.error('Номер недели вне 1–53', sheet, number, COLUMN_TITLES['week'])
        ok = False
    for key in ('amount', 'trips'):
        if numbers.get(key, 0) < 0:
            issues.error('Отрицательное число', sheet, number, COLUMN_TITLES[key])
            ok = False
    for key, ceiling in (('position', _INT32_MAX), ('trips', _INT32_MAX), ('amount', _AMOUNT_MAX)):
        if numbers.get(key, 0) > ceiling:
            issues.error('Слишком большое число', sheet, number, COLUMN_TITLES[key])
            ok = False

    row_date = None
    if not _empty(cells['date']):
        row_date = parse_date(cells['date'])
        if row_date is None:
            issues.error('Дата не распознана: «%s»' % clean_text(cells['date'])[:40], sheet, number,
                         COLUMN_TITLES['date'])
            ok = False

    texts = {key: clean_text(cells[key]) for key in _LIMITS}
    for key, limit in _LIMITS.items():
        if len(texts[key]) > limit:
            issues.error('Слишком длинное значение (%d знаков)' % len(texts[key]), sheet, number,
                         COLUMN_TITLES[key])
            ok = False
    if not ok:
        return None

    has_prize, amount_of_prize, recognised = prize_amount(cells['prize_name'])
    if not recognised:
        issues.warn('prize', '«%s»' % texts['prize_name'][:60], sheet, number, COLUMN_TITLES['prize_name'])

    license_value = texts['license']
    key_of_license = license_key(license_value)
    if not is_standard_license(key_of_license):
        issues.warn('license', '«%s»' % license_value[:40] if license_value else 'Номера нет',
                    sheet, number, COLUMN_TITLES['license'])
    key_of_driver = driver_key(texts['driver_id'])
    if not is_standard_driver_id(key_of_driver):
        issues.warn('driver_id', '«%s»' % texts['driver_id'][:40], sheet, number,
                    COLUMN_TITLES['driver_id'])

    raw_date = cells['date']
    date_text = clean_text(raw_date) if isinstance(raw_date, str) else row_date.strftime('%d.%m.%y')
    return {
        'sheet_order': sheet_order,
        'row_number': number,
        'zachet': sheet,
        'position': numbers['position'],
        'date_text': date_text[:16],
        'row_date': row_date,
        'week_number': numbers['week'],
        'driver_name': texts['driver_name'],
        'prize_name': texts['prize_name'],
        'prize_amount': amount_of_prize,
        'has_prize': has_prize,
        'amount': numbers['amount'],
        'trips': numbers['trips'],
        'city': texts['city'],
        'park': texts['park'],
        'license': license_value,
        'license_key': key_of_license,
        'driver_id': texts['driver_id'],
        'driver_key': key_of_driver,
        'search_text': search_text(texts['driver_name'], license_value, texts['driver_id']),
    }


def _read_sheet(worksheet, *, sheet_order, issues, rows, budget):
    """Читает лист в `rows`. Возвращает число строк листа или None (пропущен)."""
    title = clean_text(worksheet.title)[:64]
    try:
        # Тег dimension в файле бывает неверным, и поток тогда обрывается на
        # первой строке; сброс заставляет читать лист до конца.
        worksheet.reset_dimensions()
    except AttributeError:
        pass
    rows_iter = enumerate(worksheet.iter_rows(values_only=True), start=1)
    header_row, columns, duplicates, data_like = _find_header(rows_iter)
    if header_row is None:
        # Лист с данными, но без шапки, — ошибка, а не пропуск: иначе неделя
        # легла бы без этого зачёта, а при замене недели прежние строки зачёта
        # удалились бы, а новые не легли. Пропускаем только пустые листы и
        # листы с заметками.
        if data_like:
            issues.error('Нет строки с названиями колонок («Позиция», «Водитель», «Сумма»…)', title)
            return 0
        issues.warn('sheet_skipped', 'Нет таблицы — лист пропущен', title)
        return None
    if duplicates:
        issues.error('Колонка встречается дважды: %s' % ', '.join('«%s»' % name for name in duplicates),
                     title, header_row)
        return 0
    missing = [title_ for key, title_ in COLUMNS if key not in columns]
    if missing:
        issues.error('Нет колонок: %s' % ', '.join('«%s»' % name for name in missing), title, header_row)
        return 0

    count = 0
    for number, values in rows_iter:
        cells = {key: _cell(values, index) for key, index in columns.items()}
        if all(_empty(value) for value in cells.values()):
            continue
        if budget[0] <= 0:
            issues.error('Строк больше %d — это не неделя Байги' % MAX_ROWS, title, number)
            raise _TooManyRows()
        budget[0] -= 1
        row = _row_from_cells(cells, sheet=title, sheet_order=sheet_order, number=number, issues=issues)
        if row is not None:
            rows.append(row)
        count += 1
    if not count:
        issues.warn('sheet_skipped', 'На листе нет строк', title)
        return None
    return count


class _TooManyRows(Exception):
    pass


def parse_workbook(content, file_name):
    """Разбирает файл недели. Ничего не пишет.

    Возвращает словарь: `ok`, `period`, `sheets`, `stats`, `errors`,
    `warnings` (по видам) и `rows` — строки для записи. `rows` наружу (в ответ
    API) не отдаётся: там ФИО и номера ВУ, а экрану предпросмотра нужны только
    итоги. Ошибка формата файла — ParseError.
    """
    workbook = _open_workbook(content)
    issues = _Issues()
    rows = []
    sheets = []
    budget = [MAX_ROWS]
    try:
        for order, worksheet in enumerate(workbook.worksheets):
            if len(sheets) >= 64:
                issues.error('Листов больше 64 — это не неделя Байги')
                break
            try:
                count = _read_sheet(worksheet, sheet_order=order, issues=issues, rows=rows, budget=budget)
            except _TooManyRows:
                break
            if count is not None:
                sheets.append({'name': clean_text(worksheet.title)[:64], 'order': order, 'rows': count})
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        # В потоковом режиме лист разбирается лениво, при чтении строк: битый
        # XML листа или неверная ячейка всплывают здесь, а не при открытии книги.
        raise ParseError('Файл не открылся как книга Excel (.xlsx) — он повреждён', 'BAIGA_FILE_BROKEN') from exc
    finally:
        workbook.close()

    if not sheets and not issues.errors_total:
        issues.error('В файле нет таблицы Байги: на листах нет строки с названиями колонок '
                     '(«Позиция», «Водитель», «Сумма»…)')

    duplicates = {}
    for row in rows:
        first = duplicates.get(row['driver_key'])
        if first is None:
            duplicates[row['driver_key']] = (row['zachet'], row['row_number'])
            continue
        issues.warn('duplicate', 'Уже есть: лист «%s», строка %d' % first, row['zachet'],
                    row['row_number'], COLUMN_TITLES['driver_id'])

    period = None
    if rows:
        period, problem = resolve_period(file_name, [row['row_date'] for row in rows if row['row_date']])
        if problem:
            issues.error(problem)
        elif period['source'] == 'data':
            issues.warn('period_from_data',
                        'В имени файла нет дат — неделя считается по колонке «Дата»: %s – %s'
                        % (_fmt(period['start']), _fmt(period['end'])))
    if period:
        weeks = Counter(row['week_number'] for row in rows)
        period['week'] = weeks.most_common(1)[0][0] if weeks else None
        label = '%s–%s' % (period['start'].strftime('%d.%m'), period['end'].strftime('%d.%m.%Y'))
        for row in rows:
            if row['row_date'] and not period['start'] <= row['row_date'] <= period['end']:
                issues.warn('date_period', '%s — вне периода %s' % (row['date_text'], label),
                            row['zachet'], row['row_number'], COLUMN_TITLES['date'])

    stats = {
        'rows': len(rows),
        'drivers': len({row['driver_key'] for row in rows}),
        'prize_rows': sum(1 for row in rows if row['has_prize']),
        'prize_total': sum(row['prize_amount'] or 0 for row in rows),
    }
    return {
        'ok': not issues.errors_total and bool(rows) and period is not None,
        'campaign': CAMPAIGN,
        'file_name': str(file_name or '')[:255],
        'file_size': len(content),
        'sha256': file_sha256(content),
        'period': period,
        'sheets': sheets,
        'stats': stats,
        'errors': issues.errors,
        'errors_total': issues.errors_total,
        'warnings': issues.warning_groups(),
        'warnings_total': issues.warnings_total(),
        'rows': rows,
    }


def public_summary(result):
    """То, что видит экран предпросмотра: без строк и с датами строками ISO."""
    summary = {key: value for key, value in result.items() if key != 'rows'}
    period = result.get('period')
    if period:
        summary['period'] = {
            'start': period['start'].isoformat(),
            'end': period['end'].isoformat(),
            'source': period['source'],
            'week': period.get('week'),
        }
    return summary
