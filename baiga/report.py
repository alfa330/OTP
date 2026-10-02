# -*- coding: utf-8 -*-
"""Выгрузка выборки «Списков Байги» в xlsx (постановка #356, п. 6).

Выгружается текущая выборка со всеми фильтрами, в тех же 11 колонках, что
исходник. Два режима:

* «по зачётам» — лист на зачёт, как в исходнике: без фильтров выгрузка недели
  совпадает с загруженным файлом построчно и загружается обратно без ошибок
  (п. 8). Поэтому листа «Контекст», который ставят первым другие выгрузки
  портала, здесь НЕТ: лишний лист сломал бы обратную загрузку. Что выгружено,
  видно по имени файла (период) и по журналу выгрузок раздела;
* «одним листом» — все строки подряд в порядке экрана; зачёт уходит
  последней, двенадцатой колонкой, чтобы первые 11 остались колонками файла.

ID водителя и номер ВУ — ТЕКСТОМ: числом 10-значный номер ВУ теряет вид, а
32-значный ID Excel округлил бы. «Дата» — текстом «ДД.ММ.ГГ», как в исходнике.
Зелёные уголки «Число сохранено как текст» и «Дата с двузначным годом»
гасятся тегом <ignoredErrors> на каждом листе: значок на каждой строке —
шум, который владелец считает браком.

Книга пишется в потоковом режиме (write_only): до ста тысяч строк. Ловушки
режима — как в parcels/report.py: шапка и закрепление задаются до первого
append, объекты стиля создаются один раз модульными константами.
"""

import re
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .parse import COLUMNS

# Потолок строк в одном файле (п. 6). Больше — отказ с просьбой сузить выборку,
# а не молча обрезанный файл, который читался бы как полный.
EXPORT_LIMIT = 100000

MODE_SHEETS = 'sheets'
MODE_SINGLE = 'single'
MODES = (MODE_SHEETS, MODE_SINGLE)

SINGLE_SHEET_TITLE = 'Список байги'
ZACHET_TITLE = 'Зачёт'
FILE_PREFIX = 'Список байги'

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

_WIDTHS = {
    'position': 10, 'date': 11, 'week': 9, 'driver_name': 30, 'prize_name': 20, 'amount': 12,
    'trips': 10, 'city': 18, 'park': 28, 'license': 15, 'driver_id': 36,
}
_ZACHET_WIDTH = 32

# Колонки, которые обязаны остаться текстом (см. шапку модуля).
TEXT_KEYS = ('date', 'license', 'driver_id')

_ROW_VALUE = {
    'position': lambda row: row['position'],
    'date': lambda row: row['date_text'],
    'week': lambda row: row['week_number'],
    'driver_name': lambda row: row['driver_name'],
    'prize_name': lambda row: row['prize_name'],
    'amount': lambda row: row['amount'],
    'trips': lambda row: row['trips'],
    'city': lambda row: row['city'],
    'park': lambda row: row['park'],
    'license': lambda row: row['license'],
    'driver_id': lambda row: row['driver_id'],
}

_SHEET_PATH = re.compile(r'^xl/worksheets/sheet(\d+)\.xml$')
# Узлы листа, которые по схеме OOXML идут ПОСЛЕ <ignoredErrors>: перед первым из
# них и вставляем, иначе Excel объявит книгу повреждённой (тот же список, что у
# _excel_suppress_number_as_text_warning в bot_schedule2.py).
_TAGS_AFTER = (
    '<smartTags', '<drawing', '<legacyDrawing', '<picture', '<oleObjects',
    '<controls', '<webPublishItems', '<tableParts', '<extLst',
)


def file_name(period_start=None, period_end=None):
    """Имя файла. Одна неделя — тем же видом, что исходник: тогда обратная
    загрузка выгрузки берёт период из имени, как у настоящего файла."""
    if period_start and period_end:
        return '%s - %s - %s.xlsx' % (FILE_PREFIX, period_start.strftime('%d.%m.%Y'),
                                      period_end.strftime('%d.%m.%Y'))
    return '%s - все недели.xlsx' % FILE_PREFIX


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def _header(worksheet, titles):
    cells = []
    for title in titles:
        cell = WriteOnlyCell(worksheet, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cells.append(cell)
    return cells


def _text_cell(worksheet, value, text_format=False):
    """Текст строго ТЕКСТОМ.

    openpyxl превращает любую строку, начинающуюся с «=», в формулу — даже в
    ячейке с форматом «@». ФИО «=1+2» или приз «=HYPERLINK(…)» в выгрузке
    стали бы живой формулой, которую Excel посчитает при открытии файла с
    персональными данными, а обратная загрузка прочла бы пустую ячейку. Тип
    «строка» задаётся после значения, поверх догадки openpyxl.
    """
    cell = WriteOnlyCell(worksheet, value=value)
    cell.data_type = 's'
    if text_format:
        cell.number_format = '@'
    return cell


def _row_cells(worksheet, row, with_zachet):
    values = []
    for key, _title in COLUMNS:
        value = _clean(_ROW_VALUE[key](row))
        if key in TEXT_KEYS:
            values.append(_text_cell(worksheet, '' if value is None else str(value), text_format=True))
        elif isinstance(value, str):
            values.append(_text_cell(worksheet, value))
        else:
            values.append(value)
    if with_zachet:
        values.append(_text_cell(worksheet, _clean(row['zachet'])))
    return values


def _write_sheet(workbook, title, rows, with_zachet=False):
    worksheet = workbook.create_sheet(title=title[:31])
    titles = [column_title for _key, column_title in COLUMNS] + ([ZACHET_TITLE] if with_zachet else [])
    for index, (key, _title) in enumerate(COLUMNS, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = _WIDTHS[key]
    if with_zachet:
        worksheet.column_dimensions[get_column_letter(len(titles))].width = _ZACHET_WIDTH
    worksheet.freeze_panes = 'A2'
    last = '%s%d' % (get_column_letter(len(titles)), len(rows) + 1)
    worksheet.auto_filter.ref = 'A1:%s' % last
    worksheet.append(_header(worksheet, titles))
    for row in rows:
        worksheet.append(_row_cells(worksheet, row, with_zachet))
    return 'A1:%s' % last


def group_by_zachet(rows):
    """[(зачёт, строки)] в порядке первого появления. Строки приходят в порядке
    файла (неделя, лист, строка), так что и листы встают в порядке файла.

    Зачёты сравниваются без учёта регистра — как имена листов в Excel: «Мото
    Байга» одной недели и «МОТО БАЙГА» другой — один лист, а не два, второй из
    которых openpyxl переименовал бы в несуществующий зачёт «… Байга1».
    """
    groups, titles = {}, {}
    for row in rows:
        key = str(row['zachet'] or '').casefold()
        titles.setdefault(key, row['zachet'])
        groups.setdefault(key, []).append(row)
    return [(titles[key], group) for key, group in groups.items()]


def _sheet_titles(names):
    """Имена листов: не длиннее 31 знака (предел Excel) и без повторов после
    обрезки. Суффикс повтора умещается в те же 31 знак — иначе Excel счёл бы
    книгу повреждённой."""
    used, titles = set(), []
    for name in names:
        base = (name or 'Без зачёта')[:31]
        title, number = base, 1
        while title.casefold() in used:
            number += 1
            suffix = ' (%d)' % number
            title = base[:31 - len(suffix)] + suffix
        used.add(title.casefold())
        titles.append(title)
    return titles


def build(rows, mode=MODE_SHEETS):
    """Книга выгрузки (bytes)."""
    workbook = Workbook(write_only=True)
    ranges = []
    if mode == MODE_SINGLE or not rows:
        ranges.append(_write_sheet(workbook, SINGLE_SHEET_TITLE, rows, with_zachet=mode == MODE_SINGLE))
    else:
        groups = group_by_zachet(rows)
        for title, (_zachet, group) in zip(_sheet_titles([zachet for zachet, _ in groups]), groups):
            ranges.append(_write_sheet(workbook, title, group))
    stream = BytesIO()
    workbook.save(stream)
    return suppress_text_warnings(stream.getvalue(), ranges)


_CHUNK = 1024 * 1024
_TAIL = 64 * 1024
_SHEET_END = b'</worksheet>'


def _patch_sheet(source, target, item, node):
    """Переписывает лист ПОТОКОМ и вставляет узел перед первым из _TAGS_AFTER
    (их в листах этой выгрузки нет) или перед </worksheet>.

    Лист «одним листом» на сто тысяч строк — сотня мегабайт XML (кириллица у
    openpyxl — числовые сущности). Прочитать его целиком, раскодировать и
    склеить заново стоило бы полгигабайта памяти единственного процесса
    портала; поэтому в памяти держится только хвост — последние килобайты.
    """
    tags = [tag.encode('ascii') for tag in _TAGS_AFTER]
    with source.open(item) as reader, target.open(item.filename, 'w', force_zip64=True) as writer:
        tail = b''
        while True:
            chunk = reader.read(_CHUNK)
            if not chunk:
                break
            buffer = tail + chunk
            if len(buffer) > _TAIL:
                writer.write(buffer[:-_TAIL])
                tail = buffer[-_TAIL:]
            else:
                tail = buffer
        if b'<ignoredErrors' in tail:
            writer.write(tail)
            return
        positions = [pos for pos in (tail.find(tag) for tag in tags) if pos != -1]
        insert_at = min(positions) if positions else tail.rfind(_SHEET_END)
        if insert_at == -1:
            writer.write(tail)
            return
        writer.write(tail[:insert_at] + node + tail[insert_at:])


def suppress_text_warnings(content, ranges):
    """Дописывает <ignoredErrors> в каждый лист: ranges[i] — диапазон листа i+1.

    openpyxl этот узел не пишет (write_tail его пропускает), поэтому он
    вставляется в уже собранный xlsx перепаковкой zip.
    """
    source = ZipFile(BytesIO(content))
    patched = BytesIO()
    try:
        with ZipFile(patched, 'w', ZIP_DEFLATED) as target:
            for item in source.infolist():
                match = _SHEET_PATH.match(item.filename)
                if match and int(match.group(1)) <= len(ranges):
                    node = ('<ignoredErrors><ignoredError sqref="%s" numberStoredAsText="1" '
                            'twoDigitTextYear="1"/></ignoredErrors>'
                            % ranges[int(match.group(1)) - 1]).encode('ascii')
                    _patch_sheet(source, target, item, node)
                else:
                    target.writestr(item, source.read(item.filename))
    finally:
        source.close()
    return patched.getvalue()
