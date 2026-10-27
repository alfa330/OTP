# -*- coding: utf-8 -*-
"""Выгрузка ночного отчёта в xlsx (ТЗ #398 п. 4).

Колонки — ровно столбцы ТЗ в его порядке; «Статус аренды» только в отчёте за 01:00 (п. 3).
Второй лист «Контекст» — за какой слот отчёт, когда собран и кто из источников не ответил:
без этого через неделю по файлу не понять, полный это отчёт или половина.

Соглашения те же, что у выгрузок «Учёта воды» и «Посылок»: время — настоящей датой с
форматом, строки из внешних систем — явным текстом (имя объекта или адрес с «=» в начале
не станет формулой), книга пишется потоково.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import rules

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
DATETIME_FORMAT = 'DD.MM.YYYY HH:MM'

SHEET_REPORT = 'Отчёт'
SHEET_CONTEXT = 'Контекст'

# (ключ, заголовок ТЗ, ширина). rent_status — только в отчёте за 01:00.
COLUMNS = (
    ('gps_site', 'GPS-сайт', 14),
    ('plate', 'Гос. номер', 13),
    ('rent_status', 'Статус аренды', 15),
    ('gps_status', 'Статус GPS', 12),
    ('state', 'Состояние', 14),
    ('signal_at', 'Последнее время отправки трекинга', 20),
    ('address', 'Местоположение / Адрес', 60),
)

SOURCE_NAMES = {'yume': 'Yume', 'wialon': 'Wialon', 'glonass': 'ГЛОНАСС', 'efir': 'Efir'}


def columns_for(kind):
    return [column for column in COLUMNS if column[0] != 'rent_status' or kind == rules.KIND_NIGHT]


def report_title(report):
    slot = report['slot_at']
    return 'Отчёт по арендным машинам — %s %s' % (slot.strftime('%d.%m.%Y'), rules.KIND_LABELS[report['kind']])


def report_filename(report):
    """«Аренда авто 28.10.2026 01-00.xlsx» — двоеточие в имени файла Windows не пускает."""
    slot = report['slot_at']
    return 'Аренда авто %s %s.xlsx' % (slot.strftime('%d.%m.%Y'), rules.KIND_LABELS[report['kind']].replace(':', '-'))


def _text(sheet, value):
    text = ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value
    cell = WriteOnlyCell(sheet, value=text)
    if text is not None:
        cell.data_type = 's'
    return cell


def cell_text(key, row):
    if key == 'gps_site':
        return row.get('gps_site') or '—'
    if key == 'rent_status':
        return rules.RENT_LABELS.get(row.get('rent_status')) or '—'
    if key == 'gps_status':
        return rules.GPS_LABELS.get(row.get('gps_status')) or ''
    if key == 'state':
        return rules.STATE_LABELS.get(row.get('state')) or ''
    return row.get(key)


def _header_row(sheet, titles):
    row = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        row.append(cell)
    return row


def build_workbook(report, rows, *, generated_by='', generated_at=None):
    """Книга отчёта. Возвращает (поток, сколько строк записано)."""
    columns = columns_for(report['kind'])
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(SHEET_REPORT)
    for index, (_key, _title, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.append(_header_row(sheet, [title for _key, title, _width in columns]))
    for row in rows:
        cells = []
        for key, _title, _width in columns:
            if key == 'signal_at':
                cell = WriteOnlyCell(sheet, value=row.get('signal_at'))
                cell.number_format = DATETIME_FORMAT
            else:
                cell = _text(sheet, cell_text(key, row))
            cells.append(cell)
        sheet.append(cells)

    context = workbook.create_sheet(SHEET_CONTEXT)
    context.column_dimensions['A'].width = 26
    context.column_dimensions['B'].width = 70
    title = WriteOnlyCell(context, value=report_title(report))
    title.font = TITLE_FONT
    context.append([title])
    lines = [
        ('Собран', report['collected_at'].strftime('%d.%m.%Y %H:%M') if report.get('collected_at') else ''),
        ('Выгрузил', generated_by or ''),
        ('Выгружен', generated_at.strftime('%d.%m.%Y %H:%M') if generated_at else ''),
        ('Машин в отчёте', len(rows)),
    ]
    for name, status in (report.get('sources') or {}).items():
        if not status:
            continue
        label = SOURCE_NAMES.get(name, name)
        lines.append((label, ('ответил, машин: %d' % int(status.get('count') or 0)) if status.get('ok')
                      else 'не ответил: %s' % (status.get('error') or 'без причины')))
    if report.get('error'):
        lines.append(('Внимание', report['error']))
    for label, value in lines:
        value_cell = _text(context, value) if isinstance(value, str) else WriteOnlyCell(context, value=value)
        value_cell.alignment = WRAP_TOP
        context.append([label, value_cell])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream, len(rows)
