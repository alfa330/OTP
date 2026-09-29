# -*- coding: utf-8 -*-
"""Выгрузка жалоб в xlsx — «Администратор должен иметь возможность выгружать
данные и фильтровать их» (ТЗ задачи #297).

Три листа: «Контекст», «Жалобы», «По сотрудникам». Контекст — первым, как у
остальных выгрузок портала («Посылки», «Касания»): жалобы живые, итоги и работа
с сотрудником дописываются неделями, и через месяц по файлу иначе не понять, на
какую дату он собран и что в нём отобрано.

Соглашения те же, что в parcels/report.py:

* телефон и ID / ВУ — ТЕКСТОМ (числом номер теряет ведущие нули), зелёный
  уголок гасится тегом <ignoredErrors>; функция приходит аргументом, потому
  что живёт в монолите;
* даты — настоящими датами, чтобы сортировались и шли в сводную;
* книга потоковая (write_only): ширины и закрепление задаются до первой строки,
  стили — модульными константами.

Колонок ровно столько, сколько видит разбирающий в карточке, и не больше:
внутренние детали работы с сотрудником (что ему сказали) в выгрузку не идут —
только факты «ОС проведена / тренинг проведён», те же, что в аналитике.
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import catalog

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
TITLE_FONT = Font(bold=True, size=13)
LABEL_FONT = Font(bold=True)
WRAP = Alignment(vertical='top', wrap_text=True)
DATE_FORMAT = 'DD.MM.YYYY HH:MM'

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

COLUMNS = [
    # (заголовок, ключ, ширина)
    ('№', 'id', 7),
    ('Создана', 'created_at', 17),
    ('На кого / на что', 'target_title', 20),
    ('Причина', 'reason_title', 34),
    ('Подразделение', 'unit_name', 26),
    ('Сотрудник', 'employee_name', 28),
    ('Город', 'city', 14),
    ('Водитель', 'driver_name', 26),
    ('Телефон', 'driver_phone', 16),
    ('ID / ВУ', 'driver_ref', 14),
    ('Когда произошло', 'event_at', 17),
    ('Описание', 'description', 60),
    ('В группу', 'processing', 13),
    ('Статус', 'status_title', 13),
    ('Итог проверки', 'result_title', 30),
    ('Принятые меры', 'result_note', 40),
    ('Кто проверил', 'result_by_name', 24),
    ('Работа с сотрудником', 'work_title', 24),
    ('ОС проведена', 'feedback', 12),
    ('Тренинг проведён', 'training', 12),
    ('Принял жалобу', 'created_by_name', 26),
    ('Отработана', 'closed_at', 17),
]

TEXT_COLUMNS = ('driver_phone', 'driver_ref')

WORK_TITLES = {
    None: '',
    catalog.WORK_UNASSIGNED: 'Сотрудник не определён',
    catalog.WORK_PENDING: 'В работе',
    catalog.WORK_DONE: 'Завершена',
}


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def _as_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str) and value:
        try:
            return datetime.fromisoformat(value.replace('Z', ''))
        except ValueError:
            return value
    return None


def row_values(item):
    """Жалоба → значения строки листа «Жалобы». Отдельной функцией, чтобы тест
    проверял содержимое, не разбирая готовый xlsx."""
    work = WORK_TITLES.get(item.get('work_state'), '')
    if item.get('work_state') == catalog.WORK_PENDING and item.get('training_required'):
        work = 'Требуется тренинг'
    values = dict(item)
    values.update({
        'processing': 'Да' if item.get('requires_processing') else 'Зафиксирована',
        'status_title': 'Отработана' if item.get('status') == 'closed' else 'В работе',
        'work_title': work,
        'feedback': 'Да' if item.get('feedback_done') else '',
        'training': 'Да' if item.get('training_done') else '',
        'result_title': item.get('result_title') or '',
    })
    return [values.get(key) for _title, key, _width in COLUMNS]


def _header(sheet, titles):
    cells = []
    for title in titles:
        cell = WriteOnlyCell(sheet, value=title)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, HEADER_ALIGN
        cells.append(cell)
    sheet.append(cells)


def build_workbook(items, *, filters_note, generated_by, generated_at=None,
                   text_warning_patch=None):
    """Книга целиком. Возвращает BytesIO, готовый к send_file."""
    generated_at = generated_at or datetime.now()
    book = Workbook(write_only=True)

    context = book.create_sheet('Контекст')
    context.column_dimensions['A'].width = 34
    context.column_dimensions['B'].width = 90
    title = WriteOnlyCell(context, value='Жалобы — выгрузка')
    title.font = TITLE_FONT
    context.append([title])
    for label, value in (
        ('Собрано', generated_at.strftime('%d.%m.%Y %H:%M')),
        ('Кто выгрузил', generated_by or ''),
        ('Отбор', filters_note or 'без фильтров'),
        ('Строк', len(items)),
        ('Что в листе «Жалобы»', 'Каждая строка — одна жалоба. «В группу: Зафиксирована» — '
                                 'жалоба не уходила на разбор (Яндекс, часть жалоб на парк).'),
        ('Что в листе «По сотрудникам»', 'Жалобы, где сотрудник определён: сколько всего, '
                                         'сколько подтверждено, где проведены ОС и тренинг.'),
    ):
        head = WriteOnlyCell(context, value=label)
        head.font = LABEL_FONT
        body = WriteOnlyCell(context, value=_clean(value))
        body.alignment = WRAP
        context.append([head, body])

    sheet = book.create_sheet('Жалобы')
    for index, (_title, _key, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'B2'
    _header(sheet, [title for title, _key, _width in COLUMNS])
    date_keys = {'created_at', 'event_at', 'closed_at'}
    for item in items:
        cells = []
        for (_title, key, _width), value in zip(COLUMNS, row_values(item)):
            if key in date_keys:
                value = _as_datetime(value)
            cell = WriteOnlyCell(sheet, value=_clean(value if value is not None else ''))
            if key in date_keys and isinstance(value, datetime):
                cell.number_format = DATE_FORMAT
            if key in TEXT_COLUMNS:
                cell.number_format = '@'
            if key == 'description':
                cell.alignment = WRAP
            cells.append(cell)
        sheet.append(cells)
    if items:
        sheet.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(COLUMNS)), len(items) + 1)

    people = book.create_sheet('По сотрудникам')
    for index, width in enumerate((30, 26, 10, 12, 12, 12, 12), start=1):
        people.column_dimensions[get_column_letter(index)].width = width
    people.freeze_panes = 'A2'
    _header(people, ['Сотрудник', 'Подразделение', 'Жалоб', 'Подтверждено',
                     'ОС проведена', 'Тренинг', 'Повторная'])
    for person in employee_summary(items):
        people.append([_clean(person['name']), _clean(person['unit']), person['total'],
                       person['confirmed'], person['feedback'], person['training'],
                       'Да' if person['total'] > 1 else ''])

    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    if text_warning_patch and items:
        first = get_column_letter(1 + [key for _t, key, _w in COLUMNS].index('driver_phone'))
        last = get_column_letter(1 + [key for _t, key, _w in COLUMNS].index('driver_ref'))
        try:
            stream = text_warning_patch(stream, '%s2:%s%d' % (first, last, len(items) + 1),
                                        sheet_path='xl/worksheets/sheet2.xml')
            stream.seek(0)
        except Exception:  # noqa: BLE001 — уголок в Excel не повод терять выгрузку
            stream.seek(0)
    return stream


def employee_summary(items):
    """Жалобы по сотрудникам — те же числа, что в аналитике раздела."""
    confirmed_codes = set(catalog.CONFIRMED_RESULTS)
    people = {}
    for item in items:
        if not item.get('employee_id'):
            continue
        person = people.setdefault(item['employee_id'], {
            'name': item.get('employee_name') or '', 'unit': item.get('unit_name') or '',
            'total': 0, 'confirmed': 0, 'feedback': 0, 'training': 0,
        })
        person['total'] += 1
        person['confirmed'] += 1 if item.get('result_code') in confirmed_codes else 0
        person['feedback'] += 1 if item.get('feedback_done') else 0
        person['training'] += 1 if item.get('training_done') else 0
    return sorted(people.values(), key=lambda row: (-row['total'], row['name']))


def filename(date_from, date_to):
    def ru(value):
        try:
            return date.fromisoformat(str(value)).strftime('%d.%m.%Y')
        except ValueError:
            return str(value or '')
    if date_from and date_to:
        return 'Жалобы %s–%s.xlsx' % (ru(date_from), ru(date_to))
    return 'Жалобы.xlsx'
