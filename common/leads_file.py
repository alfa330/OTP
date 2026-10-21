# -*- coding: utf-8 -*-
"""Разбор файла «ФИО + телефон» (CSV / Excel) — общий для всех баз обзвона.

Формат один на весь проект: две колонки с шапкой `fio` и `phone` (шапка
необязательна — тогда первая колонка имя, вторая телефон). Телефон в Excel
часто хранится числом, поэтому приводится к строке без экспоненты.
"""
import csv
import io

from common.kz_phone import normalize_kz_phone

# Заголовки, которыми подписывают колонки (файл приходит как fio/phone).
FIO_HEADERS = {'fio', 'фио', 'имя', 'name', 'full_name', 'водитель', 'driver'}
PHONE_HEADERS = {'phone', 'телефон', 'номер', 'phone_number', 'msisdn', 'тел'}
# ИИН водителя — колонка базы «Обзвона» (with_iin=True): по нему проверяется
# подписание документов. Другим базам (TEZ) она не нужна и не ищется.
IIN_HEADERS = {'iin', 'иин', 'iin_driver', 'driver_iin', 'иин_водителя', 'iin_водителя'}
IIN_MISSING_MESSAGE = ('В файле нет колонки «iin» с ИИН водителя. Добавьте её: '
                       'без ИИН нельзя проверить подписание документов.')
# Город водителя — необязательная колонка базы «Удаленного КЦ» (with_city=True) и обязательная
# у списка городов (parse_city_list): по ней строится отчёт «Подписание водителей по городам».
CITY_HEADERS = {'город', 'city', 'населенный_пункт', 'населённый_пункт'}
CITY_MISSING_MESSAGE = 'В файле нет колонки «Город».'
IIN_COLUMN_MISSING_MESSAGE = 'В файле не нашлось колонки с ИИН.'
MAX_LEAD_ROWS = 50000


def _iin_cell(value):
    """Ячейка ИИН → строка. Excel хранит ИИН числом и теряет ведущий ноль у
    родившихся в 2000-х (050312… → 50312…): число дополняем до 12 цифр. Строку
    отдаём как есть — её проверит правило ИИН у вызывающего."""
    if isinstance(value, bool):
        return ''
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return str(value)
        number = int(value)
        return f"{number:012d}" if 0 < number < 10 ** 12 else str(number)
    return str(value or '').strip()


def _norm_header(value):
    return str(value or '').strip().lower().replace(' ', '_')


def _rows_from_csv(raw_bytes):
    text = raw_bytes.decode('utf-8-sig', errors='replace')
    sample = text[:4096]
    delimiter = ';' if sample.count(';') > sample.count(',') else ','
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _rows_from_xlsx(raw_bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_rows(raw_bytes, file_ext):
    ext = str(file_ext or '').lower()
    if ext in ('.xlsx', '.xlsm'):
        raw_rows = _rows_from_xlsx(raw_bytes)
    elif ext == '.csv':
        raw_rows = _rows_from_csv(raw_bytes)
    else:
        raise ValueError('Поддерживаются только .csv, .xlsx и .xlsm')
    raw_rows = [r for r in raw_rows if any(str(c or '').strip() for c in r)]
    if not raw_rows:
        raise ValueError('Файл пустой')
    return raw_rows


def parse_leads_file(raw_bytes, file_ext, max_rows=None, with_iin=False, with_city=False):
    """Разбирает файл в строки (row_number, ФИО, сырой телефон, phone_norm).

    max_rows — потолок строк с данными (по умолчанию MAX_LEAD_ROWS); файл крупнее
    отклоняется целиком, а не обрезается молча.

    with_iin — колонка ИИН обязательна («Обзвон»): строки становятся
    (row_number, ФИО, сырой телефон, phone_norm, ИИН как в файле). Нет колонки —
    ValueError: без неё база бесполезна, и молча загрузить её нельзя. Без шапки
    ИИН — третья колонка.

    with_city (только вместе с with_iin) — к строке добавляется город из колонки с шапкой
    из CITY_HEADERS; колонки нет — пустая строка: город необязателен."""
    limit = int(max_rows) if max_rows is not None else MAX_LEAD_ROWS
    raw_rows = _read_rows(raw_bytes, file_ext)

    fio_idx, phone_idx, iin_idx, city_idx, start_at = 0, 1, None, None, 0
    header = [_norm_header(c) for c in raw_rows[0]]
    known = FIO_HEADERS | PHONE_HEADERS | (IIN_HEADERS if with_iin else set())
    if any(h in known for h in header):
        for idx, name in enumerate(header):
            if name in FIO_HEADERS:
                fio_idx = idx
            elif name in PHONE_HEADERS:
                phone_idx = idx
            elif with_iin and name in IIN_HEADERS:
                iin_idx = idx
            elif with_iin and with_city and name in CITY_HEADERS and city_idx is None:
                city_idx = idx
        start_at = 1
    elif with_iin and max(len(r) for r in raw_rows) >= 3:
        iin_idx = 2
    if with_iin and iin_idx is None:
        raise ValueError(IIN_MISSING_MESSAGE)

    rows = []
    for offset, raw in enumerate(raw_rows[start_at:], start=start_at + 1):
        fio = str(raw[fio_idx] or '').strip() if fio_idx < len(raw) else ''
        phone_cell = raw[phone_idx] if phone_idx < len(raw) else ''
        if isinstance(phone_cell, float):
            phone_cell = f"{phone_cell:.0f}"
        phone_raw = str(phone_cell or '').strip()
        iin_raw = _iin_cell(raw[iin_idx]) if with_iin and iin_idx < len(raw) else ''
        if not fio and not phone_raw and not iin_raw:
            continue
        if len(rows) >= limit:
            raise ValueError(
                f'В файле больше {limit} строк с данными. '
                'Разделите его на несколько файлов.'
            )
        if with_iin and with_city:
            city = str(raw[city_idx] or '').strip() if city_idx is not None and city_idx < len(raw) else ''
            rows.append((offset, fio, phone_raw, normalize_kz_phone(phone_raw), iin_raw, city))
        elif with_iin:
            rows.append((offset, fio, phone_raw, normalize_kz_phone(phone_raw), iin_raw))
        else:
            rows.append((offset, fio, phone_raw, normalize_kz_phone(phone_raw)))
    if not rows:
        raise ValueError('В файле не нашлось ни одной строки с данными')
    return rows


def _looks_like_iin(value):
    text = _iin_cell(value)
    return text.isdigit() and len(text) in (11, 12)


def parse_city_list(raw_bytes, file_ext, max_rows=None):
    """Список «ИИН — город» (например, СМЗ-ИП за месяц) → [(row_number, ИИН как в файле, город, ФИО)].

    Город — колонка с шапкой из CITY_HEADERS (обязательна). ИИН — колонка с шапкой из
    IIN_HEADERS, а нет такой — та, где ИИН больше всего (сводная таблица Excel подписывает её
    «Названия строк»). ФИО — необязательная колонка с шапкой из FIO_HEADERS ('' — нет). Остальные
    колонки (статусы по датам и прочее) не читаются."""
    limit = int(max_rows) if max_rows is not None else MAX_LEAD_ROWS
    raw_rows = _read_rows(raw_bytes, file_ext)
    header = [_norm_header(c) for c in raw_rows[0]]
    city_idx = next((idx for idx, name in enumerate(header) if name in CITY_HEADERS), None)
    if city_idx is None:
        raise ValueError(CITY_MISSING_MESSAGE)
    fio_idx = next((idx for idx, name in enumerate(header) if name in FIO_HEADERS), None)
    iin_idx = next((idx for idx, name in enumerate(header) if name in IIN_HEADERS), None)
    data = raw_rows[1:]
    if iin_idx is None:
        sample = data[:200]
        width = max((len(r) for r in raw_rows), default=0)
        scores = [(sum(1 for r in sample if idx < len(r) and _looks_like_iin(r[idx])), idx)
                  for idx in range(width) if idx not in (city_idx, fio_idx)]
        best = max(scores, default=(0, None))
        if not sample or best[0] * 2 < len(sample):
            raise ValueError(IIN_COLUMN_MISSING_MESSAGE)
        iin_idx = best[1]
    rows = []
    for offset, raw in enumerate(data, start=2):
        iin_raw = _iin_cell(raw[iin_idx]) if iin_idx < len(raw) else ''
        city = str(raw[city_idx] or '').strip() if city_idx < len(raw) else ''
        fio = str(raw[fio_idx] or '').strip() if fio_idx is not None and fio_idx < len(raw) else ''
        if not iin_raw and not city:
            continue
        if len(rows) >= limit:
            raise ValueError(f'В файле больше {limit} строк с данными. Разделите его на несколько файлов.')
        rows.append((offset, iin_raw, city, fio))
    if not rows:
        raise ValueError('В файле не нашлось ни одной строки с данными')
    return rows
