# -*- coding: utf-8 -*-
"""Выгрузка показателей табло Тез КЦ за период в .xlsx.

Отдельный модуль по той же причине, что и `tez_wallboard_source`: bot_schedule2 из тестов
не импортируется (он на старте поднимает пул к боевой базе), а раскладку книги — какие
колонки, где итоги, что считается прочерком — сторожить тестами надо. Здесь чистая
функция: payload на входе, байты файла на выходе, ни Flask, ни database.

Что в книге и почему именно так:
  * лист «Поимённо» — строка на сотрудника и день. Это главный лист: на табло видно
    «сейчас», а вопрос «кто как отработал неделю» закрывается только здесь;
  * лист «По дням» — те же величины, свёрнутые в день. Он нужен, чтобы период читался
    глазами без сводных таблиц;
  * колонки РАЗНЫЕ у ТП и ОП, как и на самом табло: у отдела продаж входящих нет вовсе
    (проверено на живом дне — incomingSuccess = 0 у всех), и «Принято» с «Пропущено»
    стояли бы там пустыми столбцами на весь период.

Длительности — числами в СЕКУНДАХ, а не строками «2:11». Строка «2:11» в Excel и
сортируется как текст, и зажигает зелёный уголок «Число сохранено как текст», который
владелец читает как брак файла (см. память excel-export-text-number-warning). Секунды
суммируются, сортируются и не требуют лечения тега <ignoredErrors>.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Колонки листа «Поимённо» по направлениям: (заголовок, ключ в stats, ширина).
# Ключ None — значение подставляет сборщик строки (дата, имя).
_COMMON_COLUMNS = (
    ('Дата', None, 12),
    ('Сотрудник', None, 28),
)

_TP_COLUMNS = (
    ('Принято', 'served', 11),
    ('Пропущено', 'missed', 13),
    ('Разговор всего, сек', 'talk_seconds', 20),
    ('Ср. разговор, сек', 'avg_talk_seconds', 18),
    ('Ср. ожидание, сек', 'avg_wait_seconds', 18),
    ('Набрано', 'outgoing_total', 11),
    ('Дозвонились', 'outgoing_success', 14),
)

_OP_COLUMNS = (
    ('Набрано', 'outgoing_total', 11),
    ('Дозвонились', 'outgoing_success', 14),
    ('Дозвон, %', 'outgoing_success_percent', 12),
    ('Исходящий разговор, сек', 'outgoing_talk_seconds', 24),
    ('Ср. исходящий, сек', 'avg_outgoing_talk_seconds', 20),
)


def columns_for(direction):
    """Колонки направления: у ТП — про входящую линию, у ОП — про исходящие."""
    return _COMMON_COLUMNS + (_OP_COLUMNS if direction == 'op' else _TP_COLUMNS)


def _int(value):
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def person_row_values(direction, stats):
    """Значения показателей одного человека за день в порядке колонок направления.

    None остаётся None и попадёт в ячейку прочерком: «среднего нет» и «среднее ноль» —
    разные вещи, и ноль в колонке средней длительности читался бы как мгновенный разговор.
    """
    stats = stats or {}
    out = []
    for _, key, _width in columns_for(direction)[len(_COMMON_COLUMNS):]:
        value = stats.get(key)
        if key in ('avg_talk_seconds', 'avg_wait_seconds', 'avg_outgoing_talk_seconds',
                   'outgoing_success_percent'):
            out.append(None if value is None else round(float(value), 1))
        else:
            out.append(_int(value))
    return out


def day_totals(direction, people):
    """Свёртка дня: счётчики складываются, средние считаются ЗАНОВО от сумм.

    Среднее средних здесь было бы просто неверным: у одного оператора три разговора, у
    другого сорок, и их средние нельзя складывать пополам. Поэтому сумма секунд делится
    на сумму звонков — ровно как считает плитка на табло."""
    served = sum(_int((p.get('stats') or {}).get('served')) for p in people)
    missed = sum(_int((p.get('stats') or {}).get('missed')) for p in people)
    talk = sum(_int((p.get('stats') or {}).get('talk_seconds')) for p in people)
    wait = sum(_int((p.get('stats') or {}).get('wait_seconds')) for p in people)
    out_total = sum(_int((p.get('stats') or {}).get('outgoing_total')) for p in people)
    out_success = sum(_int((p.get('stats') or {}).get('outgoing_success')) for p in people)
    out_talk = sum(_int((p.get('stats') or {}).get('outgoing_talk_seconds')) for p in people)
    return {
        'served': served,
        'missed': missed,
        'talk_seconds': talk,
        # Усекаем, а не округляем: то же правило, что на табло — кабинет показывает 02:09
        # при 129,74 с, и округление разошлось бы с ним на каждой смене.
        'avg_talk_seconds': (talk // served) if served else None,
        'avg_wait_seconds': (wait // served) if served else None,
        'outgoing_total': out_total,
        'outgoing_success': out_success,
        'outgoing_talk_seconds': out_talk,
        'avg_outgoing_talk_seconds': (out_talk // out_success) if out_success else None,
        'outgoing_success_percent': (round(out_success * 100.0 / out_total, 1)
                                     if out_total else None),
        'people': len([p for p in people if _int((p.get('stats') or {}).get('served'))
                       or _int((p.get('stats') or {}).get('outgoing_total'))]),
    }


def build_workbook(payload):
    """Книга .xlsx (bytes) по собранному периоду."""
    payload = payload or {}
    direction = 'op' if payload.get('direction') == 'op' else 'tp'
    days = [day for day in (payload.get('days') or []) if isinstance(day, dict)]
    columns = columns_for(direction)

    thin = Side(style='thin', color='E2E8F0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill('solid', fgColor='1F2937')
    total_fill = PatternFill('solid', fgColor='F1F5F9')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    title_font = Font(bold=True, size=15, color='0F172A')
    hint_font = Font(color='94A3B8', size=10)
    note_font = Font(color='B45309')
    total_font = Font(bold=True, color='0F172A')
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(vertical='center')

    workbook = Workbook()

    def _head(sheet, title, subtitle, headers, widths):
        sheet['A1'] = title
        sheet['A1'].font = title_font
        sheet['A2'] = subtitle
        sheet['A2'].font = hint_font
        for index, (header, width) in enumerate(zip(headers, widths), start=1):
            cell = sheet.cell(row=4, column=index, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A5'

    label = payload.get('direction_label') or ('ОП' if direction == 'op' else 'ТП')
    period = payload.get('date_from') or ''
    if payload.get('date_to') and payload.get('date_to') != payload.get('date_from'):
        period = '%s — %s' % (payload.get('date_from'), payload.get('date_to'))
    subtitle = 'Источник: кабинет Binotel · собрано %s' % (payload.get('generated_at') or '')

    # --- лист «Поимённо» ---------------------------------------------------------------
    people_sheet = workbook.active
    people_sheet.title = 'Поимённо'
    _head(people_sheet, 'Табло Тез КЦ · %s · %s' % (label, period), subtitle,
          [column[0] for column in columns], [column[2] for column in columns])

    row_index = 5
    for day in days:
        people = [person for person in (day.get('people') or []) if isinstance(person, dict)]
        if day.get('error'):
            # День, который кабинет не отдал, НЕ превращаем в нули: пустой день и день без
            # данных на отчёте выглядели бы одинаково, а решения по ним разные.
            people_sheet.cell(row=row_index, column=1, value=day.get('day')).border = border
            cell = people_sheet.cell(row=row_index, column=2,
                                     value='Кабинет не отдал данные: %s' % day['error'])
            cell.font = note_font
            cell.border = border
            row_index += 1
            continue
        for person in sorted(people, key=lambda item: str(item.get('name') or '')):
            values = [day.get('day'), person.get('name') or '—'] + person_row_values(
                direction, person.get('stats'))
            for column_index, value in enumerate(values, start=1):
                cell = people_sheet.cell(row=row_index, column=column_index, value=value)
                cell.border = border
                cell.alignment = left if column_index <= len(_COMMON_COLUMNS) else center
            row_index += 1

    # --- лист «По дням» ----------------------------------------------------------------
    day_headers = ['Дата', 'Людей со звонками'] + [column[0] for column in columns[len(_COMMON_COLUMNS):]]
    day_widths = [12, 20] + [column[2] for column in columns[len(_COMMON_COLUMNS):]]
    day_sheet = workbook.create_sheet('По дням')
    _head(day_sheet, 'Табло Тез КЦ · %s · итоги по дням' % label, subtitle, day_headers, day_widths)

    row_index = 5
    period_people = []
    for day in days:
        people = [person for person in (day.get('people') or []) if isinstance(person, dict)]
        if day.get('error'):
            day_sheet.cell(row=row_index, column=1, value=day.get('day')).border = border
            cell = day_sheet.cell(row=row_index, column=2,
                                  value='Кабинет не отдал данные: %s' % day['error'])
            cell.font = note_font
            cell.border = border
            row_index += 1
            continue
        period_people.extend(people)
        totals = day_totals(direction, people)
        values = [day.get('day'), totals['people']] + person_row_values(direction, totals)
        for column_index, value in enumerate(values, start=1):
            cell = day_sheet.cell(row=row_index, column=column_index, value=value)
            cell.border = border
            cell.alignment = left if column_index == 1 else center
        row_index += 1

    if period_people:
        # Итог периода считается от ВСЕХ строк периода, а не сложением дневных средних —
        # та же причина, по которой день считается от людей, а не от их средних.
        totals = day_totals(direction, period_people)
        values = ['Итого', totals['people']] + person_row_values(direction, totals)
        for column_index, value in enumerate(values, start=1):
            cell = day_sheet.cell(row=row_index, column=column_index, value=value)
            cell.border = border
            cell.fill = total_fill
            cell.font = total_font
            cell.alignment = left if column_index == 1 else center

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_file_name(payload):
    """Имя файла: направление и период — чтобы две выгрузки не путались в «Загрузках»."""
    payload = payload or {}
    label = 'op' if payload.get('direction') == 'op' else 'tp'
    date_from = (payload.get('date_from') or '').replace('-', '')
    date_to = (payload.get('date_to') or '').replace('-', '')
    period = date_from if (not date_to or date_to == date_from) else '%s-%s' % (date_from, date_to)
    return 'tez_wallboard_%s_%s.xlsx' % (label, period or 'period')
